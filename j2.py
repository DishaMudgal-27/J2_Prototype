"""J2: the five functions run one department until Identity's goal is met (Native docs, system/department).
The flow follows the J2 diagram (docs/j2-flow.jpg):

  J2.a  Identity reads the Requester's brief: enough information to work? If not, it asks the Requester specific
        questions and the run waits; it asks as many rounds as it needs, then records the charter. It invents
        nothing. A fact the Requester can't supply yet is recorded as "to be confirmed" against its done-when item.
  J2.b  Adaptation proposes workflows (or disputes an Audit finding or a Coordination HOLD), using the most
        relevant lessons from earlier runs (memory/lessons.jsonl)
  J2.c  Priority prices the proposal: GRANT, COUNTER or REFUSE. The bargain is bounded by four hops, fifteen
        minutes and its spend (AI calls); past any of them it goes to Identity
  J2.d  Any two agents that disagree get 4 attempts to settle it themselves; then Identity resolves it from the
        charter, or, if the charter doesn't decide it, asks the Requester one decision-ready question and rules
  J2.e  Validator (plain code) checks and registers the agreed workflows; each one gets its own workflow agent
        ("W-1 Builder") that owns one marked section of the site file
  J2.f  Coordination schedules; Priority checks each run's envelope; the workflow's agent writes its section,
        retries within its envelope and reports back to Coordination
  J2.g  Audit checks the result against the charter's done-when list and files its findings to Identity. A
        missing fact only the Requester knows becomes Identity's question to the Requester
  J2.h  Identity: is the site workable (see workable())? Yes: Identity tells the Requester the goal is reached.
        No: back to J2.b. When Audit confirms a fix worked, Adaptation saves a lesson

Only Identity talks to the Requester: any other agent that addresses the Requester is redirected to Identity.
Each function's instructions come from agents/j2/team.md + agents/j2/<function>.md, read fresh on every call.
Speaking order is dynamic: every reply names "next", and code allows only the routes in ROUTES (an illegal pick
is redirected with a "Route check" note). Code also enforces the limits: 4 attempts per disagreement, budget,
max cycles, max turns per cycle and the same failure twice in a row. A limit never stops the run silently:
Identity asks the Requester to add budget, drop the failing items or ship as is. A run that waits on the
Requester stops and resumes on the reply, so it can't loop on its own.

J2_FIXES (default 0) picks the mode, so findings can compare the two:
  0  baseline, as close to the diagram as possible: an Identity ruling goes straight to the validator (J2.d ->
     J2.e), and a change request after done goes straight to Adaptation (J2.b).
  1  fixes: (a) an Identity ruling that costs more than Priority granted goes back to Priority (J2.c) once for
     pricing; Priority's answer is final, and if the ruling is still over budget Identity asks the Requester.
     (b) a change request after done goes to Identity at J2.a first, then J2.b.
"""
import copy
from functools import reduce
import html
import html.parser
import itertools
import json
import math
import os
import re
import threading
import time

import j2_artifacts
import j2_board
import j2_check
import j2_map
import j2_realization
import j2_schema
import j2_send
import llm

HERE = os.path.dirname(os.path.abspath(__file__))
J2_DIR = os.path.join(HERE, "agents", "j2")
RUNS_DIR = os.getenv("J2_RUNS_DIR") or os.path.join(HERE, "runs")
BUDGET = int(os.getenv("J2_BUDGET", "100"))           # starting budget
BUDGET_CAP = int(os.getenv("J2_BUDGET_CAP", str(2 * BUDGET)))   # most the Requester's approvals can raise it to
MAX_ATTEMPTS = 4           # attempts two disagreeing agents get before Identity rules (priority.md: four exchanges)
BARGAIN_MINUTES = float(os.getenv("J2_BARGAIN_MINUTES", "15"))   # J2.c bargain: four hops, fifteen minutes...
BARGAIN_CALLS = int(os.getenv("J2_BARGAIN_CALLS", "8"))          # ...and its spend, in AI calls (4 hops x 2 sides)
MAX_CYCLES = int(os.getenv("J2_MAX_CYCLES", "6"))   # plan -> build -> audit loops before Identity asks the Requester
LIMIT_ADD = int(os.getenv("J2_LIMIT_ADD", "30"))    # budget the "keep working" option adds when the budget ran out
MEMORY_FILE = os.getenv("J2_MEMORY_FILE") or os.path.join(HERE, "memory", "lessons.jsonl")
MAX_LESSONS = 5            # lessons Adaptation reads at J2.b
PLAYBOOK_FILE = os.getenv("J2_PLAYBOOK_FILE") or os.path.join(HERE, "memory", "playbooks.jsonl")
MAX_PLAYBOOKS = 2          # playbooks Adaptation reads at its first J2.b of a brief
# Audit's design and usability rubric, checked on the screenshot (360px and 1280px) and code's design measurements.
# A failing item is a normal finding: the site isn't workable until every item passes.
RUBRIC = (("R-1", "Consistent styling: the whole page follows the style guide (same fonts, colours, button and card "
                  "styles)"),
          ("R-2", "Readable contrast: text is easy to read against its background (WCAG AA)"),
          ("R-3", "Clear hierarchy: one main heading, headings in order, the most important things stand out"),
          ("R-4", "One obvious main action: a visitor can tell at once what to do next"),
          ("R-5", "Easy navigation: a visitor can reach every section quickly, on a phone too"),
          ("R-6", f"Tap targets are big enough on phones (at least {j2_check.TAP_MIN}x{j2_check.TAP_MIN}px)"),
          ("R-7", "Images have alt text"))
WORKABLE_OUTCOME = ("Done: the site is workable. Every done-when item passes with Audit's evidence and the render "
                    "check is clean.")
MAX_TURNS = int(os.getenv("J2_MAX_TURNS", "30"))       # agent turns per cycle
CHAT_FULL_CHARS = int(os.getenv("J2_CHAT_FULL_CHARS", "60000"))   # above this, summary + recent messages
CHAT_TAIL = 40             # messages always shown in full once the chat is summarised
SUMMARY_REFRESH = 20       # re-summarise once this many more messages have scrolled out of the tail
SITE_LIMIT = 150_000       # characters of the site the workflow agents and Audit see
SITE_WARN_BYTES = 60_000   # builder.md asks for a single page under ~60 KB
DELAY = 0 if llm.MOCK else float(os.getenv("J2_DELAY", "2"))   # spacing for free-tier rate limits

NAMES = {"validator": "Validator", "identity": "Identity", "adaptation": "Adaptation",
         "priority": "Priority", "coordination": "Coordination", "audit": "Audit", "user": "Requester (you)",
         "system": "System"}
FUNCTIONS = ("identity", "adaptation", "priority", "coordination", "audit")   # the five, as opposed to the builders
# Who may hand the turn to whom (J2 routes). Code narrows these further by situation. Only Identity reaches the
# Requester ("user"); an agent that addresses the Requester anyway is redirected to Identity.
ROUTES = {"identity": ("user", "adaptation", "priority", "coordination", "audit"),
          "adaptation": ("priority", "identity", "audit", "coordination"),    # Audit/Coordination: to dispute
          "priority": ("adaptation", "identity", "coordination"),   # Coordination: envelope check (HOLD)
          "coordination": ("priority", "adaptation", "identity"),
          "audit": ("adaptation", "identity")}
# Workflow agents ("W-1" -> "W-1 Builder") are added to NAMES and ROUTES when their workflow is registered.
# ponytail: module-level registry; ids restart at W-1 each run, so names stay correct for one run at a time.


def add_workflow_agent(wid):
    NAMES[wid] = f"{wid} Builder"
    ROUTES[wid] = ("coordination",)


def is_wf_agent(fn):
    return bool(re.fullmatch(r"W-\d+", str(fn or "")))
# J2 steps, labelled as in the diagram, for the progress strip in the UI
STEPS = {1: "J2.a Enough information to work?", 2: "J2.b Adaptation proposes workflows",
         3: "J2.c Priority prices the proposal", 4: "J2.d Identity resolves conflict",
         5: "J2.e Validator registers workflow", 6: "J2.f Coordination schedules, workflow agents build",
         7: "J2.g Audit checks result", 8: "J2.h Is the site workable?"}
TOPICS = {"proposal": "what to fund", "finding": "an Audit finding", "hold": "a Coordination HOLD"}
PAUSES = {"info": "J2.a: not enough information to work yet",
          "charter_drop": "the charter leaves out items from the brief",
          "conflict": "J2.d: a conflict the charter doesn't decide",
          "question": "a question the charter can't answer",
          "approval": "approval for a run over its envelope",
          "ruling_budget": "J2_FIXES=1: Identity's ruling is over budget after Priority's final answer",
          "audit": "J2.h: a done-when item only the Requester can close",
          "fact": "J2.g: a fact only the Requester knows",
          "limit": "a safety limit was hit: add budget, drop the failing items or ship as is"}
TO_USER = re.compile(r"@(?:requester|user)\b", re.I)


def system_prompt(fn, reply_format, builder="builder.md"):
    """team.md + the function's role file. builder: a workflow agent's role, a file in agents/j2 or the text itself
    (an output type's own builder prompt)."""
    parts = []
    for name in ("team.md", builder if is_wf_agent(fn) else f"{fn}.md"):
        path = os.path.join(J2_DIR, name)
        if name.endswith(".md") and os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                parts.append(j2_realization.PLANNING_PART.sub("", f.read()).strip())   # the planning part is the planner's only
        elif not name.endswith(".md"):
            parts.append(name.strip())
    parts.append(f"You are {NAMES[fn]}.\n\n## Reply format (required)\n{reply_format}")
    return "\n\n".join(parts)


def wf_cost(w):
    """Official proposals put cost in estimated_resources.money; the mock uses "cost"."""
    est = w.get("estimated_resources") or {}
    try:
        return int(float(w.get("cost") or est.get("money") or 0))
    except (TypeError, ValueError):
        return 0


def cost_of(wfs):
    return sum(wf_cost(w) for w in wfs or [])


ENVELOPE_DEFAULTS = {"model_call_limit": 2, "retry_limit": 1}   # used only when Priority leaves them out


def as_int(x, default=0):
    try:
        return int(float(x))
    except (TypeError, ValueError):
        return default


RATES_ON = os.getenv("J2_RATES", "1") != "0"
RATES = {}      # agents/j2/rates.json, reloaded at each reset so edits apply without a restart


def load_rates(output="website", card=True):
    """The rate card for the run's output type (rates.json: "rates" for a website, "rates_prd" for a PRD). card=False:
    no rate card, so each workflow costs Adaptation's own estimate (estimated_resources.money), as with J2_RATES=0."""
    global RATES, OUTPUT
    RATES, OUTPUT = {}, output
    path = os.path.join(J2_DIR, "rates.json")
    if card and RATES_ON and os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            RATES = {k: int(v) for k, v in (json.load(f).get(OUTPUTS[output]["rates"]) or {}).items()}


def _rate_card(key):
    path = os.path.join(J2_DIR, "rates.json")
    if RATES_ON and os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return {k: int(v) for k, v in (json.load(f).get(key) or {}).items()}
    return {}


def RATES_WEB():
    """The website rate card, whatever output is current (the cost of a stage still to come)."""
    return _rate_card("rates")


def RATES_DOC():
    return _rate_card("rates_prd")


def RATES_DEED():
    """Price of each send deed (rates.json "rates_deed"); empty, so free, with J2_RATES=0."""
    return _rate_card("rates_deed")


def price(w):
    """(rate-card cost, unknown parts) for a workflow's "parts"; (None, []) when the rate card is off or parts are missing."""
    parts = w.get("parts")
    if not RATES or not isinstance(parts, list) or not parts:
        return None, []
    unknown = [str(p) for p in parts if str(p) not in RATES]
    return sum(RATES.get(str(p), 0) for p in parts), unknown


# Cost follows the work: each done-when item a workflow covers has a minimum price, the rate-card part that kind of
# item needs. ponytail: keyword match on the item text; first match wins, else a page section.
ITEM_KINDS = (("form", r"\bform\b|\border|\bbook(?:ing)?\b|enquir|inquir|appointment|reserv|sign[- ]?up|subscribe"),
              ("interactive_feature", r"filter|search|sort|slider|carousel|calculator|toggle|\btabs?\b|interactive|animat"),
              ("styling", r"viewport|mobile|phone|responsive|horizontal scroll|\b360\s*px|\b1280\s*px|\bcolou?rs?\b|\bfonts?\b"),
              ("page_section", r""))


DOC_KINDS = (("requirements_section", r"requirement|components?\b|pages and sections"), ("doc_section", r""))   # a document's items


def item_kind(criterion):
    return next(k for k, pat in (DOC_KINDS if OUTPUTS[OUTPUT]["sections"] else ITEM_KINDS)
                if re.search(pat, str(criterion), re.I))


def item_min(criterion):
    return RATES.get(item_kind(criterion), 0)


# Leaving items out for budget: the no-errors gate and the brief's core facts always stay in. Identity marks core
# items at J2.a (goal["core_items"]); the keywords are only the fallback for a charter without that mark.
NO_ERRORS_GATE = re.compile(r"\b(?:javascript|js) errors\b", re.I)    # the gate item names JavaScript errors
GATE_TEXT = ("The page has no JavaScript errors, no horizontal overflow at 360px and 1280px, and every in-page link "
             "points to an existing section.")     # fixed by code at J2.a: Identity never adds features to it
CORE_FACT = re.compile(r"\bprices?\b|\bhours\b|timings?|address|location|\bmaps?\b|contact|whatsapp|e-?mail|"
                       r"phone number|\bcall\b|\border", re.I)
BUDGET_TALK = re.compile(r"budget|\bcosts?\b|afford|expensive|cheap|\bcut|\bdrop|leave (?:\w+ )?out|fund|credits?\b|"
                         r"\bunits?\b|rate[- ]card|minimums?", re.I)


def must_keep(criterion, goal):
    """True for a done-when item that can't be left out for budget: the no-errors gate, or a core fact from the brief
    (business name and location, contact details, prices, hours, ordering). Identity's mark from J2.a decides; a
    charter without it falls back to keywords."""
    g, c = goal or {}, str(criterion)
    if NO_ERRORS_GATE.search(c):
        return True
    if "core_items" in g:
        return c in g["core_items"]
    name = str(g.get("site_name") or "").lower()
    return bool(CORE_FACT.search(c) or (name and name in c.lower()))


def granted(w):
    """Rate-card price when the rate card is on; otherwise the money Priority granted (envelope.money_limit),
    and Adaptation's estimate only if no limit was given."""
    c, _ = price(w)
    if c is not None:
        return c
    m = as_int((w.get("envelope") or {}).get("money_limit"))
    return m if m > 0 else wf_cost(w)


def granted_total(wfs):
    return sum(granted(w) for w in wfs or [])


def norm(wfs):
    """Keep only workflow objects and give each a plain "cost" (the UI and validator read it):
    the rate-card price of its parts when the rate card is on, else the agent's estimate."""
    out = []
    for w in wfs or []:
        if isinstance(w, dict):
            c, _ = price(w)
            est = w.get("estimate", wf_cost(w))   # the agent's own estimate, kept for the run report
            out.append({**w, "estimate": est, "cost": c if c is not None else wf_cost(w)})
    return out


EXAMPLE_WORKFLOW = {   # shown to Adaptation in the code-side prompt wrapper (J2Run.ATOM_RULES); a test registers it as is
    "name": "Contact section", "goal_gap": "Visitors have no way to message the business", "outputs": ["#contact"],
    "completion_check": "A contact form with a send button is on the page", "parts": ["form"], "covers": [1],
    "atoms": [{"id": "a1", "goal": "Visitors can send the business a message from a contact form.",
               "constraints": ["Other sections stay unchanged", "Create only #contact and #contact-send"],
               "context": ["charter item 1"],
               "how": {"kind": "AI", "agent": "Builder", "steps": "Create section #contact holding a name field, an "
                       "email field and a send button #contact-send"},
               "expected_output": [{"text": "The contact section", "kind": "section", "target": "#contact"},
                                   {"text": "The send button", "kind": "element", "target": "#contact-send"}],
               "verification": [{"text": "One #contact section is on the page", "type": "element-count",
                                 "target": "#contact", "expect": 1},
                                {"text": "The form has a submit button", "type": "named-test",
                                 "target": "form-has-submit", "expect": True}]}]}

ATOM_FIELDS = ("goal", "constraints", "context", "how", "expected_output", "verification")


def atom_missing(w):
    """J2.e's work-atom check: every field each atom lacks, by name ("atoms" when the workflow has none). An empty
    expected_output or verification list counts as missing."""
    atoms = w.get("atoms")
    if not isinstance(atoms, list) or not atoms:
        return ["atoms"]
    return [f"{f} in work atom {n}" for n, a in enumerate(atoms, 1) for f in ATOM_FIELDS
            if not (isinstance(a, dict) and a.get(f))]


# ---- work-atom checks: the only check types Adaptation may write. Code runs them; no agent reports a result.
CHECK_TYPES = ("file-exists", "grep-count", "element-count", "named-test", "facts-present", "tap-targets",
               "behavior-steps")
DOC_CHECK_TYPES = ("facts-present", "heading-present", "section-not-empty", "ids-unique")   # a document: no browser
COUNTED = ("grep-count", "element-count", "tap-targets")     # target: a section id or class the atom says to create
NAMED_TESTS = ("page-parses", "links-resolve", "form-has-submit")
MAX_CHECK_ATTEMPTS = 3
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
SPECIFIC = re.compile(r"\d|" + EMAIL.pattern)       # a time, price, number or email in a goal
IDENT = re.compile(r"^[#.]?[A-Za-z][\w-]*$")         # a section id or class name
# J2_REALIZATION: an id range in how or constraints ("#pad-1 to #pad-9", "cell-0-0..cell-8-8", "#cell-r-c (r and c
# 0 to 8)") tells the builder to create every id in it; a name with a digit in it (#pad-4, j2-btn) is not a count.
ID_RANGE = re.compile(r"#?\b([a-z][\w-]*?)(\d+(?:-\d+)*)\s*(?:to|through|\.\.|–|—)\s*#?\1(\d+(?:-\d+)*)\b")
RC_RANGE = re.compile(r"#?\b([a-z][\w-]*?-)r-c\b[^.;]{0,40}?\b(\d+)\s*(?:to|through|\.\.|–|—)\s*(\d+)\b")
NAME_WITH_DIGIT = re.compile(r"[#.]?\b[a-z][\w-]*\d[\w-]*", re.I)
MAX_RANGE_IDS = 400
# A goal that describes an action needs a behaviour item proving it works; a rule that depends on a date needs a
# control on the page that sets the date, used by the steps (never today's date).
ACTION_GOAL = re.compile(r"\b(?:(?:re)?calculat\w*|self-updating|on change|shows?\s*(?:/|and|or)\s*hides?|"
                         r"toggl\w*|highlight\w*|validat\w*|reveal\w*|hides?|shown only when|hidden unless|"
                         r"required only (?:when|if))\b", re.I)
DATE_RULE = re.compile(r"early[- ]?bird|deadline|\b(?:up to|until|till|before|by|after)\s+(?:the\s+)?\d{1,2}(?:st|nd|rd|th)?"
                       r"\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)", re.I)
DATE_CONTROL = re.compile(r"date|day|today|early", re.I)


def human_problem(t):
    """Why a builder's human_output is unusable (None when it is fine): 1 to 3 sentences, at least 5 words."""
    t = str(t or "").strip()
    n = len([x for x in re.split(r"(?<=[.!?])\s+", t) if x])
    return "it is empty or too short (\"done.\" alone is rejected)" if len(t.split()) < 5 else \
        "it is longer than 3 sentences" if n > 3 else None


INTERNAL = re.compile(r"\bW-\d+(?:\.a\d+)?\b|\ba\d\b|(?<![\w&])#[A-Za-z][\w-]*|(?<=\s)\.[A-Za-z][\w-]+|\bunits?\b|[<>]|"
                      r"\b(?:id|class)=|J2:", re.I)


def atom_human_problem(t):
    """Why an atom's human_output is unusable (None when it is fine): the message-level wording rules (human_problem:
    plain, at least 5 words) but 1 or 2 sentences, and no internal id, selector or "units" (the Requester reads it)."""
    p = human_problem(t)
    if p:
        return p
    t = str(t).strip()
    if len([x for x in re.split(r"(?<=[.!?])\s+", t) if x]) > 2:
        return "it is longer than 2 sentences"
    m = INTERNAL.search(t)
    return f"it has an internal id, selector or \"units\" ({m.group(0)!r})" if m else None


def atom_done_problem(a):
    """A work atom is done only with both outputs: a human_output that passes atom_human_problem and an agent_output."""
    h, ag = a.get("human_output"), a.get("agent_output")
    return [x for x, bad in (("human_output", atom_human_problem((h or {}).get("text")) if isinstance(h, dict) else True),
                             ("agent_output", not (isinstance(ag, dict) and ag.get("receipt")))) if bad]


def created_in(html):
    """The element ids and class names a section's HTML creates, in order, each once."""
    ids = list(dict.fromkeys(re.findall(r"""\bid=["']([^"'\s]+)["']""", html or "")))
    cls = list(dict.fromkeys(c for v in re.findall(r"""\bclass=["']([^"']*)["']""", html or "") for c in v.split()))
    return ids, cls


def element_by_id(page, eid):
    """The outer HTML of the element with this id in a page (balanced tags), or None."""
    m = re.search(rf"""<([a-zA-Z][\w-]*)\b[^>]*\bid=["']{re.escape(eid)}["'][^>]*>""", page or "")
    if not m:
        return None
    tag, depth = m.group(1).lower(), 0
    for t in re.finditer(rf"<(/?){tag}\b[^>]*?(/?)>", page[m.start():], re.I):
        depth += -1 if t.group(1) else 0 if t.group(2) else 1
        if depth == 0:
            return page[m.start():m.start() + t.end()]
    return None


def preview_doc(page, eid):
    """A small page for a sandboxed frame: the built page's head (the style guide) and one element by its id."""
    el = element_by_id(page, eid)
    head = re.search(r"<head\b[^>]*>(.*?)</head>", page or "", re.S | re.I)
    return None if el is None else ("<!doctype html><html><head><meta charset=\"utf-8\">" + (head.group(1) if head else "")
                                    + "</head><body>" + el + "</body></html>")


def fallback_human(a, noun):
    """Code's own human_output when the builder's is unusable twice: the atom's expected output in plain words."""
    parts = [re.sub(INTERNAL, "", x.get("text") or "").strip(" .") for x in a.get("expected_output") or []
             if isinstance(x, dict)]
    text = f"It created {and_list([p for p in parts if p][:3])}." if any(parts) else ""
    return text if not atom_human_problem(text) else f"It adds the part of the {noun} this step was asked to build."


def build_agent_output(w, a, attempt, body, receipts, last_audit, n_sites, doc, out_file):
    """What other agents read instead of searching: built by code from the section as written and the receipts of
    this attempt (no model). Files, section names, element ids and selectors created, the checks and their results,
    the audit's N/P conditions, the receipt id (receipts.jsonl: atom and attempt)."""
    ids, cls = created_in(body)
    heads = [re.sub(r"<[^>]+>|^#+\s*", "", h).strip() for h in re.findall(r"<h[1-3][^>]*>.*?</h[1-3]>|^#{1,3} .*$", body or "",
                                                                           re.S | re.M)]
    want = [str(x.get("target")) for x in a.get("expected_output") or [] if isinstance(x, dict) and x.get("kind") == "section"]
    anchor = next((t[1:] for t in want if t.startswith("#") and t[1:] in ids), None) or (ids[0] if ids else None)
    recs = [r for r in receipts if r["atom"] == a["id"] and r["attempt"] == attempt and "check" in r]
    return {"atom": a["id"], "attempt": attempt, "receipt": f"{a['id']}:{attempt}",
            "files": [f"site/v{n_sites}.{'md' if doc else 'html'}", out_file or "site/index.html"],
            "section": w["section"], "section_names": [w["name"]] + [h for h in heads if h][:6], "anchor": anchor,
            "ids": ids[:40], "selectors": [f"#{i}" for i in ids[:40]] + [f".{c}" for c in cls[:40]],
            **({"kind": "doc", "file": out_file} if doc else {}),
            "checks": [{"check": r["check"], "type": r["type"], "target": r["target"], "pass": r["pass"]} for r in recs],
            "conditions": [{"id": c["id"], "met": c["met"]} for c in (last_audit or {}).get("conditions") or []]}


def slim():
    """On unless J2_SLIM=0: prompts carry finished atoms' agent_output instead of their attempts and receipts, and builders
    see their own section plus a manifest instead of the whole page. Audit still reads the full page. J2_SLIM=0 is the old
    path, prompt for prompt. Read per call so tests can switch it."""
    return os.environ.get("J2_SLIM") != "0"


def slim_atom(a):
    return {k: a[k] for k in ("id", "status", "agent_output")} if a.get("status") == "passed" and a.get("agent_output") else a


def squash(t):
    """Case-insensitive, whitespace collapsed; no other rewriting."""
    return re.sub(r"\s+", " ", str(t)).strip().lower()


def strings(x):
    """Every string inside a JSON-like value, for the facts-present source text."""
    if isinstance(x, dict):
        return [t for v in x.values() for t in strings(v)]
    if isinstance(x, list):
        return [t for v in x for t in strings(v)]
    return [x] if isinstance(x, str) else []


def target_problem(target, told, source):
    """Why a counted check or a behaviour step can't use this selector (None when it can): only an id or class
    the atom's how or constraints tells the builder to create."""
    name = str(target or "").lstrip("#.")
    if not IDENT.match(str(target or "")) or EMAIL.search(str(target)) or \
            (re.search(r"\d", name) and squash(name) in source):
        return (f"targets {target!r}: only a section id or class name, never a charter fact, an email or a generic "
                "string (use facts-present for those)")
    if name.lower() not in told:
        return f"targets {target!r}, which the atom's how or constraints doesn't tell the builder to create"
    return None


def allowed_numbers(source, steps):
    """Numbers an expect-number expression may use: those in the charter and the Requester's answers (a percentage
    also as its fraction and its remainder: 10% -> 0.1 and 0.9), and the inputs the steps themselves set."""
    nums = set(j2_check.numbers_in(source))
    for p in re.findall(r"(\d+(?:\.\d+)?)\s*(?:%|per ?cent)", source):
        nums |= {float(p) / 100, 1 - float(p) / 100}
    for st in steps:
        if str(st).split(None, 1)[:1] == ["set"]:
            nums |= set(j2_check.numbers_in(str(st).split(None, 2)[2:]))
    return nums


def behavior_problems(steps, told, source, about):
    """J2.e's rules for one behaviour check's steps. about: the atom's goal, how and the item text, squashed."""
    if not isinstance(steps, list) or not steps:
        return ["has no steps"]
    if len(steps) > j2_check.MAX_STEPS:
        return [f"has {len(steps)} steps; at most {j2_check.MAX_STEPS}"]
    out, parsed = [], []
    for i, st in enumerate(steps, 1):
        try:
            parsed.append(j2_check.parse_step(st))
        except ValueError as e:
            out.append(str(e))
    if out:
        return out
    allowed = allowed_numbers(source, steps)
    for i, p in enumerate(parsed, 1):
        if p["sel"] and (bad := target_problem(p["sel"], told, source)):
            out.append(f"step {i} {bad}")
        if p["verb"] == "expect-number":
            try:
                lits = j2_check.safe_number(p["arg"])[1]
            except ValueError:
                out.append(f"step {i} expect-number {p['arg']!r} isn't arithmetic over numbers (+ - * / and brackets "
                           "only)")
                continue
            new = [x for x in lits if not any(abs(x - y) < 1e-9 for y in allowed)]
            if new:
                out.append(f"step {i} expect-number uses {', '.join(f'{x:g}' for x in new)}, not in the charter, the "
                           "Requester's answers or the inputs the steps set (no invented totals)")
        elif p["verb"] in ("expect-changed", "expect-style-changed"):
            need = "wait" if p["verb"] == "expect-changed" else "click"
            if not any(q["verb"] == need for q in parsed[:i - 1]):
                out.append(f"step {i} {p['verb']} needs a {need} step before it")
    if DATE_RULE.search(about) and any(p["verb"] in ("expect-number", "expect-text") for p in parsed) and \
            not any(p["verb"] == "set" and DATE_CONTROL.search(p["sel"]) for p in parsed):
        out.append("its rule depends on a date: the atom must create a control that sets the date (e.g. #reg-date) "
                   "and the steps must set it, never rely on today's date")
    return out


def range_ids(told):
    """Every id an id range in this squashed text names, space-joined ("" when there is none): "pad-1 to pad-9" ->
    pad-1 ... pad-9; "cell-0-0 to cell-8-8" and "cell-r-c (r and c 0 to 8)" -> cell-0-0 ... cell-8-8. A range of
    more than MAX_RANGE_IDS ids is left out."""
    out = []
    for pre, lo, hi in ID_RANGE.findall(told):
        lo, hi = [int(x) for x in lo.split("-")], [int(x) for x in hi.split("-")]
        if len(lo) == len(hi) and all(x <= y for x, y in zip(lo, hi)) and \
                math.prod(y - x + 1 for x, y in zip(lo, hi)) <= MAX_RANGE_IDS:
            out += [pre + "-".join(map(str, p)) for p in itertools.product(*[range(x, y + 1) for x, y in zip(lo, hi)])]
    for pre, lo, hi in RC_RANGE.findall(told):
        lo, hi = int(lo), int(hi)
        if lo <= hi and (hi - lo + 1) ** 2 <= MAX_RANGE_IDS:
            out += [f"{pre}{r}-{c}" for r in range(lo, hi + 1) for c in range(lo, hi + 1)]
    return " ".join(out)


def check_problems(a, source, types=CHECK_TYPES, upstream="", ranges=False):
    """J2.e's rules for one atom's checks. source: the charter and the Requester's answers, squashed. types: the
    check types the output allows (a website's, or DOC_CHECK_TYPES for a document). upstream: the how and constraints
    of the workflows this one depends on, squashed: an id one of them is told to create may be checked here too.
    ranges (J2_REALIZATION): an id range in how or constraints names every id in it, and a digit inside a name in a
    check's text (#pad-4, j2-btn) is not a count."""
    out, checks = [], a.get("verification") if isinstance(a.get("verification"), list) else []
    # (tap-targets: expect is the minimum size in px, so a label like "at least 44x44px" must expect 44)
    told = squash(json.dumps([a.get("how"), a.get("constraints")], ensure_ascii=False)) + (" " + upstream if upstream else "")
    if ranges:
        told += " " + range_ids(told)
    for n, c in enumerate(checks, 1):
        lab = f"expected output item {c['item']}" if isinstance(c, dict) and c.get("item") else f"check {n}"
        if not isinstance(c, dict):
            out.append(f"{lab} isn't an object with text, type, target and expect")
            continue
        t, target, label = c.get("type"), c.get("target"), str(c.get("text") or "")
        if t == "lld-slice":          # written by code at registration (the LLD slice is in "spec"), never by Adaptation
            continue
        if t in HEADING_CHECKS:
            target = heading_target(target)
        if t not in types:
            out.append(f"{lab} has type {t!r}; allowed: {', '.join(types)}")
        elif t in HEADING_CHECKS and not str(target or "").strip():
            out.append(f"{lab} {t} needs the heading text as its target")
        elif t == "section-not-empty" and as_int(c.get("expect")) < 1:
            out.append(f"{lab} section-not-empty expects {c.get('expect')!r}; expect is the minimum word count (1 or more)")
        elif t == "ids-unique" and not re.fullmatch(r"[A-Z]{1,3}", str(target or "")):
            out.append(f"{lab} ids-unique targets {target!r}: only an id prefix such as \"R\" (for R1, R2, ...)")
        elif t == "named-test" and target not in NAMED_TESTS:
            out.append(f"{lab} names test {target!r}; allowed: {', '.join(NAMED_TESTS)}")
        elif t == "facts-present":
            facts = target if isinstance(target, list) else [target]
            bad = [f for f in facts if not str(f or "").strip() or squash(f) not in source]
            if bad:
                out.append(f"{lab} facts-present lists {bad}, not found in the charter or the Requester's answers")
        elif t == "behavior-steps":
            about = squash(json.dumps([a.get("goal"), a.get("how"), label], ensure_ascii=False))
            out += [f"{lab} behavior {x}" for x in behavior_problems(target, told, source, about)]
        elif t in COUNTED:
            if bad := target_problem(target, told, source):
                out.append(f"{lab} {t} {bad}")
            nums = [int(x) for x in re.findall(r"\d+", NAME_WITH_DIGIT.sub(" ", label) if ranges else label)]
            if t == "tap-targets":      # "tappable at 360px": the screen width is never the count or the size
                nums = [x for x in nums if x not in j2_check.WIDTHS.values()]
            if nums and as_int(c.get("expect"), None) not in nums:
                out.append(f"{lab} says {label!r} but expects {c.get('expect')!r}")
    if SPECIFIC.search(str(a.get("goal") or "")) and not any(
            isinstance(c, dict) and c.get("type") in ("facts-present", "named-test", "behavior-steps") for c in checks):
        out.append("its goal names a time, price, number or email but no check is facts-present or a named test")
    if "behavior-steps" in types and (m := ACTION_GOAL.search(str(a.get("goal") or ""))) and not any(
            isinstance(x, dict) and x.get("kind") == "behavior" for x in a.get("expected_output") or []):
        out.append(f"its goal describes an action ({m.group(0)!r}) but no expected output item is a behavior that "
                   "proves it works")
    return out


VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source", "track", "wbr"}
OPTIONAL_END = {"p", "li", "dt", "dd", "tr", "td", "th", "option", "thead", "tbody", "tfoot", "colgroup", "html", "head",
                "body"}


class PageScan(html.parser.HTMLParser):
    """One pass over the built page: every tag with its attributes, unclosed and stray tags, forms with a submit
    control, and the visible text."""
    def __init__(self, page):
        super().__init__(convert_charrefs=True)
        self.tags, self.stack, self.stray, self.text, self.forms, self.in_form = [], [], [], [], [], False
        self.unclosed, self.form_attrs, self.visible, self.anc, self.open_idx = [], [], [], [], []
        self.feed(page)
        self.close()
        self.unclosed += [t for t in self.stack if t not in OPTIONAL_END]

    def handle_starttag(self, tag, attrs, alone=False):
        a = {k: v or "" for k, v in attrs}
        self.anc.append(tuple(self.open_idx))      # indices in tags of the elements this one is inside
        self.tags.append((tag, a))
        if tag == "form":
            self.forms.append(0)
            self.form_attrs.append(a)
            self.in_form = True
        elif self.in_form and (tag == "button" and a.get("type", "submit").lower() == "submit"
                               or tag == "input" and a.get("type", "").lower() in ("submit", "image")):
            self.forms[-1] += 1
        if tag not in VOID and not alone:
            self.stack.append(tag)
            self.open_idx.append(len(self.tags) - 1)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs, alone=True)

    def handle_endtag(self, tag):
        if tag == "form":
            self.in_form = False
        if tag not in self.stack:
            if tag not in OPTIONAL_END and tag not in VOID:
                self.stray.append(tag)
            return
        while self.stack:
            t = self.stack.pop()
            self.open_idx.pop()
            if t == tag:
                break
            if t not in OPTIONAL_END:
                self.unclosed.append(t)

    def handle_data(self, d):
        self.text.append(d)
        if not (self.stack and self.stack[-1] in ("script", "style")):
            self.visible.append(d)

    def matching(self, target):
        name, kind = str(target).lstrip("#."), str(target)[:1]
        return [t for t, a in self.tags if (kind != "." and a.get("id") == name)
                or (kind != "#" and name in a.get("class", "").split())]


def run_check(c, site_dir, page, doc=False):
    """{actual, pass, evidence} for one check against the built site (doc=True: page is a markdown document). Code
    only."""
    t, target, expect = c.get("type"), c.get("target"), c.get("expect")
    if t == "deed-send":      # reads the send receipt in this run's receipts.jsonl; never sends
        return j2_send.check(c, os.path.dirname(site_dir))
    if doc:
        return doc_check(c, page)
    try:
        if t == "file-exists":
            path = os.path.normpath(os.path.join(site_dir, str(target)))
            ok = path.startswith(os.path.normpath(site_dir) + os.sep) and os.path.isfile(path)
            want = True if expect is None else bool(expect)
            return {"actual": ok, "pass": ok == want, "evidence": f"{target} {'exists' if ok else 'not found'} in the build"}
        if t == "grep-count":
            path = os.path.normpath(os.path.join(site_dir, str(c.get("file") or "index.html")))
            text = open(path, encoding="utf-8").read() if path.startswith(os.path.normpath(site_dir)) \
                and os.path.isfile(path) else ""
            n = text.count(str(target).lstrip("#."))
            return {"actual": n, "pass": n == as_int(expect, -1),
                    "evidence": f"{target!r} found {n} time(s) in {c.get('file') or 'index.html'}, expected {expect}"}
        if t == "lld-slice":
            ok, ev = lld_check(page, c.get("spec") or {})
            return {"actual": ok, "pass": ok, "evidence": ev}
        if t == "behavior-steps":
            return j2_check.run_steps(target if isinstance(target, list) else [], lambda ops: j2_check.behavior(page, ops))
        scan = PageScan(page)
        if t == "element-count":
            n = len(scan.matching(target))
            return {"actual": n, "pass": n == as_int(expect, -1),
                    "evidence": f"{n} element(s) match {target!r}, expected {expect}"}
        if t == "tap-targets":
            r, want = j2_check.tap_targets(page, str(target), as_int(expect, j2_check.TAP_MIN)), as_int(expect, 44)
            if not r.get("available"):      # no browser: the check didn't run, which says nothing about the build
                return {"actual": None, "pass": None, "evidence": f"could not run: {r.get('note')}"}
            ok = r["roots"] > 0 and r["checked"] > 0 and not r["small_count"]
            return {"actual": {k: r[k] for k in ("roots", "checked", "small_count", "small")}, "pass": ok,
                    "evidence": (f"nothing on the page matches {target!r}" if not r["roots"] else
                                 f"no buttons or links inside {target!r}" if not r["checked"] else
                                 f"{r['small_count']} of {r['checked']} tap target(s) in {target!r} under {want}x{want}px "
                                 f"at 360px: {', '.join(r['small'])}" if r["small_count"] else
                                 f"all {r['checked']} tap target(s) in {target!r} at least {want}x{want}px at 360px")}
        if t == "facts-present":
            text = squash(" ".join(scan.text))
            facts = target if isinstance(target, list) else [target]
            missing = [f for f in facts if squash(f) not in text]
            return {"actual": {"missing": missing}, "pass": not missing,
                    "evidence": f"missing from the page: {missing}" if missing else f"all {len(facts)} fact(s) on the page"}
        if target == "page-parses":
            body = any(x == "body" for x, _ in scan.tags)
            bad = scan.unclosed + [f"/{x}" for x in scan.stray]
            return {"actual": {"body": body, "unbalanced": bad}, "pass": body and not bad,
                    "evidence": ("page parses, tags balanced" if body and not bad else
                                 f"{'no <body>; ' if not body else ''}unbalanced tags: {bad[:8]}")}
        if target == "links-resolve":
            ids = {a.get("id") for _, a in scan.tags}
            broken = [a["href"] for _, a in scan.tags if a.get("href", "").startswith("#") and len(a["href"]) > 1
                      and a["href"][1:] not in ids]
            return {"actual": broken, "pass": not broken,
                    "evidence": f"in-page links with no target: {broken[:8]}" if broken else "every in-page link resolves"}
        if target == "form-has-submit":
            ok = bool(scan.forms) and all(scan.forms)
            return {"actual": scan.forms, "pass": ok,
                    "evidence": (f"{len(scan.forms)} form(s), each with a submit control" if ok else
                                 "no form on the page" if not scan.forms else
                                 f"{scan.forms.count(0)} of {len(scan.forms)} form(s) have no submit control")}
    except (OSError, ValueError) as e:
        return {"actual": None, "pass": False, "evidence": f"the check couldn't run: {e}"}
    return {"actual": None, "pass": False, "evidence": f"unknown check {t} {target}"}


# ---- documents (a PRD): markdown read by code, no browser
MD_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")


def strip_md(t):
    """Text without markdown emphasis and code marks, for matching facts and counting words."""
    return re.sub(r"[*_`]+", "", str(t))


def md_sections(doc):
    """[(level, heading, body)] for every markdown heading; a body runs to the next heading of the same or a higher
    level, so it includes its subsections."""
    lines = str(doc).splitlines()
    heads = [(i, len(m[1]), strip_md(m[2]).strip()) for i, ln in enumerate(lines) if (m := MD_HEADING.match(ln))]
    return [(lv, h, "\n".join(lines[i + 1:next((j for j, l2, _ in heads[k + 1:] if l2 <= lv), len(lines))]))
            for k, (i, lv, h) in enumerate(heads)]


def md_section(doc, heading):
    """The body under the first heading that starts with `heading` (case and spacing ignored), or None."""
    want = squash(strip_md(heading))
    return next((b for _, h, b in md_sections(doc) if want and squash(h).startswith(want)), None)


def md_replace_section(doc, heading, text):
    """The document with the body under `heading` replaced by `text` (the heading line is kept, the rest of the
    document is untouched). The document unchanged when it has no such heading."""
    lines, want = str(doc).splitlines(), squash(strip_md(heading))
    heads = [(i, len(m[1]), squash(strip_md(m[2]))) for i, ln in enumerate(lines) if (m := MD_HEADING.match(ln))]
    for k, (i, lv, h) in enumerate(heads):
        if want and h.startswith(want):
            end = next((j for j, l2, _ in heads[k + 1:] if l2 <= lv), len(lines))
            body = str(text).strip("\n")
            return "\n".join(lines[:i + 1] + ["", body, ""] + lines[end:])
    return str(doc)


def md_words(text):
    return len(re.findall(r"\w+", strip_md(text)))


def md_ids(text, prefix):
    """Ids defined in the text (R1, R2, ...): an id that starts a line, a list item or a table row, never one
    mentioned inside a sentence. Returned as numbers, in order."""
    return [int(n) for n in re.findall(rf"(?m)^\s*(?:[-*+]\s+|\d+[.)]\s+|\|\s*)?(?:\*\*|__)?{re.escape(prefix)}-?(\d+)\b",
                                       str(text))]


HEADING_CHECKS = ("heading-present", "section-not-empty")


def heading_target(t):
    """A heading check's target as heading text: leading "#" characters and spaces stripped ("## Problem" -> "Problem")."""
    return re.sub(r"^[#\s]+", "", str(t or "")).strip()


# An expected-output item names one concrete thing that must exist when the atom is done; its kind says which check
# type verifies it. An output allows the kinds whose check type it runs (a website's or a document's).
EXPECT_KINDS = {"section": "element-count", "element": "element-count", "facts": "facts-present",
                "file": "file-exists", "test": "named-test", "tap": "tap-targets",
                "heading": "heading-present", "ids": "ids-unique", "behavior": "behavior-steps"}


def expect_kinds(types):
    return [k for k, t in EXPECT_KINDS.items() if t in types]


def expected_check(item, n):
    """The check code derives from expected-output item n (1-based), or None when its kind has no check."""
    kind, target = item.get("kind"), item.get("target")
    t = EXPECT_KINDS.get(kind)
    if not t:
        return None
    if t == "behavior-steps":
        target = list(item.get("steps") or [])
    expect = (as_int(item.get("count"), 1) if t == "element-count" else
              as_int(item.get("count"), j2_check.TAP_MIN) if t == "tap-targets" else True)
    return {"text": str(item.get("text") or ""), "type": t, "expect": expect, "item": n,
            "target": heading_target(target) if t in HEADING_CHECKS else target}


def expected_problems(a, types):
    """Expected-output items no allowed check can verify: not an object with text, kind and target, or a kind
    whose check type this output doesn't run."""
    out, kinds = [], expect_kinds(types)
    for n, x in enumerate(a.get("expected_output") if isinstance(a.get("expected_output"), list) else [], 1):
        if not isinstance(x, dict) or not str(x.get("text") or "").strip() or \
                not x.get("steps" if x.get("kind") == "behavior" else "target"):
            out.append(f"expected output item {n} isn't an object with text, kind and target (steps for a behavior)")
        elif x.get("kind") not in kinds:
            out.append(f"expected output item {n} has kind {x.get('kind')!r}, which no check can verify; allowed: "
                       + ", ".join(kinds))
    return out


def cover_expected(a):
    """Every expected-output item covered by at least one check, in place: a check already tagged with the item's
    number, or one with the same type and target (tagged now), covers it; otherwise code adds the derived check.
    ponytail: a corrected check that drops its item number and changes the target gets the derived one back."""
    items = a.get("expected_output") if isinstance(a.get("expected_output"), list) else []
    ver = a["verification"] = a.get("verification") if isinstance(a.get("verification"), list) else []
    for n, x in enumerate(items, 1):
        d = expected_check(x, n) if isinstance(x, dict) else None
        if not d or any(isinstance(c, dict) and c.get("item") == n for c in ver):
            continue
        same = next((c for c in ver if isinstance(c, dict) and "item" not in c and c.get("type") == d["type"]
                     and json.dumps(c.get("target")) == json.dumps(d["target"])), None)
        if same:
            same["item"] = n
        else:
            ver.append(d)
    return a


def normalise_checks(atoms):
    """Every heading check's target in these atoms as heading text, and every expected-output item covered by a
    check, in place, so workflows.json and the receipts store the form the check runs on."""
    for a in atoms or []:
        for c in (a.get("verification") if isinstance(a, dict) and isinstance(a.get("verification"), list) else []):
            if isinstance(c, dict) and c.get("type") in HEADING_CHECKS:
                c["target"] = heading_target(c.get("target"))
        if isinstance(a, dict) and a.get("expected_output"):
            cover_expected(a)
    return atoms


def corrected_check(c, doc):
    """(corrected check, reason) when a failed heading or text check fails only because of how its target is written:
    read as heading text (no #, emphasis, trailing colon or the word "section") it passes, or its text is in the
    document written slightly differently (spacing, punctuation, case). None when the check fails on its merits."""
    t, target = c.get("type"), c.get("target")
    if t in HEADING_CHECKS:
        loose = re.sub(r"\s*\b(?:section|heading)\s*$", "", re.sub(r"[*_`]+|[:.]+$", "", heading_target(target)),
                       flags=re.I).strip()
        new = {**c, "target": loose}
        if loose and loose != target and doc_check(new, doc)["pass"]:
            return new, f"{t} target {target!r} passes as the heading text {loose!r}"
    if t == "facts-present":
        facts, text, fixed = target if isinstance(target, list) else [target], strip_md(doc), []
        for f in facts:
            core = re.sub(r"[^a-z0-9]", "", str(f).lower())
            m = None if squash(strip_md(f)) in squash(text) or not core else \
                re.search(r"[^a-z0-9]*".join(map(re.escape, core)), text, re.I)
            if not m and squash(strip_md(f)) not in squash(text):
                return None
            fixed.append(" ".join(m.group(0).split()) if m else f)
        new = {**c, "target": fixed}
        if fixed != facts and doc_check(new, doc)["pass"]:
            return new, f"facts-present target {facts!r} is in the document as {fixed!r}"
    return None


def doc_check(c, doc):
    """{actual, pass, evidence} for one check against a markdown document. Code only."""
    t, target, expect = c.get("type"), c.get("target"), c.get("expect")
    if t in HEADING_CHECKS:
        target = heading_target(target)
    if t == "facts-present":
        text = squash(strip_md(doc))
        facts = target if isinstance(target, list) else [target]
        missing = [f for f in facts if squash(strip_md(f)) not in text]
        return {"actual": {"missing": missing}, "pass": not missing,
                "evidence": f"missing from the document: {missing}" if missing else f"all {len(facts)} fact(s) in the document"}
    if t == "heading-present":
        ok = md_section(doc, target) is not None
        return {"actual": ok, "pass": ok, "evidence": f"heading {target!r} {'found' if ok else 'not found'} in the document"}
    if t == "section-not-empty":
        body, need = md_section(doc, target), max(1, as_int(expect, 1))
        n = md_words(body) if body is not None else 0
        return {"actual": n, "pass": body is not None and n >= need,
                "evidence": f"no heading {target!r}" if body is None else f"section {target!r} has {n} word(s), needs {need}"}
    if t == "ids-unique":
        ids = md_ids(doc, str(target))
        dup = sorted({i for i in ids if ids.count(i) > 1})
        return {"actual": ids, "pass": bool(ids) and not dup,
                "evidence": (f"no {target}1, {target}2, ... ids in the document" if not ids else
                             f"duplicate id(s): {', '.join(f'{target}{i}' for i in dup)}" if dup else
                             f"{len(ids)} unique id(s): {', '.join(f'{target}{i}' for i in ids)}")}
    return {"actual": None, "pass": False, "evidence": f"check type {t!r} doesn't run on a document"}


def md_blocks(doc):
    """The document as blocks for the preview and the .docx: ("h", level, text), ("ol"|"ul", [items]),
    ("table", [[cells]]) and ("p", text). ponytail: the markdown the builders write (headings, lists, pipe tables,
    bold); a markdown library if they need more."""
    out, para = [], []

    def flush():
        if para:
            out.append(("p", " ".join(para)))
            para.clear()
    for ln in str(doc).splitlines():
        s = ln.strip()
        if m := MD_HEADING.match(ln):
            flush()
            out.append(("h", len(m[1]), m[2].strip()))
        elif s.startswith("|"):
            flush()
            cells = [x.strip() for x in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", x) for x in cells if x):
                continue                      # the header separator row
            if out and out[-1][0] == "table":
                out[-1][1].append(cells)
            else:
                out.append(("table", [cells]))
        elif m := re.match(r"(?:(\d+)[.)]|[-*+])\s+(.*)", s):
            flush()
            kind = "ol" if m[1] else "ul"
            if out and out[-1][0] == kind:
                out[-1][1].append(m[2])
            else:
                out.append((kind, [m[2]]))
        elif s:
            para.append(s)
        else:
            flush()
    flush()
    return out


def md_inline(t):
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", html.escape(str(t)))


def md_html(doc, title="Document"):
    """The document rendered as a page, for the Preview tab."""
    parts = []
    for b in md_blocks(doc):
        if b[0] == "h":
            parts.append(f"<h{b[1]}>{md_inline(b[2])}</h{b[1]}>")
        elif b[0] in ("ol", "ul"):
            parts.append(f"<{b[0]}>" + "".join(f"<li>{md_inline(x)}</li>" for x in b[1]) + f"</{b[0]}>")
        elif b[0] == "table":
            rows = b[1]
            parts.append("<table><tr>" + "".join(f"<th>{md_inline(x)}</th>" for x in rows[0]) + "</tr>"
                         + "".join("<tr>" + "".join(f"<td>{md_inline(x)}</td>" for x in r) + "</tr>" for r in rows[1:])
                         + "</table>")
        else:
            parts.append(f"<p>{md_inline(b[1])}</p>")
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" '
            f'content="width=device-width, initial-scale=1"><title>{html.escape(title)}</title><style>body{{font:16px/1.5 '
            "system-ui;max-width:760px;margin:0 auto;padding:24px;color:#1f2933}table{border-collapse:collapse;width:100%}"
            "th,td{border:1px solid #cbd2d9;padding:6px 8px;text-align:left;vertical-align:top}th{background:#f5f7f9}"
            "</style></head><body>" + "".join(parts) + "</body></html>")


def export_docx(doc, path):
    """Write the markdown document as a .docx: headings, numbered and bulleted lists, tables and bold come through.
    Returns None when written, else why not (python-docx missing); the markdown file is kept either way."""
    try:
        import docx
    except ImportError as e:
        return f"python-docx isn't installed ({e}), so {os.path.basename(path)} wasn't written"

    def runs(p, text):
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):     # odd parts were inside **...**
            if part:
                p.add_run(strip_md(part)).bold = bool(i % 2)
    d = docx.Document()
    for b in md_blocks(doc):
        if b[0] == "h":
            d.add_heading(strip_md(b[2]), level=min(b[1] - 1, 9))       # "#" is the title, "##" Heading 1
        elif b[0] in ("ol", "ul"):
            for x in b[1]:
                runs(d.add_paragraph(style="List Number" if b[0] == "ol" else "List Bullet"), x)
        elif b[0] == "table":
            rows = b[1]
            t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for k, cell in enumerate(row):
                    runs(t.cell(r, k).paragraphs[0], cell)
        else:
            runs(d.add_paragraph(), b[1])
    d.save(path)
    return None


# Necessary conditions (conditions/<type>.json): each names a code check below (or a list of them) and its parameters.
CONDITIONS_DIR = os.path.join(HERE, "conditions")
TBC_WORDS = re.compile(r"to be confirmed|\bTB[CD]\b", re.I)


def load_conditions(name):
    with open(os.path.join(CONDITIONS_DIR, name), encoding="utf-8") as f:
        return json.load(f)["conditions"]


PHONE = re.compile(r"\+?\d[\d\s().-]{6,}\d")


def doc_numbers(text):
    """Numbers and emails a document states, leaving out ids (R1, G2, D-1) and list numbering. A phone number (10 or
    more digits, or 8 or more after a +) counts as its digits, however it is spaced."""
    phones = {re.sub(r"\D", "", m) for m in PHONE.findall(str(text))
              if len(re.sub(r"\D", "", m)) >= (8 if m.startswith("+") else 10)}
    text = PHONE.sub(lambda m: " " if re.sub(r"\D", "", m.group(0)) in phones else m.group(0), str(text))
    text = re.sub(r"\b[A-Z]{1,3}-?\d+\b", " ", strip_md(text))
    text = re.sub(r"(?m)^\s*\d+[.)]\s", " ", text)
    return set(re.findall(r"\d+(?:[.,]\d+)*", text)) | set(EMAIL.findall(text)) | phones


ASSUMPTION = re.compile(r"\bAssumption:\s*(.+)", re.I)
ASSUMPTIONS_HEADING = "Assumptions to confirm"
NO_ACTIONS = "have no actions in this version"
WINDOW = re.compile(r"\b(?:per|a|an|each|every|over|within|during|in|by|after|for)\s+(?:the\s+)?(?:first\s+|next\s+)?"
                    r"(?:\d+\s+)?(?:days?|weeks?|months?|terms?|semesters?|quarters?|years?|hours?|sessions?|launch)\b|"
                    r"\b(?:daily|weekly|monthly|quarterly|yearly|annually)\b|measurement window not set", re.I)


def doc_assumptions(doc):
    """The text of every "Assumption:" line a builder wrote (not the closing section code makes from them)."""
    body = re.split(rf"(?m)^##\s+{ASSUMPTIONS_HEADING}\s*$", str(doc))[0]
    return [m.group(1).strip() for ln in strip_md(body).splitlines() if (m := ASSUMPTION.search(ln))]


VERSION_NOTE = re.compile(r"(?m)^\s*(?:[-*+]\s*)?Version \d+: (?:Q-\d+, )?.*$")      # written by code, not by an agent


def without_assumptions(doc):
    """The document without its Assumption: lines, the closing assumptions section and the version notes code writes
    when a board answer changes a section: what must come from the brief."""
    body = re.split(rf"(?m)^##\s+{ASSUMPTIONS_HEADING}\s*$", str(doc))[0]
    body = VERSION_NOTE.sub("", body)
    return "\n".join(ln for ln in body.splitlines() if not ASSUMPTION.search(strip_md(ln)))


def doc_names(text):
    """Proper names a document states: capitalised words that don't start a sentence, line, list item or table cell.
    Ids (R1, G2), section headings and possessives are left out. ponytail: a capitalised common word mid-sentence
    ("the Requirements section") counts unless it is a section word; a real NER pass is the upgrade."""
    skip = {w.lower() for s in ALL_DOC_SECTIONS + (ASSUMPTIONS_HEADING,) for w in s.split()} | {"i", "assumption"}
    out = set()
    for ln in strip_md(text).splitlines():
        if MD_HEADING.match(ln) or re.fullmatch(r"\s*\|?[\s|:-]*", ln):
            continue
        for seg in re.split(r"(?<=[.!?:;])\s+|\|", ln):
            words = re.findall(r"[A-Za-z][\w'-]*", re.sub(r"^\s*(?:[-*+]|\d+[.)])\s*", "", seg))
            cap = [w[0].isupper() and w.lower() not in skip and not re.fullmatch(r"[A-Z]{1,3}-?\d+", w) for w in words]
            # a sentence's first word is a name only when the next is capitalised too ("Northfield College")
            out |= {re.sub(r"'s$", "", w) for i, w in enumerate(words) if cap[i] and (i or len(cap) > 1 and cap[1])}
    return out


def user_groups(goal):
    """The user groups the charter names: user_groups ([str] or [{"name"}]), else split from "users"."""
    g = goal or {}
    raw = [x.get("name") if isinstance(x, dict) else x for x in g.get("user_groups") or []]
    if not raw and g.get("users"):
        raw = re.split(r",|;|\band\b", re.sub(r"\([^)]*\)", "", str(g["users"])))
    return [re.sub(r"^(?:the|its|our)\s+", "", str(x).strip().rstrip("."), flags=re.I) for x in raw if str(x or "").strip(" .")]


def condition_met(cond, doc, goal, source, sections, page=None, mech=None):
    """(met, evidence) for one necessary condition, checked by code. doc: the document, or a website's visible text;
    page: a website's HTML; mech: its render check (j2_check.check). source: the charter and the Requester's own
    words, squashed; sections: the output's required section headings."""
    k, g = cond.get("check"), goal or {}
    if isinstance(k, list):             # several checks, all must hold
        rs = [condition_met({**cond, "check": x}, doc, goal, source, sections, page, mech) for x in k]
        return all(m for m, _ in rs), "; ".join(e for _, e in rs)
    if k in WEB_CONDITIONS:
        return WEB_CONDITIONS[k](cond, page or "", mech or {})
    if k == "sections-present":
        need = max(1, as_int(cond.get("min_words"), 1))
        bad = [f"{s} ({'missing' if md_section(doc, s) is None else f'{md_words(md_section(doc, s))} words'})"
               for s in sections if md_section(doc, s) is None or md_words(md_section(doc, s)) < need]
        return not bad, f"missing or empty: {', '.join(bad)}" if bad else f"all {len(sections)} sections present, none empty"
    if k == "ids-sequential":
        body, p = md_section(doc, cond["section"]), cond["prefix"]
        ids = md_ids(body or "", p)
        ok = bool(ids) and ids == list(range(1, len(ids) + 1))
        return ok, (f"no {cond['section']} section" if body is None else f"no {p}1, {p}2, ... ids" if not ids else
                    f"{p}1 to {p}{len(ids)}, unique and in order" if ok else
                    f"ids {', '.join(f'{p}{i}' for i in ids)}: not unique and sequential from {p}1")
    if k == "facts-present-all":
        listed = [str(f) for f in g.get("facts") or [] if str(f or "").strip()]
        facts = listed or [str(f) for f in [g.get("product_name") or g.get("site_name")] if str(f or "").strip()]
        text = squash(strip_md(doc))
        missing = [f for f in facts if squash(strip_md(f)) not in text]
        if not listed:      # the site-name fallback: say so, never "all 1 brief fact(s)"
            return bool(facts) and not missing, ("no facts list and no site name" if not facts else
                                                 f"no facts list; the site name {facts[0]!r} "
                                                 + ("is missing" if missing else "appears"))
        return not missing, f"missing: {missing}" if missing else f"all {len(facts)} brief fact(s) appear"
    if k == "groups-covered":
        groups, req = user_groups(g), squash(strip_md(md_section(doc, cond.get("section", "Requirements")) or ""))
        said = squash(strip_md(doc))
        stem = lambda x: squash(x)[:-1] if squash(x).endswith("s") else squash(x)     # "students" -> "student"
        miss = [x for x in groups if stem(x) not in req and not re.search(
            rf"{re.escape(squash(x))} ha(?:ve|s) no actions in this version", said)]
        return not miss, (f"no requirement and no \"{NO_ACTIONS}\" statement for: {', '.join(miss)}" if miss else
                          f"each of the {len(groups)} user group(s) has a requirement or a \"no actions\" statement"
                          if groups else "the charter names no user groups")
    if k == "metric-window":
        body = md_section(doc, cond.get("section", "Success metrics")) or ""
        rows = [ln for ln in strip_md(body).splitlines() if re.match(r"\s*(?:[-*+]|\d+[.)]|\|)", ln)
                and not re.fullmatch(r"\s*\|?[\s|:-]*", ln)]
        num = [r for r in rows if re.search(r"\d", re.sub(r"\b[A-Z]{1,3}-?\d+\b", "", r))]
        bad = [" ".join(r.strip(" |-").split())[:60] for r in num if not WINDOW.search(r)]
        return not bad, (f"no measurement window: {'; '.join(bad)}" if bad else
                         f"each of the {len(num)} numeric metric(s) states a window" if num else "no numeric metrics")
    if k == "no-invented-facts":
        src = doc_numbers(source)
        if cond.get("names"):           # a document: Assumption: lines may go beyond the brief; proper names count too
            doc = without_assumptions(doc)
            names = sorted(x for x in doc_names(doc) if not re.search(rf"\b{re.escape(x.lower())}\b", source))
            extra = sorted(x for x in doc_numbers(doc) if x.lower() not in src)
            bad = extra + names
            return not bad, (f"not in the charter or the Requester's answers, and not marked Assumption: "
                             f"{', '.join(bad[:8])}" if bad else "every number, email and proper name comes from the "
                             "charter or the Requester's answers, or sits in an Assumption: line")
        if page is not None:      # a website: also the numbers and emails in its tel:, mailto:, wa.me links and form actions
            doc = " ".join([doc] + [v for t, a in PageScan(page).tags for n, v in a.items()
                                    if n in ("href", "action") and re.match(r"tel:|mailto:|https?://wa\.me/", v, re.I)])
        extra = sorted(x for x in doc_numbers(doc) if x.lower() not in src)
        return not extra, (f"not in the charter or the Requester's answers: {', '.join(extra[:8])}" if extra else
                           f"every number{', phone' if page is not None else ''} and email comes from the charter or "
                           "the Requester's answers")
    if k == "nothing-to-be-confirmed":
        m, open_ = TBC_WORDS.search(str(doc)), g.get("to_be_confirmed") or []
        return not m and not open_, (f"the document says {m.group(0)!r}" if m else
                                     f"{len(open_)} fact(s) still to be confirmed" if open_ else "nothing to be confirmed")
    if k == "ids-covered":
        p, src_body, in_body = cond["prefix"], md_section(doc, cond["from"]), md_section(doc, cond["in"]) or ""
        ids = md_ids(src_body or "", p)
        miss = [f"{p}{i}" for i in ids if not re.search(rf"\b{p}-?{i}\b", in_body)]
        return bool(ids) and not miss, (f"no {p}1, {p}2, ... ids in {cond['from']}" if not ids else
                                        f"no {cond['in']} entry for {', '.join(miss)}" if miss else
                                        f"each of {', '.join(f'{p}{i}' for i in ids)} has a {cond['in']} entry")
    if k == "headings-unique":
        heads = [squash(h) for lv, h, _ in md_sections(doc) if lv == 2]
        dup = sorted(set(h for h in heads if heads.count(h) > 1))
        return not dup, f"repeated heading(s): {', '.join(dup)}" if dup else f"{len(heads)} section headings, all unique"
    if k == "chain-ids":      # every id the earlier document defines appears in a section of this one
        ids = [str(x) for x in (g.get("chain") or {}).get(cond["key"]) or []]
        body = md_section(doc, cond["in"]) or ""
        miss = [i for i in ids if not re.search(rf"\b{re.escape(i)}\b", strip_md(body))]
        return not miss, (f"no {cond['in']} entry for {', '.join(miss)}" if miss else
                          f"each of {', '.join(ids)} appears in {cond['in']}" if ids else
                          f"no {cond['label']} to trace (that document wasn't requested)")
    if k == "components-have-responsibility":
        comps = parse_hld(doc)
        bad = [c["id"] for c in comps if md_words(c["responsibility"]) < 3]
        return bool(comps) and not bad, ("no C1, C2, ... components" if not comps else
                                         f"no responsibility for {', '.join(bad)}" if bad else
                                         f"each of the {len(comps)} component(s) states its responsibility")
    if k in ("lld-ids-unique", "lld-elements", "lld-fields"):
        return lld_condition(k, doc)
    return False, f"no code check named {k!r}"


def check_conditions(conds, doc, goal, source, sections, page=None, mech=None):
    """Each condition's result. "severity": "warn" in the conditions file: shown, never blocks workable(). P4 with
    "names" also lists the document's Assumption: lines."""
    return [{"id": c["id"], "text": c["text"], "check": c["check"],
             **dict(zip(("met", "evidence"), condition_met(c, doc, goal, source, sections, page, mech))),
             **({"severity": "warn"} if c.get("severity") == "warn" else {}),
             **({"assumptions": doc_assumptions(doc)} if c.get("names") else {})} for c in conds]


def blocks(c):
    """An unmet condition that stops workable() (a warning never does)."""
    return not c["met"] and c.get("severity") != "warn"


def _page_parses(cond, page, mech):
    r = run_check({"type": "named-test", "target": "page-parses"}, "", page)
    return bool(r["pass"]), r["evidence"]


def _rendered(mech):
    return None if mech.get("available") else f"could not run: {mech.get('note') or 'no render check'}"


def _no_js_errors(cond, page, mech):
    if _rendered(mech):
        return False, _rendered(mech)
    e = mech.get("js_errors") or []
    return not e, f"{len(e)} JS error(s) on load: {'; '.join(e[:3])}" if e else "no JS errors on load"


def _no_overflow(cond, page, mech):
    if _rendered(mech):
        return False, _rendered(mech)
    widths = cond.get("widths") or list(j2_check.WIDTHS.values())
    seen = {(mech.get(n) or {}).get("width", w): (mech.get(n) or {}).get("overflow_px") for n, w in j2_check.WIDTHS.items()}
    bad = [f"{w}px {'not measured' if seen.get(w) is None else f'overflows by {seen[w]}px'}" for w in widths if seen.get(w) != 0]
    return not bad, "; ".join(bad) if bad else f"no horizontal overflow at {' and '.join(f'{w}px' for w in widths)}"


TEL = re.compile(r"tel:\+?[\d\-.() ]+", re.I)
WA_ME = re.compile(r"https?://wa\.me/\d{8,15}(?:\?.*)?", re.I)


def _links_well_formed(cond, page, mech):
    scan = PageScan(page)
    ids, hrefs = {a.get("id") for _, a in scan.tags}, [a["href"].strip() for t, a in scan.tags if t == "a" and a.get("href")]
    bad = [f"{h} (no such id)" for h in hrefs if h.startswith("#") and len(h) > 1 and h[1:] not in ids]
    bad += [h for h in hrefs if h.lower().startswith("tel:") and not (TEL.fullmatch(h) and len(re.sub(r"\D", "", h)) >= 7)]
    bad += [h for h in hrefs if re.match(r"https?://(?:www\.)?wa\.me", h, re.I) and not WA_ME.fullmatch(h)]
    n = sum(h.startswith("#") or re.match(r"tel:|https?://(?:www\.)?wa\.me", h, re.I) is not None for h in hrefs)
    return not bad, f"broken or malformed: {', '.join(bad[:8])}" if bad else \
        f"all {n} in-page, tel: and wa.me link(s) resolve and are well formed"


def _forms_have_destination(cond, page, mech):
    scan, bad = PageScan(page), []
    for i, (subs, a) in enumerate(zip(scan.forms, scan.form_attrs), 1):
        act = a.get("action", "").strip()
        if not subs:
            bad.append(f"form {i} has no submit control")
        if not act or act.startswith("#") or act.lower().startswith("javascript:"):
            bad.append(f"form {i} has no destination (action or mailto)")
    return not bad, "; ".join(bad) if bad else (f"{len(scan.forms)} form(s), each with a submit control and a "
                                                "destination" if scan.forms else "no forms on the page")


def _nav_and_main_action(cond, page, mech):
    """A <nav> whose in-page links reach every <section> outside the hero, and a link or button in the hero (not a nav
    link). The hero: the element whose id or class says hero, else the first header or section."""
    scan = PageScan(page)
    tags, anc = scan.tags, scan.anc
    hero = next((i for i, (t, a) in enumerate(tags) if "hero" in f"{a.get('id', '')} {a.get('class', '')}".lower()), None)
    if hero is None:
        hero = next((i for i, (t, _) in enumerate(tags) if t in ("header", "section")), None)
    navs = {i for i, (t, _) in enumerate(tags) if t == "nav"}
    in_nav = lambda i: bool(navs & set(anc[i]))
    linked = {a["href"][1:] for i, (t, a) in enumerate(tags) if t == "a" and a.get("href", "").startswith("#") and in_nav(i)}
    secs = [(i, a.get("id")) for i, (t, a) in enumerate(tags) if t == "section" and i != hero and hero not in anc[i]]
    unreached = [f"#{sid}" if sid else f"section {n} (no id)" for n, (i, sid) in enumerate(secs, 1) if sid not in linked]
    acts = [i for i, (t, a) in enumerate(tags) if hero is not None and hero in anc[i] and not in_nav(i)
            and (t == "a" and a.get("href") or t == "button")]
    bad = (["no <nav> on the page"] if not navs else
           [f"the nav doesn't reach {', '.join(unreached[:8])}"] if unreached else []) \
        + (["no hero"] if hero is None else ["no link or button in the hero"] if not acts else [])
    return not bad, "; ".join(bad) if bad else (f"the nav reaches all {len(secs)} section(s); the hero has "
                                                f"{len(acts)} main action(s)")


WEB_CONDITIONS = {"page-parses": _page_parses, "no-js-errors": _no_js_errors, "no-overflow": _no_overflow,
                  "links-well-formed": _links_well_formed, "forms-have-destination": _forms_have_destination,
                  "nav-and-main-action": _nav_and_main_action}
RENDER_CONDITIONS = ("no-js-errors", "no-overflow")     # the render check's own problems already list these


def mock_atom(w):
    """The work atom the mock Adaptation (and code's cheapest-fit proposal) writes for a workflow. A workflow with a
    form part builds the section that holds buttons, so it also gets R-6 as a tap-targets check. A goal that describes
    an action gets one behaviour item: a button that shows a result (mock_section builds both)."""
    goal = w.get("completion_check") or w.get("goal_gap") or w.get("name") or "The workflow's change is live"
    a = {"goal": goal,
         "constraints": ["Other workflows' sections stay unchanged", "The page loads with no console errors"],
         "context": [f"charter item {i}" for i in w.get("covers") or []] + [f"finding {f}" for f in w.get("fixes") or []]
         + ([f"section {w['revises']}"] if w.get("revises") else []) or ["charter goal"],
         "how": {"kind": "AI", "agent": "Builder", "steps": "Write the workflow's section of index.html"},
         "expected_output": [{"text": "The site file index.html", "kind": "file", "target": "index.html"},
                             {"text": "A page whose tags are balanced", "kind": "test", "target": "page-parses"}],
         "verification": [{"text": "The site file index.html exists after the build.", "type": "file-exists",
                           "target": "index.html", "expect": True},
                          {"text": "The page parses with balanced tags.", "type": "named-test",
                           "target": "page-parses", "expect": True}]}
    if "form" in (w.get("parts") or []):
        a["how"]["steps"] += "; the order form goes in a section with class order-form"
        a["verification"].append({"text": f"R-6: every button and field in the order form is at least "
                                          f"{j2_check.TAP_MIN}x{j2_check.TAP_MIN}px on a phone", "type": "tap-targets",
                                  "target": ".order-form", "expect": j2_check.TAP_MIN})
        a["expected_output"].append({"text": "The order form's buttons and fields are big enough to tap on a phone",
                                     "kind": "tap", "target": ".order-form"})
    if ACTION_GOAL.search(str(goal)):
        slug = (re.sub(r"[^a-z0-9]+", "-", str(w.get("name") or "act").lower()).strip("-") or "act")[:24]
        a["how"]["steps"] += f"; the interactive part has a button with id {slug}-try and a result with id {slug}-result"
        a["expected_output"].append({"text": "Using the interactive part shows its result", "kind": "behavior",
                                     "steps": [f"click #{slug}-try", f"expect-visible #{slug}-result"]})
    return a


def with_atoms(wfs, atom=None):
    """Workflows with a work atom each (the mock's, or code's cheapest-fit proposal); atom: the output type's."""
    return [{**w, "atoms": w.get("atoms") or [(atom or mock_atom)(w)]} if isinstance(w, dict) else w for w in wfs or []]


# ---- artifact workflows: a workflow whose output is a file Adaptation names (any path and type), not the site file
# or a section of it. The workflows after it read the file, kept under the run's artifacts/ folder. Code runs its work
# atom's text checks on the file.
FILE_PATH = re.compile(r"(?!/)(?!.*\.\.)[\w./-]+\.\w+")
SITE_FILES = ("index.html", "site/index.html")


def is_artifact(w):
    outs = [str(o).strip() for o in (w or {}).get("outputs") or []]
    return bool(outs) and all(FILE_PATH.fullmatch(o) and o.lower() not in SITE_FILES for o in outs)


def artifact_path(w):
    """Where a file workflow's output is kept: artifacts/<the path Adaptation named>."""
    f = str(w["outputs"][0]).strip()
    return f if f.startswith("artifacts/") else "artifacts/" + f


def key(name):
    s = str(name or "").strip().lstrip("@").strip().lower()
    s = s.split()[0] if s else ""
    return "user" if s in ("requester", "user") else s


def address(text, who):
    """Make the message's @name match the agent that actually speaks next."""
    tag, text = "@Requester" if who == "user" else "@" + NAMES[who], (text or "").strip()
    if tag.lower() in text.lower():
        return text
    return f"{tag} " + re.sub(r"^@\w+[,:]?\s*", "", text)


def next_field(allowed, hint=""):
    return f'"next": "{"|".join("Requester" if a == "user" else NAMES[a] for a in allowed)}"' + (f" ({hint})" if hint else "")


def line(m):
    return f"[{m['id']}] {NAMES.get(m['who'], m['who'])}: {m['text']}"


def clean_questions(raw):
    """[{question, options, budget?}] from whatever shape the model returned. An option may be a string or
    {"text", "budget_increase"}; "budget" (one int per option) is kept only when some option adds budget."""
    out = []
    for x in raw or []:
        if isinstance(x, dict):
            q = str(x.get("question") or x.get("summary") or "").strip()
            raw_opts = [o for o in x.get("options") or [] if str(o.get("text") if isinstance(o, dict) else o).strip()]
            opts = [str(o.get("text") if isinstance(o, dict) else o).strip() for o in raw_opts]
            budget = [max(0, as_int(o.get("budget_increase"))) if isinstance(o, dict) else 0 for o in raw_opts]
        else:
            q, opts, budget = str(x or "").strip(), [], []
        if q:
            out.append({"question": q, "options": opts, **({"budget": budget} if any(budget) else {})})
    return out


def numbered(qs, options_only=False):
    """The question(s) and options as text. One question isn't numbered; options_only (a re-ask) leaves it out."""
    def opt(q, j, o):
        add = q["budget"][j - 1] if q.get("budget") else 0
        return f"\n   Option {j}: {o}" + (f" [adds {add} work credits]" if add else "")
    if options_only:
        return "".join(opt(qs[0], j, o) for j, o in enumerate(qs[0]["options"], 1)).lstrip("\n")
    return "\n".join(("" if len(qs) == 1 else f"{i}. ") + q["question"]
                     + "".join(opt(q, j, o) for j, o in enumerate(q["options"], 1)) for i, q in enumerate(qs, 1))




def option_picked(q, answer):
    """Index of the one option the answer names (by number or by its full text), or None if it isn't clear."""
    nums = {int(n) for n in re.findall(r"(?<!\d)(?<!\d\.)([1-9])(?!\d)(?!\.\d)", str(answer))} & set(range(1, len(q["options"]) + 1))
    texts = {j for j, o in enumerate(q["options"], 1) if o and o.lower() in str(answer).lower()}
    hits = nums | texts
    return hits.pop() - 1 if len(hits) == 1 else None


DECIDE_FMT = ('"needs_decision": bool (true when the Requester must choose to add budget, leave items out or finish '
              'as is), "budget_needed": int|null (work credits needed to finish, your estimate)')
DECIDE_RULE = ("Never offer adding budget, leaving items out or finishing as is in your own question: set needs_decision "
               "instead and code asks the Requester with the exact amounts it can apply.")
DECISION_FMT = ('"decision": "continue"|"drop"|"ship"|"stop"|null (null if the answer doesn\'t say), "add": int|null (work credits the '
                'Requester asked to add, as they typed it), "drop_items": [int] (done-when item numbers they asked to '
                'leave out)')
BUILT_CLAIM = re.compile(r"\b(?:now has|already has|so far the (?:page|site) has|the (?:page|site) has|is built|"
                         r"(?:we|I)(?:'ve| have)? (?:built|added)|has been (?:built|added))\b", re.I)
DECISION_PAUSES = ("limit", "conflict", "ruling_budget")    # the Requester decides; info/fact/question only inform
ONLY_PICK = re.compile(r"\s*(option\s*)?[1-9]\s*[.:)]?\s*", re.I)


REPLY_FMT = ('"to_requester": str (to the Requester only, in plain words: what you understood or decided and what '
             'happens next; at most 3 short sentences; tagged @Requester or no tag; nothing for the team; "" when the '
             'Requester needn\'t hear anything now), "to_agents": str (the instruction for the team, tagged with the '
             'agent; the Requester never sees it)')
AT_REQUESTER = re.compile(r"^\s*@Requester\b[,:]?\s*", re.I)
# Agent content the Requester must never get, and user-facing content agents must never get (checked in call()).
AGENT_CONTENT = re.compile(r"@(?!Requester\b)[A-Za-z]|\b[WFRC]-\d+\b|\bre-?pric|\b\d+\s*units?\b|\b(?:Adaptation|"
                           r"Priority|Coordination|Validator|Audit)\b|\b(?:workflows?|GRANT|COUNTER|REFUSE|envelope|"
                           r"rate[- ]card)\b", re.I)
USER_CONTENT = re.compile(r"@Requester\b|\bthank(?:s| you)\b|\bI(?:'ll| will) start\b|\bwould you like\b|"
                          r"\bdo you want\b|\byour (?:site|website|shop|store|business|brief)\b", re.I)
PLAIN_WORDS = ((r"\bunits\b", "work credits"), (r"\bunit\b", "work credit"), (r"\bbuilding time\b", "work"), (r"\bdone[-_ ]when items?\b", "checklist items"),
               (r"\bdone[-_ ]when\b", "checklist"), (r"\bcycles\b", "rounds"), (r"\bcycle\b", "round"),
               (r"\bcaps?\b", "limit"), (r"\s*\((?:my |a |rough )?estimated?\)", ""), (r"\bestimated\b", "about"),
               (r"\bestimates?\b", "rough figure"))


GAP_KINDS = (("(not built yet)", "not built yet"), ("(not built)", "not built yet"), ("(unfunded)", "unfunded"),
             ("(not checked yet)", "built but not checked yet"), ("(built, not checked yet)", "built but not checked yet"))
STILL = {"not built": "still not built", "unfunded": "still not funded", "built but failing": "still failing",
         "built, not checked yet": "not checked yet", "failing": "still failing"}


def gap_kind(g):
    """How a "Still not right" entry stands, in plain words (the list's header counts these)."""
    g = str(g)
    return next((k for end, k in GAP_KINDS if g.endswith(end)), "built but failing")


JOINING = ("a", "an", "the", "and", "or", "of", "for", "with", "to", "in", "on", "at", "by", "from", "is", "are")


def short(text, n=12, clause=True):
    """A checklist item's short name, cut at a natural break: its first clause (the whole text with clause=False) when
    that is at most n words, else its first n whole words, never ending on a joining word, plus "...". A longer name,
    made to tell two items apart, keeps its last word ("Hall A")."""
    part = (re.split(r"[,:;(]| - ", str(text))[0] if clause else str(text)).split()
    words = part[:n]
    cut = len(part) > n
    while cut and len(words) > 1 and words[-1].lower().strip(".") in JOINING:
        words.pop()             # never end on a joining word ("about the beans and the")
    name = " ".join(words).rstrip(" .,;:-") if words else str(text)[:40]
    return name + "..." if cut else name


def shorts(texts, n=12):
    """Short names that are unique within one list: items that share a short name get longer ones (past the first
    clause if need be) until they differ. Identical texts are the same item and keep the same name."""
    texts = list(texts)
    out, k = [short(t, n) for t in texts], n
    longest = max((len(str(t).split()) for t in texts), default=0)

    def clashes():       # names shared by different texts
        return {o for o in out if len({t for t, x in zip(texts, out) if x == o}) > 1}
    while (c := clashes()) and k <= max(longest, n):
        out = [short(t, k, clause=False) if o in c else o for t, o in zip(texts, out)]
        k += 1
    return out


def and_list(xs):
    xs = list(xs)
    return ", ".join(xs[:-1]) + " and " + xs[-1] if len(xs) > 1 else "".join(xs)


def gap_list(gaps, title=None):
    """The "Still not right" list after Identity's words: one gap per line; the header's counts add up to the list.
    title: another header for a list that isn't about failures (what the Requester asked for)."""
    if not gaps:
        return ""
    if title:
        return f"\n{title} ({len(gaps)}):\n" + "\n".join(f"- {g}" for g in gaps)
    kinds = {}
    for g in gaps:
        k = gap_kind(g)
        kinds[k] = kinds.get(k, 0) + 1
    out = (f"\nStill not right ({len(gaps)}" + (": " + ", ".join(f"{n} {k}" for k, n in kinds.items())
                                                if len(kinds) > 1 else "") + "):\n" + "\n".join(f"- {g}" for g in gaps))
    if "built but not checked yet" in kinds:
        out += "\nSome parts were built after the last check, so they're not checked yet."
    return out


def for_agents(m):
    """Test notes are for humans; to_requester replies are for the Requester. Agents get neither."""
    return not m.get("test") and not m.get("requester_only")


NUMBER_ASK = re.compile(r"\b(?:number|target|metric|how many|how much|percent(?:age)?|rate)\b|%", re.I)
INTERNAL_ID = re.compile(r"\b(?:[FRDCPN]-\d+|[GRPN]\d{1,2})\b")     # findings, rubric, conflicts, conditions, PRD ids


def requester_words(text, keep=""):
    """Text for the Requester without agent-facing wording: no "Alternatives: ..." sentence, no parenthesis that names
    an internal id (G2, R5, F-3, R-6, N8), and no sentence that still names one. keep: the charter and the Requester's
    words (squashed); an id that appears there is their own word (a room "G2") and stays."""
    ids = lambda t: [m for m in INTERNAL_ID.findall(t) if not re.search(rf"\b{m.lower()}\b", keep)]
    t = re.sub(r"\s*\bAlternatives?\s*:[^.?!]*[.?!]?", "", str(text or ""))
    t = re.sub(r"\s*\([^()]*\)", lambda m: "" if ids(m.group(0)) else m.group(0), t)
    return " ".join(x for x in re.split(r"(?<=[.!?;])\s+", t) if not ids(x)).strip()


def to_requester(text, question=True, keep=""):
    """What the Requester reads (without its @Requester tag, which the caller adds): none of the internal words or ids,
    at most 3 short sentences plus one question (none when the question is asked separately below it: question=False).
    keep: see requester_words."""
    t = re.sub(r"(?m)^\s*\d+[.)]\s.*$", "", AT_REQUESTER.sub("", str(text or "")))   # numbered lines: code's job
    t = requester_words(" ".join(t.split()), keep)
    for pat, w in PLAIN_WORDS:
        t = re.sub(pat, w, t, flags=re.I)
    parts = [x for x in re.split(r"(?<=[.!?])\s+", t) if x.strip() and not re.fullmatch(r"\d+[.)]", x.strip())]
    said = [x for x in parts if not x.endswith("?")][:3]
    asked = [x for x in parts if x.endswith("?")][:1] if question else []
    return " ".join(said + asked)


def parse_decision(answer, q, choices, crit, nums=None):
    """The Requester's answer to the decision question, as code reads it: {"decision", "add", "drop_items"}.
    "2. Add 25 more. Leave out the lookbook and reviews" -> continue, add 25, drop the two items by name.
    nums: each item's run-long number (what the Requester sees); drop_items are those numbers, never positions."""
    answer, nums = str(answer or ""), nums or list(range(1, len(crit) + 1))
    i = option_picked(q, answer) if q.get("options") else None
    picked = choices[i] if i is not None and i < len(choices) else None
    if ONLY_PICK.fullmatch(answer):
        return {"decision": picked, "add": None, "drop_items": []}
    amount = re.search(r"\badd\s+(\d+)|(\d+)\s*(?:more|units?|extra|(?:work )?credits?)\b", answer, re.I)  # "add 25"
    add = int(amount.group(1) or amount.group(2)) if amount else None
    drops = {int(n) for n in re.findall(r"\bitems? (\d+)|\bF-(\d+)", answer, re.I) for n in n if n}
    m = re.search(r"(?:leave out|drop|without|skip|remove)\s+(.*)", answer, re.I)
    for part in re.split(r",|\band\b|\.", m.group(1)) if m else []:
        part = re.sub(r"^\s*(the|a|an)\s+", "", part.strip(), flags=re.I).strip().lower()
        drops |= {nums[j] for j, c in enumerate(crit) if len(part) > 2 and part in c.lower()}
    drops = sorted(n for n in drops if n in nums)
    if "stop" in choices and re.search(r"\bstop\b", answer, re.I) and not add:
        decision = "stop"
    elif re.search(r"\bship\b|as is|\bfinish|go live", answer, re.I) and not add:
        decision = "ship"
    elif add or picked == "continue":
        decision = "continue"
    elif drops or m or picked == "drop":
        decision = "drop"
    else:
        decision = picked
    return {"decision": decision, "add": add, "drop_items": drops}


def brief_list(pos):
    """A final position as one line: workflow names and costs, or the text itself."""
    if isinstance(pos, list):
        return ", ".join(f"{w.get('name')} ({w.get('cost')})" for w in pos) or "nothing"
    return str(pos or "")[:300]


# ---- mock mode helpers (scripted functions, no model): a generic charter and page built from the text
BUSINESS_TYPES = (("florist", r"florist|flower|bouquet"), ("photographer", r"photograph"), ("bakery", r"baker|cake"),
                  ("cafe", r"cafe|coffee|espresso"), ("clinic", r"clinic|dental|doctor|hospital"),
                  ("boutique", r"boutique|kurta|saree|dupatta|handloom"))
NEEDS = {"florist": ["Browse bouquets with prices", "Order for delivery or pickup", "Delivery areas and times"],
         "photographer": ["Portfolio of past work", "Packages and prices", "Booking or enquiry form"],
         "bakery": ["Menu with prices", "Order a cake", "Opening hours and address"],
         "cafe": ["Menu with prices", "Opening hours", "Where to find us"],
         "clinic": ["Services offered", "Doctors and timings", "Book an appointment"],
         "boutique": ["Products shown by category", "Filter products by size and price", "Order on WhatsApp"]}
BOUTIQUE_PRODUCTS = (("Kurta, earthy red", 1499), ("Kurta, indigo block print", 1699), ("Saree, handloom cotton", 3499),
                     ("Shirt, natural dye", 1299), ("Dupatta, block print", 899))      # mock only: fake catalog, fake prices


def guess_type(text):
    """Mock only: a business type from the brief's words (live runs ask Identity)."""
    return next((t for t, pat in BUSINESS_TYPES if re.search(pat, text, re.I)), "business")


def mock_charter(text):
    t = text.lower()
    m = re.search(r"\b(?:for|called)\s+([A-Z][\w']*(?:\s+[A-Z][\w']*)*)", text)
    name = m.group(1) if m else "Our cafe"
    bt = guess_type(text)
    if bt == "boutique":
        items = ["Visitors can browse products by category", "Every product shows its id and price",
                 "Visitors can filter products by size and price",
                 "Every product has an order button that opens a WhatsApp link with the product id and price"]
        if re.search(r"\btimings?\b|\bhours?\b|\bopen\b", t):    # only when the brief gives them, never invented
            items.append("The store timings and the shop address are listed")
    else:
        items = [d for k, d in (("menu", "The menu with prices is shown"), ("hour", "Opening hours are listed"),
                                ("where", "The address and a map link are shown"),
                                ("order", "Visitors can pre-order a coffee for pickup with a form"),
                                ("bean", "A section about the beans and the farms they come from"),
                                ("loyalty", "A loyalty-card explainer (buy 9, get the 10th free)"),
                                ("calculat", "A price calculator shows the total for the chosen items")) if k in t]
        if len(items) < 2:                    # not a cafe brief: the business type's usual visitor needs
            items += [f"{n} is on the page" for n in NEEDS.get(bt, ["What the business offers",
                                                                     "How to get in touch"])]
    items.append("Layout works on a phone")
    return {"statement": f"A one-page website for {name} where visitors find what they need and come in",
            "site_name": name, "audience": "Local customers, mostly on phones",
            "standards": ["Warm, modern feel", "Fast first screen on a phone"],
            "rules": ["Plain HTML/CSS/JS in one file", "No real customer data"],
            "done_when": items, "core": [bool(CORE_FACT.search(d)) for d in items], "to_be_confirmed": [],
            "dropped_from_brief": [], "facts": [name] if name in text else [],
            "business_type": guess_type(text), "visitor_needs": NEEDS.get(guess_type(text), [])}


# ---- documents as the blueprint: the LLD's sections, elements and form fields are the website builders' spec
def lld_slice(spec, ids):
    """The part of the LLD (parse_lld data) for the given section ids: those sections and their elements and fields."""
    ids = [str(i) for i in ids or []]
    return {"sections": [s for s in spec.get("sections") or [] if s["id"] in ids],
            "elements": [e for e in spec.get("elements") or [] if e["section"] in ids],
            "fields": [f for f in spec.get("fields") or [] if f["section"] in ids]}


def _has_selector(page, sel):
    sel = str(sel or "")
    if sel.startswith("#"):
        return element_by_id(page, sel[1:]) is not None
    if sel.startswith("."):
        return bool(re.search(rf"""\bclass=["'][^"']*\b{re.escape(sel[1:])}\b""", page or ""))
    return False


def lld_check(page, sl):
    """(ok, evidence) for a page against an LLD slice: every section id, element selector (and the selectors its
    behaviour names) exists, and every form field exists with its type and required flag. Code only, no browser.
    ponytail: structure, not what a click does; the behaviour steps (behavior-steps) are the upgrade."""
    bad = []
    for s in sl.get("sections") or []:
        if not s.get("anchor") or element_by_id(page, s["anchor"]) is None:
            bad.append(f"{s['id']} #{s.get('anchor')} isn't on the page")
    for e in sl.get("elements") or []:
        for sel in [e.get("selector")] + re.findall(r"`([#.][\w-]+)`", e.get("behaviour") or ""):
            if sel and not _has_selector(page, sel):
                bad.append(f"{e['id']} {sel} isn't on the page")
    for f in sl.get("fields") or []:
        scope = element_by_id(page, f["form"]) if f.get("form") else page
        tag = re.search(rf"""<(?:input|select|textarea)\b[^>]*\bname=["']{re.escape(str(f.get('name')))}["'][^>]*>""",
                        scope or "", re.I)
        if not tag:
            bad.append(f"{f['id']} field {f.get('name')} isn't in {'#' + f['form'] if f.get('form') else 'the page'}")
            continue
        ty = re.search(r"""\btype=["']([\w-]+)["']""", tag[0], re.I)
        req = bool(re.search(r"\brequired\b", tag[0], re.I))
        if (ty[1].lower() if ty else "text") != f.get("type"):
            bad.append(f"{f['id']} field {f['name']} is type {ty[1] if ty else 'text'}, the spec says {f.get('type')}")
        if req != bool(f.get("required")):
            bad.append(f"{f['id']} field {f['name']} is {'required' if req else 'optional'}, the spec says "
                       f"{'required' if f.get('required') else 'optional'}")
    n = len(sl.get("sections") or []) + len(sl.get("elements") or []) + len(sl.get("fields") or [])
    return not bad, "; ".join(bad) if bad else f"all {n} spec item(s) are on the page as written"


def lld_markup(sl, title=""):
    """Mock only: the HTML a scripted builder writes for an LLD slice (sections, elements that show their result,
    forms with their fields)."""
    out = ""
    for s in sl.get("sections") or []:
        inner = f"<h2>{html.escape(re.sub(r'[0-9]', '', title))}</h2>"
        for e in [x for x in sl.get("elements") or [] if x["section"] == s["id"]]:
            sel = e.get("selector") or ""
            attr = f'id="{sel[1:]}"' if sel.startswith("#") else f'class="{sel[1:]}"'
            inner += (f'<button {attr} type="button" style="min-height:44px;min-width:44px">Try it</button>')
            for r in re.findall(r"`#([\w-]+)`", e.get("behaviour") or ""):
                inner += f'<p id="{r}" hidden>Result shown here.</p>'
        forms = {}
        for f in [x for x in sl.get("fields") or [] if x["section"] == s["id"]]:
            forms.setdefault(f.get("form"), []).append(f)
        for form, fs in forms.items():
            rows = "".join(f'<label>{html.escape(f["type"])} field <input name="{f["name"]}" type="{f["type"]}"'
                           f'{" required" if f.get("required") else ""} style="display:block;min-height:44px"></label>'
                           for f in fs)
            fid = f'id="{form}" ' if form else ""
            inner += (f'<form {fid}action="/send">{rows}<button type="submit" '
                      'style="min-height:44px;min-width:44px">Send</button></form>')
        out += f'<section id="{s["anchor"]}">{inner}</section>'
    return out


def mock_section(w, goal, sections, lld=None, anchors=()):
    """mock_page's (head, body) plus the elements the workflow's behaviour steps use (a button for a click, a
    paragraph for anything read)."""
    head, body = mock_page(w, goal, sections, anchors)
    body += lld_markup(lld, w["name"]) if lld else ""
    made, extra = set(), ""
    for a in w.get("atoms") or []:
        for x in a.get("expected_output") or []:
            for st in x.get("steps") or [] if isinstance(x, dict) and x.get("kind") == "behavior" else []:
                p = j2_check.parse_step(st)
                if p["sel"] and p["sel"].startswith("#") and p["sel"] not in made:
                    made.add(p["sel"])
                    extra += (f'<button id="{p["sel"][1:]}" type="button" style="min-height:44px;min-width:44px">Try '
                              "it</button>" if p["verb"] == "click" else f'<p id="{p["sel"][1:]}">Result shown here.</p>')
    return head, body + (f'<div class="j2-try">{extra}</div>' if extra else "")


def mock_page(w, goal, sections, anchors=()):
    """(head, body) a scripted workflow agent writes: the scaffold styles the page, a viewport fix (or a workflow that
    builds a phone item) adds the meta tag,
    anything else adds the done-when items no other section covers yet (or a plain block, not a section). The first
    section built (the scaffold, or any section while no other has a nav) also holds the hero: the nav to every
    item's section and one main action. A later section the hero's nav doesn't reach (an item added by a change
    request) links its own items from a small nav of its own."""
    g, n = goal or {}, w["name"].lower()
    others = "".join(x["body"] for sid, x in sections.items() if sid != w["section"])
    orig = g.get("charter_items") or g.get("done_when") or []      # charter numbers: stable after drops
    num = lambda i, d: orig.index(d) + 1 if d in orig else i
    shown = [(num(i, d), d) for i, d in enumerate(g.get("done_when") or [], 1) if "phone" not in str(d).lower()]
    main = next(((i, d) for i, d in shown if item_kind(d) == "form"), shown[0] if shown else None)
    link = lambda i, d, c="": (f'<a{c} href="#item-{i}" style="display:inline-block;min-height:44px;min-width:44px;'
                               f'line-height:44px;padding:0 8px">{html.escape(short(d))}</a>')     # 44px tap targets
    hero = "" if "<nav" in others else (
        '<header id="hero"><nav>' + " ".join(link(i, d) for i, d in shown)
        + "".join(f' <a href="#{a}" style="display:inline-block;min-height:44px;min-width:44px;line-height:44px;padding:0 8px">More</a>'
                  for a in anchors)
        + f"</nav><h1>{html.escape(g.get('site_name') or 'Our site')}</h1><p>{html.escape(g.get('statement') or '')}</p>"
        + (link(*main, ' class="j2-btn"') if main else "") + "</header>")
    if "scaffold" in n:
        return ("<style>body{font:16px system-ui;margin:0 auto;max-width:900px;padding:24px;background:#fbf7f2;"
                "color:#2b211b}section{padding:16px 0;border-top:1px solid #e6dccf}</style>", hero)
    if "viewport" in n or any("phone" in str(t).lower() for t in w.get("covers_text") or []):
        return '<meta name="viewport" content="width=device-width, initial-scale=1">', ""
    if re.search(r"timings?|hours|address", n):
        # The timings/address section (and a later revise of it) writes only its own items, so the page records
        # where the timings are and a board answer can name that section instead of saying "not recorded".
        shown = [(i, d) for i, d in shown if re.search(r"timings?|hours|address|where to find", str(d), re.I)] or shown
    if "product" in n:        # fake catalog, id and price recorded as data-id/data-price so a board answer can name them
        idx, d = next(((i, d) for i, d in shown if "price" in str(d).lower()), shown[0] if shown else (1, ""))
        cards = "".join(f'<article class="j2-card" data-id="p{k}" data-price="{p}" style="max-width:100%;'
                        f'box-sizing:border-box"><h3>{html.escape(name)}</h3><p>Rs {p}</p></article>'
                        for k, (name, p) in enumerate(BOUTIQUE_PRODUCTS, 1))
        return "", (hero + f'<section id="item-{idx}" data-item="{idx}"><h2>{html.escape(str(d))}</h2>'
                   f'<div class="j2-products" style="max-width:100%">{cards}</div></section>')

    def item(i, d):
        if item_kind(d) == "form":
            return (f'<section class="order-form" id="item-{i}" data-item="{i}"><h2>{html.escape(str(d))}</h2>'
                    '<form action="/orders">'
                    '<label>Name <input name="name" style="display:block;min-height:44px;width:100%"></label>'
                    '<button class="j2-btn" type="submit" style="min-height:44px;min-width:44px">Send</button>'
                    "</form></section>")
        return (f'<section id="item-{i}" data-item="{i}"><h2>{html.escape(str(d))}</h2><p>Sample content (replace '
                "before launch).</p></section>")
    new = [(i, d) for i, d in shown if f'data-item="{i}"' not in others]
    more = [(i, d) for i, d in new if not hero and f'href="#item-{i}"' not in others]
    items = "".join(item(i, d) for i, d in new)
    return "", hero + ("<nav>" + " ".join(link(i, d) for i, d in more)
                       + "</nav>" if more else "") + (items or '<div class="j2-extra"><p>Sample content.</p></div>')


PRD_SECTIONS = ("Problem", "Users", "Goals", "Requirements", "Success metrics", "Out of scope")


def mock_prd_charter(text):
    """Mock only: a PRD charter from a brief that labels its facts ("Users: ...", "Problem: ...")."""
    m = re.search(r"\b(?:for|called)\s+([A-Z][\w']*(?:\s+[A-Z][\w']*)*)", text)
    name = m.group(1) if m else "The product"

    def field(label, default=""):
        x = re.search(rf"(?:{label})\s*:\s*(.+?)(?:\.\s|\.$|$)", text, re.I | re.S)
        return " ".join(x.group(1).split()) if x else default
    features = [f.strip() for f in re.split(r",|;|\band\b", field(r"must-have features|features")) if f.strip()]
    success = [f.strip() for f in field(r"success").split(";") if f.strip()]
    users = field("users")
    if not features:
        # A brief that labels nothing: the mock reads the brief's own sentences about what the page and the form do,
        # and its "must" sentences as the measures. Every word is the brief's; nothing is invented.
        sents = [" ".join(s.split()).split(":")[0].strip(" .") for s in re.split(r"(?<=\.)\s+|\n+", text)]
        features = [s for s in sents if re.match(r"The (?:page|form|weekly|class)\b", s)]
        success = success or [s[:1].lower() + s[1:] for s in sents if " must " in f" {s} "][:2]
        users = users or ("visitors" if features else "")
    groups = user_groups({"users": users})     # mock: the first group does the features, the rest say nothing
    items = [f"The Problem section states the problem {name} solves", "The Users section says who the users are",
             "The Goals section lists the goals as G1, G2, ...",
             "The Requirements section lists every must-have feature as R1, R2, ...",
             "The Success metrics section gives a measurable metric for each goal",
             "The Out of scope section says what the product won't do"]
    return {"statement": f"A PRD for {name}", "product_name": name, "site_name": name,
            "users": users or "Its users", "audience": users or "Its users",
            "user_groups": [{"name": x, "actions": features if i == 0 else []} for i, x in enumerate(groups)],
            "problem": field("problem", "The problem the brief describes"), "features": features,
            "constraints": [c.strip() for c in field("constraints").split(";") if c.strip()], "success_measures": success,
            "facts": [name] + features, "standards": ["Every requirement testable"], "rules": ["No invented facts"],
            "done_when": items, "core": [True, False, False, True, True, False], "to_be_confirmed": [],
            "dropped_from_brief": [], "business_type": "prd", "visitor_needs": []}


def mock_doc_section(w, goal):
    """Mock only: the markdown a scripted builder writes for the PRD section its workflow names."""
    g = goal or {}
    sec = next((s for s in PRD_SECTIONS if s.lower() in str(w.get("name")).lower()), PRD_SECTIONS[0])
    name, feats, succ = g.get("product_name") or "The product", g.get("features") or [], g.get("success_measures") or []
    groups = [x for x in g.get("user_groups") or [] if isinstance(x, dict)]
    who = groups[0]["name"] if groups and groups[0].get("actions") else "Users"
    idle = " ".join(f"{x['name'][:1].upper() + x['name'][1:]} {NO_ACTIONS}." for x in groups if not x.get("actions"))
    body = {"Problem": f"{g.get('problem') or ''}. {name} exists to solve this.",
            "Users": f"The users of {name} are {g.get('users') or 'its users'}." + (f"\n\n{idle}" if idle else ""),
            "Goals": "\n".join(f"- **G{i}**: {s}" for i, s in enumerate(succ, 1)) or f"- **G1**: {name} solves the problem",
            # A labelled brief's features are verb phrases ("see the menu"); a brief the mock read whole gives whole
            # sentences, which are written as they stand with the group named after them.
            "Requirements": "\n".join(f"{i}. **R{i}**: " + (f"{f} ({who})" if f[:1].isupper() else
                                                            f"{who[:1].upper() + who[1:]} can {f}")
                                      for i, f in enumerate(feats, 1)),
            "Success metrics": "| Goal | Metric |\n|---|---|\n" + "\n".join(f"| G{i} | {s} |" for i, s in enumerate(succ, 1)),
            "Out of scope": "\n".join(f"- Anything outside this constraint: {c}" for c in g.get("constraints") or [])
                            or f"- Features beyond the listed requirements of {name}"}[sec]
    return f"## {sec}\n\n{body}\n"


def mock_doc_atom(w):
    """The work atom the mock Adaptation writes for a PRD section workflow: the heading is there, the section isn't
    empty, and the Requirements ids are unique."""
    sec = next((s for s in PRD_SECTIONS if s.lower() in str(w.get("name")).lower()), PRD_SECTIONS[0])
    checks = [{"text": f"The document has the heading {sec}.", "type": "heading-present", "target": sec, "expect": True},
              {"text": f"The {sec} section isn't empty (at least 5 words).", "type": "section-not-empty", "target": sec,
               "expect": 5}]
    if sec == "Requirements":
        checks.append({"text": "Requirement ids R1, R2, ... are unique.", "type": "ids-unique", "target": "R", "expect": True})
    return {"goal": f"The PRD's {sec} section is written from the charter", "constraints": ["Markdown only, never HTML",
            "Other sections stay unchanged", "No invented facts"],
            "context": [f"charter item {i}" for i in w.get("covers") or []] + [f"finding {f}" for f in w.get("fixes") or []]
            or ["charter goal"],
            "how": {"kind": "AI", "agent": "Builder", "steps": f"Write the \"## {sec}\" section of the PRD in markdown"},
            "expected_output": [{"text": f"The {sec} heading", "kind": "heading", "target": sec}]
            + ([{"text": "Requirement ids R1, R2, ... with no repeats", "kind": "ids", "target": "R"}]
               if sec == "Requirements" else []),
            "verification": checks}


# ---- design documents (opt-in): PRD, then HLD, then LLD, built as a chain; the HLD and LLD reuse the PRD path
# (a document output type: markdown sections, checks by code, no browser, exported to .docx).
DOC_SET = ("prd", "hld", "lld")
_DOC_SET_WORDS = re.compile(r"\bdesign\s+(?:docs?|documents?|documentation)\b|\bfull\s+documentation\b", re.I)
_DOC_WORDS = {"prd": re.compile(r"\bPRDs?\b|\bproduct\s+requirements?\b", re.I),
              "hld": re.compile(r"\bHLDs?\b|\bhigh[- ]level\s+design\b|\barchitecture\s+overview\b", re.I),
              "lld": re.compile(r"\bLLDs?\b|\blow[- ]level\s+design\b|\bdetailed\s+design\b", re.I)}
_SITE_WORDS = re.compile(r"\b(?:web\s?)?sites?\b|\bweb\s?pages?\b|\blanding\s+page\b", re.I)


def detect_documents(*texts):
    """The documents the Requester explicitly asked for, in chain order (PRD, HLD, LLD); [] when none: never a
    default. A set request ("design docs", "full documentation") means all three. Other uses of the word design
    ("modern design", "design a logo") match nothing."""
    text = " ".join(str(t or "") for t in texts)
    if _DOC_SET_WORDS.search(text):
        return list(DOC_SET)
    return [k for k in DOC_SET if _DOC_WORDS[k].search(text)]


def wants_site(*texts):
    """A document request that also names a website: the documents are its blueprint and the site is built after."""
    return bool(_SITE_WORDS.search(" ".join(str(t or "") for t in texts)))


HLD_SECTIONS = ("Purpose and scope", "Pages", "Components", "Data flow", "Hosting and deployment", "Risks")
LLD_SECTIONS = ("Pages and sections", "Interactive elements", "Form fields", "File structure",
                "Responsive and accessibility rules")
ALL_DOC_SECTIONS = PRD_SECTIONS + HLD_SECTIONS + LLD_SECTIONS

HLD_RUBRIC = (("D-1", "Every component has one clear responsibility"),
              ("D-2", "The data flow names every form and integration the pages need"),
              ("D-3", "Hosting, deployment and risks are explicit"),
              ("D-4", "Consistent with the PRD: every requirement has a component"),
              ("D-5", "No invented facts: every fact comes from the brief or the Requester's answers or the PRD, or is "
                      "labelled \"Assumption:\"; the evidence quotes each unlabelled claim not in them"))
LLD_RUBRIC = (("D-1", "Every page section has an id and implements at least one component"),
              ("D-2", "Every interactive element has a selector, a trigger and a behaviour"),
              ("D-3", "Every form field has a type, a required flag and a validation rule"),
              ("D-4", "Consistent with the HLD: every component is built by a section"),
              ("D-5", "No invented facts: every fact comes from the brief or the Requester's answers or the earlier "
                      "documents, or is labelled \"Assumption:\"; the evidence quotes each unlabelled claim"))

HLD_BUILDER = """# Workflow agent for an HLD section (one per registered workflow, e.g. "W-7 Builder")

- You exist for one registered workflow: one section of the high-level design. Write that section only, in markdown,
  starting with its "## <Section>" heading. Never HTML. Code assembles the HLD from every section, in section order.
- Read the PRD's agent_output in the charter (its requirements R1, R2, ...), not any other text. Use only facts from
  the charter and the Requester's answers; copy names and numbers exactly.
- Components are a list, one per line: "- **C1** <Name>: <one-sentence responsibility>. Covers R1, R2." Every PRD
  requirement is covered by at least one component. Never start another line with a component id.
- Data flow names every form and integration. Hosting and deployment says where the site lives and how it ships.
- Any statement beyond the charter and the answers starts with "Assumption:" on its own line or list item. Code
  collects these lines into "Assumptions to confirm" at the end.
- If your reply can't be used, code tells you why and you retry, within your envelope. Then report to
  @Coordination what you changed."""
LLD_BUILDER = """# Workflow agent for an LLD section (one per registered workflow, e.g. "W-9 Builder")

- You exist for one registered workflow: one section of the low-level design. Write that section only, in markdown,
  starting with its "## <Section>" heading. Never HTML. Code assembles the LLD from every section, in section order.
- Read the HLD's agent_output in the charter (its components C1, C2, ...), not any other text. Use only facts from the
  charter and the Requester's answers.
- Pages and sections: one line each: "- **S1** `#section-id` (page index.html): <what it shows>. Implements C1, C2."
- Interactive elements: "- **E1** in S1: `#selector` trigger <click|change|submit|...>; behaviour <what happens>."
- Form fields: "- **F1** in S1, form `#form-id`: field `name`, type <text|email|...>, required <yes|no>, validation <rule>."
- File structure, responsive and accessibility rules are short lists. Every id is unique. The page builders will
  implement exactly this, so write only what must exist.
- Any statement beyond the charter and the answers starts with "Assumption:" on its own line or list item.
- If your reply can't be used, code tells you why and you retry, within your envelope. Then report to
  @Coordination what you changed."""
HLD_PLAN = ("\nTHE OUTPUT IS AN HLD DOCUMENT (markdown), not a website. Plan exactly one workflow per HLD section, in "
            f"this order: {', '.join(HLD_SECTIONS)}; each has one work atom whose how.steps names the section heading it "
            "writes (e.g. \"## Components\"). covers lists the done-when items about that section. The Components "
            "workflow must name every requirement id it covers (R1, R2, ...). Price them with the document rate card: "
            "doc_section, or requirements_section for Components; fixes with 'fix'. Components are C1, C2, ....")
LLD_PLAN = ("\nTHE OUTPUT IS AN LLD DOCUMENT (markdown), not a website. Plan exactly one workflow per LLD section, in "
            f"this order: {', '.join(LLD_SECTIONS)}; each has one work atom whose how.steps names the section heading it "
            "writes (e.g. \"## Form fields\"). covers lists the done-when items about that section. The Pages and "
            "sections workflow must name every component id it implements (C1, C2, ...). Price them with the document "
            "rate card: doc_section, or requirements_section for Pages and sections; fixes with 'fix'. Sections are S1, "
            "S2, ...; elements E1, ...; fields F1, ....")


def _md_list(doc, section):
    return [ln for ln in strip_md(md_section(doc, section) or "").splitlines() if ln.strip()]


def parse_hld(doc):
    """The HLD's components: [{"id", "name", "responsibility", "covers": [R ids]}], read from the Components section."""
    out = []
    for ln in _md_list(doc, "Components"):
        m = re.match(r"\s*(?:[-*+]\s+)?(C\d+)\b\s*[:.-]?\s*(.*)$", ln)
        if m:
            name, _, rest = m[2].partition(":")
            out.append({"id": m[1], "name": name.strip() if rest else "",
                        "responsibility": re.split(r"\bCovers\b", rest or name)[0].strip(" ."),
                        "covers": re.findall(r"\bR\d+\b", ln)})
    return out


def parse_lld(doc):
    """The LLD as data: {"sections": [{id, anchor, page, implements}], "elements": [{id, section, selector, trigger,
    behaviour}], "fields": [{id, section, form, name, type, required, validation}]}. A missing part stays None so the
    L-checks can name it."""
    raw = lambda s: [ln for ln in (md_section(doc, s) or "").splitlines() if ln.strip()]
    tick = lambda ln: re.findall(r"`([^`]+)`", ln)
    secs, els, flds = [], [], []
    for ln in raw("Pages and sections"):
        if m := re.match(r"\s*(?:[-*+]\s+)?(?:\*\*)?(S\d+)(?:\*\*)?\b(.*)$", ln):
            t = tick(m[2])
            page = re.search(r"page\s+`?([\w./-]+)`?", m[2])
            secs.append({"id": m[1], "anchor": next((x[1:] for x in t if x.startswith("#")), None),
                         "page": page[1] if page else None, "implements": list(dict.fromkeys(re.findall(r"\bC\d+\b", m[2])))})
    for ln in raw("Interactive elements"):
        if m := re.match(r"\s*(?:[-*+]\s+)?(?:\*\*)?(E\d+)(?:\*\*)?\s+in\s+(S\d+)\s*:?\s*(.*)$", ln):
            t = tick(m[3])
            tr, bh = re.search(r"trigger\s+([^;]+)", m[3]), re.search(r"behaviou?r\s+(.+)", m[3])
            els.append({"id": m[1], "section": m[2], "selector": t[0] if t else None,
                        "trigger": tr[1].strip() if tr else None, "behaviour": bh[1].strip(" .") if bh else None})
    for ln in raw("Form fields"):
        if m := re.match(r"\s*(?:[-*+]\s+)?(?:\*\*)?(F\d+)(?:\*\*)?\s+in\s+(S\d+)\s*,?\s*(.*)$", ln):
            t = tick(m[3])
            fld = re.search(r"field\s+`([^`]+)`", m[3])
            ty, rq, va = (re.search(r"type\s+(\w+)", m[3]), re.search(r"required\s+(yes|no|true|false)\b", m[3], re.I),
                          re.search(r"validation\s+(.+)", m[3]))
            flds.append({"id": m[1], "section": m[2], "form": next((x[1:] for x in t if x.startswith("#")), None),
                         "name": fld[1] if fld else None, "type": ty[1].lower() if ty else None,
                         "required": (rq[1].lower() in ("yes", "true")) if rq else None,
                         "validation": va[1].strip(" .") if va else None})
    return {"sections": secs, "elements": els, "fields": flds}


def doc_agent_output(kind, doc):
    """The document-level agent_output the next document reads instead of the whole text (built by code)."""
    if kind == "prd":
        body = md_section(doc, "Requirements") or ""
        return {"kind": "prd", "requirements": [{"id": f"R{m[1]}", "text": m[2].strip()} for m in
                                               re.finditer(r"(?m)^\s*(?:[-*+]\s+|\d+[.)]\s+)?(?:\*\*)?R(\d+)(?:\*\*)?\s*[:.-]?\s*(.+)$",
                                                           strip_md(body))]}
    if kind == "hld":
        return {"kind": "hld", "components": parse_hld(doc)}
    return {"kind": "lld", **parse_lld(doc)}


def _no_dups(items):
    return [x for i, x in enumerate(items) if x in items[:i]]


def lld_condition(k, doc):
    """(met, evidence) for the LLD's own checks (L2 ids, L3 elements, L4 fields)."""
    p = parse_lld(doc)
    if k == "lld-ids-unique":
        ids = [x["id"] for x in p["sections"] + p["elements"] + p["fields"]]
        sel = [x["anchor"] for x in p["sections"] if x["anchor"]] + [x["selector"] for x in p["elements"] if x["selector"]]
        dup = _no_dups(ids) + _no_dups(sel)
        bad = [x["id"] for x in p["sections"] if not x["anchor"]]
        return bool(p["sections"]) and not dup and not bad, (
            "no S1, S2, ... section entries" if not p["sections"] else
            f"repeated: {', '.join(dict.fromkeys(dup))}" if dup else f"section(s) without an id: {', '.join(bad)}" if bad else
            f"{len(ids)} entry ids and {len(sel)} selectors, all unique")
    if k == "lld-elements":
        bad = [f"{x['id']} (no {', '.join(f for f in ('selector', 'trigger', 'behaviour') if not x[f])})"
               for x in p["elements"] if not (x["selector"] and x["trigger"] and x["behaviour"])]
        return not bad, (f"incomplete: {'; '.join(bad)}" if bad else
                         f"each of the {len(p['elements'])} interactive element(s) has a selector, a trigger and a behaviour")
    bad = [f"{x['id']} (no {', '.join(f for f, v in (('name', x['name']), ('type', x['type']), ('required flag', x['required'])) if v is None)})"
           for x in p["fields"] if x["name"] is None or x["type"] is None or x["required"] is None]
    return not bad, (f"incomplete: {'; '.join(bad)}" if bad else
                     f"each of the {len(p['fields'])} form field(s) has a name, a type and a required flag")


def mock_hld_section(w, goal):
    """Mock only: the markdown a scripted builder writes for the HLD section its workflow names, from the charter."""
    g = goal or {}
    sec = next((s for s in HLD_SECTIONS if s.lower() in str(w.get("name")).lower()), HLD_SECTIONS[0])
    name, reqs = g.get("product_name") or "The product", ((g.get("chain") or {}).get("prd") or {}).get("requirements") or []
    reqs = reqs or [{"text": f} for f in g.get("features") or []]      # no PRD asked for: the charter's features
    feat = lambda r: (r["text"].split(" can ")[-1] if " can " in r["text"] else r["text"]).strip(" .").lower()
    body = {"Purpose and scope": f"{name} solves this: {g.get('problem') or 'the problem in the brief'}. This design covers "
                                 "every requirement in the PRD and nothing else.",
            "Pages": "- The product is one page, index.html, with one section for each component below.",
            "Components": "\n".join(f"- **C{i}** {feat(r)}: builds and shows {feat(r)}." + (f" Covers {r['id']}." if r.get("id") else "")
                                    for i, r in enumerate(reqs, 1)) or "- **C1** the page: shows the product.",
            "Data flow": "- Forms send what visitors type to the site owner by email; nothing is stored in the page.\n"
                         "- No outside service is called when the page loads.",
            "Hosting and deployment": "The page is a single static file. It is uploaded to any static host and replaced "
                                      "on every release.",
            "Risks": "- A form may send nothing if the email address is wrong; the owner checks the first message.\n"
                     "- A section may not fit a small phone screen; each section is checked at phone width."}[sec]
    return f"## {sec}\n\n{body}\n"


def mock_lld_section(w, goal):
    """Mock only: the markdown for the LLD section a workflow names, one page section per HLD component."""
    g = goal or {}
    sec = next((s for s in LLD_SECTIONS if s.lower() in str(w.get("name")).lower()), LLD_SECTIONS[0])
    comps = ((g.get("chain") or {}).get("hld") or {}).get("components") or [{"id": "C1"}]
    name = g.get("product_name") or "The product"
    n = len(comps)
    body = {"Pages and sections": "\n".join(f"- **S{i}** `#sec-{i}` (page index.html): shows {c['id']}. Implements {c['id']}."
                                            for i, c in enumerate(comps, 1)),
            "Interactive elements": "\n".join(f"- **E{i}** in S{i}: `#act-{i}` trigger click; behaviour shows the result in "
                                              f"`#res-{i}`." for i in range(1, n + 1)),
            "Form fields": "\n".join(f"- **F{i}** in S{i}, form `#form-{i}`: field `name-{i}`, type text, required yes, "
                                     "validation not empty." for i in range(1, n + 1, 2)),
            "File structure": f"- `index.html`: the single page of {name}, with every section in the order above.",
            "Responsive and accessibility rules": "- Every button and field is big enough to tap on a phone.\n"
                                                 "- Every field has a visible label."}[sec]
    return f"## {sec}\n\n{body}\n"


def doc_atom_for(secs, ids_in):
    """mock_doc_atom's twin for any document: heading and not-empty checks, plus unique ids where the document has them."""
    def make(w):
        sec = next((s for s in secs if s.lower() in str(w.get("name")).lower()), secs[0])
        checks = [{"text": f"The document has the heading {sec}.", "type": "heading-present", "target": sec, "expect": True},
                  {"text": f"The {sec} section isn't empty (at least 5 words).", "type": "section-not-empty", "target": sec,
                   "expect": 5}]
        exp = [{"text": f"The {sec} heading", "kind": "heading", "target": sec}]
        if ids_in.get(sec):
            checks.append({"text": f"Ids {ids_in[sec]}1, {ids_in[sec]}2, ... are unique.", "type": "ids-unique",
                           "target": ids_in[sec], "expect": True})
            exp.append({"text": f"Ids {ids_in[sec]}1, {ids_in[sec]}2, ... with no repeats", "kind": "ids", "target": ids_in[sec]})
        return {"goal": f"The {sec} section is written from the charter",
                "constraints": ["Markdown only, never HTML", "Other sections stay unchanged", "No invented facts"],
                "context": [f"charter item {i}" for i in w.get("covers") or []] + [f"finding {f}" for f in w.get("fixes") or []]
                or ["charter goal"],
                "how": {"kind": "AI", "agent": "Builder", "steps": f"Write the \"## {sec}\" section in markdown"},
                "expected_output": exp, "verification": checks}
    return make


def export_hld(doc, run_dir):
    return export_docx(doc, os.path.join(run_dir, "hld.docx"))


def export_lld(doc, run_dir):
    return export_docx(doc, os.path.join(run_dir, "lld.docx"))


def stage_goal(kind, base, chain):
    """Code's charter for the HLD or LLD: the earlier charter's facts, one checkable item per section of the new
    document, and the earlier documents' agent_output under "chain" (never their text)."""
    name = base.get("product_name") or base.get("site_name") or "The product"
    reqs = [r["id"] for r in (chain.get("prd") or {}).get("requirements") or []]
    comps = [c["id"] for c in (chain.get("hld") or {}).get("components") or []]
    items = []
    for s in OUTPUTS[kind]["sections"]:
        extra = (f", and names every requirement ({', '.join(reqs)}) it covers" if s == "Components" and reqs else
                 f", and names every component ({', '.join(comps)}) it implements" if s == "Pages and sections" and comps
                 else "")
        items.append(f"The {s} section is written for {name}{extra}")
    keep = {k: base[k] for k in ("product_name", "users", "audience", "problem", "features", "user_groups", "constraints",
                                 "success_measures") if k in base}
    return {**keep, "statement": f"An {kind.upper()} for {name}", "site_name": name, "product_name": name,
            "business_type": kind, "facts": [name], "standards": ["Every item traceable to the PRD"],
            "rules": ["No invented facts"], "done_when": items, "core": [True] * len(items), "to_be_confirmed": [],
            "dropped_from_brief": [], "visitor_needs": [], "chain": {**chain, "prd_ids": reqs, "hld_ids": comps}}


# ---- what the department builds: one registry entry per output type. The flow, agents, budget, rulings, receipts
# and the 3-attempt rule are shared; an entry says what is built and how it is checked. A document type (sections
# set) is markdown assembled in section order, checked by code with no browser, and exported at the end.
DOC_RUBRIC = (("D-1", "Every requirement is specific and testable"),
              ("D-2", "Goals and success metrics are measurable"),
              ("D-3", "Users and problem are consistent with the requirements"),
              ("D-4", "Out of scope is explicit"),
              ("D-5", "No invented facts: every fact comes from the brief or the Requester's answers, or is labelled "
                      "\"Assumption:\"; the evidence quotes each unlabelled claim not in the brief"))
PRD_BUILDER = """# Workflow agent for a PRD section (one per registered workflow, e.g. "W-1 Builder")

- You exist for one registered workflow: one section of the PRD. Write that section only, in markdown, starting
  with its "## <Section>" heading. Never HTML. Code assembles the PRD from every section, in section order.
- Use only facts from the charter and the Requester's answers; copy names and numbers exactly. Never write
  "to be confirmed": a missing fact is Identity's question to the Requester.
- Requirements are a numbered list, one per line: "1. **R1**: ...". Goals are "- **G1**: ...". Every success
  metric names the goal it measures (a table row "| G1 | metric |" or a line "- **G1**: ..."). Mention an id
  inside a sentence anywhere else, never at the start of a line.
- Each requirement is specific and testable; goals and metrics are measurable; out of scope is explicit.
- Every user group the charter names either has requirements that name it, or, when the charter records no actions
  for it, this line in the Users section: "<Group> have no actions in this version." Never invent what a group does.
- A numeric success metric states its measurement window ("per week", "within 3 months"); if the charter gives
  none, write "measurement window not set".
- Any statement that goes beyond the charter and the Requester's answers starts with "Assumption:" on its own line
  or list item. Nothing else may add a claim. Code collects these lines into "Assumptions to confirm" at the end.
- Never write an "Assumptions to confirm" heading or list yourself, and never repeat a claim without the label: every
  line that carries an assumption starts with "Assumption:", wherever it is.
- Never start a sentence with a capitalised word followed by "PRD" or another capitalised word that is not in the
  charter ("This PRD ..."): the invented-names check reads it as a made-up name. Write "This document ...".
- If your reply can't be used, code tells you why and you retry, within your envelope. Then report to
  @Coordination what you changed."""
# A documents run that ends with the website: its details are asked once, before the first document (the answers go into the
# charter every stage uses), and the website stage asks again only for what is genuinely missing.
SITE_FIRST = ("\nA website will be built from these documents. In this same round, also ask the Requester (one question) "
              "for the content details only they know that the website and its forms need and the brief lacks: for example "
              "menu items and prices, opening hours, the address, the phone number and email, and the fields a form should "
              "ask for. Ask for them once, now: they go into the charter every document and the website use. Never invent "
              "them. If the brief already gives them, ask nothing about them.")
SITE_LATER = ("\nThe documents are written and the Requester was already asked for the website's details before them (see "
              "EARLIER QUESTIONS AND ANSWERS). Ask again only for what is genuinely missing; never repeat a question.")
PRD_INTAKE = ("J2.a for a PRD: is there enough information to write it? You need these product facts: the product "
              "name, who the users are, the problem, the must-have features, constraints, and how success is measured. "
              "Ask the Requester only for the ones the brief lacks (set enough=false, 1-4 specific questions); ask about "
              "nothing else, and never invent a fact. Enough: set enough=true and record the charter from the brief and "
              "the answers only. Name every user group in user_groups. For each group whose actions the brief doesn't give, ask "
              "the Requester what that group does in the product; if they don't say, record its actions as []. When a "
              "success measure is a number, ask for its measurement window (\"over what time?\") in the same question "
              "as the number, unless the brief gives one. done_when is the checklist for the PRD: checkable items for its sections (Problem, "
              "Users, Goals, Requirements, Success metrics, Out of scope), each checkable by reading the document; mark "
              "core=true for the problem, the requirements and the success metrics. facts: the exact strings from the "
              "brief or the answers that the PRD must contain (product name, numbers, feature names). Done-when may not "
              "refer to the Requester, Audit, verification, approval or launch. Keep everything the brief asks for; if "
              "the charter leaves any of it out, list it in dropped_from_brief with the reason. Don't plan costs.")
PRD_CHARTER_FMT = ('{"statement": str, "product_name": str, "users": str, "problem": str, "features": [str] (the '
                   'must-have features), "user_groups": [{"name": str, "actions": [str] (what the group does, from the '
                   'brief or the answers; [] if the Requester didn\'t say)}], "constraints": [str], "success_measures": [str], "facts": [str] (exact strings '
                   'from the brief or the answers the PRD must contain), "standards": [str], "rules": [str], '
                   '"done_when": [str], "core": [bool] (one per done-when item, same order), '
                   '"to_be_confirmed": [{"item": int (done-when item number), "fact": str}], "dropped_from_brief": [str]}')
PRD_PLAN = ("\nTHE OUTPUT IS A PRD DOCUMENT (markdown), not a website. Plan exactly one workflow per PRD section, in this "
            f"order: {', '.join(PRD_SECTIONS)}; each has one work atom whose how.steps names the section heading it "
            "writes (e.g. \"## Problem\"). covers lists the done-when items about that section. Price them with the "
            "document rate card: doc_section, or requirements_section for Requirements; fixes with 'fix'. Requirements "
            "are numbered R1, R2, ...; goals G1, G2, ...; each success metric names the goal it measures.")
PRD_ATOM_RULES = ("\nWORK ATOMS: every workflow, fixes included, carries atoms: exactly one work atom for now, id \"a1\" "
                  "(code renames it W-n.a1 at registration). goal: what must be true when it is done, one sentence. "
                  "constraints: short strings, what must never break. context: references only, charter item numbers "
                  "and section ids, never copied text. how: {kind: \"AI\" or \"code\", agent: who builds it, steps: which "
                  "\"## <Section>\" it writes}. expected_output: written now, before any build: a short list of the "
                  "concrete things that must exist in the PRD when the atom is done, each {text: the thing in plain "
                  "words, kind, target}: heading (target: the section heading text, no # characters), ids (target: the "
                  "id prefix, \"R\" for requirement ids R1, R2, ...), facts (target: a list of strings copied exactly "
                  "from the charter or the Requester's answers). Code turns each item into a check and reports it met "
                  "or not met; an item no check can verify, or an empty list, is rejected. "
                  "verification: the checks code runs on the PRD after the builder writes it "
                  "(the builder gets 3 attempts), each {text: the check in plain words, type, target, expect}. Only these "
                  "types: heading-present (target: the section heading text only, with no # characters, e.g. \"Problem\"; "
                  "expect true); section-not-empty (target: the heading text only, no # characters; expect: the minimum "
                  "number of words); ids-unique (target: an id prefix, \"R\" for R1, R2, "
                  "...; expect true); facts-present (target: a list of strings that must all appear in the document; "
                  "expect true). No browser checks. Each facts-present string is copied exactly from the charter or the "
                  "Requester's answers (no invented facts). A goal that names a number needs a facts-present check. A "
                  "workflow whose atom breaks a rule or lacks a field is rejected.")
DOC_NOTE = ("\nOUTPUT TYPE: {label}. The department writes a markdown document, not a website: where a role file says "
            "site, website, page or HTML, read the document. There is no browser check, style guide or design rubric; "
            "Audit uses the document rubric and code checks the necessary conditions.")
WF_TYPES = '"file-exists"|"grep-count"|"element-count"|"named-test"|"facts-present"'
WF_KINDS = '"section"|"element"|"facts"|"file"|"test"|"tap"|"behavior"'


def export_prd(doc, run_dir):
    """prd.docx next to prd.md. None when written, else why not."""
    return export_docx(doc, os.path.join(run_dir, "prd.docx"))


OUTPUTS = {
    "website": {"label": "Website", "noun": "site", "sections": (), "conditions": "website.json", "rubric": RUBRIC,
                "rates": "rates", "check_types": CHECK_TYPES, "builder": "builder.md", "export": None,
                "done": WORKABLE_OUTCOME},
    "prd": {"label": "PRD document", "noun": "PRD", "sections": PRD_SECTIONS, "conditions": "prd.json",
            "rubric": DOC_RUBRIC, "rates": "rates_prd", "check_types": DOC_CHECK_TYPES, "builder": PRD_BUILDER,
            "export": export_prd, "file": "prd.md", "intake": PRD_INTAKE, "charter_fmt": PRD_CHARTER_FMT,
            "plan": PRD_PLAN, "atom_rules": PRD_ATOM_RULES, "mock_charter": mock_prd_charter,
            "mock_section": mock_doc_section, "mock_atom": mock_doc_atom,
            "done": "Done: the PRD is workable. Every done-when item, rubric item and necessary condition passes."},
}
OUTPUTS["hld"] = {"label": "HLD document", "noun": "HLD", "sections": HLD_SECTIONS, "conditions": "hld.json",
                  "rubric": HLD_RUBRIC, "rates": "rates_prd", "check_types": DOC_CHECK_TYPES, "builder": HLD_BUILDER,
                  "export": export_hld, "file": "hld.md", "plan": HLD_PLAN,
                  "intake": PRD_INTAKE.replace("PRD", "HLD"), "charter_fmt": PRD_CHARTER_FMT,
                  "mock_charter": mock_prd_charter,
                  "atom_rules": PRD_ATOM_RULES.replace("PRD", "HLD"), "mock_section": mock_hld_section,
                  "mock_atom": doc_atom_for(HLD_SECTIONS, {"Components": "C"}),
                  "done": "Done: the HLD is workable. Every done-when item, rubric item and necessary condition passes."}
OUTPUTS["lld"] = {"label": "LLD document", "noun": "LLD", "sections": LLD_SECTIONS, "conditions": "lld.json",
                  "rubric": LLD_RUBRIC, "rates": "rates_prd", "check_types": DOC_CHECK_TYPES, "builder": LLD_BUILDER,
                  "export": export_lld, "file": "lld.md", "plan": LLD_PLAN,
                  "intake": PRD_INTAKE.replace("PRD", "LLD"), "charter_fmt": PRD_CHARTER_FMT,
                  "mock_charter": mock_prd_charter,
                  "atom_rules": PRD_ATOM_RULES.replace("PRD", "LLD"), "mock_section": mock_lld_section,
                  "mock_atom": doc_atom_for(LLD_SECTIONS, {"Pages and sections": "S", "Interactive elements": "E",
                                                           "Form fields": "F"}),
                  "done": "Done: the LLD is workable. Every done-when item, rubric item and necessary condition passes."}
OUTPUT = "website"     # the current run's output type (set by load_rates at each reset; one run at a time, like NAMES)


def output_type(name=None):
    """A known output type: name, else J2_OUTPUT, else website."""
    t = str(name or os.getenv("J2_OUTPUT") or "website").strip().lower()
    return t if t in OUTPUTS else "website"


# ---- the style guide: written by the first workflow that runs, followed by every workflow agent
STYLE_FMT = ('{"palette": {"primary": hex, "secondary": hex, "accent": hex, "background": hex, "surface": hex, "text": hex, '
             '"muted": hex}, "fonts": {"heading": css font-family, "body": css font-family, "google_fonts_url": '
             'https://fonts.googleapis.com/... or null}, "spacing": {"unit": "8px", "section": "64px"}, "radius": "12px", '
             '"buttons": {"primary": css declarations, "secondary": css declarations}, "card": css declarations, '
             '"notes": str (one line: why this suits the business)}')
HEX = re.compile(r"#(?:[0-9a-fA-F]{3}){1,2}")
CSS_SAFE = re.compile(r'^[^<>{}"\\]*$')   # values only: nothing that can close a rule, a tag or an attribute


def parse_style(raw):
    """The style guide from a ```style block; ValueError with the reason when it can't be used."""
    m = re.search(r"```style\s*(\{.*?\})\s*```", raw, re.S)
    if not m:
        raise ValueError("the reply has no ```style block with the style guide (you run first, so you write it)")
    try:
        g = json.loads(m.group(1))
    except ValueError:
        raise ValueError("the ```style block isn't valid JSON") from None
    pal, fonts = g.get("palette") or {}, g.get("fonts") or {}
    missing = [k for k in ("primary", "background", "text") if not HEX.fullmatch(str(pal.get(k) or ""))]
    if missing:
        raise ValueError(f"the style guide's palette needs hex colours for {', '.join(missing)}")
    flat = [str(v) for v in list(pal.values()) + [fonts.get("heading"), fonts.get("body"), g.get("radius"),
                                                  g.get("card")] + list((g.get("spacing") or {}).values())
            + list((g.get("buttons") or {}).values()) if v]
    if not fonts.get("heading") or not fonts.get("body") or not all(CSS_SAFE.match(v) for v in flat):
        raise ValueError('the style guide needs heading and body fonts, and plain CSS values (no <, >, {, }, " or \\)')
    url = str(fonts.get("google_fonts_url") or "")
    fonts["google_fonts_url"] = url if url.startswith("https://fonts.googleapis.com/") and CSS_SAFE.match(url) else None
    return g


def style_css(g):
    """Head tags for the style guide: CSS variables (--j2-primary, --j2-font-body, ...) and the .j2-btn,
    .j2-btn-secondary and .j2-card classes every workflow agent uses."""
    pal, fonts, sp = g["palette"], g["fonts"], g.get("spacing") or {}
    var = "".join(f"--j2-{k}:{v};" for k, v in pal.items() if HEX.fullmatch(str(v)))
    var += (f"--j2-font-heading:{fonts['heading']};--j2-font-body:{fonts['body']};--j2-space:{sp.get('unit', '8px')};"
            f"--j2-section:{sp.get('section', '64px')};--j2-radius:{g.get('radius', '8px')};")
    btn = g.get("buttons") or {}
    link = f'<link rel="stylesheet" href="{html.escape(fonts["google_fonts_url"])}">' if fonts.get("google_fonts_url") else ""
    return (link + "<style>:root{" + var + "}body{margin:0;font-family:var(--j2-font-body);color:var(--j2-text);"
            "background:var(--j2-background)}h1,h2,h3,h4{font-family:var(--j2-font-heading)}"
            f".j2-btn{{display:inline-block;{btn.get('primary', '')}}}.j2-btn-secondary{{display:inline-block;"
            f"{btn.get('secondary', '')}}}.j2-card{{{g.get('card', '')}}}</style>")


MOCK_STYLE = {"palette": {"primary": "#1f5f8b", "secondary": "#e8eef3", "accent": "#b4531a", "background": "#ffffff",
                          "surface": "#f5f7f9", "text": "#1f2933", "muted": "#52606d"},
              "fonts": {"heading": "Georgia, serif", "body": "system-ui, sans-serif", "google_fonts_url": None},
              "spacing": {"unit": "8px", "section": "56px"}, "radius": "10px",
              "buttons": {"primary": "background:#1f5f8b;color:#fff;padding:12px 20px;border-radius:10px;min-height:44px",
                          "secondary": "background:#e8eef3;color:#1f2933;padding:12px 20px;border-radius:10px"},
              "card": "background:#f5f7f9;border-radius:10px;padding:16px", "notes": "Calm, readable defaults (mock)"}


# ---- the site file: one marked section per workflow agent, assembled by code
def assemble(title, sections):
    head = "".join(f"\n<!-- J2:{sid} head start (owner {x['owner']}) -->\n{x['head']}\n<!-- J2:{sid} head end -->"
                   for sid, x in sections.items() if x.get("head"))
    body = "".join(f"\n<!-- J2:{sid} start (owner {x['owner']}) -->\n{x['body']}\n<!-- J2:{sid} end -->"
                   for sid, x in sections.items() if x.get("body"))
    return (f'<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n<title>{html.escape(title)}</title>'
            f"{head}\n</head>\n<body>{body}\n</body>\n</html>\n")


def parse_section(raw):
    """(head, body) from a workflow agent's reply; ValueError with the reason when it can't be used."""
    hm = re.search(r"```head\s*(.*?)```", raw, re.S | re.I)
    bm = re.search(r"```(?:section|html)\s*(.*?)```", raw, re.S | re.I)
    head, body = (hm.group(1).strip() if hm else ""), (bm.group(1).strip() if bm else "")
    if not head and not body:
        raise ValueError("the reply has no ```section or ```head block")
    if re.search(r"<\s*(?:!doctype|html|body|head)\b", body, re.I):
        raise ValueError("the section contains <html>, <head> or <body> tags; write only your section's markup")
    if re.search(r"<\s*(?:!doctype|html|body)\b", head, re.I):
        raise ValueError("the head block may hold only meta, link, style and script tags")
    for form in re.findall(r"<form\b.*?</form>", body, re.S | re.I):    # the browser check's rule, before it ships
        if not re.search(r"<button\b(?![^>]*\btype\s*=\s*[\"']?(?:button|reset)\b)|<input\b[^>]*\btype\s*=\s*[\"']?"
                         r"(?:submit|image)\b", form, re.I):
            raise ValueError('a <form> has no submit button; add <button type="submit"> inside it')
    return head, body


def parse_doc(raw):
    """("", markdown) from a document section agent's reply; ValueError with the reason when it can't be used."""
    m = re.search(r"```(?:markdown|md)\s*(.*?)```", raw, re.S | re.I)
    body = m.group(1).strip() if m else ""
    if not body:
        raise ValueError("the reply has no ```markdown block with your section")
    if re.search(r"<\s*/?\s*[a-z][a-z0-9]*\b[^>]*>", body, re.I):
        raise ValueError("the section contains HTML tags; write markdown only")
    if not MD_HEADING.match(body.splitlines()[0]):
        raise ValueError('the section must start with its "## <Section>" heading')
    return "", body


def assemble_doc(title, sections, order):
    """The document: a title, then every section's markdown in the output's section order (by each section's first
    heading); a section that names none goes last."""
    def rank(x):
        h = next((m[2] for ln in x["body"].splitlines() if (m := MD_HEADING.match(ln))), "")
        return next((i for i, s in enumerate(order) if squash(strip_md(h)).startswith(s.lower())), len(order))
    bodies = sorted((x for x in sections.values() if x.get("body")), key=rank)    # stable: ties keep build order
    doc = f"# {title}\n\n" + "\n\n".join(x["body"].strip() for x in bodies) + "\n"
    asm = doc_assumptions(doc)          # code's closing section, made from the builders' Assumption: lines
    return doc + (f"\n## {ASSUMPTIONS_HEADING}\n\n" + "\n".join(f"- {a}" for a in asm) + "\n" if asm else "")


FORCED_NOTES = ((r"price|menu", "Two items on the menu have no price next to them, so the menu with prices isn't complete."),
                (r"hour|open", "Sunday is missing from the opening hours list, so the hours aren't complete."),
                (r"address|map|where|direction", "The address is shown, but the map link has no destination, so it "
                                                 "opens an empty map."),
                (r"form|order|book", "After the form is submitted nothing tells the visitor it went through: no "
                                     "confirmation message is shown."),
                (r"loyal|card|stamp", "The loyalty explainer doesn't say when the free item is given, so the rule "
                                      "isn't clear."))


def forced_note(criterion):
    """A concrete, plausible issue for J2_FORCE_AUDIT_FAIL, chosen from the item's words; J2_FORCE_AUDIT_NOTE sets it.
    Kept to content gaps the render check can't contradict (no overflow or JS claims)."""
    if os.getenv("J2_FORCE_AUDIT_NOTE"):
        return os.getenv("J2_FORCE_AUDIT_NOTE")
    return next((n for pat, n in FORCED_NOTES if re.search(pat, criterion, re.I)),
                f"The section for this item is on the page but leaves out details a visitor needs, so \"{criterion}\" "
                "isn't fully met.")


FINDING_ID = re.compile(r"F-\d+")


def workable(audit):
    """Code's verdict on the latest audit, no model: what stops the site being workable ([] = workable).
    Every done-when item passes Audit (a "to be confirmed" item never does), every necessary condition is met
    (conditions/<type>.json; for a website that includes nothing to be confirmed), and for a website the render check
    found no JS errors, no horizontal overflow at 360px or 1280px, no broken form and no broken in-page link."""
    if not audit:
        return ["no site has been audited yet"]
    num = lambda r: r["id"][2:] if FINDING_ID.fullmatch(str(r["id"])) else r["item"]    # F-n is charter item n
    out = [f"done-when item {num(r)} ({r['id']} {r['verdict']}): {r['criterion']}" for r in audit["results"]
           if not r["pass"]]
    if audit.get("output", "website") != "website":
        # A document: Audit accepts (every item and document rubric item passes) and every necessary condition is met.
        return (out + [f"document rubric {r['id']} ({r['verdict']}): {r['criterion']}" for r in audit.get("rubric", [])
                       if not r["pass"]]
                + [f"necessary condition {c['id']} not met: {c['text']} ({c['evidence']})"
                   for c in audit.get("conditions") or [] if blocks(c)])
    out += [f"design rubric {r['id']} ({r['verdict']}): {r['criterion']}" for r in audit.get("rubric", []) if not r["pass"]]
    out += [f"necessary condition {c['id']} not met: {c['text']} ({c['evidence']})"
            for c in audit.get("conditions") or [] if blocks(c)]
    mech = audit.get("mechanical") or {}
    if not mech.get("available"):
        return out + [f"render check didn't run, so JS errors, overflow and forms are unverified: {mech.get('note')}"]
    return out + j2_check.problems(mech)


def check_summary(audit, what="the document"):
    """What the last audit found, in plain words, written by code from its results (never by a model): how many checklist
    items and conditions are met, and which are not."""
    if not audit:
        return f"Nothing has been checked in {what} yet."
    res, rub = audit.get("results") or [], audit.get("rubric") or []
    conds = [c for c in audit.get("conditions") or [] if c.get("severity") != "warn"]
    bad = [c for c in conds if not c["met"]]
    out = [f"{sum(r['pass'] for r in res)} of {len(res)} checklist items pass"]
    if rub:
        out.append(f"{sum(r['pass'] for r in rub)} of {len(rub)} quality checks pass")
    out.append(f"{len(conds) - len(bad)} of {len(conds)} required conditions are met")
    txt = f"Check results for {what}: " + ", ".join(out) + "."
    if bad:
        txt += " Not met: " + "; ".join(c["text"] for c in bad) + "."
    unmet = [r["criterion"] for r in res + rub if not r["pass"]]
    if unmet:
        txt += " Still failing: " + "; ".join(unmet) + "."
    return txt


def failure_signature(audit):
    return sorted(f"{r['criterion']}|{r['verdict']}" for r in audit["results"] + audit.get("rubric", []) if not r["pass"]) \
        + sorted(j2_check.problems(audit.get("mechanical") or {})) \
        + sorted(f"{c['id']}|{c['evidence']}" for c in audit.get("conditions") or [] if blocks(c))


# ---- Adaptation's memory: lessons from fixes Audit confirmed, kept across runs in memory/lessons.jsonl
def memory_on():
    return os.getenv("J2_MEMORY", "1") != "0"


def realization_on():
    """On by default: Adaptation's first website plan goes through the realization axis with certainty (the planning
    part of agents/j2/adaptation.md, confidence-judge.md; j2_realization.py). J2_REALIZATION=0 turns it off."""
    return os.getenv("J2_REALIZATION", "1") != "0"


def load_lessons():
    if not memory_on() or not os.path.exists(MEMORY_FILE):
        return []
    out = []
    with open(MEMORY_FILE, encoding="utf-8") as f:
        for ln in f:
            try:
                out.append(json.loads(ln))
            except ValueError:
                pass
    return out


def save_lesson(lesson):
    if not memory_on():
        return
    lesson = {**lesson, **({"mock": True} if llm.MOCK else {})}   # mock memory never reaches a live run
    os.makedirs(os.path.dirname(MEMORY_FILE), exist_ok=True)
    with open(MEMORY_FILE, "a", encoding="utf-8") as f:
        f.write(json.dumps(lesson, ensure_ascii=False) + "\n")


def clear_lessons():
    """Delete memory/lessons.jsonl; returns how many lessons it held."""
    if not os.path.exists(MEMORY_FILE):
        return 0
    with open(MEMORY_FILE, encoding="utf-8") as f:
        n = sum(1 for ln in f if ln.strip())
    os.remove(MEMORY_FILE)
    return n


def load_playbooks():
    if not memory_on() or not os.path.exists(PLAYBOOK_FILE):
        return []
    with open(PLAYBOOK_FILE, encoding="utf-8") as f:
        return [json.loads(ln) for ln in f if ln.strip().startswith("{")]


def save_playbook(pb):
    """One playbook per run: a later Done in the same run (after a change request) replaces it."""
    if not memory_on():
        return
    pb = {**pb, **({"mock": True} if llm.MOCK else {})}       # mock memory never reaches a live run
    keep = [x for x in load_playbooks() if x.get("id") != pb["id"]]
    os.makedirs(os.path.dirname(PLAYBOOK_FILE), exist_ok=True)
    with open(PLAYBOOK_FILE, "w", encoding="utf-8") as f:
        f.writelines(json.dumps(x, ensure_ascii=False) + "\n" for x in keep + [pb])


def clear_playbooks():
    if not os.path.exists(PLAYBOOK_FILE):
        return 0
    with open(PLAYBOOK_FILE, encoding="utf-8") as f:
        n = sum(1 for ln in f if ln.strip())
    os.remove(PLAYBOOK_FILE)
    return n


def btype(x):
    return " ".join(str(x or "").lower().split())


def relevant_playbooks(pbs, business_type, text, n=MAX_PLAYBOOKS, include_test=False):
    """Playbooks for the same business type only (a cafe's plan doesn't fit a clothing shop); among those, the ones
    sharing most words with the brief, newest first on a tie. Test playbooks (J2_FORCE_AUDIT_FAIL runs) only with
    include_test; playbooks written by mock runs never reach a live run."""
    bt, want = btype(business_type), set(WORD.findall(str(text).lower()))
    if bt in ("", "business"):
        return []
    scored = [(len(want & set(WORD.findall(str(x.get("brief")).lower()))), i, x) for i, x in enumerate(pbs)
              if btype(x.get("business_type")) == bt and (include_test or not x.get("test"))
              and (llm.MOCK or not x.get("mock"))]
    return [x for *_, x in sorted(scored, key=lambda t: (-t[0], -t[1]))][:n]


WORD = re.compile(r"[a-z]{4,}")


def relevant_lessons(lessons, text, n=MAX_LESSONS, include_test=False, output="website"):
    """The n lessons sharing the most words with the current gap; newest first on a tie. Test lessons (from a
    J2_FORCE_AUDIT_FAIL run) are left out unless include_test (J2_MEMORY_TEST_LESSONS=1, demo runs only).
    ponytail: word overlap, swap for embeddings if the lesson file grows past a few hundred lines."""
    want = set(WORD.findall(text.lower()))
    lessons = [x for x in lessons if (include_test or not x.get("test")) and (llm.MOCK or not x.get("mock"))
               and x.get("output", "website") == output]      # a PRD's lessons never reach a website, and back
    scored = [(len(want & set(WORD.findall(" ".join(str(x.get(k, "")) for k in ("problem", "cause", "fix")).lower()))),
               i, x) for i, x in enumerate(lessons)]
    return [x for sc, i, x in sorted(scored, key=lambda t: (-t[0], -t[1])) if sc > 0][:n]


# ---- behind-the-scenes view: one lane per agent, live or from a saved run
FUNCTIONS = ("identity", "adaptation", "priority", "coordination", "audit", "validator")


def lanes(chat, workflows, thinking, turn, status, paused, step, write_lock):
    """Each agent's status (working / waiting / blocked / idle) and what it is on, worked out by code from the run."""
    ids = list(FUNCTIONS) + [w["id"] for w in workflows]
    last = {}
    for i, m in enumerate(chat):
        if m["who"] in ids:
            last[m["who"]] = i
    live = status == "running" or bool(paused)
    out = []
    for fn in ids:
        w = next((x for x in workflows if x["id"] == fn), None)
        i = last.get(fn)
        m = chat[i] if i is not None else None
        lane = {"id": fn, "name": NAMES.get(fn, f"{fn} Builder" if is_wf_agent(fn) else fn), "status": "idle",
                "doing": (m["text"][:140] if m else ("code: checks and registers workflows" if fn == "validator" else "")),
                "last_msg": m["id"] if m else None}
        nxt = m.get("next") if m else None
        answered = nxt and any(x["who"] == nxt for x in chat[i + 1:])
        if live and (thinking == fn or (turn == fn and status == "running")):
            lane.update(status="working", doing=(f"writing section {w['section']}: {w['name']}" if w else STEPS.get(step, "")))
        elif w and w["status"] in ("held", "unfunded", "failed"):
            lane.update(status="blocked", doing=f"{w['status']}: {w.get('last_check') or w['name']}")
        elif fn == "identity" and paused:
            lane.update(status="waiting", on="Requester", doing=PAUSES.get(paused["kind"], ""))
        elif live and nxt and nxt not in (fn, "user") and not answered:
            lane.update(status="waiting", on=NAMES.get(nxt, nxt))
        tl = [x["lesson_used"] for x in chat if x["who"] == fn and x.get("lesson_test")]
        if tl:
            lane["test_lesson"] = tl[-1]
        if w:
            env = w.get("envelope") or {}
            lane["progress"] = {"workflow": w["name"], "section": w.get("section"), "workflow_status": w["status"],
                                "attempts": w.get("calls_used", 0), "attempt_limit": as_int(env.get("model_call_limit"), 1),
                                "retries": w.get("retries_used", 0), "retry_limit": as_int(env.get("retry_limit")),
                                "last_check": w.get("last_check"), "write_lock": write_lock == fn}
        out.append(lane)
    return out


SAVED = re.compile(r"j2-[\w-]+")


def saved_dir(name):
    """A saved run's folder, or None. Only names like j2-20261007-101500 inside RUNS_DIR (no path tricks)."""
    if not SAVED.fullmatch(str(name or "")):
        return None
    d = os.path.join(RUNS_DIR, name)
    return d if os.path.isdir(d) else None


def _read(d, name, default=None):
    try:
        with open(os.path.join(d, name), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return default


def _jsonl(d, name):
    out = []
    try:
        with open(os.path.join(d, name), encoding="utf-8") as f:
            for ln in f:
                try:
                    out.append(json.loads(ln))
                except ValueError:
                    pass
    except OSError:
        pass
    return out


def list_runs(limit=60):
    out = []
    for name in sorted(os.listdir(RUNS_DIR), reverse=True) if os.path.isdir(RUNS_DIR) else []:
        d = saved_dir(name)
        if not d or not os.path.exists(os.path.join(d, "chat.jsonl")):
            continue
        s = _read(d, "summary.json", {})
        first = next((m for m in _jsonl(d, "chat.jsonl") if m.get("who") == "user"), {})
        out.append({"id": name, "status": s.get("status") or "unfinished", "brief": (s.get("brief") or first.get("text")
                                                                                   or "")[:120]})
        if len(out) >= limit:
            break
    return out


def saved_snapshot(name, since=0):
    """The page's state for a saved run, read from its folder (read-only)."""
    d = saved_dir(name)
    if not d:
        return None
    rows = _jsonl(d, "chat.jsonl")
    head = rows[0] if rows and rows[0].get("meta") else {}
    chat = [m for m in rows if not m.get("meta")]
    s, wfs = _read(d, "summary.json", {}), _read(d, "workflows.json", [])
    sections = _read(d, "sections.json", {})
    audits = [a for a in _read(d, "audits.json", []) if (a or {}).get("kind") != "recheck"]   # a board re-check is not an audit
    site = os.path.join(d, "site")
    versions = len([f for f in os.listdir(site) if re.fullmatch(r"v\d+\.(?:html|md)", f)]) if os.path.isdir(site) else 0
    docs = _read(d, "documents.json", {})
    out = output_type(s.get("output_type") or saved_output(d)) if docs else saved_output(d)
    calls = [{k: v for k, v in c.items() if k not in ("system", "prompt", "raw")} for c in _jsonl(d, "calls.jsonl")]
    b = s.get("budget") or {}
    return {"saved": name, "chat": chat[since:], "total": len(chat), "status": "saved", "thinking": None, "step": 8,
            "steps": STEPS, "goal": _read(d, "goal.json"), "stamped": True, "workflows": wfs,
            "budget": b.get("granted") or BUDGET, "spent": b.get("used") or 0, "conflicts": _read(d, "conflicts.json", []),
            "audits": audits, "versions": versions, "cycle": s.get("cycles") or 0, "calls": s.get("ai_calls") or len(calls),
            "atoms_v1": {v["id"]: v for v in (_read(d, "atoms.json", {}).get("atoms") or j2_schema.run_views(wfs, ended=True)[0])},
            "error": None, "done": True, "turn": None, "cycle_turns": 0, "max_turns": MAX_TURNS, "paused": None,
            "mock": head.get("mock"), "model": head.get("model"), "fixes": head.get("j2_fixes", 0), "placeholders": [],
            "force_audit_fail": head.get("force_audit_fail"), "test_lessons": bool(head.get("test_lessons")),
            "force_question": head.get("force_question"), "has_board": os.path.isfile(os.path.join(d, "board.json")),
            "report": _read(d, "report.json"), "outcome": s.get("outcome"),
            "sections": [{"id": k, "owner": x["owner"], "workflow": x.get("workflow"),
                          "bytes": len((x["head"] + x["body"]).encode("utf-8"))} for k, x in sections.items()],
            "write_lock": None, "agents": {w["id"]: w.get("agent") or f"{w['id']} Builder" for w in wfs},
            "lanes": lanes(chat, wfs, None, None, "saved", None, 8, None), "call_log": calls,
            "workable": s.get("workable") if s else (workable(audits[-1]) if audits else None),
            "memory": False, "lessons": [], "lessons_used": (s.get("lessons") or {}).get("used", []),
            "lessons_saved": (s.get("lessons") or {}).get("saved", []), "style": _read(d, "style.json"), "rubric": RUBRIC,
            "playbooks": [], "playbooks_used": (s.get("playbooks") or {}).get("used", []),
            "playbook_saved": (s.get("playbooks") or {}).get("saved"), **output_info(out, d, audits, list(docs))}


def saved_output(d):
    """A saved run's output type, from the first line of its chat.jsonl (runs from before output types: website)."""
    head = (_jsonl(d, "chat.jsonl")[:1] or [{}])[0] if d else {}
    return output_type(head.get("output") or "website")


DOC_NAMES = {"prd": "product requirements", "hld": "high-level design", "lld": "low-level design"}   # plain words


def output_info(out, run_dir, audits, docset=None):
    """What the page needs about the output type: its label, whether it is a document, the latest necessary
    conditions and which exported files exist. A design-document run lists every document it has written."""
    o = OUTPUTS[out]
    docset = docset or []
    files = [f for f in ([o["file"], o["file"][:-3] + ".docx"] if o.get("file") else []) if os.path.exists(os.path.join(run_dir, f))]
    if docset:
        files = [f for k in docset for f in (f"{k}.md", f"{k}.docx") if os.path.exists(os.path.join(run_dir, f))]
    return {"output": out, "output_label": o["label"], "doc": bool(o["sections"]),
            "rubric": o["rubric"], "conditions": (audits[-1].get("conditions") if audits else None), "files": files,
            "docset": docset, "docset_files": [{"kind": k, "label": DOC_NAMES[k], "md": f"{k}.md", "docx": f"{k}.docx"}
                                               for k in docset if os.path.exists(os.path.join(run_dir, f"{k}.md"))]}


def atom_section_body(d, wfs, atom_id):
    """The markdown (or HTML) section an atom's workflow wrote, from the run folder: sections.json for the current
    stage, stage-<document>.json for the earlier ones."""
    w = next((x for x in wfs if any(a["id"] == atom_id for a in x.get("atoms") or [])), None)
    if not w or not d:
        return None
    sources = [_read(d, "sections.json", {})] + [(_read(d, f, {}) or {}).get("sections") or {}
                                                  for f in sorted(os.listdir(d)) if re.fullmatch(r"stage-\w+\.json", f)]
    return next((x["body"] for src in sources for k, x in src.items() if k == w.get("section") and x.get("body")), None)


def call_record(d, cid):
    """One AI call in full (system prompt, prompt, raw reply) from a run folder's calls.jsonl."""
    return next((c for c in _jsonl(d, "calls.jsonl") if c.get("id") == cid), None) if d else None


class J2Run:
    def __init__(self):
        self.lock = threading.Lock()
        self.reset()

    # ------------------------------------------------------------------ state
    def reset(self):
        self.run_id = time.strftime("%Y%m%d-%H%M%S")
        self.dir = os.path.join(RUNS_DIR, "j2-" + self.run_id)
        n = 2
        while os.path.exists(self.dir):          # two runs started in the same second get separate folders
            self.dir = os.path.join(RUNS_DIR, f"j2-{self.run_id}-{n}")
            n += 1
        self.run_id = os.path.basename(self.dir)[3:]     # and separate ids (lesson and playbook ids use it)
        self.chat, self.workflows, self.conflicts, self.audits, self.sites = [], [], [], [], []
        # design documents (opt-in, code-detected): the stages of a run (PRD, HLD, LLD, then the site), each document's
        # agent_output, and what earlier stages left behind. Empty for every other run.
        self.docset, self.stages, self.stage_i, self.chain, self.docs = [], [], 0, {}, {}
        self.deed_log, self.deed_seen = [], set()      # send deeds done (note, result) and the notices already given
        self.milestone_log = []      # milestone emails (J2_MILESTONE_EMAILS=1, or the brief asks for updates by email): one per goal reached
        self.open_docs = []          # documents the Requester accepted as they were, with what was still not right
        self.given = {}              # workflow id -> {"edges": [{atom, why}], "whole": bool}: what its builder's prompt held (atom/1)
        self.board = j2_board.new_board()   # board/1 (j2_board.py): questions and answers, behind J2_BOARD=1
        self.board_read = {}         # function -> board ids it has just read: post() tags its next message "[read Q-n]"
        self.rechecks, self.doc_changes = [], []    # board/1 part C: recheck entries (audits.json) and doc changes made
        self.board_failed = []       # atoms a board change left failing: Adaptation takes them after this turn
        self.board_asks = []         # a PRD change a board answer asked for: Identity takes it after this turn
        self.past_workflows, self.stage_archive, self.site_tag, self.stage_checked = [], [], "", True
        self.blueprint, self.proposed_lld, self.lld_revisions, self.lld_conflicts = None, {}, [], {}   # the LLD as the spec
        self.stage_calls0 = 0         # AI calls before the current stage (a document's call range, for the cost estimate)
        self.goal, self.stamped, self.request = None, False, ""
        self.spent, self.cycle, self.calls, self.step = 0, 0, 0, 0
        self.budget, self.raises = BUDGET, []
        self.idle_cycles, self.cycle_registered = 0, True
        self.status, self.thinking, self.error, self.done = "idle", None, None, False
        self.feedback = None
        self.atom, self.proposed_atoms, self.receipts = None, {}, []        # atom id of the build call in progress; atoms by workflow name
        self.turn, self.handoff, self.outcome = None, {}, None
        self.paused = None          # {"kind", "questions", "to", "handoff", "msg"} while the run waits on the Requester
        self.intake, self.decisions = [], []   # J2.a question rounds; every Requester answer to Identity
        self.drop_answer, self.relay, self.a_msg = None, None, None
        self.fixes = os.getenv("J2_FIXES", "0") == "1"   # read per run, so one process can test both modes
        # Test switch, off unless set: in cycle 1, Audit's result for this done-when item is forced to NEEDS_REVISION
        self.force_fail = as_int(os.getenv("J2_FORCE_AUDIT_FAIL"), 0) if os.getenv("J2_FORCE_AUDIT_FAIL") else 0
        self.forced = False
        # Test switch, off unless set: a comma list of askers (identity|adaptation|priority|coordination|audit, or a
        # workflow id or role word like products|order|filter matched against its name). In mock mode each named
        # asker asks one question, once per run (message board, J2_BOARD=1 only).
        self.force_question = os.getenv("J2_FORCE_QUESTION") or ""
        self.force_qs = [x.strip().lower() for x in self.force_question.split(",") if x.strip()]
        self.question_forced = False     # the builders' one forced question
        self.forced_asks = set()         # the functions that have asked theirs
        self.test_lessons = os.getenv("J2_MEMORY_TEST_LESSONS") == "1"   # demo only: J2.b may read test lessons
        self.p_refused = {}         # workflow name -> why: what Priority refused in the current negotiation
        self.rejections = []        # every validator rejection, with its exact reason
        self.events, self.agent_turns, self.report, self.started = [], {}, None, None   # for the run report
        self.export_note = None     # a document's export at the end: None when written, else why not
        self.set_output(output_type())       # J2_OUTPUT, read per run (also loads its rate card)
        self.cycle_limit = self.cycle_turns = 0
        self.summary, self.summary_upto = "", 0
        self.sections, self.write_lock, self.bargain = {}, None, None   # section id -> {owner, head, body, workflow}
        self.memory = memory_on()
        self.call_log, self.pending_calls = [], {}     # every AI call (calls.jsonl); calls not yet shown in a message
        self.style = None                              # the style guide, written by the first workflow that runs
        self.playbooks_used, self.playbook_saved = [], None
        self.lessons_used, self.lessons_saved = [], []
        self.fail_floor = 0          # audits before this index don't count for "same failure twice in a row"
        self.last_limit = None       # the last limit question: state then and the Requester's decision
        self.adapt_idle = None       # _adapt_state() when Adaptation last proposed nothing for lack of budget
        # J2_REALIZATION=1: the judged realization plan, its decisions, whether its blocking facts were asked, the record
        self.realization = realization_on()
        self.real, self.real_decisions, self.real_asked, self.real_summary = None, {}, False, None
        self._begin_cycle(advance=False)

    def set_output(self, out):
        """Pick what the department builds (OUTPUTS) before the first message: J2_OUTPUT, or the start page.
        False once the run has started. Only what is built and checked changes; the flow is the same."""
        if self.chat or self.status == "running":
            return False
        self._apply_output(out)
        return True

    def _apply_output(self, out):
        """Switch what is built (set_output before the first message; a design-document run at each stage)."""
        self.output = output_type(out)
        self.out, self.doc = OUTPUTS[self.output], bool(OUTPUTS[self.output]["sections"])
        # J2_REALIZATION: no rate card; Adaptation estimates each workflow's budget (estimated_resources.money)
        load_rates(self.output, card=not (getattr(self, "realization", None) or realization_on()))
        types = "|".join(f'"{t}"' for t in self.out["check_types"])
        # A document's agents get its own check types, finding ids and atom rules; a website's prompts stay as they were.
        kinds = "|".join(f'"{k}"' for k in expect_kinds(self.out["check_types"]))
        doc_fmt = lambda f: f.replace(WF_TYPES, types).replace(WF_KINDS, kinds).replace(
            ' (section/element only, default 1)', ' (leave out)').replace(', "steps": [str] (behavior only)', '').replace('"R-6"', f'"{self.out["rubric"][0][0]}", "P2"')
        self.WF_FMT = doc_fmt(J2Run.WF_FMT) if self.doc else J2Run.WF_FMT
        self.GRANT_FMT = doc_fmt(J2Run.GRANT_FMT) if self.doc else J2Run.GRANT_FMT
        self.ATOM_RULES = self.out.get("atom_rules") or J2Run.ATOM_RULES

    def _begin_cycle(self, advance=True):
        if advance:
            self.cycle += 1
        self.cycle_turns = self.exchanges = self.hold_bounces = 0
        self.conflict = None                 # the Adaptation <-> Priority negotiation in progress, once they disagree
        self.a_pos, self.p_pos = [], []      # Adaptation's and Priority's latest positions this cycle
        self.a_named = []                    # done-when items Adaptation's last proposal covered or left out
        self.deferred = {}                   # done-when item text -> rounds (cycles) it was left out or refused
        self.pending_left = None             # items a ruling left out: Identity asks after this cycle's Audit

    def _ship_outcome(self, why):
        """Outcome of "finish as it is". "Done with open items" only with items to name: a workable site is done."""
        if not self.sites:
            return f"Stopped {why}: nothing was built."
        fails = workable(self.audits[-1]) if self.audits else ["nothing has been checked"]
        fails = self.open_docs + fails
        if not fails:
            return self.out["done"]
        return f"Done with open items: shipped as is {why}. Still not workable: {'; '.join(fails)}."

    def _turn_cap(self):
        """Turns allowed this cycle, scaled with the work: each workflow registered this cycle needs its own turns."""
        return max(MAX_TURNS, 10 + 4 * sum(w.get("cycle") == self.cycle for w in self.workflows))

    def _atoms_v1(self):
        try:
            return {v["id"]: v for v in j2_schema.run_views(self._all_wfs(), self.given)[0]}
        except Exception:
            return {}

    def snapshot(self, since=0):
        return {"chat": self.chat[since:], "total": len(self.chat), "status": self.status,
                "thinking": self.thinking, "step": self.step, "steps": STEPS, "goal": self.goal,
                "stamped": self.stamped, "workflows": self._all_wfs(), "budget": self.budget, "spent": self.spent,
                "conflicts": self.conflicts, "audits": self.audits, "versions": len(self.sites),
                "cycle": self.cycle, "calls": self.calls, "error": self.error, "done": self.done,
                "atoms_v1": self._atoms_v1(),
                "turn": self.turn, "cycle_turns": self.cycle_turns, "max_turns": self._turn_cap(),
                "paused": ({k: self.paused[k] for k in ("kind", "questions", "msg")}
                           | {"label": PAUSES.get(self.paused["kind"], ""),
                              "choices": self.paused["handoff"].get("choices")}) if self.paused else None,
                "outcome": self.outcome,          # read-only, for the Simple view
                "style": self.style, "rubric": RUBRIC,
                "playbooks": load_playbooks() if self.memory else [], "playbooks_used": self.playbooks_used,
                "playbook_saved": self.playbook_saved,
                "mock": llm.MOCK, "model": None if llm.MOCK else llm.MODEL, "fixes": int(self.fixes),
                "force_audit_fail": self.force_fail or None, "test_lessons": self.test_lessons,
                "force_question": self.force_question or None, "has_board": bool(self.board["entries"]),
                "placeholders": self._placeholders(), "report": self.report,
                "sections": [{"id": k, "owner": x["owner"], "workflow": x.get("workflow"),
                              "bytes": len((x["head"] + x["body"]).encode("utf-8"))} for k, x in self.sections.items()],
                "write_lock": self.write_lock, "agents": {w["id"]: NAMES.get(w["id"]) for w in self._all_wfs()},
                "lanes": lanes(self.chat, self._all_wfs(), self.thinking, self.turn, self.status, self.paused, self.step,
                               self.write_lock),
                "call_log": self.call_log,
                "workable": workable(self.audits[-1]) if self.audits else None,
                "memory": self.memory, "lessons": load_lessons() if self.memory else [],
                "lessons_used": self.lessons_used, "lessons_saved": [x["id"] for x in self.lessons_saved],
                **output_info(self.output, self.dir, self.audits, [k for k in self.docset if k in self.docs])}

    def _placeholders(self):
        out = []
        for fn in ("team", "builder", "identity", "adaptation", "priority", "coordination", "audit"):   # builder.md: workflow agents
            p = os.path.join(J2_DIR, f"{fn}.md")
            if os.path.exists(p) and "PLACEHOLDER" in open(p, encoding="utf-8").read(300):
                out.append(fn)
        return out

    def _event(self, agent, kind, detail, **extra):
        """Something the run report counts: redirects, invalid moves, refusals, holds, limits, pauses, conflicts."""
        self.events.append({"agent": agent, "type": kind, "detail": detail, "msg": len(self.chat) + 1,
                            "cycle": self.cycle, **extra})

    def _raise_budget(self, add, why):
        """Only after the Requester approved more budget in answer to Identity's question; capped at BUDGET_CAP."""
        if add <= 0:
            return
        new = min(self.budget + add, BUDGET_CAP)
        if new <= self.budget:
            self.post("system", f"The Requester approved +{add}, but the budget is already at the cap ({BUDGET_CAP}). "
                                "No change.")
            return
        old, self.budget = self.budget, new
        self.raises.append({"cycle": self.cycle, "asked": add, "from": old, "to": new, "why": why})
        self._event("identity", "budget_raise", f"{old} -> {new} (asked +{add}, cap {BUDGET_CAP}; {why})")
        self.post("system", f"Budget raised on the Requester's approval: {old} -> {new}"
                  + (f" (asked +{add}; capped at {BUDGET_CAP})" if old + add > new else "")
                  + f". Remaining: {self.budget - self.spent}.")

    def post(self, who, text, **meta):
        if who == "identity" and meta.get("to_user") and meta.get("act") != "ASK_REQUESTER" and not meta.get("shaped"):
            if re.match(r"\s*@(?!Requester\b)\w", str(text or "")):
                # Addressed to an agent but sent to the Requester: never re-tag it, record it (call() should have
                # caught it).
                self.events.append({"agent": who, "type": "misaddressed", "detail": str(text)[:120],
                                    "msg": len(self.chat) + 1, "cycle": self.cycle})
            else:
                text = "@Requester " + to_requester(text, keep=self._fact_source())   # _ask_user shapes its own
        if who == "identity" and meta.get("to_user"):
            text = self._requester_noun(text)
        if self.board_read.get(who) and not meta.get("to_user"):
            # Item 5: a function's message carries the same "[read Q-n]" tag a builder's does, once.
            ids = self.board_read.pop(who)
            text, meta = f"{text} [read {', '.join(ids)}]", {**meta, "board_read": ids}
        m = {"id": len(self.chat) + 1, "t": time.strftime("%H:%M:%S"), "who": who, "text": text, "step": self.step, **meta}
        if self.pending_calls.get(who):          # the AI calls behind this message, for the detail view
            m["calls"] = self.pending_calls.pop(who)
        if meta.get("warning"):
            self.events.append({"agent": "system", "type": "warning", "detail": text, "msg": m["id"], "cycle": self.cycle})
        self.chat.append(m)
        os.makedirs(self.dir, exist_ok=True)
        with open(os.path.join(self.dir, "chat.jsonl"), "a", encoding="utf-8") as f:
            if m["id"] == 1:   # first line of the file: which mode and model produced this run
                f.write(json.dumps({"meta": True, "run": self.run_id, "output": self.output, "j2_fixes": int(self.fixes),
                                    "force_audit_fail": self.force_fail or None,
                                    "force_question": self.force_question or None,
                                    "test_lessons": self.test_lessons or None,
                                    "realization": int(self.realization),
                                    "provider": None if llm.MOCK else llm.PROVIDER, "model": None if llm.MOCK else llm.MODEL,
                                    "mock": llm.MOCK}) + "\n")
            f.write(json.dumps(m, ensure_ascii=False) + "\n")
        return m

    def _all_wfs(self):
        """Every workflow of the run: the earlier stages' (documents) and the current stage's."""
        return self.past_workflows + self.workflows

    def save(self, name, data):
        if name == "workflows.json" and self.past_workflows:
            data = self.past_workflows + data
        os.makedirs(self.dir, exist_ok=True)
        with open(os.path.join(self.dir, name), "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)

    COST_KEYS = ("cost", "estimate", "envelope", "charged", "calls_used", "retries_used")

    def state(self, fn=None):
        """Structured state every agent gets on every call. Identity decides purpose, not resources, so its view
        leaves out the budget, the rate card and workflow costs (resources are Adaptation's and Priority's)."""
        blind = fn == "identity"
        conflicts = [{k: c.get(k) for k in ("id", "cycle", "topic", "between", "subject", "status", "resolved_by",
                                            "outcome")} for c in self.conflicts]
        audits = [{"cycle": a["cycle"], "version": a["version"], "all_pass": a["all_pass"],
                   "results": [{k: r.get(k) for k in ("id", "criterion", "verdict", "severity", "note")} for r in a["results"]]}
                  for a in self.audits]
        # Agents see every checklist item under its charter number (never its position, which shifts after a drop).
        charter = self.goal and {**{k: v for k, v in self.goal.items() if k != "charter_items"},
                                 "done_when": [{"item": self._no(c), "criterion": c}
                                               for c in self.goal.get("done_when") or []]}
        return json.dumps({"requester_brief": self.request, "requester_feedback": self.feedback, "charter": charter,
                           "requester_decisions": self.decisions,
                           "budget": None if blind else {"total": self.budget, "spent": self.spent,
                                                         "remaining": self.budget - self.spent},
                           "cycle": self.cycle, "negotiation": {"exchanges_used": self.exchanges,
                                                                "max_exchanges": MAX_ATTEMPTS},
                           "workflows": self._wfs_state(blind), "conflicts": conflicts, "audits": audits,
                           "site_versions": len(self.sites),
                           "site_sections": [{"section": k, "owner": x["owner"], "workflow": x.get("workflow")}
                                             for k, x in self.sections.items()],
                           "style_guide": self.style,
                           "rate_card": ({"prices": RATES, "rule": "a workflow costs the sum of its parts' prices; "
                                          "estimated_resources.money is ignored; list a part once per unit",
                                          "item_minimums": [{"item": self._no(c), "kind": item_kind(c),
                                                             "min": item_min(c)}
                                                            for c in (self.goal or {}).get("done_when") or []],
                                          "fix_rule": "one 'fix' part per Audit finding fixed"}
                                         if RATES and not blind else None)}, ensure_ascii=False)

    def _wfs_state(self, blind):
        wfs = [{k: v for k, v in w.items() if k not in self.COST_KEYS} for w in self.workflows] if blind else self.workflows
        return [{**w, "atoms": [slim_atom(a) for a in w.get("atoms") or []]} if w.get("atoms") else w
                for w in wfs] if slim() else wfs

    def chat_view(self):
        """Whole chat while it's short; otherwise a running summary of the older part plus the recent messages.
        Messages flagged test (J2_FORCE_AUDIT_FAIL, test lessons) are for humans only and never reach an agent."""
        full = "\n".join(line(m) for m in self.chat if for_agents(m))
        if len(full) <= CHAT_FULL_CHARS:
            return full
        older = self.chat[:-CHAT_TAIL]
        if not older:
            return full
        if not self.summary or len(older) - self.summary_upto >= SUMMARY_REFRESH:
            self._summarize(older)
        recent = [m for m in self.chat[self.summary_upto:] if for_agents(m)]
        return (f"(Long chat: messages 1-{self.summary_upto} are summarised, the last {len(recent)} are in full.)\n"
                f"SUMMARY OF EARLIER MESSAGES:\n{self.summary}\n\nRECENT MESSAGES:\n"
                + "\n".join(line(m) for m in recent))

    def _summarize(self, older):
        if self.calls and DELAY:
            time.sleep(DELAY)
        self.calls += 1
        new = [m for m in older[self.summary_upto:] if for_agents(m)]
        self.summary = llm.ask(
            "You keep a running summary of a team's group chat. Write at most 12 short bullets: decisions, workflows "
            "agreed or refused, conflicts and how they were settled, audit results, open issues. Facts only.",
            f"PREVIOUS SUMMARY:\n{self.summary or '(none)'}\n\nNEW MESSAGES:\n" + "\n".join(line(m) for m in new),
            lambda: f"- Messages 1-{len(older)}: charter recorded, workflows negotiated, built and audited (mock summary).",
            max_tokens=600, effort="low").strip()
        self.summary_upto = len(older)
        os.makedirs(self.dir, exist_ok=True)
        with open(os.path.join(self.dir, "summary.txt"), "w", encoding="utf-8") as f:
            f.write(f"Messages 1-{self.summary_upto}\n{self.summary}\n")

    def context(self, fn=None):
        return f"STATE:\n{self.state(fn)}\n\nCHAT:\n{self.chat_view()}"

    REASON = '"reason": str (one short sentence: why you decided this)'
    # Every Identity call gets this: anything Identity writes to the Requester must read plainly at the source.
    # (The Simple view's word list is only a fallback.) Messages to other agents keep their precise terms.
    PLAIN_TO_REQUESTER = (
        "\nWHEN YOU WRITE TO THE REQUESTER (any say addressed to them, any question or option for them): they are a "
        "college student with no tech background. Use everyday words. Never use internal terms: charter, done-when, "
        "checklist, standards, reserve, fix, effort, budget, units, workflow, cycle, finding, audit, ruling, envelope, "
        "content types such as \"warmth content\", or the names of the functions (Identity, Adaptation, Priority, "
        "Coordination, Audit). At most two short sentences of context. Each option says what they get and what it costs, "
        "in plain words, e.g. \"Add a gallery of past work (takes a bit longer)\" / \"Keep it simple, no gallery\". "
        "No recommendation.")
    MOCK_REASON = "Scripted mock reply: no model reasoning."

    REPLY_DEFAULTS = ("Thanks, got it. We're working on it now.", "Go ahead with the Requester's latest answer.")

    def call(self, fn, reply_format, task, mock, max_tokens=900, json_reply=True, images=None, defaults=None,
             _retry=False, answer_only=False):
        """One agent's AI call. Every JSON reply format gets a "reason" field. Each call is logged to calls.jsonl
        (inputs, reply, reason, tokens, time) and linked to the agent's next chat message for the detail view.

        With J2_BOARD=1 every function's format gets the board fields the same way (board_question, board_used) and
        every function's prompt gets the answered entries (_board_context); answer_only is the board's own answer
        call, which may not ask back."""
        board = j2_board.enabled() and json_reply and fn in FUNCTIONS and not answer_only
        if board and reply_format.rstrip().endswith("}") and '"board_question"' not in reply_format:
            reply_format = reply_format.rstrip()[:-1] + ", " + self.FN_BOARD_FMT + "}"
        if json_reply and reply_format.rstrip().endswith("}") and '"reason"' not in reply_format:
            reply_format = reply_format.rstrip()[:-1] + ", " + self.REASON + "}"
        user = f"{self.context(fn)}\n\nYOUR TURN:\n{task}" + (self.PLAIN_TO_REQUESTER if fn == "identity" else "") \
            + (self._board_context() if board else "") \
            + (DOC_NOTE.format(label=self.out["label"]) if self.doc else "")
        system = system_prompt(fn, reply_format, self.out["builder"])
        if self.calls and DELAY:
            time.sleep(DELAY)
        self.calls += 1
        self.thinking = fn
        t0 = time.time()
        try:
            raw = llm.ask(system, user, mock, max_tokens=max_tokens, effort="low" if json_reply else "high",
                          json_mode=json_reply, images=images)
        finally:
            self.thinking = None
        reply = raw
        if json_reply:
            try:
                reply = llm.extract_json(raw)
            except ValueError:
                reply = {"say": raw[:400]}
            if llm.MOCK and not reply.get("reason"):
                reply["reason"] = self.MOCK_REASON
        self._log_call(fn, task, system, user, raw, reply if json_reply else None, time.time() - t0)
        if images:
            self.call_log[-1]["images"] = [os.path.relpath(p, self.dir) for p in images]
        if fn == "identity" and json_reply and '"to_requester"' in reply_format and reply.get("enough") is not False:
            reply = self._check_audiences(reply, reply_format, task, mock, max_tokens, defaults, _retry)
        if board and isinstance(reply, dict):
            self._board_turn(fn, reply)
        return reply

    def _check_audiences(self, r, fmt, task, mock, max_tokens, defaults, retried):
        """One message, one addressee: to_requester may not carry agent content (instructions, unit counts, workflow
        ids, an @agent), to_agents may not read as user-facing (thanks, I'll start, a question to the user). A mixed
        reply is regenerated once with the reason; if it is still mixed, the bad part becomes a template."""
        req = str(r.get("to_requester") or "")
        ag = str(r["to_agents"] if "to_agents" in r else r.get("say") or "")
        bad = ([f"to_requester has agent content ({m.group(0)!r})"] if (m := AGENT_CONTENT.search(req)) else []) \
            + ([f"to_agents reads as user-facing ({m.group(0)!r})"] if (m := USER_CONTENT.search(ag)) else [])
        if not bad:
            return r
        self._event("identity", "audience_mix", "; ".join(bad))
        if not retried:
            return self.call("identity", fmt, task + "\nYOUR LAST REPLY MIXED THE TWO AUDIENCES: " + "; ".join(bad)
                             + ". Write it again: to_requester only for the Requester (plain words, nothing for the "
                             "team), to_agents only the instruction for the team (no thanks, no questions to the "
                             "Requester).", mock, max_tokens=max_tokens, defaults=defaults, _retry=True)
        d_req, d_ag = defaults or self.REPLY_DEFAULTS
        r = dict(r)
        if AGENT_CONTENT.search(req) or (USER_CONTENT.search(ag) and not req):
            r["to_requester"] = d_req
        if USER_CONTENT.search(ag):
            r["to_agents"] = d_ag
        self._event("identity", "audience_template", "; ".join(bad))
        return r

    def _inputs(self, fn, task):
        """What the agent read, summarised by code: [label, text] pairs. The full prompt is in calls.jsonl."""
        g, out = self.goal or {}, []
        if g:
            out.append(["charter", f"{g.get('site_name') or 'site'}: {len(g.get('done_when') or [])} done-when items"
                        + (f", {len(g['to_be_confirmed'])} to be confirmed" if g.get("to_be_confirmed") else "")])
        if fn != "identity":
            out.append(["ledger", f"budget {self.spent}/{self.budget} used, {self.budget - self.spent} left"
                        + (f"; rate card on ({len(RATES)} parts)" if RATES else "")])
        if self.workflows:
            out.append(["workflows", f"{len(self.workflows)} registered: "
                        + ", ".join(f"{w['id']} {w['status']}" for w in self.workflows)])
        if self.audits:
            a = self.audits[-1]
            out.append(["findings", f"audit of v{a['version']}: {sum(r['pass'] for r in a['results'])}/{len(a['results'])} "
                        "pass" + "".join(f"; {r['id']} {r['verdict']}" for r in a["results"] if not r["pass"])])
        for label, mark in (("proposal", "ON THE TABLE:"), ("offer", "PRIORITY OFFERS:"), ("proposal", "RUN THIS WORKFLOW"),
                            ("findings", "AUDIT'S FINDINGS"), ("lessons", "LESSONS FROM EARLIER WORK"),
                            ("workflow", "YOUR WORKFLOW:"), ("conflict", "CONFLICT "), ("dispute", "DISAGREEMENT ")):
            i = task.find(mark)
            if i >= 0:
                out.append([label, task[i:i + 240].split("\n")[0] + ("…" if len(task) > i + 240 else "")])
        out.append(["chat", f"{len(self.chat)} messages" + (f" (1-{self.summary_upto} summarised)" if self.summary else "")])
        out.append(["task", task[:280] + ("…" if len(task) > 280 else "")])
        return out

    def _log_call(self, fn, task, system, user, raw, reply, secs):
        meta = dict(llm.LAST_META)
        rec = {"id": self.calls, "fn": fn, "step": self.step, "cycle": self.cycle, "t": time.strftime("%H:%M:%S"),
               "seconds": round(secs, 2), "model": None if llm.MOCK else meta.get("model") or llm.MODEL,
               "tokens": ({"input": meta.get("input_tokens"), "output": meta.get("output_tokens")}
                          if meta.get("input_tokens") or meta.get("output_tokens") else None),
               "chars": {"prompt": len(system) + len(user), "reply": len(raw)},
               "inputs": self._inputs(fn, task),
               "reply": {k: v for k, v in reply.items() if k != "reason"} if isinstance(reply, dict) else None,
               "reason": reply.get("reason") if isinstance(reply, dict) else None,
               **({"atom": self.atom} if self.atom else {})}
        self.call_log.append(rec)
        self.pending_calls.setdefault(fn, []).append(rec["id"])
        os.makedirs(self.dir, exist_ok=True)
        with open(os.path.join(self.dir, "calls.jsonl"), "a", encoding="utf-8") as f:
            f.write(json.dumps({**rec, "system": system, "prompt": user, "raw": raw}, ensure_ascii=False) + "\n")
        return rec

    # ------------------------------------------------------------------ routing
    def _pick(self, frm, wanted, allowed, fallback, why=""):
        """Return (next agent, redirect note or None). Only J2 routes are ever allowed."""
        assert all(a in ROUTES[frm] for a in allowed) and fallback in allowed, (frm, allowed, fallback)
        want = key(wanted)
        if want in allowed:
            return want, None
        if want == "user" and frm != "identity":
            # Only Identity talks to the Requester: _say sends this message to Identity first (see _run).
            self.relay = self.relay or {"from": frm}
            self._event(frm, "redirect", "addressed the Requester, routed to Identity")
            return fallback, (f"Route check: {NAMES[frm]} addressed the Requester. Only Identity talks to the "
                              "Requester, so the message goes to Identity first.")
        label = NAMES.get(want, str(wanted)) if wanted else "no one"
        self._event(frm, "redirect", f"picked {label}, routed to {NAMES[fallback]}" + (f" ({why})" if why else ""))
        return fallback, (f"Route check: {NAMES[frm]} picked {label} to speak next, which isn't an allowed J2 route "
                          f"here{' (' + why + ')' if why else ''}. Passing to {NAMES[fallback]}.")

    def _say(self, who, text, nxt=None, note=None, **meta):
        text = text or ""
        if who in ("adaptation", "priority", "coordination"):
            # The same message as this agent's last one adds nothing: don't post it again, Identity takes the turn.
            # Not Audit: its message reports a fresh check, and "the same result again" is news.
            norm = lambda t: " ".join(re.sub(r"^@[\w-]+[,:]?\s*", "", t or "").lower().split())
            last = next((m for m in reversed(self.chat) if m["who"] == who), None)
            if last and norm(last["text"]) == norm(text):
                self._event(who, "repeat", "same message as its last one, not posted again")
                self.post("system", f"{NAMES[who]} repeated its last message unchanged, so it isn't posted again. "
                                    "Identity takes the turn.")
                self.relay = {"from": who, "text": f"You repeated your last message unchanged: {text}"}
                return last
        if who in ROUTES and who != "identity" and (self.relay or TO_USER.search(text)):
            if not self.relay:
                self._event(who, "redirect", "addressed the Requester in the message, routed to Identity")
                note = (f"Route check: {NAMES[who]} addressed the Requester. Only Identity talks to the Requester, so "
                        "the message goes to Identity first.")
            self.relay = {"from": who, "text": text}
            nxt = "identity"
        m = self.post(who, address(text, nxt) if nxt else text, **({"next": nxt} if nxt else {}), **meta)
        if note:
            if not note.startswith("Route check"):
                self._event(who, "invalid_move", note)
            self.post("system", note, redirect=True)
        return m

    def _reply(self, r, nxt, note=None, **meta):
        """Identity after a Requester answer: to_requester goes to the Requester only, to_agents to the team only."""
        if r.get("to_requester"):
            self.post("identity", r["to_requester"], to_user=True, requester_only=True, audience="user")
        return self._say("identity", r.get("to_agents") or r.get("say") or "", nxt, note,
                         **({"audience": nxt} if nxt else {}), **meta)

    def _handoff(self, to, **h):
        self.turn, self.handoff = to, h

    def _wf(self, wid):
        return next((w for w in self.workflows if w["id"] == wid), None)

    def _conflict(self, cid):
        return next(c for c in self.conflicts if c["id"] == cid)

    def _save_conflicts(self):
        self.save("conflicts.json", self.conflicts)

    def _ask_user(self, say, questions, kind, to="identity", handoff=None, gaps=None, options_only=False, gaps_title=None,
                  **meta):
        """Identity asks the Requester and the run pauses. The reply resumes the run at `to` with `handoff`. `gaps`
        (the decision question) is appended by code after Identity's words, so the list and its count are one list."""
        src = self._fact_source()
        say = to_requester(say, question=not questions, keep=src)
        questions = [{**q, "question": self._requester_noun(requester_words(q["question"], src) or q["question"]),
                      "options": [self._requester_noun(requester_words(o, src) or o) for o in q.get("options") or []]}
                     for q in questions]
        if gaps is not None:
            # At most 2 sentences, then the full list, one gap per line; the header's counts add up to the list.
            say = " ".join(re.split(r"(?<=[.!?])\s+", say)[:2])
            say += gap_list(gaps, gaps_title)
        if len(questions) > 1:
            # One question per message: several become one list of what to send (options can't be kept).
            items = [re.sub(r"[?.]$", "", q["question"].strip()) for q in questions]
            questions = [{"question": "Please send these: " + "; ".join(items) + "."}]
        questions = [{**q, "options": q.get("options", [])[:3],
                      **({"budget": q["budget"][:3]} if q.get("budget") else {})} for q in questions]
        if kind != "limit":
            # Adding budget is offered only by the decision question (_identity_limit), whose options code builds
            # and applies. Approval and ruling_budget apply their own amount after the answer.
            questions = [{k: v for k, v in q.items() if k != "budget"} for q in questions]
        m = self.post("identity", "@Requester " + say + ("\n" + numbered(questions, options_only) if questions else ""),
                      act="ASK_REQUESTER", to_user=True, next="user", pause=kind, **meta)
        self.paused = {"kind": kind, "questions": questions, "to": to, "handoff": handoff or {}, "msg": m["id"]}
        self._event("identity", "pause", f"{kind}: {len(questions)} question(s) for the Requester", pause=kind)
        self.outcome = (f"Paused: Identity is waiting on the Requester "
                        f"({questions[0]['question'] if kind == 'limit' else PAUSES[kind]}).")
        self.turn = None

    # ------------------------------------------------------------------ entry
    def user_says(self, text):
        with self.lock:
            if self.status == "running":
                return False
            key = os.getenv("J2_OVERRIDE_KEY")
            if key and str(text).strip().startswith(key):
                return self._override(str(text).strip()[len(key):].strip())   # the key is never posted or stored
            self.post("user", text, **({"change": True} if self.done and not self.paused else {}))
            if self.done:
                self.done = False
                if not self.paused:          # an answer to Identity's questions is not a change request
                    self.feedback = text
            self.status = "running"
            threading.Thread(target=self._run_safe, daemon=True).start()
            return True

    def _override(self, cmd):
        """Owner override (J2_OVERRIDE_KEY, in .env only): code applies the command directly and posts a visible note.
        "budget +N" ignores the cap, "ship" ends the run as it is, "drop <item>" (charter number or words) leaves an item out."""
        g = self.goal or {}
        crit = g.get("done_when") or []
        note = None
        if m := re.fullmatch(r"budget\s*\+\s*(\d+)", cmd, re.I):
            old, self.budget = self.budget, self.budget + int(m[1])
            self.raises.append({"cycle": self.cycle, "asked": int(m[1]), "from": old, "to": self.budget,
                                "why": "owner override", "override": True})
            note = f"budget +{m[1]} ({old} -> {self.budget}; the cap doesn't apply)"
        elif re.fullmatch(r"ship", cmd, re.I):
            self.paused, self.turn, self.done = None, None, True
            self.outcome = self._ship_outcome("on an owner override")
            note = "ship: the run ends with the site as it is"
        elif m := re.fullmatch(r"drop\s+(.+)", cmd, re.I):
            want = m[1].strip()
            hit = [c for c in crit if want == str(self._no(c)) or (len(want) > 2 and want.lower() in c.lower())]
            if hit:
                g["done_when"] = [c for c in crit if c not in hit]
                g["dropped_by_requester"] = (g.get("dropped_by_requester") or []) + hit
                g["dropped_reasons"] = (g.get("dropped_reasons") or []) + [{"criterion": c, "reason": "Owner override"}
                                                                            for c in hit]
                self.save("goal.json", g)
                note = "dropped " + "; ".join(hit)
        if not note:
            self.post("system", "Owner override not applied: use budget +N, ship, or drop <item number or words>.",
                      override=True)
            return True
        self._event("owner", "override", note)
        self.post("system", f"Owner override applied: {note}.", override=True)
        if self.outcome and self.outcome.startswith(("Done", "Stopped")) and self.done and not self.paused:
            self._write_summary()
        return True

    def _run_safe(self):
        with self.lock:
            try:
                self._run()
            except Exception as e:  # noqa: BLE001
                self.error = f"{type(e).__name__}: {str(e)[:300]}"
                self.post("system", f"Stopped: {self.error}")
                self.outcome = f"Crashed: {self.error}"
                try:
                    self._write_summary()      # code only, before the slow reviewer call
                    self._make_report()
                except Exception as e2:  # noqa: BLE001
                    self.post("system", f"Run report failed: {type(e2).__name__}: {str(e2)[:200]}")
                self.turn = None
                self.status = "waiting"

    def _run(self):
        last_user = next(m["text"] for m in reversed(self.chat) if m["who"] == "user")
        self.started = self.started or time.time()
        self.outcome, p = None, self.paused
        if p:
            # The Requester answered Identity: resume where the run stopped. Answers don't renew the cycle allowance.
            self.paused = None
            self.cycle_turns = 0       # the answer is new information: without this the turn cap re-asked at once
            d = {"kind": p["kind"], "asked": [q["question"] for q in p["questions"]], "answer": last_user}
            self.decisions.append(d)
        if p:
            h = dict(p["handoff"])
            if h.get("why"):
                h["why"] += f" The Requester answered: {last_user}"
            self._handoff(p["to"], **h, answer=last_user, asked=p["questions"])
            self._plan_stages()
        else:
            if not self.goal:
                self.request = last_user
                self._plan_stages()
                self._handoff("identity", kind="intake")
            elif self.fixes:   # fix (b): a change request is checked at J2.a before Adaptation works on it
                self._handoff("identity", kind="intake", change=self.feedback, mark=len(self.intake))
            else:
                self._handoff("adaptation", kind="propose", new_cycle=True,
                              why=f"The Requester asked for a change: {self.feedback}")
            self.cycle_limit = self.cycle + MAX_CYCLES   # a new brief or change request gets a fresh allowance
        self.idle_cycles, self.cycle_registered = 0, True                    # a new message is new information
        while self.turn:
            # Safety limits never stop the run silently: Identity asks the Requester what to do (_identity_limit).
            if self.handoff.get("new_cycle"):
                self.idle_cycles = 0 if self.cycle_registered else self.idle_cycles + 1
                self.cycle_registered = False
                cheapest = min(RATES.values()) if RATES else 1
                limit = None
                if self.idle_cycles >= 2:
                    self._event("system", "stuck", f"cycles {self.cycle - 1} and {self.cycle} registered nothing")
                    limit = (f"Cycles {self.cycle - 1} and {self.cycle} both registered nothing, so another cycle would "
                             "repeat the same blocker.")
                elif self.cycle >= self.cycle_limit:
                    limit = f"The department used its {MAX_CYCLES} cycles (J2_MAX_CYCLES) and the site isn't workable yet."
                elif self.goal and self.budget - self.spent < cheapest:
                    self._event("system", "out_of_budget", f"{self.budget - self.spent} left, cheapest part {cheapest}")
                    limit = (f"Out of budget: {self.budget - self.spent} of {self.budget} work credits left, less than the "
                             f"cheapest rate-card part ({cheapest}), so no new workflow can be funded.")
                if limit:
                    self.idle_cycles, self.cycle_registered = 0, True
                    self._handoff("identity", kind="limit", reason=limit, out_of_budget=limit.startswith("Out of budget"))
                    continue
                self._begin_cycle()
            # J2.h after a passing Audit always runs: the cap never asks "keep going?" about a finished site.
            j2h = self.turn == "identity" and self.handoff.get("kind") == "audit" and self.audits \
                and not workable(self.audits[-1])
            if self.cycle_turns >= self._turn_cap() and self.handoff.get("kind") != "limit" and not j2h:
                self._handoff("identity", kind="limit", reason=f"Cycle {self.cycle} used {self._turn_cap()} turns (the "
                                                               "per-cycle cap) without finishing.")
                continue
            self.cycle_turns += 1
            self.agent_turns[self.turn] = self.agent_turns.get(self.turn, 0) + 1
            if is_wf_agent(self.turn):
                self._turn_workflow_agent(self.handoff)
            else:
                getattr(self, "_turn_" + self.turn)(self.handoff)
            if self.board_failed and self.turn != "adaptation":
                # A board change left a built atom failing its checks (board/1 part C): the existing atom_failed
                # path takes it, and its own ending hands the turn on.
                h = self.board_failed.pop(0)
                self.post("system", f"{', '.join(h['atoms'])} {h['owner_note']}. Adaptation reads the receipts and "
                                    "prices the fix.", warning=True)
                self._handoff("adaptation", kind="atom_failed", **h)
            if self.board_asks and self.turn != "identity":
                # A board answer asked for a PRD change (board/1 part C): Identity decides it as a change request,
                # on the same path as any question the charter can't answer, and the flow resumes after it.
                q = self.board_asks.pop(0)
                self._handoff("identity", kind="question", frm=q["frm"], question=q["question"],
                              resume={"to": self.turn, "handoff": self.handoff} if self.turn else None)
            if self.relay:
                # An agent addressed the Requester: Identity takes it first, then the flow resumes where it was going.
                r, self.relay = self.relay, None
                if self.turn != "identity":
                    resume = {"to": self.turn, "handoff": self.handoff} if self.turn else None
                    self._handoff("identity", kind="question", frm=r["from"], question=r.get("text") or "", resume=resume)
        self.turn, self.done = None, True
        if not self.paused:
            self.step = 8
        if self.doc and self.sites and not self.paused:
            self._export()
        self.post("system", (self.outcome or "Stopped.") + f" {self.out['noun'] if self.doc else 'Site'} "
                  f"v{len(self.sites)} is in the preview. "
                  + ("Answer Identity's question above to continue." if self.paused else
                     "Tell us if you want changes."))
        # summary.json first (code only): the reviewer call in _make_report can take minutes or hang, and the run
        # must have its summary the moment it ends.
        if not self.paused:
            self._write_summary()
        self._make_report(review=not self.paused)
        self.feedback = None
        self.status = "waiting"

    # ------------------------------------------------------------------ Identity
    def _turn_identity(self, h):
        return {"intake": self._identity_intake, "stage_decision": self._stage_decision, "resolve": self._identity_resolve, "approval": self._identity_approval,
                "audit": self._identity_audit, "requester_answered": self._identity_hears_requester,
                "ruling_budget": self._identity_ruling_budget, "limit": self._identity_limit,
                }.get(h.get("kind"), self._identity_question)(h)

    CHARTER_FMT = ('{"statement": str, "site_name": str, "audience": str, "standards": [str], "rules": [str], '
                   '"done_when": [str], "core": [bool] (one per done-when item, same order: true for the business '
                   'name and location, contact details, prices, hours, ordering and the no-errors check; these are '
                   'never left out for budget), "business_type": str (one or two words, e.g. "florist"), "visitor_needs": [str] '
                   '(what this business\'s visitors need, as confirmed), '
                   '"to_be_confirmed": [{"item": int (done-when item number), "fact": str}], '
                   '"dropped_from_brief": [str], "facts": [str] (every concrete fact the site must show: dates, '
                   'times, prices, ages, addresses, names, emails, phone numbers, counts; each copied exactly, '
                   'character for character, from the brief or the Requester\'s answers, never paraphrased or '
                   'invented)}')

    def _fact_problems(self, goal, change=None):
        """Charter facts that aren't a substring of the brief or the Requester's answers (case and spacing ignored):
        paraphrased or invented. A website charter only; a PRD keeps its own rules."""
        if self.doc:
            return []
        src = squash(" ".join([self.request, str(change or "")] + [x["answer"] for x in self.intake]
                              + [m["text"] for m in self.chat if m["who"] == "user"]))
        return [f for f in goal.get("facts") or [] if not str(f or "").strip() or squash(f) not in src]

    # ---- design documents: stages. PRD, HLD and LLD run one after another through the same flow (each its own output
    # type, charter, plan, build and audit); the site follows when it was asked for. Plain runs never get here.
    SITE_ESSENTIALS = ("styling", "page_section", "page_section", "page_section", "form")   # parts the site needs first

    def _plan_stages(self):
        """Once, before the charter: the documents the brief or the answers ask for by name (detect_documents). None
        asked: nothing changes. The first document's output type is picked here; a site follows only if it was named."""
        if self.stages or self.goal:
            return
        said = [self.request] + [x["answer"] for x in self.intake]
        docs = detect_documents(*said)
        deeds = j2_send.detect_deeds(*said)
        for k in deeds:      # the brief never chooses the recipient: say so once, in plain words
            if j2_send.names_other(k, *said) and k not in self.deed_seen:
                self.deed_seen.add(k)
                self.post("system", f"You named a different {'address' if k == 'send_email' else 'Slack channel or person'}. "
                          f"Test mode sends only to the configured test {'address' if k == 'send_email' else 'Slack target'}, "
                          "so nothing goes to the one you named.", warning=True)
        if deeds and not docs and "none" not in self.deed_seen:
            self.deed_seen.add("none")
            self.post("system", "A send needs a document to send. Ask for a PRD, HLD or LLD as well; nothing will be sent.",
                      warning=True)
        if not docs:
            return
        self.docset = docs
        self.stages = docs + deeds + (["website"] if wants_site(*said) else [])
        self.stage_i = 0
        self._apply_output(self.stages[0])
        self.site_tag = self.stages[0] + "-" if len(self.stages) > 1 else ""      # one document keeps its plain file names
        self.stage_checked = False       # the funding of the whole chain is checked once, when the first stage's plan starts
        names = " and ".join(f"the {k.upper()}" for k in docs) if len(docs) < 3 else "the PRD, the HLD and the LLD"
        self.post("system", f"Documents asked for: {names}" + (", then the website" if "website" in self.stages else "")
                  + (f", then {and_list(j2_send.DEEDS[k]['words'] for k in deeds)} of the last document" if deeds else "")
                  + ". They are written in that order, each from the one before.")

    @property
    def stage_kind(self):
        return self.stages[self.stage_i] if self.stages else None

    def _stage_artifacts(self):
        """The finished document stage: exported, its agent_output (what the next document reads), its record."""
        kind = self.stage_kind
        self._export()
        ao = doc_agent_output(kind, self.sites[-1])
        self.chain[kind] = ao
        self.docs[kind] = {"file": self.out["file"], "docx": self.out["file"][:-3] + ".docx", "agent_output": ao,
                           "calls": {"from": self.stage_calls0 + 1, "to": self.calls},
                           "workflows": [w["id"] for w in self.workflows], "atoms": [a["id"] for w in self.workflows
                                                                                      for a in w.get("atoms") or []]}
        self.save("documents.json", self.docs)

    def _next_stage(self):
        """A stage passed its audit. True when another starts (the caller returns), False after the last."""
        if not self.stages:
            return False
        if self.doc:
            self._stage_artifacts()
        base = self.goal or {}
        self.stage_archive.append({"stage": self.stage_kind, "goal": base, "audits": self.audits, "sections": self.sections})
        self.save(f"stage-{self.stage_kind}.json", self.stage_archive[-1])
        return self._advance(self.stage_i + 1)

    # ---- funding the chain. Documents the Requester asked for are not optional, and the LLD is the site's blueprint:
    # nothing is skipped by code. When the work credits left can't cover the stages still to do, the Requester gets the
    # decision question (add credits, leave out, finish as it is), with each document and the site as an item.
    def _stage_cost(self, kind):
        """The least a stage can cost on the rate card: every section a doc_section and one requirements_section, or the
        site's essentials (a scaffold, three sections, a form). 0 when the rate card is off."""
        if kind in j2_send.DEEDS:
            return RATES_DEED().get(kind, 0)
        if kind == "website":
            r, parts = RATES_WEB(), self.SITE_ESSENTIALS
            return sum(r.get(k, 0) for k in parts)
        r, n = RATES_DOC(), len(OUTPUTS[kind]["sections"])
        return r.get("doc_section", 0) * (n - 1) + r.get("requirements_section", 0)

    STAGE_WORDS = {"prd": "the product requirements", "hld": "the high-level design", "lld": "the low-level design",
                   "website": "the website", **{k: f"the {d['words']}" for k, d in j2_send.DEEDS.items()}}

    def _advance(self, i):
        """Start stage i, unless the credits left can't cover it and the rest (then the Requester is asked). False when
        no stage is left."""
        if i >= len(self.stages):
            return False
        self.stage_i = i - 1 if i else 0
        if self._stage_gate(i):
            return True
        self.stage_i = i
        self._begin_stage(self.stages[i], self.stage_archive[0]["goal"])
        return True

    def _stage_gate(self, i, started=False):
        """Ask the decision question when stages i.. cost more than the credits left. started: stage i is already
        running (its own funding is the usual flow's), so only the stages after it can be left out. True when asked."""
        left, todo = self.budget - self.spent, self.stages[i:]
        short, run, rows = [], left, []
        for n, k in enumerate(todo):
            c = self._stage_cost(k)
            ok = not short and (c <= run or (started and n == 0))     # in chain order: the first that doesn't fit blocks the rest
            run -= c if ok else 0
            rows.append((k, c, ok))
            if not ok:
                short.append(k)
        if not short:
            return False
        need = sum(c for _, c, _ in rows) - left
        room = BUDGET_CAP - self.budget
        add = min(max(need, 1), room) if room > 0 else 0
        opts, budget, choices = [], [], []
        if add:
            opts.append(f"Add {add} work credits{' (the most allowed)' if add < need else ''} and keep going: you get more "
                        f"of what you asked for built. Or type your own number, up to {room}")
            budget.append(add)
            choices.append("continue")
        names = and_list(self.STAGE_WORDS[k] for k in short)
        opts.append(f"Leave out {names}: you get the rest without {'it' if len(short) == 1 else 'them'}.")
        budget.append(0)
        choices.append("drop")
        if self.docs:
            opts.append("Finish as it is: you get what is already written, and the rest is not made")
            budget.append(0)
            choices.append("ship")
        ask = {"continue": "add more and keep going", "drop": "leave out what doesn't fit", "ship": "finish as it is"}
        ask = [ask[c] for c in choices]
        ask = ", ".join(ask[:-1]) + (", or " if len(ask) > 2 else " or ") + ask[-1] if len(ask) > 1 else ask[0]
        q = {"question": f"The work you approved doesn't cover everything you asked for. {ask[0].upper() + ask[1:]}?",
             "options": opts, **({"budget": budget} if any(budget) else {})}
        shown = [f"{self.STAGE_WORDS[k]} (needs at least {c} work credits{'' if ok else '; not covered'})" for k, c, ok in rows]
        self._event("system", "limit", f"stages {', '.join(short)} can't be funded: {left} left, {need} short")
        self.post("system", f"The work credits left ({left}) can't cover {names}. The Requester decides; nothing is skipped.",
                  warning=True)
        self._ask_user(f"You asked for {and_list(self.STAGE_WORDS[k] for k in todo)}, which needs at least "
                       f"{sum(c for _, c, _ in rows)} work credits; {left} are left.", [q], "limit", gaps=shown,
                       gaps_title="What you asked for",
                       handoff={"kind": "stage_decision", "stage_i": i, "started": started, "choices": choices,
                                "unfunded": short, "limit": "stages"})
        return True

    def _stage_decision(self, h):
        """The Requester's answer to the stage question: add credits (then the chain goes on), leave out the named
        stages, or finish as it is. Code applies it, typed or clicked, and says so."""
        self.step = 8
        answer, choices, i, short = h.get("answer") or "", h.get("choices") or [], h.get("stage_i", 0), h.get("unfunded") or []
        q = (h.get("asked") or [{}])[0]
        labels = [f"{self.STAGE_WORDS[k]} {k}" for k in self.stages[i:]]
        d = parse_decision(answer, q, choices, labels, list(range(1, len(labels) + 1)))
        decision = d.get("decision")
        if decision not in choices:
            decision = "continue" if "continue" in choices else "drop"
        add = d.get("add") or (q.get("budget")[choices.index("continue")] if decision == "continue" and q.get("budget") else None)
        named = [self.stages[i + n - 1] for n in d.get("drop_items") or [] if 1 <= n <= len(self.stages[i:])] or \
            (short if decision == "drop" else [])
        self.decisions[-1].update(decision=decision, **({"budget_added": add} if add else {}),
                                  **({"left_out": named} if decision == "drop" else {}))
        if decision == "ship":
            left = and_list(self.STAGE_WORDS[k] for k in self.stages[i:])
            self.post("identity", f"Got it. We'll finish with what is already written. Not made: {left}.", to_user=True,
                      next="user", done=True)
            self.outcome = f"Done with open items: finished as it is on the Requester's decision. Not made: {left}."
            self.turn = None
            return
        if decision == "drop":
            gone = [k for k in dict.fromkeys(named) if k in self.stages[i + (1 if h.get("started") else 0):]]
            self.stages = [k for k in self.stages if k not in gone]
            self.docset = [k for k in self.docset if k not in gone]
            self.post("identity", f"Got it. We'll leave out {and_list(self.STAGE_WORDS[k] for k in gone)}"
                      + (f" and add {add} work credits" if add else "") + ". Everything else stays as it is.", to_user=True,
                      next="user")
            self.post("system", f"Left out on the Requester's decision: {and_list(self.STAGE_WORDS[k] for k in gone)}.")
            self.decisions[-1]["left_out"] = gone
        if add:
            self._raise_budget(add, f"the Requester's answer: {answer}")
        if decision == "continue":
            self.post("identity", "Got it. We'll keep going." + (f" {add} work credits added." if add else ""), to_user=True,
                      next="user")
        if h.get("started"):                         # stage i is already running: its plan goes on
            self.stage_checked = True
            return self._handoff("adaptation", kind="propose", new_cycle=True, why=f"The Requester decided: {answer}")
        if not self._advance(i) and self.turn is None and not self.paused:
            self._stages_finished()

    def _stages_finished(self):
        """The last stage left in the chain is done (the rest were left out): the same closing as a normal Done."""
        self.post("identity", "Your design documents are ready." if len(self.docset) > 1 else
                  "Your document is ready." if self.docset else "Your website is ready.", done=True, to_user=True, next="user")
        self.outcome = ("Done: the design documents are workable. Every done-when item, rubric item and necessary condition "
                        "passes in each." if len(self.docset) > 1 else self.out["done"])
        self.turn = None

    def _begin_stage(self, kind, base):
        """Fresh working state for the next stage; the earlier stage's workflows and documents are kept for the record.
        A send deed (j2_send.DEEDS) is the one stage with no agents: code does it all in _deed_stage."""
        if kind in j2_send.DEEDS:
            return self._deed_stage(kind)
        self.past_workflows += self.workflows
        self.workflows, self.sites, self.sections, self.audits, self.style = [], [], {}, [], None
        self.goal, self.stamped, self.proposed_atoms, self.p_refused = None, False, {}, {}
        self.fail_floor, self.feedback, self.drop_answer, self.adapt_idle, self.last_limit = 0, None, None, None, None
        self.write_lock = self.atom = None
        self.site_tag = kind + "-"
        self.stage_calls0 = self.calls
        self._apply_output(kind)
        if kind == "website":
            self.post("system", "The documents are done. Now the website, built from them.")
            self.proposed_lld = {}
            if (self.chain.get("lld") or {}).get("sections"):       # only an LLD is a blueprint; an HLD alone is a summary
                self.blueprint = copy.deepcopy(self.chain["lld"])
                self.WF_FMT = self.WF_FMT.replace('"covers": [int]', '"lld_sections": [str] (the LLD section ids this '
                                                  'workflow implements, e.g. "S1"), "covers": [int]')
                self.GRANT_FMT = self.GRANT_FMT.replace('"covers": [int]', '"lld_sections": [str], "covers": [int]')
            return self._handoff("identity", kind="intake")
        self.goal = stage_goal(kind, self.stage_archive[0]["goal"], dict(self.chain))
        if self.intake:      # what the Requester told Identity before the first document (details, form fields) is in every stage's charter
            self.goal["requester_answers"] = [x["answer"] for x in self.intake]
        self.goal["charter_items"] = list(self.goal["done_when"])
        self.goal["core_items"] = list(self.goal["done_when"])
        self.goal.pop("core", None)
        self.save("goal.json", self.goal)
        self.stamped = True
        self.post("system", f"Next: the {kind.upper()}, written from the {'PRD' if kind == 'hld' else 'HLD'}'s agent_output.")
        self._handoff("adaptation", kind="propose", new_cycle=True,
                      why=f"The {self.stages[self.stage_i - 1].upper()} is done. Propose the first workflows for the {kind.upper()}.")

    def _deed_stage(self, kind):
        """A send deed after the documents: one workflow, one atom, no charter, plan, builder or audit. Code builds the
        message from the last finished document, sends it (dry-run or mock; recipient only from the environment),
        writes the receipt, and runs the atom's one check, which only reads that receipt. Then the chain goes on."""
        self.past_workflows += self.workflows
        self.workflows = []
        d, cost = j2_send.DEEDS[kind], RATES_DEED().get(kind, 0)
        wid = f"W-{len(self.past_workflows) + 1}"
        src = next(reversed(self.docs)) if self.docs else None
        art = self.docs[src] if src else None
        text = open(os.path.join(self.dir, art["file"]), encoding="utf-8").read() if art else ""
        aid = f"{wid}.a1"
        atom = {"id": aid, "goal": f"Send the {d['words']} about {art['file'] if art else 'the document'}",
                "how": {"agent": "code"}, "constraints": [f"only to the target in {d['var']}", "built by code, no model"],
                "context": [art["file"]] if art else [],
                "expected_output": [{"text": f"The {d['words']} reaches the configured test target", "kind": "send",
                                     "target": kind}],
                "verification": [{"text": f"The {d['words']} was sent to the configured test target", "type": "deed-send",
                                  "target": kind, "expect": True, "atom": aid, "item": 1}]}
        wf = {"id": wid, "proposal_id": None, "name": d["label"], "covers": [], "fixes": [], "covers_text": [],
              "purpose": f"The Requester asked for a {d['words']}", "cost": cost, "estimate": cost,
              "envelope": {"money_limit": cost, "model_call_limit": 0, "retry_limit": 0}, "charged": True,
              "calls_used": 0, "retries_used": 0, "outputs": None, "completion_check": "", "dependencies": [],
              "status": "running", "cycle": self.cycle, "agent": "Code", "section": None, "revises": None, "atoms": [atom]}
        self.workflows = [wf]
        self._recalc_spent()
        self.given[wid] = {"edges": [{"atom": x, "why": f"the message is built from {art['file']}"}
                                     for x in art["atoms"]] if art else [], "whole": False}
        rec = j2_send.send(kind, art["file"], text, art["file"]) if art else \
            {"deed": kind, "channel": d["channel"], "mode": j2_send.mode(), "sent": False, "result": "refused",
             "refused": "there is no finished document to send", "artifact": None}
        os.makedirs(self.dir, exist_ok=True)
        with open(os.path.join(self.dir, "receipts.jsonl"), "a", encoding="utf-8") as f:
            f.write(json.dumps({"atom": aid, "type": "deed-send-record", **rec, "t": time.strftime("%Y-%m-%dT%H:%M:%S"),
                                "by": "code"}, ensure_ascii=False) + "\n")
        self._run_checks(wf, 1, "", atoms=[atom])
        ok = atom["results"][0]["pass"]
        note = ("Not sent (dry run)." if ok is None else f"Not sent: {rec['refused'] or atom['results'][0]['evidence']}."
                if ok is False else f"The {d['words']} was sent in test mode to the mock recorder only.")
        atom["human_output"] = {"text": note}
        atom["deed"] = {k: rec.get(k) for k in ("channel", "mode", "sent")}
        wf["status"] = "failed" if ok is False else "done"
        self.save("workflows.json", self.workflows)
        self.deed_log.append({"kind": kind, "pass": ok, "note": note})
        self.post("system", f"{d['label']}: {note}", warning=ok is False)
        if self._advance(self.stage_i + 1):
            return
        self._deeds_finished()

    GOAL_TAGS = {"prd": "PRD", "hld": "HLD", "lld": "LLD", "website": "Website"}

    def _milestone(self):
        """A goal reached (a document filed and checked, or the site workable): one email about it, built and sent by code,
        when milestone emails are on (J2_MILESTONE_EMAILS=1, or the brief asks for updates by email). Dry-run or mock only,
        to J2_TEST_EMAIL only. It is a deed atom (verb do) whose parent is the goal it reports. An email never stops the
        run: when the credits can't cover it, it is skipped and the reason recorded."""
        said = [self.request] + [x["answer"] for x in self.intake]
        if os.getenv("J2_MILESTONE_EMAILS", "0") != "1" and not j2_send.asks_updates(*said):
            return
        kind = self.stage_kind or self.output
        tag = self.GOAL_TAGS.get(kind)
        if not tag or any(x["goal"] == tag for x in self.milestone_log):
            return
        cost = RATES_DEED().get("send_email", 0)
        if cost > self.budget - self.spent:
            why = f"it needs {cost} work credits and {max(self.budget - self.spent, 0)} are left"
            self.milestone_log.append({"goal": tag, "atom": None, "result": "skipped", "why": why})
            self._event("code", "milestone_skipped", f"{tag}: {why}")
            return self.post("system", f"Milestone email for {tag} skipped: {why}. The run goes on.", warning=True)
        base = (self.stage_archive[0]["goal"] if self.stage_archive else self.goal) or {}
        project = base.get("site_name") or base.get("product_name") or "Project"
        res = [r for w in self.workflows for a in w.get("atoms") or [] for r in a.get("results") or []]
        checks = (sum(1 for r in res if r.get("pass") is True), len(res))
        art = self.out["file"] if self.doc else "site/index.html"
        path = os.path.join(self.dir, art)
        text = open(path, encoding="utf-8").read() if os.path.isfile(path) else ""
        if self.doc:      # the document's first section: its first line of text
            sec = re.split(r"^##\s+.*$", text, flags=re.M)
            summary = next((ln.strip(" -*#") for ln in (sec[1] if len(sec) > 1 else text).splitlines() if ln.strip(" -*#")), "")
        else:
            summary = re.split(r"(?<=[.!?])\s", str(base.get("statement") or ""))[0]
        built = j2_send.milestone(project, tag, len(self.sites) or None, checks, summary)
        wid = f"W-{len(self.past_workflows) + len(self.workflows) + 1}"
        aid = f"{wid}.a1"
        atom = {"id": aid, "goal": f"Email: {tag} done", "how": {"agent": "code"}, "constraints": ["only to the target in J2_TEST_EMAIL", "built by code, no model"],
                "context": [art] if text else [],
                "expected_output": [{"text": f"The milestone email for {tag} reaches the configured test address", "kind": "send", "target": "milestone_email"}],
                "verification": [{"text": f"The milestone email for {tag} was sent to the configured test address", "type": "deed-send",
                                  "target": "milestone_email", "expect": True, "atom": aid, "item": 1}]}
        wf = {"id": wid, "proposal_id": None, "name": f"Email · {tag} done", "covers": [], "fixes": [], "covers_text": [],
              "purpose": f"A milestone email: {tag} done", "cost": cost, "estimate": cost,
              "envelope": {"money_limit": cost, "model_call_limit": 0, "retry_limit": 0}, "charged": True,
              "calls_used": 0, "retries_used": 0, "outputs": None, "completion_check": "", "dependencies": [],
              "status": "running", "cycle": self.cycle, "agent": "Code", "section": None, "revises": None, "atoms": [atom]}
        self.given[wid] = {"edges": [{"atom": a["id"], "why": f"the email reports {tag} done"} for w in self.workflows for a in w.get("atoms") or []], "whole": False}
        rec = j2_send.send("milestone_email", art, text, art if text else None, built=built)
        os.makedirs(self.dir, exist_ok=True)
        with open(os.path.join(self.dir, "receipts.jsonl"), "a", encoding="utf-8") as f:
            f.write(json.dumps({"atom": aid, "type": "deed-send-record", **rec, "t": time.strftime("%Y-%m-%dT%H:%M:%S"), "by": "code"}, ensure_ascii=False) + "\n")
        self._run_checks(wf, 1, "", atoms=[atom])
        ok = atom["results"][0]["pass"]
        note = "Not sent (dry run)." if ok is None else f"Not sent: {rec['refused'] or atom['results'][0]['evidence']}." if ok is False \
            else "Sent in test mode to the mock recorder only."
        atom["human_output"] = {"text": note}
        atom["deed"] = {k: rec.get(k) for k in ("channel", "mode", "sent")}
        wf["status"] = "failed" if ok is False else "done"
        self.past_workflows.append(wf)
        self._recalc_spent()
        self.save("workflows.json", self.workflows)
        self.milestone_log.append({"goal": tag, "atom": aid, "result": rec["result"], "why": rec.get("refused")})
        self.post("system", f"Milestone email · {tag} done: {note}", warning=ok is False)

    def _deeds_finished(self):
        """The last stage was a send: the closing message and outcome, with what each send did in plain words."""
        words = " ".join(x["note"] for x in self.deed_log)
        self.post("identity", f"{'Your design documents are ready.' if len(self.docset) > 1 else 'Your document is ready.'} {words}",
                  done=True, to_user=True, next="user")
        base = ("Done: the design documents are workable. Every done-when item, rubric item and necessary condition "
                "passes in each." if len(self.docset) > 1 else self.out["done"])
        self.outcome = f"Done with open items: {words}" if any(x["pass"] is False for x in self.deed_log) else f"{base} {words}"
        self.turn = None

    def _identity_intake(self, h):
        """J2.a: is there enough information to work? If not, Identity asks the Requester and the run waits; the
        answer comes back here. There is no round cap: Identity asks until it has enough to write the charter.
        A fact the Requester can't supply yet is recorded as to be confirmed; its done-when item is not met."""
        self.step = 1
        feedback = h.get("feedback") or ""
        if h.get("drop"):
            self.drop_answer = h.get("answer")
            feedback = (f"The Requester answered about the brief items the charter left out: {self.drop_answer}. Put "
                        "back what they want kept; list in dropped_from_brief only what they agreed to drop.")
        elif h.get("answer"):
            self.intake.append({"asked": [q["question"] for q in h.get("asked") or []], "answer": h["answer"]})
        change, mark = h.get("change"), h.get("mark", 0)      # J2_FIXES=1: a change request after done
        last = bool(h.get("force"))      # only after a reply with neither questions nor a charter
        if change:
            feedback += (f"\nTHE REQUESTER ASKED FOR A CHANGE AFTER THE GOAL WAS REACHED: {change}\nJ2.a for the change: "
                         "is there enough information to make it? If not, ask. If yes, return the updated charter: keep "
                         "everything already in it and add or adjust done-when items so the change is checkable. "
                         f"CURRENT CHARTER: {json.dumps(self.goal, ensure_ascii=False)}")

        def mock():
            if change:
                # J2.a for a change request: Identity asks the workflow that already built that part of the page
                # where it is, instead of assuming the change needs a new section (board, J2_BOARD=1).
                return json.dumps({"say": "@Adaptation Charter updated for the change request.", "enough": True,
                                   "next": "Adaptation", "goal": {**self.goal, "done_when": list(self.goal["done_when"])
                                                                  + [f"Change: {change}"]},
                                   **self._mock_board_ask("identity",
                                                          self._board_owner_atom("timings", "hours", "address",
                                                                                 "location"),
                                                          "where the store timings are",
                                                          "Are store timings already on the page, and where?")})
            text = " ".join([self.request] + [x["answer"] for x in self.intake])
            if self.doc and not self.intake and len(self.request.split()) < 25 and not last:
                return json.dumps({"say": "Before we start, a few quick questions so the document has the facts.",
                                   "enough": False, "next": "Requester", "questions": [
                    {"question": "What is the product called, who are its users, and what does each group do in "
                                 "it?", "options": []},
                    {"question": "What problem does it solve, and what are the must-have features?", "options": []},
                    {"question": "Any constraints, and how will you measure success? If it's a number, over what "
                                 "time?", "options": []}]})
            if self.doc:
                return json.dumps({"to_requester": "Thanks, I have everything I need. I'm starting the document now.",
                                   "to_agents": "@Adaptation Charter recorded for the PRD.", "enough": True,
                                   "next": "Adaptation", "goal": self.out["mock_charter"](text)})
            if not self.intake and len(self.request.split()) < 25 and not last:
                kind = guess_type(self.request)
                needs = (f"For a {kind} site, visitors usually need: {', '.join(NEEDS[kind]).lower()}. Is that right, "
                         "or what would you change?") if kind in NEEDS else "What must visitors be able to see or do on the site?"
                return json.dumps({"say": "Before we start, a few quick questions so we build the right thing.",
                                   "enough": False, "next": "Requester", "questions": [
                    {"question": "What is the business called, and where is it?", "options": []},
                    {"question": needs, "options": []},
                    {"question": "How should it feel, and must it work well on phones?", "options": []}]})
            return json.dumps({"to_requester": "Thanks, I have everything I need. I'm starting the build now.",
                               "to_agents": "@Adaptation Charter recorded from the brief"
                                      + (" and the Requester's answers." if self.intake else "; it had enough to work."),
                               "enough": True, "next": "Adaptation", "goal": mock_charter(text)})
        r = self.call("identity", '{"say": str (when you ask: to the Requester), ' + REPLY_FMT + ' (when enough is '
                      'true), "enough": bool, "questions": [{"question": str, "options": [str] '
                      '(0-4)}] (when enough is false), "goal": ' + self.out.get("charter_fmt", self.CHARTER_FMT)
                      + '|null (when enough is true), '
                      + next_field(("user", "adaptation"), "Requester when you ask; Adaptation once the charter is "
                                   "recorded") + "}",
                      f"REQUESTER'S BRIEF: {self.request}\nEARLIER QUESTIONS AND ANSWERS: {json.dumps(self.intake)}\n"
                      f"{feedback}\n" + (self.out["intake"] + (SITE_FIRST if "website" in self.stages and not change else "") if self.doc else
                      "J2.a: is there enough information to work? You need enough to write the charter: "
                      "goal, audience, standards, rules and a checkable done-when list. Never invent a missing fact.\n"
                      "- Work out what kind of business this is and what its visitors usually need (for example a "
                      "florist: browse bouquets, order, delivery areas; a photographer: portfolio, packages, booking; a "
                      "dental clinic: services, prices, appointments). Don't assume: unless the brief already says what "
                      "visitors must see or do, list what you think they need in one question and ask the Requester to "
                      "confirm or change it. Record business_type and the confirmed visitor_needs in the charter.\n"
                      "- Not enough: set enough=false and ask the Requester 1-4 specific questions about what the brief "
                      "leaves open (what the site is for, who it serves, what visitors must be able to see or do, hard "
                      "rules). Offer options only for a choice between known alternatives. Don't ask about cost.\n"
                      "- Enough: set enough=true and record the charter from the brief and the Requester's answers only. "
                      "List in facts every concrete fact from the brief and the answers: dates, times, prices, ages, "
                      "addresses, names, emails, phone numbers, counts, each copied exactly (code checks each one is in "
                      "the brief or the answers, and that the site shows it exactly). "
                      "Every done-when item must be checkable by reading the site file itself. Content facts only the "
                      "Requester knows (phone numbers, prices, addresses) are questions for the Requester: ask for them. "
                      "Only if the Requester has said they can't supply a fact yet, record it in to_be_confirmed with the "
                      "number of the done-when item it blocks: that item is not met until they confirm it. Done-when "
                      "may not refer to the Requester, Audit, verification, approval or launch. Keep everything the brief "
                      "asks for; if the charter leaves any of it out, list it in dropped_from_brief with the reason. "
                      "Don't plan costs or resources.") + (SITE_LATER if self.docset and not self.doc and not change else "")
                      + ("\nRecord the charter now (enough=true) from what you have." if last else ""), mock,
                      defaults=("Thanks, I have everything I need. I'm starting the build now.",
                                "Charter is ready. Propose the first workflows."))
        qs = clean_questions(r.get("questions"))
        if not r.get("enough") and not last and qs:
            return self._ask_user(r.get("say", ""), qs, "info", handoff={"kind": "intake", "change": change, "mark": mark})
        goal = r.get("goal") if isinstance(r.get("goal"), dict) else None
        if not goal:
            self._event("identity", "invalid_move", "J2.a reply had neither questions nor a charter")
            if not last:
                self.post("system", "Identity's reply had neither questions nor a charter. Asking it once more to "
                                    "record the charter.", redirect=True)
                return self._handoff("identity", kind="intake", force=True, change=change, mark=mark)
            self.outcome = "Stopped: Identity couldn't record a charter from the brief and the answers."
            self.turn = None
            return
        # The STATE shows each done-when item as {"item": n, "criterion": text}; a charter that copies that shape is read as the text
        # (run j2-20261023-102259: the objects reached a set and stopped the run). Everything after this sees strings.
        # The item number is a position in the charter, so the objects are put in item order first: item n stays checklist item n.
        dw = goal.get("done_when") or []
        if any(isinstance(c, dict) for c in dw):
            dw = sorted(dw, key=lambda c: as_int(c.get("item")) if isinstance(c, dict) and as_int(c.get("item")) > 0 else 10**6)
        goal["done_when"] = [str(c.get("criterion") or c.get("text") or "").strip() if isinstance(c, dict) else c for c in dw]
        goal["done_when"] = [c for c in goal["done_when"] if c != ""]
        bad = self._fact_problems(goal, change)
        if bad:
            why = (f"charter facts not found in the brief or the Requester's answers: {bad}; each fact must be copied "
                   "exactly from them, never paraphrased or invented")
            self.rejections.append({"cycle": self.cycle, "workflow": "charter", "reason": why})
            self._event("validator", "charter_facts", why)
            if not h.get("facts_reasked"):
                self.post("system", f"Validator rejected the charter: {why}. Identity sends it again.", rejected="charter")
                return self._handoff("identity", kind="intake", facts_reasked=True, change=change, mark=mark,
                                     force=last, feedback=feedback + "\nTHE VALIDATOR REJECTED YOUR CHARTER: " + why
                                     + ". Send the whole charter again with only exact facts.")
            goal["facts"] = [f for f in goal.get("facts") or [] if f not in bad]
            self.post("system", f"The charter still lists facts not in the brief or the answers; code drops them: "
                                f"{bad}.", warning=True)
        goal["to_be_confirmed"] = self._tbc(goal.get("to_be_confirmed"), goal.get("done_when") or [])
        goal.pop("open_inputs", None)
        if self.doc:      # code reads site_name (titles, core facts) and business_type (playbooks) for every output
            goal.setdefault("site_name", goal.get("product_name") or "Document")
            goal["business_type"] = self.output
            if self.output in ("hld", "lld"):      # a design document asked for without a PRD: its own sections and checks
                goal = stage_goal(self.output, goal, self.chain)
        # The no-errors gate has fixed wording: Identity can't extend it with features (runs/j2-20261009-133203
        # item 16 added filters and category browsing, so it couldn't pass once those were dropped).
        goal["done_when"] = [GATE_TEXT if NO_ERRORS_GATE.search(str(c)) else
                             (c.get("criterion") or c.get("text") or str(c)) if isinstance(c, dict) else c
                             for c in goal.get("done_when") or []]
        crit, core = goal["done_when"], goal.pop("core", None)
        # Every item keeps its number for the whole run: a new charter (a change request) keeps the earlier numbers
        # and numbers its new items after them. Audit may still use these numbers after items are dropped. But if no
        # work has been registered against the old numbers yet (the charter was rejected and replaced before
        # Adaptation proposed anything), there is nothing to keep stable: the replacement renumbers from 1.
        old = (self.goal or {}).get("charter_items") or [] if self.workflows or self.past_workflows else []
        goal["charter_items"] = old + [c for c in crit if c not in old]
        if isinstance(core, list) and len(core) == len(crit):          # stored by text: drops renumber the list
            goal["core_items"] = [c for c, k in zip(crit, core) if k is True]
        elif "core_items" in goal:                                       # a change request kept the old marks
            goal["core_items"] = [c for c in goal["core_items"] if c in crit]
        self.goal = goal
        self.save("goal.json", self.goal)
        ok, why = self._stamp_check()
        dropped = [str(d) for d in self.goal.get("dropped_from_brief") or [] if str(d).strip()]
        if ok and dropped and not self.drop_answer:
            return self._ask_user("The charter leaves out these items from your brief: " + "; ".join(dropped)
                                  + ". Work starts after your answer.",
                                  [{"question": "Keep them in scope, or drop them?",
                                    "options": ["Keep them in scope", "Drop them"]}],
                                  "charter_drop", handoff={"kind": "intake", "force": True, "drop": True})
        nxt, note = self._pick("identity", r.get("next"), ("adaptation",), "adaptation",
                               "after the charter, Adaptation proposes")
        self._reply(r, nxt, note, artifact="charter")
        start = dict(kind="propose", new_cycle=True, why=f"The Requester asked for a change: {change}. Identity checked it at "
                     "J2.a and updated the charter." if change else "The charter is recorded. Propose the first workflows.")
        if ok:
            self.stamped = True
            self._say("validator", f"Charter valid: statement present, {len(self.goal['done_when'])} checkable "
                                   "done-when items. Work can start."
                      + (f" Dropped with the Requester's agreement: {'; '.join(dropped)}." if dropped else ""),
                      "adaptation", stamp=True)
            return self._handoff("adaptation", **start)
        if h.get("attempt", 1) < 2:
            self._event("validator", "refused", f"charter: {why}")
            self._say("validator", f"Charter rejected: {why}.", "identity", stamp=False)
            return self._handoff("identity", kind="intake", force=True, attempt=2, change=change, mark=mark,
                                 feedback=f"The validator rejected the charter: {why}. Fix only that.")
        self.stamped = True
        self.post("system", "Charter accepted after two attempts; the validator's concern is on record.")
        self._handoff("adaptation", **start)

    @staticmethod
    def _tbc(raw, crit, keep=()):
        """To-be-confirmed facts keyed by the criterion text (item numbers shift when items are dropped)."""
        out = list(keep)
        for x in raw or []:
            if not isinstance(x, dict):
                continue
            n = as_int(x.get("item"))
            c = crit[n - 1] if 1 <= n <= len(crit) else str(x.get("criterion") or "")
            if c in crit and str(x.get("fact") or "").strip() and all(t["criterion"] != c for t in out):
                out.append({"criterion": c, "fact": str(x["fact"]).strip()})
        return out

    # Done-when must be checkable from the site itself. These words mark items that depend on someone's later
    # action or on the audit process.
    NOT_CHECKABLE = re.compile(r"\b(requester|audit\w*|launch\w*|go(?:es|ing)? live|verif\w*|approv\w*|sign(?:ed)?[- ]?off)\b",
                               re.I)

    def _stamp_check(self):
        g = self.goal or {}
        if not g.get("statement"):
            return False, "goal has no statement"
        if len(g.get("done_when") or []) < 3:
            return False, "fewer than 3 'done when' criteria"
        # Quoted text is a label the site shows (e.g. 'sample, replace before launch'), not a dependency.
        unquoted = lambda d: re.sub(r"'[^']*'|\"[^\"]*\"|‘[^’]*’|“[^”]*”", "", str(d))
        bad = [d for d in g["done_when"] if self.NOT_CHECKABLE.search(unquoted(d))]
        if bad:
            return False, ("these done-when items can't be checked by reading the site (they depend on the Requester, "
                           "Audit, verification, approval or launch); rephrase them as something the site shows, and "
                           "record a fact only the Requester can give in to_be_confirmed: " + json.dumps(bad))
        return True, "passes code checks"

    def _positions(self, c):
        """Each side's final position: for a proposal, its workflow list; otherwise its last attempt's position."""
        if c["topic"] == "proposal":
            return {"adaptation": c.get("adaptation_final") or [], "priority": c.get("priority_final") or []}
        last = {}
        for a in c["attempts"]:
            last[a["who"]] = a.get("position") or a["text"]
        return {w: last.get(w, "") for w in c["between"]}

    def _identity_resolve(self, h):
        """J2.d: rule on a conflict two agents couldn't settle, from the charter. If the charter doesn't decide it,
        ask the Requester one decision-ready question; the answer comes back here and Identity rules with it."""
        self.step = 4
        c = self._conflict(h.get("conflict"))
        a, b = c["between"]
        pos, answer, proposal = self._positions(c), h.get("answer"), c["topic"] == "proposal"
        if answer:
            c["user_answer"] = answer
        rulings = [a, b] + (["split"] if proposal else [])
        fmt = ('{' + REPLY_FMT + ', "clause": str, "facts": [str], "resolved": bool, "ruling": "' + '"|"'.join(rulings) + '" '
               '(whose final position stands), '
               + ('"workflows": [str] (for split: names of the workflows to register, taken as they are from either '
                  'final position), ' if proposal else '')
               + '"instruction": str (what happens next, one sentence)'
               + ('' if answer else ', "user_question": {"question": str, "options": [str] (2-4, each text with its '
                  'consequence and risk; no recommendation; never adding budget, leaving items out or finishing as '
                  'is: only code offers those)}|null') + "}")
        attempts = "\n".join(f"  attempt {x['n']} {NAMES[x['who']]}: {x['text']}" for x in c["attempts"])
        task = (f"CONFLICT {c['id']} ({TOPICS[c['topic']]}) between {NAMES[a]} and {NAMES[b]}: {c['subject']}\n"
                f"{c.get('detail') or ''}\nTHEIR ATTEMPTS:\n{attempts}\n"
                f"{NAMES[a].upper()}'S FINAL POSITION: {json.dumps(pos[a], ensure_ascii=False)}\n"
                f"{NAMES[b].upper()}'S FINAL POSITION: {json.dumps(pos[b], ensure_ascii=False)}\n")
        if answer:
            task += (f"You asked the Requester: {json.dumps(c.get('user_question'))}\nTHE REQUESTER ANSWERED: {answer}\n"
                     "Rule now, following their answer (resolved=true). Don't rewrite either position.")
        else:
            task += ("J2.d: resolve it using the charter. Name the clause and separate facts from assumptions. If the "
                     "charter decides it, set resolved=true and rule. If it doesn't (purpose, risk, an outside commitment "
                     "or a value trade-off the charter leaves open), set resolved=false and put ONE decision-ready "
                     "question for the Requester in user_question. Don't rewrite either position.")

        def mock():
            names = [w["name"] for w in pos.get("priority") or []] if proposal else []
            extra = [w["name"] for w in pos.get("adaptation") or [] if w["name"] not in names] if proposal else []
            split = {"ruling": "split", "workflows": names + extra}
            beans = next((d for d in (self.goal or {}).get("done_when", []) if re.search(r"bean|farm", d, re.I)), None)
            if not proposal:
                return json.dumps({"to_agents": f"Clause: done_when. {NAMES[b]}'s position stands.", "clause": "done_when",
                                   "resolved": True, "ruling": b, "instruction": "Carry on."})
            if answer:
                fund = re.search(r"\b1\b|fund|include|add a section", answer, re.I)
                return json.dumps({"to_agents": "The Requester chose. " + ("The stories section is funded as Adaptation "
                                   "proposed; Priority's other cuts stand." if fund else "Priority's position stands: "
                                   "no stories section, the reserve stays."), "clause": "Requester's decision",
                                   "resolved": True, "instruction": "Register and build.",
                                   **(split if fund else {"ruling": "priority"})})
            if beans:
                return json.dumps({"to_agents": f"Clause: done_when \"{beans}\". The stories cover that item, so they're in; "
                                          "Priority's cut to the animations stands.", "clause": f"done_when: {beans}",
                                   "facts": ["The charter names the beans and farms content"], "resolved": True,
                                   "instruction": "Register the split.", **split})
            return json.dumps({"to_agents": "", "to_requester": "@Requester One choice is up to you. A story section makes the site feel more "
                                      "personal, but it takes longer to build.", "clause": "done_when; standards",
                               "resolved": False, "ruling": None, "instruction": "", "user_question": {
                                   "question": "Do you want a section about your coffee's story?",
                                   "options": ["Add a section about your coffee's story (takes a bit longer)",
                                               "Keep it simple, no story section"]}})
        r = self.call("identity", fmt, task, mock)
        ruling = key(r.get("ruling"))
        c["identity"].append({"say": r.get("to_agents") or r.get("say"), "clause": r.get("clause"), "resolved": bool(r.get("resolved")),
                              "ruling": ruling or None, "after_user": bool(answer)})
        if r.get("resolved") and ruling in rulings:
            c["resolved_by"] = "identity_after_user" if answer else "identity"
            self._reply(r, None, conflict=c["id"], decision=True)
            return self._apply(c, ruling, r)
        if not answer:
            q = clean_questions([r.get("user_question")]) if isinstance(r.get("user_question"), dict) else []
            if not q or len(q[0]["options"]) < 2:
                self._event("identity", "invalid_move", f"{c['id']}: no decision-ready question; code offered both "
                                                        "final positions")
                q = [{"question": f"{NAMES[a]} and {NAMES[b]} disagree about {c['subject']}, and the charter doesn't "
                                  "decide it. Which position should stand?",
                      "options": [f"{NAMES[a]}'s: {brief_list(pos[a])}", f"{NAMES[b]}'s: {brief_list(pos[b])}"]}]
            c["user_question"], c["escalated_to_user"] = q[0], True
            self._save_conflicts()
            return self._ask_user(r.get("to_requester") or r.get("say", ""), q, "conflict", handoff={"kind": "resolve", "conflict": c["id"]},
                                  conflict=c["id"])
        self._event("identity", "invalid_move", f"{c['id']}: no ruling after the Requester's answer")
        self.post("system", f"Identity didn't rule on {c['id']} after the Requester answered. The default for "
                            f"{TOPICS[c['topic']]} applies.", redirect=True)
        c["resolved_by"] = "default"
        self._apply(c, None, r)

    def _identity_question(self, h):
        """A question for Identity: a policy question from Adaptation, or a message another agent addressed to the
        Requester. Identity answers from the charter, or asks the Requester when only they can decide."""
        self.step = 4
        frm = h.get("frm") or "adaptation"
        resume = h.get("resume") or {"to": "adaptation", "handoff": {"kind": "propose", "why": "Identity answered."}}
        r = self.call("identity", '{"say": str (address @' + NAMES[frm] + '), "clause": str, "answer": str|null, '
                      '"user_question": {"question": str, "options": [str] (0-4, no recommendation)}|null, '
                      + DECIDE_FMT + '}',
                      f"QUESTION FROM {NAMES[frm]}: {h.get('question')}\nAnswer it from the charter, naming the clause. "
                      "Only if the charter can't decide it and only the Requester can (a fact, an approval or a value "
                      "trade-off), put one decision-ready question in user_question instead. " + DECIDE_RULE,
                      lambda: json.dumps({"say": f"@{NAMES[frm]} The charter's rules cover this; proceed inside them.",
                                          "clause": "rules", "answer": "Proceed inside the charter's rules.",
                                          "user_question": None}))
        cheapest = min(RATES.values()) if RATES else 1
        if r.get("needs_decision") or (frm == "adaptation" and self.budget - self.spent < cheapest):
            # Budget, leaving items out and finishing are the code's decision question, never a free-form one.
            return self._identity_limit({"reason": f"{NAMES[frm]} asked: {h.get('question')}",
                                         "out_of_budget": self.budget - self.spent < cheapest,
                                         "budget_needed": as_int(r.get("budget_needed")) or None, "resume": resume})
        q = clean_questions([r.get("user_question")]) if isinstance(r.get("user_question"), dict) else []
        if q:
            return self._ask_user(r.get("say", ""), q, "question", to=resume["to"], handoff=resume["handoff"])
        self._say("identity", r.get("say") or r.get("answer") or "", frm if frm in ROUTES["identity"] else None)
        h2 = dict(resume["handoff"])
        if h2.get("kind") == "propose":
            h2["why"] = f"{h2.get('why') or ''} Identity answered {NAMES[frm]}: {r.get('answer') or r.get('say')}"
        self._handoff(resume["to"], **h2)

    def _identity_approval(self, h):
        """Priority requested approval for a run over its envelope. Identity may not raise an envelope, so it denies
        from the charter or asks the Requester; a yes raises the budget (capped) and the envelope check runs again."""
        self.step = 6
        w, answer = self._wf(h.get("wid")), h.get("answer")
        need = max(0, w["cost"] - (self.budget - self.spent))
        if answer:
            r = self.call("identity", '{' + REPLY_FMT + ', "approved": bool, "budget_increase": int|null}',
                          f"YOU ASKED THE REQUESTER about running {w['id']} {w['name']} ({w['cost']} work credits, {need} more "
                          f"than is left). THE REQUESTER ANSWERED: {answer}\nRecord their decision. Add budget only if "
                          "they approved it.",
                          lambda: json.dumps({"say": "The Requester approved it.", "approved": True,
                                              "budget_increase": need}) if re.search(r"\b1\b|yes|approve", answer, re.I)
                          else json.dumps({"say": "The Requester declined; it stays parked.", "approved": False}))
            self._reply(r, "priority" if r.get("approved") else "coordination")
            if r.get("approved"):
                self._raise_budget(as_int(r.get("budget_increase")) or need, f"{w['id']} over its envelope")
                return self._handoff("priority", kind="envelope", wid=w["id"])
        else:
            r = self.call("identity", '{"say": str, "clause": str, "decision": "DENY"|"ASK", "user_question": '
                          '{"question": str, "options": [str]}|null}',
                          f"Priority requests approval to run {w['id']} {w['name']} over its envelope: "
                          f"{'; '.join(h.get('over') or [])}. You may not raise an envelope or the budget yourself. DENY "
                          "if a charter clause decides it; otherwise ASK the Requester one decision-ready question.",
                          lambda: json.dumps({"say": "Adding budget is the Requester's call.", "clause": "authority",
                                              "decision": "ASK", "user_question": None}))
            if str(r.get("decision")).upper() == "ASK":
                q = clean_questions([r.get("user_question")]) if isinstance(r.get("user_question"), dict) else []
                if not q or len(q[0]["options"]) < 2:
                    q = [{"question": f"Run {w['name']} over the budget?",
                          "options": [f"Approve: add {need} work credits so it runs", "Don't run it; it stays parked"]}]
                return self._ask_user(r.get("say", ""), q, "approval", handoff={"kind": "approval", "wid": w["id"]})
            self._say("identity", r.get("say", ""), "coordination")
        w["status"] = "unfunded"
        self.save("workflows.json", self.workflows)
        self._handoff("coordination", kind="schedule")

    def _findings(self, audit=None):
        """Audit's full findings for Identity: every id, verdict, severity and text, done-when items and extra issues."""
        a = audit or (self.audits[-1] if self.audits else None)
        keys = ("id", "item", "criterion", "verdict", "severity", "confidence", "note", "needs_requester_fact")
        return [{**{k: r.get(k) for k in keys if r.get(k) is not None},
                 **({"item": self._no(r["criterion"])} if r in a["results"] else {})}
                for r in (a["results"] + a.get("rubric", []) + a.get("extra_issues", []) if a else [])]

    def _identity_hears_requester(self, h):
        """The Requester answered Identity at J2.g/J2.h (a fact, a limit, or another question Identity asked).
        Facts they supply are recorded; facts they can't supply yet become to be confirmed (that item isn't met).
        After the decision question (_identity_limit): code applies what the answer asks, typed or clicked: add
        budget, leave items out (both at once is fine), keep working, or finish as is."""
        self.step = 8
        limit, choices, answer = h.get("limit"), h.get("choices") or [], h.get("answer") or ""
        q = (h.get("asked") or [{}])[0] if limit else {}
        g = self.goal or {}
        tbc = g.get("to_be_confirmed") or []
        crit = g.get("done_when") or []
        nums, at = [self._no(c) for c in crit], self._item_at()     # item numbers are run-long, never positions

        def mock():
            ship = re.search(r"\bship\b|as is", answer, re.I)
            drop = re.search(r"\bdrop\b", answer, re.I)
            unknown = re.search(r"don'?t know|not yet|later|no idea|can'?t", answer, re.I)
            asked_items = [as_int(n) for n in re.findall(r"item (\d+)", json.dumps(h.get("asked")))]
            items = asked_items or [self._no(t["criterion"]) for t in tbc if t["criterion"] in crit]
            facts = [{"item": n, "fact": answer} for n in items]
            d = parse_decision(answer, q, choices, crit, nums) if limit else {}
            return json.dumps({**d, "to_requester": "Thanks, got it." if not unknown else "Thanks. We'll mark that as "
                               "still to be confirmed.", "to_agents": "Recorded." if not unknown else "Recorded as to "
                               "be confirmed.",
                               "clause": "done_when", **({} if limit else {"decision": "ship" if ship else "drop" if
                                                                             drop else "continue"}),
                               "confirmed": [] if unknown or ship or drop or d.get("add") or d.get("drop_items")
                               or d.get("decision") in ("ship", "drop", "stop") else facts,
                               "still_to_be_confirmed": facts if unknown else [],
                               "instruction_to_adaptation": f"Put the Requester's answer in: {answer}",
                               "next": "Adaptation"})
        r = self.call("identity", '{' + REPLY_FMT + ', "clause": str, ' + DECISION_FMT + ', "confirmed": '
                      '[{"item": int, "fact": str}] (facts the Requester supplied, by done-when item), '
                      '"still_to_be_confirmed": [{"item": int, "fact": str}] (facts they said they can\'t supply yet), '
                      '"instruction_to_adaptation": str|null, '
                      + next_field(("audit", "adaptation"), "Audit to re-check the current site when the site needs no "
                                   "change; Adaptation when the site must change (e.g. real details to put in)") + "}",
                      f"YOU ASKED THE REQUESTER: {json.dumps(h.get('asked'), ensure_ascii=False)}\nTHE REQUESTER "
                      f"ANSWERED: {answer}\nDONE-WHEN (by item number; numbers never change, F-n is item n): "
                      f"{json.dumps(dict(zip(nums, crit)), ensure_ascii=False)}\nTO BE CONFIRMED NOW: "
                      f"{json.dumps(tbc, ensure_ascii=False)}\nAUDIT FINDINGS: {json.dumps(self._findings(), ensure_ascii=False)}"
                      "\nRecord what the answer settles against the charter, without changing done-when. A fact they "
                      "gave goes in confirmed; a fact they say they can't give yet goes in still_to_be_confirmed (that "
                      "done-when item is not met until they confirm it)."
                      + (f"\nThis answers the decision question ({limit}). Read what they asked for: decision "
                         "continue (keep working), drop (leave out the failing items) or ship (finish as it is); add: "
                         "work credits they asked to add; drop_items: the item numbers they named to leave out. Adding and "
                         "leaving out can come together. Code applies it. The Requester's decision is final: you may "
                         "say what won't fit in the budget, never refuse it." if limit else ""), mock)
        for k in ("confirmed", "still_to_be_confirmed"):          # item numbers -> current positions
            r[k] = [{**x, "item": at.get(as_int(x.get("item")), 0)} for x in r.get(k) or [] if isinstance(x, dict)]
        decision, add, drops = "continue", None, []
        named = [t for t in h.get("drop_texts") or [] if t in crit]   # what the "leave out" option named
        if limit:
            if ONLY_PICK.fullmatch(answer):          # a click or a bare option number: exactly that option
                r.update(parse_decision(answer, q, choices, crit, nums))
            decision = str(r.get("decision") or "").lower() or None
            add = max(0, as_int(r.get("add"))) or None
            drops = sorted({at[as_int(n)] for n in r.get("drop_items") or [] if as_int(n) in at})
            if decision == "continue" and not add and q.get("budget") and "continue" in choices:
                add = q["budget"][choices.index("continue")] or None       # the option's own amount
            if decision == "drop" and not drops:
                # the option itself: exactly the items it named (older pauses carry positions only)
                drops = [crit.index(t) + 1 for t in named] if "drop_texts" in h else \
                    [i for i in h.get("failing_items") or [] if 1 <= i <= len(crit)]
            if decision == "continue" and "continue" not in choices and drops:
                decision = "drop"                     # keep working wasn't offered (nothing to fund it): leave out
            if decision not in choices and not h.get("reasked"):
                if self.decisions:
                    self.decisions[-1]["unclear"] = True
                self._event("identity", "unclear_answer", f"answer to '{q.get('question')}' names no option")
                n = len(q.get("options") or [])
                nums = ", ".join(str(i) for i in range(1, n)) + f" or {n}" if n > 1 else "1"
                self._ask_user(f"I couldn't tell which option you chose. Reply {nums}, or type what to add or leave "
                               "out.", [q], "limit", auto=True, options_only=True,
                               handoff={**{k: h[k] for k in ("kind", "limit", "choices", "failing_items", "resume",
                                                             "drop_texts", "drop_named") if k in h},
                                        "reasked": True})
                return
            if decision not in choices:
                decision = "continue" if "continue" in choices else "ship"
        confirmed = self._tbc(r.get("confirmed"), crit)
        done_facts = {t["criterion"] for t in confirmed}
        g["confirmed_facts"] = (g.get("confirmed_facts") or []) + confirmed
        g["to_be_confirmed"] = self._tbc(r.get("still_to_be_confirmed"), crit,
                                         keep=[t for t in tbc if t["criterion"] not in done_facts])
        g["to_be_confirmed"] = [t for t in g["to_be_confirmed"] if t["criterion"] not in done_facts]
        self.save("goal.json", g)
        self.fail_floor = len(self.audits)       # the Requester's answer is new information
        if self.decisions:
            self.decisions[-1].update(decision=decision, **({"budget_added": add} if add else {}),
                                      **({"dropped_items": drops} if drops else {}))
        if limit and self.last_limit:
            self.last_limit["decision"] = decision
        gone = [crit[i - 1] for i in drops]
        label = {**self._names(crit), **dict(zip(h.get("drop_texts") or [], h.get("drop_named") or []))}
        names = and_list(dict.fromkeys(label[c] for c in gone))   # the option's own names (a section once)
        if limit and decision == "drop" and ONLY_PICK.fullmatch(answer) and "drop_texts" in h:
            # The check: the items the option named are exactly the items dropped.
            if sorted(gone) != sorted(named):
                self._event("system", "drop_mismatch", f"named {named}, dropping {gone}")
                gone, names = list(named), and_list(dict.fromkeys(label[c] for c in named))
        if self.decisions and limit:
            self.decisions[-1].update(named=h.get("drop_named") if decision == "drop" and "drop_texts" in h else [],
                                      dropped=[label[c] for c in gone])
        if gone:
            # Code writes what was left out: Identity's own words could name different items than code drops.
            r["to_requester"] = (f"Got it. We'll leave out {names}" + (f" and add {add} work credits" if add and
                                                                       decision != "ship" else "")
                                 + ". Everything else stays as it is.")
            r["to_agents"] = (f"The Requester left out {names}; {'it is' if len(gone) == 1 else 'they are'} off the "
                              "checklist now. Work on what's left.")
            g["done_when"] = [c for c in crit if c not in gone]
            g["dropped_by_requester"] = (g.get("dropped_by_requester") or []) + gone
            g["dropped_reasons"] = (g.get("dropped_reasons") or []) + [
                {"criterion": c, "reason": f"The Requester chose to leave it out ({limit}): {answer}"} for c in gone]
            g["to_be_confirmed"] = [t for t in g["to_be_confirmed"] if t["criterion"] not in gone]
            self.save("goal.json", g)
            self.post("system", f"Left out on the Requester's decision: {names}.")
        if add and decision != "ship":
            self._raise_budget(add, f"the Requester's answer: {answer}")
        if decision == "ship" and limit and self.doc and self.stages and self.stage_i + 1 < len(self.stages):
            # A document in a chain: accept it as it is (open items recorded) and go on with the next stage.
            fails = workable(self.audits[-1]) if self.audits else []
            self.open_docs.append(f"{self.STAGE_WORDS[self.stage_kind]}: {'; '.join(fails)}" if fails else
                                  self.STAGE_WORDS[self.stage_kind])
            self.post("identity", f"Got it. We'll keep {self.STAGE_WORDS[self.stage_kind]} as it is and go on with "
                      f"{and_list(self.STAGE_WORDS[k] for k in self.stages[self.stage_i + 1:])}.", to_user=True, next="user")
            self.post("system", f"Accepted as it is on the Requester's decision: {self.STAGE_WORDS[self.stage_kind]}.")
            if self._next_stage():
                return
        if decision in ("ship", "stop"):
            rest = and_list(self.STAGE_WORDS[k] for k in self.stages[self.stage_i + 1:]) if decision == "stop" else ""
            self.post("identity", r.get("to_requester") or ("We'll finish the site as it is now." if not self.doc else
                                                            f"We'll finish with what is already written."
                                                            + (f" Not made: {rest}." if rest else "")), to_user=True,
                      next="user", done=True)
            self.outcome = self._ship_outcome(f"on the Requester's decision ({limit or 'their answer'})")
            if rest:
                self.outcome += f" Not made: {rest}."
            self.turn = None
            return
        if h.get("after_ruling") and (gone or decision == "drop"):
            # Left out after Priority's ruling: added credits fund a new proposal, else the funded work is built next.
            if add:
                self._reply(r, "adaptation")
                return self._handoff("adaptation", kind="propose", why=f"The Requester decided: {answer}")
            if self._waiting_work():
                self._reply(r, "coordination")
                return self._handoff("coordination", kind="schedule")
            self._reply(r, "adaptation")
            return self._handoff("adaptation", kind="propose", new_cycle=True,
                                 why=f"The Requester left items out after Priority's ruling: {answer}")
        if gone or decision == "drop":
            # Audit re-checks against the new done-when first: every later step (Adaptation's plan, the next
            # question) works from a fresh audit, never one that still lists the dropped items.
            if decision == "continue":
                self.cycle_limit = self.cycle + MAX_CYCLES
            self._reply(r, "audit")
            self.post("system", "Audit re-checks the current site against the new done-when.")
            return self._handoff("audit", kind="audit")
        if limit and h.get("resume"):
            # The decision answered an agent's question mid-cycle (Adaptation's policy question): it goes on there.
            self._reply(r, h["resume"]["to"])
            h2 = dict(h["resume"]["handoff"])
            h2["why"] = f"{h2.get('why') or ''} The Requester decided: {answer}"
            return self._handoff(h["resume"]["to"], **h2)
        if limit:
            self.cycle_limit = self.cycle + MAX_CYCLES      # keep working: a fresh cycle allowance
            nxt, note = "adaptation", None
        else:
            nxt, note = self._pick("identity", r.get("next"), ("audit", "adaptation"), "adaptation")
        self._reply(r, nxt, note)
        if confirmed:
            self.post("system", "Confirmed by the Requester: " + "; ".join(f"{t['criterion']} = {t['fact']}"
                                                                          for t in confirmed))
        if nxt == "audit":
            return self._handoff("audit", kind="audit")
        self._handoff("adaptation", kind="propose", new_cycle=bool(limit),
                      why=f"The Requester answered Identity: {answer}. Instruction: {r.get('instruction_to_adaptation')}"
                          + (f" Facts to put in: {json.dumps(confirmed, ensure_ascii=False)}" if confirmed else ""))

    def _no(self, text):
        """A checklist item's number for the whole run (charter_items): never renumbered after a drop; an item added
        later gets the next number."""
        g = self.goal if self.goal is not None else {}
        orig = g.setdefault("charter_items", list(g.get("done_when") or []))
        if text not in orig:
            orig.append(text)
        return orig.index(text) + 1

    def _requester_noun(self, text):
        """A document run's Requester reads "document", never "site" or "website"."""
        if not self.doc or "website" in self.stages:      # a website was asked for as well: it is named as one
            return text
        return re.sub(r"\b(?:web)?site(s?)\b", lambda m: ("Document" if m.group(0)[0].isupper() else "document")
                      + m.group(1), str(text), flags=re.I)

    def _fact_question(self, x):
        """Identity's question for a fact only the Requester knows (an Audit finding's needs_requester_fact), in the
        Requester's words: the item's charter number and short text, never an internal id or agent-facing wording."""
        crit, src = (self.goal or {}).get("done_when") or [], self._fact_source()
        n = self._no_at(x.get("item"))
        text = crit[as_int(x.get("item")) - 1] if n else None
        ask = requester_words(x["needs_requester_fact"], src).rstrip(" .") or \
            (f"The detail needed for {short(text)}" if text else "The detail this needs")
        if self.doc and NUMBER_ASK.search(ask) and not re.search(r"over what time|window", ask, re.I):
            ask += ", and over what time it is measured (for example per week or per term)"   # one question, P8
        return {"question": ask + (f" (checklist item {n}: {short(text)})" if n else "")
                + ". If you can't give it yet, say so and it's marked to be confirmed.", "options": []}

    def _no_at(self, i):
        """The run-long number of the item at current position i (None when there is no such item)."""
        crit = (self.goal or {}).get("done_when") or []
        return self._no(crit[as_int(i) - 1]) if 1 <= as_int(i) <= len(crit) else None

    def _item_at(self):
        """Run-long number -> current position in done_when."""
        return {self._no(c): i for i, c in enumerate((self.goal or {}).get("done_when") or [], 1)}

    def _item_status(self, text, gap):
        """Audit's current status of an unfinished checklist item, in plain words."""
        if gap_kind(gap) == "built but not checked yet":
            return "built, not checked yet"
        crit = (self.goal or {}).get("done_when") or []
        st = {w["status"] for w in self.workflows if text in [crit[i - 1] for i in self._wf_items(w)]}
        waiting = bool(st & {"registered", "held", "running", "failed"})     # funded work not built (yet)
        if gap_kind(gap) == "not built yet":                                 # an unfinished workflow covers it
            return "not built" if waiting else "unfunded"
        return "built but failing" if "done" in st else "not built" if waiting else "unfunded"

    def _names(self, texts):
        """{item text: "item n short name"} (n: its charter number) with short names unique among these texts. The
        Requester reads these, so never an internal id (F-n)."""
        texts = list(dict.fromkeys(texts))
        return {t: f"item {self._no(t)} {x}" for t, x in zip(texts, shorts(texts))}

    def _gap_rows(self, gaps, last, left):
        """The unfinished list as rows: name ("item n" with the item's run-long number and its short text, or a rubric
        item's or condition's short text, no id), Audit's status, and whether "leave out" drops it. The question, its options and the drop that code
        applies are all made from these rows, so the message and the outcome can't disagree."""
        crit = (self.goal or {}).get("done_when") or []
        rubric = {r["criterion"]: r["id"] for r in (last or {}).get("rubric", [])}
        names = self._names(t for g in gaps for t in [re.sub(r" \((?:not built yet|not checked yet)\)$", "", g)]
                            if t in crit)
        rows = []
        for g in gaps:
            t = re.sub(r" \((?:not built yet|not checked yet)\)$", "", g)
            if t in crit:
                st = self._item_status(t, g)
                rows.append({"name": names[t], "text": t, "status": st,
                             "drop": st != "built, not checked yet" and not (left and must_keep(t, self.goal))})
            else:
                rows.append({"name": short(t) if t in rubric else t, "text": t, "status": "failing",
                             "drop": False})
        if self.doc:      # a document: one row per section, named by the section; its items leave out together
            by = {}
            for r in rows:
                sec = next((x for x in self.out["sections"] if x.lower() in r["text"].lower()), None)
                k = sec or r["text"]
                if k in by:
                    by[k]["texts"].append(r["text"])
                    by[k]["drop"] = by[k]["drop"] and r["drop"]
                else:
                    by[k] = {**r, "name": f"the {sec} section" if sec else r["name"], "texts": [r["text"]]}
            rows = list(by.values())
        return rows

    def _identity_limit(self, h):
        """A safety limit was hit (budget, cycles, turns, the same failure twice in a row) or only the Requester can
        close what's left. Identity asks: keep working (adding budget if it ran out), drop the failing items, or ship
        as is. Options are fixed by code; Identity writes the message. The answer goes to _identity_hears_requester."""
        self.step = 8
        reason, last = h.get("reason") or "", self.audits[-1] if self.audits else None
        fails = workable(last)
        items = [r for r in (last or {}).get("results", []) if not r["pass"]]
        left = h.get("left_out")          # after Priority's ruling: the done-when items nothing will build
        gaps = [f"{c} (not built yet)" for c in left] if left else self._open_list(last, items)
        cheapest = min(RATES.values()) if RATES else 1
        room = BUDGET_CAP - self.budget
        short = self.budget - self.spent < cheapest       # whatever the limit was, nothing new can be funded
        want = max(as_int(h.get("budget_needed")) or LIMIT_ADD, cheapest)
        add = min(want, room) if short or h.get("budget_needed") else 0
        # Keep working only when it can do something: with no money left it needs a raise the cap allows. And after
        # "keep working" changed nothing (same site, audits, spend), the same option isn't offered again.
        state = (len(self.sites), len(self.workflows), self.spent)   # an audit alone changes nothing
        stale = self.last_limit and self.last_limit["state"] == state and self.last_limit["decision"] == "continue"
        opts, budget, choices = [], [], []
        dw = (self.STAGE_WORDS.get(self.stage_kind) or f"the {self.out['noun']}") if self.doc else None   # a document in plain words
        later = self.stages[self.stage_i + 1:] if self.doc and self.stages else []
        if (not short or self.budget + add - self.spent >= cheapest) and not stale and (add or not h.get("budget_needed")):
            more = f"the team works on what is not right in {dw}" if dw else "you get more of the missing parts built"
            opts.append((f"Add {add} work credits{' (the most allowed)' if add < want else ''} and keep going: {more}. "
                         f"Or type your own number, up to {room}" if add else
                         f"Keep going: {more}")
                        + (" (send the missing details with your answer)" if h.get("tbc") else ""))
            budget.append(add)
            choices.append("continue")
        crit = (self.goal or {}).get("done_when") or []
        rows = self._gap_rows(gaps, last, left)
        named = [r for r in rows if r["drop"]]
        if self.doc:      # never an option that leaves the document with one section or none
            gone = {t for r in named for t in r.get("texts", [r["text"]])}
            kept = {x for x in self.out["sections"] for c in crit if c not in gone and x.lower() in c.lower()}
            named = named if len(kept) > 1 else []
        drop = [crit.index(t) + 1 for r in named for t in r.get("texts", [r["text"]])]
        shown = [f"{r['name']} ({r['status']})" for r in rows]      # the list under Identity's words
        if named:
            opts.append(f"Leave out {and_list(r['name'] for r in named)}: you get {dw or 'the site'} without "
                        f"{'it' if len(named) == 1 else 'them'}."
                        + "".join(f" {r['name']} stays as it is, {STILL[r['status']]}." for r in rows if not r["drop"]))
            budget.append(0)
            choices.append("drop")
        if self.sites:                    # finishing needs a site to finish
            opts.append(f"Finish as it is: you get {dw or 'the site'} as it is now" + (f", with {len(gaps)} thing"
                        f"{'s' if len(gaps) > 1 else ''} still not right: {and_list(r['name'] for r in rows)}"
                        if gaps else "") + (f". Then the team goes on with {and_list(self.STAGE_WORDS[k] for k in later)}"
                                            if later else ""))
            budget.append(0)
            choices.append("ship")
            if later:         # a chain: "finish as it is" accepts this document; stopping the whole run is its own option
                opts.append(f"Stop the whole run here: you get what is already written, and "
                            f"{and_list(self.STAGE_WORDS[k] for k in later)} {'is' if len(later) == 1 else 'are'} not made")
                budget.append(0)
                choices.append("stop")
        if not self.sites and len(choices) < 2:
            opts.append("Stop here: nothing is built yet, so you get no site")
            budget.append(0)
            choices.append("ship")
        lead = ("The last round produced no new work, so keeping going as before would change nothing." if stale else
                "Some parts were put off twice." if h.get("deferred") else
                "Some parts don't fit in the work you approved." if left else"The team paused after a round of work." if re.match(r"Cycle \d+ used|The department used", reason) else
                "The work you approved is used up." if short else
                "The same problem came back twice." if reason.startswith("The same failure") else
                "Only you can give what's still missing." if reason.startswith("Only you") else
                "The team needs your decision.")
        say_choices = {"continue": "add more and keep going" if add else "keep going",
                       "drop": "leave out the missing parts" if not dw else "leave out what is not right",
                       "ship": "finish as it is" if self.sites else "stop here", "stop": "stop the whole run here"}
        ask = [say_choices[c] for c in choices]
        ask = ", ".join(ask[:-1]) + (", or " if len(ask) > 2 else " or ") + ask[-1] if len(ask) > 1 else ask[0]
        q = {"question": f"{lead} {ask[0].upper() + ask[1:]}?", "options": opts,
             **({"budget": budget} if any(budget) else {})}
        self._event("system", "limit", reason)
        plain_missing = f"{dw[0].upper() + dw[1:]} isn't finished yet." if dw else "The site isn't finished yet."
        passed, funded = self._built_and_funded(last)
        r = self.call("identity", '{"say": str (to the Requester, in plain words: what has been done so far, at most 2 '
                      'short sentences; don\'t list what is missing, code adds that list; no question, it is asked '
                      'below; no recommendation; never the words "building time", "cap", "cycles", "done-when" or '
                      '"estimate")}',
                      f"PASSED BY AUDIT (the only things you may describe as built or on the site): "
                      f"{json.dumps(passed, ensure_ascii=False)}\nFUNDED, NOT BUILT OR NOT CHECKED YET (say the team "
                      f"will build these, never that the site has them): {json.dumps(funded, ensure_ascii=False)}\n"
                      f"LIMIT: {reason}\nWHAT IS STILL NOT RIGHT (code lists it after your words): "
                      f"{json.dumps(shown, ensure_ascii=False)}\nTHE QUESTION AND OPTIONS CODE WILL OFFER: "
                      f"{json.dumps(q, ensure_ascii=False)}\nTell the Requester.",
                      lambda: json.dumps({"say": plain_missing}))
        say = r.get("say") or plain_missing
        if self.doc:      # a document: the check results are written by code, whatever Identity said
            say = check_summary(last, dw)
        if not passed and BUILT_CLAIM.search(say):
            # Nothing has passed Audit, so nothing may be described as built (runs/j2-20261009-133203 msg 29).
            self._event("identity", "invalid_move", "described unbuilt work as built")
            say = (f"Nothing has passed the checks yet; the team will build {len(funded)} funded part"
                   f"{'s' if len(funded) != 1 else ''} first." if funded else plain_missing)
        self.last_limit = {"state": state, "decision": None}
        resume = h.get("resume") or ({"to": "adaptation", "handoff": {"kind": "propose", "why": "The Requester added "
                                      f"credits for the items left out: {json.dumps(left, ensure_ascii=False)}"}}
                                     if left else None)
        handoff = {"kind": "requester_answered", "limit": reason, "choices": choices, "resume": resume,
                   "failing_items": drop, "drop_texts": [t for r in named for t in r.get("texts", [r["text"]])],
                   "drop_named": [r["name"] for r in named for _ in r.get("texts", [r["text"]])],
                   **({"after_ruling": True} if left and not h.get("at_audit") else {})}
        if len(choices) < 2:
            # Never a one-option question. The only lone option is "finish as it is" (with no site, "stop here" is
            # always added, and money is never added without asking): code applies it and says so.
            self.post("system", f"Only one option was possible ({opts[0]}), so it is applied without asking.")
            self.decisions.append({"kind": "limit", "asked": [q["question"]], "answer": "1", "only_option": True,
                                   "decision": "ship"})
            return self._finish_only_option(reason, gaps, room, stale, shown)
        self._ask_user(say, [q], "limit", gaps=shown, handoff=handoff)

    def _built_and_funded(self, last):
        """(done-when items the last Audit passed on the current site, items funded but not built or not checked)."""
        crit = (self.goal or {}).get("done_when") or []
        fresh = last and last["version"] == len(self.sites) and self.sites
        passed = [r["criterion"] for r in (last or {}).get("results", []) if fresh and r["pass"] and r["criterion"] in crit]
        funded = [crit[i - 1] for w in self.workflows if w["status"] in ("registered", "held", "running", "done")
                  for i in self._wf_items(w) if crit[i - 1] not in passed]
        return passed, list(dict.fromkeys(funded))

    def _finish_only_option(self, reason, gaps, room, stale, shown=None):
        """The decision question had one option, "finish as it is": code applies it and tells the Requester why
        (which limit was reached, which core items can't be paid for or left out) and what is still missing."""
        crit = (self.goal or {}).get("done_when") or []
        core = [t for g in gaps for t in [re.sub(r" \((?:not built yet|not checked yet)\)$", "", g)]
                if t in crit and must_keep(t, self.goal)]
        limit = ("you've approved the most work credits allowed" if room <= 0 else
                 "keeping going last time changed nothing" if stale else "the work you approved is used up")
        why = f"We've reached a limit: {limit}"
        why += (f", and {'; '.join(core)} can't be paid for and can't be left out, because "
                f"{'it is' if len(core) == 1 else 'they are'} core to your site." if core else ".")
        say = to_requester(why + (" So we're finishing the site as it is now." if self.sites else
                                  " So we're stopping here; nothing is built yet."), question=False,
                           keep=self._fact_source())
        self.post("identity", "@Requester " + say + gap_list(shown or gaps), to_user=True, next="user", done=True,
                  shaped=True)
        self.outcome = self._ship_outcome(f"as the only option left ({reason})")
        self.turn = None

    def _open_list(self, last, items):
        """One list for the decision message and its count. (1) What the last audit found: failing checklist items,
        design checks, browser-check problems. (2) Every checklist item an unfinished workflow covers (registered,
        held, running, failed or unfunded): mid-cycle, the audit can't know about them. One entry per checklist item.
        Items built after the last audit are marked "(not checked yet)"."""
        crit = (self.goal or {}).get("done_when") or []
        out = {}
        for r in items:                     # by text: an item dropped since this audit is gone, the rest renumbered
            if r["criterion"] in crit:
                out[("item", crit.index(r["criterion"]) + 1)] = r["criterion"]
        for w in self.workflows:
            if w["status"] != "done":
                for i in self._wf_items(w):
                    out.setdefault(("item", i), f"{crit[i - 1]} (not built yet)")
        for r in (last or {}).get("rubric", []) if not self.doc else []:     # a document lists its sections only
            if not r["pass"]:
                out.setdefault(("rubric", r["id"]), r["criterion"])
        mech = (last or {}).get("mechanical") or {}
        for p in (j2_check.problems(mech) if mech.get("available") else ["the browser check didn't run"]) \
                if last and not self.doc else []:
            out.setdefault(("check", p), p)
        for c in (last or {}).get("conditions") or [] if not self.doc else []:   # by their text (the Requester reads it)
            if blocks(c) and not (not self.doc and c.get("check") in RENDER_CONDITIONS):
                out.setdefault(("condition", c["id"]), c["text"])
        if last and last["version"] < len(self.sites):
            for w in self.workflows:
                if w["status"] == "done" and w.get("version", 0) > last["version"]:
                    for i in self._wf_items(w):
                        text = re.sub(r" \((?:not built yet|not checked yet)\)$", "", out.get(("item", i), crit[i - 1]))
                        out[("item", i)] = f"{text} (not checked yet)"
        return list(out.values()) or ([] if last else ["nothing has been checked yet"])

    def _identity_audit(self, h):
        """J2.h: is the site workable? Code decides that (workable()); Identity declares done with Audit's evidence.
        Not workable: back to J2.b, unless Audit needs a fact only the Requester knows (Identity asks), only
        to-be-confirmed items are left, or the same failure came back twice in a row (Identity asks what to do)."""
        self.step = 8
        last = self.audits[-1] if self.audits else None
        fails, findings = workable(last), self._findings(last)
        if self.pending_left and not self._waiting_work():
            # The cycle's funded work is built and audited: now Identity asks about what the ruling left out.
            p, self.pending_left = self.pending_left, None
            crit = (self.goal or {}).get("done_when") or []
            passed = {r["criterion"] for r in (last or {}).get("results", []) if r["pass"]}
            left = [c for c in p["left_out"] if c in crit and c not in passed]
            if left:
                need = sum(item_min(c) for c in left) - (self.budget - self.spent)
                return self._identity_limit({**p, "left_out": left, "at_audit": True,
                                             "budget_needed": need if need > 0 else None})
        tbc = {t["criterion"] for t in (self.goal or {}).get("to_be_confirmed") or []}
        facts = [x for x in findings if x.get("needs_requester_fact") and x.get("verdict") not in ("PASS", "PASS_WITH_WARNING")
                 and x.get("criterion") not in tbc]
        recent = self.audits[max(self.fail_floor, len(self.audits) - 2):]
        same = len(recent) == 2 and failure_signature(recent[0]) == failure_signature(recent[1]) and fails
        # Only the Requester can help when every remaining failure is a to-be-confirmed item: no other done-when item,
        # no rubric item, no render problem. (Checking done-when alone sent a lone rubric failure to the Requester as
        # a limit question: runs/j2-20261007-184146.) Anything else the department can fix loops back to J2.b.
        open_items = [r for r in (last or {}).get("results", []) + (last or {}).get("rubric", []) if not r["pass"]]
        only_tbc = bool(fails and open_items and all(r["verdict"] == "TO_BE_CONFIRMED" for r in open_items)
                        and not j2_check.problems(last.get("mechanical") or {})
                        and ((last.get("mechanical") or {}).get("available") or self.doc))
        if fails and not facts and (same or only_tbc):
            return self._identity_limit({"reason": "Only you can close what's left: these done-when items are to be "
                                                   "confirmed." if only_tbc else
                                         "The same failure came back twice in a row, so another cycle would likely "
                                         "repeat it.", "tbc": bool(only_tbc)})
        cheapest = min(RATES.values()) if RATES else 1
        if fails and self.budget - self.spent < cheapest:
            # Not workable and nothing can be funded: only the Requester's decision moves it (runs/j2-20261008-123848
            # msgs 22 and 48: Identity offered money and "finish as is" in its own question; code applied neither).
            return self._identity_limit({"reason": f"Out of budget: {self.budget - self.spent} of {self.budget} work credits "
                                                   "left, less than the cheapest rate-card part, and the site isn't "
                                                   "workable yet.", "out_of_budget": True})

        def mock():
            if not fails and self.doc:
                return json.dumps({"to_requester": f"@Requester Your {self.out['label']} is ready. Every section you "
                                                   "asked for is in it.", "to_agents": "", "clause": "done_when",
                                   "done": True, "next": "Requester"})
            if not fails:
                return json.dumps({"to_requester": "@Requester Your website is ready. Everything you asked for is on "
                                                   "the page, and it works on phones and computers.",
                                   "to_agents": "", "clause": "done_when", "done": True, "next": "Requester"})
            return json.dumps({"to_requester": "", "to_agents": "@Adaptation Not workable yet; the failing findings "
                                                                "need rework.",
                               "clause": "done_when", "done": False, "needs_requester": bool(facts),
                               "requester_questions": [], "next": "Requester" if facts else "Adaptation"})
        r = self.call("identity", '{' + REPLY_FMT + ', "clause": str, "done": bool, "needs_requester": bool, '
                      '"requester_questions": [str], "instruction_to_adaptation": str|null, ' + DECIDE_FMT + ', '
                      + next_field(("user", "adaptation"), "Requester when done=true or needs_requester=true; "
                                   "Adaptation for rework") + "}",
                      f"AUDIT'S FINDINGS (all, in full): {json.dumps(findings, ensure_ascii=False)}\n"
                      f"WORKABLE CHECK BY CODE: {'workable' if not fails else 'not workable: ' + json.dumps(fails, ensure_ascii=False)}\n"
                      f"FACTS ONLY THE REQUESTER KNOWS (from Audit): {json.dumps([x.get('needs_requester_fact') for x in facts], ensure_ascii=False)}\n"
                      f"{h.get('why', '')}\nJ2.h: the site is done only when it is workable: every done-when item passes "
                      "Audit, no JS errors, no horizontal overflow at 360px and 1280px, forms work, no broken in-page "
                      "link, every necessary condition met, nothing to be confirmed. Code checks this. Workable: declare done=true and tell the Requester "
                      "the goal is reached. Not workable: it goes back to Adaptation (J2.b); give the instruction. If a "
                      "fix needs a fact only the Requester knows (facts above, or any you see), set needs_requester=true "
                      "and ask for it in requester_questions; the run pauses until they answer. " + DECIDE_RULE, mock,
                      defaults=("Your website is ready.", "") if not fails else
                      ("We're still working on your site.", "Not workable yet: rework the failing findings."))
        if not fails and r.get("done"):
            self._milestone()                  # a goal reached: its email (when milestone emails are on); never stops the run
            if self.stages:                    # a design-document run: the next stage starts, or the last one is recorded
                self.post("system", f"The {self.out['noun']} is workable: every check passes.")
                if self._next_stage():
                    return
            self.post("identity", ("We're finished. Kept as they were: " + "; ".join(self.open_docs) + "." if self.open_docs else
                                   "Your design documents are ready." if self.docset and self.stage_kind != "website" and
                                   len(self.docset) > 1 else
                                   r.get("to_requester") or r.get("say") or ("Your document is ready." if self.doc else
                                                                             "Your website is ready.")), done=True,
                      to_user=True, next="user")
            self.outcome = (self._ship_outcome("after the Requester accepted a document as it is") if self.open_docs else
                            "Done: the design documents are workable. Every done-when item, rubric item and necessary "
                            "condition passes in each." if self.docset and self.doc and len(self.docset) > 1
                            else self.out["done"])
            self.turn = None
            if self.memory and not self.docset:
                self._save_playbook()
            return
        if fails and r.get("needs_decision"):
            return self._identity_limit({"reason": "Identity: the site can't be finished as asked without your "
                                                   "decision.", "budget_needed": as_int(r.get("budget_needed")) or None})
        questions = clean_questions(r.get("requester_questions"))
        questions += [self._fact_question(x) for x in facts]
        if questions and (facts or r.get("needs_requester")):
            return self._ask_user(r.get("to_requester") or r.get("say", ""), questions, "fact",
                                  handoff={"kind": "requester_answered"}, done=False)
        note0 = ("Identity declared done, but the site isn't workable, so done-when can't hold. Back to J2.b."
                 if r.get("done") and fails else
                 "The site is workable, but Identity didn't declare done. Back to J2.b with Identity's reason."
                 if not fails else None)
        nxt, note = self._pick("identity", r.get("next"), ("adaptation",), "adaptation")
        self._reply(r, nxt, note0 or note, done=False)
        self._handoff("adaptation", kind="propose", new_cycle=True,
                      why=f"Identity at J2.h: {r.get('instruction_to_adaptation') or r.get('to_agents') or r.get('say')} "
                          "Not workable yet: "
                          f"{json.dumps(fails, ensure_ascii=False)}")

    # ------------------------------------------------------------------ conflicts between any two agents
    def _open_conflict(self, topic, between, about, subject, detail=None):
        c = {"id": f"C-{len(self.conflicts) + 1}", "cycle": self.cycle, "topic": topic, "between": between,
             "about": about, "subject": subject, "detail": detail, "attempts": [], "status": "open",
             "escalated": False, "attempts_before_escalation": None, "resolved_by": None, "identity": [],
             "escalated_to_user": False, "user_question": None, "user_answer": None, "winner": None, "outcome": None}
        self.conflicts.append(c)
        self._event(between[0], "dispute", f"{c['id']}: {NAMES[between[0]]} vs {NAMES[between[1]]} about {subject}")
        self._save_conflicts()
        return c

    def _attempt(self, c, who, text, position, msg, agree=None, n=None):
        n = n if n is not None else sum(1 for x in c["attempts"] if x["n"] > 0) + 1
        c["attempts"].append({"n": n, "who": who, "msg": msg["id"] if msg else None, "text": text or "",
                              "position": position, "agree": agree})
        self._save_conflicts()

    def _escalate(self, c, why=None):
        """No agreement after MAX_ATTEMPTS (or the J2.c bargain hit its time or spend bound): both final positions
        go to Identity (J2.d)."""
        a, b = c["between"]
        c["status"], c["escalated"] = "escalated", True
        c["attempts_before_escalation"] = sum(1 for x in c["attempts"] if x["n"] > 0)
        why = why or f"didn't settle it in {MAX_ATTEMPTS} attempts"
        c["escalated_because"] = why
        self.bargain = None
        self._event("system", "conflict", f"{c['id']}: {NAMES[a]} and {NAMES[b]}: {why}")
        self.post("system", f"{NAMES[a]} and {NAMES[b]} on {c['id']}: {why}. Passing both final positions to "
                            "Identity (J2.d).")
        self._save_conflicts()
        self._handoff("identity", kind="resolve", conflict=c["id"])

    def _settle(self, c, winner, outcome, by="agents"):
        c["status"], c["winner"], c["outcome"] = "settled", winner, outcome
        c["resolved_by"] = c["resolved_by"] or by
        c["attempts_used"] = sum(1 for x in c["attempts"] if x["n"] > 0)
        self._save_conflicts()

    def _can_dispute(self, other, about):
        """(ok, subject id or reason): Adaptation may dispute a failing Audit finding or a held workflow, once each."""
        about = str(about or "").strip()
        if other == "audit":
            last = self.audits[-1] if self.audits else None
            fails = [r for r in (last["results"] + last.get("rubric", []) + last.get("extra_issues", []) if last else [])
                     if not r.get("pass")]
            num = lambda r: self._no(r["criterion"]) if r in last["results"] else r.get("item")    # charter number
            f = next((r for r in fails if about.lower() in (str(r.get("id")).lower(), f"item {num(r)}",
                                                             str(num(r)))), None)
            if not f:
                return False, f"no failing Audit finding '{about}'"
            sid = f["id"]
        elif other == "coordination":
            w = self._wf(about.upper())
            if not w or w["status"] != "held":
                return False, f"no held workflow '{about}'"
            sid = w["id"]
        else:
            return False, "Adaptation can dispute only an Audit finding or a Coordination HOLD"
        done = next((c for c in self.conflicts if c["about"] == sid and c["topic"] != "proposal"), None)
        return (False, f"{sid} was already disputed in {done['id']}") if done else (True, sid)

    def _turn_dispute(self, fn, h):
        """One attempt in a disagreement between two agents (not Adaptation <-> Priority, which bargain over
        workflows): agree with the other side, or hold a position with a reason. After MAX_ATTEMPTS, Identity rules."""
        c = self._conflict(h.get("conflict"))
        other = next(x for x in c["between"] if x != fn)
        n = sum(1 for x in c["attempts"] if x["n"] > 0) + 1

        def mock():
            if fn == "audit":
                return json.dumps({"say": "@Adaptation The finding stands: without a viewport meta tag, phones render "
                                          "the page at desktop width and shrink it, whatever the CSS does.",
                                   "agree": False, "position": "Finding stands: add the viewport meta tag."})
            return json.dumps({"say": f"@{NAMES[other]} Fair point, I accept that. I'll propose the fix.",
                               "agree": True, "position": f"Accepts {NAMES[other]}'s position."})
        tried = "\n".join(f"  attempt {x['n']} {NAMES[x['who']]}: {x['text']}" for x in c["attempts"])
        r = self.call(fn, '{"say": str (address @' + NAMES[other] + '), "agree": bool (true if you accept their '
                      'position), "position": str (your position in one sentence)}',
                      f"DISAGREEMENT {c['id']}: {NAMES[c['between'][0]]} vs {NAMES[c['between'][1]]} about "
                      f"{c['subject']}\n{c.get('detail') or ''}\nATTEMPTS SO FAR:\n{tried}\nThis is attempt {n} of "
                      f"{MAX_ATTEMPTS}. Settle it with {NAMES[other]} if you can: agree when their reason is better (say "
                      "what you accept), or keep your position with a reason tied to the charter or the evidence. After "
                      f"{MAX_ATTEMPTS} attempts it goes to Identity with both final positions."
                      + (f"\nRECEIPTS (written by code when it ran the work atoms' checks; the latest per check): "
                         f"{json.dumps(list({(x['atom'], x['check']): x for x in self.receipts}.values())[-30:], ensure_ascii=False)}"
                         if self.receipts else ""), mock)
        agree = bool(r.get("agree"))
        last = n >= MAX_ATTEMPTS and not agree
        m = self._say(fn, r.get("say", ""), None if agree else ("identity" if last else other), None,
                      conflict=c["id"], attempt=n, agree=agree)
        self._attempt(c, fn, r.get("say"), r.get("position"), m, agree, n)
        if agree:
            return self._apply(c, other, {"instruction": r.get("say")})
        if last:
            return self._escalate(c)
        self._handoff(other, kind="dispute", conflict=c["id"])

    def _apply(self, c, winner, r=None):
        """Carry out a settled conflict. winner: the agent whose position stands, "split" (proposal only), or None
        (no ruling: the default for the topic applies)."""
        r = r or {}
        topic, sid = c["topic"], c.get("about")
        by = c["resolved_by"] or "agents"
        if topic == "proposal":
            pa, pp = c.get("adaptation_final") or [], c.get("priority_final") or []
            if winner is None:
                self._settle(c, None, "Parked: no ruling, nothing registered from this negotiation.", by)
                self.conflict, self.exchanges, self.a_pos, self.p_pos = None, 0, [], []
                return self._parked(f"{c['id']} ended without a ruling")
            if winner == "split":
                names = [str(x) for x in r.get("workflows") or []]
                wfs = [next((x for x in pp + pa if x.get("name") == n), None) for n in names]
                missing = [n for n, x in zip(names, wfs) if x is None]
                if missing:
                    self._event("identity", "invalid_move", f"{c['id']}: split names unknown workflows {missing}")
                    self.post("system", f"Split check: {', '.join(missing)} isn't in either final position, so it isn't "
                                        "registered.", redirect=True)
                wfs = [x for x in wfs if x]
            else:
                wfs = pa if winner == "adaptation" else pp
            label = "Split" if winner == "split" else f"{NAMES[winner]}'s position stands"
            if self.fixes and not c.get("repriced") and granted_total(wfs) > granted_total(pp):
                # Fix (a): a ruling that costs more than Priority granted goes back to Priority once for pricing.
                c.update(repriced=True, status="repricing", pending_winner=winner, ruled_wfs=wfs)
                self._save_conflicts()
                self.post("system", f"J2_FIXES=1: Identity's ruling on {c['id']} costs {granted_total(wfs)}, more than "
                                    f"Priority granted ({granted_total(pp)}). It goes back to Priority once for pricing; "
                                    "Priority's answer is final.")
                return self._handoff("priority", kind="ruling", table=wfs, conflict=c["id"])
            self._settle(c, winner, f"{label}: {brief_list(wfs)} goes to the validator.", by)
            return self._agree(wfs)
        if topic == "hold":
            w = self._wf(sid)
            if winner == "adaptation" and w:
                w["hold_override"] = True
                self._settle(c, winner, f"HOLD on {sid} lifted: Coordination schedules it.", by)
                return self._handoff("coordination", kind="schedule")
            self._settle(c, winner or "coordination", f"HOLD on {sid} stands: Adaptation reshapes or parks the work.", by)
            return self._handoff("adaptation", kind="propose", why=f"{c['id']}: the HOLD on {sid} stands. "
                                 f"{r.get('instruction') or ''} Reshape the work so the missing condition is met, or park it.")
        # finding
        audit = self.audits[-1]
        f = next((x for x in audit["results"] + audit.get("rubric", []) + audit.get("extra_issues", []) if x.get("id") == sid), None)
        if winner == "adaptation" and f:
            f.update({"pass": True, "verdict": "WITHDRAWN", "note": f"{f.get('note', '')} (withdrawn in {c['id']})"})
            audit["all_pass"] = all(x.get("pass") for x in audit["results"] + audit.get("rubric", []))
            self._save_audits()
            self._settle(c, winner, f"Finding {sid} withdrawn.", by)
            if audit["all_pass"]:
                return self._handoff("identity", kind="audit")
            return self._handoff("adaptation", kind="propose", why=f"{c['id']}: finding {sid} withdrawn; other "
                                                                   f"findings still fail: {self._gaps()}")
        self._settle(c, winner or "audit", f"Finding {sid} stands: Adaptation proposes the rework.", by)
        self._handoff("adaptation", kind="propose", why=f"{c['id']}: Audit's finding {sid} stands. "
                                                        f"{r.get('instruction') or ''} Propose the rework.")

    def _bargain_over(self):
        """J2.c: the Adaptation <-> Priority bargain is bounded by four hops (MAX_ATTEMPTS exchanges, checked by the
        callers), fifteen minutes and its spend (AI calls). Returns the bound it passed, or None."""
        b = self.bargain
        if not b:
            return None
        mins = (time.time() - b["start"]) / 60
        if mins >= BARGAIN_MINUTES:
            return f"the bargain ran {mins:.1f} min, past its {BARGAIN_MINUTES:g}-minute bound"
        if self.calls - b["calls"] >= BARGAIN_CALLS:
            return f"the bargain spent {self.calls - b['calls']} AI calls, its spend bound is {BARGAIN_CALLS}"
        return None

    def _parked(self, why):
        """Adaptation reshapes or parks it -> J2.b. Work already registered is built first. If the ruling left done-when
        items out, Identity asks the Requester first."""
        self.bargain = None
        if left := self._left_by_ruling():
            if not self._waiting_work():
                self.post("system", f"Nothing new registered ({why}).")
                return self._handoff("identity", kind="limit", **left)
            self._ask_after_audit(left)
        if self._waiting_work():
            self.post("system", f"Nothing new registered ({why}). Building the work already registered.")
            return self._handoff("coordination", kind="schedule")
        self.post("system", f"Nothing registered in cycle {self.cycle}: {why}. The next cycle starts from the open gaps.")
        self._handoff("adaptation", kind="propose", new_cycle=True, why=f"Nothing was registered last cycle ({why}).")

    # ------------------------------------------------------------------ J2.b / J2.c: Adaptation <-> Priority
    WF_FMT = ('[{"id": str, "name": str, "goal_gap": str, "expected_outcome": str, "outputs": [str], '
              '"completion_check": str, "dependencies": [str], "parts": [rate-card item, repeated for quantity], '
              '"covers": [int] (the done-when item numbers this workflow builds), "fixes": [str] (the Audit finding ids '
              'it fixes, e.g. "F-3", "R-6"), '
              '"revises": str|null (registered workflow id whose site section this one rewrites; null adds a new '
              'section), "estimated_resources": {"money": int, "calls": int, "time_minutes": int}, '
              '"atoms": [{"id": str, "goal": str, "constraints": [str], "context": [str], "how": {"kind": "AI"|"code", '
              '"agent": str, "steps": str (what to build, naming every section id and class name it creates)}, '
              '"expected_output": [{"text": str, "kind": "section"|"element"|"facts"|"file"|"test"|"tap"|"behavior", '
              '"target": str|[str], "count": int (section/element only, default 1), "steps": [str] (behavior only)}], '
              '"verification": [{"text": str, "type": "file-exists"|"grep-count"|"element-count"|"named-test"|'
              '"facts-present", "target": str|[str], "expect": int|bool}]}]}]')
    ATOM_RULES = ("\nWORK ATOMS: every workflow, fixes included, carries atoms: exactly one work atom for now, id \"a1\" "
                  "(code renames it W-n.a1 at registration). goal: what must be true when it is done, one sentence. "
                  "constraints: short strings, what must never break. context: references only, charter item numbers "
                  "and section ids (e.g. \"charter item 3\", \"section W-2\"), never copied text. how: {kind: \"AI\" or "
                  "\"code\", agent: who builds it, steps: what to build, naming every section id and class name it "
                  "creates}. expected_output: written now, before any build: a short list of the concrete things that "
                  "must exist on the page when the atom is done, each {text: the thing in plain words, kind, target}: "
                  "section (target: the \"#id\" of the section the steps create; one must exist), element (target: "
                  "\"#id\" or \".class\" the steps create; count: how many), facts (target: a list of strings shown "
                  "on the page, copied exactly from the charter or the Requester's answers), file (target: a path in "
                  "the build), test (target: page-parses, links-resolve or form-has-submit), tap (target: a section id "
                  "or class whose buttons and fields must be big enough to tap), behavior (proves the page works, not "
                  "only that things exist: text says what happens in plain words, e.g. \"choosing T-shirt yes shows "
                  "the size field\"; steps: at most 8, each one of: \"set <selector> <value>\" (type into a field, "
                  "pick an option, tick a box), \"click <selector>\", \"scroll-to <selector>\", \"wait <seconds>\" "
                  "(at most 3), \"expect-visible <selector>\", \"expect-hidden <selector>\", \"expect-text <selector> "
                  "<string>\", \"expect-number <selector> <expression>\" (arithmetic over numbers from the charter, "
                  "the Requester's answers or the values the steps set, e.g. 500*0.9+2*250; never a total you worked "
                  "out yourself), \"expect-changed <selector>\" (its text differs from before the last wait), "
                  "\"expect-style-changed <selector> <css property>\" (computed value differs from before the last "
                  "click), \"expect-class <selector> <class>\". Selectors: only \"#id\" or \".class\" names the "
                  "atom's how.steps tells the builder to create. A rule that depends on a date (an early-bird price) "
                  "needs a control the atom creates to set the date (e.g. #reg-date), and the steps set it; never "
                  "rely on today's date). Every goal that describes an action (calculate, show/hide, validate, update, "
                  "highlight, toggle) needs at least one behavior item. When the charter has them, write behavior "
                  "items for: a calculator's total for a stated input; a field shown only when its trigger is set; a "
                  "field required only when a rule holds; countdown text that changes after a wait; a menu link that "
                  "gets the active class after scroll-to its section; a contrast toggle that changes the computed "
                  "background-color of the body and of one section. Code turns each item into a check and "
                  "reports it met or not met; an item no check can verify, or an empty list, is rejected. "
                  "verification: the checks code runs on the built site after the builder writes it (the "
                  "builder gets 3 attempts), each {text: the check in plain words, type, target, expect}. Only these "
                  "types: file-exists (target: a path in the build, e.g. index.html; expect true); grep-count (target: a "
                  "section id or class name; expect: how many times it appears in index.html); element-count (target: "
                  "\"#id\" or \".class\"; expect: how many elements on the page match); named-test (target: page-parses, "
                  "links-resolve or form-has-submit; expect true); facts-present (target: a list of strings that must all "
                  "appear on the page; expect true); tap-targets (target: a section id or class; code measures every "
                  "button, link and field inside it in a headless browser at 360px; expect: the minimum size in px, "
                  "44). DESIGN RUBRIC: also add, as checks of these types, the rubric rules (R-1 to R-6) that apply to "
                  "the atom's section: tap-targets on every section that holds buttons, links or form fields (R-6); "
                  "element-count of the classes the style guide gives buttons and cards (j2-btn, j2-card) when the "
                  "section uses them (R-1); links-resolve for navigation (R-5). Never add a rubric check code can't run "
                  "(contrast, hierarchy, one main action): Audit checks those. Rules the validator enforces: each facts-present string is copied "
                  "exactly from the charter or the Requester's answers (it can't be found otherwise, so no invented "
                  "facts); prefer facts-present over counts for names, times, prices, numbers and emails. grep-count and "
                  "element-count may only target a section id or class name that your how.steps or constraints tells "
                  "the builder to create, never a generic string, a charter fact or an email. A count whose text "
                  "contains a number expects exactly that number. A goal that names a time, price, number or email needs "
                  "at least one facts-present check or named test. Every claim in the goal needs at least one check that "
                  "would fail if the claim were false: a check that only proves a section exists is not enough for a goal "
                  "that names specific content. A workflow whose atom breaks a rule or lacks a field is rejected.")

    ATOM_RULES += ("\nTHE VALIDATOR REQUIRES (a reply that lacks any of this is rejected and costs a retry): every "
                   "workflow has a non-empty name, goal_gap, outputs and completion_check; every atom has a non-empty "
                   "goal, constraints, context, how, expected_output and verification. Every \"#id\" or \".class\" in a "
                   "count check or a behaviour step must appear in that atom's how.steps or constraints, written out. "
                   "EXAMPLE (one workflow, one atom, passes):\n" + json.dumps(EXAMPLE_WORKFLOW, ensure_ascii=False))

    GRANT_FMT = WF_FMT[:-2] + (', "envelope": {"money_limit": int, "model_call_limit": int, "time_limit_minutes": int, '
                               '"retry_limit": int, "stop_loss": str}}]')

    def _gaps(self):
        return [r["criterion"] for r in (self.audits[-1]["results"] if self.audits else []) if not r["pass"]]

    def _turn_adaptation(self, h):
        self.step = 2
        if h.get("kind") == "dispute":
            return self._turn_dispute("adaptation", h)
        if h.get("kind") == "respond":
            return self._adaptation_responds(h)
        if h.get("kind") == "atom_failed":
            return self._adaptation_atom_failed(h)
        if self.stages and not self.stage_checked and not self.workflows:
            self.stage_checked = True            # once, before the first plan: can the work credits cover the whole chain?
            if self._stage_gate(self.stage_i, started=True):
                return
        if self.adapt_idle and self.adapt_idle == self._adapt_state():
            # Adaptation already said it can propose nothing at this budget, and nothing changed since (no raise, no
            # drop, no new answer from the Requester). Asking again only repeats it (runs/j2-20261008-110419, msgs
            # 93-103): Identity takes the turn and asks the Requester.
            self.agent_turns["adaptation"] -= 1
            self._event("system", "adaptation_idle", "nothing changed since Adaptation could propose nothing")
            self.post("system", "Adaptation can't propose anything at this budget and nothing has changed since it said "
                                "so, so it isn't asked again. Identity asks the Requester.")
            return self._handoff("identity", kind="limit", out_of_budget=True,
                                 reason=f"Out of budget: {self.budget - self.spent} of {self.budget} work credits left, less "
                                        "than the cheapest rate-card part, so Adaptation can propose nothing.")
        last = self.audits[-1] if self.audits else None
        fails = [{**{k: r.get(k) for k in ("id", "item", "criterion", "verdict", "note")},
                  **({"item": self._no(r["criterion"])} if last and r in last["results"] else {})}   # charter numbers
                 for r in (last["results"] + last.get("rubric", []) + last.get("extra_issues", []) if last else [])
                 if not r.get("pass")]
        fails += [{"id": c["id"], "criterion": c["text"], "verdict": "NOT_MET", "note": c["evidence"]}
                  for c in (last or {}).get("conditions") or [] if blocks(c)]        # necessary conditions
        held = [{"id": w["id"], "name": w["name"]} for w in self.workflows if w["status"] == "held"]
        lessons = relevant_lessons(load_lessons(), " ".join([str(h.get("why") or ""), json.dumps(fails),
                                                             json.dumps((self.goal or {}).get("done_when"))]),
                                   include_test=self.test_lessons, output=self.output) if self.memory else []
        offered = [{k: x.get(k) for k in ("id", "problem", "cause", "fix", "test") if x.get(k) is not None}
                   for x in lessons]
        g = self.goal or {}
        pbs = relevant_playbooks(load_playbooks(), g.get("business_type"), " ".join([self.request, str(self.feedback or "")]),
                                 include_test=self.test_lessons) \
            if self.memory and (not self.workflows or self.feedback) else []
        pb_offer = [{k: x.get(k) for k in ("id", "business_type", "sections", "build_order", "watch_out", "advice",
                                           "outcome", "test") if x.get(k) is not None} for x in pbs]

        crit = (self.goal or {}).get("done_when") or []

        def items(pat, kinds=None, skip=r"$^"):             # charter numbers, as an agent writes them
            return [self._no(c) for c in crit if re.search(pat, c, re.I) and not re.search(skip, c, re.I)
                    and (kinds is None or item_kind(c) in kinds)]

        def mock_plain():
            # The mock declares honest covers/fixes: each workflow's parts pay for the items it covers (J2.e checks).
            if self.doc:
                return json.dumps(self._mock_doc_plan(fails))
            if self.blueprint and not self.workflows and not self.feedback:
                return json.dumps(self._mock_blueprint_plan(crit))
            if self.feedback:
                # A change to something an earlier workflow already built is a revise of that workflow's section,
                # priced as a fix. Which workflow that is comes off the board (Q-n), not from a guess.
                owner = self._board_owner("timings", "hours", "address", "location") \
                    if j2_board.enabled() and re.search(r"\bhours?\b|\btimings?\b|\baddress\b",
                                                        str(self.feedback), re.I) else None
                if owner:
                    return json.dumps({"say": f"@Priority The Requester asked: \"{self.feedback}\". {owner['id']} "
                                              f"already owns that section ({owner['name']}), so this is a revise of "
                                              "it, priced as a fix.", "act": "PROPOSE", "next": "Priority",
                                       "workflows": [{"name": "Visiting hours", "goal_gap": self.feedback,
                                                      "parts": ["fix"], "revises": owner["id"], "covers": [],
                                                      "outputs": ["index.html"],
                                                      "completion_check": "The visiting hours are in that section"}]})
                return json.dumps({"say": f"@Priority The Requester asked: \"{self.feedback}\". One workflow covers it.",
                                   "act": "PROPOSE", "next": "Priority",
                                   "workflows": [{"name": "Requester change", "goal_gap": self.feedback,
                                                  "parts": ["page_section"], "covers": [self._no(crit[-1])],
                                                  "outputs": ["index.html"], "completion_check": "Change visible"}]})
            phone = next((f for f in fails if "phone" in str(f["criterion"]).lower()      # the done-when item, not R-6
                          and not str(f["id"]).startswith("R-")), None)
            if phone and self._can_dispute("audit", phone["id"])[0]:
                return json.dumps({"say": f"@Audit I disagree with {phone['id']}: the layout uses fluid widths and the "
                                          "mechanical check found no overflow at 360px.", "act": "DISPUTE",
                                   "next": "Audit", "dispute": {"with": "Audit", "about": phone["id"],
                                                                "reason": "Fluid layout, no overflow at 360px"}})
            others = [{"name": f"Cover item {f['item']}", "goal_gap": f["criterion"], "covers": [f["item"]],
                       "parts": [item_kind(f["criterion"])], "outputs": ["index.html"], "completion_check": f["criterion"]}
                      if f.get("verdict") == "UNPAID" else
                      {"name": f"Fix item {f['item']}", "goal_gap": f["criterion"], "parts": ["fix"], "revises": "W-2",
                       "fixes": [f["id"]],
                       "outputs": ["index.html"], "completion_check": f["criterion"]} for f in fails
                      if f is not phone and f.get("item") and f.get("verdict") != "TO_BE_CONFIRMED"]
            if phone:
                return json.dumps({"say": f"@Priority {phone['id']} stands. One small fix closes it: add the viewport "
                                          "meta tag." + (" Plus a fix for each other failing item." if others else ""),
                                   "act": "PROPOSE", "next": "Priority",
                                   "workflows": [{"name": "Viewport fix", "goal_gap": phone["criterion"], "parts": ["fix"],
                                                  "fixes": [phone["id"]],
                                                  "outputs": ["index.html"],
                                                  "completion_check": "viewport meta tag present"}] + others})
            if others and self.cycle > 1:
                return json.dumps({"say": "@Priority One fix per failing item.", "act": "PROPOSE", "next": "Priority",
                                   "workflows": others})
            if self.cycle == 1 and not self.workflows and (self.goal or {}).get("business_type") == "boutique":
                timings = items(r"timings|address", {"page_section"})    # only when the brief gave them
                return json.dumps({"say": "@Priority Proposal for cycle 1: scaffold and styling, products, order "
                                          "on WhatsApp, and filter."
                                          + (" Plus the store timings and address." if timings else ""),
                                   "act": "PROPOSE", "next": "Priority", "workflows": [
                                       {"name": "Scaffold and styling", "goal_gap": "No site exists",
                                        "parts": ["page_section", "styling"], "covers": [self._no(crit[0])],
                                        "outputs": ["index.html"],
                                        "completion_check": "index.html loads with a nav to every section and a main "
                                                            "action in the hero"},
                                       {"name": "Products", "goal_gap": "Each product's id and price are shown",
                                        "parts": ["page_section"], "covers": items(r"id and price", {"page_section"}),
                                        "outputs": ["index.html"],
                                        "completion_check": "Every product shows its id and price"}]
                                   # before the other sections, so the timings sit in their own section
                                   + ([{"name": "Store timings and address",
                                        "goal_gap": "Visitors can't see when the shop is open or where it is",
                                        "parts": ["page_section"], "covers": timings, "outputs": ["index.html"],
                                        "completion_check": "The timings and the address are on the page"}]
                                      if timings else []) + [
                                       {"name": "Order on WhatsApp", "goal_gap": "Order button opens a WhatsApp link",
                                        "parts": ["form"], "covers": items(r"whatsapp", {"form"}),
                                        "outputs": ["index.html"],
                                        "completion_check": "The order button opens WhatsApp with the product id "
                                                            "and price"},
                                       {"name": "Filter", "goal_gap": "Filter by size and price",
                                        "parts": ["interactive_feature"], "covers": items(r"filter", {"interactive_feature"}),
                                        "outputs": ["index.html"], "completion_check": "The filters work together"}]})
            if self.cycle == 1 and not self.workflows:
                extra = r"bean|farm|loyal|phone"
                core = [i for i in items(r".", {"page_section"}, extra) if i != self._no(crit[0])][:3] \
                    + items(r".", {"form"}, extra)[:1]
                return json.dumps({"say": "@Priority Proposal for cycle 1: scaffold and styling, the charter's core "
                                          "sections, a stories section for warmth, and subtle animations.",
                                   "act": "PROPOSE", "next": "Priority", "workflows": [
                                       {"name": "Scaffold and styling", "goal_gap": "No site exists",
                                        "parts": ["page_section", "styling"], "covers": [self._no(crit[0])], "outputs": ["index.html"],
                                        "completion_check": "index.html loads with a nav to every section and a main "
                                                            "action in the hero"},
                                       {"name": "Core sections", "goal_gap": "Every done-when section",
                                        "parts": ["page_section", "page_section", "page_section", "form"], "covers": core,
                                        "outputs": ["index.html"], "completion_check": "All done-when sections present"},
                                       {"name": "Stories section", "goal_gap": "Warm content: the beans, the farms and "
                                        "the people behind the cafe", "parts": ["blog_post"] * 3, "covers": items(r"bean|farm"),
                                        "outputs": ["stories section"], "completion_check": "3 stories visible"},
                                       {"name": "Animations", "goal_gap": "Scroll and hover effects",
                                        "parts": ["interactive_feature", "interactive_feature"], "outputs": ["CSS"],
                                        "covers": items(r"loyal")[:1],
                                        "completion_check": "Effects run on scroll"}]
                                   + [{"name": f"Interactive item {self._no(c)}", "goal_gap": c, "covers": [self._no(c)],
                                       "parts": ["interactive_feature"], "outputs": ["index.html"],
                                       "completion_check": c} for c in crit if ACTION_GOAL.search(str(c))][:1]})
            if not self.audits and self._cheapest_fit()["workflows"]:
                return json.dumps(self._cheapest_fit())         # items left out earlier, now that credits were added
            return json.dumps({"say": "@Priority Nothing left to propose.", "act": "PROPOSE", "next": "Priority",
                               "workflows": []})

        def mock_real():
            # J2_REALIZATION's first proposal: the judged plan's workflows as written. The mock gives a workflow that
            # covers no item a file output (artifacts/<name>.md) with a text check, and estimates every budget.
            wfs = []
            for x in self.real_summary["proposal_brief"]:
                w_ = {"name": x["name"], "goal_gap": x.get("goal") or x["name"], "covers": x["covers"],
                      "parts": [item_kind(crit[i - 1]) for i in x["covers"] if 0 < i <= len(crit)] or ["page_section"],
                      "outputs": ["index.html"], "completion_check": x.get("done_test") or x["name"],
                      "dependencies": x["after"], "estimated_resources": {"money": 10, "calls": 2, "time_minutes": 5}}
                if not x["covers"]:
                    f = "artifacts/" + re.sub(r"[^a-z0-9]+", "-", x["name"].lower()).strip("-") + ".md"
                    w_.update(outputs=[f], parts=[], estimated_resources={"money": 5, "calls": 1, "time_minutes": 3},
                              atoms=[{"goal": f"{f} states what {x['name']} sets", "constraints": [f"Write only {f}"],
                                      "context": ["charter goal"], "how": {"kind": "AI", "agent": "Builder",
                                                                           "steps": f"Write {f}"},
                                      "expected_output": [{"text": f"The {x['name']} heading", "kind": "heading",
                                                           "target": x["name"]},
                                                          {"text": f"{f} is written", "kind": "file", "target": f}],
                                      "verification": [{"text": f"The file has the heading {x['name']}.",
                                                        "type": "heading-present", "target": x["name"],
                                                        "expect": True, "item": 1},
                                                       {"text": f"{f} exists.", "type": "file-exists", "target": f,
                                                        "expect": True, "item": 2}]}])
                wfs.append(w_)
            return json.dumps({"say": f"@Priority The judged plan: {len(wfs)} workflows, in run order.", "act": "PROPOSE",
                               "next": "Priority", "workflows": wfs})

        def mock():
            first_real = bool(real and (self.real_summary or {}).get("proposal_brief") and not self.workflows)
            r = json.loads(mock_real() if first_real else mock_plain())
            r["workflows"] = with_atoms(r.get("workflows"), self.out.get("mock_atom"))
            if r.get("act") == "PROPOSE" and not self.workflows and sum(item_min(c) for c in crit) > self.budget - self.spent:
                r = self._cheapest_fit()                         # the charter costs more than the budget left
            if offered and r.get("act") == "PROPOSE":
                r["say"] += f" Using lesson {offered[0]['id']}: {offered[0]['fix']}"
                r["lesson_used"] = offered[0]["id"]
            if pb_offer and r.get("act") == "PROPOSE":
                r["say"] += f" Using playbook {pb_offer[0]['id']} ({pb_offer[0]['business_type']}): same build order, " \
                            "sections adjusted to this brief."
                r.update(playbook_used=pb_offer[0]["id"], playbook_changes="Sections adjusted to this brief.")
            return json.dumps(r)
        g = self.goal or {}
        keep = [self._no(c) for c in crit if must_keep(c, g)]
        fmt = ('{"say": str, "act": "PROPOSE"|"POLICY_QUESTION"|"DISPUTE", "workflows": '
                      + self.WF_FMT + ', "left_out": [{"item": int, "reason": str}] (done-when items this proposal '
                      'leaves out because the budget left can\'t pay for them), "question": str|null, "dispute": {"with": "Audit"|"Coordination", "about": str '
                      '(finding id or workflow id), "reason": str}|null, "playbook_used": str|null (id of the playbook '
                      'you followed), "playbook_changes": str|null (what you changed from it, one sentence), '
                      '"lesson_used": str|null (id of the lesson '
                      'you applied, if any), '
                      + next_field(("priority", "identity", "audit", "coordination"), "Priority to get your proposal "
                                   "funded; Identity only for a policy question, put in \"question\"; Audit or "
                                   "Coordination only when you DISPUTE their finding or HOLD") + "}")
        task = (f"WHY YOU ARE UP: {h.get('why')}\nOPEN AUDIT FINDINGS: {json.dumps(fails)}\nHELD WORK: "
                      f"{json.dumps(held)}\nPropose the smallest workflows that close the gap, or ask Identity a policy "
                      "question. If you disagree with an open Audit finding or a Coordination HOLD, DISPUTE it with your "
                      f"reason instead: you and they get {MAX_ATTEMPTS} attempts to settle it before Identity rules. "
                      "Workflows change the site; checking it is Audit's job, so don't propose check, review or report "
                      "workflows. Cost follows the work: list in covers every done-when item a workflow builds; each "
                      "newly covered item costs at least its rate-card minimum (see STATE.rate_card.item_minimums), so "
                      "bundling many items into one cheap workflow is rejected. A workflow that fixes Audit findings "
                      "lists them in fixes and is priced with one 'fix' part per finding, never as a cheaper part. "
                      "After the first Audit, every workflow must cover an item or fix a finding. "
                      "Each registered workflow gets its own agent that owns one section of the site file "
                      "(STATE.site_sections): to change what's inside an earlier workflow's section, set revises to that "
                      "workflow's id; otherwise the workflow adds a new section. A revision (revises: \"W-n\", one 'fix' "
                      "part, corrected checks in its atom) re-runs W-n's builder; it needs no Audit finding id."
                      + (f"\nLESSONS FROM EARLIER WORK (most relevant first): {json.dumps(offered, ensure_ascii=False)}\n"
                         "If one applies, use it, set lesson_used to its id and say in one clause which one you used."
                         + (" A lesson marked test is a demo lesson; if you use one, call it a test lesson."
                            if any(x.get("test") for x in offered) else "")   # only with J2_MEMORY_TEST_LESSONS=1
                         if offered else "")
                      + (f"\nPLAYBOOKS FROM EARLIER RUNS FOR SIMILAR BUSINESSES (closest first): "
                         f"{json.dumps(pb_offer, ensure_ascii=False)}\nIf one fits, follow its sections and build order "
                         "and adapt it to this brief: set playbook_used to its id, say in one clause which one you used, "
                         "and put what you changed in playbook_changes. Its watch_out items are mistakes to avoid."
                         if pb_offer else "") + (" Price every workflow with rate-card parts (see STATE.rate_card); its cost is "
                                      "fixed by them." if RATES else "")
                      + "\nPriority decides what gets funded, never you: don't judge affordability and never ask about "
                      "budget, cost or cutting. When the done-when items cost more than the budget left, still PROPOSE "
                      "the best set that fits and list each item you leave out in left_out with its reason. Never leave "
                      f"out these items (the no-errors check and the brief's core facts): {keep}."
                      + (self.out["plan"] if self.doc else "") + self._blueprint_plan() + self.ATOM_RULES
                      + f"\nNECESSARY CONDITIONS (code checks every one on the built {self.out['noun']}; it is workable "
                      "only when all are met, so your plan must build what each needs):\n"
                      + "\n".join(f"{c['id']}: {c['text']}" for c in load_conditions(self.out["conditions"])))
        real = None
        if self._realization_applies():
            real = self._realization_plan(h)
            if real is None:         # a blocking Requester fact went to Identity first
                return
            task += real
        # A proposal that follows the realization plan carries more workflows and atoms than the 900-token default
        r = self.call("adaptation", fmt, task, mock, **({"max_tokens": 8000} if real else {}))
        if self._budget_question(r):
            # Affordability is Priority's call (J2.c): a budget question is rejected, asked once more, then replaced.
            self._event("adaptation", "invalid_move", "budget POLICY_QUESTION rejected")
            self.post("system", "Adaptation's question about budget was rejected: Priority decides what gets funded. "
                                "Adaptation proposes the best set that fits instead.")
            r = self.call("adaptation", fmt, task + "\nYOUR LAST REPLY WAS REJECTED: it asked about budget, cost or "
                          "cutting. PROPOSE the best set that fits the budget left, with left_out.", mock)
            if self._budget_question(r) or str(r.get("act", "")).upper() != "PROPOSE" or not norm(r.get("workflows")):
                r = self._cheapest_fit()
                self.post("system", "Adaptation asked about budget again, so code proposes the cheapest done-when "
                                    "items that fit the budget left.")
        r = self._usable(r, fmt, task, mock)
        r = self._atoms_checked(r, fmt, task, mock)
        r = self._plan_checked(r, fmt, task, mock)
        act = str(r.get("act", "")).upper()
        note = None
        if act == "DISPUTE":
            d = r.get("dispute") if isinstance(r.get("dispute"), dict) else {}
            other = key(d.get("with") or r.get("next"))
            ok, sid = self._can_dispute(other, d.get("about"))
            if ok:
                topic = "finding" if other == "audit" else "hold"
                src = next((x for x in fails if x["id"] == sid), None) if topic == "finding" else self._wf(sid)
                c = self._open_conflict(topic, ["adaptation", other], sid,
                                        f"{'finding' if topic == 'finding' else 'HOLD on'} {sid}",
                                        detail=json.dumps(src, ensure_ascii=False))
                m = self._say("adaptation", r.get("say", ""), other, None, act="DISPUTE", conflict=c["id"], attempt=1,
                              agree=False)
                self._attempt(c, "adaptation", r.get("say"), d.get("reason"), m, False, 1)
                return self._handoff(other, kind="dispute", conflict=c["id"])
            note = f"Dispute isn't valid: {sid}. Treated as a proposal."
            self._event("adaptation", "invalid_move", note)
        wfs = norm(r.get("workflows"))
        if real and self.real_summary is not None:       # did the proposal follow the plan? (names match, by design)
            planned = [x["name"] for x in self.real_summary["proposal_brief"]]
            got = [w.get("name") for w in wfs]
            self.real_summary["proposed"] = {"planned": planned, "followed": [n for n in planned if n in got],
                                             "not_proposed": [n for n in planned if n not in got],
                                             "extra": [n for n in got if n not in planned]}
        st = self._adapt_state()
        self.adapt_idle = st if st[0] and (not wfs or act == "POLICY_QUESTION") else None
        if act == "POLICY_QUESTION":
            self._say("adaptation", r.get("say", ""), "identity", note, act="POLICY_QUESTION")
            return self._handoff("identity", kind="question", frm="adaptation", question=r.get("question") or r.get("say"),
                                 resume={"to": "adaptation", "handoff": {"kind": "propose", "why": h.get("why")}})
        nxt, note2 = self._pick("adaptation", r.get("next") if act != "DISPUTE" else "priority", ("priority",),
                                "priority")
        used = str(r.get("lesson_used") or "")
        used = used if used in {x["id"] for x in offered} else None
        used_test = bool(used and next(x for x in offered if x["id"] == used).get("test"))
        if used:
            self.lessons_used.append({"cycle": self.cycle, "lesson": used, "offered": [x["id"] for x in offered],
                                      **({"test": True} if used_test else {})})
        pb = str(r.get("playbook_used") or "")
        pb = pb if pb in {x["id"] for x in pb_offer} else None
        if pb:
            self.playbooks_used.append({"cycle": self.cycle, "playbook": pb, "changes": r.get("playbook_changes"),
                                        "offered": [x["id"] for x in pb_offer]})
        self.a_msg = self._say("adaptation", r.get("say", ""), nxt, note or note2, proposal=wfs, cycle=self.cycle,
                               **({"lesson_used": used} if used else {}), **({"lesson_test": True} if used_test else {}),
                               **({"lessons_offered": [x["id"] for x in offered]} if offered else {}),
                               **({"playbook_used": pb, "playbook_changes": r.get("playbook_changes")} if pb else {}),
                               **({"playbooks_offered": [x["id"] for x in pb_offer]} if pb_offer else {}))
        if used_test:
            self._event("adaptation", "test_lesson", f"used test lesson {used}")
            self.post("system", f"TEST LESSON: Adaptation used {used}, a lesson from a forced test failure "
                                "(J2_MEMORY_TEST_LESSONS=1, demo runs only). It is not evidence from a real fix.",
                      warning=True, test=True)
        at = self._item_at()          # left_out names charter numbers
        out = [at.get(as_int(x.get("item") if isinstance(x, dict) else x)) for x in r.get("left_out") or []]
        self.a_named = [c for i, c in enumerate(crit, 1) if i in out or any(i in self._covers(w) for w in wfs)]
        if not wfs and (left := self._left_by_ruling()):
            if self._waiting_work():
                self._ask_after_audit(left)
                return self._handoff("coordination", kind="schedule")
            self.post("system", "Nothing fits the budget left, so there is nothing for Priority to fund. Identity asks "
                                "the Requester.")
            return self._handoff("identity", kind="limit", **left)
        if not wfs:
            # Audit re-checks first, so J2.h never judges a stale audit (e.g. after the Requester confirmed a fact).
            self.post("system", "Adaptation proposed no workflows, so there is nothing to fund or register. Audit "
                                f"re-checks site v{len(self.sites)}, then Identity decides J2.h.")
            return self._handoff("audit", kind="audit")
        self.a_pos = wfs
        self.bargain = {"start": time.time(), "calls": self.calls}      # J2.c starts: hops, minutes, spend
        self._handoff("priority", kind="proposal", table=wfs)

    REAL_NOTE = ("\nTHIS PLAN RUNS IN J2. Every workflow becomes one funded J2 workflow with one output, run in "
                 "dependency order. A workflow's output is either one section of the single-page site (index.html), "
                 "checked by code (file, element and text checks, named tests, behaviour steps in a browser), or one "
                 "file that later workflows read; every workflow that lists it in inputs_from, directly or through "
                 "another, reads that file. Its done_test names what the output must contain. Checking the whole site "
                 "is Audit's job: never plan check, review, QA, test or report workflows.")

    def _realization_applies(self):
        """J2_REALIZATION=1 shapes the first plan of a website run (not documents, a blueprint build or a change request)."""
        return (self.realization and self.goal and not self.doc and not self.blueprint and not self.feedback
                and not self._all_wfs())

    def _realization_plan(self, h):
        """Adaptation's first plan through the realization axis (j2_realization): factors, stages, workflows with a
        certainty zone each, checked by code; a separate judge scores every workflow and Adaptation revises below 0.80.
        Research workflows are off: J2's calls have no web access. A blocking Requester fact goes to Identity once,
        before the proposal (returns None). Otherwise returns the text added to Adaptation's task: every workflow to
        propose, in run order, each with its output (a site section or an artifact file), and the items to leave out.
        Records realization.json and realization.md."""
        if self.real is None:
            self.post("system", "Realization is on: Adaptation plans by stage and certainty, and a separate judge scores "
                                "every workflow before anything is built.")
            r = j2_realization.plan_goal(self.goal, web=False, log=lambda m: None, extra=self.REAL_NOTE)
            self.real = r
            self.calls += r["calls"]
            self.save("realization.json", r)
            with open(os.path.join(self.dir, "realization.md"), "w", encoding="utf-8") as f:
                f.write(j2_realization.report(r))
            rows, _ = j2_realization.proposal_brief(r)
            self.real_summary = {
                "plan_confidence": r["plan_confidence"],
                "stages": [s.get("id") for s in r["plan"].get("stages") or []],
                "workflows": len(rows), "rounds": len(r["rounds"]) - 1,
                "plan_faults": r["plan_faults"], "waiting_on_identity": r["waiting_on_identity"],
                "not_ready": r["not_ready"], "calls": r["calls"], "seconds": r["seconds"]}
            self._event("adaptation", "realization", f"plan confidence {r['plan_confidence']:.2f}, "
                                                     f"{len(rows)} workflows")
            low = [x["workflow"] for x in r["not_ready"]]
            self.post("adaptation", f"Plan by realization and certainty: stages {', '.join(self.real_summary['stages'])}; "
                                    f"{len(rows)} workflows; the judge's plan confidence is {r['plan_confidence']:.2f}"
                                    + (f"; not ready: {', '.join(low)}" if low else "") + ". Details in realization.md.",
                      realization=True)
        r = self.real
        if r["waiting_on_identity"] and not self.real_asked:
            self.real_asked = True
            asks = "; ".join(str(a) for x in r["waiting_on_identity"] for a in x["ask"])
            q = f"Before anything is built, the plan needs facts only the Requester knows: {asks}"
            self._say("adaptation", f"@Identity {q}", "identity", None, act="POLICY_QUESTION")
            self._handoff("identity", kind="question", frm="adaptation", question=q,
                          resume={"to": "adaptation", "handoff": {"kind": "propose", "why": h.get("why")}})
            return None
        rows, left = j2_realization.proposal_brief(r, answered=self.real_asked)
        self.real_summary["proposal_brief"] = rows
        self.real_summary["left_out"] = left
        return (f"\nREALIZATION PLAN (judged before any build, plan confidence {r['plan_confidence']:.2f}; "
                "realization.md): propose exactly these workflows, in this order, one J2 workflow each, with the same "
                "name, the same covers, and dependencies naming the workflows in after. Give each one output in "
                "outputs: index.html when it builds part of the site, or one file that later workflows read (you "
                "choose its path, name and type, e.g. a .md or .json path; keep the same path in every reply). A file "
                "workflow builds no part of the "
                "site; its work atom's checks run on the file as text, so use only the check types file-exists, "
                "facts-present, heading-present and section-not-empty (expected_output kinds file, facts and heading). "
                "Estimate each "
                "workflow's budget yourself in estimated_resources.money: there is no rate card. Every work atom must "
                "prove its workflow's done_test; an experiment workflow builds only the probe it describes.\n"
                + json.dumps(rows, ensure_ascii=False)
                + (f"\nLEAVE OUT these items, in left_out with this reason: {json.dumps(left, ensure_ascii=False)}"
                   if left else ""))

    def _mock_doc_plan(self, fails):
        """Mock Adaptation for a document: one workflow per section in the first cycle; then one fix per failing
        checklist item, rewriting the section that holds it. ponytail: no mock for change requests to a document."""
        crit, secs = (self.goal or {}).get("done_when") or [], self.out["sections"]
        sec_of = lambda text: next((x for x in secs if x.lower() in str(text).lower()), None)
        if not self.workflows:
            link = {"Components": "prd_ids", "Pages and sections": "hld_ids"}      # the ids the plan must name (a chain link)
            ids = lambda x: ((self.goal or {}).get("chain") or {}).get(link.get(x)) or []
            wfs = [{"name": f"Write {x}", "goal_gap": f"The {self.out['noun']} has no {x} section yet", "parts": [item_kind(x)],
                    "covers": [self._no(c) for c in crit if sec_of(c) == x], "outputs": [self.out["file"]],
                    "completion_check": f"## {x} is present and not empty" + (f", and covers {', '.join(ids(x))}"
                                                                                if ids(x) else "")} for x in secs]
            return {"say": f"@Priority One workflow per PRD section: {', '.join(secs)}.", "act": "PROPOSE",
                    "next": "Priority", "workflows": wfs}
        fix = []
        for f in fails:
            x = sec_of(f.get("criterion"))
            owner = next((w["id"] for w in reversed(self.workflows) if x and x.lower() in w["name"].lower()), None)
            if x and owner and f.get("verdict") != "TO_BE_CONFIRMED":
                fix.append({"name": f"Fix {x}", "goal_gap": f["criterion"], "parts": ["fix"], "fixes": [f["id"]],
                            "revises": owner, "outputs": [self.out["file"]], "completion_check": f["criterion"]})
        return {"say": "@Priority " + ("One fix per failing item." if fix else "Nothing left to propose."),
                "act": "PROPOSE", "next": "Priority", "workflows": fix}

    def _atoms_checked(self, r, fmt, task, mock):
        """J2.e's work-atom check on Adaptation's workflows: one whose atom lacks a field is rejected with the field
        named and Adaptation gets one retry; what is still incomplete after it is dropped from the proposal."""
        for retry in (True, False):
            pool = norm(r.get("workflows"))
            bad = {w.get("name") or "?": m for w in pool if (m := self._atom_problems(w, pool))}
            if not bad:
                break
            why = [self._reject(n, "; ".join(m)) for n, m in bad.items()]
            if retry:
                r = self.call("adaptation", fmt, task + "\nTHE VALIDATOR REJECTED: " + "; ".join(why) + ". Send your "
                              "reply again with every field of every work atom filled in.", mock)
            else:
                r = {**r, "workflows": [w for w in pool if not self._atom_problems(w, pool)]}
        for w in norm(r.get("workflows")):
            self.proposed_atoms[w.get("name")] = normalise_checks(w["atoms"])
            if self.blueprint:
                self.proposed_lld[w.get("name")] = w.get("lld_sections")
        return r

    REVISION_EXAMPLE = {"say": "@Priority One revision of W-1: its builder rewrites the section with the corrected check.",
                        "act": "PROPOSE", "next": "Priority", "left_out": [], "workflows": [
                            {"name": "Revise W-1", "goal_gap": "W-1's check failed", "revises": "W-1", "covers": [],
                             "parts": ["fix"], "outputs": ["the section W-1 owns"], "completion_check": "its checks pass",
                             "atoms": ["<one work atom, as in ATOM RULES>"]}]}

    def _unusable(self, r):
        """Why Adaptation's reply can't be acted on (None when it can): no next speaker, or a proposal whose workflows
        aren't a list (a reply cut short becomes a bare "say", runs/j2-20261013-185235 msgs 24-46)."""
        act = str(r.get("act") or "").upper()
        if not r.get("next"):
            return "it names no next speaker (\"next\")"
        if act in ("", "PROPOSE") and not isinstance(r.get("workflows"), list):
            return "it has no \"workflows\" list"
        return None

    def _usable(self, r, fmt, task, mock):
        """An unusable reply is re-asked once with the exact format and an example. Unusable again: code builds the fix
        itself (_code_fix), so recovery always makes progress."""
        why = self._unusable(r)
        if not why:
            return r
        self._event("adaptation", "invalid_move", f"unusable reply: {why}")
        self.post("system", f"Adaptation's reply can't be used: {why}. Adaptation is asked once more, with the exact "
                            "format.")
        r = self.call("adaptation", fmt, task + f"\nYOUR LAST REPLY COULD NOT BE USED: {why}. Reply with ONE complete "
                      f"JSON object in exactly this format: {fmt}\nEXAMPLE (a revision of W-1):\n"
                      f"{json.dumps(self.REVISION_EXAMPLE, ensure_ascii=False)}", mock)
        why = self._unusable(r)
        if not why:
            return r
        self._event("adaptation", "invalid_move", f"unusable reply again: {why}")
        r = self._code_fix()
        self.post("system", f"Adaptation's reply couldn't be used twice ({why}), so code proposes the fix itself: "
                            + (", ".join(w["name"] for w in r["workflows"]) or "nothing fits") + ".")
        return r

    def _code_fix(self):
        """Code's proposal when Adaptation can't give one: a revision of each workflow whose atom failed its checks
        (escalated) and isn't being redone yet, re-running its builder with the checks corrected (heading targets as
        heading text, targets read as the document writes them), one 'fix' part each (the rate-card minimum). No such
        workflow: the cheapest done-when items that fit (_cheapest_fit)."""
        redone = {x.get("revises") for x in self.workflows if x["status"] in ("registered", "held", "running", "done")}
        doc = self.sites[-1] if self.sites else ""
        wfs = []
        for w in self.workflows:
            esc = [a for a in w.get("atoms") or [] if a.get("status") == "escalated"]
            if not esc or w["id"] in redone or w["section"] not in self.sections:
                continue
            atoms = [{**{k: v for k, v in a.items() if k not in ("id", "results", "status", "attempts")},
                      "expected_output": [{k: v for k, v in x.items() if k != "met"}
                                          for x in a.get("expected_output") or [] if isinstance(x, dict)],
                      "verification": [(corrected_check(c, doc) or (c,))[0] if self.doc else c
                                       for c in normalise_checks([a])[0]["verification"]]} for a in esc]
            wfs.append({"name": f"Revise {w['id']} {w['name']}", "goal_gap": f"{', '.join(a['id'] for a in esc)} failed "
                        "its checks", "revises": w["id"], "covers": [], "parts": ["fix"], "outputs": w.get("outputs"),
                        "completion_check": w.get("completion_check") or "its checks pass", "atoms": atoms})
        if not wfs:
            return self._cheapest_fit()
        return {"say": f"@Priority Code's proposal: {and_list(x['name'] for x in wfs)}, one 'fix' part each.",
                "act": "PROPOSE", "next": "Priority", "workflows": wfs, "left_out": [], "code_fix": True}

    def _plan_gaps(self, wfs):
        """Necessary conditions that need their own work ("plan" in conditions/<type>.json: words a workflow must
        mention), not met at the last audit, that no proposed or registered workflow mentions."""
        last = self.audits[-1] if self.audits else {}
        met = {c["id"] for c in last.get("conditions") or [] if c["met"]}
        text = json.dumps(list(wfs) + [w for w in self.workflows if w["status"] in ("registered", "held", "running", "done")],
                          ensure_ascii=False)
        gaps = [c for c in load_conditions(self.out["conditions"])
                if c.get("plan") and c["id"] not in met and not re.search(c["plan"], text, re.I)
                and (not c.get("key") or (self.goal or {}).get("chain", {}).get(c["key"]))]   # a link needs the earlier document
        return gaps + self._lld_gaps(wfs)

    def _mock_blueprint_plan(self, crit):
        """Mock Adaptation for a site with an LLD: the scaffold, then one workflow per LLD section; the charter's items
        are spread over them so each pays its rate-card minimum, the last one takes what is left."""
        secs, items = self.blueprint["sections"], [c for c in crit if "phone" not in str(c).lower()]
        wfs = [{"name": "Scaffold and styling", "goal_gap": "No site exists", "parts": ["styling"], "covers": [],
                "lld_sections": [], "outputs": ["index.html"], "completion_check": "index.html loads with a nav to every "
                                                                                 "section and a main action in the hero"}]
        for n, sec in enumerate(secs):
            mine = items[n:n + 1] if n < len(secs) - 1 else items[n:]
            wfs.append({"name": f"Build section {sec['id']}", "goal_gap": f"LLD section {sec['id']} is not on the page",
                        "parts": [item_kind(c) for c in mine] or ["interactive_feature"],
                        "covers": [self._no(c) for c in mine], "lld_sections": [sec["id"]], "outputs": ["index.html"],
                        "completion_check": f"The page matches LLD section {sec['id']}"})
        return {"say": "@Priority One workflow per LLD section, after the scaffold.", "act": "PROPOSE", "next": "Priority",
                "workflows": wfs}

    def _blueprint_plan(self):
        if not self.blueprint:
            return ""
        return ("\nTHE LLD IS THE BLUEPRINT: plan one workflow per LLD section and list the LLD section ids it implements "
                "in lld_sections (a styling or fix workflow may list none). Every LLD section needs a workflow. LLD "
                f"SECTIONS: {json.dumps([{'id': x['id'], 'anchor': x['anchor'], 'implements': x['implements']} for x in self.blueprint['sections']])}")

    def _lld_gaps(self, wfs):
        """Before building the site from the LLD: every LLD section has a workflow, and every workflow that builds
        something cites the LLD sections it implements. Pseudo-conditions, worded for the plan flag."""
        if not self.blueprint:
            return []
        ids = [x["id"] for x in self.blueprint["sections"]]
        live = [w for w in self.workflows if w["status"] in ("registered", "held", "running", "done")]
        cited = {str(i) for w in list(wfs) + live for i in w.get("lld_sections") or []}
        out = [{"id": i, "text": "an LLD section", "why": f"no workflow implements LLD section {i}"} for i in ids if i not in cited]
        out += [{"id": w.get("name"), "text": "no LLD citation", "why": f"workflow {w.get('name')!r} cites no LLD section "
                 "(lld_sections)"} for w in wfs if not w.get("lld_sections") and not w.get("revises")
                and not set(w.get("parts") or ["x"]) <= {"styling", "fix"}]
        out += [{"id": w.get("name"), "text": "unknown LLD section", "why": f"workflow {w.get('name')!r} cites "
                 f"{[i for i in w.get('lld_sections') or [] if str(i) not in ids]}, which are not LLD sections"}
                for w in wfs if any(str(i) not in ids for i in w.get("lld_sections") or [])]
        return out

    def _plan_checked(self, r, fmt, task, mock):
        """J2.e before building: a plan that leaves a necessary condition with no workflow is flagged and Adaptation
        gets one retry. If the gap is still there, the plan goes ahead with a warning (Audit's check will show it)."""
        for retry in (True, False):
            gaps = self._plan_gaps(norm(r.get("workflows"))) if norm(r.get("workflows")) else []
            if not gaps:
                return r
            why = "; ".join(c.get("why") or f"no workflow builds {c['id']} ({c['text']})" for c in gaps)
            self.rejections.append({"cycle": self.cycle, "workflow": "plan", "reason": why})
            self._event("validator", "plan_gap", why)
            if retry:
                self.post("system", f"Validator flagged the plan before building: {why}. Adaptation sends it again.",
                          rejected="plan")
                r = self._atoms_checked(self.call("adaptation", fmt, task + "\nTHE VALIDATOR FLAGGED YOUR PLAN: " + why
                                                  + ". Put that work in a workflow and send your whole reply again.", mock),
                                        fmt, task, mock)
            else:
                self.post("system", f"The plan still leaves {', '.join(c['id'] for c in gaps)} with no workflow. It goes "
                                    "ahead; the necessary-conditions check after the build will show the gap.", warning=True)
        return r

    def _fact_source(self, goal=None):
        """What a facts-present string must come from: the charter and the Requester's own words, squashed. goal: an
        earlier stage's charter, for a re-check of a document that is already filed."""
        return squash(" ".join(strings(goal or self.goal or {}) + [self.request] + [m["text"] for m in self.chat
                                                                                    if m["who"] == "user"]))

    def _atom_problems(self, w, pool=()):
        """J2.e's work-atom check: missing fields first; then the check rules (check_problems) for every atom. pool: the
        other workflows proposed with this one; an id a workflow it depends on is told to create counts as known."""
        miss = atom_missing(w)
        if miss:
            return ["missing " + ", ".join(miss)]
        # a file workflow's checks run on its file as text, and "the file exists" (code confirms the file it wrote)
        src, types = self._fact_source(), (DOC_CHECK_TYPES + ("file-exists",)) if self._file_wf(w) else self.out["check_types"]
        real = bool(getattr(self, "realization", False))
        up = self._upstream_told(w, pool) if real else ""      # base J2 stays as it was
        return [f"work atom {n}: {x}" for n, a in enumerate(w["atoms"], 1)
                for x in expected_problems(a, types) or check_problems(cover_expected(copy.deepcopy(a)), src, types, up,
                                                                       ranges=real)]

    def _priority_file_note(self, table):
        """J2_REALIZATION: what Priority needs to know about file workflows on the table ("" when there are none, so other
        runs' prompts are unchanged). They cover no done-when item by design; the build workflows after them read them."""
        files = [w.get("name") for w in table if isinstance(w, dict) and self._file_wf(w)]
        if not files:
            return ""
        return (f" FILE WORKFLOWS ON THE TABLE: {json.dumps(files, ensure_ascii=False)}. They come from the judged "
                "realization plan (realization.md): each writes one file (a product spec, design guide or architecture) "
                "that the build workflows after it read and follow, so they cover no done-when item by design and are "
                "not optional extras. Judge them like any workflow: fund them when the build workflows that depend on "
                "them are funded, and question only their cost. Folding their work into the build workflows' steps "
                "is not the plan.")

    def _file_wf(self, w):
        """A workflow whose output is a file later workflows read (J2_REALIZATION website runs only; a document run's
        workflows write sections of its document as before)."""
        return bool(getattr(self, "realization", False)) and not self.doc and is_artifact(w)

    def _next_wf_id(self, name, plan_ix=None):
        """The id a newly registered workflow gets: W-<its place in the realization plan> when it is planned and that id
        is free; else the next number after every id in use and every planned place."""
        plan_ix = plan_ix or {}
        used = {x["id"] for x in self._all_wfs()}
        if name in plan_ix and f"W-{plan_ix[name] + 1}" not in used:
            return f"W-{plan_ix[name] + 1}"
        n = max([len(self.past_workflows) + len(self.workflows), len(plan_ix)] + [as_int(str(i)[2:]) for i in used]) + 1
        while f"W-{n}" in used:
            n += 1
        return f"W-{n}"

    def _upstream_told(self, w, pool=()):
        """The how and constraints of every workflow w depends on, directly or through others, squashed: from the
        workflows proposed with it and those already registered."""
        by = {}
        for x in list(self._all_wfs()) + [x for x in pool if isinstance(x, dict)]:
            for k in (x.get("id"), x.get("name")):
                if k:
                    by[k] = x
        seen, stack, told = set(), list(w.get("dependencies") or []), []
        while stack:
            d = by.get(str(stack.pop()))
            if not d or id(d) in seen or d is w:
                continue
            seen.add(id(d))
            stack += d.get("dependencies") or []
            for a in d.get("atoms") or self.proposed_atoms.get(d.get("name")) or []:
                if isinstance(a, dict):
                    told.append([a.get("how"), a.get("constraints")])
        return squash(json.dumps(told, ensure_ascii=False)) if told else ""

    def _budget_question(self, r):
        return str(r.get("act", "")).upper() == "POLICY_QUESTION" and \
            bool(BUDGET_TALK.search(f"{r.get('question') or ''} {r.get('say') or ''}"))

    def _cheapest_fit(self):
        """Adaptation's proposal when it won't propose within the budget: one workflow per done-when item no workflow
        builds yet, must-keep items first, then the cheapest, while the budget left pays. The rest go in left_out."""
        g = self.goal or {}
        crit = g.get("done_when") or []
        built = {i for w in self.workflows if w["status"] in ("registered", "held", "running", "done")
                 for i in self._wf_items(w)}
        todo = sorted((i for i in range(1, len(crit) + 1) if i not in built),
                      key=lambda i: (not must_keep(crit[i - 1], g), item_min(crit[i - 1]), i))
        left, wfs, out = self.budget - self.spent, [], []
        for i in todo:
            c, cost = crit[i - 1], item_min(crit[i - 1])
            if cost > left:
                out.append({"item": self._no(c), "reason": f"costs {cost}, only {left} left after the items before it"})
                continue
            left -= cost
            first = not self.sites and not self.workflows and not wfs     # the first workflow lays out the page
            wfs.append({"name": f"{'Scaffold and item' if first else 'Build item'} {i}", "goal_gap": c,
                        "parts": [item_kind(c)], "covers": [self._no(c)], "outputs": ["index.html"],
                        "completion_check": c})
        wfs = with_atoms(wfs, self.out.get("mock_atom"))
        say = f"@Priority The cheapest done-when items that fit the {self.budget - self.spent} left: " + \
            (", ".join(w["name"] for w in wfs) or "none") + "." + \
            (f" Left out: items {', '.join(str(x['item']) for x in out)}." if out else "")
        return {"say": say, "act": "PROPOSE", "next": "Priority", "workflows": wfs, "left_out": out}

    def _left_by_ruling(self):
        """After Priority's ruling: done-when items Adaptation's proposal covered or left out that no workflow builds.
        Identity asks the Requester about them (_identity_limit with left_out) when the budget left can't pay for them,
        or when one of them was already left out or refused in an earlier round."""
        crit = (self.goal or {}).get("done_when") or []
        built = {crit[i - 1] for w in self.workflows if w["status"] in ("registered", "held", "running", "done")
                 for i in self._wf_items(w)}
        left, self.a_named = [c for c in self.a_named if c in crit and c not in built], []
        for c in left:
            self.deferred.setdefault(c, set()).add(self.cycle)
        twice = [c for c in left if len(self.deferred[c]) >= 2]
        need = sum(item_min(c) for c in left) - (self.budget - self.spent)
        if twice and need <= 0:
            # Budget remains, but the same item was put off in 2 rounds: the Requester decides, not a third round.
            return {"reason": f"The same checklist item(s) were left out or refused in 2 rounds: {'; '.join(twice)}.",
                    "left_out": left, "deferred": True}
        if not left or need <= 0:
            return None             # the budget left still pays for them: Adaptation proposes them next (J2.b)
        return {"reason": f"Priority's ruling leaves out {len(left)} checklist item(s) the budget left can't pay for.",
                "left_out": left, "budget_needed": need}

    def _ask_after_audit(self, left):
        """Funded work is built first: Identity's question about the items left out waits for this cycle's Audit,
        so it is asked about a real site and never while registered work is unbuilt."""
        self.pending_left = left
        self.post("system", f"{len(left['left_out'])} checklist item(s) aren't funded. The funded work is built and "
                            "checked first; then Identity asks the Requester about them.")

    def _adapt_state(self):
        """(can't fund a part, budget, done-when, Requester decisions): Adaptation is asked again once any changes."""
        cheapest = min(RATES.values()) if RATES else 1
        return (self.budget - self.spent < cheapest, self.budget, tuple((self.goal or {}).get("done_when") or []),
                len(self.decisions))

    def _proposal_conflict(self):
        """The Adaptation <-> Priority negotiation becomes a logged conflict at Priority's first COUNTER or REFUSE."""
        if not self.conflict:
            self.conflict = self._open_conflict("proposal", ["adaptation", "priority"], None,
                                                f"what to fund in cycle {self.cycle}")
            m = self.a_msg
            self._attempt(self.conflict, "adaptation", m["text"] if m else "", "opening proposal", m, None, 0)
        return self.conflict

    def _adaptation_responds(self, h):
        offer, dec = h.get("offer") or [], h.get("decision") or ""
        c = self._proposal_conflict()

        def mock():
            if dec == "refuse" or self.exchanges in (1, 3) and self.cycle == 1:
                beans = [self._no(c) for c in (self.goal or {}).get("done_when") or [] if re.search(r"bean|farm", c, re.I)]
                stories = {"name": "Stories section", "goal_gap": "Warm content: the beans, the farms and the people",
                           "covers": beans,
                           "parts": ["blog_post"] * 2, "outputs": ["stories section"], "completion_check": "2 stories"}
                say = ("@Priority The stories are what make a small cafe feel warm and personal. I can shrink them to "
                       "two (20)." if self.exchanges == 1 else
                       "@Priority Without any content beyond listings the page feels empty. I can't accept zero.")
                return json.dumps({"say": say, "act": "REVISE", "workflows": with_atoms(offer + [stories], self.out.get("mock_atom")),
                                   "next": "Priority"})
            return json.dumps({"say": "@Priority Accepted.", "act": "ACCEPT", "workflows": [], "next": "Priority"})
        fmt = ('{"say": str, "act": "ACCEPT"|"REVISE"|"PARK", "workflows": ' + self.WF_FMT
               + ' (your revision, for REVISE), ' + next_field(("priority", "identity"), "Priority to answer "
               "your revision; Identity only for a policy question") + "}")
        task = (f"PRIORITY'S DECISION: {dec.upper()}\nPRIORITY OFFERS: {json.dumps(offer)}\nACCEPT only if the "
                "offer is non-empty and still closes the gap. After a REFUSE you must REVISE (reshape it: a smaller "
                "or different proposal) or PARK it (drop it for now). Exchange "
                f"{self.exchanges + 1} of {MAX_ATTEMPTS}; after that Identity rules on both final positions."
                + self.ATOM_RULES)
        r = self._atoms_checked(self.call("adaptation", fmt, task, mock), fmt, task, mock)
        act, wfs = str(r.get("act") or "").upper(), norm(r.get("workflows"))
        act = "PARK" if act == "WITHDRAW" else act
        if act not in ("ACCEPT", "REVISE", "PARK"):
            act = "ACCEPT" if r.get("accept") and offer else ("REVISE" if wfs else "PARK")
        note = None
        if act == "ACCEPT" and (not offer or dec == "refuse"):
            act = "REVISE" if wfs else "PARK"
            note = (f"Accept isn't valid: nothing is on the table after Priority's {dec.upper() or 'answer'}. "
                    f"Treated as {act}.")
        elif act == "REVISE" and not wfs:
            act, note = "PARK", "Adaptation revised to an empty proposal. Treated as PARK."
        self.exchanges += 1
        meta = dict(exchange=self.exchanges, cycle=self.cycle, conflict=c["id"])
        if act == "ACCEPT":
            m = self._say("adaptation", r.get("say", ""), None, note, accept=True, **meta)
            self._attempt(c, "adaptation", r.get("say"), "accepts Priority's offer", m, True, self.exchanges)
            self._settle(c, "priority", f"Adaptation accepted Priority's {dec.upper()}: {brief_list(offer)} goes to "
                                        "the validator.")
            return self._agree(offer)
        if act == "PARK":
            m = self._say("adaptation", r.get("say", ""), None, note, act="PARK", **meta)
            self._attempt(c, "adaptation", r.get("say"), "parks the proposal", m, True, self.exchanges)
            self._settle(c, "priority", "Adaptation parked the proposal; nothing registered.")
            self.conflict, self.exchanges, self.a_pos, self.p_pos = None, 0, [], []
            return self._parked("Adaptation parked its proposal after Priority's " + dec.upper())
        self.a_pos = c["adaptation_final"] = wfs
        over = self._bargain_over()
        if self.exchanges >= MAX_ATTEMPTS or over:
            m = self._say("adaptation", r.get("say", ""), "identity", note, accept=False, proposal=wfs, **meta)
            self._attempt(c, "adaptation", r.get("say"), brief_list(wfs), m, False, self.exchanges)
            return self._escalate(c, over)
        nxt, note2 = self._pick("adaptation", r.get("next"), ("priority", "identity"), "priority")
        m = self._say("adaptation", r.get("say", ""), nxt, note or note2, accept=False, proposal=wfs, **meta)
        self._attempt(c, "adaptation", r.get("say"), brief_list(wfs), m, False, self.exchanges)
        if nxt == "identity":
            return self._handoff("identity", kind="question", frm="adaptation", question=r.get("say"),
                                 resume={"to": "priority", "handoff": {"kind": "proposal", "table": wfs}})
        self._handoff("priority", kind="proposal", table=wfs)

    def _turn_priority(self, h):
        if h.get("kind") == "dispute":
            return self._turn_dispute("priority", h)
        if h.get("kind") == "envelope":
            return self._priority_envelope(h)
        if h.get("kind") == "ruling":
            return self._priority_prices_ruling(h)
        self.step = 3
        table, remaining = h.get("table") or [], self.budget - self.spent

        def env(w):
            return {**w, "envelope": {"money_limit": w["cost"], "model_call_limit": 2, "time_limit_minutes": 10,
                                      "retry_limit": 1, "stop_loss": "stop if the build fails twice"}}

        def mock():
            if self.cycle == 1 and len(table) >= 4 and any("stories" in w["name"].lower() for w in table):
                offer = [env({**w, "parts": ["interactive_feature"], "cost": 10} if "anim" in w["name"].lower() else w)
                         for w in table if "stories" not in w["name"].lower()]
                say = ("@Adaptation Scaffold and core sections: granted. The stories serve no done-when item, refused. "
                       "Animations cut to 10. I'm keeping the rest in reserve for Audit fixes."
                       if self.exchanges == 0 else
                       "@Adaptation Still no on the stories. The charter's sections are in; the reserve stays.")
                # The first counter deliberately names Identity, to show the route check in mock mode.
                return json.dumps({"say": say, "decision": "COUNTER", "workflows": offer,
                                   "next": "Identity" if self.exchanges == 0 else "Adaptation"})
            # Pricing a revise: Priority asks Adaptation whether this is new work or a fix before it grants
            # (board, J2_BOARD=1), instead of reading the price off the parts alone.
            revises = next((str(w.get("revises")) for w in table if w.get("revises")), "")
            ask = self._mock_board_ask("priority", "adaptation", "new work or a fix",
                                       f"Is this new work or a fix of {revises}?") if revises else {}
            return json.dumps({"say": f"@Adaptation Granted ({cost_of(table)} of {remaining} left).",
                               "decision": "GRANT", "workflows": [env(w) for w in table], "next": "Adaptation",
                               **ask})
        p = self.call("priority", '{"say": str (address @Adaptation), "decision": "GRANT"|"COUNTER"|"REFUSE", '
                      '"workflows": ' + self.GRANT_FMT + ' (what you fund, each with its envelope; for COUNTER, your '
                      'smaller version), '
                      + next_field(("adaptation", "identity"), f"Adaptation normally; Identity only once {MAX_ATTEMPTS} "
                                   "exchanges are used") + "}",
                      f"ON THE TABLE: {json.dumps(table)} (cost {cost_of(table)}, remaining budget {remaining}). "
                      f"COST PER DONE-WHEN ITEM: {json.dumps(self._per_item(table), ensure_ascii=False)}. Question any "
                      "plan whose cost per item is far below the rate-card minimums (many items bundled into one cheap "
                      "workflow, or a fix priced as a cheaper part): COUNTER or REFUSE with that reason. "
                      "A COUNTER names in say what you cut and why. "
                      f"Exchange {self.exchanges + 1} of {MAX_ATTEMPTS}; after that Identity rules on both final "
                      "positions." + (" Prices are fixed by the rate card: to spend less, remove parts or whole "
                                      "workflows, not the price." if RATES else "")
                      + self._priority_file_note(table), mock)
        dec = str(p.get("decision") or "").lower()
        wfs = norm(p.get("workflows")) or (table if dec == "grant" else [])
        note = None
        if dec not in ("grant", "counter", "refuse"):
            dec, wfs, note = "refuse", [], f"Priority's decision '{p.get('decision')}' isn't GRANT/COUNTER/REFUSE. Treated as REFUSE."
        elif dec == "grant" and granted_total(wfs) > remaining:
            dec, wfs, note = ("refuse", [], f"GRANT of {granted_total(wfs)} exceeds the remaining budget ({remaining}). "
                                            "Treated as REFUSE.")
        elif dec == "counter" and not wfs:
            dec, note = "refuse", "COUNTER with nothing in it. Treated as REFUSE."
        self.exchanges += 1
        self.p_pos = wfs
        self._track_refused(table, wfs, dec, f"at exchange {self.exchanges}")
        meta = dict(decision=dec, exchange=self.exchanges, cycle=self.cycle)
        if dec == "grant":
            m = self._say("priority", p.get("say", ""), None, note, **meta, **({"conflict": self.conflict["id"]}
                                                                              if self.conflict else {}))
            if self.conflict:
                self._attempt(self.conflict, "priority", p.get("say"), "grants Adaptation's revision", m, True,
                              self.exchanges)
                self._settle(self.conflict, "adaptation", f"Priority granted Adaptation's revision: {brief_list(wfs)} "
                                                          "goes to the validator.")
            return self._agree(wfs)
        c = self._proposal_conflict()
        c["priority_final"] = wfs
        c.setdefault("adaptation_final", self.a_pos)
        meta["conflict"] = c["id"]
        over = self._bargain_over()
        if self.exchanges >= MAX_ATTEMPTS or over:
            m = self._say("priority", p.get("say", ""), "identity", note, **meta)
            self._attempt(c, "priority", p.get("say"), brief_list(wfs), m, False, self.exchanges)
            return self._escalate(c, over)
        nxt, note2 = self._pick("priority", p.get("next"), ("adaptation",), "adaptation",
                                f"Identity only after {MAX_ATTEMPTS} exchanges; this was exchange {self.exchanges}")
        m = self._say("priority", p.get("say", ""), nxt, note or note2, **meta)
        self._attempt(c, "priority", p.get("say"), f"{dec.upper()}: {brief_list(wfs)}", m, False, self.exchanges)
        self._handoff("adaptation", kind="respond", offer=wfs, decision=dec)

    def _per_item(self, table):
        """For Priority: each proposed workflow's price against the rate-card minimum for the items it covers."""
        crit = (self.goal or {}).get("done_when") or []
        out = []
        for w in table:
            covers = self._covers(w)
            rc = price(w)[0]
            mins = sum(item_min(crit[i - 1]) for i in covers)
            out.append({"workflow": w.get("name"), "covers": [self._no_at(i) for i in covers],
                        "fixes": w.get("fixes") or [], "price": rc,
                        "rate_card_minimum_for_items": mins,
                        "per_item": round(rc / len(covers), 1) if covers and rc is not None else None})
        return out

    def _track_refused(self, table, offered, dec, when):
        """What Priority refused in this negotiation: everything on a REFUSE, or what a COUNTER left out."""
        names = {w.get("name") for w in offered}
        self.p_refused = {w.get("name"): f"Priority {'refused it' if dec == 'refuse' else 'left it out of its ' + dec.upper()} "
                                         f"{when}" for w in table if w.get("name") not in names}

    def _priority_prices_ruling(self, h):
        """J2_FIXES=1, fix (a): Priority prices Identity's ruling once and its answer is final. If the ruling is still
        over the remaining budget, Identity asks the Requester: add budget, or go with Priority's answer."""
        self.step = 3
        c, table, remaining = self._conflict(h.get("conflict")), h.get("table") or [], self.budget - self.spent

        def mock():
            keep = list(table)
            while keep and granted_total(keep) > remaining:
                keep.pop()
            env = lambda w: {**w, "envelope": {"money_limit": w["cost"], "model_call_limit": 2, "retry_limit": 1}}
            return json.dumps({"say": f"@Adaptation Identity's ruling priced: {granted_total(keep)} of {remaining} left.",
                               "decision": "GRANT" if len(keep) == len(table) else "COUNTER",
                               "workflows": [env(w) for w in keep]})
        p = self.call("priority", '{"say": str, "decision": "GRANT"|"COUNTER"|"REFUSE", "workflows": ' + self.GRANT_FMT
                      + ' (what you fund, each with its envelope)}',
                      f"IDENTITY RULED ON {c['id']}: {(c['identity'] or [{}])[-1].get('say')}\nTHE RULING TO PRICE: "
                      f"{json.dumps(table)} (cost {granted_total(table)}, remaining budget {remaining}).\nPrice it once: "
                      "GRANT, COUNTER (a smaller version you will fund) or REFUSE with a resource reason. Your answer is "
                      "final.", mock)
        dec = str(p.get("decision") or "").lower()
        wfs = norm(p.get("workflows")) or (table if dec == "grant" else [])
        if dec not in ("grant", "counter", "refuse") or granted_total(wfs) > remaining or (dec == "counter" and not wfs):
            dec, wfs = "refuse", []
        m = self._say("priority", p.get("say", ""), None, None, decision=dec, exchange="ruling", conflict=c["id"])
        c["repricing"] = {"decision": dec.upper(), "say": p.get("say"), "msg": m["id"], "workflows": wfs}
        self._track_refused(table, wfs, dec, f"when pricing Identity's ruling on {c['id']}")
        if dec == "grant":
            self._settle(c, c["pending_winner"], f"Priority granted Identity's ruling on repricing: {brief_list(wfs)} goes "
                                                 "to the validator.")
            return self._agree(wfs)
        cost = granted_total(table)
        if cost > remaining:
            need = cost - remaining
            q = {"question": f"Identity's ruling on {c['id']} costs {cost}, but only {remaining} is left. What should happen?",
                 "options": [f"Add {need} work credits so the ruling is funded in full: {brief_list(table)}",
                             f"Go with Priority's final answer ({dec.upper()}): {brief_list(wfs)}"]}
            c.update(user_question=q, escalated_to_user=True)
            self._save_conflicts()
            return self._ask_user(f"Priority's final answer on Identity's ruling was {dec.upper()}, and the ruling is over "
                                  "budget.", [q], "ruling_budget", handoff={"kind": "ruling_budget", "conflict": c["id"],
                                                                            "need": need},
                                  conflict=c["id"], auto=True)
        self._settle(c, "priority", f"Priority's final answer on Identity's ruling: {dec.upper()} {brief_list(wfs)}.")
        if wfs:
            return self._agree(wfs)
        self.conflict, self.exchanges, self.a_pos, self.p_pos = None, 0, [], []
        self._parked(f"Priority refused Identity's ruling on {c['id']} when pricing it")

    def _identity_ruling_budget(self, h):
        """J2_FIXES=1: the Requester answered whether to add budget for Identity's ruling."""
        self.step = 4
        c, answer, need = self._conflict(h.get("conflict")), h.get("answer"), as_int(h.get("need"))
        c["user_answer"] = answer
        r = self.call("identity", '{' + REPLY_FMT + ', "choice": 1|2}',
                      f"YOU ASKED THE REQUESTER: {json.dumps(c.get('user_question'))}\nTHE REQUESTER ANSWERED: {answer}\n"
                      "Which option did they choose? Say what happens next.",
                      lambda: json.dumps({"say": "The Requester added the budget; the ruling is registered in full."
                                          if re.search(r"\b1\b|add|yes|fund", answer, re.I) else
                                          "The Requester kept Priority's final answer.",
                                          "choice": 1 if re.search(r"\b1\b|add|yes|fund", answer, re.I) else 2}))
        self._reply(r, None, conflict=c["id"], decision=True)
        c["resolved_by"] = "identity_after_user"
        if as_int(r.get("choice")) == 1:
            self._raise_budget(need, f"{c['id']} ruling")
            self.p_refused = {}              # the Requester funded the ruling, so nothing in it counts as refused
            self._settle(c, c["pending_winner"], "The Requester added budget: Identity's ruling goes to the validator.")
            return self._agree(c["ruled_wfs"])
        wfs = (c.get("repricing") or {}).get("workflows") or []
        self._settle(c, "priority", f"The Requester kept Priority's final answer: {brief_list(wfs)}.")
        if wfs:
            return self._agree(wfs)
        self.conflict, self.exchanges, self.a_pos, self.p_pos = None, 0, [], []
        self._parked(f"the Requester kept Priority's refusal on {c['id']}")

    # ------------------------------------------------------------------ J2.e: validator
    def _agree(self, wfs):
        self.step = 5
        registered, refused, defaults = self._register(wfs)
        self.cycle_registered = self.cycle_registered or bool(registered)
        if self.conflict:
            self.conflict["registered"] = [w["name"] for w in registered]
            self._save_conflicts()
        # Agreement closes this negotiation; a later one in the same cycle (e.g. a HOLD fix) starts fresh.
        self.exchanges, self.conflict, self.a_pos, self.p_pos, self.p_refused = 0, None, [], [], {}
        self.bargain = None
        if registered:
            self.feedback = None             # the Requester's change request is now being worked on
            msg = f"Registered {', '.join(w['id'] + ' ' + w['name'] for w in registered)}."
            if refused:
                msg += f" Rejected: {'; '.join(refused)}."
            if defaults:
                msg += f" Envelope defaults applied: {'; '.join(defaults)}."
            msg += " Workflow agents: " + "; ".join(
                f"{w['agent']} owns section {w['section']}" + (f" (rewrites it, was {self.sections[w['section']]['owner']})"
                                                               if w["section"] in self.sections else "")
                for w in registered) + "."
            left = self._left_by_ruling()
            self._say("validator", msg + f" Budget: {self.spent}/{self.budget} used.", "coordination", register=True,
                      charged=sum(w["cost"] for w in registered))
            if left:
                self._ask_after_audit(left)
            return self._handoff("coordination", kind="schedule")
        self._say("validator", f"Rejected: {'; '.join(refused) or 'the agreed set was empty'}. Adaptation reshapes or "
                               "parks it.", "adaptation", register=False)
        self._parked("the validator registered nothing")

    def _waiting_work(self):
        return any(w["status"] in ("registered", "held") for w in self.workflows)

    CHECK_WORK = re.compile(r"\b(audit\w*|check|self[- ]check|verif\w*|qa|review|report|test)\b", re.I)

    REQUIRED = ("name", "goal_gap", "outputs", "completion_check")

    def _reject(self, name, reason):
        """A validator rejection: counted, saved for summary.json, and posted as a system note with the exact reason."""
        self.rejections.append({"cycle": self.cycle, "workflow": name, "reason": reason})
        self._event("validator", "role_refusal" if "Audit's job" in reason else "refused", f"{name}: {reason}")
        self.post("system", f"Validator rejected '{name}': {reason}.", rejected=name)
        return f"{name} ({reason})"

    def _register(self, agreed):
        """J2.e, plain code: every proposal is checked before it is registered, including one from an Identity ruling.
        Checks: required fields present, priced from the rate card, nothing Priority refused, cost fits what's left."""
        remaining, registered, refused, defaults = self.budget - self.spent, [], [], []
        covered = {i for x in self.workflows if x["status"] in ("registered", "held", "running", "done")
                   for i in self._wf_items(x)}
        batch = norm(agreed)
        # J2_REALIZATION: register in the plan's run order, each under the id of its place in the plan (W-1 is the first
        # to run), even when it is registered later than others; other workflows take ids after the plan's
        plan_ix = {x["name"]: i for i, x in enumerate((self.real_summary or {}).get("proposal_brief") or [])}
        if plan_ix:
            batch.sort(key=lambda x: plan_ix.get(x.get("name"), 10 ** 6))
        for w in batch:
            c, name = granted(w), w.get("name") or "?"
            missing = [f for f in self.REQUIRED if not (w.get(f) or (f == "goal_gap" and w.get("purpose")))]
            atoms = normalise_checks(w.get("atoms") or self.proposed_atoms.get(name))   # Priority or a ruling may drop them
            no_atom = self._atom_problems({"atoms": atoms, "outputs": w.get("outputs"),
                                           "dependencies": w.get("dependencies")}, batch)
            lld = [str(i) for i in (w.get("lld_sections") or self.proposed_lld.get(name) or [])] if self.blueprint else []
            unknown_lld = [i for i in lld if i not in {x["id"] for x in self.blueprint["sections"]}] if self.blueprint else []
            rc, unknown = price(w)
            reason = ("a check or report on the site is Audit's job, not a workflow" if self.CHECK_WORK.search(str(name))
                      else f"required field(s) missing: {', '.join(missing)}" if missing
                      else "; ".join(no_atom) if no_atom
                      else f"it cites LLD section(s) that don't exist: {unknown_lld}" if unknown_lld
                      else f"{self.p_refused[name]}, so it can't be registered" if name in self.p_refused
                      else "no rate-card parts, so it has no price" if RATES and rc is None
                      else f"unknown rate-card parts {unknown}; allowed: {', '.join(RATES)}" if unknown
                      else "its cost is 0" if c <= 0
                      else self._underpriced(w, rc, covered) if RATES
                      else None) or (f"it costs {c}, over the {remaining} left in the budget" if c > remaining else None)
            if reason:
                refused.append(self._reject(name, reason))
                continue
            remaining -= c
            env = {k: (w.get("envelope") or {}).get(k) for k in ("money_limit", "model_call_limit", "time_limit_minutes",
                                                               "retry_limit", "stop_loss")}
            env["money_limit"] = c
            for k, d in ENVELOPE_DEFAULTS.items():
                if env[k] is None:
                    env[k] = d
                    defaults.append(f"{name}: {k}={d}")
            if not RATES and as_int((w.get("envelope") or {}).get("money_limit")) <= 0:
                defaults.append(f"{name}: no money_limit granted, charged the estimate ({c})")
            covered |= set(self._covers(w))
            wf = {"id": self._next_wf_id(name, plan_ix), "proposal_id": w.get("id"), "name": w.get("name", "workflow"),
                  "covers": [self._no_at(i) for i in self._covers(w)], "fixes": [str(f) for f in w.get("fixes") or []],
                  "covers_text": [((self.goal or {}).get("done_when") or [])[i - 1] for i in self._covers(w)],
                  "purpose": w.get("purpose") or w.get("goal_gap", ""), "cost": c, "estimate": w.get("estimate", wf_cost(w)),
                  "envelope": env, "charged": True, "calls_used": 0, "retries_used": 0, "outputs": w.get("outputs"),
                  "completion_check": w.get("completion_check", ""), "dependencies": w.get("dependencies") or [],
                  "status": "registered", "cycle": self.cycle}
            wf["atoms"] = [{**a, "id": f"{wf['id']}.a{n}"} for n, a in enumerate(atoms, 1)]
            if lld:       # the LLD slice is this workflow's spec: code adds the check that the page matches it
                wf["lld_sections"] = lld
                wf["atoms"][0]["verification"] = list(wf["atoms"][0].get("verification") or []) + [self._lld_check_for(lld)]
            # J2.e: the workflow gets its own agent, which owns one marked section of the site file. "revises" hands
            # an earlier workflow's section to it (only a section that exists and isn't already being rewritten).
            rv = str(w.get("revises") or "").upper()
            src = self._wf(rv)
            taken = {x["section"] for x in self.workflows if x["status"] in ("registered", "held", "running")}
            wf["section"] = src["section"] if src and src["section"] in self.sections and src["section"] not in taken \
                else wf["id"]
            wf["revises"] = rv if wf["section"] != wf["id"] else None
            if self.real_summary:       # J2_REALIZATION: the plan's stage (a fix keeps the stage of what it revises)
                st = {x["name"]: x.get("stage") for x in self.real_summary.get("proposal_brief") or []}.get(wf["name"])
                if st or (src or {}).get("stage"):
                    wf["stage"] = st or src["stage"]
            add_workflow_agent(wf["id"])
            wf["agent"] = NAMES[wf["id"]]
            self.workflows.append(wf)
            registered.append(wf)
        self._recalc_spent()
        self.save("workflows.json", self.workflows)
        return registered, refused, defaults

    def _lld_check_for(self, ids):
        return {"text": f"The page matches LLD section(s) {', '.join(ids)}: ids, selectors and form fields as specified",
                "type": "lld-slice", "target": ", ".join(ids), "expect": True, "spec": lld_slice(self.blueprint, ids)}

    def _wf_items(self, w):
        """The current done-when item numbers a registered workflow covers, matched by the item text stored at
        registration (covers_text): a drop renumbers the list, so numbers alone could point at the wrong item, and a
        dropped item is simply gone. Runs from before covers_text fall back to the numbers."""
        crit = (self.goal or {}).get("done_when") or []
        if "covers_text" in w:
            return [crit.index(t) + 1 for t in w["covers_text"] if t in crit]
        return self._covers(w)

    def _covers(self, w):
        """Current positions of the items a workflow covers. covers holds charter numbers (what agents see and write),
        so after a drop covers [5] still means charter item 5; a dropped item is simply gone."""
        at = self._item_at()
        return sorted({at[as_int(i)] for i in w.get("covers") or [] if as_int(i) in at})

    def _underpriced(self, w, rc, covered):
        """Cost follows the work (J2.e). None when the price holds, else the exact reason:
        - every newly covered done-when item costs at least its rate-card minimum (bundling can't make items cheaper);
        - a fix (it lists findings, or rewrites an existing section) has one 'fix' part per finding: a fix can't be
          relabelled as a cheaper part type;
        - a workflow with page_section or form parts (the parts that build done-when sections) must list the items it
          covers; after the first Audit every workflow must cover an item or fix a finding."""
        crit = (self.goal or {}).get("done_when") or []
        covers, fixes, parts = self._covers(w), [str(f) for f in w.get("fixes") or []], [str(p) for p in w.get("parts") or []]
        target = self._wf(str(w.get("revises") or "").upper())
        # rewriting a built section (done, or failed its checks) is a fix and needs no finding id; an unbuilt one is new
        is_fix = bool(fixes) or bool(target and (target["status"] == "done" or target["section"] in self.sections))
        if not covers and not fixes and not is_fix and (self.audits or {"page_section", "form"} & set(parts)):
            return ("it covers no done-when item and fixes no Audit finding: list the items it builds in covers, or "
                    "the findings it fixes in fixes" + (" (page_section and form parts build done-when items, so "
                                                       "they must say which)" if not self.audits else ""))
        need_fix = max(1, len(fixes)) if is_fix else 0
        if parts.count("fix") < need_fix:
            return (f"a fix is priced as a fix: it needs {need_fix} 'fix' part(s) "
                    f"({', '.join(fixes) or 'it rewrites ' + str(w.get('revises'))}), but its parts are {parts}")
        new = [i for i in covers if i not in covered]
        minimum = sum(item_min(crit[i - 1]) for i in new) + need_fix * RATES.get("fix", 0)
        if rc < minimum:
            each = "; ".join(f"item {i} {item_kind(crit[i - 1])} {item_min(crit[i - 1])}" for i in new)
            return (f"it costs {rc}, below the rate-card minimum of {minimum} for the work it covers ({each}"
                    + (f"; {need_fix} fix {need_fix * RATES.get('fix', 0)}" if need_fix else "") + ")")
        return None

    def _recalc_spent(self):
        """Budget used = money of every workflow currently charged (held work is released, not charged)."""
        self.spent = sum(w["cost"] for w in self._all_wfs() if w.get("charged"))

    # ------------------------------------------------------------------ J2.f: Coordination, envelope check, workflow agents
    def _turn_coordination(self, h):
        if h.get("kind") == "dispute":
            return self._turn_dispute("coordination", h)
        self.step = 6
        todo = [w for w in self.workflows if w["status"] in ("registered", "held")]
        if not todo:
            self.agent_turns["coordination"] -= 1
            self.cycle_turns -= 1
            self.post("system", f"Queue empty: no registered work left in cycle {self.cycle}. Hand-off by code: Audit "
                                f"checks site v{len(self.sites)}.")
            return self._handoff("audit", kind="audit")
        ids = [w["id"] for w in todo]
        finished = {x[k] for x in self.workflows if x["status"] == "done" for k in ("id", "name") if x.get(k)}
        waits = {w["id"]: [d for d in w["dependencies"] if d not in finished] for w in todo}
        sched = self.call("coordination", '{"say": str, "rule": str, "order": [workflow id], '
                          '"holds": [{"id": str, "missing": str}], ' + next_field(("priority", "adaptation"),
                          "Priority to check the run envelope of the first workflow in your order; Adaptation only if "
                          "everything is on HOLD") + "}",
                          f"REGISTERED AND WAITING: {json.dumps(todo)}\nDONE OR FAILED: "
                          f"{json.dumps([[w['id'], w['status']] for w in self.workflows if w not in todo])}\n"
                          "Schedule them. Priority checks each run's envelope, then the workflow's own agent runs it "
                          "and reports back to you. One writer at a time on the site file: you hold the write lock.",
                          lambda: json.dumps({"say": "@Priority Schedule: " + " → ".join(f"{w['id']} {w['name']}" for w in todo
                                                                                       if not waits[w["id"]])
                                              + ". Each one writes index.html, so they run one at a time. Please check "
                                                "the first one's envelope.",
                                              "rule": "dependencies first, then registration order; one writer at a "
                                                      "time on index.html", "order": ids, "next": "Priority",
                                              "holds": [{"id": i, "missing": f"waits for {', '.join(d)} to finish"}
                                                        for i, d in waits.items() if d]}))
        order = [i for i in sched.get("order") or [] if i in ids] + [i for i in ids if i not in (sched.get("order") or [])]
        held = {x.get("id"): x.get("missing") for x in sched.get("holds") or [] if x.get("id") in ids}
        lifted = [i for i in held if (self._wf(i) or {}).get("hold_override")]
        for i in lifted:
            held.pop(i)
        released = []
        for w in todo:
            if w["id"] in held and w.get("charged"):
                w["charged"] = False
                released.append(f"{w['cost']} from {w['id']}")
            w["status"] = "held" if w["id"] in held else "registered"
        self._recalc_spent()
        ready = [i for i in order if i not in held]
        text = sched.get("say", "") + (f" Rule: {sched.get('rule')}." if sched.get("rule") else "")
        text += "".join(f"\nHOLD {i}: {why}" for i, why in held.items())
        for i, why in held.items():
            self._event("coordination", "hold", f"{i}: {why}")
        self.save("workflows.json", self.workflows)
        notes = [f"{', '.join(lifted)}: the HOLD was lifted in a settled conflict, so it can't be held again."] if lifted else []
        if released:
            notes.append(f"Released {', '.join(released)} while on HOLD; charged again only when it runs. "
                         f"Budget: {self.spent}/{self.budget} used.")
        if ready:
            nxt, note = self._pick("coordination", sched.get("next"), ("priority",), "priority",
                                   "Priority checks the run envelope before a workflow runs")
            self._say("coordination", text, nxt, note, schedule=ready)
            for n in notes:
                self.post("system", n)
            return self._handoff("priority", kind="envelope", wid=ready[0])
        nxt, note = self._pick("coordination", sched.get("next"), ("adaptation",), "adaptation",
                               "everything is on HOLD, so Adaptation must address the missing conditions")
        self._say("coordination", text, nxt, note, schedule=[])
        for n in notes:
            self.post("system", n)
        self.hold_bounces += 1
        if self.hold_bounces > 1:
            self.post("system", "Everything is still on HOLD after Adaptation's attempt. Audit checks what was built.")
            return self._handoff("audit", kind="audit")
        self._handoff("adaptation", kind="propose", why="Coordination put workflows on HOLD: "
                      + "; ".join(f"{i}: {why}" for i, why in held.items())
                      + ". Address the missing conditions, or DISPUTE a HOLD you think is wrong.")

    def _priority_envelope(self, h):
        """J2.f: Priority checks the run envelope before the workflow runs. The limits are numbers, so code checks
        them and posts Priority's 'allowed' without an AI call; Priority's model is asked only when the run is over a
        limit: HOLD it (it doesn't run), or request approval from Identity (once per workflow)."""
        self.step = 6
        w = self._wf(h.get("wid"))
        if not w:
            self.post("system", f"Priority was handed unknown workflow {h.get('wid')}. Back to Coordination.")
            return self._handoff("coordination", kind="schedule")
        env, remaining = w.get("envelope") or {}, self.budget - self.spent
        calls = as_int(env.get("model_call_limit"), 1)
        over = ([f"money: needs {w['cost']}, {remaining} left"] if not w.get("charged") and w["cost"] > remaining else []) \
            + ([f"model calls: {w['calls_used']} of {calls} used"] if w["calls_used"] >= calls else [])
        if not over:
            recharged = 0
            if not w.get("charged"):
                w["charged"], recharged = True, w["cost"]
                self._recalc_spent()
            self._say("priority", f"Envelope check for {w['id']}: within limits (money {w['cost']}, model calls "
                                  f"{w['calls_used']}/{calls}, retries {w['retries_used']}/{as_int(env.get('retry_limit'))}"
                                  f"). Budget {self.spent}/{self.budget} used. Runs now.", w["id"], None,
                      auto=True, envelope="allowed", charged=recharged)
            return self._handoff(w["id"], kind="build", wid=w["id"])
        asked = w.get("approval_asked")
        p = self.call("priority", '{"say": str, "decision": "HOLD"|"REQUEST_APPROVAL", "reason": str, '
                      + next_field(("coordination", "identity"), "Coordination for HOLD; Identity to request approval")
                      + "}",
                      f"RUN ENVELOPE CHECK for {json.dumps(w)}: over limit ({'; '.join(over)}). Budget "
                      f"{self.spent}/{self.budget} used. HOLD it (it doesn't run; Coordination carries on with other "
                      "work) or REQUEST_APPROVAL from Identity." + (" Approval was already asked for this workflow, so "
                                                                    "HOLD." if asked else ""),
                      lambda: json.dumps({"say": f"@Identity {w['id']} is over its envelope; requesting approval.",
                                          "decision": "HOLD" if asked else "REQUEST_APPROVAL",
                                          "next": "Coordination" if asked else "Identity"}))
        self._event("priority", "hold", f"{w['id']}: over its envelope ({'; '.join(over)})")
        if str(p.get("decision")).upper() == "REQUEST_APPROVAL" and not asked:
            w["approval_asked"] = True
            self._say("priority", p.get("say", ""), "identity", None, envelope="over limit")
            return self._handoff("identity", kind="approval", wid=w["id"], over=over)
        w["status"], w["charged"] = "unfunded", False
        self._recalc_spent()
        self.save("workflows.json", self.workflows)
        self._say("priority", p.get("say", ""), "coordination", None, envelope="hold")
        self._handoff("coordination", kind="schedule")

    def _turn_workflow_agent(self, h):
        """J2.f: the workflow's own agent (e.g. "W-1 Builder") writes only its marked section of the site file, while
        Coordination's write lock is on it. It retries within its envelope (model_call_limit, retry_limit) when a reply
        can't be used. Code assembles the file from every section and runs the work atom's checks (receipts.jsonl);
        a failed check goes back to the builder with its evidence, up to MAX_CHECK_ATTEMPTS attempts. Then it reports
        to Coordination, or, after the last failed attempt, code hands the receipts to Adaptation."""
        self.step = 6
        w = self._wf(h.get("wid") or self.turn)
        if not w:
            self.post("system", f"{self.turn} has no registered workflow. Back to Coordination.")
            return self._handoff("coordination", kind="schedule")
        fn, sid = w["id"], w["section"]
        w["status"] = "running"
        self.write_lock = fn
        env = w.get("envelope") or {}
        max_calls, max_retries = as_int(env.get("model_call_limit"), 1), as_int(env.get("retry_limit"), 0)
        brief = {k: w.get(k) for k in ("id", "name", "purpose", "outputs", "completion_check", "envelope", "revises",
                                        "covers", "fixes")}
        brief["atoms"] = [{k: v for k, v in a.items() if k not in ("results", "status", "attempts")}
                          for a in w.get("atoms") or []]
        brief["covers_items"] = [((self.goal or {}).get("done_when") or [])[i - 1] for i in self._wf_items(w)]
        self.atom = (w.get("atoms") or [{}])[0].get("id")     # its AI calls are tagged with the atom they build
        if self._file_wf(w):
            return self._run_artifact(w, brief, max_calls)
        attempt, failed, evidence, prev_sig, early, spec_stop = 0, [], "", set(), None, None
        while True:
            attempt += 1
            prev = self.sites[-1] if self.sites else ""
            if len(prev) > SITE_LIMIT:
                self.post("system", f"Warning: site v{len(self.sites)} is {len(prev):,} characters. {fn} sees only the "
                                    f"first {SITE_LIMIT:,}, so the end of the file may be lost.", warning=True)
            mine = self.sections.get(sid)
            # Before any section is built, the first workflow that runs writes the style guide; every later agent
            # follows it.
            first = self.style is None and not self.doc      # a document has no style guide

            def mock():
                board_extra = self._mock_board_extra(w, attempt) if j2_board.enabled() else {}
                if self.doc:
                    return (json.dumps({"say": f"@Coordination {fn} done: {w['name']}. Section {sid} written.",
                                        "human_output": f"I wrote the {w['name']} section ({sid}) of the document.",
                                        "atom_human_output": "This adds a written part of the document that readers can use.",
                                        "next": "Coordination", **board_extra})
                            + f"\n```markdown\n{self.out['mock_section'](w, self.goal)}```")
                head, body = mock_section(w, self.goal, self.sections, self._lld_slice_for(w),
                                          [x["anchor"] for x in self.blueprint["sections"]
                                           if x["id"] in {i for o in self.workflows for i in o.get("lld_sections") or []}]
                                          if self.blueprint else ())
                return (json.dumps({"say": f"@Coordination {fn} done: {w['name']}. Section {sid} written.",
                                    "human_output": f"I built the {w['name']} section ({sid}) of the page.",
                                    "atom_human_output": "This adds a part of the page that visitors can read and use.",
                                    "next": "Coordination", **board_extra})
                        + (f"\n```style\n{json.dumps(MOCK_STYLE)}\n```" if first else "")
                        + (f"\n```head\n{head}\n```" if head else "") + (f"\n```section\n{body}\n```" if body else ""))
            style_task = ("\nYOU RUN FIRST, SO YOU ALSO WRITE THE STYLE GUIDE for this brief, before your section: "
                          "palette, fonts, spacing, button and card styles that suit this business and its visitors, with "
                          "readable contrast (WCAG AA) and buttons at least 44px tall. Code turns it into CSS variables "
                          "(--j2-primary, --j2-font-body, --j2-radius, ...) and the classes j2-btn, j2-btn-secondary and "
                          "j2-card; use them in your section." if first else
                          f"\nSTYLE GUIDE (follow it exactly): {json.dumps(self.style, ensure_ascii=False)}\nUse its CSS "
                          "variables (var(--j2-primary), var(--j2-font-heading), ...) and the classes j2-btn, "
                          "j2-btn-secondary and j2-card. Don't add other colours or fonts.")
            # A failed check buys one more call than the envelope's model_call_limit, so all attempts can run.
            limit = max_calls + attempt - 1

            def one_call():
                """One model call for this attempt (retried within the envelope on an unusable reply). Pulled out
                of line so a board Answer (item 3) can trigger a second pass within the same attempt: resuming A
                (item 4) is just calling this again once self.board/self.given carry the answer, since _page_context
                and _doc_prompt already read that state fresh on every call."""
                got, raw, stop, guide, last_err = None, "", None, None, None
                while got is None:
                    if w["calls_used"] >= limit:
                        stop = f"it used its envelope's model_call_limit ({limit})" + (f"; last problem: {last_err}"
                                                                                       if last_err else "")
                        break
                    w["calls_used"] += 1
                    raw = self.call(fn, *(self._doc_prompt(w, brief, sid, mine, attempt, evidence, last_err, prev)
                                          if self.doc else (
                                    'First line: a JSON object {"say": str (one short line reporting to @Coordination: '
                                    'what you changed; code runs your checks and reports them, so don\'t claim they pass), '
                                    '"human_output": str (1 to 3 plain sentences telling a person what you built; \"done.\" alone is rejected), "atom_human_output": str (1 or 2 plain sentences telling the Requester what this part created, in everyday words: no ids, no selectors, no \"units\"), "next": "Coordination", ' + (self.SPEC_FMT if self.blueprint else '') + (self.BOARD_FMT if j2_board.enabled() else '') + self.REASON + '}. '
                                    + ('Then a ```style block with the style guide as JSON: ' + STYLE_FMT + '. ' if first
                                       else '')
                                    + 'Then, '
                                    'only if your section needs page-head tags, a ```head block (meta, link, style or script '
                                    'tags only). Then a ```section block: your section\'s HTML only, no <html>, <head> or '
                                    '<body> tags (inline <style> and <script> are fine).',
                                    f"YOU ARE {NAMES[fn]}, the agent for one registered workflow.\nYOUR WORKFLOW: "
                                    f"{json.dumps(brief, ensure_ascii=False)}\nYOU OWN SECTION {sid} of the site file. Code "
                                    "puts your head block in the page head and your section in the body, in section order, "
                                    "between markers. Write the whole section each time; don't repeat other sections. Build "
                                    "only the done-when items your workflow covers and fix only its findings; other items "
                                    "belong to other workflows. After you write it, code runs every check in your work atoms' "
                                    f"verification on the whole page ({MAX_CHECK_ATTEMPTS} attempts)." + style_task
                                    + "\nDESIGN RUBRIC (Audit checks every section against it; follow it):\n"
                                    + "\n".join(f"{rid} {txt}" for rid, txt in RUBRIC[:6])
                                    + (f"\nYOUR SECTION NOW (last written by {mine['owner']}):\n```head\n{mine['head']}\n```\n"
                                       f"```section\n{mine['body']}\n```" if mine else "")
                                    + (f"\nATTEMPT {attempt} OF {MAX_CHECK_ATTEMPTS}. CODE RAN YOUR CHECKS ON YOUR LAST "
                                       f"ATTEMPT AND THESE FAILED:\n{evidence}\nFix the section so they pass." if evidence
                                       else "")
                                    + (f"\nYOUR LAST ATTEMPT WAS REJECTED: {last_err}. Fix that." if last_err else "")
                                    + self._artifact_context(w)
                                    + self._page_context(w, prev))),
                                    mock, max_tokens=16000, json_reply=False)
                    try:
                        parts = parse_doc(raw) if self.doc else parse_section(raw)
                        guide = parse_style(raw) if first else None
                        got = parts      # only once both parse: a missing style guide must be retried, not accepted
                        w["last_check"] = ("style guide, " if first else "") + "usable section" + (" and head tags"
                                                                                                   if got[0] else "")
                    except ValueError as e:
                        last_err = w["last_check"] = str(e)
                        if w["retries_used"] >= max_retries:
                            stop = f"{last_err}, and it used its envelope's retry_limit ({max_retries})"
                            break
                        w["retries_used"] += 1
                        self.post("system", f"{fn} attempt {w['calls_used']} can't be used: {last_err}. Retry "
                                            f"{w['retries_used']} of {max_retries} (within its envelope).")
                head = raw.split("```")[0].strip()
                try:
                    meta = llm.extract_json(head)
                except ValueError:
                    meta = {"say": head}
                log = next((c for c in reversed(self.call_log) if c["fn"] == fn), None)
                if log:   # the structured part of a workflow agent's reply, for the detail view
                    log["reply"] = {"say": meta.get("say"), "next": meta.get("next"), "check": w.get("last_check"),
                                    "head_chars": len(got[0]) if got else 0, "section_chars": len(got[1]) if got else 0}
                    log["reason"] = meta.get("reason") or (self.MOCK_REASON if llm.MOCK else None)
                return got, raw, guide, meta, stop

            got, raw, guide, meta, stop = one_call()
            if got is None:      # code writes the failure report, so the unusable reply isn't route-checked
                self.write_lock = self.atom = None
                w["status"] = "failed"
                self._say(fn, f"@Coordination {fn} failed: {stop}. My section is unchanged.", "coordination", None,
                          workflow=fn)
                self._event(fn, "limit_fail", f"{fn}: {stop}")
                self.save("workflows.json", self.workflows)
                return self._handoff("coordination", kind="schedule")
            if j2_board.enabled() and isinstance(meta.get("question"), dict) and meta["question"].get("to"):
                entry = self._ask_board(w, meta["question"], attempt)
                if entry and entry["status"] == "answered":
                    got2, raw2, guide2, meta2, stop2 = one_call()
                    if got2 is not None:      # keep the original reply if the resume reply can't be used
                        got, raw, guide, meta, stop = got2, raw2, guide2, meta2, stop2
            board_read = self._credit_board_used(meta) if j2_board.enabled() else []
            if first:     # the style guide goes first in the head, so every section can use its variables and classes
                self.style = guide
                self.sections = {"style": {"owner": fn, "head": style_css(guide), "body": "", "workflow": "Style guide"},
                                 **self.sections}
                self.save("style.json", guide)
                w["style_guide"] = True
            self.sections[sid] = {"owner": fn, "head": got[0], "body": got[1], "workflow": w["name"]}
            title = (self.goal or {}).get("site_name") or "Site"
            html_out = assemble_doc(title, self.sections, self.out["sections"]) if self.doc else assemble(title, self.sections)
            self.sites.append(html_out)
            os.makedirs(os.path.join(self.dir, "site"), exist_ok=True)
            for name in ((f"site/{self.site_tag}v{len(self.sites)}.md", self.out.get("file")) if self.doc else
                         (f"site/v{len(self.sites)}.html", "site/index.html")):
                with open(os.path.join(self.dir, name), "w", encoding="utf-8") as f:
                    f.write(html_out)
            self.save("sections.json", self.sections)
            if self.blueprint and (verdict := self._spec_conflict(w, meta)):
                if verdict == "revised":
                    for a in w["atoms"]:      # the check now holds the revised slice
                        a["verification"] = [self._lld_check_for(w["lld_sections"]) if c.get("type") == "lld-slice" else c
                                             for c in a.get("verification") or []]
                else:
                    spec_stop = meta["spec_conflict"]
            failed = self._run_checks(w, attempt, html_out)
            if spec_stop:
                early = f"the builder reports a second conflict with LLD section {spec_stop.get('section')}"
                break
            if not failed or attempt >= MAX_CHECK_ATTEMPTS:
                break
            # Another attempt can't help when the same check fails on the same evidence again, or when the builder
            # says the cause is in another atom's section: Adaptation gets it now (runs/j2-20261015-114723, W-6).
            sig = {(r["check"], r["evidence"]) for r in failed}
            early = ("the same check failed with the same evidence twice" if sig & prev_sig else
                     "the builder says the cause is in another atom's section" if self._blames_other(meta.get("say"), fn)
                     else None)
            prev_sig = sig
            if early:
                break
            evidence = "\n".join(f"- {r['check']} ({r['type']} {r['target']}, expect {r['expect']}): {r['evidence']}"
                                 for r in failed)
            self.post("system", f"Code ran {fn}'s checks on attempt {attempt} of {MAX_CHECK_ATTEMPTS}: {len(failed)} "
                                f"failed. {fn} gets the evidence and tries again.")
        self.write_lock = self.atom = None
        if spec_stop:
            for a in w["atoms"]:
                a["status"] = "failed"
        nxt, note = "coordination", None      # a builder only reports to Coordination: code sets it, whatever "next" says
        human = self._human_output(w, meta, sid)
        self._atom_outputs(w, meta, attempt)
        w["status"], w["version"] = "done", len(self.sites)
        self.save("workflows.json", self.workflows)
        self._board_forward(w)      # built after a board change to a document it depends on: marked re-checked
        self._say(fn, f"@Coordination {human}" + (f" [read {', '.join(board_read)}]" if board_read else ""), nxt, note,
                  workflow=fn, human_output=human, version=len(self.sites), section=sid, workflow_cost=w["cost"],
                  **({"board_read": board_read} if board_read else {}))
        size = len(html_out.encode("utf-8"))
        if size > SITE_WARN_BYTES:
            self.post("system", f"Warning: site v{len(self.sites)} is {size / 1000:.0f} KB, over builder.md's ~60 KB "
                                "single-page rule.", warning=True)
        if failed or spec_stop:
            bad = sorted({r["atom"] for r in failed}) or [a["id"] for a in w["atoms"]]
            for a in w["atoms"]:
                if a["id"] in bad:
                    a["status"] = "failed"
            owner = (self._lld_owner_note(spec_stop) if spec_stop else self._owner_note(failed, sid)) if early else ""
            self.post("system", f"{', '.join(bad)} failed its checks on " + (f"attempt {attempt}: {early}, so no more "
                                f"attempts. {owner}" if early else f"all {MAX_CHECK_ATTEMPTS} attempts.") + " Code sends "
                                "the receipts to Adaptation, which decides whether the check or the build is wrong.")
            return self._handoff("adaptation", kind="atom_failed", wid=fn, atoms=bad, attempt=attempt, owner_note=owner)
        self._handoff("coordination", kind="schedule")

    def _artifact_inputs(self, w):
        """The finished artifact workflows this workflow depends on, directly or through others (a fix also reads
        what the workflow it revises read), in id order."""
        by = {k: x for x in self._all_wfs() for k in (x.get("id"), x.get("name")) if k}
        src = by.get(str(w.get("revises") or ""))
        stack, seen, out = list(w.get("dependencies") or []) + list((src or {}).get("dependencies") or []), set(), []
        while stack:
            d = by.get(str(stack.pop()))
            if not d or d["id"] in seen or d["id"] == w["id"]:
                continue
            seen.add(d["id"])
            stack += d.get("dependencies") or []
            if self._file_wf(d) and d.get("status") == "done":
                out.append(d)
        return sorted(out, key=lambda x: as_int(str(x["id"])[2:]))

    def _artifact_context(self, w):
        """The text of every artifact this workflow reads, for its prompt ("" when it reads none)."""
        parts = []
        for d in self._artifact_inputs(w):
            f = artifact_path(d)
            try:
                with open(os.path.join(self.dir, f), encoding="utf-8") as fh:
                    parts.append(f"--- {f} (written by {d['id']} {d['name']}) ---\n{fh.read()[:6000]}")
            except OSError:
                continue
        return ("\nARTIFACTS YOU READ (written by the workflows you depend on; follow their 'Rules for later "
                "workflows'):\n" + "\n".join(parts)) if parts else ""

    def _run_artifact(self, w, brief, max_calls):
        """An artifact workflow's agent writes one file (its output, named by Adaptation) instead of a section of the
        site. Code runs its work atom's checks on the file as text, up to MAX_CHECK_ATTEMPTS attempts, then reports to
        Coordination like any workflow (or hands the receipts to Adaptation after the last failed attempt)."""
        fn, f = w["id"], artifact_path(w)
        path = os.path.join(self.dir, f)
        fmt = ('First line: a JSON object {"say": str (one short line to @Coordination: what you wrote; code runs your '
               'checks, so don\'t claim they pass), "human_output": str (1 to 3 plain sentences telling a person what '
               'this file contains), "atom_human_output": str (1 or 2 plain sentences in everyday words: no ids, no '
               'selectors), "next": "Coordination"}. Then the whole file between a line <<<FILE and a line FILE>>>.')
        g = self.goal or {}
        charter = {k: g.get(k) for k in ("statement", "audience", "standards", "rules", "done_when", "to_be_confirmed")}

        def mock():
            return (json.dumps({"say": f"@Coordination {fn} wrote {f}.", "human_output": f"I wrote {f} for the "
                                "workflows after it.", "atom_human_output": "This sets what the later parts of the "
                                "page follow.", "next": "Coordination"})
                    + f"\n<<<FILE\n# {w['name']}\n\nWhat {w['name']} sets for the workflows after it, from the charter.\n"
                      "FILE>>>")
        attempt, evidence, text, meta, failed = 0, "", None, {}, []
        while attempt < MAX_CHECK_ATTEMPTS and w["calls_used"] < max_calls + attempt:
            attempt += 1
            w["calls_used"] += 1
            raw = self.call(fn, fmt, f"YOU ARE {NAMES[fn]}, the agent for one registered workflow.\nYOUR WORKFLOW: "
                            f"{json.dumps(brief, ensure_ascii=False)}\nYOUR OUTPUT IS ONE FILE, {f}, which the "
                            "workflows after you read and follow. It builds no part of the site. Write what your "
                            "workflow's goal and work atom ask for, in the format the file's name says. Mark anything "
                            "the charter doesn't give as an assumption. Code runs your work atom's checks on the file."
                            f"\nCHARTER: {json.dumps(charter, ensure_ascii=False)}" + self._artifact_context(w)
                            + (f"\nATTEMPT {attempt} OF {MAX_CHECK_ATTEMPTS}. THESE CHECKS FAILED ON YOUR LAST "
                               f"ATTEMPT:\n{evidence}\nFix the file so they pass." if evidence else ""),
                            mock, max_tokens=8000, json_reply=False)
            body = re.search(r"^<<<FILE[ \t]*\n(.*?)\n?^FILE>>>", raw or "", re.S | re.M)
            head = str(raw or "").split("<<<FILE")[0].strip()
            try:
                meta = llm.extract_json(head)
            except ValueError:
                meta = {"say": head}
            if not body:
                evidence = "- the reply had no file between a line <<<FILE and a line FILE>>>"
                continue
            text = body[1].strip() + "\n"
            os.makedirs(os.path.dirname(path), exist_ok=True)
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(text)
            failed = self._run_checks(w, attempt, text, doc=True)
            if not failed:
                break
            evidence = "\n".join(f"- {r['check']} ({r['type']} {r['target']}, expect {r['expect']}): {r['evidence']}"
                                 for r in failed)
        self.write_lock = self.atom = None
        if text is None:
            w["status"] = "failed"
            self._say(fn, f"@Coordination {fn} failed: no usable artifact after {attempt} attempt(s). {f} was not "
                          "written.", "coordination", None, workflow=fn)
            self._event(fn, "limit_fail", f"{fn}: no usable artifact")
            self.save("workflows.json", self.workflows)
            return self._handoff("coordination", kind="schedule")
        human = self._human_output(w, meta, w["section"])
        self._atom_outputs(w, meta, attempt)
        w["status"], w["version"], w["artifact"] = "done", len(self.sites), f
        self.save("workflows.json", self.workflows)
        self._say(fn, f"@Coordination {human}", "coordination", None, workflow=fn, human_output=human,
                  version=len(self.sites), section=w["section"], workflow_cost=w["cost"], artifact=f)
        if failed:
            bad = sorted({r["atom"] for r in failed})
            for a in w["atoms"]:
                if a["id"] in bad:
                    a["status"] = "failed"
            self.post("system", f"{', '.join(bad)} failed its checks on all {attempt} attempts. Code sends the receipts "
                                "to Adaptation, which decides whether the check or the artifact is wrong.")
            return self._handoff("adaptation", kind="atom_failed", wid=fn, atoms=bad, attempt=attempt, owner_note="")
        self._handoff("coordination", kind="schedule")

    def _human_output(self, w, meta, sid):
        """1 to 3 plain sentences saying what the builder built. "done." alone is rejected and asked for once more
        (a short call about the section just written); if that fails too, code writes the sentence."""
        fn, last = w["id"], str(meta.get("human_output") or "")
        if human_problem(last):
            self._event(fn, "invalid_move", f"human_output: {human_problem(last)}")
            sec = (self.sections.get(sid) or {}).get("body") or ""
            r = self.call(fn, '{"human_output": str (1 to 3 plain sentences)}',
                          f"YOUR REPORT '{last}' CAN'T BE USED: {human_problem(last)}. Tell a person in 1 to 3 plain "
                          f"sentences what you built in section {sid} ({w['name']}). Never just \"done.\"\nYOUR "
                          f"SECTION:\n{sec[:3000]}",
                          lambda: json.dumps({"human_output": f"I built the {w['name']} section ({sid}) of the "
                                                              f"{self.out['noun']}."}))
            last = str(r.get("human_output") or "") if isinstance(r, dict) else ""
        if human_problem(last):
            self._event(fn, "invalid_move", f"human_output again: {human_problem(last)}")
            last = f"{fn} built {w['name']} in section {sid}. Code ran its checks; see the receipts."
        return last

    def _atom_outputs(self, w, meta, attempt):
        """Every atom's human_output (the builder's words, checked by atom_human_problem, asked once more if bad, else
        written by code from the atom's expected output), then, for an atom whose checks passed, its agent_output."""
        for n, a in enumerate(w.get("atoms") or []):
            text = str(meta.get("atom_human_output") or "") if n == 0 else ""
            if atom_human_problem(text):
                if n == 0:
                    self._event(w["id"], "invalid_move", f"atom human_output: {atom_human_problem(text)}")
                    r = self.call(w["id"], '{"atom_human_output": str (1 or 2 plain sentences)}',
                                  f"YOUR atom_human_output '{text}' CAN'T BE USED: {atom_human_problem(text)}. Tell the "
                                  "Requester in 1 or 2 plain sentences what this part created. No ids, no selectors, "
                                  f"no \"units\". The part: {a.get('goal')}",
                                  lambda: json.dumps({"atom_human_output": "This adds a part of the page that visitors "
                                                                           "can read and use."}))
                    text = str(r.get("atom_human_output") or "") if isinstance(r, dict) else ""
                if atom_human_problem(text):
                    text = fallback_human(a, self.out["noun"])
            a["human_output"] = {"atom": a["id"], "attempt": attempt, "text": text}
        self._seal_atoms(w, w.get("atoms") or [], attempt)

    def _agent_output(self, w, a, attempt):
        """What other agents read instead of searching: built by code from the section as written and the receipts of
        this attempt. Files, section names, element ids and selectors created, the checks and their results, the
        receipt id (receipts.jsonl: atom and attempt). An artifact workflow: its file and headings."""
        if self._file_wf(w):
            f = artifact_path(w)
            try:
                with open(os.path.join(self.dir, f), encoding="utf-8") as fh:
                    text = fh.read()
            except OSError:
                text = ""
            recs = [r for r in self.receipts if r["atom"] == a["id"] and r["attempt"] == attempt and "check" in r]
            return {"atom": a["id"], "attempt": attempt, "receipt": f"{a['id']}:{attempt}", "kind": "artifact",
                    "file": f, "files": [f], "section_names": [w["name"]] + [m[2] for ln in text.splitlines()
                                                                             if (m := MD_HEADING.match(ln))][:6],
                    "checks": [{"check": r["check"], "type": r["type"], "target": r["target"], "pass": r["pass"]}
                               for r in recs]}
        return build_agent_output(w, a, attempt, (self.sections.get(w["section"]) or {}).get("body") or "", self.receipts,
                                  self.audits[-1] if self.audits else None, len(self.sites), self.doc, self.out.get("file"))

    def _seal_atoms(self, w, atoms, attempt):
        """An atom whose checks passed gets its agent_output; it stays "passed" only with both outputs (else "failed")."""
        for a in atoms:
            if a.get("status") != "passed":
                continue
            a["agent_output"] = self._agent_output(w, a, attempt)
            if (miss := atom_done_problem(a)):
                a["status"] = "failed"
                self._event(w["id"], "refused", f"{a['id']} can't be done without {', '.join(miss)}")
                continue
            os.makedirs(self.dir, exist_ok=True)
            with open(os.path.join(self.dir, "outputs.jsonl"), "a", encoding="utf-8") as f:
                f.write(json.dumps({"atom": a["id"], "attempt": attempt, "human_output": a["human_output"],
                                    "agent_output": a["agent_output"]}, ensure_ascii=False) + "\n")
        self.save("workflows.json", self.workflows)

    SPEC_FMT = ('"spec_conflict": null|{"section": str (the LLD section id), "reason": str, "use": {"#spec-selector": '
                '"#selector-you-built"}, "fields": {"F4": null|"text"} (optional: null removes that spec field because the charter has no '
                'such field, a string gives its new type)} (only when your spec can\'t be built as written; never deviate '
                'silently), ')

    BOARD_FMT = ('"question": null|{"to": str (an atom id from OTHER SECTIONS or BOARD), "about": str, "text": str} '
                 '(optional: the one atom that knows; at most one question per atom, never repeat an already-answered '
                 'question), "board_used": null|[str] (optional: the board ids, e.g. ["Q-1"], whose answer you used), ')
    # The same two fields for the five functions (added to every function format in call()). The question field is
    # called board_question because every function format already uses "question"/"questions" for the Requester or
    # for a policy question; code reads either key but only ever acts on the {to, about, text} shape.
    FN_BOARD_FMT = ('"board_question": null|{"to": str (the workflow id, atom id or function name that knows), '
                    '"about": str, "text": str} (optional: ask the one that knows instead of guessing; at most one '
                    'question per turn, and never one the BOARD already answers), "board_used": null|[str] '
                    '(optional: the board ids, e.g. ["Q-1"], whose answer you used)')

    def _hld_summary(self):
        comps = (self.chain.get("hld") or {}).get("components")
        return ("\nHLD SUMMARY (components): " + json.dumps([{k: c[k] for k in ("id", "name", "responsibility")} for c in comps],
                                                            ensure_ascii=False)) if comps else ""

    def _page_context(self, w, prev):
        """What a builder sees of the rest of the site. Plain runs: the whole page (J2_SLIM=1: the manifest). With an
        LLD: its own slice of it (its spec) and a short HLD summary instead of the page; J2_SLIM=1 adds the manifest."""
        manifest = ("\nOTHER SECTIONS (their atoms' agent_output; the rest of the page isn't shown): "
                    f"{json.dumps(self._manifest(w), ensure_ascii=False)}")
        lld = bool(self.blueprint and w.get("lld_sections"))
        self._record_given(w, (self._slice_edges(w) if lld else []) + self._hld_edges()
                           + (self._manifest_edges(w) if slim() else []) + self._board_edges(),
                           whole=not lld and not slim())
        board = self._board_context()
        if self.blueprint and w.get("lld_sections"):
            return ("\nYOUR SPEC (the LLD sections you implement: build exactly these ids, selectors and form fields with "
                    "their behaviour; if one can't be built as written, report it in spec_conflict, never deviate "
                    f"silently): {json.dumps(self._lld_slice_for(w), ensure_ascii=False)}" + self._hld_summary()
                    + (manifest if slim() else "") + board)
        if slim():
            return manifest + self._hld_summary() + board
        return f"\nTHE WHOLE SITE FILE NOW (v{len(self.sites)}, read only):\n{prev[:SITE_LIMIT]}" + self._hld_summary() + board

    # ---- atom/1 edges: what a builder's prompt actually held, recorded where the prompt is made (no prompt changes)
    def _record_given(self, w, edges, whole=False):
        self.given[w["id"]] = {"edges": edges, "whole": whole}

    def _manifest_edges(self, w):
        return [{"atom": x["id"], "why": "its section, anchor and selectors were listed in this builder's prompt"}
                for o in self.workflows if o["id"] != w["id"] for x in o.get("atoms") or [] if x.get("agent_output")]

    def _doc_atoms(self, kind, heading):
        ids = (self.docs.get(kind) or {}).get("workflows") or []
        return [a["id"] for o in self._all_wfs() if o["id"] in ids and heading.lower() in str(o.get("name")).lower()
                for a in o.get("atoms") or []]

    def _slice_edges(self, w):
        sl = self._lld_slice_for(w) or {}
        ids = ", ".join(w.get("lld_sections") or [])
        return [{"atom": a, "why": f"LLD slice {ids}: its {h.lower()} were in this builder's prompt"}
                for key, h in (("sections", "Pages and sections"), ("elements", "Interactive elements"), ("fields", "Form fields"))
                if sl.get(key) for a in self._doc_atoms("lld", h)]

    def _hld_edges(self):
        return [{"atom": a, "why": "the HLD component summary was in this builder's prompt"}
                for a in (self._doc_atoms("hld", "Components") if (self.chain.get("hld") or {}).get("components") else [])]

    def _chain_edges(self):
        """A document builder is given the charter, which carries the earlier documents' agent_output."""
        ch = (self.goal or {}).get("chain") or {}
        return [{"atom": a, "why": why} for kind, h, why in (("prd", "Requirements", "the PRD's requirements were in this builder's "
                                                              "prompt (charter)"),
                                                             ("hld", "Components", "the HLD's components were in this builder's "
                                                              "prompt (charter)")) if ch.get(kind) for a in self._doc_atoms(kind, h)]

    def _lld_slice_for(self, w):
        return lld_slice(self.blueprint, w.get("lld_sections") or []) if self.blueprint else None

    def _manifest(self, w):
        """J2_SLIM: other atoms' agent_output, short (ids, sections, selectors), so a builder needn't read the page."""
        return [{k: (x.get("agent_output") or {}).get(k) for k in ("atom", "section", "anchor", "selectors")}
                for o in self.workflows if o["id"] != w["id"] for x in o.get("atoms") or [] if x.get("agent_output")]

    # ---- message board (board/1, j2_board.py). Askers: the five functions (their format gets FN_BOARD_FMT in call())
    # and the builders (BOARD_FMT). Targets: a workflow, including a done one from an earlier cycle, or another
    # function. _ask_board is Ask and Answer, _board_turn is one function's turn, _credit_board_used is Reuse.
    def _board_context(self):
        if not j2_board.enabled():
            return ""
        answered = [e for e in self.board["entries"] if e["status"] == "answered"]
        if not answered:
            return ""
        return ("\nBOARD (answered questions; check here before asking): "
                + json.dumps([{"id": e["id"], "to": e["to"], "about": e["about"], "answer": e["answer"]}
                              for e in answered], ensure_ascii=False))

    def _board_edges(self):
        """depends_on from every atom a board answer came from, recorded the same way manifest/HLD/LLD edges are:
        what the prompt held, not whether it was used (used_by only follows a named board_used). Only atoms: an
        answer from a function is not an atom, so it carries no atom/1 edge."""
        if not j2_board.enabled():
            return []
        atoms = self._atoms_v1()
        return [{"atom": e["to"], "why": f"board {e['id']}"} for e in self.board["entries"]
                if e["status"] == "answered" and e["to"] in atoms]

    def _board_target(self, to):
        """What a question names: ("function", name) for one of the five, ("workflow", workflow) for a workflow id or
        one of its atom ids (any cycle of this run, and the earlier stages' workflows), else (None, None)."""
        to = str(to or "")
        if to.lower() in FUNCTIONS:
            return "function", to.lower()
        w = next((x for x in self._all_wfs() if x["id"] == to.split(".")[0]), None)
        return ("workflow", w) if w else (None, None)

    def _ask_board(self, asker, question, attempt, role=None):
        """Items 2 (Ask) and 3 (Answer), for a function (asker: its name) or a builder (asker: its workflow). The
        question is refused, with the reason posted, when it names no target, when a builder's atom already asked,
        when the board already answers it (60% of its content words, j2_board.duplicate_of) or when the target isn't
        a built workflow or another function. A valid question gets one short answer-only call, which costs no
        credits: to the target's builder, given its recorded section and agent_output, or to the target function.
        If a builder's answer corrects its own output, _apply_board_update bumps the site version ("Updated by Q-n")."""
        builder = isinstance(asker, dict)
        role = role or ("builder" if builder else "function")      # what the entry records as asked_by_role
        asked_by = asker["id"] if builder else asker
        asking = self.atom if builder else asked_by      # the entry's own side: an atom id, or the function's name
        label = NAMES.get(asked_by, asked_by)
        to = question.get("to")
        about, text = str(question.get("about") or ""), str(question.get("text") or "")
        kind, target = self._board_target(to)
        atoms_v1 = self._atoms_v1()
        dup = j2_board.duplicate_of(text, self.board["entries"])
        reason = ("max one question per atom; this atom already asked"
                      if builder and any(e["atom"] == asking for e in self.board["entries"]) else
                  "max one question per turn; this function already asked this turn"
                      if not builder and any(e["asked_by"] == asked_by and e["cycle"] == self.cycle
                                             and e["attempt"] == attempt for e in self.board["entries"]) else
                  "no target named" if not to else
                  f"already answered in {dup}" if dup else
                  "the target isn't a workflow of this run or one of the five functions" if not target else
                  "a function can't ask itself" if kind == "function" and target == asked_by else
                  "the target atom isn't in the manifest" if kind == "workflow" and "." in str(to) and to not in atoms_v1 else
                  "the target atom isn't built yet"
                      if kind == "workflow" and "." in str(to) and atoms_v1[to]["status"] != "closed" else
                  "the target workflow isn't built yet"
                      if kind == "workflow" and "." not in str(to) and target.get("status") != "done" else None)
        entry = j2_board.new_entry(f"Q-{len(self.board['entries']) + 1}", asked_by, asking, to, about, text,
                                   self.cycle, attempt, asked_by_role=role)
        self.board["entries"].append(entry)
        if reason:
            entry["status"] = "refused"
            j2_board.write(self.dir, self.board)
            j2_board.log_event(self.dir, {"event": "refused", "id": entry["id"], "reason": reason})
            self.post("system", f"Board {entry['id']}: {asked_by}'s question to {to} refused — {reason}.")
            return entry
        j2_board.write(self.dir, self.board)
        j2_board.log_event(self.dir, {"event": "asked", "id": entry["id"], "asked_by": asked_by, "role": role,
                                      "atom": asking, "to": to, "about": about})
        answerer = target if kind == "function" else target["id"]
        self.post(asked_by, f"@{NAMES.get(answerer, answerer)} {entry['id']}: {text}")
        reply = self._board_answer(entry, label, kind, target)
        answer = str((reply or {}).get("answer") or "").strip()
        if not answer or answer.lower() in ("not mine", "n/a", "none"):
            entry["status"] = "unanswered"      # the existing Adaptation path takes it from here, same as any block
        else:
            entry["status"], entry["answer"], entry["answered_by"] = "answered", answer, answerer
            entry["answered_from"] = ({"function": answerer} if kind == "function" else
                                      {"workflow": target["id"], "version": target.get("version")})
            update = (reply or {}).get("artifact_update")
            doc_up = kind == "workflow" and isinstance(update, dict) and update.get("doc")
            updated = kind == "workflow" and isinstance(update, dict) and not doc_up
            if doc_up:
                # An answer that changes a document (item 1): only the owner of that section may, and a PRD change
                # is a change request for Identity, never applied here.
                why = self._board_doc_update(entry, target, update)
                if why:
                    entry["status"] = "refused"
                    j2_board.write(self.dir, self.board)
                    j2_board.log_event(self.dir, {"event": "refused", "id": entry["id"], "reason": why})
                    self.post("system", f"Board {entry['id']}: {answerer}'s change to the "
                                        f"{str(update.get('doc')).upper()} refused — {why}.")
                    return entry
            if updated:
                entry["artifact_update"] = update
                self._apply_board_update(target["id"], update, entry["id"])
            self.post(answerer, f"@{label} {entry['id']}: {answer}"
                      + (" (updated my section)" if updated else
                         f" (changed the {str(update.get('doc')).upper()})" if doc_up and entry["artifact_update"] else
                         " (sent to Identity)" if doc_up else ""))
        j2_board.write(self.dir, self.board)
        j2_board.log_event(self.dir, {"event": entry["status"], "id": entry["id"],
                                      "answered_from": entry["answered_from"]})
        return entry

    def _board_answer(self, entry, label, kind, target):
        """The Answer call (item 3): one short answer-only call, no question field, no credits, counted as an AI
        call. A workflow answers from its recorded section and agent_output; a function answers from the run's
        record. Nothing recorded means "not recorded", never a guess."""
        ask = f"{label} ASKS YOU ABOUT {entry['about']}: {entry['question']}\n"
        rules = ('Answer from the record only, in one or two short sentences: say "not recorded" for anything this '
                 'run has not recorded, and "not mine" if this isn\'t your area. No credits are charged for this '
                 "reply, and you may not ask anything back.")
        if kind == "function":
            return self.call(target, '{"answer": str (plain words; say "not mine" if this isn\'t your area)}',
                             ask + rules, lambda: json.dumps(self._mock_board_answer_fn(target, entry)),
                             answer_only=True) or {}
        sec = self.sections.get(target.get("section")) or {}
        recorded = {"workflow": target["id"], "name": target.get("name"), "section": target.get("section"),
                    "status": target.get("status"), "version": target.get("version"), "cycle": target.get("cycle"),
                    "atoms": [{"atom": a["id"], "agent_output": a.get("agent_output"),
                               "human_output": a.get("human_output")} for a in target.get("atoms") or []]}
        # A document's section: its own text, read back from the run's record (the current stage's sections.json, or
        # the stage file of a document that is already filed).
        kind = self._doc_of_workflow(target["id"])
        body = atom_section_body(self.dir, self._all_wfs(), (target.get("atoms") or [{}])[0].get("id")) if kind else None
        return self.call(target["id"],
                         '{"answer": str (plain words; say "not mine" if this isn\'t your area), "artifact_update": '
                         + (f'null|{{"doc": "{kind}", "section": str (the heading you own), "text": str (the whole new '
                            'section in markdown, from the record only), "reason": str} (only if your own section of '
                            'the document needs correcting)}' if kind else
                            'null|{"body": str, "head": str} (only if your own section needs correcting)}'),
                         ask + f"YOUR RECORDED WORK: {json.dumps(recorded, ensure_ascii=False)}\n"
                         + (f"YOUR SECTION OF {kind.upper()} AS RECORDED:\n```markdown\n{body}\n```\n" if kind and body
                            else f"YOUR SECTION AS RECORDED ({target.get('section')}):\n```head\n{sec.get('head') or ''}"
                            f"\n```\n```section\n{sec.get('body') or ''}\n```\n" if sec and not kind else
                            "YOUR SECTION AS RECORDED: not recorded.\n") + rules,
                         lambda: json.dumps(self._mock_board_answer(target["id"], entry)),
                         answer_only=True) or {}

    def _board_turn(self, fn, reply):
        """One function's board turn, from call(): any board id it names is credited (and its next message tagged
        "[read Q-n]"), then its one question is asked. At most one question per turn, by the shape of the format."""
        if llm.MOCK and not reply.get("board_used"):
            reply.update(self._mock_board_reuse(fn))
        q = reply.get("board_question") or reply.get("question")
        if isinstance(q, dict) and q.get("to"):
            # Asked first, so the tag below lands on the function's own message. The turn it asked in is recorded as
            # the entry's attempt, which is how "at most one question per turn" is checked.
            self._ask_board(fn, q, self.agent_turns.get(fn, 0), role="function")
        read = self._credit_board_used(reply, fn)
        if read:
            self.board_read[fn] = sorted(set((self.board_read.get(fn) or []) + read))

    def _apply_board_update(self, wid, update, entry_id):
        """An answer's artifact_update that changes the target's own section (website builds only; a document
        section's update is recorded on the board but not re-applied, and a past stage's atom is out of reach since
        _wf only finds the current stage's workflows): reassembles the page with the correction, bumping the site
        version, and notes why on the atom's human_output."""
        w = self._wf(wid)
        if not w or self.doc:
            return
        sid = w["section"]
        cur = self.sections.get(sid) or {}
        body, head = update.get("body", cur.get("body")), update.get("head", cur.get("head"))
        if body == cur.get("body") and head == cur.get("head"):
            return
        self.sections[sid] = {**cur, "body": body, "head": head}
        title = (self.goal or {}).get("site_name") or "Site"
        html_out = assemble(title, self.sections)
        self.sites.append(html_out)
        os.makedirs(os.path.join(self.dir, "site"), exist_ok=True)
        with open(os.path.join(self.dir, f"site/v{len(self.sites)}.html"), "w", encoding="utf-8") as f:
            f.write(html_out)
        with open(os.path.join(self.dir, "site/index.html"), "w", encoding="utf-8") as f:
            f.write(html_out)
        self.save("sections.json", self.sections)
        for a in w.get("atoms") or []:
            if a.get("agent_output"):
                ho = a.get("human_output") or {}
                a["human_output"] = {**ho, "text": f"{ho.get('text') or ''} Updated by {entry_id}.".strip()}
        self.save("workflows.json", self.workflows)

    # ---- a board answer that changes a document (board/1 part C): only the owner of that section may change it, a
    # PRD change is a change request for Identity, and the change is re-checked and propagated by recorded edges.
    def _doc_of_workflow(self, wid):
        """The document a workflow wrote (documents.json), or None for a website workflow."""
        return next((k for k, d in self.docs.items() if wid in (d.get("workflows") or [])), None)

    def _doc_of_atom(self, aid):
        """The document an atom wrote (documents.json), or None for a website atom."""
        return next((k for k, d in self.docs.items() if aid in (d.get("atoms") or [])), None)

    def _doc_version(self, kind):
        """A document's latest version number, from the version files the run wrote (site/hld-v3.md, or v3.md in a
        one-document run). 0 when the run recorded none."""
        site = os.path.join(self.dir, "site")
        ns = [int(m[1]) for f in (os.listdir(site) if os.path.isdir(site) else [])
              for p in (rf"{kind}-v(\d+)\.md", r"v(\d+)\.md") if (m := re.fullmatch(p, f))]
        return max(ns, default=0)

    def _board_doc_owner(self, aid, kind, section):
        """Whether the answering atom owns that section: the atom is one of that document's atoms and its workflow
        writes that heading (documents.json and workflows.json, never a guess)."""
        d, want = self.docs.get(kind) or {}, squash(strip_md(section))
        if not want or aid not in (d.get("atoms") or []):
            return False
        w = next((x for x in self._all_wfs() if any(a["id"] == aid for a in x.get("atoms") or [])), None)
        return bool(w) and w["id"] in (d.get("workflows") or []) and squash(w.get("name") or "").endswith(want)

    def _board_doc_update(self, entry, target, update):
        """Item 1: the reason the change can't be applied, or None once it has been applied. A PRD change is never
        applied here: it becomes a change request for Identity (the existing relay path) and is recorded as such."""
        kind, section = str(update.get("doc") or "").lower(), str(update.get("section") or "")
        text, why = str(update.get("text") or ""), str(update.get("reason") or "").strip()
        aid = (target.get("atoms") or [{}])[0].get("id")
        if kind not in self.docs:
            return f"{kind or 'that document'} is not a document of this run"
        if section not in OUTPUTS[kind]["sections"]:
            return f"{section!r} is not a section of the {kind.upper()}"
        if not self._board_doc_owner(aid, kind, section):
            return f"{aid} doesn't own the {kind.upper()}'s {section} section"
        if not text.strip():
            return "the change has no new text"
        if kind == "prd":
            # The PRD is the Requester's own charter of requirements: Identity decides, the board never rewrites it.
            entry["artifact_update"] = None
            entry["sent_to_identity"] = {"doc": "prd", "section": section, "reason": why or "not recorded"}
            self.board_asks.append({"frm": target["id"],
                                    "question": f"{entry['id']}: the PRD's {section} section needs changing "
                                                f"({why or 'no reason recorded'}). The board doesn't change the PRD, so "
                                                "this is a change request for you."})
            self.post("system", f"Board {entry['id']}: a PRD change is never applied from the board. It goes to "
                                f"Identity as a change request ({section}).", warning=True)
            return None
        v0, v1 = self._revise_doc(kind, section, text, why or f"answered {entry['id']}", entry["id"])
        if v1 is None:
            return f"{kind}.md is not in the run folder"
        entry["artifact_update"] = {"doc": kind, "section": section, "from_version": v0, "to_version": v1}
        self._board_note(aid, entry["id"])
        passed, sent = self._board_propagate(entry, aid, kind, v1)
        self.post("system", f"{kind.upper()} v{v0} -> v{v1} ({entry['id']}). Re-checked: {passed} passed, "
                            f"{sent} sent to Adaptation.", checks=target["id"])
        return None

    def _revise_doc(self, kind, section, text, reason, qid):
        """Item 2: replace one section of a filed document, append its version note, export it again and refresh
        everything that reads it. Returns (from_version, to_version), or (None, None) when the file isn't there."""
        path = os.path.join(self.dir, f"{kind}.md")
        if not os.path.isfile(path):
            return None, None
        with open(path, encoding="utf-8") as f:
            old = f.read()
        v0 = self._doc_version(kind)
        v1 = v0 + 1
        body = re.sub(r"^\s*#{1,6}\s+[^\n]*\n?", "", str(text).strip("\n"), count=1)      # the heading line stays
        new = md_replace_section(old, section, body.strip("\n") + f"\n\nVersion {v1}: {qid}, {reason}.")
        if new == old:
            return None, None
        self._write_doc(kind, new, version=v1)
        self.doc_changes.append({"id": qid, "doc": kind, "section": section, "from": v0, "to": v1, "reason": reason})
        return v0, v1

    def _write_doc(self, kind, text, version=None):
        """Write a document the run has already filed: <kind>.md, its new version file, the export, and the records
        that read it (documents.json's agent_output, the chain, the charter's chain, the LLD blueprint)."""
        with open(os.path.join(self.dir, f"{kind}.md"), "w", encoding="utf-8") as f:
            f.write(text)
        if version:
            os.makedirs(os.path.join(self.dir, "site"), exist_ok=True)
            tag = f"{kind}-" if len(self.stages) > 1 else ""
            with open(os.path.join(self.dir, "site", f"{tag}v{version}.md"), "w", encoding="utf-8") as f:
                f.write(text)
        if OUTPUTS[kind].get("export"):
            OUTPUTS[kind]["export"](text, self.dir)
        ao = doc_agent_output(kind, text)
        if kind in self.docs:
            self.docs[kind] = {**self.docs[kind], "agent_output": ao}
            self.save("documents.json", self.docs)
        if kind in self.chain:
            self.chain[kind] = ao
        ch = (self.goal or {}).get("chain")
        if isinstance(ch, dict) and kind in ch:      # the charter every later builder reads
            ch[kind] = ao
            ch[f"{kind}_ids"] = [x["id"] for x in (ao.get("requirements") or ao.get("components") or [])]
            self.save("goal.json", self.goal)
        if kind == "lld" and self.blueprint:
            self.blueprint = copy.deepcopy(ao)
        return ao

    def _board_note(self, aid, qid):
        """"Updated by Q-k" on the atom's own note (atoms.json reads it), once."""
        for w in self._all_wfs():
            for a in w.get("atoms") or []:
                if a["id"] == aid:
                    ho = a.get("human_output") or {}
                    if f"Updated by {qid}" not in str(ho.get("text") or ""):
                        a["human_output"] = {**ho, "text": f"{ho.get('text') or ''} Updated by {qid}.".strip()}
                    self.save("workflows.json", self.workflows)

    def _board_propagate(self, entry, aid, kind, version):
        """Items 3 and 4: re-run the changed document's conditions and the owner atom's checks, then follow the
        recorded used_by edges (nothing else) and re-check every atom that is already closed. A failing atom goes to
        Adaptation on the existing atom_failed path. Returns (passed, sent to Adaptation)."""
        qid = entry["id"]
        self._doc_conditions(kind, qid, version)
        views = {v["id"]: v for v in j2_schema.run_views(self._all_wfs(), self.given)[0]}
        seen, queue, impact = set(), [aid], []
        while queue:                       # transitively, by recorded edges only
            cur = queue.pop(0)
            if cur in seen:
                continue
            seen.add(cur)
            queue += [e["atom"] for e in (views.get(cur) or {}).get("used_by") or []]
        results, sent = [], []
        for x in [aid] + [i for i in seen if i != aid]:
            a, w = self._atom_of(x)
            if not a or a.get("status") != "passed":      # not built yet: its own build checks it (_board_forward)
                continue
            a["board_stale"] = qid
            page, doc = self._recheck_page(x)
            failed = self._run_checks(w, (a.get("attempts") or 1), page, atoms=[a], doc=doc, trigger=qid)
            ok = not failed
            if ok:
                a.pop("board_stale", None)
                a["board_rechecked"] = qid
            else:
                sent.append((w, a))
            if x != aid:      # impact is what the change reached; the owner's own re-check is item 3
                impact.append({"atom": x, "result": "passed" if ok else "failed"})
            results += (a.get("results") or [])
        entry["impact"] = impact
        self.audits_recheck(qid, kind, version, results)
        self.save("workflows.json", self.workflows)
        j2_board.write(self.dir, self.board)
        for w, a in sent:
            # The existing path, after this turn (_run): Adaptation reads the receipts and prices a revise as a fix.
            self._event(w["id"], "board_recheck_failed", f"{a['id']} failed its checks after {qid}")
            self.board_failed.append({"wid": w["id"], "atoms": [a["id"]], "attempt": a.get("attempts") or 1,
                                      "owner_note": f"its checks failed after the board change {qid}"})
        return sum(1 for x in impact if x["result"] == "passed"), len(sent)

    def _atom_of(self, aid):
        """(atom, workflow) for an atom id of this run, else (None, None)."""
        for w in self._all_wfs():
            for a in w.get("atoms") or []:
                if a["id"] == aid:
                    return a, w
        return None, None

    def _recheck_page(self, aid):
        """(page, doc) an atom's own checks run against: the document its stage wrote, or what is built now."""
        kind = self._doc_of_atom(aid)
        if kind:
            path = os.path.join(self.dir, f"{kind}.md")
            if os.path.isfile(path):
                with open(path, encoding="utf-8") as f:
                    return f.read(), True
        return (self.sites[-1] if self.sites else ""), self.doc

    def _doc_conditions(self, kind, qid, version):
        """Item 3: the changed document's own necessary conditions (P1-P8 / H1-H6 / L1-L6), re-run by code on the new
        text with that stage's charter. Saved as conditions-<kind>.json; the live run's conditions.json is the
        current stage's and is left alone."""
        path = os.path.join(self.dir, f"{kind}.md")
        if not os.path.isfile(path):
            return []
        with open(path, encoding="utf-8") as f:
            text = f.read()
        goal = next((s["goal"] for s in reversed(self.stage_archive) if s.get("stage") == kind), self.goal)
        conds = check_conditions(load_conditions(OUTPUTS[kind]["conditions"]), text, goal,
                                 self._fact_source(goal), OUTPUTS[kind]["sections"])
        self.save(f"conditions-{kind}.json", conds)
        self.post("system", f"{kind.upper()} v{version} conditions after {qid} (checked by code): " + "; ".join(
            f"{c['id']} " + ("met" if c["met"] else f"NOT MET ({c['evidence']})") for c in conds), mechanical=True)
        return conds

    def _save_audits(self):
        """audits.json: the audits, then the board's re-check entries (board/1 part C). Readers of the file skip
        "kind": "recheck"; with the board off the file is the audits alone, as before."""
        self.save("audits.json", self.audits + self.rechecks)

    def audits_recheck(self, qid, kind, version, results):
        """Item 3: the recheck's own entry in audits.json. Kept apart from the audits themselves (self.audits), so
        nothing that reads the latest audit sees it; readers of the file skip "kind": "recheck"."""
        self.rechecks.append({"kind": "recheck", "trigger": qid, "doc": kind, "version": version,
                              "t": time.strftime("%Y-%m-%dT%H:%M:%S"),
                              "results": [{k: r.get(k) for k in ("atom", "check", "pass", "evidence")} for r in results]})
        self._save_audits()

    def _board_forward(self, w):
        """Item 4 for an atom built after the change: when its recorded edges reach an atom a board answer changed,
        its own checks have just run against the changed document, so it is marked re-checked and listed in the
        entry's impact. Recorded edges only."""
        if not (j2_board.enabled() and self.doc_changes):
            return
        views = {v["id"]: v for v in j2_schema.run_views(self._all_wfs(), self.given)[0]}

        def reaches(aid, want, seen=None):
            seen = seen if seen is not None else set()
            if aid in seen:
                return False
            seen.add(aid)
            deps = [e["atom"] for e in (views.get(aid) or {}).get("depends_on") or []]
            return want in deps or any(reaches(d, want, seen) for d in deps)
        for a in w.get("atoms") or []:
            if a.get("status") != "passed":
                continue
            for e in self.board["entries"]:
                qid, up = e["id"], e.get("artifact_update") or {}
                if not up.get("doc") or a.get("board_rechecked") == qid or any(x["atom"] == a["id"]
                                                                              for x in e.get("impact") or []):
                    continue
                if reaches(a["id"], e["atom"]):
                    a["board_rechecked"] = qid
                    e["impact"] = (e.get("impact") or []) + [{"atom": a["id"], "result": "passed"}]
                    self._receipt(a, {"text": f"built after {qid} and its checks passed", "type": "recheck",
                                      "target": e["atom"]}, a.get("attempts") or 1,
                                  {"actual": True, "pass": True,
                                   "evidence": f"{a['id']} depends on {e['atom']}, which {qid} changed"}, trigger=qid)
                    j2_board.write(self.dir, self.board)
                    self.save("workflows.json", self.workflows)

    def _credit_board_used(self, meta, reader=None):
        """Item 4 (Reuse): used_by is only recorded when the reply names the board id it used. The reader is the
        atom that is building (a builder) or the function's name. Returns the ids actually credited, so the caller
        can tag its own chat message ("read Q-n")."""
        read, changed, who = [], False, reader or self.atom
        for qid in meta.get("board_used") or []:
            e = next((e for e in self.board["entries"] if e["id"] == qid and e["status"] == "answered"), None)
            if e and who and who not in e["used_by"]:
                e["used_by"].append(who)
                changed = True
            if e:
                read.append(qid)
        if changed:
            j2_board.write(self.dir, self.board)
        return read

    # Topics the mock reuse rule matches on (the keywords of this demo's brief and change request; timings and hours
    # are the same topic). An entry is reused only when a topic is in both its question and the reader's own work.
    BOARD_TOPICS = (("timings", "timing", "hours", "hour"), ("price", "prices", "priced", "pricing"),
                    ("product", "products"), ("order", "orders"), ("filter", "filters", "filtering"))

    def _board_reader(self, reader):
        """(what the reader is called, the items it is working on, its own asked_by id). A builder: its name and the
        done-when items it covers. A function: its name, the Requester's change request and the workflows of the
        cycle it is working on — plus, for Audit, every done-when item, since it checks them all."""
        crit = (self.goal or {}).get("done_when") or []
        if isinstance(reader, dict):
            covered = [c for c in crit if self._no(c) in (self._wf_items(reader) or [])]
            return str(reader.get("name") or ""), covered + [str(reader.get("goal_gap") or "")], reader["id"]
        working = [str(self.feedback or "")] + [str(w.get("name") or "") for w in self.workflows
                                                if w.get("cycle") == self.cycle]
        return NAMES.get(reader, reader), (crit if reader == "audit" else []) + working, reader

    def _mock_board_reuse(self, reader):
        """Mock mode's Reuse (item 4): the reader names the answered entries that share a topic (BOARD_TOPICS) with
        its own name or covered items, never every entry on the board. A builder doesn't reuse what its own atom
        asked (it got that as an Answer); a function reads its own earlier entry back off the board on a later
        turn, which is the point of the board."""
        if not j2_board.enabled():
            return {}
        name, covered, mine = self._board_reader(reader)
        mytext = " ".join([name] + [str(c) for c in covered]).lower()
        used = [e["id"] for e in self.board["entries"] if e["status"] == "answered"
                and not (isinstance(reader, dict) and e["asked_by"] == mine)
                and any(any(k in mytext for k in topic) and any(k in str(e["question"]).lower() for k in topic)
                        for topic in self.BOARD_TOPICS)]
        return {"board_used": used} if used else {}

    def _mock_board_ask(self, fn, to, about, text):
        """Mock mode's Ask (item 2) for a function, via J2_FORCE_QUESTION (a comma list: identity, priority, audit,
        coordination, adaptation, or a builder's id or role word). The named function asks this one question once
        per run. The switch is for humans: it is posted as a test note and never reaches an agent's prompt."""
        if not (j2_board.enabled() and llm.MOCK) or fn not in self.force_qs or fn in self.forced_asks:
            return {}
        if not to or not self._board_target(to)[1]:
            return {}
        self.forced_asks.add(fn)
        self.post("system", f"TEST SWITCH J2_FORCE_QUESTION={self.force_question} (humans only, agents don't see "
                            f"this): {NAMES[fn]} asks {to} about {about}.", warning=True, test=True)
        return {"board_question": {"to": to, "about": about, "text": text}}

    def _board_owner(self, *words):
        """The done workflow that owns this subject: the one whose name says so, else the one whose gap or covered
        items do, latest first. Who to ask, read off the run's own record; None when nothing built it."""
        crit = (self.goal or {}).get("done_when") or []
        done = [w for w in reversed(self._all_wfs()) if w.get("status") == "done"]
        for name_only in (True, False):
            for w in done:
                text = " ".join([str(w.get("name") or "")] if name_only else
                                [str(w.get("goal_gap") or "")]
                                + [c for c in crit if self._no(c) in (self._wf_items(w) or [])]).lower()
                if any(k in text for k in words):
                    return w
        return None

    def _board_owner_atom(self, *words):
        """The atom of that workflow (what a question names), or None."""
        w = self._board_owner(*words) or {}
        return next((a["id"] for a in w.get("atoms") or [] if a.get("agent_output")), None)

    # The document question each role word scripts (mock mode): the builder of this section asks the atom of the
    # earlier document its own prompt was built from (_chain_edges), about the subject that section must settle.
    DOC_ASKS = {"lld": ("Interactive elements", "hld", "Components", "how the booking is sent",
                        "How is the booking sent: mailto or a server endpoint?"),
                "hld": ("Components", "prd", "Requirements", "which requirement the booking form covers",
                        "Which requirement covers the booking form?")}

    def _mock_board_extra(self, w, attempt):
        """Mock mode only, for a builder. Ask (item 2, via J2_FORCE_QUESTION): a question only on the matched
        workflow's first attempt, at most once per run. Reuse (item 4): the answered entries that share a topic with
        this workflow, so later builders reuse them without a real model."""
        ask = self.DOC_ASKS.get(self.output) if self.doc else None
        if ask and attempt == 1 and not self.question_forced and self.output in self.force_qs \
                and ask[0].lower() in str(w["name"]).lower():
            to = next(reversed(self._doc_atoms(ask[1], ask[2])), None)
            if to:
                self.question_forced = True
                self.post("system", f"TEST SWITCH J2_FORCE_QUESTION={self.force_question} (humans only, agents don't "
                                    f"see this): {w['id']} asks {to} about {ask[3]}.", warning=True, test=True)
                return {"question": {"to": to, "about": ask[3], "text": ask[4]}}
        words = [x for x in self.force_qs if x not in FUNCTIONS]
        if words and attempt == 1 and not self.question_forced \
                and any(x == w["id"].lower() or x in str(w["name"]).lower() for x in words):
            manifest = self._manifest(w)
            if manifest:
                self.question_forced = True
                self.post("system", f"TEST SWITCH J2_FORCE_QUESTION={self.force_question} (humans only, agents don't "
                                    f"see this): {w['id']} asks {manifest[-1]['atom']} about its data attributes.",
                          warning=True, test=True)
                return {"question": {"to": manifest[-1]["atom"], "about": "data attributes",
                                      "text": "Which attributes hold this item's id and price?"}}
        return self._mock_board_reuse(w)

    def _mock_board_answer(self, wid, entry=None):
        """Mock mode's Answer (item 3) from a workflow: read out of its own recorded section. A question about an
        attribute or a price names the section's data-* attributes (and offers the correction that records them); a
        question about where something is names the section and its anchor. "not recorded" when the section holds
        neither, never invented."""
        w = next((x for x in self._all_wfs() if x["id"] == wid), None)
        if self._doc_of_workflow(wid):
            return self._mock_board_answer_doc(wid, w, entry)
        sid = (w or {}).get("section")
        body = (self.sections.get(sid) or {}).get("body") or ""
        asked = " ".join([str((entry or {}).get("about") or ""), str((entry or {}).get("question") or "")]).lower()
        attrs = sorted(set(re.findall(r"\bdata-[\w-]+(?==)", body)))
        if not re.search(r"where|which section|on the page", asked) and attrs:
            note = f"<!-- board: {', '.join(attrs)} confirmed -->"
            return {"answer": f"This section's items use {', '.join(attrs)}.",
                    "artifact_update": None if note in body else {"body": body.rstrip() + "\n" + note}}
        # "where": the section of its own body whose heading answers the question, by anchor and heading.
        secs = re.findall(r'<section[^>]*\bid="([\w-]+)"[^>]*>\s*(?:<h[23][^>]*>([^<]*)</h[23]>)?', body)
        want = j2_board.content_words(asked)
        pick = next((s for s in secs if any(k in (s[1] or "").lower() for k in want)), secs[0] if secs else None)
        if sid and pick:
            return {"answer": f"Yes: it is in section {sid}, anchor #{pick[0]}"
                              + (f', headed "{pick[1]}"' if pick[1] else "") + ".", "artifact_update": None}
        if attrs:
            return {"answer": f"This section's items use {', '.join(attrs)}.", "artifact_update": None}
        return {"answer": "not recorded", "artifact_update": None}

    def _mock_board_answer_doc(self, wid, w, entry):
        """Mock mode's Answer from a document section (board/1 part C): read out of the section the run recorded. A
        question about how something is sent is answered from the address the brief gives, with the correction to its
        own section; a question about which requirement covers something names the requirement. The words are the
        brief's and the document's; nothing is invented, and "not recorded" when the run holds neither."""
        kind = self._doc_of_workflow(wid)
        section = next((s for s in OUTPUTS[kind]["sections"] if squash(w.get("name") or "").endswith(squash(s))), None)
        aid = ((w or {}).get("atoms") or [{}])[0].get("id")
        body = (atom_section_body(self.dir, self._all_wfs(), aid) or "") if aid else ""
        body = re.sub(r"^\s*#{1,6}\s+[^\n]*\n?", "", body.strip("\n"), count=1).strip("\n")
        asked = " ".join([str((entry or {}).get("about") or ""), str((entry or {}).get("question") or "")]).lower()
        said = " ".join([self.request] + [m["text"] for m in self.chat if m["who"] == "user"])
        if not section or not body.strip():
            return {"answer": "not recorded", "artifact_update": None}
        if re.search(r"\bsent\b|\bsend\b|mailto|endpoint|server", asked):
            mail = next(iter(sorted(x.rstrip(".") for x in set(EMAIL.findall(said)))), None)
            if not mail:
                return {"answer": "not recorded", "artifact_update": None}
            add = f"The booking is sent by mailto to {mail}; there is no server endpoint."
            lines = body.splitlines()
            # into the responsibility of the component that holds the form (before its "Covers R..", so the ids and
            # the chain checks still read as they did)
            i = next((k for k, ln in enumerate(lines) if re.search(r"\bform\b|\bbook", ln, re.I)), len(lines) - 1)
            m = re.search(r"\s*\bCovers\b", lines[i])
            lines[i] = lines[i][:m.start()] + " " + add + lines[i][m.start():] if m else lines[i].rstrip() + " " + add
            return {"answer": f"mailto to {mail}, no server.",
                    "artifact_update": {"doc": kind, "section": section, "text": "\n".join(lines),
                                        "reason": f"the form sends the booking by mailto to {mail}, with no server"}
                    if add not in body else None}
        if re.search(r"requirement|covers", asked):
            m = next((ln for ln in body.splitlines() if re.search(r"\bform\b|\bbook", ln, re.I)), None)
            rid = re.search(r"\bR(\d+)\b", m) if m else None
            if not rid:
                return {"answer": "not recorded", "artifact_update": None}
            mail = next(iter(sorted(x.rstrip(".") for x in set(EMAIL.findall(said)))), None)
            lines = body.splitlines()                 # the same requirement, saying where the booking goes
            lines[lines.index(m)] = m.rstrip() + (f" The booking goes to {mail}." if mail else "")
            return {"answer": f"R{rid[1]} covers it: {squash(strip_md(m))[:110]}",
                    "artifact_update": {"doc": kind, "section": section, "text": "\n".join(lines),
                                        "reason": f"R{rid[1]} should say where the booking goes"} if mail else None}
        return {"answer": f"That is in the {kind.upper()}'s {section} section, as recorded.", "artifact_update": None}

    def _mock_board_answer_fn(self, fn, entry):
        """Mock mode's Answer (item 3) from a function: read out of the run's own record. Adaptation answers from the
        proposal it has on the table (self.a_pos), so it can say whether the work is new or a revise priced as a
        fix while Priority is still pricing it. Anything the run hasn't recorded is "not recorded"."""
        text = str((entry or {}).get("question") or "").lower()
        if fn == "adaptation" and re.search(r"new work|fix|revis", text):
            table = self.a_pos or []
            named = (re.search(r"\bW-\d+\b", text, re.I) or [""])[0].upper()
            w = next((x for x in table if str(x.get("revises") or "").upper() == named and named), None) \
                or next((x for x in table if x.get("revises")), None)
            if w:
                return {"answer": f"A fix: {w.get('name')} revises {w['revises']}'s section, priced with a 'fix' "
                                  f"part ({', '.join(w.get('parts') or []) or 'not recorded'})."}
            if table:
                return {"answer": f"New work: {table[0].get('name')} revises nothing and covers "
                                  f"{table[0].get('covers') or 'not recorded'}."}
        return {"answer": "not recorded"}

    def _spec_conflict(self, w, meta):
        """A builder says its LLD slice can't be built as written (spec_conflict). The first time for a section, that
        LLD section is revised once, with a version note and the reason (_revise_lld); the second time it is
        escalated. None when there is no usable conflict."""
        c = meta.get("spec_conflict")
        sec = str(c.get("section")) if isinstance(c, dict) else ""
        if sec not in {x["id"] for x in (self.blueprint or {}).get("sections") or []}:
            return None
        self.lld_conflicts[sec] = self.lld_conflicts.get(sec, 0) + 1
        self._event(w["id"], "spec_conflict", f"{sec}: {c.get('reason')}")
        if self.lld_conflicts[sec] > 1:
            return "escalate"
        self._revise_lld(w, sec, c)
        return "revised"

    def _revise_lld(self, w, sec, c):
        """Revise one LLD section once: apply the builder's selector changes to the spec and to lld.md, add a version
        note, and record the reason in the LLD's agent_output (document and owner atom)."""
        use = {str(k): str(v) for k, v in (c.get("use") or {}).items() if str(k).startswith(("#", ".")) and str(v)[:1] in "#."}
        spec, v = self.blueprint, 2
        fld = {str(k): x for k, x in (c.get("fields") or {}).items()} if isinstance(c.get("fields"), dict) else {}
        fix = lambda t: reduce(lambda a, kv: str(a).replace(kv[0], kv[1]), use.items(), t) if t else t
        for x in spec["sections"]:
            if x["id"] == sec:
                x["version"] = v
                x["anchor"] = (fix("#" + str(x["anchor"])) or "#")[1:]
        for x in spec["elements"]:
            if x["section"] == sec:
                x["selector"], x["behaviour"] = fix(x["selector"]), fix(x["behaviour"])
        for x in spec["fields"]:
            if x["section"] == sec and x.get("form"):
                x["form"] = (fix("#" + x["form"]))[1:]
        spec["fields"] = [f for f in spec["fields"] if not (f["section"] == sec and f["id"] in fld and fld[f["id"]] is None)]
        for f in spec["fields"]:      # a field the builder retyped (a string) is retyped in the spec
            if f["section"] == sec and isinstance(fld.get(f["id"]), str) and fld[f["id"]]:
                f["type"] = fld[f["id"]]
        note = {"section": sec, "version": v, "reason": str(c.get("reason") or "").strip(), "changes": use, "by": w["id"],
                **({"fields": fld} if fld else {})}
        self.lld_revisions.append(note)
        lld = self.docs.get("lld") or {}
        path = os.path.join(self.dir, "lld.md")
        if os.path.exists(path):
            text = fix(open(path, encoding="utf-8").read())
            for fid, new in fld.items():      # a removed field's line goes; a retyped field's line says its new type
                text = re.sub(rf"(?m)^\s*[-*+]\s*\*\*{re.escape(fid)}\*\*.*\n?", "" if new is None else
                              lambda m: re.sub(r"type [\w-]+", f"type {new}", m.group(0)), text)
            text = re.sub(rf"(?m)^(\s*[-*+]\s*\*\*{sec}\*\*.*?)$", lambda m: m.group(1) + f" Version {v}: {note['reason']}.", text, count=1)
            open(path, "w", encoding="utf-8").write(text)
            export_lld(text, self.dir)
        if lld:
            lld["agent_output"] = {**lld["agent_output"], **{k: spec[k] for k in ("sections", "elements", "fields")},
                                   "revisions": self.lld_revisions}
        for o in self.past_workflows:        # the LLD atom that wrote the sections carries the reason in its agent_output
            if "Pages and sections" in o.get("name", ""):
                for a in o.get("atoms") or []:
                    if a.get("agent_output"):
                        a["agent_output"] = {**a["agent_output"], "revisions": self.lld_revisions}
        self.save("documents.json", self.docs)
        self.post("system", f"{w['id']} says LLD section {sec} can't be built as written ({note['reason']}). The section is "
                            f"revised once, to version {v}; the checks now use the revised section.", warning=True)

    def _lld_owner_note(self, c):
        o = next((x for x in reversed(self.past_workflows) if "Pages and sections" in x.get("name", "")), None)
        a = ((o or {}).get("atoms") or [{}])[0].get("id")
        return (f"Owner atom: {a or 'the LLD'} (the LLD section {c.get('section')} was already revised once; "
                f"reason: {c.get('reason')}).")

    def _blames_other(self, say, fn):
        """The builder's own report puts the cause in another workflow's section (names one besides its own)."""
        say = str(say or "")
        return bool(set(re.findall(r"\bW-\d+\b", say)) - {fn}) and bool(re.search(r"can(?:no|')t|unable|not (?:in|mine)|"
                                                                                    r"outside|belongs?|their|its own", say, re.I))

    def _owner_note(self, failed, sid):
        """Evidence for Adaptation: which other workflow's atom holds the selector a failed check targets."""
        for r in failed:
            name = str(r.get("target") or "").lstrip("#.")
            for k, x in self.sections.items():
                if k not in (sid, "style") and name and re.search(rf"""(?:id=["']{re.escape(name)}["']|class=["'][^"']*\b"""
                                                                  rf"""{re.escape(name)}\b)""", x.get("body") or ""):
                    o = self._wf(str(x["owner"]).split()[0])
                    a = ((o or {}).get("atoms") or [{}])[0].get("id")
                    return (f"Owner atom: {a or x['owner']} (section {k}, {x['owner']}) also contains {r.get('target')}; "
                            f"the cause is in its section, not the failing atom's.")
        return ""

    def _doc_prompt(self, w, brief, sid, mine, attempt, evidence, last_err, prev):
        """(reply format, task) for a document section's workflow agent: markdown for its own section only."""
        fn = w["id"]
        self._record_given(w, self._chain_edges() + (self._manifest_edges(w) if slim() else []) + self._board_edges(),
                           whole=not slim())
        return ('First line: a JSON object {"say": str (one short line reporting to @Coordination: what you changed; '
                'code runs your checks and reports them, so don\'t claim they pass), "human_output": str (1 to 3 plain sentences telling a person what you built; \"done.\" alone is rejected), "atom_human_output": str (1 or 2 plain sentences telling the Requester what this part created, in everyday words: no ids, no selectors, no \"units\"), "next": "Coordination", '
                + (self.BOARD_FMT if j2_board.enabled() else '')
                + self.REASON + '}. Then a ```markdown block: your section only, starting with its "## <Section>" '
                'heading. Markdown only, never HTML.',
                f"YOU ARE {NAMES[fn]}, the agent for one registered workflow.\nYOUR WORKFLOW: "
                f"{json.dumps(brief, ensure_ascii=False)}\nYOU OWN SECTION {sid} of the {self.out['label']}. Code "
                f"assembles the document from every section in this order: {', '.join(self.out['sections'])}. Write "
                "the whole section each time; don't repeat other sections. After you write it, code runs every check "
                f"in your work atoms' verification on the whole document ({MAX_CHECK_ATTEMPTS} attempts).\n"
                f"FACTS YOU MAY USE (the charter; copy names and numbers exactly, invent none): "
                f"{json.dumps(self.goal, ensure_ascii=False)}\nDOCUMENT RUBRIC (Audit checks it; follow it):\n"
                + "\n".join(f"{rid} {txt}" for rid, txt in self.out["rubric"])
                + (f"\nYOUR SECTION NOW (last written by {mine['owner']}):\n```markdown\n{mine['body']}\n```" if mine else "")
                + (f"\nATTEMPT {attempt} OF {MAX_CHECK_ATTEMPTS}. CODE RAN YOUR CHECKS ON YOUR LAST ATTEMPT AND THESE "
                   f"FAILED:\n{evidence}\nFix the section so they pass." if evidence else "")
                + (f"\nYOUR LAST ATTEMPT WAS REJECTED: {last_err}. Fix that." if last_err else "")
                + (f"\nOTHER SECTIONS (their atoms' agent_output; the rest isn't shown): "
                   f"{json.dumps(self._manifest(w), ensure_ascii=False)}" if slim() else
                   f"\nTHE WHOLE DOCUMENT NOW (v{len(self.sites)}, read only):\n{prev[:SITE_LIMIT]}")
                + self._board_context())

    def _receipt(self, a, c, attempt, res, **extra):
        """One line of receipts.jsonl: written by code only, never by an agent."""
        rec = {"atom": a["id"], "check": c.get("text"), "type": c.get("type"), "target": c.get("target"),
               "expect": c.get("expect"), "actual": res["actual"], "pass": res["pass"], "evidence": res["evidence"],
               "attempt": attempt, "of": MAX_CHECK_ATTEMPTS, "t": time.strftime("%Y-%m-%dT%H:%M:%S"), "by": "code",
               **extra}
        self.receipts.append(rec)
        os.makedirs(self.dir, exist_ok=True)
        with open(os.path.join(self.dir, "receipts.jsonl"), "a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
        return rec

    def _expected_receipts(self, a, ver, recs, attempt, **extra):
        """Per expected-output item: met (all its checks passed), not met (one failed) or None (could not run), stored
        on the item and written to receipts.jsonl by code. These lines carry "expected", not "check"."""
        for n, x in enumerate([x for x in a.get("expected_output") or [] if isinstance(x, dict)], 1):
            res = [r["pass"] for c, r in zip(ver, recs) if c.get("item") == n]
            x["met"] = None if not res or (None in res and False not in res) else False not in res
            rec = {"atom": a["id"], "expected": x.get("text"), "kind": x.get("kind"), "target": x.get("target"),
                   "met": x["met"], "checks": [c.get("text") for c in ver if c.get("item") == n], "attempt": attempt,
                   "of": MAX_CHECK_ATTEMPTS, "t": time.strftime("%Y-%m-%dT%H:%M:%S"), "by": "code", **extra}
            with open(os.path.join(self.dir, "receipts.jsonl"), "a", encoding="utf-8") as f:
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    def _run_checks(self, w, attempt, page, atoms=None, doc=None, **extra):
        """Code runs every check of the workflow's atoms on the built site (doc=True: page is a markdown document, as
        for an artifact workflow); each result is a receipt and the atom keeps its latest ones. Returns the failed
        receipts."""
        site_dir, failed, doc = os.path.join(self.dir, "site"), [], self.doc if doc is None else doc

        def one(c):
            if doc and c.get("type") == "file-exists" and self._file_wf(w):      # a file workflow: the file it wrote
                f = artifact_path(w)
                ok = os.path.isfile(os.path.join(self.dir, f))
                return {"actual": ok, "pass": ok, "evidence": f"{f} {'exists' if ok else 'not found'} in the run"}
            return run_check(c, site_dir, page, doc)
        for a in atoms or w.get("atoms") or []:
            ver = [c for c in a.get("verification") or [] if isinstance(c, dict)]
            recs = [self._receipt(a, c, attempt, one(c), **extra) for c in ver]
            a["results"], a["attempts"] = recs, attempt
            self._expected_receipts(a, ver, recs, attempt, **extra)
            # pass None = could not run (no browser): never a failed attempt, never escalated
            a["status"] = "failed" if any(r["pass"] is False for r in recs) else \
                "could not run" if any(r["pass"] is None for r in recs) else "passed"
            failed += [r for r in recs if r["pass"] is False]
        if w.get("atoms"):
            n = sum(len(a.get("results") or []) for a in atoms or w["atoms"])
            self.post("system", f"Code ran {n} check(s) for {w['id']} (attempt {attempt} of {MAX_CHECK_ATTEMPTS}): "
                                + (f"{len(failed)} failed: " + "; ".join(f"{r['check']} ({r['evidence']})"
                                                                        for r in failed) if failed else
                                   "all the others passed; could not run: " + "; ".join(
                                       r["check"] for a in atoms or w["atoms"] for r in a["results"] if r["pass"] is None)
                                   if any(a.get("status") == "could not run" for a in atoms or w["atoms"]) else
                                   "all passed")
                                + ".", checks=w["id"])
        self.save("workflows.json", self.workflows)
        return failed

    def _adaptation_atom_failed(self, h):
        """E: a work atom failed its checks on every attempt. Adaptation reads the receipts (all attempts, the check
        text, Audit's latest view) and decides: the check is wrong (it rewrites it and code re-runs it on the same
        build), or the build is wrong. Only a wrong build the budget left can't fix goes to Identity. An atom with no
        fix after its attempts is ESCALATED."""
        self.step = 2
        w = self._wf(h.get("wid")) or next((x for x in self._all_wfs() if x["id"] == h.get("wid")), None)
        atoms = [a for a in (w or {}).get("atoms") or [] if a["id"] in (h.get("atoms") or [])]
        if not atoms:
            return self._handoff("coordination", kind="schedule")
        recs = [r for r in self.receipts if r["atom"] in {a["id"] for a in atoms}]
        last = self.audits[-1] if self.audits else None
        view = ([{k: r.get(k) for k in ("id", "criterion", "verdict", "note")} for r in last["results"]
                 if r.get("item") in self._wf_items(w) or not r.get("pass")] if last else "no Audit yet")
        cheapest = RATES.get("fix", min(RATES.values()) if RATES else 1)
        a0 = atoms[0]
        # A document: code checks the obvious case first. A failed heading or text check that passes once its target
        # is read as written in the document means the check's target is wrong, not the build.
        corr, why = {}, []
        for a in atoms if self.doc and self.sites else []:
            page = self._recheck_page(a["id"])[0]      # its own document: a filed one when the atom wrote that
            fixed = [(c, corrected_check(c, page) if not doc_check(c, page)["pass"] else None)
                     for c in a.get("verification") or [] if isinstance(c, dict)]
            if any(x for _, x in fixed):
                corr[a["id"]] = [x[0] if x else c for c, x in fixed]
                why += [f"{a['id']}: {x[1]}" for _, x in fixed if x]

        def mock():
            bad = [r for r in recs if r["attempt"] == h.get("attempt", MAX_CHECK_ATTEMPTS) and not r["pass"]]
            if bad and all(r["type"] in ("grep-count", "element-count") for r in bad):
                fixed = {(r["type"], str(r["target"])): r["actual"] for r in bad}
                ver = [{**c, "expect": fixed[k], "text": re.sub(r"\d+", str(fixed[k]), c.get("text") or "")}
                       if (k := (c.get("type"), str(c.get("target")))) in fixed else c for c in a0["verification"]]
                return json.dumps({"say": f"@Coordination The check is wrong, not the build: the page has "
                                          f"{bad[0]['actual']} and that is what the goal needs. Check corrected.",
                                   "verdict": "CHECK_WRONG", "atom": a0["id"], "verification": ver})
            return json.dumps({"say": "@Coordination The build is wrong: the section misses what the check asks for.",
                               "verdict": "BUILD_WRONG", "atom": a0["id"],
                               "affordable": self.budget - self.spent >= cheapest})
        r = self.call("adaptation", '{"say": str, "verdict": "CHECK_WRONG"|"BUILD_WRONG", "atom": str (the atom id), '
                      '"verification": [check] (CHECK_WRONG only: the corrected checks, same rules; keep each check\'s '
                      '"item" number), "affordable": bool '
                      '(BUILD_WRONG only: can the budget left pay for a fix workflow?), "build_wrong_reason": str|null '
                      '(BUILD_WRONG only: exactly what the build lacks)}',
                      f"WORK ATOM FAILED: {', '.join(a['id'] for a in atoms)} failed its checks on "
                      + (f"attempt {h['attempt']}; code stopped further attempts. {h.get('owner_note') or ''}"
                         if h.get("attempt", MAX_CHECK_ATTEMPTS) < MAX_CHECK_ATTEMPTS else
                         f"all {MAX_CHECK_ATTEMPTS} attempts.") +
                      f"\nTHE ATOMS: {json.dumps(atoms, ensure_ascii=False)}\nRECEIPTS "
                      f"(written by code, every attempt): {json.dumps(recs, ensure_ascii=False)}\nAUDIT'S LATEST VIEW: "
                      f"{json.dumps(view, ensure_ascii=False)}\nDecide whether the check or the build is wrong. "
                      "CHECK_WRONG: the check doesn't test the goal correctly; send the corrected verification and code "
                      "re-runs it on the same build. BUILD_WRONG: the section doesn't meet the goal; say whether the "
                      f"budget left ({self.budget - self.spent}) can pay for a fix."
                      + (f"\nCODE CHECKED THE OBVIOUS CASE FIRST: the check's target looks wrong: {'; '.join(why)}. "
                         "The decision is CHECK_WRONG with the corrected checks unless you give a specific reason the "
                         "build is wrong in build_wrong_reason." if why else "") + self.ATOM_RULES, mock)
        verdict = str(r.get("verdict") or "").upper()
        a = next((x for x in atoms if x["id"] == r.get("atom")), a0)
        self._say("adaptation", r.get("say", ""), "coordination", None, act=verdict or "BUILD_WRONG", atom=a["id"])
        if corr and not (verdict == "BUILD_WRONG" and str(r.get("build_wrong_reason") or "").strip()):
            a = a if a["id"] in corr else next(x for x in atoms if x["id"] in corr)
            if verdict != "CHECK_WRONG":
                r = {**r, "verification": corr[a["id"]]}
            verdict = "CHECK_WRONG"
            self._event("validator", "check_wrong_default", "; ".join(why))
            self.post("system", f"Code checked the obvious case first: the check's target looks wrong ({'; '.join(why)}). "
                                "Adaptation named no specific way the build is wrong, so the decision is CHECK_WRONG and "
                                f"code re-runs the corrected check on the same {self.out['noun']}.")
        if verdict == "CHECK_WRONG":
            new = normalise_checks([{**copy.deepcopy(a), "verification": r.get("verification")}])[0]
            probs = atom_missing({"atoms": [new]}) or check_problems(new, self._fact_source(), self.out["check_types"])
            if probs and a["id"] in corr:        # Adaptation's own correction breaks a rule: code's correction
                new, probs = {**a, "verification": corr[a["id"]]}, []
            if probs:
                self._reject(f"{a['id']} corrected check", "; ".join(probs))
            else:
                a["verification"] = new["verification"]
                if not self._run_checks(w, MAX_CHECK_ATTEMPTS, *self._recheck_page(a["id"])[:1], atoms=[a],
                                        doc=self._recheck_page(a["id"])[1],
                                        recheck="check corrected by Adaptation"):
                    self._seal_atoms(w, [a], MAX_CHECK_ATTEMPTS)
                    if a.get("status") == "passed":
                        return self._handoff("coordination", kind="schedule")
        a["status"] = "escalated"
        self._event("adaptation", "atom_escalated", f"{a['id']}: {verdict or 'no verdict'}")
        self.save("workflows.json", self.workflows)
        if verdict == "BUILD_WRONG" and r.get("affordable") is False:
            self.post("system", f"{a['id']} is ESCALATED: the build is wrong and Adaptation says the budget left can't "
                                "pay for a fix. Identity decides.")
            return self._handoff("identity", kind="question", frm="adaptation",
                                 question=f"{a['id']} ({w['name']}) fails its checks and a fix doesn't fit the budget "
                                          f"left: {r.get('say')}",
                                 resume={"to": "coordination", "handoff": {"kind": "schedule"}})
        self.post("system", f"{a['id']} is ESCALATED after {MAX_CHECK_ATTEMPTS} failed attempts with no fix. Audit "
                            f"checks the {self.out['noun']}; Adaptation proposes the fix from its findings.")
        self._handoff("coordination", kind="schedule")

    # ------------------------------------------------------------------ J2.g: Audit
    def _audit_findings(self, findings, crit):
        """One result per done-when item, matched by item number. A missing item, or a PASS with no evidence, becomes
        BLOCKED_BY_EVIDENCE (audit.md: when evidence is absent, don't guess). Other findings are extra issues."""
        by_item, extras = {}, []
        g = self.goal or {}
        orig, dropped = g.get("charter_items") or crit, g.get("dropped_by_requester") or []
        for i, f in enumerate(x for x in findings if isinstance(x, dict)):
            row = {"id": f.get("id") or f"F-{i + 1}", "criterion": f.get("criterion", ""),
                   "verdict": str(f.get("verdict", "")).upper(), "note": str(f.get("evidence") or "").strip(),
                   "severity": f.get("severity"), "confidence": f.get("confidence"),
                   "needs_requester_fact": str(f.get("needs_requester_fact") or "").strip() or None}
            n = self._finding_item(as_int(f.get("item"), 0), row["criterion"], crit, orig, dropped)
            if n == -1:     # a dropped item is out of scope: never a finding, never an extra issue
                self._event("audit", "invalid_move", f"{row['id']} is about a dropped item: ignored")
                continue
            if 1 <= n <= len(crit) and n not in by_item:
                by_item[n] = row
            else:
                extras.append(row)
        results = []
        for n, c in enumerate(crit, 1):
            row = by_item.get(n)
            if not row:
                self._event("audit", "invalid_move", f"no finding for done-when item {n}")
                row = {"id": f"F-missing-{n}", "verdict": "BLOCKED_BY_EVIDENCE", "note": "Audit filed no finding for "
                       "this item.", "severity": "medium", "confidence": "high"}
            elif row["verdict"] in ("PASS", "PASS_WITH_WARNING") and not row["note"]:
                self._event("audit", "invalid_move", f"{row['id']} passed item {n} without evidence")
                row = {**row, "verdict": "BLOCKED_BY_EVIDENCE", "note": "Marked PASS with no evidence; treated as "
                       "BLOCKED_BY_EVIDENCE."}
            tbc = next((t for t in (self.goal or {}).get("to_be_confirmed") or [] if t["criterion"] == c), None)
            if tbc:   # J2.g: a fact the Requester hasn't confirmed yet: this item is not met, whatever the page shows
                row = {**row, "verdict": "TO_BE_CONFIRMED", "note": f"Waiting on the Requester: {tbc['fact']}. "
                                                                    f"(Audit's own note: {row.get('note') or 'none'})"}
            # A finding is named after its item's run-long number (F-n is item n, also after drops); Audit's own id is
            # kept as audit_id.
            fid = f"F-{self._no(c)}"
            row = {**row, "id": fid, **({"audit_id": row["id"]} if row["id"] not in (fid, f"F-missing-{n}") else {})}
            results.append({**row, "item": n, "criterion": c, "pass": row["verdict"] in ("PASS", "PASS_WITH_WARNING")})
        used = {r["id"] for r in results}
        for x in extras:                 # an extra issue never shares an id with a checklist item
            if x["id"] in used:
                x["audit_id"], x["id"] = x["id"], next(f"F-{k}" for k in range(len(g.get("charter_items") or crit) + 1,
                                                                             10 ** 4) if f"F-{k}" not in used)
            used.add(x["id"])
        return results, extras

    @staticmethod
    def _finding_item(n, text, crit, orig, dropped):
        """The current item number a finding is about, -1 for a dropped item, 0 for an extra issue. After a drop
        Audit may still use the charter's first numbering (runs/j2-20261009-133203: items 9, 15, 16 for 5, 6, 7), so
        the number is read against both lists and the finding's own text picks between them."""
        def like(item, finding):        # share of the finding's words that the item's text contains
            words = set(re.findall(r"[a-z0-9]{4,}", str(finding).lower()))
            return len(words & set(re.findall(r"[a-z0-9]{4,}", str(item).lower()))) / len(words) if words else 0
        cands = []                       # the charter number first: that's the numbering Audit is given
        if 1 <= n <= len(orig):
            cands.append(orig[n - 1])
        if orig != crit and 1 <= n <= len(crit):
            cands.append(crit[n - 1])
        if not cands:                    # no usable number: the closest current or dropped item by text
            best = max(crit + dropped, key=lambda c: like(c, text), default=None)
            cands = [best] if best and like(best, text) >= 0.6 else []
        if not cands:
            return 0
        pick = max(cands, key=lambda c: like(c, text)) if text else cands[0]
        return crit.index(pick) + 1 if pick in crit else -1 if pick in dropped else 0

    def _audit_tally(self, results, extras):
        """A code-written tally under Audit's message, so the chat and the report always agree."""
        by = {}
        for x in results:
            by.setdefault(x["verdict"], []).append(x["id"])       # F-n names item n (run-long numbers)
        parts = [f"{len(v)} {k}" + ("" if k == "PASS" else f" ({', '.join(v)})") for k, v in by.items()]
        return (f"Findings (code tally): {'; '.join(parts)}"
                + (f"; extra issues: {', '.join(x['id'] + ' ' + x['verdict'] + ' (' + short(x.get('criterion') or x.get('note') or 'no description', 10) + ')' for x in extras)}"
                   if extras else "") + ".")

    def _turn_audit(self, h):
        if h.get("kind") == "dispute":
            return self._turn_dispute("audit", h)
        self.step = 7
        site = self.sites[-1] if self.sites else ""
        crit = (self.goal or {}).get("done_when", [])
        extra = {}
        if self.doc:      # a document: code checks the necessary conditions; no browser, screenshot or design rubric
            r, extra = self._audit_doc(site, crit)
            mech, shot, measured = {}, None, {}
        else:
            r, mech, shot, measured, conds = self._audit_site(site, crit)
            extra = {"conditions": conds}
        results, extras = self._audit_findings(r.get("findings") or r.get("results") or [], crit)
        forced = self._force_fail(results)
        self._flag_unpaid(results)
        rubric = self._audit_rubric(r.get("rubric") or [], measured, self.out["rubric"])
        audit = {"cycle": self.cycle, "version": len(self.sites), "results": results, "extra_issues": extras,
                 "rubric": rubric, "mechanical": mech, "screenshot": os.path.relpath(shot, self.dir) if shot else None,
                 "all_pass": bool(results) and all(x.get("pass") for x in results + rubric)
                 and not any(blocks(c) for c in extra.get("conditions") or []), **extra}
        if self.blueprint and not self.doc:
            audit["lld"] = self._lld_compare(site)
        self.audits.append(audit)
        self._save_audits()
        nxt, note = self._pick("audit", r.get("next"), ("identity",), "identity",
                               "Audit files its findings to Identity, who decides J2.h")
        fails = workable(audit)
        if forced:   # Audit's own line may say "all pass": code writes it from the findings as filed
            r["say"] = f"@Identity Site v{len(self.sites)} checked. " + " ".join(
                f"{x['id']} item {x['item']} {x['verdict']}: {x['note']}" for x in results + rubric if not x["pass"])
        self._say("audit", r.get("say", "") + "\n" + self._audit_tally(results + rubric, extras)
                  + (" Workable (code): yes." if not fails else f" Workable (code): no, {len(fails)} problem(s)."),
                  nxt, note, audit=audit["all_pass"], version=len(self.sites), workable=not fails,
                  **({"forced_test": forced} if forced else {}))
        if self.memory and len(self.audits) >= 2:
            self._learn(self.audits[-2], audit)
        failed = [f"workflow {w['id']} {w['name']} failed (envelope limits)" for w in self.workflows
                  if w["status"] == "failed" and w["cycle"] == self.cycle]
        self._handoff("identity", kind="audit", why=f"Failed workflows this cycle: {failed}" if failed else "")

    def _lld_compare(self, page):
        """The finished page against the whole LLD, by code: which sections match, which are on the page but differ
        (a mismatch), and which no workflow built (dropped by Priority: a warning, not a failure)."""
        built = {i for w in self.workflows if w["status"] in ("registered", "held", "running", "done")
                 for i in w.get("lld_sections") or []}
        out = {"matches": [], "mismatches": [], "dropped": []}
        for x in self.blueprint["sections"]:
            ok, ev = lld_check(page, lld_slice(self.blueprint, [x["id"]]))
            (out["matches"] if ok else out["mismatches"] if x["id"] in built else out["dropped"]).append(
                x["id"] if ok else {"section": x["id"], "evidence": ev})
        return out

    def _lld_audit_text(self, page):
        if not self.blueprint:
            return ""
        c = self._lld_compare(page)
        return ("LLD COMPARISON (checked by code: the page against the LLD; report what matches and what does not; "
                "sections no workflow built were dropped by Priority and are a warning only): "
                f"{json.dumps(c, ensure_ascii=False)}\n")

    def _chain_record(self):
        """requirement -> HLD component -> LLD section -> built atom, from the documents' agent_output and the site's
        workflows (each cites the LLD sections it implements). Code only."""
        reqs = (self.chain.get("prd") or {}).get("requirements") or []
        comps = (self.chain.get("hld") or {}).get("components") or []
        secs = (self.chain.get("lld") or {}).get("sections") or []
        atoms_of = lambda sid: [a["id"] for w in self._all_wfs() if sid in (w.get("lld_sections") or [])
                                for a in w.get("atoms") or []]
        rows = []
        for r in reqs or [{"id": None}]:
            cs = [c for c in comps if r["id"] in c.get("covers", [])] or [{"id": None}]
            for c in cs:
                ss = [x for x in secs if c["id"] in x.get("implements", [])] or [{"id": None}]
                for x in ss:
                    rows.append({"requirement": r["id"], "component": c["id"], "lld_section": x["id"],
                                 "atoms": atoms_of(x["id"]) if x["id"] else []})
        return rows

    def _audit_site(self, site, crit):
        """J2.g for a website: the render check and screenshot in headless Chrome, then Audit's call.
        Returns (reply, render check, screenshot path, measured design problems, necessary conditions)."""
        mech = j2_check.check(site) if site else {"available": False, "note": "no site built yet"}
        if site:
            self.post("system", j2_check.summary(mech) + f" (site v{len(self.sites)})", mechanical=True)
            for p in j2_check.problems(mech):
                self.events.append({"agent": "system", "type": "mechanical", "detail": p, "msg": len(self.chat),
                                    "cycle": self.cycle})
        if len(site) > SITE_LIMIT:
            self.post("system", f"Warning: site v{len(self.sites)} is {len(site):,} characters. Audit sees only the "
                                f"first {SITE_LIMIT:,}, so findings about the end of the file may be wrong.", warning=True)
        shot = None   # live runs: Audit looks at both widths; mock replies don't read images, so no Chrome run for them
        if site and not llm.MOCK:
            os.makedirs(os.path.join(self.dir, "site", "shots"), exist_ok=True)
            shot = j2_check.screenshot(site, os.path.join(self.dir, "site", "shots", f"v{len(self.sites)}.png"))
        measured = j2_check.design_problems(mech)
        conds = self._conditions(site, mech)

        def mock():
            fs = []
            for c in crit:
                phone, i = "phone" in c.lower(), self._no(c)
                ok = 'name="viewport"' in site if phone else f'data-item="{i}"' in site
                fs.append({"id": f"F-{i}", "item": i, "criterion": c, "verdict": "PASS" if ok else "NEEDS_REVISION",
                           "severity": "low" if ok else "high", "confidence": "high",
                           "evidence": ("Viewport meta tag present; layout is fluid" if ok else "No viewport meta tag, so "
                                        "phones render the page at desktop width") if phone else
                                       ("Section present in the page" if ok else "No section for this item")})
            n = sum(f["verdict"] == "NEEDS_REVISION" for f in fs)
            rub = [{"id": rid, "verdict": "PASS", "evidence": "Mock: looks fine in the measurements"} for rid, _ in RUBRIC]
            # Checking a price item: Audit asks the workflow that built the products which attribute holds the
            # price, rather than guessing from the markup (board, J2_BOARD=1).
            changed = any(m.get("change") for m in self.chat)     # the demo's order: after the change request
            ask = self._mock_board_ask("audit", self._board_owner_atom("product"), "the price attribute",
                                       "Which attribute holds the price?") if changed else {}
            return json.dumps({"say": "@Identity All criteria pass; done-when may hold." if not n else
                               f"@Identity {n} finding needs revision.", "findings": fs, "rubric": rub,
                               "next": "Identity", **ask})
        r = self.call("audit", '{"say": str, "findings": [{"id": "F-1", "item": int|null (done-when item number; null '
                      'for an extra issue), "criterion": str, "verdict": "PASS"|"PASS_WITH_WARNING"|'
                      '"NEEDS_REVISION"|"BLOCKED_BY_EVIDENCE"|"POLICY_CONCERN", "severity": "low"|"medium"|"high"|'
                      '"critical", "confidence": "low"|"medium"|"high", "evidence": str (one line: what you observed '
                      'in the code or the mechanical check), "needs_requester_fact": str|null (the fact only the '
                      'Requester knows that this item needs, e.g. "the business phone number"; null otherwise)}] '
                      '(exactly one finding per done-when item, plus one per extra issue; a PASS needs evidence), '
                      '"rubric": [{"id": "R-1".."R-7", "verdict": "PASS"|"PASS_WITH_WARNING"|"NEEDS_REVISION", '
                      '"evidence": str (one line: what you saw on the screenshot or in the measurements)}] (exactly '
                      'one per rubric item), '
                      + next_field(("identity",), "Identity: you file every finding to Identity, who decides J2.h") + "}",
                      f"CRITERIA (file each finding under its item number; numbers never change, also after items are "
                      f"dropped): {json.dumps([{'item': self._no(c), 'criterion': c} for c in crit], ensure_ascii=False)}\n"
                      + (f"DROPPED BY THE REQUESTER (out of scope: file no finding for these): "
                         f"{json.dumps((self.goal or {}).get('dropped_by_requester'), ensure_ascii=False)}\n"
                         if (self.goal or {}).get("dropped_by_requester") else "")
                      + f"MECHANICAL CHECK (run by code in headless Chrome; treat as evidence, and "
                      f"use BLOCKED_BY_EVIDENCE only for what neither it nor the code can show): {json.dumps(mech)}\n"
                      f"NECESSARY CONDITIONS (checked by code on the page; evidence): {json.dumps(conds, ensure_ascii=False)}\n"
                      f"DESIGN AND USABILITY RUBRIC (judge each item; a failing item is a normal finding): "
                      f"{json.dumps(dict(RUBRIC), ensure_ascii=False)}\nSTYLE GUIDE: {json.dumps(self.style, ensure_ascii=False)}"
                      f"\nDESIGN MEASUREMENTS (code, both widths): {json.dumps(mech.get('design'), ensure_ascii=False)}\n"
                      f"MEASURED PROBLEMS (code; you can't PASS these items): {json.dumps(measured, ensure_ascii=False)}\n"
                      + ("SCREENSHOT ATTACHED: the first screen at 360px (left) and 1280px (right). Judge R-1 to R-5 on "
                         "what you see.\n" if shot else "No screenshot this time: judge from the code and the measurements.\n")
                      + self._lld_audit_text(site) + f"SITE v{len(self.sites)}:\n{site[:SITE_LIMIT]}", mock,
                      max_tokens=1600, images=[shot] if shot else None)
        return r, mech, shot, measured, conds

    def _conditions(self, built, mech=None):
        """Code checks the output's necessary conditions on the latest build, saves conditions.json and posts the list.
        A website is checked on its HTML (visible text for the facts) and its render check."""
        page = None if self.doc else built
        text = built if self.doc else " ".join(PageScan(built).visible)
        conds = check_conditions(load_conditions(self.out["conditions"]), text, self.goal, self._fact_source(),
                                 self.out["sections"], page, mech)
        self.save("conditions.json", conds)
        if built:
            self.post("system", "Necessary conditions (checked by code): " + "; ".join(
                f"{c['id']} " + ("met" if c["met"] else f"{'WARNING' if c.get('severity') == 'warn' else 'NOT MET'} "
                                                       f"({c['evidence']})") for c in conds)
                + f" ({self.out['noun']} v{len(self.sites)})", mechanical=True)
        return conds

    def _audit_doc(self, doc, crit):
        """J2.g for a document: code checks the necessary conditions on the text (conditions.json), then Audit judges
        the done-when items and the document rubric from the text and says what it read. Returns (reply, the audit's
        extra fields)."""
        conds = self._conditions(doc)
        read = "; ".join(f"{h} ({md_words(b)} words)" for lv, h, b in md_sections(doc) if lv == 2) or "no sections"

        def mock():
            fs = []
            for c in crit:
                i, sec = self._no(c), next((x for x in self.out["sections"] if x.lower() in c.lower()), None)
                body = md_section(doc, sec) if sec else doc
                ok = bool(body and md_words(body))
                fs.append({"id": f"F-{i}", "item": i, "criterion": c, "verdict": "PASS" if ok else "NEEDS_REVISION",
                           "severity": "low" if ok else "high", "confidence": "high",
                           "evidence": f"{sec or 'The document'}: {md_words(body or '')} words" if ok else
                           f"No {sec} section with text"})
            n = sum(f["verdict"] != "PASS" for f in fs)
            return json.dumps({"say": "@Identity All criteria pass." if not n else f"@Identity {n} finding needs revision.",
                               "read": f"Read {self.out['file']} v{len(self.sites)}: {read}", "findings": fs,
                               "rubric": [{"id": rid, "verdict": "PASS", "evidence": "Mock: read from the text"}
                                          for rid, _ in self.out["rubric"]], "next": "Identity"})
        r = self.call("audit", '{"say": str, "read": str (what you read: each section and roughly how long it is), '
                      '"findings": [{"id": "F-1", "item": int|null (done-when item number; null for an extra issue), '
                      '"criterion": str, "verdict": "PASS"|"PASS_WITH_WARNING"|"NEEDS_REVISION"|"BLOCKED_BY_EVIDENCE"|'
                      '"POLICY_CONCERN", "severity": "low"|"medium"|"high"|"critical", "confidence": "low"|"medium"|"high", '
                      '"evidence": str (one line: what the text says), "needs_requester_fact": str|null}] (exactly one '
                      'finding per done-when item, plus one per extra issue; a PASS needs evidence), "rubric": [{"id": '
                      f'"{self.out["rubric"][0][0]}".."{self.out["rubric"][-1][0]}", "verdict": "PASS"|"PASS_WITH_WARNING"|'
                      '"NEEDS_REVISION", "evidence": str (one line: what in the text shows it; for D-5, quote each '
                      'claim not in the brief that isn\'t labelled "Assumption:")}] (exactly one per rubric item), ' + next_field(("identity",), "Identity: you file every finding to Identity") + "}",
                      f"CRITERIA (file each finding under its item number; numbers never change): "
                      f"{json.dumps([{'item': self._no(c), 'criterion': c} for c in crit], ensure_ascii=False)}\n"
                      + (f"DROPPED BY THE REQUESTER (out of scope: file no finding for these): "
                         f"{json.dumps((self.goal or {}).get('dropped_by_requester'), ensure_ascii=False)}\n"
                         if (self.goal or {}).get("dropped_by_requester") else "")
                      + f"NECESSARY CONDITIONS (checked by code on the text; evidence): {json.dumps(conds, ensure_ascii=False)}\n"
                      f"DOCUMENT RUBRIC (judge each item from the text; a failing item is a normal finding): "
                      f"{json.dumps(dict(self.out['rubric']), ensure_ascii=False)}\nFACTS FROM THE BRIEF AND THE REQUESTER "
                      f"(anything else stated as fact is invented): {json.dumps(self.request, ensure_ascii=False)} "
                      f"{json.dumps([m['text'] for m in self.chat if m['who'] == 'user'], ensure_ascii=False)}\n"
                      f"Say in read what you read. {self.out['label'].upper()} v{len(self.sites)} (markdown):\n"
                      f"{doc[:SITE_LIMIT]}", mock, max_tokens=1600)
        if not str(r.get("read") or "").strip():
            self._event("audit", "invalid_move", "didn't say what it read")
        return r, {"output": self.output, "conditions": conds, "read": str(r.get("read") or "").strip() or None}

    def paid_items(self):
        """Done-when items some built workflow paid for: its covers, plus the items of the findings it fixed."""
        crit = (self.goal or {}).get("done_when") or []
        # a finding's item by its text (an older audit row without one: its number)
        by_finding = {r["id"]: (crit.index(r["criterion"]) + 1 if r.get("criterion") in crit else
                                None if r.get("criterion") else r.get("item"))
                      for a in self.audits for r in a["results"]}
        return {i for w in self.workflows if w["status"] == "done"
                for i in self._wf_items(w) + [by_finding.get(f) for f in w.get("fixes") or []] if i}

    def _flag_unpaid(self, results):
        """Cost follows the work, checked on the page: a done-when item that passes although no workflow paid for it
        (an agent built it without declaring it) becomes a normal finding, UNPAID. It loops back: a workflow covering
        it at its rate-card minimum makes it paid."""
        if not RATES:
            return
        paid = self.paid_items()
        for row in results:
            if row["pass"] and row["item"] not in paid:
                kind = item_kind(row["criterion"])
                row.update(verdict="UNPAID", **{"pass": False}, severity="medium",
                           note=f"Built on the page, but no workflow covers item {row['item']}, so nobody paid for it. "
                                f"Register a workflow that covers it at its rate-card minimum ({kind}, "
                                f"{RATES.get(kind, 0)}). (Audit's own note: {row.get('note') or 'none'})")
                self._event("validator", "unpaid", f"item {row['item']} passes but no workflow covers it")

    def _audit_rubric(self, rows, measured, rubric=RUBRIC):
        """One result per rubric item. Missing, or a PASS without evidence, is BLOCKED_BY_EVIDENCE; an item code
        measured as failing (contrast, hierarchy, tap targets, alt text) is NEEDS_REVISION whatever Audit said."""
        by = {str(x.get("id")).upper(): x for x in rows if isinstance(x, dict)}
        out = []
        for rid, crit in rubric:
            x = by.get(rid) or {}
            verdict, note = str(x.get("verdict") or "").upper(), str(x.get("evidence") or "").strip()
            if not x:
                verdict, note = "BLOCKED_BY_EVIDENCE", "Audit filed no result for this rubric item."
            elif verdict in ("PASS", "PASS_WITH_WARNING") and not note:
                verdict, note = "BLOCKED_BY_EVIDENCE", "Marked PASS with no evidence."
            if measured.get(rid):
                if verdict in ("PASS", "PASS_WITH_WARNING"):
                    self._event("audit", "invalid_move", f"{rid} passed although code measured: {measured[rid][0]}")
                verdict, note = "NEEDS_REVISION", "Measured by code: " + "; ".join(measured[rid])
            out.append({"id": rid, "item": rid, "criterion": crit, "verdict": verdict, "note": note, "rubric": True,
                        "severity": "medium", "pass": verdict in ("PASS", "PASS_WITH_WARNING")})
        return out

    def _force_fail(self, results):
        """J2_FORCE_AUDIT_FAIL=<item>: a test switch for live runs. In cycle 1 only (Audit's first check), the item's
        result is overridden to NEEDS_REVISION with a plausible note (forced_note()), so agents treat it as a real
        finding; Audit's own verdict is kept in real_verdict, for humans only. Everything after runs normally, and later audits are real. Returns the item number when applied."""
        n = self.force_fail
        if not n or self.forced or self.cycle != 1:
            return None
        self.forced = True
        row = next((x for x in results if x["item"] == n), None)
        if not row:
            self.post("system", f"TEST SWITCH J2_FORCE_AUDIT_FAIL={n} ignored: the charter has no done-when item {n} "
                                f"(it has {len(results)}).", warning=True, test=True)
            return None
        real = f"{row['verdict']}: {row.get('note') or 'no note'}"
        # Agents see only a normal-looking finding. The test labels (forced_test, real_verdict, the TEST SWITCH note,
        # the UI tag) are for humans: state(), _findings() and chat_view() never pass them to an agent.
        row.update(verdict="NEEDS_REVISION", severity="medium", forced_test=True, real_verdict=real,
                   note=forced_note(row["criterion"]), **{"pass": False})
        self._event("system", "forced_test", f"item {n} forced to NEEDS_REVISION (Audit's own verdict: {real})")
        self.post("system", f"TEST SWITCH J2_FORCE_AUDIT_FAIL={n} (humans only, agents don't see this): Audit's "
                            f"result for done-when item {n} (\"{row['criterion']}\") is overridden to NEEDS_REVISION "
                            f"with the note \"{row['note']}\". Audit's own verdict: {real}. Later audits are real.",
                  warning=True, test=True)
        return n

    def _save_playbook(self):
        """J2.h learns plans: a run that ended Done leaves a playbook for its kind of business. Code records what
        happened (workflows in build order, costs, what failed and what fixed it, the measured outcome); Adaptation
        writes the reusable part (sections, mistakes to avoid, advice). One per run; a later Done replaces it."""
        g = self.goal or {}
        built = sorted((w for w in self.workflows if w["status"] == "done"), key=lambda w: w.get("version", 0))
        fixed = []
        for prev, cur in zip(self.audits, self.audits[1:]):
            now_ok = {r["criterion"] for r in cur["results"] + cur.get("rubric", []) if r["pass"]}
            ran = [w["name"] for w in self.workflows if prev["version"] < w.get("version", 0) <= cur["version"]]
            fixed += [{"problem": f"{r['criterion']}: {r.get('note')}", "fixed_by": ran} for r in
                      prev["results"] + prev.get("rubric", []) if not r["pass"] and r["criterion"] in now_ok and ran]
        first = self.audits[0] if self.audits else None
        outcome = {"cycles": self.cycle, "audits": len(self.audits), "ai_calls": self.calls,
                   "first_pass_audit": bool(first) and not workable(first), "budget_used": self.spent,
                   "change_requests": [m["text"] for m in self.chat if m["who"] == "user" and m.get("change")]}

        def mock():
            return json.dumps({"say": f"Playbook saved for a {g.get('business_type') or 'business'} site.",
                               "sections": [w["name"] for w in built], "watch_out": [x["problem"] for x in fixed][:5],
                               "advice": "Build the scaffold and style guide first, then the main sections."})
        r = self.call("adaptation", '{"say": str (one line for the chat), "sections": [str] (the sections a site for this '
                      'kind of business needs, in build order, named generally), "watch_out": [str] (what failed and '
                      'what fixed it, one line each, reusable on another site), "advice": str (one or two sentences for '
                      'the next run of this kind)}',
                      f"THE RUN ENDED DONE. Write the playbook for a {g.get('business_type') or 'business'} site.\n"
                      f"WORKFLOWS IN BUILD ORDER: {json.dumps([{k: w.get(k) for k in ('name', 'purpose', 'cost', 'cycle')} for w in built], ensure_ascii=False)}"
                      f"\nWHAT FAILED AND WHAT FIXED IT (from Audit): {json.dumps(fixed, ensure_ascii=False)}\n"
                      f"OUTCOME (measured by code): {json.dumps(outcome, ensure_ascii=False)}", mock)
        pb = {"id": f"P-{self.run_id}", "run": self.run_id, "t": time.strftime("%Y-%m-%d %H:%M:%S"),
              "business_type": str(g.get("business_type") or "business").lower(), "brief": self.request[:300],
              "visitor_needs": g.get("visitor_needs") or [], "sections": r.get("sections") or [w["name"] for w in built],
              "build_order": [w["name"] for w in built],
              "workflows": [{k: w.get(k) for k in ("name", "cost", "cycle", "status", "revises")} for w in self._all_wfs()],
              "costs": {"used": self.spent, "budget": self.budget}, "what_failed": fixed,
              "watch_out": r.get("watch_out") or [], "advice": r.get("advice") or "", "outcome": outcome,
              **({"test": True} if self.force_fail else {}),
              **({"used_playbook": self.playbooks_used[0]["playbook"]} if self.playbooks_used else {})}
        save_playbook(pb)
        self.playbook_saved = pb["id"]
        self.agent_turns["adaptation"] = self.agent_turns.get("adaptation", 0) + 1
        self._say("adaptation", (r.get("say") or "Playbook saved.") + f" Saved {pb['id']} ({pb['business_type']}).", None,
                  None, act="PLAYBOOK", playbook=pb["id"])

    def _learn(self, prev, cur):
        """J2.h learns: Audit confirmed a fix worked (a problem in the previous audit is gone, and a workflow ran in
        between), so Adaptation writes a lesson (problem, cause, fix) and code saves it to memory/lessons.jsonl."""
        now_ok = {r["criterion"] for r in cur["results"] + cur.get("rubric", []) if r["verdict"] in ("PASS", "PASS_WITH_WARNING")}
        fixed = [{"id": r["id"], "criterion": r["criterion"], "verdict": r["verdict"], "note": r.get("note")}
                 for r in prev["results"] + prev.get("rubric", []) if not r["pass"] and r["criterion"] in now_ok]
        test = any(r.get("forced_test") for r in prev["results"] if r["criterion"] in now_ok)   # kept out of the prompt
        gone = set(j2_check.problems(prev.get("mechanical") or {})) - set(j2_check.problems(cur.get("mechanical") or {}))
        fixed += [{"id": "mechanical", "criterion": p, "verdict": "render check", "note": p} for p in sorted(gone)
                  if (cur.get("mechanical") or {}).get("available")]
        ran = [{"id": w["id"], "name": w["name"], "purpose": w["purpose"], "completion_check": w["completion_check"]}
               for w in self.workflows if w["status"] == "done" and prev["version"] < w.get("version", 0) <= cur["version"]]
        if not fixed or not ran:
            return

        def mock():
            return json.dumps({"say": f"Audit confirmed {len(fixed)} fix(es); lesson saved.", "lessons": [
                {"problem": f"{x['criterion']}: {x['note']}", "cause": "The first build left it out",
                 "fix": ", ".join(w["name"] for w in ran)} for x in fixed]})
        r = self.call("adaptation", '{"say": str (one line for the chat), "lessons": [{"problem": str, "cause": str, '
                      '"fix": str}] (one per fixed problem)}',
                      f"AUDIT CONFIRMED THESE FIXES WORKED (failing in v{prev['version']}, passing in v{cur['version']}): "
                      f"{json.dumps(fixed, ensure_ascii=False)}\nWORKFLOWS THAT RAN IN BETWEEN: "
                      f"{json.dumps(ran, ensure_ascii=False)}\nWrite one lesson per fixed problem for future runs on any "
                      "website: the problem in general terms, its cause, and the fix that worked. One sentence each.", mock)
        saved = []
        for x in [x for x in r.get("lessons") or [] if isinstance(x, dict) and x.get("problem") and x.get("fix")][:len(fixed)]:
            lesson = {"id": f"L-{self.run_id}-{len(self.lessons_saved) + 1}", "problem": str(x["problem"]),
                      **({"output": self.output} if self.doc else {}),
                      **({"test": True} if test else {}),
                      "cause": str(x.get("cause") or ""), "fix": str(x["fix"]), "run": self.run_id, "cycle": self.cycle,
                      "versions": [prev["version"], cur["version"]], "workflows": [w["id"] for w in ran],
                      "t": time.strftime("%Y-%m-%d %H:%M:%S")}
            save_lesson(lesson)
            self.lessons_saved.append(lesson)
            saved.append(lesson["id"])
        self.agent_turns["adaptation"] = self.agent_turns.get("adaptation", 0) + 1
        if saved:
            self._say("adaptation", (r.get("say") or "Lesson saved.") + f" Saved {', '.join(saved)} to memory.", None,
                      None, act="LESSON", lessons=saved)

    # ------------------------------------------------------------------ run report
    REVIEW_FMT = ('{"verdict": str (one sentence on how the department did), '
                  '"agents": {"identity"|"adaptation"|"priority"|"coordination"|"audit"|<workflow agent id, e.g. "W-1">: '
                  '{"score": 1-5, "note": str}}, "went_well": [str], "problems": [str], "learnings": [str]}')

    def _metrics(self):
        """Facts counted by code: no model judgment in here."""
        ev = lambda t, a=None: [e for e in self.events if e["type"] == t and (a is None or e["agent"] == a)]
        decisions = [m.get("decision") for m in self.chat if m["who"] == "priority" and m.get("decision")]
        agents = {fn: {"turns": self.agent_turns.get(fn, 0), "redirects": len(ev("redirect", fn)),
                       "invalid_moves": len(ev("invalid_move", fn))}
                  for fn in ("identity", "adaptation", "priority", "coordination", "audit")
                  + tuple(w["id"] for w in self._all_wfs())}
        est_vs_price = [{"id": w["id"], "name": w["name"], "agent_estimate": w.get("estimate"), "charged": w["cost"]}
                        for w in self._all_wfs()]
        audits = [{"cycle": a["cycle"], "version": a["version"], "all_pass": a["all_pass"],
                   "pass": sum(r["verdict"].upper() == "PASS" for r in a["results"]),
                   "warn": sum(r["verdict"].upper() == "PASS_WITH_WARNING" for r in a["results"]),
                   "fail": sum(not r["pass"] for r in a["results"]),
                   "rubric_fail": sum(not r["pass"] for r in a.get("rubric", [])),
                   "mechanical": j2_check.problems(a.get("mechanical") or {}),
                   "mechanical_ran": bool((a.get("mechanical") or {}).get("available"))} for a in self.audits]
        cs = self.conflicts
        pauses = [e.get("pause") for e in ev("pause")]
        return {"outcome": self.outcome, "brief": self.request, "model": None if llm.MOCK else llm.MODEL,
                "mock": llm.MOCK, "cycles": self.cycle, "ai_calls": self.calls,
                "minutes": round((time.time() - self.started) / 60, 1) if self.started else 0,
                "budget": {"total": self.budget, "start": BUDGET, "cap": BUDGET_CAP, "raises": self.raises,
                           "spent": self.spent, "rate_card": bool(RATES)},
                "site_versions": len(self.sites), "site_bytes": len(self.sites[-1].encode("utf-8")) if self.sites else 0,
                "agents": agents,
                "negotiation": {"proposals": sum(1 for m in self.chat if m["who"] == "adaptation" and m.get("proposal")),
                                "grants": decisions.count("grant"), "counters": decisions.count("counter"),
                                "refuses": decisions.count("refuse"),
                                "refuses_converted_by_code": sum("Treated as REFUSE" in e["detail"]
                                                                 for e in ev("invalid_move", "priority")),
                                "parks": sum(1 for m in self.chat if m.get("act") == "PARK")},
                "conflicts": {"total": len(cs),
                              "settled_by_agents": sum(c["resolved_by"] == "agents" for c in cs),
                              "resolved_by_identity": sum(c["resolved_by"] == "identity" for c in cs),
                              "asked_requester": sum(bool(c["escalated_to_user"]) for c in cs),
                              "open": sum(c["status"] != "settled" for c in cs),
                              "list": [{"id": c["id"], "cycle": c["cycle"], "topic": c["topic"], "between": c["between"],
                                        "attempts": sum(1 for x in c["attempts"] if x["n"] > 0),
                                        "escalated": c["escalated"], "resolved_by": c["resolved_by"],
                                        "user_question": (c["user_question"] or {}).get("question"),
                                        "user_answer": c["user_answer"], "outcome": c["outcome"]} for c in cs]},
                "requester_questions": {"total": len(pauses), "by_kind": {k: pauses.count(k) for k in set(pauses)}},
                "workflows": {"registered": len(self._all_wfs()),
                              "done": sum(w["status"] == "done" for w in self._all_wfs()),
                              "failed": sum(w["status"] == "failed" for w in self._all_wfs()),
                              "held_ever": len({e["detail"].split(":")[0] for e in ev("hold")}),
                              "estimate_vs_charged": est_vs_price},
                "validator_refusals": len(ev("refused")) + len(ev("role_refusal")),
                "role_refusals": len(ev("role_refusal")), "pauses": len(pauses), "stuck": len(ev("stuck")),
                "limit_failures": len(ev("limit_fail")), "warnings": len(ev("warning")),
                "audits": audits, "to_be_confirmed": (self.goal or {}).get("to_be_confirmed") or [],
                "workable": workable(self.audits[-1]) if self.audits else None,
                "lessons": {"memory": self.memory, "saved": [x["id"] for x in self.lessons_saved],
                            "used": self.lessons_used}}

    def _links(self):
        """Hand-offs in the chat, "who->next": workflow agents count as "W", the Requester as "requester"."""
        def node(x):
            x = str(x).lower()
            return "W" if re.match(r"w-\d+", x) else "requester" if x in ("user", "requester") else x
        out = {}
        for m in self.chat:
            if m.get("next") and m["who"] not in ("system", "user"):
                k = f"{node(m['who'])}->{node(m['next'])}"
                out[k] = out.get(k, 0) + 1
        return out

    def _export(self):
        """A document's export when the run ends (prd.docx next to prd.md). If it can't be written, say so; the
        markdown file is kept."""
        self.export_note = self.out["export"](self.sites[-1], self.dir) if self.out.get("export") else None
        f = self.out["file"][:-3] + ".docx"
        self.post("system", f"Exported {f} from {self.out['file']}." if not self.export_note else
                  f"Export failed: {self.export_note}. {self.out['file']} is kept.", **({} if not self.export_note
                                                                                        else {"warning": True}))

    def _write_summary(self):
        """runs/<run>/summary.json when a run ends (done, stopped, out of budget or crashed). Code only, no AI call."""
        o = self.outcome or ""
        status = next((k for p, k in (("Done with open items", "done_with_open_items"), ("Done", "done"), ("Crashed", "crashed"))
                       if o.startswith(p)), "stopped")
        last = self.audits[-1] if self.audits else None
        conds = (last or {}).get("conditions")
        if self.docset:
            self.save("chain.json", self._chain_record())
        try:      # atom/1 views: a new file; a failure here never changes the run
            views, warns = j2_schema.run_views(self._all_wfs(), self.given, ended=True)
            self.save("atoms.json", {"schema": j2_schema.SCHEMA, "atoms": views, "warnings": warns})
        except Exception as e:
            self._event("system", "atoms_json", f"atoms.json not written: {e}")
        self.save("summary.json", {
            "run": self.run_id, "output_type": self.output, "j2_fixes": int(self.fixes),
            **({"realization": {"applied": self.real is not None, **(self.real_summary or {})}} if self.realization else {}),
            **({"design_documents": {"requested": self.docset, "stages": self.stages, "files": {k: v["file"] for k, v in self.docs.items()},
                                     "chain": self._chain_record(), "lld_revisions": self.lld_revisions,
                                     "lld_compare": (self.audits[-1].get("lld") if self.audits else None)}}
               if self.docset else {}),
            "provider": None if llm.MOCK else llm.PROVIDER,
            "sections": [{"heading": h, "words": md_words(b)} for lv, h, b in md_sections(self.sites[-1]) if lv == 2]
            if self.doc and self.sites else list(self.sections),
            "conditions": {"met": sum(c["met"] for c in conds), "total": len(conds),
                           "warnings": sum(c.get("severity") == "warn" and not c["met"] for c in conds),
                           "list": [{k: c[k] for k in ("id", "met", "evidence")} for c in conds]} if conds else None,
            "necessary_conditions": {"met": sum(c["met"] for c in conds), "total": len(conds),
                                     "warnings": sum(c.get("severity") == "warn" and not c["met"] for c in conds),
                                     "list": [{k: c[k] for k in ("id", "text", "met", "evidence", "severity",
                                                                 "assumptions") if k in c} for c in conds]}
            if conds else None,
            **({"export": {"file": self.out["file"][:-3] + ".docx", "written": not self.export_note,
                           "note": self.export_note}} if self.doc else {}),
            "force_audit_fail": self.force_fail or None, "test_lessons": self.test_lessons,
            "force_question": self.force_question or None, "has_board": bool(self.board["entries"]),
            "model": None if llm.MOCK else llm.MODEL, "mock": llm.MOCK, "brief": self.request,
            "change_requests": [m["text"] for m in self.chat if m["who"] == "user" and m.get("change")],
            "status": status, "outcome": o,
            "budget": {"start": BUDGET, "granted": self.budget, "used": self.spent, "raises": self.raises},
            "ai_calls": self.calls, "cycles": self.cycle,
            "links": self._links(),
            "verification": {**{k: sum(a.get("status") == k for w in self._all_wfs() for a in w.get("atoms") or [])
                                for k in ("passed", "failed", "escalated")},
                             "could_not_run": sum(a.get("status") == "could not run" for w in self._all_wfs()
                                                  for a in w.get("atoms") or []),
                             "checks_run": len(self.receipts)},
            "identity_to_requester_decisions": sum(m["who"] == "identity" and m.get("pause") in DECISION_PAUSES
                                                   for m in self.chat),
            "identity_to_requester_facts": sum(m["who"] == "identity" and m.get("pause") in ("fact", "info")
                                               for m in self.chat),
            "priority_rulings": {d: sum(m["who"] == "priority" and m.get("decision") == d for m in self.chat)
                                 for d in ("grant", "counter", "refuse")},
            "questions_to_user": self.decisions,
            "conflicts": [{"id": c["id"], "pair": c["between"], "topic": c["topic"], "subject": c["subject"],
                           "attempts": sum(1 for x in c["attempts"] if x["n"] > 0), "escalated": c["escalated"],
                           "resolved_by": c["resolved_by"], "repriced": bool(c.get("repriced")),
                           "user_question": (c["user_question"] or {}).get("question"), "user_answer": c["user_answer"],
                           "outcome": c["outcome"]} for c in self.conflicts],
            "validator_rejections": self.rejections,
            "holds": [{"cycle": e["cycle"], "by": e["agent"], "detail": e["detail"]} for e in self.events
                      if e["type"] == "hold"],
            "audits": [{"cycle": a["cycle"], "version": a["version"], "all_pass": a["all_pass"],
                        "verdicts": [{"id": r.get("id"), "item": r.get("item"), "verdict": r.get("verdict")}
                                     for r in a["results"] + a.get("rubric", []) + a.get("extra_issues", [])]}
                       for a in self.audits],
            "site_versions": len(self.sites),
            "workable": workable(last) if last else None,
            "to_be_confirmed": (self.goal or {}).get("to_be_confirmed") or [],
            "dropped_by_requester": (self.goal or {}).get("dropped_by_requester") or [],
            "workflow_agents": [{"id": w["id"], "agent": w.get("agent"), "section": w.get("section"),
                                 "revises": w.get("revises"), "status": w["status"], "calls": w["calls_used"],
                                 "retries": w["retries_used"]} for w in self._all_wfs()],
            "limits_hit": [e["detail"] for e in self.events if e["type"] == "limit"],
            **({"milestone_emails": self.milestone_log} if self.milestone_log else {}),
            "lessons": {"memory": self.memory, "saved": [x["id"] for x in self.lessons_saved], "used": self.lessons_used},
            "playbooks": {"saved": self.playbook_saved, "used": self.playbooks_used},
            "business_type": (self.goal or {}).get("business_type"), "style_guide": bool(self.style),
            "rubric": [{"id": r["id"], "pass": r["pass"], "verdict": r["verdict"]} for r in (last or {}).get("rubric", [])],
            "done_when": [{"item": self._no(c), "criterion": c,
                           "pass": next((r["pass"] for r in (last or {}).get("results", []) if r.get("criterion") == c),
                                        None)}
                          for c in (self.goal or {}).get("done_when") or []],
            "route_check_violations": [{"agent": e["agent"], "detail": e["detail"], "msg": e["msg"], "cycle": e["cycle"]}
                                       for e in self.events if e["type"] == "redirect"]})
        self._write_artifacts()

    def _write_artifacts(self):
        """artifacts.json: a new file listing what the run made (artifacts/types.json). Written with the summary and
        again once the report exists; a failure here never changes the run."""
        try:
            j2_artifacts.write(self.dir)
        except Exception as e:
            self._event("system", "artifacts_json", f"artifacts.json not written: {e}")
        try:      # map.json (schema map/1), built from the files above; same rule: never changes the run
            j2_map.write(self.dir)
        except Exception as e:
            self._event("system", "map_json", f"map.json not written: {e}")

    def _review(self, metrics):
        """One extra model call, outside J2: a reviewer reads the run and writes learnings. Its scores are judgment."""
        roles = []
        for name in ("team.md", "identity.md", "adaptation.md", "priority.md", "coordination.md", "audit.md",
                     "builder.md"):
            path = os.path.join(J2_DIR, name)
            if os.path.exists(path):
                with open(path, encoding="utf-8") as f:
                    roles.append(f"### {name}\n{j2_realization.PLANNING_PART.sub('', f.read()).strip()}")

        def mock():
            c = metrics["conflicts"]
            return json.dumps({
                "verdict": f"Mock run: {metrics['outcome']}",
                "agents": {fn: {"score": 4 if a["redirects"] + a["invalid_moves"] == 0 else 3,
                                "note": f"{a['turns']} turns, {a['redirects']} redirects, {a['invalid_moves']} invalid moves"}
                           for fn, a in metrics["agents"].items()},
                "went_well": [f"{c['total']} conflict(s): {c['settled_by_agents']} settled by the agents, "
                              f"{c['resolved_by_identity']} by Identity, {c['asked_requester']} needed the Requester."],
                "problems": [f"{sum(a['redirects'] for a in metrics['agents'].values())} illegal route pick(s) were redirected."],
                "learnings": ["Mock replies are scripted, so these learnings only show the report works."]})
        if self.calls and DELAY:
            time.sleep(DELAY)
        self.calls += 1
        raw = llm.ask(
            "You review one test run of J2, a department of AI agents with strict roles (role files below). You are not "
            "part of the department. Judge from the evidence only: the code-counted metrics, the event log and the chat. "
            "Cite chat message ids like [12]. Score each agent 1-5 on staying in its role and moving the goal forward. "
            "Learnings are about J2's design and about this test system: what to change, keep or test next, each tied to "
            "evidence. Say so when the evidence is thin. Plain English, short items, no hype.\n\n"
            "## Reply format (required)\n" + self.REVIEW_FMT + "\n\n## Role files\n" + "\n\n".join(roles),
            f"METRICS:\n{json.dumps(metrics, ensure_ascii=False)}\n\nEVENTS:\n{json.dumps(self.events, ensure_ascii=False)}"
            f"\n\nCHARTER:\n{json.dumps(self.goal, ensure_ascii=False)}\n\nCHAT:\n{self.chat_view()}",
            mock, max_tokens=2000, effort="high", json_mode=True)
        try:
            return llm.extract_json(raw)
        except ValueError:
            return {"verdict": "The reviewer's reply wasn't valid JSON.", "raw": raw[:2000]}

    def _make_report(self, review=True):
        """Metrics always; the reviewer call only when the run ends (not on every pause for the Requester)."""
        metrics = self._metrics()
        self.thinking = None
        try:
            review = self._review(metrics) if review else {"verdict": "Paused for the Requester; the reviewer runs "
                                                                      "when the run ends."}
        except Exception as e:  # noqa: BLE001  the metrics are still worth saving
            review = {"verdict": f"Reviewer call failed: {type(e).__name__}: {str(e)[:200]}"}
        metrics["ai_calls"] = self.calls
        self.report = {"metrics": metrics, "review": review, "events": self.events,
                       "made": time.strftime("%Y-%m-%d %H:%M:%S")}
        self.save("report.json", self.report)
        with open(os.path.join(self.dir, "report.md"), "w", encoding="utf-8") as f:
            f.write(self._report_md())
        self._write_artifacts()

    def _report_md(self):
        m, r = self.report["metrics"], self.report["review"]
        n, cf = m["negotiation"], m["conflicts"]
        out = [f"# J2 run report: {self.run_id}", "", f"Brief: {m['brief']}", "",
               f"**Outcome:** {m['outcome']}", f"**Reviewer verdict:** {r.get('verdict', '')}", "",
               "## Measured (counted by code)",
               f"- {m['ai_calls']} AI calls, {m['cycles']} cycles, {m['minutes']} min, model {m['model'] or 'mock'}",
               f"- Budget {m['budget']['spent']}/{m['budget']['total']} (started {m['budget']['start']}, "
               f"{len(m['budget']['raises'])} raise(s) approved by the Requester, cap {m['budget']['cap']}; rate card "
               f"{'on' if m['budget']['rate_card'] else 'off'}), {m['site_versions']} site versions, final site "
               f"{m['site_bytes'] / 1000:.0f} KB",
               f"- Negotiation: {n['proposals']} proposals, {n['grants']} GRANT / {n['counters']} COUNTER / "
               f"{n['refuses']} REFUSE ({n['refuses_converted_by_code']} converted by code), {n['parks']} parked",
               f"- Conflicts {cf['total']}: {cf['settled_by_agents']} settled by the agents, {cf['resolved_by_identity']} "
               f"resolved by Identity from the charter, {cf['asked_requester']} needed the Requester",
               f"- Questions to the Requester {m['requester_questions']['total']} {m['requester_questions']['by_kind']}",
               f"- Validator refusals {m['validator_refusals']} (role {m['role_refusals']}), "
               f"limit failures {m['limit_failures']}, warnings {m['warnings']}", "",
               "| Agent | Turns | Redirects | Invalid moves | Reviewer score | Note |", "|---|---|---|---|---|---|"]
        for fn, a in m["agents"].items():
            sc = (r.get("agents") or {}).get(fn) or {}
            out.append(f"| {NAMES.get(fn, fn)} | {a['turns']} | {a['redirects']} | {a['invalid_moves']} | "
                       f"{sc.get('score', '')} | {str(sc.get('note', '')).replace('|', '/')} |")
        out += ["", "## Conflicts (full log in conflicts.json)"] + (
            [f"- {c['id']} (cycle {c['cycle']}, {c['topic']}): {' vs '.join(NAMES[x] for x in c['between'])}, "
             f"{c['attempts']} attempt(s), " + ("escalated to Identity, " if c["escalated"] else "")
             + f"resolved by {c['resolved_by']}"
             + (f"; asked the Requester \"{c['user_question']}\", answer \"{c['user_answer']}\"" if c["user_question"] else "")
             + f". Outcome: {c['outcome']}" for c in cf["list"]] or ["- None"])
        out += ["", "## Audit by cycle"] + [f"- Cycle {a['cycle']}, v{a['version']}: {a['pass']} pass, {a['warn']} warn, "
                                             f"{a['fail']} fail; mechanical: "
                                             + ("not run" if not a.get("mechanical_ran") else
                                                ", ".join(a["mechanical"]) or "clean") for a in m["audits"]]
        for title, k in (("What went well", "went_well"), ("Problems", "problems"), ("Learnings", "learnings")):
            out += ["", f"## {title} (reviewer)"] + [f"- {x}" for x in r.get(k) or []]
        return "\n".join(out) + "\n"
