"""Offline check of the J2 flow in mock mode: python j2_flow_test.py [--fast [--browser]] [--chat]

--fast stubs the headless-browser measurement (the render check and tap targets read the page without Chrome) and
skips the tests that drive the page in Chrome; --fast --browser runs those and the real render check, with only tap
targets stubbed. Without --fast everything uses Chrome.

Each scenario runs with J2_FIXES=0 (baseline) and J2_FIXES=1 (fixes):
1. A vague brief: J2.a asks the Requester, a conflict the charter can't decide goes to the Requester, Audit and
   Adaptation settle a finding between themselves, the run ends Done. Then a change request after done.
   Run twice: the Requester keeps Priority's position, or picks the stories option (baseline: the validator
   rejects it because Priority refused it; fixes: Priority reprices the ruling and funds it).
2. A full brief (Brew Lab): J2.a passes without questions; Identity resolves the conflict from the charter.
   Baseline: the ruling goes straight to the validator, which rejects what Priority refused.
   Fixes: the ruling goes back to Priority once; with too little budget, Identity asks the Requester.
3. Only Identity talks to the Requester: an agent that addresses the Requester is redirected to Identity.
B. One agent per workflow: each owns a marked section; it retries within its envelope; "revises" hands over a section.
C. Workable loop: the same failure twice, the cycle limit and to-be-confirmed facts each pause for the Requester
   (keep working, drop, ship); the J2.c bargain's time bound escalates to Identity; workable() unit checks.
D. Memory: a confirmed fix saves a lesson, the next run's Adaptation uses it, J2_MEMORY=0 reads and writes nothing.
E. The page: Agents and Lessons tabs, the state they read, and the clear-memory endpoint.
I. Cost follows the work (runs/j2-20261007-190411: 13 items for 35 by bundling, a fix relabelled as styling): each
   covered item has a rate-card minimum, a fix is priced as a fix, Priority sees cost per item.
H. Any business, better each run: J2.a proposes visitor needs for the business type; the first workflow writes the
   style guide everyone follows; Audit's design rubric loops a measured failure back; playbooks saved at Done and
   reused for a similar business; no cafe examples in live prompts.
G. The Simple view (default page): rendered in headless Chrome for a start screen, questions, a limit question and a
   finished run; its visible text never has agent names, J2 codes, "budget" or other internals.
F. J2_FORCE_AUDIT_FAIL=<item>: a forced cycle-1 failure that agents can't tell from a real finding (no test label in
   any prompt), then a real fix, a real re-check and a saved lesson. summary.json exists before the reviewer call.
   J2_MEMORY_TEST_LESSONS=1 lets J2.b read test lessons (off by default), flagged in the chat and the lanes.
Runs are written to a temp folder (J2_RUNS_DIR), never the real runs/; each ended run has summary.json. Memory is off (J2_MEMORY=0) except in
the memory test, which uses its own lessons file.
"""
import inspect
import json
import os
import re
import sys
import time

import hashlib
import tempfile

os.environ["STUDIO_MOCK"] = "1"
os.environ["J2_MEMORY"] = "0"        # clean runs: the memory tests turn it on
os.environ["J2_REALIZATION"] = "0"   # the classic path (realization is on by default; realization_test.py covers it)
# Every memory file of this test session lives in a temp folder, set before j2 is imported, so no test can write
# lessons or playbooks into the real memory/ (it once wrote 109 mock cafe playbooks there; a live run reused one).
SESSION_MEMORY = tempfile.mkdtemp(prefix="j2-test-memory-")
os.environ["J2_MEMORY_FILE"] = os.path.join(SESSION_MEMORY, "lessons.jsonl")
os.environ["J2_PLAYBOOK_FILE"] = os.path.join(SESSION_MEMORY, "playbooks.jsonl")
os.environ["J2_RUNS_DIR"] = os.path.join(SESSION_MEMORY, "runs")   # runs/ holds real runs only: report evidence
REAL_MEMORY = [os.path.join(os.path.dirname(os.path.abspath(__file__)), "memory", f) for f in ("lessons.jsonl", "playbooks.jsonl")]


def fingerprint(paths):
    return {p: hashlib.sha256(open(p, "rb").read()).hexdigest() if os.path.exists(p) else None for p in paths}


REAL_BEFORE = fingerprint(REAL_MEMORY)
import j2  # noqa: E402
for _k in ("J2_SEND_MODE", *j2.j2_send.SMTP_VARS, "J2_SEND_CAP"):      # .env is never used here: the tests never send a real email
    os.environ.pop(_k, None)
ATOM = [j2.mock_atom({})]            # a complete work atom for workflows a test registers directly

BREW_LAB = ("A one-page website for Brew Lab, a small specialty coffee shop in Indiranagar, Bengaluru. Visitors should "
            "see the menu with prices, the opening hours, and where to find us, and be able to pre-order a coffee for "
            "pickup. It should feel warm and modern and work well on a phone. Include a section about our beans and the "
            "farms they come from, and a simple loyalty-card explainer (buy 9, get the 10th free).")
BOUTIQUE_BRIEF = ("make a website for my cousin's clothing boutique called Thread & Loom in Pune. it sells handloom "
                  "kurtas (1499, 1699), sarees (3499), cotton shirts (1299) and dupattas (899). people should be able "
                  "to browse products by category, filter by size and price, see a size guide, check new arrivals, "
                  "and order on whatsapp (+91 00000 11111). it should feel earthy and elegant and work well on phones")
# The same boutique with its timings and address, so an earlier workflow owns that part of the page: the message
# board demo (board_functions). Every price, timing and number here is in the brief; the phone numbers are fake.
BOUTIQUE_HOURS_BRIEF = (BOUTIQUE_BRIEF.replace(" it should feel earthy",
                        " the shop is open 11am to 8pm every day except monday, at 14 Laxmi Road in Pune, and the "
                        "phone is +91 00000 22222. it should feel earthy"))
SHOW = "--chat" in sys.argv
FAST = "--fast" in sys.argv
BROWSER = not FAST or "--browser" in sys.argv        # the tests that drive the page in Chrome
REAL_FIND, REAL_TAP, REAL_CHECK = j2.j2_check.find_chrome, j2.j2_check.tap_targets, j2.j2_check.check


def fake_tap_targets(html, selector, min_px=j2.j2_check.TAP_MIN):
    """--fast: no browser. The inline width/height (or min-width/min-height) in px of the buttons, links and fields
    inside the elements matching selector; a control with no inline size counts as big enough."""
    name, kind = selector.lstrip("#."), selector[:1]
    attr = "id" if kind == "#" else "class"
    roots, small, checked = 0, [], 0
    for m in re.finditer(rf'<(\w+)\b[^>]*\b{attr}="(?:[^"]*\s)?{re.escape(name)}(?:\s[^"]*)?"[^>]*>', html):
        roots += 1
        tag, depth, end = m.group(1), 0, len(html)
        for t in re.finditer(rf"<(/?){tag}\b[^>]*>", html[m.start():]):
            depth += -1 if t.group(1) else 1
            if not depth:
                end = m.start() + t.end()
                break
        for c in re.finditer(r'<(button|a|input|select|textarea)\b([^>]*)>([^<]*)', html[m.start():end]):
            if c.group(1) == "a" and "href" not in c.group(2) or 'type="hidden"' in c.group(2):
                continue
            checked += 1
            style = dict(re.findall(r"(min-height|min-width|height|width)\s*:\s*(\d+)px", c.group(2)))
            h = max(int(style.get("height", 0)), int(style.get("min-height", 0))) or min_px
            w = max(int(style.get("width", 0)), int(style.get("min-width", 0))) or min_px
            if h < min_px or w < min_px:
                small.append(f'{c.group(1)} "{c.group(3).strip()[:25]}" {w}x{h}')
    return {"available": True, "roots": roots, "checked": checked, "small_count": len(small), "small": small[:6]}


def fake_check(html):
    """--fast: the render check without a browser: static checks, no overflow, no JS errors, forms submit."""
    st = j2.j2_check.static_checks(html)
    return {"available": True, **st, "phone": {"width": 360, "overflow_px": 0},
            "desktop": {"width": 1280, "overflow_px": 0}, "js_errors": [], "design": {},
            "form_results": [{"index": i + 1, "has_submit": True, "submitted": True, "invalid": [], "errors_during": [],
                              "ok": True} for i in range(st["forms"])]}


class FakeDriver:
    """A behaviour-step driver with no browser: elements by selector ({visible, text, classes, style}) and handlers
    keyed ("set"|"click"|"scroll", selector) or "wait" that change them, as a page's script would."""
    def __init__(self, els, on=None):
        self.els, self.on = els, on or {}

    def __call__(self, ops):
        out = []
        for o in ops:
            el = self.els.get(o.get("sel"))
            if o.get("do") == "wait":
                self.on.get("wait", lambda e, o: None)(self.els, o)
                out.append({"v": True})
            elif el is None:
                out.append({"missing": True})
            elif o.get("do"):
                if el.get("error"):
                    out.append({"error": el["error"]})
                    continue
                self.on.get((o["do"], o["sel"]), lambda e, o: None)(self.els, o)
                out.append({"v": True})
            else:
                r = o["read"]
                out.append({"v": el.get("visible", True) if r == "visible" else o["cls"] in el.get("classes", ())
                            if r == "class" else el.get("style", {}).get(o["prop"], "") if r == "style"
                            else el.get("text", "")})
        return {"available": True, "results": out}


def fake_behavior(html, ops):
    """--fast: behaviour steps on the static page (no script runs): ids and classes, hidden or not, their text."""
    els = {}
    for m in re.finditer(r"<(\w+)\b([^>]*)>([^<]*)", html):
        attrs, cls = m.group(2), re.search(r'\bclass="([^"]+)"', m.group(2))
        el = {"visible": not re.search(r"\shidden\b|display:\s*none", attrs), "text": m.group(3).strip(),
              "classes": set(cls.group(1).split()) if cls else set()}
        for sel in ([f"#{x}" for x in re.findall(r'\bid="([^"]+)"', attrs)] + [f".{c}" for c in el["classes"]]):
            els.setdefault(sel, el)
    return FakeDriver(els)(ops)


if FAST:
    j2.j2_check.tap_targets = fake_tap_targets
    if not BROWSER:
        j2.j2_check.behavior = fake_behavior
    if not BROWSER:     # --fast --browser keeps the real render check: the rubric tests need its measurements
        j2.j2_check.check = fake_check
        j2.j2_check.find_chrome = lambda: None      # the page-driving tests skip; screenshots aren't taken


def say(run, text):
    assert run.user_says(text)
    while run.status != "waiting":
        time.sleep(0.02)


def show(run):
    if SHOW:
        print(f"\n===== {run.dir}")
        for m in run.chat:
            print(f"[{m['id']}] {j2.NAMES.get(m['who'], m['who'])}: {m['text']}")


def finish(run):
    """Answer what a person would, until the run ends: a conflict question with the simpler option, a budget
    limit with "keep working"."""
    for _ in range(6):
        if not run.paused:
            return
        say(run, "Option 2" if run.paused["kind"] == "conflict" else "1")


def add_budget_if_asked(run):
    """J2_FIXES=1 funds Identity's ruling (the stories), so the viewport fix later finds the budget spent and Identity
    asks for more. Answer "1" (keep working, +30), only for that reason."""
    while run.paused and run.paused["kind"] == "limit":
        assert run.paused["handoff"]["limit"].startswith("Out of budget"), run.paused["handoff"]["limit"]
        say(run, "1")


def only_identity_addresses_user(run):
    for m in run.chat:
        if m.get("next") == "user" or m.get("to_user"):
            assert m["who"] == "identity", m


def summary(run):
    with open(os.path.join(run.dir, "summary.json"), encoding="utf-8") as f:
        s = json.load(f)
    with open(os.path.join(run.dir, "chat.jsonl"), encoding="utf-8") as f:
        head = json.loads(f.readline())
    assert head["meta"] and head["j2_fixes"] == s["j2_fixes"] == int(run.fixes), (head, s["j2_fixes"])
    return s


def vague(fixes):
    run = j2.J2Run()
    say(run, "a website for my cafe")
    assert run.paused and run.paused["kind"] == "info" and not run.goal, run.outcome
    say(run, "It's called Little Fig, in Koramangala, Bengaluru. Visitors should see our menu with prices, opening "
             "hours and where to find us. Cozy and simple, and it must work on phones.")
    assert run.goal and run.paused and run.paused["kind"] == "conflict", run.outcome
    c = run.conflicts[0]
    assert c["escalated"] and c["attempts_before_escalation"] == 4 and c["user_question"], c
    say(run, "Option 2: skip the stories, keep the reserve for fixes.")
    c = run.conflicts[0]
    assert c["resolved_by"] == "identity_after_user" and c["user_answer"] and c["winner"] == "priority", c
    f = next(x for x in run.conflicts if x["topic"] == "finding")
    assert f["resolved_by"] == "agents" and not f["escalated"] and f["winner"] == "audit", f
    assert run.outcome.startswith("Done"), run.outcome
    s = summary(run)
    assert s["status"] == "done" and len(s["questions_to_user"]) == 2 and len(s["conflicts"]) == 2, s
    n = len(run.chat)
    say(run, "Add a section for visiting hours")
    first = next(m for m in run.chat[n:] if m["who"] not in ("user", "system"))
    assert first["who"] == ("identity" if fixes else "adaptation"), first   # fix (b): J2.a before J2.b
    assert run.outcome.startswith("Done"), run.outcome
    assert summary(run)["change_requests"] == ["Add a section for visiting hours"]
    only_identity_addresses_user(run)
    show(run)


def vague_stories(fixes):
    """J2.d asks the Requester and they pick the stories option. Baseline: Identity rules it in, but the ruling goes
    straight to the validator, which rejects what Priority refused (the Requester's choice doesn't override it).
    Fixes: the ruling costs more than Priority granted, so Priority reprices it once and funds it."""
    run = j2.J2Run()
    say(run, "a website for my cafe")
    say(run, "It's called Little Fig, in Koramangala, Bengaluru. Visitors should see our menu with prices, opening "
             "hours and where to find us. Cozy and simple, and it must work on phones.")
    assert run.paused and run.paused["kind"] == "conflict", run.outcome
    say(run, "Option 1: fund the stories section.")
    c = run.conflicts[0]
    assert c["resolved_by"] == "identity_after_user" and c["winner"] == "split" and c["user_answer"], c
    stories = [w for w in run.workflows if w["name"] == "Stories section"]
    if fixes:
        assert c["repriced"] and c["repricing"]["decision"] == "GRANT" and stories and not run.rejections, c
    else:
        assert not c.get("repriced") and not stories, c
        assert [r["workflow"] for r in run.rejections] == ["Stories section"] and "Priority" in run.rejections[0]["reason"]
    assert run.outcome.startswith("Done"), run.outcome
    s = summary(run)
    assert s["conflicts"][0]["user_answer"].startswith("Option 1") and s["validator_rejections"] == run.rejections, s
    only_identity_addresses_user(run)
    show(run)


def full(fixes):
    run = j2.J2Run()
    say(run, BREW_LAB)
    assert not any(e["type"] == "pause" and e.get("pause") == "info" for e in run.events), "J2.a asked questions"
    c = run.conflicts[0]
    assert c["resolved_by"] == "identity" and not c["escalated_to_user"] and c["winner"] == "split", c
    stories = [w for w in run.workflows if w["name"] == "Stories section"]
    if fixes:   # repriced by Priority, granted, registered
        assert c["repriced"] and c["repricing"]["decision"] == "GRANT" and stories, c
    else:       # straight to the validator, which rejects what Priority refused
        assert not c.get("repriced") and not stories, c
        assert any(r["workflow"] == "Stories section" and "Priority" in r["reason"] for r in run.rejections), run.rejections
        assert any(m["who"] == "system" and m["text"].startswith("Validator rejected 'Stories section'") for m in run.chat)
    assert run.outcome.startswith("Done"), run.outcome
    s = summary(run)
    assert s["status"] == "done" and all(d["pass"] for d in s["done_when"]) and s["route_check_violations"], s
    only_identity_addresses_user(run)
    show(run)


def ruling_over_budget():
    """Fixes only: the ruling (90) is over the 85 left; Priority counters; Identity asks; the Requester adds budget.
    The ruling then takes the whole budget, so the viewport fix can't be funded: the run doesn't stop, Identity asks
    (keep working with +30, drop, ship). Keep working funds the fix and the site becomes workable."""
    os.environ["J2_FIXES"] = "1"
    run = j2.J2Run()
    run.budget = 85
    say(run, BREW_LAB)
    assert run.paused and run.paused["kind"] == "ruling_budget", run.outcome
    say(run, "1, add the units")
    c = run.conflicts[0]
    assert c["repricing"]["decision"] == "COUNTER" and c["resolved_by"] == "identity_after_user", c
    assert run.budget == 90 and any(w["name"] == "Stories section" for w in run.workflows), run.budget
    assert run.paused and run.paused["kind"] == "limit" and run.paused["handoff"]["limit"].startswith("Out of budget"), \
        run.outcome
    q = run.paused["questions"][0]
    assert q["options"][0].startswith("Add 30 work credits and keep going") and q["budget"][0] == 30 and q["options"][-1].startswith("Finish as it is")
    say(run, "1")
    assert run.budget == 120 and run.outcome.startswith("Done"), (run.budget, run.outcome)
    s = summary(run)
    assert s["status"] == "done" and s["limits_hit"] and not s["workable"], s


def built_first(run):
    """runs/j2-20261009-133203: the decision question about unfunded items comes after the funded work is built and
    audited, never while registered work is unbuilt."""
    assert run.paused and run.paused["kind"] == "limit" and not run.paused["handoff"].get("after_ruling"), run.outcome
    assert run.sites and run.audits and run.audits[-1]["version"] == len(run.sites), run.outcome
    assert not any(w["status"] in ("registered", "held", "running") for w in run.workflows), run.workflows
    reg = next(m["id"] for m in run.chat if m["who"] == "validator" and m.get("register"))
    order = [m["who"] for m in run.chat if m["id"] > reg]
    assert order.index("coordination") < order.index("audit") < len(order) - 1 - order[::-1].index("identity"), order


def policy_budget(fixes):
    """runs/j2-20261008-153402 (budget 40, rate-card floor 110): Adaptation asked about affordability twice and never
    proposed, so Priority was never called and nothing was built. Now a budget POLICY_QUESTION is rejected by code and
    Adaptation is asked once more; it proposes the best set that fits, Priority rules, and only then Identity asks one
    decision question (add credits or leave out; no "finish", there's no site). An unclear answer is asked once more;
    "option 1" raises the budget by code; the rest is proposed, granted and built."""
    run = j2.J2Run()
    run.budget = 60                                          # Brew Lab's floor is 70
    orig, asked = run.call, []

    def call(fn, fmt, task, mock, **k):
        if fn == "adaptation" and "WHY YOU ARE UP" in task:
            asked.append("REJECTED" in task)
            if len(asked) == 1:
                return {"say": "@Identity Policy question: 60 can't cover the charter's sections. Can we add budget?",
                        "act": "POLICY_QUESTION", "question": "60 can't cover the charter's sections. Add budget?"}
        assert not (fn == "identity" and task.startswith("QUESTION FROM")), "a budget question reached Identity"
        return orig(fn, fmt, task, mock, **k)
    run.call = call
    say(run, BREW_LAB)
    assert asked[:2] == [False, True], asked                 # regenerated once, with the rejection in the task
    assert any(m["who"] == "system" and m["text"].startswith("Adaptation's question about budget was rejected")
               for m in run.chat)
    i = {m["who"]: m["id"] for m in reversed(run.chat) if m["who"] != "validator" or m.get("register")}
    assert i["adaptation"] < i["priority"] < i["validator"], i   # proposal, then Priority's ruling, then registration
    built_first(run)
    q = run.paused["questions"][0]
    assert run.paused["handoff"]["choices"] == ["continue", "drop", "ship"], run.paused["handoff"]
    assert q["options"][0].startswith("Add 10 work credits and keep going") and q["options"][1].startswith("Leave out"), q
    say(run, "the bigger one I think")                       # unclear: no guess, Identity asks once more
    assert run.paused and run.paused["handoff"].get("reasked") and run.budget == 60, run.paused
    say(run, "option 1")
    assert run.budget == 70 and run.raises[0]["asked"] == 10, run.raises
    d = [x for x in run.decisions if x["kind"] == "limit"]
    assert d[0].get("unclear") and d[1]["decision"] == "continue" and d[1]["budget_added"] == 10, d
    assert run.spent == 70 and all(w["status"] == "done" for w in run.workflows), (run.spent, run.workflows)
    assert run.outcome.startswith("Done") and not run.goal.get("dropped_by_requester"), run.outcome
    assert summary(run)["budget"]["granted"] == 70
    only_identity_addresses_user(run)
    show(run)


def connections():
    """Every J2 link works at three budgets: generous (2x the rate-card floor), tight (floor - 15, cap above the
    floor) and impossible (cap below the floor). Adaptation proposes, Priority rules on every proposal, the validator
    registers or rejects with a reason, Coordination schedules, Audit checks what was built, and every run ends with
    a site built (the Requester always picks leave-out or keep going here). summary.json counts the links."""
    j2.load_rates()
    floor = sum(j2.item_min(c) for c in j2.mock_charter(BREW_LAB)["done_when"])
    assert floor == 70, floor
    for name, budget, cap in (("generous", 2 * floor, None), ("tight", floor - 15, None),
                              ("impossible", floor - 30, floor - 30)):
        old_cap = j2.BUDGET_CAP
        try:
            if cap:
                j2.BUDGET_CAP = cap
            assert j2.BUDGET_CAP > floor if name != "impossible" else j2.BUDGET_CAP < floor, name
            run = j2.J2Run()
            run.budget = budget
            say(run, BREW_LAB)
            asked = []
            for _ in range(6):
                if not run.paused:
                    break
                p = run.paused
                asked.append(p)
                assert p["kind"] != "limit" or not any(w["status"] in ("registered", "held") for w in run.workflows), \
                    (name, "paused with registered, unbuilt work")
                if p["kind"] == "limit":
                    ch = p["handoff"]["choices"]
                    say(run, str(ch.index("drop" if "drop" in ch else "continue") + 1))
                else:
                    say(run, "Option 2" if p["kind"] == "conflict" else "1")
        finally:
            j2.BUDGET_CAP = old_cap
        s, chat = summary(run), run.chat
        links = s["links"]
        assert links.get("adaptation->priority", 0) >= 1, (name, links)                     # Adaptation proposed
        for i, m in enumerate(chat):                     # Priority ruled on every proposal sent to it
            if m["who"] == "adaptation" and m.get("next") == "priority":
                ruling = next((x for x in chat[i + 1:] if x["who"] == "priority"), None)
                assert ruling and ruling.get("decision") in ("grant", "counter", "refuse"), (name, m["text"], ruling)
        reg = [m for m in chat if m["who"] == "validator" and "register" in m]
        assert reg and all(m["register"] or "Rejected:" in m["text"] for m in reg), (name, reg)
        assert all(r["reason"] for r in run.rejections), (name, run.rejections)
        assert links.get("coordination->priority", 0) >= 1 and links.get("priority->W", 0) >= 1, (name, links)
        assert run.sites and run.audits and links.get("audit->identity", 0) >= 1, (name, run.outcome)
        decisions = [m for m in chat if m["who"] == "identity" and m.get("pause") in j2.DECISION_PAUSES]
        assert s["identity_to_requester_decisions"] == len(decisions), (name, s["identity_to_requester_decisions"])
        facts = [m for m in chat if m["who"] == "identity" and m.get("pause") in ("fact", "info")]
        assert s["identity_to_requester_facts"] == len(facts), (name, s["identity_to_requester_facts"])
        rulings = [m["decision"] for m in chat if m["who"] == "priority" and m.get("decision")]
        assert s["priority_rulings"] == {d: rulings.count(d) for d in ("grant", "counter", "refuse")} \
            and sum(s["priority_rulings"].values()) == len(rulings) >= 1, (name, s["priority_rulings"], rulings)
        limits = [p for p in asked if p["kind"] == "limit"]
        for p in limits:                                 # "finish" only once a site exists; never one option
            opts = p["questions"][0]["options"]
            assert len(opts) >= 2, (name, opts)
            built = any(re.match(r"W-\d+$", m["who"]) and m["id"] < p["msg"] for m in chat)
            assert built or not any(o.startswith("Finish as it is") for o in opts), (name, opts)   # finish needs a site
        if name == "tight":
            assert any(m.get("decision") == "counter" for m in chat) or \
                any("drop" in p["handoff"]["choices"] for p in limits), (name, links)
        if name == "impossible":
            first = limits[0] if limits else None
            assert first and "drop" in first["handoff"]["choices"] and "continue" not in first["handoff"]["choices"], first
            # the finish code applied (the only option left) tells the Requester why, and what's missing
            i = next(i for i, m in enumerate(chat) if m["text"].startswith("Only one option was possible"))
            told = chat[i + 1]
            form = next(c for c in run.goal["done_when"] if "pre-order" in c)
            assert told["who"] == "identity" and told.get("to_user") and told.get("done"), told
            assert "We've reached a limit: you've approved the most work credits allowed" in told["text"], told["text"]
            assert f"{form} can't be paid for and can't be left out" in told["text"], told["text"]
            assert f"Still not right (1):\n- item {run._no(form)} {j2.short(form)} (unfunded)" in told["text"], told["text"]
            assert run.outcome.startswith("Done with open items: shipped as is as the only option"), run.outcome
            plain_for_requester(run)
        if name == "generous":
            assert run.outcome == j2.WORKABLE_OUTCOME, run.outcome


def build_before_asking():
    """runs/j2-20261009-133203: 6 workflows were registered across two rulings, but the run went validator -> Identity's
    decision question each time and the site stayed at v0; Identity told the Requester "the site now has..." about
    unbuilt work; the gate item was extended with filters; Audit filed items 5-7 under their pre-drop numbers."""
    # (1) budget 40: the funded workflows are built and audited before the question
    run = j2.J2Run()
    run.budget = 40
    say(run, BREW_LAB)
    built_first(run)
    assert all(w["status"] == "done" for w in run.workflows) and len(run.workflows) >= 1, run.workflows
    # (2) nothing has passed Audit: Identity may not describe anything as built ("will build" instead)
    run = j2.J2Run()
    run.goal = j2.mock_charter(BREW_LAB)
    run.workflows = [{"id": "W-1", "name": "Menu", "status": "registered", "covers_text": [run.goal["done_when"][0]]}]
    run.call = lambda fn, fmt, task, mock, **k: {"say": "So far the page has the earthy look and all your prices."} \
        if fn == "identity" else json.loads(mock())
    run._identity_limit({"reason": "Priority's ruling", "left_out": run.goal["done_when"][1:], "budget_needed": 20})
    text = run.chat[-1]["text"]
    assert "page has" not in text and "will build 1 funded part" in text, text
    assert any(e["type"] == "invalid_move" and "as built" in e["detail"] for e in run.events)
    # (3) the gate's wording is fixed by code: it can't name features
    run = j2.J2Run()
    orig = run.call

    def call(fn, fmt, task, mock, **k):
        if fn == "identity" and '"core": [bool]' in fmt:
            r = json.loads(mock())
            r["goal"]["done_when"] = r["goal"]["done_when"] + [
                "The page runs with no JavaScript errors, filters and category browsing work, and every in-page link "
                "points to an existing section."]
            r["goal"]["core"] = [False] * (len(r["goal"]["done_when"]) - 1) + [True]
            return r
        return orig(fn, fmt, task, mock, **k)
    run.call = call
    run.budget = 200
    say(run, BREW_LAB)
    gate = run.goal["done_when"][-1]
    assert gate == j2.GATE_TEXT and not re.search(r"filter|categor", gate, re.I), gate
    assert gate in run.goal["core_items"] and j2.must_keep(gate, run.goal), run.goal["core_items"]
    # (5) after a drop, Audit's pre-drop numbers are read by text; dropped items get no finding; nothing is "missing"
    run = j2.J2Run()
    orig_items = [c for c, _ in THREAD_AND_LOOM]
    keep = [orig_items[i] for i in (0, 1, 6, 8, 11)]
    run.goal = {"done_when": keep, "charter_items": orig_items,
                "dropped_by_requester": [c for c in orig_items if c not in keep]}
    findings = [{"id": f"F-{i}", "item": i, "criterion": c[:70], "verdict": "PASS", "evidence": "seen"}
                for i, c in ((1, keep[0]), (2, keep[1]))]
    findings += [{"id": "F-3", "item": 7, "criterion": "Every product has an order button to WhatsApp", "verdict": "PASS",
                  "evidence": "seen"},                                   # old number 7 -> item 3
                 {"id": "F-4", "item": 9, "criterion": "Store section with address, hours and map link",
                  "verdict": "PASS", "evidence": "seen"},                # old number 9 -> item 4
                 {"id": "F-5", "item": 12, "criterion": "No JavaScript errors", "verdict": "PASS", "evidence": "seen"},
                 {"id": "F-6", "item": 8, "criterion": "Lookbook with 3 styled outfits", "verdict": "NEEDS_REVISION",
                  "evidence": "absent"},                                 # dropped: no finding
                 {"id": "F-7", "item": None, "criterion": "A reviews section shows three reviews", "verdict":
                  "NEEDS_REVISION", "evidence": "absent"}]               # dropped, unnumbered: no finding
    results, extras = run._audit_findings(findings, keep)
    # findings are named by the items' run-long numbers: 1, 2, 7, 9, 12 (never renumbered after the drop)
    assert [r["id"] for r in results] == ["F-1", "F-2", "F-7", "F-9", "F-12"] and all(r["pass"] for r in results), results
    assert not extras, extras


def below_floor():
    """runs/j2-20261008-153402: budget below the rate-card floor. Adaptation proposes the best set that fits and lists
    the rest in left_out (core facts first); Priority rules; the funded work is built and audited; then Identity asks
    one decision question with a leave-out option. A question never has one option; "finish" needs a site."""
    run = j2.J2Run()
    run.budget = 50                                          # Brew Lab's floor is 70; its core items cost 45
    say(run, BREW_LAB)
    assert run.goal["core_items"] == [c for c in run.goal["done_when"] if re.search(r"menu|hours|address|order", c)]
    a = next(m for m in run.chat if m["who"] == "adaptation")
    p = next(m for m in run.chat if m["who"] == "priority")
    assert a["id"] < p["id"] and a.get("proposal") and p.get("decision") == "grant", (a, p)
    crit = run.goal["done_when"]
    funded = {t for w in run.workflows for t in w["covers_text"]}
    assert all(c in funded for c in crit if j2.must_keep(c, run.goal)), funded    # core facts are never left out
    left = [c for c in crit if c not in funded]
    built_first(run)
    assert run.paused["handoff"]["choices"] == ["continue", "drop", "ship"], run.paused
    assert [crit[i - 1] for i in run.paused["handoff"]["failing_items"]] == left, run.paused["handoff"]
    assert len([e for e in run.events if e.get("type") == "pause"]) == 1, "more than one question"
    say(run, "2")
    assert run.goal["dropped_by_requester"] == left and not run.paused, run.goal
    assert run.sites and run.outcome == j2.WORKABLE_OUTCOME, run.outcome
    # at the cap: the funded work is built first, so the question has a site to finish (and no "add credits")
    run = j2.J2Run()
    run.budget = j2.BUDGET_CAP = 40
    try:
        say(run, BREW_LAB)
        built_first(run)
        q = run.paused["questions"][0]
        assert run.paused["handoff"]["choices"] == ["drop", "ship"] and q["options"][1].startswith("Finish as it is"), q
    finally:
        j2.BUDGET_CAP = 2 * j2.BUDGET
    # only "add credits" would be possible (every left-out item is a core fact): stop here is the second option
    run = j2.J2Run()
    run.goal = {"done_when": ["The menu with prices is shown", "Opening hours are listed"]}
    run._identity_limit({"reason": "Priority's ruling", "left_out": run.goal["done_when"], "budget_needed": 20})
    assert run.paused["handoff"]["choices"] == ["continue", "ship"], run.paused["handoff"]



THREAD_AND_LOOM = [   # runs/j2-20261008-153402's charter (12 items, floor 110), with Identity's core marks
    ("The page shows the name Thread & Loom and Pune, has a viewport meta tag, and uses an earthy colour palette "
     "defined in the CSS.", True),
    ("The site lists exactly these 10 products with their prices: Indigo block-print kurta 1,499; Mustard handloom "
     "kurta 1,699; White chikankari kurta 2,199; Ikat handloom saree 3,499; Khadi silk saree 4,999; Cotton jamdani "
     "saree 3,999; Khadi cotton shirt 1,299; Linen-blend shirt 1,599; Block-print dupatta 899; Chanderi dupatta 1,099.", True),
    ("Visitors can filter products by category (All, Kurtas, Sarees, Shirts, Dupattas), by size, and by price, and the "
     "filters work together.", False),
    ("Size options match the supplied data: Indigo block-print kurta S to XXL; Mustard handloom kurta S to XL; White "
     "chikankari kurta XS to XL; Khadi cotton shirt S to XXL; Linen-blend shirt M to XXL; all sarees and dupattas Free "
     "size.", False),
    ("A size guide table shows chest in inches: XS 32, S 34, M 36, L 38, XL 40, XXL 42.", False),
    ("A New Arrivals section shows exactly the White chikankari kurta, Ikat handloom saree and Linen-blend shirt.", False),
    ("Every product has an order button that opens a WhatsApp link to 910000011111 with the product name and chosen "
     "size in the message text.", True),
    ("A lookbook section shows 3 styled outfits, each built only from the listed products and naming them.", False),
    ("A store section shows 18 FC Road, Shivajinagar, Pune 411005; hours 11am to 9pm, Tuesday to Sunday, closed "
     "Monday; shop phone +91 00000 22222; and a map link for the address.", True),
    ("A policy section states: exchange within 7 days if unworn with tags; no returns on sale items; and delivery in "
     "Pune in 2 days for Rs 79, free above 2,999, pay on delivery.", False),
    ("A reviews section shows these three reviews: \"Beautiful fabric, fits perfectly\" - Meera K.; \"Fast delivery and "
     "lovely packing\" - Arjun S.; \"My go-to for festive kurtas\" - Priya D.", False),
    ("The page has no JavaScript errors, no horizontal overflow at 360px or 1280px, and no broken in-page links.", True)]


def thread_and_loom_core():
    """Follow-up to 3a8960c: "any number" made prices, sizes, "3 styled outfits" and "360px" core, so almost nothing in
    the Thread & Loom charter could be left out. Identity marks core items at J2.a; at budget 40 the decision question
    offers the lookbook, reviews and new arrivals as leave-out options, never a core item."""
    run = j2.J2Run()
    run.budget = 40
    orig = run.call

    def call(fn, fmt, task, mock, **k):
        if fn == "identity" and '"core": [bool]' in fmt:              # J2.a records the charter with core marks
            r = json.loads(mock())
            r["goal"].update(site_name="Thread & Loom", business_type="clothing boutique",
                             done_when=[c for c, _ in THREAD_AND_LOOM], core=[k for _, k in THREAD_AND_LOOM])
            return r
        return orig(fn, fmt, task, mock, **k)
    run.call = call
    say(run, "A website for Thread & Loom, a handloom clothing boutique in Pune. " * 3)
    crit = run.goal["done_when"]
    expect = [j2.GATE_TEXT if "JavaScript errors" in c else c for c, k in THREAD_AND_LOOM if k]
    assert run.goal["core_items"] == expect, run.goal["core_items"]
    built_first(run)
    offered = [crit[i - 1] for i in run.paused["handoff"]["failing_items"]]
    for words in ("lookbook", "reviews section", "New Arrivals"):
        assert any(words in c for c in offered), (words, offered)
    assert not any(j2.must_keep(c, run.goal) for c in offered), offered
    # the keyword fallback no longer treats a number as a core fact
    assert not j2.must_keep(THREAD_AND_LOOM[7][0], {}) and j2.must_keep(THREAD_AND_LOOM[1][0], {})


def deferred_twice():
    """The same done-when item left out or refused in 2 rounds: Identity asks the Requester even though budget remains."""
    run = j2.J2Run()
    run.goal = {"done_when": ["The menu with prices is shown", "A lookbook section shows 3 styled outfits"],
                "core_items": ["The menu with prices is shown"]}
    run.budget, look = 100, "A lookbook section shows 3 styled outfits"
    run.a_named = [look]
    run._parked("Priority refused it")                       # round 1: budget remains, Adaptation proposes again
    assert run.turn == "adaptation" and run.handoff.get("new_cycle"), run.handoff
    run.cycle += 1
    run.a_named = [look]
    run._parked("Priority refused it")                       # round 2: the Requester decides
    assert run.turn == "identity" and run.handoff["kind"] == "limit" and run.handoff["deferred"], run.handoff
    run._identity_limit(run.handoff)
    q, h = run.paused["questions"][0], run.paused["handoff"]
    assert h["choices"] == ["continue", "drop"] and h["failing_items"] == [2], h
    assert q["question"].startswith("Some parts were put off twice.") and "budget" not in q, q

def redirect():
    """Audit addresses the Requester: the message goes to Identity first, with a Route check note."""
    run = j2.J2Run()
    run.goal = {"done_when": ["a", "b", "c"]}
    run.cycle = 1
    run._handoff("audit", kind="audit")
    run._say("audit", "@Requester can you send the phone number?", "adaptation")
    assert run.relay and run.relay["from"] == "audit"
    notes = [m["text"] for m in run.chat if m["who"] == "system"]
    assert notes and notes[-1].startswith("Route check: Audit addressed the Requester"), notes
    assert run.chat[-2]["text"].startswith("@Identity"), run.chat[-2]


def workflow_agents():
    """B: every registered workflow gets its own agent and section; the site file is assembled from the sections."""
    run = j2.J2Run()
    orig, bad = run.call, []

    def call(fn, fmt, task, mock, **k):                 # W-2's first reply can't be used: it retries once
        if fn == "W-2" and not bad:
            bad.append(1)
            return '{"say": "@Coordination here is the page"}\n```html\n<!doctype html><html><body>x</body></html>\n```'
        return orig(fn, fmt, task, mock, **k)
    run.call = call
    say(run, BREW_LAB)
    assert run.outcome.startswith("Done"), run.outcome
    site = run.sites[-1]
    for w in run.workflows:
        assert w["agent"] == f"{w['id']} Builder" and w["section"] == w["id"] and w["status"] == "done", w
        assert any(m["who"] == w["id"] for m in run.chat), f"{w['id']} never spoke"
        assert f"<!-- J2:{w['id']} " in site, w["id"]
    assert not any(m["who"] == "builder" for m in run.chat)
    w2 = run._wf("W-2")
    assert w2["retries_used"] == 1 and any(m["text"].startswith("W-2 attempt 1 can't be used") for m in run.chat), w2
    assert site.index("<!-- J2:W-1 start") < site.index("<!-- J2:W-2 start"), "sections out of order"
    vp = next(w for w in run.workflows if "viewport" in w["name"].lower())
    assert f'<!-- J2:{vp["id"]} head start' in site and site.index('name="viewport"') < site.index("</head>")
    assert os.path.exists(os.path.join(run.dir, "sections.json")) and run.write_lock is None
    # revises: a new workflow rewrites W-2's section and becomes its owner
    run.budget, run.p_refused = 500, {}
    reg, refused, _ = run._register([{"name": "Core sections rework", "goal_gap": "tighter copy", "outputs": ["W-2"], "atoms": ATOM,
                                      "completion_check": "same items", "parts": ["fix"], "revises": "W-2"}])
    assert reg and not refused and reg[0]["section"] == "W-2" and reg[0]["revises"] == "W-2", (reg, refused)
    run._turn_workflow_agent({"wid": reg[0]["id"]})
    assert run.sections["W-2"]["owner"] == reg[0]["id"] and f"(owner {reg[0]['id']})" in run.sites[-1]
    # past retry_limit: the agent reports the failure to Coordination and its section is unchanged
    reg, _, _ = run._register([{"name": "Broken one", "goal_gap": "x", "outputs": ["y"], "completion_check": "z", "atoms": ATOM,
                                "parts": ["fix"], "fixes": ["F-9"]}])
    run.call = lambda fn, *a, **k: "no blocks at all"
    n = len(run.sites)
    run._turn_workflow_agent({"wid": reg[0]["id"]})
    assert reg[0]["status"] == "failed" and len(run.sites) == n and run.turn == "coordination", reg[0]
    assert run.chat[-1]["who"] == reg[0]["id"] and "failed" in run.chat[-1]["text"]
    show(run)


def same_failure():
    """C: Audit fails item 1 every time; the same failure twice in a row pauses with keep / drop / ship. Drop
    removes the item, Audit re-checks, the site is workable."""
    def failing_run():
        run = j2.J2Run()
        orig = run.call

        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            if fn == "audit" and "findings" in r:          # not Audit's dispute turns
                f = next((x for x in r["findings"] if x["item"] == 1), None)    # charter item 1, the menu
                if f and "menu" in f["criterion"].lower():
                    f.update(verdict="NEEDS_REVISION", evidence="Prices missing")
            if fn == "adaptation" and "WHY YOU ARE UP" in task and run.audits and not run.audits[-1]["results"][0]["pass"]:
                r = {"say": "@Priority Menu fix.", "act": "PROPOSE", "next": "Priority", "workflows": j2.with_atoms([
                    {"name": "Menu fix", "goal_gap": "prices", "parts": ["fix"], "fixes": ["F-1"], "outputs": ["menu"],
                     "completion_check": "prices"}])}
            return r
        run.call = call
        return run
    run = failing_run()
    say(run, BREW_LAB)
    assert run.paused and run.paused["kind"] == "limit" and "same failure" in run.paused["handoff"]["limit"], run.outcome
    opts = run.paused["questions"][0]["options"]
    assert run.paused["handoff"]["choices"] == ["continue", "drop", "ship"] and len(opts) == 3, opts
    say(run, "2, drop it")
    assert run.outcome.startswith("Done") and run.goal["dropped_by_requester"][0].startswith("The menu"), run.outcome
    assert summary(run)["dropped_by_requester"]
    only_identity_addresses_user(run)

    run2 = failing_run()
    say(run2, BREW_LAB)
    say(run2, "3, ship it as is")
    s = summary(run2)
    assert run2.outcome.startswith("Done with open items") and s["status"] == "done_with_open_items" and s["workable"], run2.outcome


def cycle_limit():
    """C: J2_MAX_CYCLES=1. Audit fails the phone item in cycle 1, so the next cycle would pass the limit: Identity
    asks. "Keep working" gives a fresh allowance and the run goes on to a workable site."""
    old, j2.MAX_CYCLES = j2.MAX_CYCLES, 1
    try:
        run = j2.J2Run()
        say(run, BREW_LAB)
        assert run.paused and run.paused["kind"] == "limit" and "1 cycles" in run.paused["handoff"]["limit"], run.outcome
        assert not run.paused["questions"][0].get("budget"), "no budget to add when the budget isn't the limit"
        say(run, "Option 1, keep going")
        assert run.outcome.startswith("Done") and run.decisions[-1]["decision"] == "continue", run.outcome
    finally:
        j2.MAX_CYCLES = old


def to_be_confirmed():
    """C / A.2: the Requester can't give the menu prices yet. Item 1 is to be confirmed, so it fails at Audit and the
    site isn't workable; once only that is left, Identity asks; the Requester gives the prices; workable."""
    run = j2.J2Run()
    orig = run.call

    def call(fn, fmt, task, mock, **k):
        r = orig(fn, fmt, task, mock, **k)
        if fn == "identity" and task.startswith("REQUESTER'S BRIEF") and r.get("goal"):
            r["goal"]["to_be_confirmed"] = [{"item": 1, "fact": "the real menu prices"}]
        return r
    run.call = call
    say(run, BREW_LAB)
    assert run.goal["to_be_confirmed"][0]["criterion"] == run.goal["done_when"][0]
    first = run.audits[0]["results"][0]
    assert first["verdict"] == "TO_BE_CONFIRMED" and not first["pass"], first
    assert run.paused and run.paused["kind"] == "limit" and run.paused["handoff"]["limit"].startswith("Only you"), \
        (run.paused or {}).get("handoff", {}).get("limit") or run.outcome
    say(run, "1. Espresso 120, latte 180")
    assert not run.goal["to_be_confirmed"] and run.goal["confirmed_facts"], run.goal
    add_budget_if_asked(run)
    assert run.outcome.startswith("Done"), run.outcome


def bargain_time():
    """C / A.4: the J2.c bargain is bounded by time too; at 0 minutes the first COUNTER goes straight to Identity."""
    old, j2.BARGAIN_MINUTES = j2.BARGAIN_MINUTES, 0
    try:
        run = j2.J2Run()
        say(run, BREW_LAB)
        c = run.conflicts[0]
        assert c["escalated"] and c["attempts_before_escalation"] == 1 and "minute" in c["escalated_because"], c
        add_budget_if_asked(run)
        assert run.outcome.startswith("Done"), run.outcome
    finally:
        j2.BARGAIN_MINUTES = old


def workable_unit():
    ok = {"results": [{"item": 1, "id": "F-1", "verdict": "PASS", "criterion": "a", "pass": True}],
          "mechanical": {"available": True, "phone": {"overflow_px": 0}, "desktop": {"overflow_px": 0}, "js_errors": [],
                         "broken_in_page_links": [], "form_results": [{"index": 1, "ok": True}]}}
    assert j2.workable(ok) == []
    bad = json.loads(json.dumps(ok))
    bad["mechanical"]["phone"]["overflow_px"] = 40
    bad["mechanical"]["form_results"] = [{"index": 1, "ok": False, "has_submit": False, "errors_during": []}]
    assert j2.workable(bad) == ["phone overflow 40px", "form 1 has no submit button"], j2.workable(bad)
    off = {**ok, "mechanical": {"available": False, "note": "Chrome not found"}}
    assert j2.workable(off)[0].startswith("render check didn't run"), "unverified must not count as workable"


def memory():
    """D: a fix Audit confirms becomes a lesson; the next run's Adaptation reads it and says which one it used;
    J2_MEMORY=0 reads and writes nothing; clear_lessons empties the file."""
    tmp = tempfile.mkdtemp()
    old, j2.MEMORY_FILE = j2.MEMORY_FILE, os.path.join(tmp, "lessons.jsonl")
    os.environ["J2_MEMORY"] = "1"
    try:
        run = j2.J2Run()
        say(run, BREW_LAB)                 # cycle 1 fails the phone item, the viewport fix closes it in cycle 2
        lessons = j2.load_lessons()
        assert lessons and lessons[0]["run"] == run.run_id and all(lessons[0][k] for k in ("problem", "cause", "fix"))
        assert any(m.get("act") == "LESSON" for m in run.chat) and run.lessons_saved
        run2 = j2.J2Run()
        say(run2, BREW_LAB)
        m = next(m for m in run2.chat if m.get("lesson_used"))
        assert m["who"] == "adaptation" and m["lesson_used"] in m["text"] and run2.lessons_used, m
        assert summary(run2)["lessons"]["used"][0]["lesson"] == m["lesson_used"]
        n = len(j2.load_lessons())
        os.environ["J2_MEMORY"] = "0"
        run3 = j2.J2Run()
        say(run3, BREW_LAB)
        assert not run3.memory and not any(x.get("lessons_offered") or x.get("act") == "LESSON" for x in run3.chat)
        os.environ["J2_MEMORY"] = "1"
        assert len(j2.load_lessons()) == n, "J2_MEMORY=0 wrote a lesson"
        assert j2.clear_lessons() == n and j2.load_lessons() == [] and j2.clear_lessons() == 0
    finally:
        j2.MEMORY_FILE = old
        os.environ["J2_MEMORY"] = "0"


def forced_fail():
    """F: J2_FORCE_AUDIT_FAIL=2 forces item 2 to NEEDS_REVISION in cycle 1 only, marked as a test. Identity sends it
    back, Adaptation proposes a fix, its workflow agent runs, Audit re-checks for real, the lesson is saved (flagged
    as a test lesson, so later runs don't read it). Off when unset; an item number the charter lacks is ignored."""
    tmp = tempfile.mkdtemp()
    old, j2.MEMORY_FILE = j2.MEMORY_FILE, os.path.join(tmp, "lessons.jsonl")
    os.environ.update(J2_MEMORY="1", J2_FORCE_AUDIT_FAIL="2")
    try:
        run = j2.J2Run()
        run.budget = 150      # room for the extra fix in both modes (J2_FIXES=1 also funds the stories)
        say(run, BREW_LAB)
        add_budget_if_asked(run)
        first, last = run.audits[0], run.audits[-1]
        r1 = next(x for x in first["results"] if x["item"] == 2)
        assert r1["verdict"] == "NEEDS_REVISION" and r1["forced_test"] and r1["real_verdict"].startswith("PASS"), r1
        assert r1["note"] == j2.forced_note(r1["criterion"]) and "Sunday" in r1["note"], r1["note"]   # item 2: hours
        # humans see the test: the red note, the tag, chat.jsonl and summary.json
        assert any(m.get("test") and m["text"].startswith("TEST SWITCH J2_FORCE_AUDIT_FAIL=2") for m in run.chat)
        audit_msg = next(m for m in run.chat if m["who"] == "audit")
        assert audit_msg.get("forced_test") == 2 and "F-2 item 2 NEEDS_REVISION" in audit_msg["text"], audit_msg["text"]
        with open(os.path.join(run.dir, "chat.jsonl"), encoding="utf-8") as f:
            assert json.loads(f.readline())["force_audit_fail"] == 2
        # agents can't tell: no prompt or reply format of any AI call carries a test label or Audit's real verdict
        with open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8") as f:
            calls = [json.loads(ln) for ln in f]
        for c in calls:
            seen = (c["system"] + c["prompt"]).lower()
            for bad in ("forced test", "test switch", "j2_force_audit_fail", "forced_test", "real_verdict", "own verdict",
                        "test failure", "not a real finding"):
                assert bad not in seen, (c["id"], c["fn"], bad)
        ident = next(c for c in calls if c["fn"] == "identity" and "AUDIT'S FINDINGS" in c["prompt"])
        assert r1["note"] in ident["prompt"], "Identity must see the forced finding as a normal one"
        fix = next(w for w in run.workflows if w["name"] == "Fix item 2")
        assert fix["status"] == "done" and fix["revises"] == "W-2" and fix["cycle"] > 1, fix
        r2 = next(x for x in last["results"] if x["item"] == 2)
        assert last["cycle"] > 1 and r2["verdict"] == "PASS" and not r2.get("forced_test"), r2   # the re-check is real
        assert sum(1 for a in run.audits for x in a["results"] if x.get("forced_test")) == 1, "forced more than once"
        assert run.outcome.startswith("Done"), run.outcome
        test_lessons = [x for x in j2.load_lessons() if x.get("test")]
        assert test_lessons and fix["id"] in test_lessons[0]["workflows"], j2.load_lessons()
        assert not any(x.get("test") for x in j2.relevant_lessons(j2.load_lessons(), "forced test failure item opening hours"))
        s = summary(run)
        assert s["force_audit_fail"] == 2 and s["status"] == "done", s
        os.environ["J2_FORCE_AUDIT_FAIL"] = "99"            # no such item: a warning, nothing forced
        run = j2.J2Run()
        say(run, BREW_LAB)
        add_budget_if_asked(run)
        assert any("J2_FORCE_AUDIT_FAIL=99 ignored" in m["text"] for m in run.chat) and run.outcome.startswith("Done")
        os.environ.pop("J2_FORCE_AUDIT_FAIL")                # unset: off
        run = j2.J2Run()
        assert run.force_fail == 0 and run.snapshot()["force_audit_fail"] is None
    finally:
        j2.MEMORY_FILE = old
        os.environ["J2_MEMORY"] = "0"
        os.environ.pop("J2_FORCE_AUDIT_FAIL", None)


def test_lessons():
    """J2_MEMORY_TEST_LESSONS=1 (demo only) lets Adaptation read test lessons at J2.b; by default they're filtered.
    A used test lesson is flagged in the chat (tag + TEST LESSON note) and on Adaptation's lane."""
    tmp = tempfile.mkdtemp()
    old, j2.MEMORY_FILE = j2.MEMORY_FILE, os.path.join(tmp, "lessons.jsonl")

    def only_test_lesson():
        with open(j2.MEMORY_FILE, "w") as f:
            f.write(json.dumps({"id": "L-test-1", "test": True, "problem": "Layout works on a phone: no viewport meta "
                                "tag", "cause": "forced test failure", "fix": "Viewport fix", "run": "x"}) + "\n")
    only_test_lesson()
    os.environ["J2_MEMORY"] = "1"
    try:
        run = j2.J2Run()                                    # default: the test lesson is never offered
        say(run, BREW_LAB)
        add_budget_if_asked(run)
        assert not run.test_lessons and not any(m.get("lessons_offered") for m in run.chat), "test lesson offered"
        os.environ["J2_MEMORY_TEST_LESSONS"] = "1"
        only_test_lesson()                                  # drop the real lesson the first run saved
        run = j2.J2Run()
        say(run, BREW_LAB)
        add_budget_if_asked(run)
        m = next(m for m in run.chat if m.get("lesson_used"))
        assert m["lesson_used"] == "L-test-1" and m["lesson_test"] and m["who"] == "adaptation", m
        assert any(x.get("test") and x["text"].startswith("TEST LESSON: Adaptation used L-test-1") for x in run.chat)
        snap = run.snapshot()
        lane = next(x for x in snap["lanes"] if x["id"] == "adaptation")
        assert snap["test_lessons"] and lane["test_lesson"] == "L-test-1", lane
        s = summary(run)
        assert s["test_lessons"] and s["lessons"]["used"][0]["test"], s["lessons"]
    finally:
        j2.MEMORY_FILE = old
        os.environ["J2_MEMORY"] = "0"
        os.environ.pop("J2_MEMORY_TEST_LESSONS", None)


def summary_before_review():
    """summary.json is written the moment a run ends, before the slow reviewer call (runs/j2-20261007-134200 had none
    for 98 s after "Done", and none at all if the server stopped then). Also when the reviewer fails or a turn crashes."""
    run = j2.J2Run()
    seen = []

    def review(metrics):
        seen.append(os.path.exists(os.path.join(run.dir, "summary.json")))
        raise RuntimeError("reviewer down")
    run._review = review
    say(run, BREW_LAB)
    add_budget_if_asked(run)
    assert seen and all(seen), seen
    assert summary(run)["status"] == "done" and "Reviewer call failed" in run.report["review"]["verdict"]
    run = j2.J2Run()                                     # a crash mid-run still leaves summary.json
    run._turn_coordination = lambda h: (_ for _ in ()).throw(RuntimeError("boom"))
    say(run, BREW_LAB)
    assert run.outcome.startswith("Crashed") and summary(run)["status"] == "crashed", run.outcome


SIMPLE_FORBIDDEN = re.compile(r"\b(Identity|Adaptation|Priority|Coordination|Audit|Validator|Builder|Requester)\b|J2|"
                              r"budget|\bW-\d|@|\bcycles?\b|AI calls?|\bmock\b|charter|done[-_ ]when|workflow|\bthe you\b",
                              re.I)
SIMPLE_FORBIDDEN_CASE = re.compile(r"\b(COUNTER|GRANT|REFUSE|Clause|NEEDS_REVISION|PASS)\b"
                                   r"|\b(?:after|before|by|from|until|than) the team\b")   # e.g. "after Audit" -> not "after the team"


def chrome_dom(cmd, capture_output=True, text=True, timeout=90):
    """Headless Chrome for a page check (j2_check.run_chrome: on a timeout Chrome and its helpers are killed). No retry: a page
    that hangs fails the test."""
    return j2.j2_check.run_chrome(cmd, timeout)


def simple_text(base):
    """Visible text of the Simple view (#sview) after the page has rendered, from headless Chrome's DOM."""
    import html.parser
    import subprocess
    chrome = j2.j2_check.find_chrome()
    out = chrome_dom([chrome, "--headless=new", "--disable-gpu", "--virtual-time-budget=6000", "--dump-dom",
                          base + "/?view=simple"], capture_output=True, text=True, timeout=90).stdout

    class P(html.parser.HTMLParser):
        depth, skip, text, buttons = 0, 0, [], []

        def handle_starttag(self, tag, attrs):
            a = dict(attrs)
            if self.depth or a.get("id") == "sview":
                self.depth += 1
                if tag in ("script", "style") or "hidden" in a:
                    self.skip += 1
                if tag == "button":
                    self.buttons.append(a)

        def handle_endtag(self, tag):
            if self.depth:
                self.depth -= 1
                if self.skip and tag in ("script", "style"):
                    self.skip -= 1

        def handle_data(self, d):
            if self.depth and not self.skip:
                self.text.append(d)
    p = P()
    p.text, p.buttons = [], []
    p.feed(out)
    body = re.search(r'<body class="([^"]*)"', out)
    return " ".join(" ".join(p.text).split()), p.buttons, (body.group(1).split() if body else [])


def simple_view():
    """G: the Simple view in five states. Plain words only; options are buttons; ready card when done."""
    import threading
    import j2_server
    if not j2.j2_check.find_chrome():
        print("Simple view test skipped: Chrome not found")
        return
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    keep = j2_server.RUN

    def check(run, *want, opts=0, once=(), site=None):
        j2_server.RUN = run
        text, buttons, cls = simple_text(base)
        bad = SIMPLE_FORBIDDEN.findall(text) + SIMPLE_FORBIDDEN_CASE.findall(text)
        assert not bad, (bad, text[:600])
        for w in want:
            assert w in text, (w, text[:600])
        assert sum(1 for b in buttons if b.get("data-opt")) >= opts, (opts, buttons)
        for w in once:                         # a question with buttons is shown once, not also as a numbered list
            assert text.count(w) == 1, (w, text.count(w), text[-500:])
        if site is not None:                   # phone: the Chat | Preview toggle only once a site exists
            assert ("has-site" in cls) == site, cls
        return text
    try:
        os.environ["J2_FIXES"] = "0"
        check(j2.J2Run(), "Describe the website you want", "A website for my cafe", "Understanding your request",
              site=False)
        run = j2.J2Run()
        say(run, "a website for my cafe")
        check(run, "Your website team", "What is the business called", site=False)          # J2.a questions
        say(run, "It's called Little Fig, in Koramangala, Bengaluru. Visitors should see our menu with prices, opening "
                 "hours and where to find us. Cozy and simple, and it must work on phones.")
        check(run, opts=2, once=("Do you want a section about your coffee's story?",
                                 "Add a section about your coffee's story (takes a bit longer)",
                                 "Keep it simple, no story section"))                       # a conflict: once, as buttons
        say(run, "Option 2")
        check(run, "Your website is ready", "Open full screen", "Download the HTML file", "Ask for a change",
              "Keep it simple, no story section", site=True)                                # the chosen label shows
        os.environ["J2_FIXES"] = "1"
        run = j2.J2Run()
        run.budget = 85
        say(run, BREW_LAB)
        say(run, "Option 1")                                                                 # ruling: add effort
        assert run.paused and run.paused["kind"] == "limit", run.outcome
        check(run, opts=2, once=("Add more and keep going, leave out the missing parts, or finish as it is?",
                                 "Finish as it is: you get the site as it is now"))
        # (c) the Simple view shows only messages for you: no agent-tagged message, no agent content
        text = check(run, opts=2)
        for m in run.chat:
            if m["who"] == "identity" and not m.get("to_user"):
                body = re.sub(r"^@[\w-]+[,:]?\s*", "", m["text"])[:40]
                assert body not in text, ("agent message in the Simple view", m["text"][:80])
        # (d) the Detailed view shows both parts of a reply, each labelled with who it is for
        import subprocess
        dom = chrome_dom([j2.j2_check.find_chrome(), "--headless=new", "--disable-gpu", "--virtual-time-budget=6000",
                              "--dump-dom", base + "/?view=detail"], capture_output=True, text=True, timeout=90).stdout
        assert "for Requester (you)" in dom and "for Adaptation" in dom, "Detailed view must label both audiences"
    finally:
        j2_server.RUN = keep
        srv.shutdown()
        os.environ["J2_FIXES"] = "0"


def work_atoms():
    """Adaptation writes each workflow's work atom (goal, constraints, context, how, verification). The validator
    rejects a workflow whose atom lacks a field, naming it, and Adaptation gets one retry; still incomplete after it,
    the workflow is dropped. Registered workflows store their atoms (W-n.a1) and each build call is tagged with one."""
    run = j2.J2Run()
    say(run, BREW_LAB)
    assert run.outcome.startswith("Done"), run.outcome
    assert run.workflows and not [x for x in run.rejections if "atom" in x["reason"]], run.rejections
    for w in run.workflows:
        assert [a["id"] for a in w["atoms"]] == [f"{w['id']}.a1"] and not j2.atom_missing(w), w
    saved = json.load(open(os.path.join(run.dir, "workflows.json"), encoding="utf-8"))
    assert all(x["atoms"] == w["atoms"] for x, w in zip(saved, run.workflows)), "atoms not in workflows.json"
    built = [c for c in run.call_log if j2.is_wf_agent(c["fn"])]
    assert built and all(c["atom"] == f"{c['fn']}.a1" for c in built), [(c["fn"], c.get("atom")) for c in built]
    assert not any("atom" in c for c in run.call_log if not j2.is_wf_agent(c["fn"]))
    assert "WORK ATOMS" in next(c["prompt"] for c in calls_of(run) if c["fn"] == "adaptation")

    def broken(field, times):
        run, orig, asked = j2.J2Run(), None, []
        orig = run.call

        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            if fn == "adaptation" and isinstance(r, dict) and r.get("workflows") and len(asked) < times:
                asked.append(task)
                a = r["workflows"][0]["atoms"][0]
                a[field] = [] if field == "verification" else a.pop(field) and None
            return r
        run.call = call
        say(run, BREW_LAB)
        return run, asked
    for f in j2.ATOM_FIELDS:
        run, asked = broken(f, 1)
        assert len(asked) == 1 and run.rejections[0]["reason"] == f"missing {f} in work atom 1", (f, run.rejections)
        retry = [c for c in calls_of(run) if c["fn"] == "adaptation" and "THE VALIDATOR REJECTED" in c["prompt"]]
        assert len(retry) == 1 and f in retry[0]["prompt"], f
        assert run.outcome.startswith("Done") and all(not j2.atom_missing(w) for w in run.workflows), f
    run, asked = broken("goal", 2)                      # still missing after the one retry: dropped, not retried again
    first = run.rejections[0]["workflow"]
    assert [x["workflow"] for x in run.rejections[:2]] == [first, first], run.rejections
    assert len([c for c in calls_of(run) if "THE VALIDATOR REJECTED" in c["prompt"]]) == 1
    # dropping the scaffold left N8 (the nav) with no workflow: the plan is flagged and re-asked, and the clean reply
    # brings it back
    assert any(x["workflow"] == "plan" and "N8" in x["reason"] for x in run.rejections), run.rejections
    assert first in [w["name"] for w in run.workflows if w["cycle"] == 1], first
    # the validator itself: no atoms at all, or an empty verification list, is rejected with the field named
    reg, refused, _ = run._register([{"name": "Bare", "goal_gap": "x", "outputs": ["y"], "completion_check": "z",
                                      "parts": ["fix"], "fixes": ["F-9"]}])
    assert not reg and refused == ["Bare (missing atoms)"], refused


def work_atom_cards():
    """Detailed view: opening a workflow lists its work atoms as cards (one atom opens expanded, two open collapsed),
    each with six plain rows and its own Run details listing only the calls tagged with that atom."""
    import html as htmllib
    import subprocess
    import threading
    import j2_server
    chrome = j2.j2_check.find_chrome()
    if not chrome:
        print("work atom cards test skipped: Chrome not found")
        return
    run = j2.J2Run()
    say(run, BREW_LAB)
    w1, w2 = run.workflows[0], run.workflows[1]
    w1["atoms"].append({**w1["atoms"][0], "id": "W-1.a2", "goal": "The second atom's goal holds"})
    other = next(c for c in run.call_log if c["fn"] == "priority")
    other["atom"] = "W-1.a2"
    mine = [c["id"] for c in run.call_log if c.get("atom") == "W-1.a1"]
    assert mine
    keep, j2_server.RUN = j2_server.RUN, run
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        dom = chrome_dom([chrome, "--headless=new", "--disable-gpu", "--virtual-time-budget=8000", "--dump-dom",
                              f"http://127.0.0.1:{srv.server_address[1]}/?view=detail"],
                             capture_output=True, text=True, timeout=120).stdout
    finally:
        j2_server.RUN = keep
        srv.shutdown()

    def block(key):              # the <details data-key=key> element's HTML (balanced on nested details)
        i = dom.index(f'data-key="{key}"')
        i, depth, j = dom.rindex("<details", 0, i), 0, i
        for m in re.finditer(r"<details\b|</details>", dom[i:]):
            depth += 1 if m.group(0) == "<details" else -1
            if not depth:
                return dom[i:i + m.end()]
    one, two = block(f"wf-{w2['id']}"), block("wf-W-1")
    assert two.count('class="atom"') == 2 and "Work atom 1" in two and "Work atom 2" in two, two[:300]
    assert one.count('class="atom"') == 1 and re.search(r'class="atom"[^>]*\bopen\b', one), "one atom opens expanded"
    assert not re.search(r'class="atom"[^>]*\bopen\b', two), "two atoms open collapsed"
    a1 = block("atom-W-1.a1")
    for key in ("goal", "how", "constraints", "context", "expected_output", "verification", "human_output", "agent_output"):
        assert f'data-block="{key}"' in a1, key
    assert re.search(r"Passed</span>: [^<]+ <span class=\"empty\">Attempt 1 of 3", a1) and "not run" not in a1 and "{" not in re.sub(r"<[^>]+>", "", a1.split("Run details")[0])
    run1, run2 = block("run-W-1.a1"), block("run-W-1.a2")
    labels = re.findall(r"<(?:h4|summary)>([^<]*)<", run1)
    assert labels == ["Run details", "Prompt sent", "Reply received", "Reason given", "Cost", "Attempts"], labels
    assert not any(n.lower() in x.lower() for x in labels for n in j2.ATOM_FIELDS), labels
    assert all(f"call {i}<" in run1 for i in mine) and f"call {other['id']}<" not in run1, run1[:400]
    assert f"call {other['id']}<" in run2 and not any(f"call {i}<" in run2 for i in mine), run2[:400]
    pane = dom.split('id="pane-workflows"')[1].split('id="pane-conflicts"')[0]
    pane = re.sub(r"<h4>Prompt sent</h4>.*?<h4>", "", pane, flags=re.S)       # the prompt is shown as it was sent
    pane = htmllib.unescape(re.sub(r"<[^>]+>", " ", pane))
    assert "Limits" in pane and not re.search(r"envelope|done.when", pane, re.I), "plain words in the Detailed view"


def atom_checks():
    """Verification is real: the validator's check rules, code running each atom's checks into receipts.jsonl, and a
    third failed attempt going to Adaptation before Identity (Identity only for a wrong build that doesn't fit)."""
    run = j2.J2Run()
    say(run, BREW_LAB)
    finish(run)
    src, a = run._fact_source(), j2.mock_atom({"covers": [1]})
    assert not j2.check_problems(a, src), j2.check_problems(a, src)
    # facts-present: every string must come from the charter or the Requester's answers (no invented facts)
    fact = {"text": "Name and area shown", "type": "facts-present", "target": ["brew  LAB", "Indiranagar"], "expect": True}
    assert not j2.check_problems({**a, "verification": [fact]}, src)
    bad = j2.check_problems({**a, "verification": [{**fact, "target": ["Brew Lab", "+91 98450 12345"]}]}, src)
    assert len(bad) == 1 and "['+91 98450 12345'], not found in the charter" in bad[0], bad
    # counts: label and expect must agree; targets only ids and classes the atom tells the builder to create
    how = {"kind": "AI", "agent": "Builder", "steps": "Write section #events with one .event-option per event"}
    opt = {"text": "8 event options", "type": "element-count", "target": ".event-option", "expect": 9}
    bad = j2.check_problems({**a, "how": how, "verification": [opt]}, src)
    assert bad == ["check 1 says '8 event options' but expects 9"], bad
    assert not j2.check_problems({**a, "how": how, "verification": [{**opt, "expect": 8}]}, src)
    # tap targets: the screen widths 360/1280 in a label are never the count or the size (runs/j2-20261026-174314)
    tap = {"text": "Controls are tappable at 360px", "type": "tap-targets", "target": "#events", "expect": 44}
    assert not j2.check_problems({**a, "how": how, "verification": [tap]}, src)
    assert not j2.check_problems({**a, "how": how, "verification": [{**tap, "text": "44px targets at 360px and 1280px"}]}, src)
    assert j2.check_problems({**a, "how": how, "verification": [{**tap, "text": "48px targets at 360px"}]}, src) == [
        "check 1 says '48px targets at 360px' but expects 44"], "a real size still has to match"
    for t in ("Day 1", "hello@brewlab.in", ".untold-class"):
        bad = j2.check_problems({**a, "how": how, "verification": [{**opt, "text": "event options", "target": t}]}, src)
        assert len(bad) == 1 and "targets" in bad[0], (t, bad)
    assert j2.check_problems({**a, "verification": [{**opt, "type": "screenshot-diff"}]}, src)[0].startswith(
        "check 1 has type 'screenshot-diff'")
    # expected output: every item is covered by at least one check; an item no check can verify is rejected, and a
    # derived check follows the same rules (facts from the charter, counts only on ids the atom creates)
    cov = j2.normalise_checks([json.loads(json.dumps(j2.mock_atom({"parts": ["form"]})))])[0]
    assert all(any(c.get("item") == n for c in cov["verification"]) for n in range(1, len(cov["expected_output"]) + 1))
    assert len(cov["verification"]) == 3, cov["verification"]          # matched by type and target, not duplicated
    item = lambda **x: run._atom_problems({"atoms": [{**a, "expected_output": [{"text": "x", **x}]}]})
    assert "has kind 'heading', which no check can verify" in item(kind="heading", target="Menu")[0]
    assert "expected output item 1 facts-present lists ['+91 98450 12345']" in item(kind="facts", target=["+91 98450 12345"])[0]
    assert "expected output item 1 element-count targets '#events'" in item(kind="section", target="#events")[0]
    assert not item(kind="facts", target=["Indiranagar"]) and run._atom_problems({"atoms": [{**a, "expected_output": []}]})
    # a goal naming times whose checks only count headings is flagged, and sent back to Adaptation once
    timed = {**a, "goal": "Opening hours 8am to 9pm are listed", "how": {**how, "steps": "h2 headings, class hours-head"},
             "expected_output": [{"text": "2 hours headings", "kind": "element", "target": ".hours-head", "count": 2}],
             "verification": [{"text": "2 hours headings", "type": "element-count", "target": ".hours-head", "expect": 2}]}
    assert j2.check_problems(timed, src) == ["its goal names a time, price, number or email but no check is "
                                             "facts-present or a named test"], j2.check_problems(timed, src)

    def patched(edit, reply=lambda r, fn: r, done=True):
        """A run whose first Adaptation proposal has its first atom changed by edit (once)."""
        run, hit = j2.J2Run(), []
        orig = run.call

        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            if fn == "adaptation" and isinstance(r, dict) and r.get("workflows") and not run.workflows and not hit:
                hit.append(1)
                edit(r["workflows"][0]["atoms"][0])
            return reply(r, fn)
        run.call = call
        say(run, BREW_LAB)
        if done:
            finish(run)
        return run
    run = patched(lambda at: at.update({k: v for k, v in timed.items() if k != "context"}))
    assert "names a time, price, number or email" in run.rejections[0]["reason"], run.rejections
    assert len([c for c in calls_of(run) if "THE VALIDATOR REJECTED" in c["prompt"]]) == 1
    assert run.outcome.startswith("Done"), run.outcome

    # the same failure on two attempts (identical evidence): no third attempt, the receipts go to Adaptation; it corrects the check and Identity is never called
    def four_cards(at):
        at["how"] = {**at["how"], "steps": "Give every card the class j2-card."}
        at["verification"].append({"text": "4 cards", "type": "element-count", "target": ".j2-card", "expect": 4})
    run = patched(four_cards)
    calls = calls_of(run)
    i = next(n for n, c in enumerate(calls) if "WORK ATOM FAILED" in c["prompt"])
    assert calls[i]["fn"] == "adaptation" and [c["fn"] for c in calls[i - 2:i]] == ["W-1"] * 2, [c["fn"] for c in calls]
    assert "ATTEMPT 2 OF 3" in calls[i - 1]["prompt"] and "0 element(s) match '.j2-card'" in calls[i - 1]["prompt"]
    nxt = next(n for n, c in enumerate(calls) if n > i and c["fn"] != "adaptation")
    assert calls[nxt]["fn"] != "identity" and not any(c["fn"] == "identity" for c in calls[i:nxt]), calls[nxt]["fn"]
    at = run._wf("W-1")["atoms"][0]
    card = [r for r in run.receipts if r["target"] == ".j2-card"]
    assert [(r["attempt"], r["expect"], r["actual"], r["pass"]) for r in card] == \
        [(1, 4, 0, False), (2, 4, 0, False), (3, 0, 0, True)], card
    assert card[-1]["recheck"] == "check corrected by Adaptation" and at["status"] == "passed", at
    assert run.outcome.startswith("Done"), run.outcome
    # receipts: written by code only, one per check run, matching each atom's checks
    lines = [json.loads(x) for x in open(os.path.join(run.dir, "receipts.jsonl"), encoding="utf-8")]
    met = [r for r in lines if "expected" in r]          # per expected-output item: met or not met
    lines = [r for r in lines if "expected" not in r]
    assert met and all(r["by"] == "code" and r["met"] in (True, False, None) and r["checks"] for r in met), met[:3]
    assert lines == run.receipts and all(r["by"] == "code" and r["of"] == 3 and r["t"] for r in lines)
    for w in run.workflows:
        for x in w["atoms"]:
            ran = [(r["check"], r["type"], r["target"], r["expect"]) for r in lines
                   if r["atom"] == x["id"] and r["attempt"] == x["attempts"]][-len(x["verification"]):]
            assert ran == [(c["text"], c["type"], c["target"], c["expect"]) for c in x["verification"]], (x["id"], ran)
            assert x["results"] == [r for r in lines if r["atom"] == x["id"]][-len(x["verification"]):]
    assert not any("receipts.jsonl" in str(c.get("raw")) for c in calls), "an agent wrote a receipt"
    v = json.load(open(os.path.join(run.dir, "summary.json"), encoding="utf-8"))["verification"]
    assert v == {"passed": len(run.workflows), "failed": 0, "escalated": 0, "could_not_run": 0,
                 "checks_run": len(lines)}, v

    # a wrong build the budget left can't fix: Adaptation first, then Identity; the atom is ESCALATED
    def missing_fact(at):
        at["verification"].append({**fact, "target": ["Indiranagar"]})

    def unaffordable(r, fn):
        if fn == "adaptation" and isinstance(r, dict) and r.get("verdict") == "BUILD_WRONG":
            r["affordable"] = False
        return r
    run = patched(missing_fact, unaffordable, done=False)
    calls = calls_of(run)
    i = next(n for n, c in enumerate(calls) if "WORK ATOM FAILED" in c["prompt"])
    assert calls[i]["fn"] == "adaptation" and not any(c["fn"] == "identity" for c in calls[i - 3:i])
    assert calls[i + 1]["fn"] == "identity" and "fails its checks" in calls[i + 1]["prompt"], calls[i + 1]["fn"]
    assert run._wf("W-1")["atoms"][0]["status"] == "escalated"
    finish(run)
    v = json.load(open(os.path.join(run.dir, "summary.json"), encoding="utf-8"))["verification"]
    assert v["escalated"] == 1, v


def leave_out_named():
    """runs/j2-20261012-170632 msgs 58-63: "Leave out the 2 unfinished checklist parts" named nothing; Identity then
    told the team fees and sponsors, while code dropped sponsors and the colour scheme. Now the question lists every
    unfinished item with its run-long number, short text and status, each option says what it drops and keeps, the
    drop is exactly what the option named, and numbers never shift after a drop."""
    six = BREW_LAB.replace(", and a simple loyalty-card explainer (buy 9, get the 10th free)", "")
    run = j2.J2Run()
    say(run, six)
    finish(run)
    crit = list(run.goal["done_when"])
    assert len(crit) == 6, crit
    hours, beans = (next(i for i, c in enumerate(crit, 1) if k in c.lower()) for k in ("hours", "beans"))
    last = run.audits[-1]
    for r in last["results"]:
        if r["item"] in (hours, beans):
            r.update(verdict="NEEDS_REVISION", note="missing", **{"pass": False})
    last["rubric"] = [{"id": "R-6", "criterion": j2.RUBRIC[5][1], "verdict": "NEEDS_REVISION", "pass": False,
                       "note": "button 30px", "rubric": True}]
    run.outcome, run.done = None, False
    run._identity_limit({"reason": "The same failure came back twice."})
    q = run.paused["questions"][0]
    text = run.chat[-1]["text"]
    for i in (hours, beans):
        assert re.search(rf"\n- item {i} {re.escape(j2.short(crit[i - 1]))} \((built but failing|not built|unfunded)\)",
                         text), text
    assert "\n- Tap targets are big enough on phones (failing)" in text, text
    drop = next(o for o in q["options"] if o.startswith("Leave out"))
    named = [f"item {hours} {j2.short(crit[hours - 1])}", f"item {beans} {j2.short(crit[beans - 1])}"]
    assert drop == (f"Leave out {named[0]} and {named[1]}: you get the site without them. Tap targets are big "
                    "enough on phones stays as it is, still failing."), drop
    ship = q["options"][-1]
    assert ship.endswith(f"still not right: {named[0]}, {named[1]} and Tap targets are big enough on phones"), ship
    assert run.paused["handoff"]["drop_named"] == named
    say(run, str(q["options"].index(drop) + 1))
    d = next(x for x in reversed(run.decisions) if x.get("decision") == "drop")
    assert d["named"] == d["dropped"] == named, d                          # the dropped set is the named set
    assert [c for c in crit if c not in run.goal["done_when"]] == [crit[hours - 1], crit[beans - 1]]
    ack = next(m for m in run.chat if m.get("to_user") and "Got it. We'll leave out" in m["text"])
    assert named[0] in ack["text"] and named[1] in ack["text"], ack["text"]
    # numbers don't shift: the items after the dropped ones keep theirs, in the next audit and the next question
    for i, c in enumerate(crit, 1):
        if c in run.goal["done_when"]:
            assert run._no(c) == i and next(r["id"] for r in run.audits[-1]["results"] if r["criterion"] == c) == f"F-{i}"
    late = run.goal["done_when"][-1]
    assert run._gap_rows([late], run.audits[-1], None)[0]["name"].startswith(f"item {crit.index(late) + 1} ")
    # a typed answer names items by those numbers too
    assert j2.parse_decision("leave out F-6", {}, ["drop"], run.goal["done_when"],
                             [run._no(c) for c in run.goal["done_when"]])["drop_items"] == [6]


def rubric_and_credits():
    """Every builder's brief carries the design rubric R-1 to R-6; the section holding buttons has a tap-targets
    check code runs at 360px (it fails on a 30px button); no user-facing string says "units"."""
    run = j2.J2Run()
    say(run, BREW_LAB)
    finish(run)
    built = [c for c in calls_of(run) if j2.is_wf_agent(c["fn"])]
    assert built and all(f"{rid} {txt}" in c["prompt"] for c in built for rid, txt in j2.RUBRIC[:6])
    form = next(i for i, c in enumerate(run.goal["done_when"], 1) if j2.item_kind(c) == "form")
    w = next(w for w in run.workflows if form in run._wf_items(w))
    tap = [c for a in w["atoms"] for c in a["verification"] if c["type"] == "tap-targets"]
    assert tap and tap[0]["target"] == ".order-form" and "order-form" in json.dumps(w["atoms"][0]["how"]), w["atoms"]
    assert 'class="order-form"' in run.sites[-1] and "<button" in run.sites[-1].split('class="order-form"')[1]
    got = [r for r in run.receipts if r["type"] == "tap-targets"]
    assert got and got[-1]["pass"] and "at 360px" in got[-1]["evidence"], got
    page = ('<!doctype html><html><body><section class="order-form"><form><button type="submit" '
            'style="height:30px;width:120px">Send</button></form></section></body></html>')
    r = j2.run_check(tap[0], run.dir, page)
    assert r["pass"] is False and "1 of 1 tap target(s)" in r["evidence"] and "120x30" in r["evidence"], r
    # "units": never in the outcome, the report, summary.json or any chat message; the page uses it only to parse
    run = j2.J2Run()
    run.budget = 12
    say(run, BREW_LAB)
    say(run, "finish as it is")
    assert run.outcome.startswith("Done with open items"), run.outcome
    texts = [run.outcome, open(os.path.join(run.dir, "report.md"), encoding="utf-8").read(),
             open(os.path.join(run.dir, "summary.json"), encoding="utf-8").read()] + [m["text"] for m in run.chat]
    assert "work credits" in " ".join(texts)
    bad = [t[:200] for t in texts if re.search(r"\bunits?\b", t)]
    assert not bad, bad
    ui = [ln for ln in open(os.path.join(j2.HERE, "j2_ui.html"), encoding="utf-8") if re.search(r"\bunits\b", ln)]
    assert all(re.search(r"/[^/]*\bunits\b[^/]*/|\(\?:units", ln) for ln in ui), ui     # regex patterns only


def charter_numbers():
    """Agents' prompts and covers use each item's charter number, never its position: after item 2 is dropped, an
    agent's covers [5] still means charter item 5, and every number an agent reads is a charter number."""
    run = j2.J2Run()
    say(run, BREW_LAB)
    finish(run)
    crit = list(run.goal["done_when"])
    run._override("drop 2")
    assert crit[1] not in run.goal["done_when"] and run.goal["done_when"][1] == crit[2]   # position 2 is now item 3
    run.budget, run.p_refused = 500, {}
    reg, refused, _ = run._register([{"name": "Beans again", "goal_gap": "x", "outputs": ["x"], "completion_check": "y",
                                      "atoms": ATOM, "parts": ["page_section"], "covers": [5]}])
    assert reg and not refused, refused
    assert reg[0]["covers"] == [5] and reg[0]["covers_text"] == [crit[4]] and run._wf_items(reg[0]) == [4], reg[0]
    assert run._covers({"covers": [2]}) == [], "a dropped charter number covers nothing"
    nums = [1, 3, 4, 5, 6, 7]
    st = json.loads(run.state("adaptation"))
    assert [x["item"] for x in st["charter"]["done_when"]] == nums and "charter_items" not in st["charter"], st["charter"]
    assert [x["item"] for x in st["rate_card"]["item_minimums"]] == nums
    assert run._per_item([{"name": "x", "parts": ["page_section"], "covers": [5]}])[0]["covers"] == [5]
    run._turn_audit({"kind": "audit"})
    prompt = calls_of(run)[-1]["prompt"]
    assert '{"item": 5, "criterion": ' + json.dumps(crit[4]) in prompt and '"item": 2,' not in prompt
    assert [r["id"] for r in run.audits[-1]["results"]] == [f"F-{n}" for n in nums], run.audits[-1]["results"]
    fails = [{**r, "item": run._no(r["criterion"])} for r in run.audits[-1]["results"]]
    assert all(f["item"] == int(f["id"][2:]) for f in fails)
    # code's own proposal (cheapest fit) writes charter numbers too
    run.workflows = []
    fit = run._cheapest_fit()
    assert {i for w in fit["workflows"] for i in w["covers"]} | {x["item"] for x in fit["left_out"]} <= set(nums), fit


def chrome_unavailable():
    """No Chrome: a browser check records "could not run"; it isn't a failed attempt and nothing escalates."""
    old_env, old_tap, old_check = os.environ.get("J2_CHROME"), j2.j2_check.tap_targets, j2.j2_check.check
    os.environ["J2_CHROME"], j2.j2_check.tap_targets = "none", REAL_TAP
    j2.j2_check.check = fake_check      # Audit's render check isn't under test here (without Chrome it pauses the run)
    try:
        run = j2.J2Run()
        say(run, BREW_LAB)
        finish(run)
    finally:
        j2.j2_check.tap_targets, j2.j2_check.check = old_tap, old_check
        if old_env is None:
            os.environ.pop("J2_CHROME", None)
        else:
            os.environ["J2_CHROME"] = old_env
    tap = [r for r in run.receipts if r["type"] == "tap-targets"]
    assert tap and all(r["pass"] is None and r["evidence"].startswith("could not run") for r in tap), tap
    assert all(r["attempt"] == 1 for r in tap), "a check that could not run isn't retried"
    a = next(a for w in run.workflows for a in w["atoms"] if any(c["type"] == "tap-targets" for c in a["verification"]))
    assert a["status"] == "could not run" and a["attempts"] == 1, a
    assert not any("WORK ATOM FAILED" in c["prompt"] for c in calls_of(run)), "it went to Adaptation"
    assert not any(e["type"] == "atom_escalated" for e in run.events)
    v = json.load(open(os.path.join(run.dir, "summary.json"), encoding="utf-8"))["verification"]
    assert v["could_not_run"] == 1 and v["failed"] == v["escalated"] == 0, v
    assert run.outcome.startswith("Done"), run.outcome


def unique_short_names():
    """Short item names are unique within a list: items sharing their first 7 words get longer names until they
    differ (past the first clause if need be), in the leave-out question too."""
    a = "Each event card shows its fee and the room for events in Hall A"
    b = "Each event card shows its fee and the room for events in Hall B"
    names = j2.shorts([a, b, "Opening hours are listed"])
    assert names == [a, b, "Opening hours are listed"], names
    g = ["A gallery: photos of the shop, outside", "A gallery: photos of the team"]
    assert len(set(j2.shorts(g))) == 2 and all(n.startswith("A gallery: photos of the") for n in j2.shorts(g))
    assert j2.shorts(["The menu with prices is shown", "Opening hours are listed"]) == \
        ["The menu with prices is shown", "Opening hours are listed"], "names that already differ stay short"
    run = j2.J2Run()
    say(run, BREW_LAB)
    finish(run)
    run.goal["done_when"] += [a, b]
    gaps = [a, b, run.goal["done_when"][0]]
    rows = run._gap_rows(gaps, run.audits[-1], None)
    shown = [r["name"].split(" ", 1)[1] for r in rows]
    assert len(set(shown)) == 3 and shown[0].endswith("Hall A") and shown[1].endswith("Hall B"), shown


def detail_stays_open():
    """The Detailed view re-renders on every poll. An opened screenshot (under Audit), the report's event log, and an
    opened message detail view (with its full text) must stay open while the run changes, until the user closes them.
    Real browser: the page opens them, then the server changes the run (a new audit and a new message) mid-test."""
    import html as htmllib
    import shutil
    import subprocess
    import threading
    import j2_server
    chrome = j2.j2_check.find_chrome()
    if not chrome:
        print("detail view test skipped: Chrome not found")
        return
    run = j2.J2Run()
    say(run, BREW_LAB)
    add_budget_if_asked(run)
    v = run.audits[-1]["version"]
    os.makedirs(os.path.join(run.dir, "site", "shots"), exist_ok=True)
    shutil.copy(os.path.join(j2.HERE, "docs", "ui-simple-start-390.png"), os.path.join(run.dir, "site", "shots", f"v{v}.png"))
    run.audits[-1]["screenshot"] = f"site/shots/v{v}.png"
    msg = next(m["id"] for m in run.chat if m.get("calls"))
    # found by their visible text, so the same test also runs against an older page
    hook = ("const pick = (pane, t) => [...document.querySelectorAll(pane + ' details')].find(d => d.textContent.includes(t));\n"
            "setTimeout(async () => { const q = s => document.querySelector(s);\n"
            "  pick('#pane-audit', 'Screenshot').open = true; pick('#pane-report', 'Event log').open = true;\n"
            f"  await openDetail({msg}); q('#full').open = true;\n"
            "  await fetch('/test/bump');   // the server changes the run: a new audit and a new message\n"
            "  setTimeout(() => { document.body.dataset.r = JSON.stringify({shot: !!pick('#pane-audit', 'Screenshot')?.open,"
            " events: !!pick('#pane-report', 'Event log')?.open, dialog: q('#detail').open, full: q('#full').open,"
            " audits: q('#pane-audit').innerText.split('Cycle ').length - 1}); }, 4000); }, 1500);")
    ui = open(os.path.join(j2.HERE, "j2_ui.html"), encoding="utf-8").read().replace("poll();\n</script>", f"poll();\n{hook}\n</script>")
    orig = j2_server.Handler.do_GET

    def do_GET(self):
        if self.path.startswith("/?"):
            return self._send(200, ui, "text/html")
        if self.path == "/test/bump":
            run.audits.append({**run.audits[-1], "version": v, "cycle": run.audits[-1]["cycle"] + 1})
            run.post("system", "A new message while the detail view is open.")
            return self._send(200, "{}")
        return orig(self)
    j2_server.Handler.do_GET = do_GET
    keep, j2_server.RUN = j2_server.RUN, run
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        before = len(run.audits)
        out = chrome_dom([chrome, "--headless=new", "--disable-gpu", "--virtual-time-budget=12000", "--dump-dom",
                              f"http://127.0.0.1:{srv.server_address[1]}/?view=detail"],
                             capture_output=True, text=True, timeout=120).stdout
        m = re.search(r'data-r="([^"]*)"', out)
        assert m, "the page never reported"
        r = json.loads(htmllib.unescape(m.group(1)))
        assert len(run.audits) == before + 1 and r["audits"] == before + 1, (r, "the Audit pane didn't re-render")
        assert r == {**r, "shot": True, "events": True, "dialog": True, "full": True}, r
    finally:
        j2_server.Handler.do_GET = orig
        j2_server.RUN = keep
        srv.shutdown()


def plain_words():
    """The Simple view's word list (fallback only): Audit becomes "our checks", never "the team", so nothing reads
    "after the team"; other function names become "the team"; budget becomes effort. Runs the page's own JS in node."""
    import shutil
    import subprocess
    if not shutil.which("node"):
        print("plain-words check skipped: node not found")
        return
    with open(os.path.join(j2.HERE, "j2_ui.html"), encoding="utf-8") as f:
        ui = f.read()
    js = ui[ui.index("const AGENT ="):ui.index("const STAGES =")]
    cases = {"only 10 units stay for fixes after Audit": "Only 10 work credits stay for fixes after our checks",
             "Every done-when item passes with Audit's evidence.": "Every checklist item passes with our checks' evidence.",
             "@Adaptation Priority's final answer was COUNTER.": "The team's final answer was a smaller offer.",
             "The Requester picked; the budget is over budget.": "You picked; the effort is more than planned."}
    out = subprocess.run(["node", "-e", js + "\nconst c = " + json.dumps(list(cases)) +
                          ";\nprocess.stdout.write(JSON.stringify(c.map(plain)));"], capture_output=True, text=True)
    got = json.loads(out.stdout or "null") or sys.exit(out.stderr)
    for (src, want), g in zip(cases.items(), got):
        assert g == want, (src, g, want)
        assert not re.search(r"\b(?:after|before|by|from|until|than) the team\b", g), g


def identity_prompt_is_plain():
    """Item 1: every Identity call carries the plain-language rule for what it writes to the Requester; no one else's."""
    run = j2.J2Run()
    say(run, BREW_LAB)
    add_budget_if_asked(run)
    with open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8") as f:
        calls = [json.loads(ln) for ln in f]
    assert any(c["fn"] == "identity" for c in calls)
    for c in calls:
        assert ("college student with no tech background" in c["prompt"]) == (c["fn"] == "identity"), (c["id"], c["fn"])


FLORIST = "A website for Petal & Stem, a florist in Jaipur"


def calls_of(run):
    with open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8") as f:
        return [json.loads(ln) for ln in f]


def style_guide():
    """H: the first workflow that runs writes the style guide (retrying within its envelope if it leaves it out);
    code puts it first in the head; every later workflow agent gets it and is told to follow it."""
    run = j2.J2Run()
    orig, dropped = run.call, []

    def call(fn, fmt, task, mock, **k):              # W-1's first reply leaves the style guide out: it retries
        raw = orig(fn, fmt, task, mock, **k)
        if fn == "W-1" and not dropped:
            dropped.append(1)
            raw = re.sub(r"```style.*?```", "", raw, flags=re.S)
        return raw
    run.call = call
    say(run, BREW_LAB)
    add_budget_if_asked(run)
    w1 = run._wf("W-1")
    assert run.style and run.style["palette"]["primary"] and w1["retries_used"] == 1 and w1.get("style_guide"), w1
    assert list(run.sections)[0] == "style" and run.sections["style"]["owner"] == "W-1"
    site = run.sites[-1]
    assert "--j2-primary" in site and site.index("<!-- J2:style head start") < site.index("</head>")
    assert os.path.exists(os.path.join(run.dir, "style.json")) and run.snapshot()["style"] == run.style
    cs = calls_of(run)
    w1_calls = [c for c in cs if c["fn"] == "W-1"]
    assert all("YOU RUN FIRST, SO YOU ALSO WRITE THE STYLE GUIDE" in c["prompt"] for c in w1_calls), "W-1 not told"
    later = [c for c in cs if j2.is_wf_agent(c["fn"]) and c["fn"] != "W-1"]
    assert later and all("STYLE GUIDE (follow it exactly)" in c["prompt"] for c in later), [c["fn"] for c in later]
    assert any("the reply has no ```style block" in m["text"] for m in run.chat), "retry reason not shown"


def rubric_loop():
    """H: W-2's first build has a tap target too small for a phone. Code measures it, Audit's R-6 can't pass, the
    site isn't workable, Identity sends it back, Adaptation proposes a fix, the re-check passes, Done."""
    if not j2.j2_check.find_chrome():
        print("rubric loop skipped: Chrome not found")
        return
    run = j2.J2Run()
    run.budget = 150      # room for the extra fix in both modes (J2_FIXES=1 also funds the stories)
    orig, planted = run.call, []

    def call(fn, fmt, task, mock, **k):
        raw = orig(fn, fmt, task, mock, **k)
        if fn == "W-2" and not planted:
            planted.append(1)
            raw = raw.replace("```section\n", '```section\n<a href="https://example.com" style="display:inline-block;'
                                              'font-size:10px">Book</a>\n', 1)
        return raw
    run.call = call
    say(run, BREW_LAB)
    add_budget_if_asked(run)
    first, last = run.audits[0], run.audits[-1]
    assert [r["id"] for r in first["rubric"]] == [rid for rid, _ in j2.RUBRIC]
    r6 = next(r for r in first["rubric"] if r["id"] == "R-6")
    assert r6["verdict"] == "NEEDS_REVISION" and r6["note"].startswith("Measured by code") and "Book" in r6["note"], r6
    assert any(x.startswith("design rubric R-6") for x in j2.workable(first))
    fix = next(w for w in run.workflows if w["name"] == "Fix item R-6")
    assert fix["status"] == "done" and fix["revises"] == "W-2", fix
    assert next(r for r in last["rubric"] if r["id"] == "R-6")["pass"] and run.outcome.startswith("Done"), run.outcome
    assert any("DESIGN AND USABILITY RUBRIC" in c["prompt"] for c in calls_of(run) if c["fn"] == "audit")


def lone_rubric_failure():
    """Bug from runs/j2-20261007-184146: the first audit's only failure was R-6 (tap targets), with budget left in
    cycle 1, and the run paused with a limit question. A normal failure with budget left must loop back to J2.b
    without asking the Requester: Adaptation proposes the fix, it runs, Audit re-checks, Done, no limit hit."""
    if not j2.j2_check.find_chrome():
        print("lone rubric failure test skipped: Chrome not found")
        return
    os.environ["J2_FIXES"], mode = "1", os.environ.get("J2_FIXES", "0")    # stories funded: the beans item is paid
    run = j2.J2Run()
    os.environ["J2_FIXES"] = mode
    run.budget = 150
    orig, done = run.call, set()

    def call(fn, fmt, task, mock, **k):
        raw = orig(fn, fmt, task, mock, **k)
        if fn == "adaptation" and isinstance(raw, dict) and raw.get("workflows") and not run.workflows:
            raw["workflows"][0]["covers"] = raw["workflows"][0].get("covers", []) + [7]   # the scaffold pays for item 7
        if fn == "W-1" and fn not in done:          # the scaffold already has the viewport tag: phones pass
            done.add(fn)
            raw = raw.replace("```head\n", '```head\n<meta name="viewport" content="width=device-width, initial-scale=1">\n', 1)
        if fn == "W-2" and fn not in done:          # and one tap target is too small: R-6 is the only failure
            done.add(fn)
            raw = raw.replace("```section\n", '```section\n<a href="https://example.com" style="display:inline-block;'
                                              'font-size:10px">Call us</a>\n', 1)
        return raw
    run.call = call
    say(run, BREW_LAB)
    first = run.audits[0]
    fails = j2.workable(first)
    assert len(fails) == 1 and fails[0].startswith("design rubric R-6"), fails   # the bug's exact situation
    assert not any(e["type"] == "limit" for e in run.events), [e["detail"] for e in run.events if e["type"] == "limit"]
    assert not any(m.get("pause") == "limit" for m in run.chat), "asked the Requester a limit question"
    assert any(w["name"] == "Fix item R-6" and w["status"] == "done" for w in run.workflows), run.workflows
    assert run.outcome.startswith("Done") and summary(run)["limits_hit"] == [], run.outcome


def pricing():
    """I: the validator ties price to the work, not to labels."""
    run = j2.J2Run()
    run.goal = j2.mock_charter(BREW_LAB)    # 1 menu, 2 hours, 3 address, 4 order (form), 5 beans, 6 loyalty, 7 phone
    run.budget, run.cycle = 500, 1
    base = {"goal_gap": "x", "outputs": ["index.html"], "completion_check": "y", "atoms": ATOM}
    kinds = [j2.item_kind(c) for c in run.goal["done_when"]]
    assert kinds == ["page_section"] * 3 + ["form"] + ["page_section"] * 2 + ["styling"], kinds
    # bundling: six items in one 20-unit workflow (as in the run) is rejected; the same work priced honestly passes
    reg, refused, _ = run._register([{**base, "name": "Shop", "parts": ["page_section", "interactive_feature"],
                                      "covers": [1, 2, 3, 4, 5, 6]}])
    assert not reg and "costs 20, below the rate-card minimum of 65" in refused[0], refused
    assert "item 4 form 15" in refused[0], refused[0]
    reg, refused, _ = run._register([{**base, "name": "Shop", "covers": [1, 2, 3, 4, 5, 6],
                                      "parts": ["page_section"] * 5 + ["form"]}])
    assert reg and reg[0]["cost"] == 65 and reg[0]["covers"] == [1, 2, 3, 4, 5, 6], refused
    # page_section/form parts must say which items they build
    reg, refused, _ = run._register([{**base, "name": "Page frame and store info", "parts": ["page_section", "styling"]}])
    assert not reg and "covers no done-when item" in refused[0], refused
    # revising a section whose workflow hasn't run yet is a new section, not a fix (runs/j2-20261009-133203 msg 27
    # rejected "Browse by category" for revising W-2 before W-2 had built anything)
    assert run._wf("W-1")["status"] == "registered"
    reg, refused, _ = run._register([{**base, "name": "Browse by category", "parts": ["page_section"], "covers": [5],
                                      "revises": "W-1"}])
    assert reg and reg[0]["cost"] == 10 and not refused, refused
    # after the first Audit: a fix relabelled as a cheaper part is rejected, with or without listing the finding
    run._wf("W-1")["status"] = "done"                        # W-1 ran: rewriting its section is a fix now
    run.audits = [{"cycle": 1, "version": 1, "results": [], "rubric": [], "mechanical": {}, "all_pass": False}]
    for w in ({"fixes": ["R-6"], "revises": "W-1"}, {"revises": "W-1"}, {"fixes": ["R-6", "F-2"]}):
        reg, refused, _ = run._register([{**base, "name": "Header tap area", "parts": ["styling"], **w}])
        assert not reg and "a fix is priced as a fix" in refused[0], (w, refused)
    reg, refused, _ = run._register([{**base, "name": "Header tap area", "parts": ["fix", "fix"], "fixes": ["R-6", "F-2"]}])
    assert reg and reg[0]["cost"] == 16, refused
    # a fix to an item already paid for isn't charged the item again; styling with no item or finding is rejected now
    reg, refused, _ = run._register([{**base, "name": "Hours fix", "parts": ["fix"], "fixes": ["F-2"], "covers": [2]}])
    assert reg and reg[0]["cost"] == 8, refused
    reg, refused, _ = run._register([{**base, "name": "Polish", "parts": ["styling"]}])
    assert not reg and "covers no done-when item" in refused[0], refused
    # Priority is told the cost per done-when item and to question plans far below the rate card
    run = j2.J2Run()
    say(run, BREW_LAB)
    finish(run)
    p = next(c for c in calls_of(run) if c["fn"] == "priority" and "ON THE TABLE" in c["prompt"])
    assert "COST PER DONE-WHEN ITEM" in p["prompt"] and "rate_card_minimum_for_items" in p["prompt"], p["prompt"][-800:]
    assert "far below the rate-card minimums" in p["prompt"]
    assert all(w.get("covers") or w.get("fixes") or w["name"] in ("Stories section", "Animations") for w in run.workflows)


def unpaid_items():
    """Cost follows the work, on the page: a done-when item that passes although no workflow paid for it is a normal
    finding (UNPAID). Baseline Brew Lab: the stories workflow (covering the beans item) is refused, yet the mock
    builds that item anyway. Audit flags it, Adaptation registers a workflow that covers it, and the run ends Done."""
    os.environ["J2_FIXES"], mode = "0", os.environ.get("J2_FIXES", "0")
    try:
        run = j2.J2Run()
        say(run, BREW_LAB)
        finish(run)
        beans = next(i for i, c in enumerate(run.goal["done_when"], 1) if "bean" in c.lower())
        first = next(r for r in run.audits[0]["results"] if r["item"] == beans)
        assert first["verdict"] == "UNPAID" and not first["pass"] and "no workflow covers" in first["note"], first
        assert any(x.startswith(f"done-when item {beans}") and "UNPAID" in x for x in j2.workable(run.audits[0]))
        pay = next(w for w in run.workflows if w.get("covers") == [beans] and w["status"] == "done")
        assert pay["cost"] >= j2.item_min(run.goal["done_when"][beans - 1]), pay
        assert beans in run.paid_items() and run.outcome.startswith("Done"), run.outcome
        assert not any(r["verdict"] == "UNPAID" for r in run.audits[-1]["results"])
        # unit: an item a fix paid for (via its finding) isn't flagged; an undeclared one is
        run = j2.J2Run()
        run.goal = {"done_when": ["A", "B"]}
        run.audits = [{"results": [{"id": "F-2", "item": 2}]}]
        run.workflows = [{"status": "done", "covers": [1]}, {"status": "done", "fixes": ["F-2"]}]
        rows = [{"item": 1, "pass": True, "criterion": "A"}, {"item": 2, "pass": True, "criterion": "B"}]
        run._flag_unpaid(rows)
        assert all(r["pass"] for r in rows), rows
        run.workflows = run.workflows[:1]
        run._flag_unpaid(rows)
        assert rows[0]["pass"] and rows[1]["verdict"] == "UNPAID", rows
    finally:
        os.environ["J2_FIXES"] = mode


def playbook_matching():
    """Playbooks: the same business type only; mock playbooks never reach a live run; tests write to a temp folder."""
    assert j2.PLAYBOOK_FILE.startswith(SESSION_MEMORY) and j2.MEMORY_FILE.startswith(SESSION_MEMORY), j2.PLAYBOOK_FILE
    cafe = {"id": "P-cafe", "business_type": "cafe", "brief": "A website for a cafe shop in Pune with a menu"}
    boutique = {"id": "P-shop", "business_type": "Clothing  Boutique", "brief": "A clothing boutique in Pune"}
    pbs = [cafe, boutique]
    # a clothing brief that shares words with the cafe brief still doesn't get the cafe playbook
    assert [x["id"] for x in j2.relevant_playbooks(pbs, "clothing boutique", "A website for a shop in Pune")] == ["P-shop"]
    assert j2.relevant_playbooks(pbs, "florist", cafe["brief"]) == []
    assert j2.relevant_playbooks(pbs, "", cafe["brief"]) == [] and j2.relevant_playbooks(pbs, "business", "x") == []
    # a playbook written by a mock run is never offered to a live run
    mock_pb = {**cafe, "id": "P-mock", "mock": True}
    live = j2.llm.MOCK
    try:
        j2.llm.MOCK = False
        assert j2.relevant_playbooks([mock_pb, cafe], "cafe", "cafe") == [cafe]
        assert j2.relevant_lessons([{"id": "L-m", "mock": True, "problem": "cafe menu", "fix": "x"}], "cafe menu") == []
    finally:
        j2.llm.MOCK = live
    # what mock runs save is marked mock
    os.environ["J2_MEMORY"] = "1"
    try:
        j2.save_playbook({"id": "P-t", "business_type": "cafe"})
        assert j2.load_playbooks()[-1]["mock"] is True
    finally:
        os.environ["J2_MEMORY"] = "0"
        j2.clear_playbooks()


def intake_needs():
    """H: at J2.a Identity proposes what this kind of business's visitors need and asks the Requester to confirm."""
    run = j2.J2Run()
    say(run, FLORIST)
    q = " ".join(x["question"] for x in run.paused["questions"])
    assert "florist" in q and "browse bouquets" in q and "Is that right" in q, q
    say(run, "Yes, and add wedding flowers. Petal & Stem, Jaipur. Elegant, must work on phones.")
    finish(run)
    assert run.goal["business_type"] == "florist" and run.goal["visitor_needs"], run.goal


def playbooks():
    """H: a Done run saves a playbook for its business type with the measured outcome; the next run for the same
    type reads it at J2.b and says which it used and what it changed; another type doesn't get it; a forced-failure
    run's playbook is a test playbook and isn't reused; J2_MEMORY=0 reads and writes nothing."""
    tmp = tempfile.mkdtemp()
    old = j2.MEMORY_FILE, j2.PLAYBOOK_FILE
    j2.MEMORY_FILE, j2.PLAYBOOK_FILE = os.path.join(tmp, "lessons.jsonl"), os.path.join(tmp, "playbooks.jsonl")
    os.environ["J2_MEMORY"] = "1"
    try:
        run = j2.J2Run()
        say(run, BREW_LAB)
        add_budget_if_asked(run)
        pb = j2.load_playbooks()[-1]
        assert pb["id"] == f"P-{run.run_id}" and pb["business_type"] == "cafe" and pb["build_order"], pb
        o = pb["outcome"]
        assert o["first_pass_audit"] is False and o["cycles"] >= 2 and o["ai_calls"] > 0 and pb["what_failed"], o
        assert any(m.get("act") == "PLAYBOOK" for m in run.chat) and summary(run)["playbooks"]["saved"] == pb["id"]
        first_id = pb["id"]
        run = j2.J2Run()                                   # a florist: the cafe playbook isn't offered
        say(run, FLORIST)
        say(run, "Yes. Petal & Stem, Jaipur. Elegant, must work on phones.")
        finish(run)
        assert run.outcome.startswith("Done"), run.outcome
        assert not any(m.get("playbooks_offered") for m in run.chat), "cafe playbook offered to a florist"
        assert j2.load_playbooks()[-1]["business_type"] == "florist"
        run = j2.J2Run()                                   # another cafe: reads it, says which, and what changed
        say(run, BREW_LAB)
        add_budget_if_asked(run)
        m = next(m for m in run.chat if m.get("playbook_used"))
        assert m["who"] == "adaptation" and m["playbook_used"] == first_id and first_id in m["text"] and m["playbook_changes"]
        assert run.playbooks_used[0]["playbook"] == first_id and summary(run)["playbooks"]["used"][0]["playbook"] == first_id
        assert j2.load_playbooks()[-1]["used_playbook"] == first_id
        n = len(j2.load_playbooks())
        os.environ["J2_FORCE_AUDIT_FAIL"] = "2"            # a test run: its playbook is flagged and not reused
        run = j2.J2Run()
        run.budget = 150
        say(run, BREW_LAB)
        add_budget_if_asked(run)
        os.environ.pop("J2_FORCE_AUDIT_FAIL")
        assert j2.load_playbooks()[-1]["test"] and len(j2.load_playbooks()) == n + 1
        assert not any(x.get("test") for x in j2.relevant_playbooks(j2.load_playbooks(), "cafe", BREW_LAB, n=10))
        os.environ["J2_MEMORY"] = "0"                      # memory off: nothing read, nothing written
        run = j2.J2Run()
        say(run, BREW_LAB)
        assert not any(m.get("playbooks_offered") or m.get("act") == "PLAYBOOK" for m in run.chat)
        os.environ["J2_MEMORY"] = "1"
        assert len(j2.load_playbooks()) == n + 1
        assert j2.clear_playbooks() == n + 1 and j2.load_playbooks() == []
    finally:
        j2.MEMORY_FILE, j2.PLAYBOOK_FILE = old
        os.environ["J2_MEMORY"] = "0"
        os.environ.pop("J2_FORCE_AUDIT_FAIL", None)


def generalist_units():
    """H: the design measurements and the screenshot on a planted page; no cafe examples in any live prompt of a
    florist run; the start screen lists a florist and a photographer."""
    if j2.j2_check.find_chrome():
        bad = ('<!doctype html><html><head><meta name="viewport" content="width=device-width"></head><body>'
               '<h1>A</h1><h1>B</h1><p style="color:#bbb">pale text</p><nav><a href="https://example.com" '
               'style="display:inline-block;font-size:10px">tiny</a></nav><img src="x.png"></body></html>')
        dp = j2.j2_check.design_problems(j2.j2_check.check(bad))
        assert set(dp) == {"R-2", "R-3", "R-6", "R-7"}, dp
        shot = os.path.join(tempfile.mkdtemp(), "s.png")
        assert j2.j2_check.screenshot(bad, shot) and os.path.getsize(shot) > 1000
    run = j2.J2Run()
    say(run, FLORIST)
    say(run, "Yes. Petal & Stem, Jaipur.")
    finish(run)
    # Our instructions carry no cafe examples: every system prompt (role files + reply formats, incl. Audit's) and the
    # plain-language rule. Chat history may still quote mock replies, which can stay.
    for bad in ("coffee's story", "cafe's phone"):
        assert bad not in j2.J2Run.PLAIN_TO_REQUESTER, bad
        for c in calls_of(run):
            assert bad not in c["system"], (c["fn"], bad)
    with open(os.path.join(j2.HERE, "j2_ui.html"), encoding="utf-8") as f:
        ui = f.read()
    assert "a florist in Jaipur" in ui and "wedding and portrait photographer" in ui


def page():
    """E: the behind-the-scenes view. Lanes with live status, a detail view per agent message (inputs, decision,
    reason, cost), the timeline's J2 steps, workflow-agent progress, saved runs opened afterwards, no secrets."""
    import threading
    import urllib.error
    import urllib.request
    import j2_server
    tmp = tempfile.mkdtemp()
    old, j2.MEMORY_FILE = j2.MEMORY_FILE, os.path.join(tmp, "lessons.jsonl")
    with open(j2.MEMORY_FILE, "w") as f:
        f.write('{"id": "L-x-1", "problem": "p", "cause": "c", "fix": "f"}\n')
    os.environ["J2_MEMORY"] = "1"
    os.environ["J2_TEST_API_KEY"] = "sk-test-not-a-real-key-123"
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    get = lambda path: urllib.request.urlopen(base + path).read().decode()
    try:
        j2_server.RUN.reset()
        say(j2_server.RUN, BREW_LAB + " Test key: sk-test-not-a-real-key-123")
        html = get("/")
        for x in ('data-tab="timeline"', 'data-tab="lessons"', 'id="lanes"', 'id="detail"', "renderLanes", "openDetail",
                  "renderTimeline", 'id="runsel"', "clear-lessons", "--wf0"):
            assert x in html, x
        raw = get("/api/state")
        st = json.loads(raw)
        assert "sk-test-not-a-real-key-123" not in raw and "[redacted]" in raw, "secret reached the page"
        assert st["agents"]["W-1"] == "W-1 Builder" and st["sections"] and st["write_lock"] is None
        assert st["workable"] == [] and st["memory"] and st["lessons"][0]["id"] == "L-x-1"
        # lanes: every function, the validator and each workflow agent; idle once done; workflow progress
        ids = [x["id"] for x in st["lanes"]]
        assert ids[:6] == list(j2.FUNCTIONS) and set(ids[6:]) == {w["id"] for w in st["workflows"]}, ids
        assert all(x["status"] == "idle" for x in st["lanes"]), [x["status"] for x in st["lanes"]]
        w1 = next(x for x in st["lanes"] if x["id"] == "W-1")["progress"]
        assert w1["section"] == "W-1" and w1["attempts"] == 1 and "usable section" in w1["last_check"], w1
        # detail view data: every AI message links its calls; each call has inputs, a reason, time; tokens or None
        ai = [m for m in st["chat"] if m.get("calls")]
        assert ai and all(m["step"] for m in st["chat"] if m["who"] not in ("user", "system")), \
            "agent messages need their J2 step for the timeline"
        log = {c["id"]: c for c in st["call_log"]}
        for m in ai:
            for cid in m["calls"]:
                c = log[cid]
                assert c["fn"] == m["who"] and c["inputs"] and c["reason"] and c["seconds"] >= 0, c
        assert all(c["tokens"] is None for c in st["call_log"]), "mock mode must not invent token counts"
        assert any(m.get("charged") for m in st["chat"] if m["who"] == "validator"), "budget charged per step"
        full = json.loads(get(f"/api/call?id={ai[0]['calls'][0]}"))
        assert full["system"] and full["prompt"] and full["raw"], full.keys()
        assert '"reason": str' in full["system"] or '"reason": str' in full["prompt"] or "reason" in full["system"]
        # a lane is "waiting" when it handed off and nobody answered yet; "blocked" for a failed workflow
        fake = [{"id": 1, "who": "adaptation", "text": "@Priority", "next": "priority"}]
        assert j2.lanes(fake, [], None, None, "running", None, 3, None)[1]["status"] == "waiting"
        wf = [{"id": "W-1", "name": "x", "status": "failed", "last_check": "no section", "envelope": {}}]
        j2.add_workflow_agent("W-1")
        assert j2.lanes([], wf, None, None, "waiting", None, 8, None)[6]["status"] == "blocked"
        # saved runs: listed, opened read-only with the same view, its site and its calls
        name = os.path.basename(j2_server.RUN.dir)
        runs = json.loads(get("/api/runs"))
        assert runs["live"] == name
        j2_server.RUN.reset()                               # the old run is now a saved run
        assert any(r["id"] == name for r in json.loads(get("/api/runs"))["runs"])
        sv = json.loads(get(f"/api/state?run={name}"))
        assert sv["saved"] == name and sv["status"] == "saved" and sv["total"] == st["total"] and sv["lanes"], sv["status"]
        assert sv["workable"] == [] and sv["call_log"] and "prompt" not in sv["call_log"][0]
        assert "J2:W-1" in get(f"/site/latest?run={name}")
        assert json.loads(get(f"/api/call?id=1&run={name}"))["id"] == 1
        for bad in ("../memory", "j2-../../etc", "nope"):
            try:
                get(f"/api/state?run={bad}")
                raise AssertionError(f"opened {bad}")
            except urllib.error.HTTPError as e:
                assert e.code == 404
        n = len(j2.load_lessons())             # the seeded lesson plus the one this run saved
        assert n == 2, n
        req = urllib.request.Request(base + "/api/lessons/clear", method="POST", data=b"{}")
        assert json.loads(urllib.request.urlopen(req).read())["cleared"] == n and not os.path.exists(j2.MEMORY_FILE)
    finally:
        srv.shutdown()
        j2.MEMORY_FILE = old
        os.environ["J2_MEMORY"] = "0"
        os.environ.pop("J2_TEST_API_KEY", None)


def limit_answers():
    """runs/j2-20261008-110419: the Requester answered a limit question while the cycle was at its turn cap, and the
    cap fired again before the answer was handled, so Identity re-asked forever. Now every answer is handled:
    "25 units" adds that amount, "2" drops the failing items with the reason, "3" ends Done with open items."""
    def at_cap():
        old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"   # fix (a) asks about a ruling first
        run = j2.J2Run()
        os.environ["J2_FIXES"] = old
        run.budget = 12                          # out of budget after the first cycle: a limit question
        say(run, BREW_LAB)
        assert run.paused and run.paused["kind"] == "limit", run.outcome
        run.cycle_turns = j2.MAX_TURNS           # the live run's state when the Requester answered
        return run
    run = at_cap()
    say(run, "i want a different option so go with 25 units")
    assert run.raises and run.raises[0]["asked"] == 25 and run.raises[0]["to"] == 37, run.raises
    assert "25 units" in run.raises[0]["why"] and run.decisions[0]["budget_added"] == 25, run.decisions
    assert not any(e["type"] == "limit" and "turns" in e["detail"] for e in run.events), "turn cap fired on the answer"

    run = at_cap()
    asked = run.paused["questions"]
    say(run, "2")
    g = run.goal
    assert g["dropped_by_requester"] and len(g["dropped_reasons"]) == len(g["dropped_by_requester"]), g
    assert all(c not in g["done_when"] for c in g["dropped_by_requester"]) and "leave it out" in g["dropped_reasons"][0]["reason"]
    assert not (run.paused and run.paused["questions"] == asked), "asked the same question again"

    run = at_cap()                           # the funded work was built first, so "finish as it is" is an option
    built_first(run)
    assert run.paused["handoff"]["choices"] == ["continue", "drop", "ship"], run.paused["handoff"]
    say(run, "Option 3")
    assert run.outcome.startswith("Done with open items") and not run.paused, run.outcome
    assert summary(run)["status"] == "done_with_open_items"


def nine_workflows():
    """runs/j2-20261008-190919: Audit passed every item at msg 40, then the 30-turn cap fired before J2.h, so Identity
    asked "keep going or finish?" about a finished site and the outcome listed no open items. The cap now scales
    with the workflows registered this cycle, J2.h always runs after a passing Audit, and finishing a workable
    site is "Done", never "Done with open items" with an empty list."""
    def nine(cap=None):
        old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"
        run = j2.J2Run()
        os.environ["J2_FIXES"] = old
        run.budget, orig = 400, run.call
        if cap:
            run._turn_cap = lambda: cap

        def call(fn, fmt, task, mock, **k):          # cycle 1: nine workflows, one item each, all of them pass
            if fn == "adaptation" and "WHY YOU ARE UP" in task and not run.workflows:
                crit = run.goal["done_when"]

                def wf(name, i, parts=("page_section",)):
                    return {"name": name, "goal_gap": crit[i - 1], "parts": list(parts), "covers": [i],
                            "outputs": ["index.html"], "completion_check": crit[i - 1]}
                phone = next(i for i, c in enumerate(crit, 1) if "phone" in c.lower())
                form = next(i for i, c in enumerate(crit, 1) if "form" in c.lower())
                rest = [i for i in range(2, len(crit) + 1) if i not in (phone, form)]
                wfs = [wf("Scaffold and styling", 1, ("page_section", "styling")), wf("Viewport fix", phone, ("fix",)),
                       wf("Pre-order form", form, ("form",))] + [wf(f"Section {i}", i) for i in rest]
                return {"say": "@Priority Nine workflows.", "act": "PROPOSE", "next": "Priority",
                        "workflows": j2.with_atoms(wfs + [wf(f"Extra {i}", i) for i in rest][:9 - len(wfs)])}
            return orig(fn, fmt, task, mock, **k)
        run.call = call
        say(run, BREW_LAB)
        return run
    run = nine()
    assert len(run.workflows) == 9 and all(w["cycle"] == 1 and w["status"] == "done" for w in run.workflows), run.workflows
    assert run.cycle_turns > j2.MAX_TURNS and run._turn_cap() == 10 + 4 * 9, (run.cycle_turns, run._turn_cap())
    assert run.outcome == j2.WORKABLE_OUTCOME and not run.paused, run.outcome
    assert not any(e["type"] == "limit" for e in run.events) and summary(run)["status"] == "done"
    # (1) the cap is reached exactly when Audit hands a passing site to Identity: J2.h still runs
    turns = run.cycle_turns
    run = nine(cap=turns - 1)
    assert run.outcome == j2.WORKABLE_OUTCOME and not any(e["type"] == "limit" for e in run.events), run.outcome
    # (3) "finish as it is" on a workable site is Done, not Done with open items and an empty list
    assert run._ship_outcome("on the Requester's decision") == j2.WORKABLE_OUTCOME


def budget_cap():
    """The cap shows in the option text; an option that adds budget the cap won't allow is never offered, and
    "keep working" isn't offered when it can't fund anything or when it already changed nothing. When that leaves
    one option, code applies it instead of asking."""
    run = j2.J2Run()
    q = {"question": "More?", "options": ["Add more building", "Ship"], "budget": [30, 0]}
    run._ask_user("x", [q], "question")             # only the decision question may carry money
    assert "budget" not in run.paused["questions"][0], run.paused
    run.sites = ["<!doctype html><p>v1</p>"]       # a site exists, so "finish as it is" is an option
    run.budget = j2.BUDGET_CAP - 40                 # Identity says 78 is needed; the cap allows 40
    run._identity_limit({"reason": "Adaptation asked", "budget_needed": 78})
    o = run.paused["questions"][0]["options"][0]
    assert o.startswith("Add 40 work credits (the most allowed)") and run.paused["questions"][0]["budget"][0] == 40, o
    def only_finish(h):                         # one option left: applied by code, never a one-option question
        n, run.paused = len(run.chat), None
        run._identity_limit(h)
        assert any(m["text"].startswith("Only one option was possible (Finish as it is") for m in run.chat[n:]), run.chat[n:]
        assert run.outcome.startswith("Done with open items: shipped as is as the only option") and not run.paused, run.outcome
    run.budget = j2.BUDGET_CAP
    only_finish({"reason": "Adaptation asked", "budget_needed": 78})
    cheapest = min(j2.RATES.values())
    run.spent = run.budget - cheapest + 1                       # less than the cheapest part left, no room to add
    only_finish({"reason": "Out of budget", "out_of_budget": True})
    run.spent = 0                                               # money left: keep working is offered once
    run._identity_limit({"reason": "turn cap"})
    assert run.paused["questions"][0]["options"][0].startswith(("Add ", "Keep going")), run.paused
    run.last_limit["decision"] = "continue"                     # ...and changed nothing: not offered again
    only_finish({"reason": "turn cap"})


def repeats():
    """An agent message the same as that agent's last one isn't posted again; Identity takes the turn."""
    run = j2.J2Run()
    run._say("adaptation", "@Identity No budget left, so I propose nothing.", "identity")
    n = len(run.chat)
    m = run._say("adaptation", "No budget left,  so I propose NOTHING.", "identity")
    assert m is run.chat[n - 1] and run.chat[n:] and all(x["who"] == "system" for x in run.chat[n:]), run.chat[n:]
    assert run.relay and run.relay["from"] == "adaptation" and any(e["type"] == "repeat" for e in run.events)
    run.relay = None
    run._say("adaptation", "@Identity A different message.", "identity")
    assert run.chat[-1]["who"] == "adaptation" and not run.relay


def adaptation_idle():
    """runs/j2-20261008-110419 msgs 93-103: with less left than the cheapest part, Adaptation said it could propose
    nothing, Identity answered, and Adaptation was asked again, four times. Once it has said so, it isn't asked again
    until something changes (budget raised, item dropped, a new Requester answer); Identity asks the Requester."""
    run = j2.J2Run()
    run.goal = {"done_when": ["The menu with prices is shown", "Opening hours are listed"]}
    run.budget, run.spent = 40, 38
    asked = []

    def call(fn, fmt, task, mock, **k):
        asked.append(fn)
        return {"say": "@Identity 2 left, the cheapest part costs more, so I propose nothing.",
                "act": "POLICY_QUESTION", "workflows": [], "next": "Identity"}
    run.call = call
    h = {"kind": "propose", "why": "Identity answered."}
    run.agent_turns["adaptation"] = 1
    run._turn_adaptation(h)       # a budget question: rejected, asked once more, then code's fallback fits nothing
    assert asked == ["adaptation"] * 2 and run.turn == "identity" and run.handoff["kind"] == "limit", run.handoff
    assert run.handoff["left_out"] == run.goal["done_when"], run.handoff
    run.agent_turns["adaptation"] = 2
    run._turn_adaptation(h)
    assert asked == ["adaptation"] * 2, "Adaptation was asked again with nothing changed"
    assert run.turn == "identity" and run.handoff["kind"] == "limit" and run.handoff["out_of_budget"], run.handoff
    assert run.chat[-1]["who"] == "system" and "isn't asked again" in run.chat[-1]["text"]
    for change in (lambda: run.decisions.append({"answer": "1"}),               # a new Requester answer
                   lambda: run.goal["done_when"].pop(),                        # an item dropped
                   lambda: setattr(run, "budget", run.budget + 1)):            # budget raised (still short)
        change()
        n = len(asked)
        run._turn_adaptation(h)
        assert len(asked) > n, "a change must let Adaptation re-plan"


def identity_decisions():
    """runs/j2-20261008-123848: when the money ran out, Identity asked its own free-form question (msgs 22, 48), so
    code applied nothing: "2. Add 25 more. Leave out ..." added nothing and the items came back (msg 24 -> 56), and
    "option 2: finish as is" led to a rejected done and the limit question again (msg 50 -> 56). Now money, leaving
    items out and finishing are only ever asked by code's decision question, and a typed answer is applied."""
    def out_of_money(built=False):
        old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"
        run = j2.J2Run()
        os.environ["J2_FIXES"] = old
        orig = run.call

        def call(fn, fmt, task, mock, **k):
            if fn == "identity" and "AUDIT'S FINDINGS" in task:   # what Identity did live: money in its own question
                return {"say": "@Requester The money is used up.", "done": False, "needs_requester": True,
                        "requester_questions": ["Option 1: Add more money so everything is built. Option 2: Add a "
                                                "smaller amount and drop some parts. Option 3: Finish as is."]}
            return orig(fn, fmt, task, mock, **k)
        run.call = call
        run.budget = 12
        say(run, BREW_LAB)
        if built:       # below the floor: the funded work is built and audited before the decision question
            built_first(run)
        assert run.paused and run.paused["kind"] == "limit", run.paused and run.paused["kind"]
        assert not any(m.get("pause") == "fact" for m in run.chat), "Identity asked about money in free form"
        return run

    # msg 24: add and leave out in one typed answer, both applied by code
    run = out_of_money()
    n = len(run.chat)
    gone = [c for c in run.goal["done_when"] if re.search(r"opening hours|loyalty|beans", c, re.I)]
    say(run, "2. Add 25 more. Leave out the opening hours, loyalty-card explainer and beans.")
    assert run.raises and run.raises[0]["asked"] == 25 and run.raises[0]["from"] == 12, run.raises
    g = run.goal
    assert len(gone) == 3 and all(c not in g["done_when"] for c in gone), g["done_when"]
    assert [x["criterion"] for x in g["dropped_reasons"]] == gone, g["dropped_reasons"]
    later = " ".join(m["text"] for m in run.chat[n + 1:] if m["who"] == "identity"
                     and not re.search(r"leave out|left out", m["text"]))     # not the drop's own acknowledgement
    assert not any(c in later for c in gone), "dropped items came back"

    # msg 50: "finish as is" ends the run, no rejected done and no limit question straight after
    run = out_of_money(built=True)
    n = len(run.chat)
    say(run, "finish as is")
    assert run.outcome.startswith("Done with open items") and not run.paused, run.outcome
    assert not any(m.get("pause") for m in run.chat[n:]), [m["text"][:80] for m in run.chat[n:] if m.get("pause")]
    assert not any("Identity declared done" in m["text"] for m in run.chat[n:])
    assert summary(run)["status"] == "done_with_open_items"


BANNED_FOR_REQUESTER = re.compile(r"building time|\bcaps?\b|\bcycles?\b|done[-_ ]when|\bestimate", re.I)


def plain_for_requester(run):
    """Everything Identity sends the Requester: no @ in front, no internal words, at most 3 sentences plus one
    question, at most one question with at most 3 options."""
    for m in run.chat:
        if m["who"] != "identity" or not m.get("to_user"):
            continue
        say = m["text"].split("\n")[0]
        assert say.startswith("@Requester ") and not re.search(r"@(?!Requester\b)\w", say), m["text"]
        assert not BANNED_FOR_REQUESTER.search(m["text"].replace("checklist item", "")), m["text"]
        parts = [x for x in re.split(r"(?<=[.!?])\s+", say[len("@Requester "):]) if x.strip()]
        assert len([x for x in parts if not x.endswith("?")]) <= 3 and len([x for x in parts if x.endswith("?")]) <= 1, say
        qs = re.findall(r"^\d+\. ", m["text"], re.M)
        assert len(qs) <= 1 and len(re.findall(r"^\s+Option \d+:", m["text"], re.M)) <= 3, m["text"]


def two_audiences():
    """Item 4, corrected: one message, one addressee, the tag matches the text. (a) Live run 8 Oct 14:06 sent "@Adaptation
    Thanks, I have everything I need. I'll start building..." as one message: it becomes one @Requester and one
    @Adaptation message. (b) A to_requester with agent content ("re-price", "W-3") is rejected and regenerated. Agents
    never get the to_requester text."""
    old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"
    run = j2.J2Run()
    os.environ["J2_FIXES"] = old
    orig, seen = run.call, []

    def call(fn, fmt, task, mock, **k):
        if fn == "identity" and task.startswith("REQUESTER'S BRIEF"):          # (a) what Identity said live
            def mixed():
                r = json.loads(mock())
                r.pop("to_requester", None), r.pop("to_agents", None)
                if r.get("enough"):
                    r["say"] = "@Adaptation Thanks, I have everything I need. I'll start building the Brew Lab site now."
                return json.dumps(r)
            return orig(fn, fmt, task, mixed, **k)
        if fn == "identity" and task.startswith("YOU ASKED THE REQUESTER"):   # (b) agent content, then a clean retry
            seen.append(task)
            def first():                       # bad only on the first attempt; the retry gets a clean reply
                r = json.loads(mock())
                if len(seen) == 1:
                    r["to_requester"] = "Got it. I'll re-price W-3 for Priority and drop those items."
                return json.dumps(r)
            return orig(fn, fmt, task, first, **k)
        return orig(fn, fmt, task, mock, **k)
    run.call = call
    run.budget = 12
    say(run, BREW_LAB)
    ch = next(i for i, m in enumerate(run.chat) if m.get("artifact") == "charter")
    pair = run.chat[ch - 1:ch + 1]
    assert pair[0]["text"] == "@Requester Thanks, I have everything I need. I'm starting the build now." and \
        pair[0]["to_user"] and pair[0]["audience"] == "user", pair
    assert pair[1]["text"].startswith("@Adaptation Charter is ready") and not pair[1].get("to_user") \
        and pair[1]["audience"] == "adaptation", pair
    assert any(e["type"] == "audience_template" for e in run.events)
    n = len(run.chat)
    say(run, "2")
    assert len(seen) == 2 and "MIXED THE TWO AUDIENCES" in seen[1] and "re-pric" in seen[1], seen[-1][-300:]
    mine = [m for m in run.chat[n:] if m.get("requester_only")]
    # the answer left items out, so code writes the reply naming them (Identity's own words could name others)
    assert len(mine) == 1 and mine[0]["text"].startswith("@Requester Got it. We'll leave out item "), mine
    assert not any("re-price" in m["text"] for m in run.chat if m.get("to_user"))
    assert "@Requester Thanks, got it." not in run.chat_view(), "agents must not get to_requester"
    plain_for_requester(run)


def simple_messages():
    """Item 5: what the Requester reads is short and plain; one question, at most 3 options; the decision question's
    options say what they get."""
    t = j2.to_requester("@Requester The building time is used up. We hit the cap after 3 cycles. The done-when "
                        "list has 4 items (estimate). One more sentence. And another. Which do you want? Really?")
    assert t == ("The work is used up. We hit the limit after 3 rounds. The checklist list has 4 items. Which do you "
                 "want?"), t
    assert j2.to_requester("Done. What now?", question=False) == "Done."
    run = j2.J2Run()
    run._ask_user("Two things. What is your phone number?", [{"question": "What is your phone number?"},
                                                              {"question": "What are your hours?"}], "info")
    q = run.paused["questions"]
    assert len(q) == 1 and q[0]["question"] == "Please send these: What is your phone number; What are your hours.", q
    assert "?" not in run.chat[-1]["text"].split("\n")[0], run.chat[-1]["text"]
    run._ask_user("Pick one.", [{"question": "Which?", "options": ["a", "b", "c", "d"]}], "question")
    assert run.paused["questions"][0]["options"] == ["a", "b", "c"], run.paused
    run.budget, run.spent = 40, 38
    run._identity_limit({"reason": "Out of budget", "out_of_budget": True})
    opts = run.paused["questions"][0]["options"]                # no site yet: "stop here", never "finish"
    assert opts[0].startswith("Add ") and "you get" in opts[0] and opts[-1].startswith("Stop here: nothing is built"), opts
    run.sites = ["<!doctype html><p>v1</p>"]
    run._identity_limit({"reason": "Out of budget", "out_of_budget": True})
    opts = run.paused["questions"][0]["options"]
    assert opts[0].startswith("Add ") and "you get" in opts[0] and opts[-1].startswith("Finish as it is: you get"), opts
    plain_for_requester(run)


def owner_override():
    """Item 6: a message starting with J2_OVERRIDE_KEY and a command is applied by code; the key is never posted,
    stored or sent to an agent, and the server redacts it."""
    import j2_server
    key = "test-override-key-123"
    old_key, os.environ["J2_OVERRIDE_KEY"] = os.environ.get("J2_OVERRIDE_KEY"), key
    old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"
    try:
        run = j2.J2Run()
        os.environ["J2_FIXES"] = old
        run.budget = 12
        say(run, BREW_LAB)
        assert run.paused and run.paused["kind"] == "limit"
        say(run, f"{key} budget +500")
        assert run.budget == 512 and run.raises[-1]["override"] and run.budget > j2.BUDGET_CAP, run.raises
        assert run.chat[-1]["text"].startswith("Owner override applied: budget +500") and run.paused, run.chat[-1]
        crit = run.goal["done_when"][1]
        say(run, f"{key} drop 2")
        assert crit not in run.goal["done_when"] and run.goal["dropped_reasons"][-1]["reason"] == "Owner override"
        say(run, f"{key} paint it blue")
        assert run.chat[-1]["text"].startswith("Owner override not applied"), run.chat[-1]
        say(run, f"{key} ship")                    # the funded work was built before the question
        assert run.outcome.startswith("Done with open items") and not run.paused and run.done, run.outcome
        assert summary(run)["status"] == "done_with_open_items"
        for name in ("chat.jsonl", "calls.jsonl", "goal.json", "summary.json"):
            path = os.path.join(run.dir, name)
            assert not os.path.exists(path) or key not in open(path, encoding="utf-8").read(), name
        assert key not in run.chat_view() and key not in json.dumps(run.chat)
        assert key not in j2_server.redact(f"typed: {key} ship"), "the server must redact the key"
    finally:
        os.environ["J2_FIXES"] = old
        if old_key is None:
            os.environ.pop("J2_OVERRIDE_KEY", None)
        else:
            os.environ["J2_OVERRIDE_KEY"] = old_key


def form_submit():
    """Item 7 (run 20261008-123848 msg 46: a form with no submit button): a section whose form has no submit button
    is rejected so the agent retries; the browser check still makes such a site not workable."""
    for body, ok in (('<form><input name="n"><button type="button">Order</button></form>', False),
                     ('<form><input name="n"></form>', False),
                     ('<form><input name="n"><button>Order on WhatsApp</button></form>', True),
                     ('<form><input type="submit" value="Send"></form>', True)):
        try:
            j2.parse_section("```section\n" + body + "\n```")
            got = True
        except ValueError as e:
            got = False
            assert "submit button" in str(e), e
        assert got == ok, body
    bad = {"results": [{"item": 1, "id": "F-1", "verdict": "PASS", "criterion": "a", "pass": True}],
           "mechanical": {"available": True, "phone": {"overflow_px": 0}, "desktop": {"overflow_px": 0},
                          "js_errors": [], "broken_in_page_links": [],
                          "form_results": [{"index": 1, "ok": False, "has_submit": False, "errors_during": []}]}}
    assert "form 1 has no submit button" in j2.workable(bad)


def audiences_everywhere():
    """Follow-up to 5ba8160. J2.h's reply has the same two parts and check: "@Adaptation Thanks, your site is almost
    ready..." is regenerated into one @Requester and one @Adaptation message. REPLY_FMT allows the @Requester tag.
    post() never re-tags an @agent message as @Requester; it records an event."""
    assert "never starts with @" not in j2.REPLY_FMT and "tagged @Requester or no tag" in j2.REPLY_FMT, j2.REPLY_FMT
    run = j2.J2Run()
    orig, seen = run.call, []

    def call(fn, fmt, task, mock, **k):
        if fn == "identity" and task.startswith("AUDIT'S FINDINGS") and len(seen) < 2:
            seen.append(task)
            def h():
                if len(seen) == 1:      # what Identity sends live: user talk addressed to an agent
                    return json.dumps({"say": "@Adaptation Thanks, your site is almost ready. Fix the phone layout next.",
                                       "clause": "done_when", "done": False, "next": "Adaptation"})
                return json.dumps({"to_requester": "Your site is almost ready. We're fixing the phone layout now.",
                                   "to_agents": "@Adaptation Fix the phone layout next.", "clause": "done_when",
                                   "done": False, "next": "Adaptation"})
            return orig(fn, fmt, task, h, **k)
        return orig(fn, fmt, task, mock, **k)
    run.call = call
    say(run, BREW_LAB)
    assert len(seen) == 2 and "MIXED THE TWO AUDIENCES" in seen[1], "J2.h's mixed reply must be regenerated"
    i = next(j for j, m in enumerate(run.chat) if "Your site is almost ready" in m["text"])
    req, ag = run.chat[i], run.chat[i + 1]
    assert req["text"] == "@Requester Your site is almost ready. We're fixing the phone layout now." and req["to_user"], req
    assert ag["text"] == "@Adaptation Fix the phone layout next." and not ag.get("to_user") \
        and ag["audience"] == "adaptation", ag
    assert not any("Thanks, your site" in m["text"] for m in run.chat)
    n = len(run.events)
    m = run.post("identity", "@Adaptation Re-price W-3.", to_user=True)
    assert m["text"] == "@Adaptation Re-price W-3." and run.events[n]["type"] == "misaddressed", (m, run.events[n:])


def listed_gaps(text):
    """The "Still not right" entries of a decision message: one per line, "- " in front."""
    return [ln[2:] for ln in text.split("\n") if ln.startswith("- ")]


def limit_wording():
    """Run 20261008-154537: msg 5 printed "1.\n1." (Identity's own numbered list plus code's); msgs 45/46 paused at
    the per-cycle turn cap with a money banner although no money option was offered, and "1 thing still not right"
    didn't match the gaps the message listed. The list is numbered once; the question and the banner say why the team
    paused and offer only what's on offer; the gap list and its count come from one list; every option says what you
    get, with the amount when money is involved."""
    run = j2.J2Run()
    run._ask_user("Thanks, that covers most of it. A few things are still missing.\n1. Which products are new?\n"
                  "2. Which sizes?", [{"question": "Which products are new arrivals?"},
                                      {"question": "Which sizes does each product come in?"}], "info")
    txt = run.chat[-1]["text"]
    assert not re.findall(r"(?m)(?:^|\s)\d+\.(?:\s|$)", txt) and txt.endswith("Which sizes does each product come in."), txt

    old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"
    run = j2.J2Run()
    os.environ["J2_FIXES"] = old
    say(run, BREW_LAB)
    run.audits[-1]["results"][0].update(verdict="NEEDS_REVISION", **{"pass": False})
    run.audits[-1]["results"][2].update(verdict="NEEDS_REVISION", **{"pass": False})
    run._identity_limit({"reason": f"Cycle {run.cycle} used {j2.MAX_TURNS} turns (the per-cycle cap) without finishing."})
    q = run.paused["questions"][0]
    assert q["question"] == "The team paused after a round of work. Keep going, leave out the missing parts, or " \
                            "finish as it is?", q["question"]
    assert "budget" not in q and "add budget" not in run.outcome and q["question"] in run.outcome, run.outcome
    gaps = j2.workable(run.audits[-1])
    listed = listed_gaps(run.chat[-1]["text"])
    assert len(listed) == len(gaps) and f"with {len(gaps)} things still not right: item " in q["options"][-1], (listed, q)
    for o in q["options"]:
        assert "you get" in o, o
    run.spent = run.budget                                    # money: the option says the amount
    run._identity_limit({"reason": "Out of budget", "out_of_budget": True})
    q = run.paused["questions"][0]
    assert q["question"].startswith("The work you approved is used up. Add more and keep going"), q["question"]
    assert re.match(r"Add \d+ work credits", q["options"][0]) and all("you get" in o for o in q["options"]), q["options"]


def open_list():
    """Follow-up to 6e5aa59 (run 20261008-154537 msg 45: the audit had 1 problem, but reviews, styling and menu links
    weren't built yet). The list combines the last audit's failures with every checklist item an unfinished workflow
    covers, one entry per item; the count is its length; items built after the last audit are marked."""
    def paused_run():
        old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"
        run = j2.J2Run()
        os.environ["J2_FIXES"] = old
        say(run, BREW_LAB)                                      # done and workable: a clean last audit
        run.audits[-1]["results"][0].update(verdict="NEEDS_REVISION", **{"pass": False})   # 1 audit failure: item 1
        return run

    def unfinished(run, covers, status="registered"):
        run.workflows.append({"id": f"W-{len(run.workflows) + 1}", "name": "x", "status": status, "covers": covers})

    def ask(run):
        run._identity_limit({"reason": f"Cycle {run.cycle} used {j2.MAX_TURNS} turns (the per-cycle cap) without "
                                       "finishing."})
        text = run.chat[-1]["text"]
        listed = listed_gaps(text)
        return text, listed, run.paused["questions"][0]["options"][-1]

    # (a) 1 audit failure + 3 unfinished workflows (registered, running, unfunded) -> 4 items, "4 things"
    run = paused_run()
    unfinished(run, [2]), unfinished(run, [3], "running"), unfinished(run, [5], "unfunded")
    text, listed, ship = ask(run)
    crit = run.goal["done_when"]
    assert len(listed) == 4 and "with 4 things still not right: " in ship, (listed, ship)
    row = lambda i, st: f"item {i} {j2.short(crit[i - 1])} ({st})"           # run-long number, short text, status
    assert listed[:2] == [row(1, "built but failing"), row(2, "not built")] and listed[3] == row(5, "unfunded"), listed

    # (b) an item both failing in the audit and covered by an unfinished workflow appears once
    run = paused_run()
    unfinished(run, [1, 2])
    text, listed, ship = ask(run)
    assert listed == [row(1, "built but failing"), row(2, "not built")] and "with 2 things still not right: " in ship, \
        listed

    # (c) the site is newer than the last audit: what was built since is marked, and the message says so
    run = paused_run()
    run.sites.append(run.sites[-1])
    run.workflows.append({"id": "W-99", "name": "y", "status": "done", "covers": [1, 4], "version": len(run.sites)})
    text, listed, ship = ask(run)
    assert listed == [row(1, "built, not checked yet"), row(4, "built, not checked yet")], listed
    assert "built after the last check, so they're not checked yet" in text and "with 2 things still not right: " in ship, text


def covers_by_text():
    """Follow-up to 0543422: a workflow stores the text of the done-when items it covers, so dropping an item (which
    renumbers the list) can't point an unfinished workflow at the wrong item, and a dropped item never shows."""
    old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"
    run = j2.J2Run()
    os.environ["J2_FIXES"] = old
    say(run, BREW_LAB)
    crit = list(run.goal["done_when"])
    run.budget, run.p_refused = 500, {}
    reg, refused, _ = run._register([
        {"name": f"Item {i}", "goal_gap": crit[i - 1], "outputs": ["x"], "completion_check": "y", "atoms": ATOM,
         "parts": ["form" if "form" in crit[i - 1].lower() else "page_section"], "covers": [i]} for i in (4, 5, 6)])
    assert len(reg) == 3 and not refused, refused
    assert [w["covers_text"] for w in reg] == [[crit[3]], [crit[4]], [crit[5]]], [w.get("covers_text") for w in reg]
    run.audits[-1]["results"][3].update(verdict="NEEDS_REVISION", **{"pass": False})     # item 4 also fails the audit
    run._override("drop 4")                                                          # renumbers: 5 -> 4, 6 -> 5
    assert run.goal["done_when"] == crit[:3] + crit[4:], run.goal["done_when"]
    items = [r for r in run.audits[-1]["results"] if not r["pass"]]
    listed = run._open_list(run.audits[-1], items)
    # (a) the unfinished workflows still name their own items, at their new numbers
    assert listed == [f"{crit[4]} (not built yet)", f"{crit[5]} (not built yet)"], listed
    assert [run._wf_items(w) for w in reg] == [[], [4], [5]], [run._wf_items(w) for w in reg]
    # (b) the dropped item never appears: not from its workflow, not from the audit that still lists it
    assert not any(crit[3] in x for x in listed), listed
    # (c) runs saved before covers_text still load and fall back to the numbers
    for w in run.workflows:
        w.pop("covers_text", None)
    run.save("workflows.json", run.workflows)
    snap = j2.saved_snapshot(os.path.basename(run.dir))
    assert snap and len(snap["workflows"]) == len(run.workflows), "an old run must load"
    # covers holds charter numbers: item 5 sits at position 4 after the drop
    assert run._wf_items(reg[1]) == [4] and run._open_list(run.audits[-1], []), "old workflows use their numbers"


def readable_decision():
    """Follow-up to bab0ca7: the decision question is readable without shortening the gaps. (1) one gap per line,
    full text, after at most 2 sentences, then the question and options; (2) no "1." before the question; (3) work
    credits, not units; (4) an unclear answer gets one short line and the options only; (5) the counts add up, and a
    difference is explained."""
    old, os.environ["J2_FIXES"] = os.environ.get("J2_FIXES"), "0"
    run = j2.J2Run()
    os.environ["J2_FIXES"] = old
    say(run, BREW_LAB)                                          # done and workable: a clean last audit
    run.spent = run.budget                                      # then out of money: add / leave out / finish
    crit = run.goal["done_when"]
    last = run.audits[-1]
    for i in (1,):
        last["results"][i].update(verdict="NEEDS_REVISION", **{"pass": False})
    last["rubric"][0].update(verdict="NEEDS_REVISION", **{"pass": False})
    for i in (3, 4, 5):
        run.workflows.append({**run.workflows[0], "id": f"W-{90 + i}", "status": "registered", "covers": [i + 1],
                              "covers_text": [crit[i]]})
    run._identity_limit({"reason": "Out of budget", "out_of_budget": True})
    text = run.chat[-1]["text"]
    lines = text.split("\n")
    gaps = listed_gaps(text)
    q = run.paused["questions"][0]
    # (1) intro (at most 2 sentences), the header, one full gap per line, then the question, then the options
    intro = lines[0][len("@Requester "):]
    assert len([x for x in re.split(r"(?<=[.!?])\s+", intro) if x]) <= 2, intro
    i = lines.index(next(ln for ln in lines if ln.startswith("Still not right (")))
    assert lines[i + 1:i + 1 + len(gaps)] == [f"- {g}" for g in gaps], lines
    # each line: the item's run-long number (a rubric item: no id), its short text and Audit's status
    assert len(gaps) == 5 and all(re.fullmatch(r"(item \d+ )?.+ \((built but failing|not built|unfunded|failing)\)", g)
                                  and not j2.INTERNAL_ID.search(g)
                                  for g in gaps), gaps
    assert lines[i + 1 + len(gaps)] == q["question"] and lines[i + 2 + len(gaps)].startswith("   Option 1: "), lines
    # (2) no "1." before the question
    assert not any(re.match(r"\d+\. ", ln) for ln in lines), lines
    # (3) work credits, never units
    assert "work credits" in q["options"][0] and not re.search(r"\bunits?\b", text), text
    # (5) counts: the header explains the mix; finish counts the whole list; leave-out explains the difference
    assert lines[i] == f"Still not right ({len(gaps)}: 2 built but failing, 3 not built yet):", lines[i]
    assert f"with {len(gaps)} things still not right: " in q["options"][-1], q["options"]
    drop = next(o for o in q["options"] if o.startswith("Leave out"))
    assert drop.startswith("Leave out item ") and drop.count("item ") == 4 and drop.endswith(
        "Consistent styling stays as it is, still failing."), drop
    # (4) unclear answer: one short line, then the options only
    say(run, "hmm, not sure")
    re_ask = next(m for m in reversed(run.chat) if m["who"] == "identity")["text"].split("\n")
    assert re_ask[0] == "@Requester I couldn't tell which option you chose. Reply 1, 2 or 3, or type what to add or " \
                        "leave out.", re_ask[0]
    assert len(re_ask) == 4 and all(ln.startswith("   Option ") for ln in re_ask[1:]), re_ask


PRD_BRIEF = ("Write a PRD for QuickQueue, a mobile app that lets clinic patients join a waiting line remotely. Users: "
             "patients at small clinics and the front-desk staff who manage the line. Problem: patients wait 45 minutes on "
             "average in crowded waiting rooms. Must-have features: join the queue from a phone, live position updates, SMS "
             "alert when 2 patients are ahead, staff dashboard to call the next patient. Constraints: works on low-end "
             "Android phones; no patient health data stored. Success: average in-room wait under 10 minutes; 60% of "
             "patients join remotely within 3 months.")
BROWSER_FNS = ("check", "screenshot", "tap_targets", "find_chrome", "design_problems")


def prd_run():
    """J2_OUTPUT=prd: a mock PRD brief reaches Done workable with prd.md and prd.docx; no browser check is ever called;
    only P-conditions and the D-rubric are listed, never a website's; the validator keeps the facts-present rule and
    the document's own check types."""
    called, saved = [], {k: getattr(j2.j2_check, k) for k in BROWSER_FNS}

    def boom(name):
        def f(*a, **k):
            called.append(name)
            raise AssertionError(f"a PRD run called the browser check {name}")
        return f
    for k in BROWSER_FNS:
        setattr(j2.j2_check, k, boom(k))
    os.environ["J2_OUTPUT"] = "prd"
    try:
        run = j2.J2Run()
        assert run.output == "prd" and run.doc and j2.RATES == {"doc_section": 10, "requirements_section": 48, "fix": 8}
        say(run, PRD_BRIEF)
    finally:
        os.environ.pop("J2_OUTPUT")
        for k, v in saved.items():
            setattr(j2.j2_check, k, v)
    assert not called, called
    assert run.outcome.startswith("Done: the PRD is workable"), run.outcome
    s = summary(run)
    with open(os.path.join(run.dir, "chat.jsonl"), encoding="utf-8") as f:
        assert json.loads(f.readline())["output"] == "prd"
    assert s["output_type"] == "prd" and s["status"] == "done" and s["workable"] == [], s
    assert [x["heading"] for x in s["sections"]] == list(j2.PRD_SECTIONS) and all(x["words"] for x in s["sections"]), s
    # P8 is a warning ("under 10 minutes" states no window): shown, and the PRD is still workable
    assert (s["conditions"]["met"], s["conditions"]["total"], s["conditions"]["warnings"]) == (7, 8, 1), s["conditions"]
    assert s["export"]["written"], s
    assert s["verification"]["passed"] == 6 and not s["verification"]["failed"] and s["verification"]["checks_run"], s
    # one workflow per section, one atom each, priced from the document rate card
    assert [w["name"] for w in run.workflows] == [f"Write {x}" for x in j2.PRD_SECTIONS], run.workflows
    assert all(len(w["atoms"]) == 1 for w in run.workflows) and run.spent == 98, run.spent
    md = open(os.path.join(run.dir, "prd.md"), encoding="utf-8").read()
    assert md.startswith("# QuickQueue") and [h for lv, h, _ in j2.md_sections(md) if lv == 2] == list(j2.PRD_SECTIONS)
    assert "1. **R1**:" in md and "<" not in md, md
    conds = json.load(open(os.path.join(run.dir, "conditions.json"), encoding="utf-8"))
    assert [c["id"] for c in conds] == [f"P{i}" for i in range(1, 9)], conds
    assert all(c["met"] for c in conds[:7]) and not conds[7]["met"] and conds[7]["severity"] == "warn", conds
    # never a website condition: no render check, no R-rubric, no tap-targets rule in any prompt
    a = run.audits[-1]
    assert [r["id"] for r in a["rubric"]] == ["D-1", "D-2", "D-3", "D-4", "D-5"] and not a["mechanical"], a["rubric"]
    assert a["output"] == "prd" and a["read"] and a["conditions"] == conds, a
    assert not any(m["text"].startswith("Mechanical check") for m in run.chat)
    calls = [json.loads(x) for x in open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8")]
    for c in calls:
        text = c["system"] + c["prompt"]
        assert "tap-targets" not in text and "Tap targets are big enough" not in text and "R-6" not in text, c["fn"]
    assert any("OUTPUT TYPE: PRD document" in c["prompt"] for c in calls)
    assert any(c["fn"].startswith("W-") and "Workflow agent for a PRD section" in c["system"] for c in calls)
    # the .docx: headings, numbered requirements and the metrics table come through
    import docx
    d = docx.Document(os.path.join(run.dir, "prd.docx"))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert heads == list(j2.PRD_SECTIONS), heads
    reqs = [p.text for p in d.paragraphs if p.style.name == "List Number"]
    assert reqs[0].startswith("R1:") and len(reqs) == 4, reqs
    assert d.tables and d.tables[0].cell(1, 0).text == "G1", [t.cell(0, 0).text for t in d.tables]
    # the validator: facts-present keeps its rule (charter or the Requester's answers only); document check types only
    atom = j2.mock_doc_atom({"name": "Write Problem"})
    fact = {"text": "Product name", "type": "facts-present", "target": ["QuickQueue", "45 minutes"], "expect": True}
    assert not run._atom_problems({"atoms": [{**atom, "verification": atom["verification"] + [fact]}]})
    bad = run._atom_problems({"atoms": [{**atom, "verification": [{**fact, "target": ["QuickQueue", "99.9% uptime"]}]}]})
    assert len(bad) == 1 and "['99.9% uptime'], not found in the charter" in bad[0], bad
    for chk, why in (({"type": "tap-targets", "target": ".x", "expect": 44}, "allowed: facts-present, heading-present"),
                     ({"type": "ids-unique", "target": "Requirement", "expect": True}, "only an id prefix"),
                     ({"type": "section-not-empty", "target": "Users", "expect": 0}, "minimum word count")):
        bad = run._atom_problems({"atoms": [{**atom, "verification": [{"text": "x", **chk}]}]})
        assert len(bad) == 1 and why in bad[0], (chk, bad)
    # a website run after it is a website again
    assert j2.J2Run().output == "website" and "page_section" in j2.RATES
    show(run)
    return run


def prd_checks():
    """heading-present, section-not-empty and ids-unique pass and fail correctly (a duplicate R2 is rejected); each
    necessary condition fails on the document that breaks it."""
    doc = "# X\n\n## Users\n\nClinic patients and front-desk staff.\n\n## Requirements\n\n1. **R1**: Join the queue\n2. **R2**: See the position\n"
    chk = lambda t, target, expect=True, d=doc: j2.doc_check({"type": t, "target": target, "expect": expect}, d)
    assert chk("heading-present", "Requirements")["pass"] and chk("heading-present", "users")["pass"]
    assert not chk("heading-present", "Goals")["pass"]
    assert chk("section-not-empty", "Users", 5)["pass"] and not chk("section-not-empty", "Users", 50)["pass"]
    assert not chk("section-not-empty", "Goals", 1)["pass"]
    assert not chk("section-not-empty", "Users", 1, "## Users\n\n## Goals\n\n- G1: x\n")["pass"]     # empty section
    assert chk("ids-unique", "R")["pass"] and chk("ids-unique", "R")["actual"] == [1, 2]
    dup = chk("ids-unique", "R", d=doc + "3. **R2**: Get an SMS\n")
    assert not dup["pass"] and "duplicate id(s): R2" in dup["evidence"], dup
    assert not chk("ids-unique", "R", d="## Requirements\n\nJoin the queue.\n")["pass"]
    assert chk("ids-unique", "R", d=doc + "\nThe SMS alert also serves R2.\n")["pass"]   # a mention, not an id
    # the same checks through run_check, as code runs a work atom's verification
    assert not j2.run_check({"type": "ids-unique", "target": "R"}, "", doc + "- **R2**: again\n", doc=True)["pass"]
    # necessary conditions, each on a document that breaks it
    conds = {c["id"]: c for c in j2.load_conditions("prd.json")}
    goal = {"product_name": "QuickQueue", "facts": ["QuickQueue", "SMS alert"], "to_be_confirmed": []}
    src = j2.squash("QuickQueue SMS alert 45 minutes 10 minutes 60% 3 months")
    good = ("# QuickQueue\n\n## Problem\n\nPatients wait 45 minutes.\n\n## Users\n\nClinic patients.\n\n## Goals\n\n- **G1**: "
            "Short waits\n\n## Requirements\n\n1. **R1**: SMS alert\n2. **R2**: Join from a phone\n\n## Success metrics\n\n"
            "| Goal | Metric |\n|---|---|\n| G1 | Under 10 minutes, measured monthly |\n\n## Out of scope\n\n- Payments\n")
    met = lambda cid, d, g=goal: j2.condition_met(conds[cid], d, g, src, j2.PRD_SECTIONS)
    assert all(met(c, good)[0] for c in conds), [(c, met(c, good)) for c in conds]
    assert not met("P1", good.replace("## Out of scope\n\n- Payments\n", "## Out of scope\n"))[0]
    assert not met("P1", good.replace("## Users", "## People"))[0]
    assert not met("P2", good.replace("phone\n", "phone\n3. **R2**: Duplicate\n"))[0] and not met("P2", good.replace("**R2**", "**R3**"))[0]
    assert not met("P3", good.replace("SMS alert", "text message"))[0]
    p4 = met("P4", good.replace("Under 10 minutes", "Under 7 minutes"))
    assert not p4[0] and "7" in p4[1], p4
    assert not met("P5", good.replace("Payments", "Payments (to be confirmed)"))[0]
    assert not met("P5", good, {**goal, "to_be_confirmed": [{"criterion": "x", "fact": "y"}]})[0]
    assert not met("P6", good.replace("| G1 | Under", "| Wait | Under"))[0]
    assert not met("P6", good.replace("- **G1**: Short waits", "- Short waits"))[0]
    # workable() for a document: conditions and the D-rubric count, the render check doesn't
    audit = {"output": "prd", "results": [], "rubric": [], "mechanical": {},
             "conditions": [{"id": "P2", "text": "ids", "met": False, "evidence": "R2 twice"}]}
    assert j2.workable(audit) == ["necessary condition P2 not met: ids (R2 twice)"], j2.workable(audit)
    assert j2.workable({**audit, "conditions": []}) == []


def prd_groups_metrics_assumptions():
    """P7 user groups, P8 metric windows (a warning), P4 with Assumption: lines and proper names, the closing
    assumptions section, the window in Identity's number question, and short item names cut at a natural break."""
    conds = {c["id"]: c for c in j2.load_conditions("prd.json")}
    goal = {"product_name": "LabSlot", "users": "Students (book slots) and lab staff", "facts": ["LabSlot"],
            "to_be_confirmed": []}
    src = j2.squash("LabSlot students lab staff book a slot 80% of bookings are made without staff help")
    doc = ("# LabSlot\n\n## Users\n\nStudents and lab staff.\n\n## Requirements\n\n1. **R1**: A student can book a slot\n\n"
           "## Success metrics\n\n- **G1**: 80% of bookings are made without staff help\n")
    met = lambda cid, d, g=goal: j2.condition_met(conds[cid], d, g, src, j2.PRD_SECTIONS)
    assert j2.user_groups(goal) == ["Students", "lab staff"], j2.user_groups(goal)
    p7 = met("P7", doc)
    assert not p7[0] and "lab staff" in p7[1] and "Students" not in p7[1], p7       # "a student" covers Students
    assert met("P7", doc.replace("lab staff.\n", "lab staff.\n\nLab staff have no actions in this version.\n"))[0]
    assert met("P7", doc, {**goal, "user_groups": [{"name": "students", "actions": ["book"]}]})[0]
    # P8: a warning, never a blocker
    p8 = met("P8", doc)
    assert not p8[0] and "80% of bookings" in p8[1], p8
    for window in (" per term", ", measured weekly", " (measurement window not set)", " within 3 months"):
        assert met("P8", doc.replace("staff help\n", f"staff help{window}\n"))[0], window
    rows = j2.check_conditions(j2.load_conditions("prd.json"), doc, goal, src, j2.PRD_SECTIONS)
    w = next(c for c in rows if c["id"] == "P8")
    assert w["severity"] == "warn" and not w["met"] and not j2.blocks(w), w
    assert all("severity" not in c for c in rows if c["id"] != "P8")             # the others stay blocking
    audit = {"output": "prd", "results": [], "rubric": [], "mechanical": {}, "conditions": [w]}
    assert j2.workable(audit) == [], j2.workable(audit)
    # P4: a number or proper name not in the charter fails outside an Assumption: line and passes inside one
    for claim in ("Bookings open 7 days ahead.", "Northfield College runs the labs."):
        bad = met("P4", doc + f"\nOut of scope: {claim}\n")
        assert not bad[0] and ("7" in bad[1] or "Northfield" in bad[1]), (claim, bad)
        assert met("P4", doc + f"\n- Assumption: {claim}\n")[0], claim
    secs = {"W-1": {"body": "## Users\n\nStudents and lab staff.\n\n- Assumption: Bookings open 7 days ahead."},
            "W-2": {"body": "## Requirements\n\n1. **R1**: A student can book a slot"}}
    md = j2.assemble_doc("LabSlot", secs, j2.PRD_SECTIONS)
    assert md.endswith("## Assumptions to confirm\n\n- Bookings open 7 days ahead.\n"), md
    assert j2.doc_assumptions(md) == ["Bookings open 7 days ahead."] and met("P4", md)[0]
    assert "Assumptions to confirm" not in j2.assemble_doc("LabSlot", {"W-2": secs["W-2"]}, j2.PRD_SECTIONS)
    p4 = next(c for c in j2.check_conditions(j2.load_conditions("prd.json"), md, goal, src, j2.PRD_SECTIONS)
              if c["id"] == "P4")
    assert p4["met"] and p4["assumptions"] == ["Bookings open 7 days ahead."], p4
    assert "quote each" in dict(j2.DOC_RUBRIC)["D-5"] or "quotes each" in dict(j2.DOC_RUBRIC)["D-5"]

    # a mock PRD run whose Out of scope builder adds an assumption: in conditions.json and the closing section
    os.environ["J2_OUTPUT"] = "prd"
    try:
        run = j2.J2Run()
        orig = run.call

        def call(fn, *a, **k):
            raw = orig(fn, *a, **k)
            if isinstance(raw, str) and "## Out of scope" in raw:
                head, _, tail = raw.rpartition("```")
                raw = head + "- Assumption: Clinics may want a printed ticket within 2 weeks.\n```" + tail
            return raw
        run.call = call
        say(run, PRD_BRIEF)
        q = run._fact_question({"needs_requester_fact": "A target for in-room wait", "item": None})["question"]
        assert "over what time" in q, q                                       # the window in the same question
    finally:
        os.environ.pop("J2_OUTPUT")
    assert run.outcome.startswith("Done: the PRD is workable"), run.outcome
    md = open(os.path.join(run.dir, "prd.md"), encoding="utf-8").read()
    assert md.endswith("## Assumptions to confirm\n\n- Clinics may want a printed ticket within 2 weeks.\n"), md[-300:]
    saved = json.load(open(os.path.join(run.dir, "conditions.json"), encoding="utf-8"))
    p4 = next(c for c in saved if c["id"] == "P4")
    assert p4["met"] and p4["assumptions"] == ["Clinics may want a printed ticket within 2 weeks."], p4
    assert next(c for c in saved if c["id"] == "P7")["met"], saved
    site = j2.J2Run()
    site.goal = {"done_when": ["x"]}
    q = site._fact_question({"needs_requester_fact": "A target for in-room wait", "item": None})["question"]
    assert "over what time" not in q, q                                       # websites: as before

    # short names: cut at a natural break, at most 12 words, "..." when cut mid-clause, never on a joining word
    texts = ["The Requirements section lists every must-have feature as R1, R2, ...",
             "Visitors can see every event of the three-day fest grouped by day with its start time and the room it is in",
             "A section about the beans and the farms they come from",
             "Opening hours are listed"]
    for t in texts:
        n = j2.short(t)
        words = n.rstrip(".").split()
        assert len(words) <= 12 and not (n.endswith("...") and words[-1].lower() in j2.JOINING), n
        first = re.split(r"[,:;(]| - ", t)[0].strip()
        assert n == first or (n.endswith("...") and first.startswith(n[:-3])), (t, n)
    assert j2.short(texts[1]).endswith("...") and j2.short(texts[3]) == texts[3]
    assert j2.short(texts[2]) == texts[2]                                     # 11 words: whole, no "..."
    names = j2.shorts([texts[1], texts[1].replace("room", "hall")])
    assert len(set(names)) == 2, names                                        # still unique within a list


def prd_recovery():
    """runs/j2-20261013-185235: heading targets are normalised ("## Problem" is "Problem"); a wrong target after three
    failures ends in CHECK_WRONG by code without asking Identity, and the corrected check re-runs on the same document;
    an unusable Adaptation reply twice gets a code-built revision and the run goes on; the PRD question lists
    sections only, says "document", never drops everything and never repeats with nothing changed."""
    assert j2.heading_target("## Problem") == j2.heading_target(" #Problem") == "Problem"
    doc = "# X\n\n## Problem\n\nBooking is done on paper and slots get double-booked.\n"
    for t in ("## Problem", "Problem"):
        assert j2.doc_check({"type": "heading-present", "target": t}, doc)["pass"], t
        assert j2.doc_check({"type": "section-not-empty", "target": t, "expect": 5}, doc)["pass"], t
    atom = j2.mock_doc_atom({"name": "Write Problem"})
    assert not j2.check_problems({**atom, "verification": [{"text": "x", "type": "heading-present", "target": "## Problem",
                                                            "expect": True}]}, "", j2.DOC_CHECK_TYPES)
    assert "no # characters" in j2.PRD_ATOM_RULES

    def prd(patch):
        """A mock PRD run whose AI calls pass through patch(fn, reply, n) (n: that function's call count)."""
        os.environ["J2_OUTPUT"] = "prd"
        try:
            run, seen = j2.J2Run(), {}
            orig = run.call

            def call(fn, *a, **k):
                seen[fn] = seen.get(fn, 0) + 1
                return patch(fn, orig(fn, *a, **k), seen[fn])
            run.call = call
            say(run, PRD_BRIEF)
            finish(run)
        finally:
            os.environ.pop("J2_OUTPUT")
        return run, seen

    def target(t):              # Adaptation's first plan writes the Problem heading checks with target t
        def patch(fn, r, n):
            if fn == "adaptation" and n == 1 and isinstance(r, dict):
                for w in r.get("workflows") or []:
                    for c in w["atoms"][0]["verification"] if "Problem" in w["name"] else []:
                        if c["type"] in j2.HEADING_CHECKS:
                            c["target"] = t
            return r
        return patch
    # (1) "## Problem" behaves exactly like "Problem", and the normalised form is stored
    run, seen = prd(target("## Problem"))
    assert run.outcome.startswith("Done: the PRD is workable"), run.outcome
    w1 = run.workflows[0]
    assert [c["target"] for c in w1["atoms"][0]["verification"] if c["type"] in j2.HEADING_CHECKS] == ["Problem"] * 2
    saved = json.load(open(os.path.join(run.dir, "workflows.json"), encoding="utf-8"))[0]
    assert all(c["target"] == "Problem" for c in saved["atoms"][0]["verification"] if c["type"] in j2.HEADING_CHECKS)
    recs = [r for r in run.receipts if r["atom"] == "W-1.a1" and r["type"] in j2.HEADING_CHECKS]
    assert recs and all(r["target"] == "Problem" and r["pass"] for r in recs) and seen["W-1"] == 1, recs

    # (2) a wrong target fails twice with the same evidence (no third attempt): code checks the obvious case, CHECK_WRONG, re-run on the same document
    run, seen = prd(target("Problem section"))
    assert run.outcome.startswith("Done: the PRD is workable"), run.outcome
    assert seen["W-1"] == 2, seen                      # no new build after the correction
    assert any(e["type"] == "check_wrong_default" for e in run.events), run.events
    note = next(m for m in run.chat if m["text"].startswith("Code checked the obvious case first"))
    assert "'Problem section' passes as the heading text 'Problem'" in note["text"], note["text"]
    re_run = [r for r in run.receipts if r["atom"] == "W-1.a1" and r.get("recheck")]
    assert re_run and all(r["pass"] for r in re_run) and {r["target"] for r in re_run if r["type"] in
                                                           j2.HEADING_CHECKS} == {"Problem"}, re_run
    assert run.workflows[0]["atoms"][0]["status"] == "passed"
    assert not any(m["who"] == "identity" and "fails its checks" in m["text"] for m in run.chat)
    call = next(c for c in calls_of(run) if "WORK ATOM FAILED" in c["prompt"])
    assert "the check's target looks wrong" in call["prompt"], call["prompt"][-600:]

    # (3) the build is really wrong (W-1 writes "## Background"), then Adaptation's reply is a bare "say" twice:
    # code builds the revision of W-1 itself and the run gets to Done
    blob = {"say": '{"say":"One proposal: a W-1 revision covering items 1 and 8, priced at 28. The builder writes'}

    def patch(fn, r, n):
        if fn == "W-1":
            return r.replace("## Problem", "## Background")
        if fn == "adaptation" and n in (3, 4):          # 1: the plan, 2: the failed atom, 3-4: the next proposals
            return dict(blob)
        return r
    run, seen = prd(patch)
    assert run.outcome.startswith("Done: the PRD is workable"), run.outcome
    assert run.workflows[0]["atoms"][0]["status"] == "escalated", run.workflows[0]["atoms"]
    retry = [c for c in calls_of(run) if "YOUR LAST REPLY COULD NOT BE USED" in c["prompt"]]
    assert len(retry) == 1 and '"revises": "W-1"' in retry[0]["prompt"], retry
    fix = next(w for w in run.workflows if w.get("revises") == "W-1")
    assert fix["name"] == "Revise W-1 Write Problem" and fix["cost"] == j2.RATES["fix"] and fix["status"] == "done", fix
    assert any(m["text"].startswith("Adaptation's reply couldn't be used twice") for m in run.chat)
    md = open(os.path.join(run.dir, "prd.md"), encoding="utf-8").read()
    assert "## Problem" in md and "## Background" not in md, md[:300]

    # (4) the PRD question: sections only, "document", never an option that drops everything, never a repeat
    run, _ = prd(lambda fn, r, n: r)
    crit, last = list(run.goal["done_when"]), run.audits[-1]
    for r in last["results"][1:]:                       # every section but Problem fails
        r.update(verdict="NEEDS_REVISION", note="missing", **{"pass": False})
    last["rubric"][0].update(verdict="NEEDS_REVISION", **{"pass": False})
    last["conditions"][0].update(met=False, evidence="missing")
    run.outcome, run.done = None, False
    run._identity_limit({"reason": "The same failure came back twice."})
    text, q = run.chat[-1]["text"], run.paused["questions"][0]
    rows = [ln[2:] for ln in text.split("\n") if ln.startswith("- ")]
    assert rows and all(re.fullmatch(r"the (Users|Goals|Requirements|Success metrics|Out of scope) section \(.+\)", x)
                        for x in rows), rows
    assert not re.search(r"\b(?:web)?site\b", text + json.dumps(q), re.I), text
    assert not re.search(r"\bP[1-8]\b|\bD-\d\b", text), text
    assert not any(o.startswith("Leave out") for o in q["options"]), q   # it would leave only Problem
    run.paused = None
    for r in last["results"][1:4]:
        r.update(verdict="PASS", **{"pass": True})
    run.last_limit = {"state": (len(run.sites), len(run.workflows), run.spent), "decision": "continue"}
    run._identity_limit({"reason": "The same failure came back twice."})
    q = run.paused["questions"][0]
    assert q["question"].startswith("The last round produced no new work"), q
    assert not any(o.startswith("Keep going") for o in q["options"]), q          # only options that change something
    assert any(o.startswith("Leave out the Success metrics section and the Out of scope section") for o in q["options"]), q


def prd_page():
    """The page: the output type is picked before the first message and shown; the preview renders prd.md; the .md
    and .docx download, for the live run and once it is saved."""
    import threading
    import urllib.error
    import urllib.request
    import j2_server
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    get = lambda path: urllib.request.urlopen(base + path).read()
    post = lambda path, body: urllib.request.urlopen(urllib.request.Request(
        base + path, json.dumps(body).encode(), {"Content-Type": "application/json"}))
    try:
        j2_server.RUN.reset()
        assert json.loads(post("/api/output", {"type": "prd"}).read())["output"] == "prd"
        st = json.loads(get("/api/state"))
        assert st["output"] == "prd" and st["output_label"] == "PRD document" and st["doc"], st["output"]
        html = get("/").decode()
        assert 'id="outsel"' in html and 'id="otype"' in html and 'id="doclinks"' in html and "/api/output" in html
        say(j2_server.RUN, PRD_BRIEF)
        try:
            post("/api/output", {"type": "website"})
            raise AssertionError("the output type changed after the run started")
        except urllib.error.HTTPError as e:
            assert e.code == 409
        st = json.loads(get("/api/state"))
        assert st["files"] == ["prd.md", "prd.docx"] and st["conditions"] and st["rubric"][0][0] == "D-1", st["files"]
        prev = get("/site/latest").decode()
        assert "<h2>Problem</h2>" in prev and "<table>" in prev and "<ol>" in prev, prev[:300]
        assert get("/file/prd.md").decode().startswith("# QuickQueue") and get("/file/prd.docx")[:2] == b"PK"
        if j2.j2_check.find_chrome():            # the Simple view in Chrome: plain words, document links
            text, _, _ = simple_text(base)
            bad = SIMPLE_FORBIDDEN.findall(text) + SIMPLE_FORBIDDEN_CASE.findall(text)
            assert not bad and "Your document team" in text and "Your document is ready" in text, (bad, text[-600:])
            assert "Download the Word file (.docx)" in text and "Download the markdown (.md)" in text, text[-400:]
            assert "Download the HTML file" not in text, text[-400:]
        name = os.path.basename(j2_server.RUN.dir)
        j2_server.RUN.set_output("website")
        j2_server.RUN.reset()                    # the PRD run is now a saved run
        sv = json.loads(get(f"/api/state?run={name}"))
        assert sv["output"] == "prd" and sv["versions"] == 6 and sv["files"] == ["prd.md", "prd.docx"], sv["output"]
        assert "<h2>Requirements</h2>" in get(f"/site/latest?run={name}").decode()
        assert get(f"/file/prd.docx?run={name}")[:2] == b"PK"
        try:
            get(f"/file/goal.json?run={name}")
            raise AssertionError("/file/ served a file that isn't the document")
        except urllib.error.HTTPError as e:
            assert e.code == 404
    finally:
        srv.shutdown()
        j2_server.RUN.reset()


def website_conditions():
    """conditions/website.json: a page that breaks each of N1 to N8 is reported not met, with evidence; a plan with no
    workflow for the nav is flagged before building; workable() needs every condition; Identity's question for a
    fact shows no "None" and no internal id."""
    conds = j2.load_conditions("website.json")
    assert [c["id"] for c in conds] == [f"N{i}" for i in range(1, 9)], conds
    goal = {"site_name": "Little Fig", "facts": ["Little Fig", "+91 98765 43210"], "to_be_confirmed": []}
    src = j2.squash("Little Fig is open 8 to 10. Call +91 98765 43210 or write to hello@littlefig.in")
    good = ('<!doctype html><html><head><title>Little Fig</title></head><body><header id="hero"><nav>'
            '<a href="#menu">Menu</a> <a href="#visit">Visit</a></nav><h1>Little Fig</h1>'
            '<a class="j2-btn" href="#visit">Visit us</a></header><section id="menu"><h2>Menu</h2><p>Open 8 to 10.</p>'
            '</section><section id="visit"><h2>Visit</h2><a href="tel:+919876543210">+91 98765 43210</a> '
            '<a href="https://wa.me/919876543210">WhatsApp</a><form action="mailto:hello@littlefig.in">'
            '<input name="n"><button type="submit">Send</button></form></section></body></html>')
    clean = {"available": True, "js_errors": [], "phone": {"width": 360, "overflow_px": 0},
             "desktop": {"width": 1280, "overflow_px": 0}}

    def check(page, mech=clean, g=goal):
        return {c["id"]: c for c in j2.check_conditions(conds, " ".join(j2.PageScan(page).visible), g, src, (),
                                                         page, mech)}
    ok = check(good)
    assert all(c["met"] for c in ok.values()), ok
    breaks = {"N1": (good.replace("</h2><p>Open", "<p>Open"), clean, goal),
              "N2": (good, {**clean, "js_errors": ["ReferenceError: menu is not defined"]}, goal),
              "N3": (good, {**clean, "phone": {"width": 360, "overflow_px": 40}}, goal),
              "N4": (good.replace("https://wa.me/919876543210", "https://wa.me/+91 98765"), clean, goal),
              "N5": (good, clean, {**goal, "to_be_confirmed": ["the opening hours"]}),
              "N6": (good.replace("<p>Open 8 to 10.</p>", "<p>Open 8 to 10. Free delivery over 500.</p>"), clean, goal),
              "N7": (good.replace(' action="mailto:hello@littlefig.in"', ""), clean, goal),
              "N8": (good.replace('<a href="#menu">Menu</a> ', ""), clean, goal)}
    for cid, (page, mech, g) in breaks.items():
        r = check(page, mech, g)[cid]
        assert not r["met"] and r["evidence"] and r["evidence"] != ok[cid]["evidence"], (cid, r)
    assert "500" in check(*breaks["N6"])["N6"]["evidence"] and "#menu" in check(*breaks["N8"])["N8"]["evidence"]
    assert "wa.me/+91" in check(*breaks["N4"])["N4"]["evidence"]
    r = check(good.replace('<a class="j2-btn" href="#visit">Visit us</a>', ""))["N8"]      # no main action in the hero
    assert not r["met"] and "hero" in r["evidence"], r
    r = check(good, {"available": False, "note": "Chrome not found"})
    assert not r["N2"]["met"] and r["N2"]["evidence"].startswith("could not run"), r["N2"]
    assert not check(good, clean, {**goal, "facts": ["Koramangala"]})["N5"]["met"]
    assert ok["N5"]["evidence"].startswith("all 2 brief fact(s) appear"), ok["N5"]
    r = check(good, clean, {**goal, "facts": []})["N5"]          # no facts list: the site name only, and it says so
    assert r["met"] and r["evidence"].startswith("no facts list; the site name 'Little Fig' appears"), r

    # workable(): Audit accepts and the render check is clean; every condition met, or not workable
    row = {"id": "F-1", "item": 1, "criterion": "Menu", "verdict": "PASS", "pass": True}
    audit = {"results": [row], "rubric": [], "mechanical": fake_check(good), "conditions": list(ok.values())}
    assert j2.workable(audit) == [], j2.workable(audit)
    bad = {**audit, "conditions": list(check(*breaks["N8"]).values())}
    assert j2.workable(bad) and all("necessary condition N8 not met" in x for x in j2.workable(bad)), j2.workable(bad)

    # the plan check: no workflow for the nav is flagged before anything is built
    run = j2.J2Run()
    assert [c["id"] for c in run._plan_gaps([{"name": "Menu section", "goal_gap": "the menu"}])] == ["N8"]
    assert run._plan_gaps([{"name": "Header", "goal_gap": "a nav bar to every section and a main action"}]) == []
    orig = run.call

    def call(fn, fmt, task, mock, **k):
        r = orig(fn, fmt, task, mock, **k)
        if fn == "adaptation" and isinstance(r, dict) and "FLAGGED" not in task and r.get("workflows"):
            r["workflows"] = json.loads(json.dumps(r["workflows"]).replace(        # the first plan names no nav
                " with a nav to every section and a main action in the hero", ""))
        return r
    run.call = call
    say(run, BREW_LAB)
    finish(run)
    i = next(n for n, m in enumerate(run.chat) if m["text"].startswith("Validator flagged the plan before building"))
    assert "N8" in run.chat[i]["text"] and not any(m["who"].startswith("W-") for m in run.chat[:i]), run.chat[i]
    retry = [c for c in calls_of(run) if c["fn"] == "adaptation" and "THE VALIDATOR FLAGGED YOUR PLAN" in c["prompt"]]
    assert len(retry) == 1 and "N8" in retry[0]["prompt"], retry
    assert "NECESSARY CONDITIONS" in retry[0]["prompt"] and "N8: A nav reaches every section" in retry[0]["prompt"]
    assert run.outcome.startswith("Done"), run.outcome
    s = summary(run)
    assert s["necessary_conditions"]["met"] == s["necessary_conditions"]["total"] == 8, s["necessary_conditions"]
    saved = json.load(open(os.path.join(run.dir, "conditions.json"), encoding="utf-8"))
    assert [c["id"] for c in saved] == [f"N{i}" for i in range(1, 9)] and all(c["met"] for c in saved), saved
    assert run.snapshot()["conditions"] == run.audits[-1]["conditions"]

    # Identity's question for a fact only the Requester knows: no "None", no internal ids, no agent-facing wording
    run = j2.J2Run()
    run.goal = {"done_when": ["Opening hours are listed", "Bookings never clash"]}
    fact = ("A measurable target for double-booked slots. The brief gives none, so any number would be invented. "
            "Alternatives: Requester supplies a target, or G2 is stated as covered by the requirement R5 instead of a "
            "metric.")
    for item, ref in ((None, None), (2, "(checklist item 2: Bookings never clash)")):
        q = run._fact_question({"needs_requester_fact": fact, "item": item})["question"]
        assert "None" not in q and not j2.INTERNAL_ID.search(q) and "Alternatives" not in q, q
        assert q.startswith("A measurable target for double-booked slots. The brief gives none"), q
        assert (ref in q) if ref else "checklist item" not in q, q
    q = run._fact_question({"needs_requester_fact": "Confirm option B (G2 covered by R5, no metric)", "item": None})
    assert q["question"].startswith("Confirm option B. If you can't"), q
    run._ask_user("Nearly done. F-3 needs one number.", [q, run._fact_question({"needs_requester_fact": fact,
                                                                                  "item": 1})], "fact")
    text = run.chat[-1]["text"]
    assert "None" not in text and not j2.INTERNAL_ID.search(text) and "Alternatives" not in text and ".." not in text, text
    assert j2.requester_words("Meet in Room G2.", j2.squash("the venue is Room G2")) == "Meet in Room G2."


def behavior_steps():
    """Behaviour items: the step interpreter with a fake driver (no Chrome), the validator's rules for behaviour
    steps, one re-ask on a rejected item, and the mock writing one behaviour item per atom with an action goal."""
    ck, run_steps = j2.j2_check, j2.j2_check.run_steps
    # the interpreter: passes, and each failure named in plain words after the last action
    page = lambda: FakeDriver({"#tshirt": {}, "#size": {"visible": False}, "#calc-go": {}, "#total": {"text": "Rs 0"},
                               "#countdown": {"text": "10 days"}, "#contrast-toggle": {}, "body": {"style": {
                                   "background-color": "rgb(255, 255, 255)"}}, "#schedule": {},
                               "#link-schedule": {"classes": set()}, "#broken": {"error": "no option XXL"}},
                              {("set", "#tshirt"): lambda e, o: e["#size"].update(visible=o["val"] == "yes"),
                               ("click", "#calc-go"): lambda e, o: e["#total"].update(text="Total: Rs 950"),
                               "wait": lambda e, o: e["#countdown"].update(text="9 days"),
                               ("click", "#contrast-toggle"): lambda e, o: e["body"]["style"].update(
                                   {"background-color": "rgb(0, 0, 0)"}),
                               ("scroll", "#schedule"): lambda e, o: e["#link-schedule"]["classes"].add("active")})
    ok = ["expect-hidden #size", "set #tshirt yes", "expect-visible #size", "click #calc-go",
          "expect-number #total 500*0.9+2*250", "wait 2", "expect-changed #countdown", "click #contrast-toggle"]
    r = run_steps(ok, page())
    assert r["pass"] is True and r["evidence"] == "all 8 step(s) behaved as expected", r
    r = run_steps(["expect-style-changed body background-color"][:0] + ["click #contrast-toggle",
                  "expect-style-changed body background-color", "scroll-to #schedule",
                  "expect-class #link-schedule active", "expect-text #countdown 10 DAYS"], page())
    assert r["pass"] is True, r
    for steps, why in ((["set #tshirt no", "expect-visible #size"], "after choosing tshirt no, size was still hidden"),
                       (["set #tshirt yes", "expect-hidden #size"], "after choosing tshirt yes, size was still visible"),
                       (["click #calc-go", "expect-number #total 500+2*250"],
                        "after clicking calc go, total showed 'Total: Rs 950', expected 1000 (500+2*250)"),
                       (["expect-text #total Rs 950"], "on load, total showed 'Rs 0', not 'Rs 950'"),
                       (["click #calc-go", "wait 1", "expect-changed #total"],
                        "after waiting 1s, total still showed 'Total: Rs 950'"),
                       (["click #calc-go", "expect-style-changed body background-color"],
                        "after clicking calc go, the background-color of body stayed rgb(255, 255, 255)"),
                       (["expect-class #link-schedule active"], "on load, link schedule did not have the class active"),
                       (["click #buy-now"], "on load, nothing on the page matches buy now (step 1: click)"),
                       (["set #broken XXL"], "on load, step 1 (set broken) could not be done: no option XXL")):
        r = run_steps(steps, page())
        assert r["pass"] is False and r["evidence"] == why and r["actual"]["step"] == steps[-1], (steps, r)
    two = FakeDriver({"#calc-go": {}, "#total": {"text": "2 runners: Rs 950"}})
    r = run_steps(["click #calc-go", "expect-number #total 500*0.9+2*250"], two)
    assert r["pass"] is False and r["evidence"].endswith("could not tell which number is the total"), r
    r = run_steps(["click #calc-go"], lambda ops: {"available": False, "note": "Chrome not found"})
    assert r["pass"] is None and r["evidence"].startswith("could not run"), r
    assert ck.safe_number("(500*0.9)+2*250") == (950.0, [500, 0.9, 2, 250])
    for bad in ("__import__('os')", "price*2", "2**8", "1;2"):
        try:
            ck.safe_number(bad)
            raise AssertionError(bad)
        except ValueError:
            pass

    # the validator: selectors the atom creates, numbers from the charter, no today's date, at most 8 steps
    src = j2.squash("Race fee Rs 500. 10% off race fees for registrations up to 31 December 2026. T-shirt Rs 250 each.")
    how = {"kind": "AI", "agent": "Builder", "steps": "Build the calculator #calc with #shirts, #reg-date, #calc-go and "
           "#total; #tshirt and #size; #countdown"}
    goal = "A fee calculator shows the total, with 10% off up to 31 December 2026"

    def probs(steps, g=goal, h=how, kind="behavior"):
        a = {"goal": g, "how": h, "constraints": [], "context": ["charter item 1"], "verification": [],
             "expected_output": [{"text": "The total for one 10K runner and 2 T-shirts", "kind": kind, "steps": steps}]}
        return j2.expected_problems(a, j2.CHECK_TYPES) or j2.check_problems(j2.cover_expected(a), src)
    good = ["set #reg-date 2026-12-01", "set #shirts 2", "click #calc-go", "expect-number #total 500*0.9+2*250"]
    assert probs(good) == [], probs(good)
    for steps, part in (([*good[:3], "expect-number #total 950"], "uses 950, not in the charter"),
                        ([*good[:3], "expect-number #total 500*0.85"], "uses 0.85"),
                        ([*good[:3], "expect-number #total price*2"], "isn't arithmetic over numbers"),
                        ([*good[:3], "expect-visible #buy-now"], "'#buy-now', which the atom's how or constraints"),
                        ([*good[:3], "expect-text #total hello@x.com"], None),
                        (good[1:], "depends on a date"),
                        (good + ["wait 1"] * 5, "has 9 steps; at most 8"),
                        (["wait 5", "expect-changed #countdown"], "wait is at most 3 seconds"),
                        (["hover #total"], "starts with 'hover'"),
                        (["expect-changed #countdown"], "expect-changed needs a wait step before it"),
                        ([], "isn't an object with text, kind and target (steps for a behavior)")):
        got = probs(steps)
        assert (got == [] if part is None else len(got) == 1 and part in got[0]), (steps, got)
    assert probs(good)[:0] == [] and "expected output item 1 behavior step 4" in probs(
        [*good[:3], "expect-number #total 950"])[0]
    # an action goal needs a behaviour item; a document has no behaviour check
    act = probs(["click #calc-go"], kind="facts")
    assert any("describes an action ('calculator')" in x for x in probs(["x"], kind=None) + [
        *j2.check_problems({"goal": "A calculator shows the total", "how": how, "expected_output": [],
                            "verification": [{"text": "x", "type": "file-exists", "target": "index.html"}]}, src)]), act
    for g in ("Shows the menu", "Update the menu section", "The menu section is updated with prices"):
        assert not j2.check_problems({"goal": g, "how": how, "expected_output": [], "verification": []}, src), g
    for g in ("Recalculates the total on change", "A self-updating countdown", "Reveals the size field",
              "Shows or hides the guardian field", "T-shirt size shown only when T-shirt is yes",
              "guardian name required only when the runner is under 18", "Guardian name required only if under 18",
              "The size field is hidden unless T-shirt is yes", "Hide the size field"):
        assert any("describes an action" in x for x in j2.check_problems(
            {"goal": g, "how": how, "expected_output": [], "verification": []}, src)), g
    assert "which no check can verify" in j2.expected_problems(
        {"expected_output": [{"text": "x", "kind": "behavior", "steps": ["click #a"]}]}, j2.DOC_CHECK_TYPES)[0]

    # one re-ask: an invented total is rejected and Adaptation sends the reply again; the run still finishes
    run, hit = j2.J2Run(), []
    orig = run.call

    def call(fn, fmt, task, mock, **k):
        r = orig(fn, fmt, task, mock, **k)
        if fn == "adaptation" and isinstance(r, dict) and r.get("workflows") and not run.workflows and not hit:
            hit.append(1)
            a = r["workflows"][0]["atoms"][0]
            a["how"] = {**a["how"], "steps": a["how"]["steps"] + "; button #calc-go shows #total"}
            a["expected_output"].append({"text": "The total", "kind": "behavior",
                                         "steps": ["click #calc-go", "expect-number #total 4321"]})
        return r
    run.call = call
    say(run, BREW_LAB)
    finish(run)
    assert "uses 4321, not in the charter" in run.rejections[0]["reason"], run.rejections
    assert len([c for c in calls_of(run) if "THE VALIDATOR REJECTED" in c["prompt"]]) == 1
    assert run.outcome.startswith("Done"), run.outcome

    # mock mode: one behaviour item per atom whose goal describes an action, built and met
    run = j2.J2Run()
    say(run, BREW_LAB.replace("and where to find us", "where to find us, and a price calculator for an order"))
    finish(run)
    atoms = [a for w in run.workflows for a in w.get("atoms") or []]
    acts = [a for a in atoms if j2.ACTION_GOAL.search(a["goal"])]
    assert acts and all(len([x for x in a["expected_output"] if x["kind"] == "behavior"]) == 1 for a in acts), [(w["name"], w.get("completion_check"), [a["goal"] for a in w["atoms"]]) for w in run.workflows]
    assert not any(x["kind"] == "behavior" for a in atoms if a not in acts for x in a["expected_output"])
    if not BROWSER or j2.j2_check.find_chrome():
        assert all(x["met"] is True for a in acts for x in a["expected_output"]), [a["expected_output"] for a in acts]
    assert run.outcome.startswith("Done"), run.outcome
    assert "behavior (proves the page works" in next(c["prompt"] for c in calls_of(run) if c["fn"] == "adaptation")


def charter_facts():
    """J2.a facts: a fact not in the brief or the answers (paraphrased or invented) is rejected and Identity is asked
    once more; still there after that, code drops it. The mock charter's facts are exact."""
    def run_with(times):
        run, hit = j2.J2Run(), []
        orig = run.call

        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            if fn == "identity" and isinstance(r, dict) and isinstance(r.get("goal"), dict):
                hit.append(task)
            if len(hit) <= times and hit and hit[-1] is task and fn == "identity":
                r["goal"]["facts"] = list(r["goal"].get("facts") or []) + ["Bengaluru's best coffee", "Indiranagar"]
            return r
        run.call = call
        say(run, BREW_LAB)
        finish(run)
        return run, hit
    run, hit = run_with(1)
    assert len(hit) == 2 and "THE VALIDATOR REJECTED YOUR CHARTER" in hit[1], hit
    rej = [x for x in run.rejections if x["workflow"] == "charter"]
    assert len(rej) == 1 and "[\"Bengaluru's best coffee\"]" in rej[0]["reason"], rej
    assert run.goal["facts"] == ["Brew Lab"] and run.outcome.startswith("Done"), (run.goal["facts"], run.outcome)
    run, hit = run_with(2)                                   # still invented after the re-ask: dropped, not asked again
    assert len(hit) == 2 and len([x for x in run.rejections if x["workflow"] == "charter"]) == 2
    assert run.goal["facts"] == ["Brew Lab", "Indiranagar"] and run.outcome.startswith("Done"), run.goal["facts"]
    assert "List in facts every concrete fact" in next(c["prompt"] for c in calls_of(run) if c["fn"] == "identity")


BEHAVIOR_PAGES = {
    "calculator": ("""<select id="race"><option value="300">5K</option><option value="500">10K</option></select>
<input id="shirts" type="number" value="0"><input id="reg-date" type="date" value="2027-01-20">
<button id="calc-go">Total</button><p id="total">Total: Rs 0</p>
<script>document.getElementById('calc-go').onclick = () => { const early = document.getElementById('reg-date').value <= '2026-12-31';
  const fee = +document.getElementById('race').value * (early && %s ? 0.9 : 1);
  document.getElementById('total').textContent = 'Total: Rs ' + (fee + 250 * document.getElementById('shirts').value); };</script>""",
                   ["set #race 10K", "set #shirts 2", "set #reg-date 2026-12-01", "click #calc-go",
                    "expect-number #total 500*0.9+2*250"], "after clicking calc go, total showed 'Total: Rs 1000'"),
    "conditional field": ("""<select id="tshirt"><option>no</option><option>yes</option></select>
<label id="size" hidden>Size <select><option>S</option><option>M</option></select></label>
<script>document.getElementById('tshirt').onchange = e => { if (%s) document.getElementById('size').hidden = e.target.value !== 'yes'; };</script>""",
                          ["expect-hidden #size", "set #tshirt yes", "expect-visible #size", "set #tshirt no",
                           "expect-hidden #size"], "after choosing tshirt yes, size was still hidden"),
    "required only when": ("""<input id="dob" type="date"><input id="guardian"><button id="reg-go">Register</button>
<p id="guardian-error" hidden>Guardian name is required for runners under 18</p>
<script>document.getElementById('reg-go').onclick = () => { const minor = document.getElementById('dob').value > '2009-02-14';
  document.getElementById('guardian-error').hidden = !(%s && minor && !document.getElementById('guardian').value); };</script>""",
                           ["set #dob 2012-05-01", "click #reg-go", "expect-visible #guardian-error",
                            "set #dob 1990-01-01", "click #reg-go", "expect-hidden #guardian-error"],
                           "after clicking reg go, guardian error was still hidden"),
    "countdown": ("""<p id="countdown">Race starts in 100 seconds</p>
<script>let n = 100; setInterval(() => { if (%s) document.getElementById('countdown').textContent = 'Race starts in ' + (--n) + ' seconds'; }, 1000);</script>""",
                  ["wait 2", "expect-changed #countdown"], "after waiting 2s, countdown still showed"),
    "active menu link": ("""<nav style="position:fixed;top:0"><a class="menu-link" id="link-races" href="#races">Races</a>
<a class="menu-link" id="link-schedule" href="#schedule">Schedule</a></nav>
<section id="races" style="height:1500px">Races</section><section id="schedule" style="height:1500px">Schedule</section>
<script>addEventListener('scroll', () => { if (!%s) return; for (const s of document.querySelectorAll('section')) {
  const r = s.getBoundingClientRect(); document.getElementById('link-' + s.id).classList.toggle('active', r.top <= 100 && r.bottom > 100); } });</script>""",
                         ["expect-class #link-races active"][:0] + ["scroll-to #schedule",
                                                                     "expect-class #link-schedule active"],
                         "after scrolling to schedule, link schedule did not have the class active"),
    "contrast toggle": ("""<style>body{background:#fff} #menu{background:#f4efe6} body.hc{background:#000} body.hc #menu{background:#111}</style>
<button id="contrast-toggle">High contrast</button><section id="menu">Menu</section>
<script>document.getElementById('contrast-toggle').onclick = () => { if (%s) document.body.classList.toggle('hc'); };</script>""",
                        ["click #contrast-toggle", "expect-style-changed body background-color",
                         "expect-style-changed #menu background-color"],
                        "after clicking contrast toggle, the background-color of body stayed"),
}


def behavior_browser():
    """The point-5 cases in headless Chrome: each page works (true) and is broken (false); the broken one fails at
    the step that shows it, in plain words. Each item also passes the validator."""
    if not BROWSER or not REAL_FIND():
        print("behaviour browser tests skipped: Chrome not found or --fast without --browser")
        return
    src = j2.squash("Race fee Rs 500. 10% off race fees up to 31 December 2026. T-shirt Rs 250.")
    for name, (body, steps, why) in BEHAVIOR_PAGES.items():
        how = {"kind": "AI", "agent": "Builder", "steps": "Creates " + " ".join(
            s.split()[1] for s in steps) + " and the body class hc"}
        a = {"goal": f"The {name} works", "how": how, "constraints": [], "context": ["charter item 1"],
             "verification": [], "expected_output": [{"text": name, "kind": "behavior", "steps": steps}]}
        assert not j2.check_problems(j2.cover_expected(a), src), (name, j2.check_problems(a, src))
        for works in (True, False):
            html = f"<!doctype html><html><body>{body % ('true' if works else 'false')}</body></html>"
            r = j2.run_check(a["verification"][0], "", html)
            assert r["pass"] is works and (works or r["evidence"].startswith(why)), (name, works, r)
    print("behaviour steps in Chrome: calculator total (early-bird date set by a control), conditional field, "
          "required only when, countdown, active menu link after scroll, contrast toggle on body and a section ok")


def run_20261015_fixes():
    """runs/j2-20261015-114723: (1) the prompt wrapper shows the validator's rules and one atom that passes, (2) the
    same failure twice goes to Adaptation after attempt 2, naming the owner atom, (3) a builder's human_output."""
    # 1: a mock Adaptation that copies the wrapper's example registers on the first submission, no rejection
    text = j2.J2Run.ATOM_RULES
    assert "THE VALIDATOR REQUIRES" in text and "goal_gap" in text and "completion_check" in text
    ex = json.loads(text.split("passes):\n", 1)[1])
    run = j2.J2Run()
    run.budget, run.goal = 500, {"done_when": ["A contact form is on the page"]}
    assert not run._atom_problems(ex), run._atom_problems(ex)
    reg, refused, _ = run._register([{**ex, "envelope": {"money_limit": 20}}])
    assert reg and not refused and not run.rejections, (refused, run.rejections)

    # 2: W-1 owns #register; W-2's check on #register fails with the same evidence twice: no attempt 3
    run = j2.J2Run()
    run.budget, run.goal = 500, {"done_when": ["a", "b"]}
    run.sections = {"W-1": {"owner": "W-1 Builder", "head": "", "body": '<section id="register">x</section>',
                            "workflow": "Placeholder"}}
    w1 = {"id": "W-1", "section": "W-1", "atoms": [{"id": "W-1.a1"}], "status": "done"}
    run.workflows, run.style = [w1], dict(j2.MOCK_STYLE)
    atom = {**ATOM[0], "verification": [{"text": "One register section on the page", "type": "grep-count", "target": "register",
                                          "expect": 1}]}
    atom["how"] = {**atom["how"], "steps": atom["how"]["steps"] + "; create section #register"}
    reg, refused, _ = run._register([{"name": "Registration", "goal_gap": "x", "outputs": ["#register"],
                                      "completion_check": "z", "atoms": [atom], "parts": ["fix"]}])
    assert reg, (refused, run.rejections)
    wid, orig, calls = reg[0]["id"], run.call, []
    run.sections[wid] = {"owner": f"{wid} Builder", "head": "", "body": "", "workflow": "Registration"}

    def call(fn, fmt, task, mock, **k):
        if fn == f"{wid}":
            calls.append(fn)
            body = '<section id="register">mine</section>'
            return ('{"say": "@Coordination written", "human_output": "I wrote the registration section.", '
                    '"atom_human_output": "This adds the sign-up part of the page.", '
                    '"next": "Coordination"}\n```section\n' + body + "\n```")
        return orig(fn, fmt, task, mock, **k)
    run.call = call
    run._turn_workflow_agent({"wid": wid})
    assert len(calls) == 2 and run.turn == "adaptation" and run.handoff["kind"] == "atom_failed", (calls, run.turn, [m["text"][:200] for m in run.chat][-6:])
    assert run.handoff["attempt"] == 2 and "W-1" in run.handoff["owner_note"] and "register" in run.handoff["owner_note"], \
        run.handoff
    assert any("same evidence twice" in m["text"] for m in run.chat if m["who"] == "system")

    # 3: "done." alone is rejected and asked for once more; the next speaker is set by code even when missing
    run = j2.J2Run()
    run.budget, run.goal = 500, {"done_when": ["a"]}
    reg, _, _ = run._register([{"name": "Hero", "goal_gap": "x", "outputs": ["y"], "completion_check": "z",
                                "atoms": ATOM, "parts": ["fix"]}])
    wid, orig, asked = reg[0]["id"], run.call, []
    run.style = dict(j2.MOCK_STYLE)

    def call3(fn, fmt, task, mock, **k):
        if fn == wid and "human_output" not in fmt.split("}")[0][:40]:
            return '{"say": "@Coordination done.", "human_output": "done.", "atom_human_output": "This adds a short intro for visitors."}\n```section\n<section id="x">x</section>\n```'
        if fn == wid:
            asked.append(task)
            return {"human_output": "I built the hero section with a headline and a button."}
        return orig(fn, fmt, task, mock, **k)
    run.call = call3
    run._turn_workflow_agent({"wid": wid})
    last = next(m for m in reversed(run.chat) if m["who"] == wid)
    assert len(asked) == 1 and "headline and a button" in last["text"] and run.turn == "coordination", (asked, last)
    assert not any("Route check" in m["text"] for m in run.chat), "route check fired"
    assert j2.human_problem("done.") and j2.human_problem("a b c d e. f. g. h.") and not j2.human_problem("I built it all today.")
    print("runs/j2-20261015-114723 fixes: wrapper example registers first try, same failure twice escalates after "
          "attempt 2 naming the owner atom, human_output regenerated once and next speaker set by code ok")


def calls_file(run):
    return [json.loads(l) for l in open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8")]


def atom_outputs():
    """A: every atom that passed carries a human_output (plain, no ids) and a code-built agent_output, both tagged with
    atom id and attempt, stored in workflows.json and outputs.jsonl; no atom is done without both."""
    for bad in ("done.", "This adds #hero for visitors.", "It builds W-1.a1 for the team now.", "It costs 3 units to build.",
                "One. Two. Three sentences here today."):
        assert j2.atom_human_problem(bad), bad
    assert not j2.atom_human_problem("This adds a hero part that visitors can read and use.")
    run = j2.J2Run()
    say(run, BREW_LAB)
    assert run.outcome.startswith("Done"), run.outcome
    passed = [(w, a) for w in run.workflows for a in w["atoms"] if a.get("status") == "passed"]
    assert passed, [a.get("status") for w in run.workflows for a in w["atoms"]]
    for w, a in passed:
        h, ag = a["human_output"], a["agent_output"]
        assert h["atom"] == ag["atom"] == a["id"] and h["attempt"] == ag["attempt"] >= 1 and not j2.atom_done_problem(a), a
        assert not j2.atom_human_problem(h["text"]), h
        assert ag["receipt"] == f"{a['id']}:{ag['attempt']}" and ag["section"] == w["section"] and ag["files"], ag
        body = run.sections[w["section"]]["body"]
        assert all(f'id="{i}"' in body for i in ag["ids"]) and all(sel[0] in "#." for sel in ag["selectors"]), ag
        assert ag["checks"] and all("pass" in c for c in ag["checks"]), ag
    rows = [json.loads(l) for l in open(os.path.join(run.dir, "outputs.jsonl"), encoding="utf-8")]
    assert {r["atom"] for r in rows} >= {a["id"] for _, a in passed}
    saved = json.load(open(os.path.join(run.dir, "workflows.json"), encoding="utf-8"))
    assert all("human_output" in x["atoms"][0] and "agent_output" in x["atoms"][0] for x in saved
               if x["atoms"][0].get("status") == "passed"), "run record lacks outputs"
    # the gate: without a human_output a passed atom becomes failed
    w, a = passed[0]
    del a["human_output"], a["agent_output"]
    assert j2.atom_done_problem(a) == ["human_output", "agent_output"]
    run._seal_atoms(w, [a], 1)
    assert a["status"] == "failed", a
    # a builder reply whose atom_human_output holds an id is asked once more, then code writes the sentence
    run = j2.J2Run()
    orig, small = run.call, []

    def call(fn, fmt, task, mock, **k):
        if fn.startswith("W-") and fmt.startswith('{"atom_human_output"'):
            small.append(fn)
            return {"atom_human_output": "It makes #hero for W-1.a1."}
        r = orig(fn, fmt, task, mock, **k)
        return re.sub(r'"atom_human_output": "[^"]*"', '"atom_human_output": "Shows #hero."', r) if isinstance(r, str) else r
    run.call = call
    say(run, BREW_LAB)
    w1 = run.workflows[0]["atoms"][0]
    assert small and not j2.atom_human_problem(w1["human_output"]["text"]) and "#" not in w1["human_output"]["text"], \
        (small, w1["human_output"])
    print("atoms: human_output (plain, no ids, one retry then code's sentence) and code-built agent_output on every "
          "passed atom, stored in workflows.json and outputs.jsonl, no done atom without both ok")


def slim_prompts():
    """C: J2_SLIM=1 sends builders their own section plus a manifest, and slims finished atoms in the shared state; the
    old path (J2_SLIM unset) still sends the whole page. Both end Done; the slim prompts are shorter."""
    sizes = {}
    for mode in ("0", "1"):
        os.environ["J2_SLIM"] = mode
        try:
            run = j2.J2Run()
            say(run, BREW_LAB)
            assert run.outcome.startswith("Done"), (mode, run.outcome)
            rows = calls_file(run)
            builders = [r for r in rows if r["fn"].startswith("W-")]
            later = [r for r in builders if r["fn"] != "W-1"]
            assert later
            if mode == "1":
                assert all("THE WHOLE SITE FILE" not in r["prompt"] and "OTHER SECTIONS" in r["prompt"] for r in later)
                st = [json.loads(r["prompt"].split("STATE:\n", 1)[1].split("\n\nCHAT:", 1)[0]) for r in rows[-3:]]
                slimmed = [a for s in st for w in s["workflows"] for a in w.get("atoms") or [] if a.get("status") == "passed"]
                assert slimmed and all(set(a) == {"id", "status", "agent_output"} for a in slimmed), slimmed[:1]
                assert "audit" in {r["fn"] for r in rows}          # Audit still ran on the full page
            else:
                assert all("THE WHOLE SITE FILE" in r["prompt"] and "OTHER SECTIONS" not in r["prompt"] for r in later)
            sizes[mode] = sum(len(r["prompt"]) for r in rows)
        finally:
            os.environ.pop("J2_SLIM", None)
    assert sizes["1"] < sizes["0"], sizes
    print(f"slim: both paths Done; prompt characters {sizes['0']:,} (old) -> {sizes['1']:,} (J2_SLIM=1) in the mock run")


def render_node(state, previews=None, html=None, docs=None):
    """j2_ui_render.js: the page's script against a stub DOM (no browser). Returns {simple, sprog, cards}."""
    import subprocess
    d = tempfile.mkdtemp(prefix="j2-render-")
    paths = []
    for name, obj in (("state.json", state), ("previews.json", previews or {}), ("docs.json", docs or {"docs": {}})):
        paths.append(os.path.join(d, name))
        json.dump(obj, open(paths[-1], "w"), default=str)
    here = os.path.dirname(os.path.abspath(__file__))
    r = subprocess.run(["node", os.path.join(here, "j2_ui_render.js"), html or os.path.join(here, "j2_ui.html"), *paths],
                       capture_output=True, text=True, timeout=60)
    assert r.returncode == 0, r.stderr[-800:]
    return json.loads(r.stdout)


def atom_cards_render():
    """B, browser-free: cards rendered from a stored run have the preview, eight blocks in three groups and three
    links; long blocks are collapsed; the preview source is the section in the saved page; Simple view is unchanged."""
    import subprocess
    import html as htmllib
    import shutil
    if not shutil.which("node"):
        print("atom cards render test skipped: node not found")
        return
    run = j2.J2Run()
    say(run, BREW_LAB)
    assert run.outcome.startswith("Done"), run.outcome
    page = open(os.path.join(run.dir, "site", "index.html"), encoding="utf-8").read()
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)          # the stored run, as the page reads it
    atoms = [(w, a) for w in snap["workflows"] for a in w["atoms"] if a.get("agent_output")]
    anchored = [a for _, a in atoms if a["agent_output"]["anchor"]]
    plain = [a for _, a in atoms if not a["agent_output"]["anchor"]]
    assert anchored, "the mock run builds sections with ids"
    previews = {a["id"]: j2.preview_doc(page, a["agent_output"]["anchor"]) for a in anchored}
    out = render_node(snap, previews)
    for w, a in atoms:
        card = out["cards"][w["id"]]
        card = card[card.index(f'data-key="atom-{a["id"]}"'):]
        card = card[:card.index("Run details")]
        blocks = re.findall(r'data-block="([a-z_]+)"', card)
        assert blocks == ["goal", "how", "constraints", "context", "expected_output", "verification", "human_output",
                          "agent_output", "result"], blocks
        assert all(f"<b>{x}</b>" in card for x in ("Notes", "Workflow output", "Result")) and "<b>Human output</b>" not in card
        assert re.findall(r'<h4 class="agroup">([^<]*)<', card) == ["1. Particles", "2. Defined output and verification",
                                                                      "3. Outputs"]
        for link in ("site", "agent", "evidence"):
            assert (f'data-link="{link}"' in card) == (link != "site" or bool(a["agent_output"]["anchor"])), (link, a["id"])
        for key in ("constraints", "context"):                      # collapsed by default behind a "show" toggle
            blk = card.split(f'data-block="{key}"')[1].split('data-block=')[0]
            assert '<div class="fold" data-foldlist="blk-' + a["id"] + "-" + key + '">' in blk and ">show</button>" in blk \
                and '<div class="foldbody" hidden>' in blk and "<details" not in blk, key
        assert 'data-foldlist="blk-' + a["id"] + '-evidence"><button type="button" class="linkbtn" data-fold="blk-' + a["id"] \
            + '-evidence" aria-expanded="false">show</button><div class="foldbody" hidden>' in card
        for key in ("goal", "expected_output", "human_output"):     # always visible
            blk = card.split(f'data-block="{key}"')[1].split('data-block=')[0]
            assert "<details" not in blk, key
        assert htmllib.escape(a["human_output"]["text"], quote=True).replace("&#x27;", "'") in htmllib.unescape(card) \
            or a["human_output"]["text"] in htmllib.unescape(card)
        # the Result block holds the preview: the section in a sandboxed frame, else the Notes text
        pre = card.split('data-block="result"')[1]
        if a["agent_output"]["anchor"]:
            src = htmllib.unescape(re.search(r'<iframe class="apreview" data-preview="frame" sandbox="" [^>]*srcdoc="([^"]*)"', pre).group(1))
            el = j2.element_by_id(page, a["agent_output"]["anchor"])
            assert src == previews[a["id"]] and el and el in page and el in src, a["id"]
            assert src.count(el) == 1 and src.endswith(el + "</body></html>")
        else:
            assert 'data-preview="text"' in pre and htmllib.unescape(a["human_output"]["text"]) in htmllib.unescape(pre)
    # no section artifact in the stored data: the Notes text takes the preview's place, nothing is fetched, no link
    a0 = anchored[0]
    saved = json.loads(json.dumps(snap))
    saved["atoms_v1"][a0["id"]]["output"] = [o for o in saved["atoms_v1"][a0["id"]]["output"] if "#" not in o["ref"]]
    for w in saved["workflows"]:
        for a in w["atoms"]:
            if a["id"] == a0["id"]:
                a["agent_output"]["anchor"] = None
    pre = render_node(saved)["cards"][next(w["id"] for w in saved["workflows"] if any(a["id"] == a0["id"] for a in w["atoms"]))]
    card0 = pre.split(f'data-key="atom-{a0["id"]}"')[1].split("Run details")[0]
    assert 'data-preview="text"' in pre and "<iframe" not in card0 and 'data-link="site"' not in card0
    # an atom with an anchor but no preview yet shows a loading note
    assert 'data-preview="wait"' in render_node(snap)["cards"][atoms[0][0]["id"]] or not atoms[0][1]["agent_output"]["anchor"]
    # Simple view output is identical to before the atom cards changed (the checkpoint commit's page, same state)
    old = subprocess.run(["git", "show", "3ab8120:j2_ui.html"], capture_output=True, text=True,
                         cwd=os.path.dirname(os.path.abspath(__file__)))
    if old.returncode == 0:
        p = os.path.join(tempfile.mkdtemp(prefix="j2-old-ui-"), "j2_ui_old.html")
        open(p, "w", encoding="utf-8").write(old.stdout)
        before = render_node(snap, html=p)
        assert before["simple"] == out["simple"] and before["sprog"] == out["sprog"] and before["simple"], "Simple view changed"
    # the page's preview endpoint serves the same source
    import threading
    import urllib.request
    import j2_server
    keep, j2_server.RUN = j2_server.RUN, run
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        got = json.load(urllib.request.urlopen(f"http://127.0.0.1:{srv.server_address[1]}/api/preview?atom={a0['id']}"))
        assert got["anchor"] == a0["agent_output"]["anchor"] and got["html"] == previews[a0["id"]], got
        # a saved run is read from its folder
        got = json.load(urllib.request.urlopen(f"http://127.0.0.1:{srv.server_address[1]}/api/preview?atom={a0['id']}"
                                               f"&run={os.path.basename(run.dir)}"))
        assert got["html"] == previews[a0["id"]], got
    finally:
        j2_server.RUN = keep
        srv.shutdown()
    print(f"atom cards (no browser): preview + eight blocks in three groups, three links, collapsed long blocks, preview "
          f"source = the saved page's section ({len(anchored)} atoms with a frame, {len(plain)} with text), "
          "Simple view identical to the checkpoint page, /api/preview ok")


def docs_tab():
    """UI: Context and Constraints in plain words, the read-only Documents tab (/api/docs and its render), and links from
    an atom card to a document item. Nothing stored changes: the same lines are shown, grouped."""
    import shutil
    import threading
    import urllib.request
    import urllib.error
    import j2_server
    if not shutil.which("node"):
        print("docs tab test skipped: node not found")
        return
    plain = j2.J2Run()
    say(plain, BREW_LAB)
    run = docs_run(DESIGN_BRIEF, 900)
    assert run.outcome.startswith("Done"), run.outcome
    keep, j2_server.RUN = j2_server.RUN, run
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    try:
        for q in ("", f"?run={os.path.basename(run.dir)}"):          # live and saved: the three files, exactly as written
            got = json.load(urllib.request.urlopen(f"{base}/api/docs{q}"))
            assert list(got) == ["prd", "hld", "lld"], got
            for k in got:
                assert got[k]["asked"] and got[k]["text"] == open(os.path.join(run.dir, f"{k}.md"), encoding="utf-8").read(), k
                assert got[k]["workflows"] == run.docs[k]["workflows"] and got[k]["workflows"], k
        for bad in ("?run=../..", "?run=nope"):                      # no path tricks: a run name only; anything else is an empty result
            got = json.load(urllib.request.urlopen(f"{base}/api/docs{bad}"))
            assert all(v == {"asked": False, "text": None, "workflows": []} for v in got.values()), (bad, got)
        j2_server.RUN = plain                                        # no document asked for: noted, no text
        got = json.load(urllib.request.urlopen(f"{base}/api/docs"))
        assert all(v == {"asked": False, "text": None, "workflows": []} for v in got.values()), got
    finally:
        j2_server.RUN = keep
        srv.shutdown()
    # render: a card's constraints grouped, its context in words, document items linked
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
    w = snap["workflows"][0]
    a = w["atoms"][0]
    a["constraints"] = ["Create only #menu and .j2-card", "Other sections stay unchanged", "No invented items or prices",
                        "Ids R1, R2 ... unique and sequential", "Placeholder for any photo"]
    a["context"] = ["charter items 1 and 2", "section S1", "PRD R2, R7", "HLD C1", "LLD S1", "finding F-3", "charter item 9"]
    w2 = snap["workflows"][1]
    snap["atoms_v1"][a["id"]]["depends_on"] = [{"atom": w2["atoms"][0]["id"], "why": "the section it wrote"}]
    snap["goal"]["charter_items"] = ["The page shows the menu", "The page shows the hours"]
    docs = {"docs": {"prd": {"asked": True, "workflows": ["W-1", "W-2"],
                             "text": "# P\n\n- **G1**: Goal one\n2. **R2**: The pre-order form\n| G1 | 30 |\n"},
                     "hld": {"asked": True, "text": "- **C1** Menu display: shows the menu\n- **C2** Form", "workflows": ["W-9"]},
                     "lld": {"asked": True, "text": "- **S1** `#menu`: the menu with prices"}}, "focus": {"kind": "prd", "item": "R2"}}
    out = render_node(snap, docs=docs)
    card = out["cards"][w["id"]]
    card = card[card.index(f'data-key="atom-{a["id"]}"'):].split("Run details")[0]
    con = card.split('data-block="constraints"')[1].split("data-block=")[0]
    assert re.findall(r"<b>([^<]*)</b><ul", con) == ["Only build", "No inventing", "Do not touch", "Format"], con
    for line in a["constraints"]:
        assert line in con, line
    assert re.search(r"</div><ul class=\"clist\"><li>Placeholder for any photo</li></ul>", con), "ungrouped line: plain list"
    ctx = card.split('data-block="context"')[1].split("data-block=")[0]
    assert "<b>Context</b> <span class=\"empty\">What it needs to know</span>" in card and "What this workflow was given" not in card
    rows = dict.fromkeys(re.findall(r'<li data-ctx="([\w-]+)">', ctx))
    assert list(rows) == ["charter", "lld", "prd", "hld", "finding", "from-workflow"], list(rows)
    assert '<li data-ctx="charter"><b class="ctype">Charter</b> <a href="#" data-charter="1">item 1</a>: The page shows the menu</li>' in ctx, ctx
    assert 'data-charter="2">item 2</a>: The page shows the hours</li>' in ctx
    assert 'item 9</li>' not in ctx and 'item 9 <span class="empty">no source on record</span>' in ctx, "a row with no source says so"
    assert '<b class="ctype">LLD</b> <a href="#" data-doc="lld" data-item="S1">S1</a>: #menu: the menu with prices</li>' in ctx, ctx
    assert '<b class="ctype">PRD</b> <a href="#" data-doc="prd" data-item="R2">R2</a>: The pre-order form</li>' in ctx, ctx
    assert '<b class="ctype">PRD</b> R7 <span class="empty">no source on record</span>' in ctx or 'data-item="R7">R7</a> <span class="empty">' in ctx, ctx
    assert '<b class="ctype">HLD</b> <a href="#" data-doc="hld" data-item="C1">C1</a>: Menu display: shows the menu</li>' in ctx, ctx
    assert '<b class="ctype">LLD</b> <a href="#" data-doc="lld" data-item="S1">S1</a>' in ctx
    assert re.search(r'<b class="ctype">Finding</b> F-3: [^<]+</li>', ctx), "F-3 is found in the run's audit"
    assert f'<b class="ctype">From workflow</b> {w2["id"]} {w2["name"]}: the section it wrote.</li>' in ctx.replace("&amp;", "&"), ctx
    assert "Never break" not in card and "Uses " not in ctx
    # no group fits: a plain list only
    a["constraints"] = ["Placeholder for any photo"]
    plain_card = render_node(snap, docs=docs)["cards"][w["id"]].split('data-block="constraints"')[1].split("data-block=")[0]
    assert "<b>" not in plain_card.split("</button>")[1] and "<li>Placeholder for any photo</li>" in plain_card
    # the tab: a jump bar, one coloured block per document (name, item count, builders), dividers, the note, the highlight
    d = out["docs"]
    assert re.findall(r'data-doc-kind="(\w+)"', d) == ["prd", "hld", "lld"]
    assert re.findall(r'data-docjump="(\w+)"', d.split('<section')[0]) == ["prd", "hld", "lld"], "jump bar first"
    heads = re.findall(r'<header class="dochead" style="background:var\(--(\w+)\)"><b>([^<]*)</b> <span class="code">(\w+)</span>(?:<span class="docmeta">([^<]*)</span>)?</header>', d)
    assert [h[:3] for h in heads] == [("prd", "Product requirements", "PRD"), ("hld", "High-level design", "HLD"),
                                      ("lld", "Low-level design", "LLD")], heads
    assert heads[0][3] == f"2 items · built by W-1 {snap['workflows'][0]['name']}, W-2 {snap['workflows'][1]['name']}", heads[0]
    assert all(not h[3].startswith(("·", " ·")) for h in heads), "no stray leading dot before the count"
    assert 'data-tabjump="charter">Charter</a>' in d.split("<section")[0], "the jump bar has a Charter entry"
    assert "2 items" in heads[1][3] and "built by W-9" in heads[1][3] and "items" in heads[2][3] and "built by" not in heads[2][3], heads
    assert re.findall(r'<span>(End of \w+|Start of \w+)</span>', d) == ["End of PRD", "Start of HLD", "End of HLD", "Start of LLD", "End of LLD"], d
    assert d.count('data-docjump="top"') == 3
    assert '<p class="dl hit" id="doc-prd-R2">' in d and '<p class="dl" id="doc-prd-G1">' in d and 'id="doc-lld-S1"' in d, d
    assert d.count('id="doc-prd-G1"') == 1 and "<h4>P</h4>" in d, "an item's first line only carries the id"
    docs["docs"]["hld"] = {"asked": False, "text": None, "workflows": []}
    nd = render_node(snap, docs=docs)["docs"]
    assert "Not asked for in this run." in nd.split('data-doc-kind="hld"')[1].split("</section>")[0] and nd.count("End of HLD") == 1
    print("docs tab: Context in words, Constraints grouped (stored text unchanged, ungrouped in a plain list), document "
          "items linked, Documents tab render, /api/docs live and saved, 404 for bad runs, not-asked note ok")


def html_page():
    return open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "j2_ui.html"), encoding="utf-8").read()


def card_content():
    """UI: card headings unchanged; a revise atom names what it revises; Context rows by type (workflow ids, condition and
    finding lookups, no source on record, never a blank row); long dependency lists collapsed; outputs vs "Saved in"."""
    import shutil
    if not shutil.which("node"):
        print("card content test skipped: node not found")
        return
    run = docs_run(DESIGN_BRIEF, 900)
    assert run.outcome.startswith("Done"), run.outcome
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
    wfs = snap["workflows"]
    w1, w2, w3 = wfs[0], wfs[1], wfs[2]
    a1, a2 = w1["atoms"][0], w2["atoms"][0]
    w2["revises"], w2["section"], w2["agent"] = w1["id"], w1["id"], f"{w2['id']} Builder"
    a2["how"] = {"kind": "AI", "agent": f"{w1['id']} Builder", "steps": "Rewrite"}
    a2["context"] = [f"section {w1['id']}", f"section {w3['id']}", "condition H5", "finding D-5", "finding F-1", "finding X-9",
                     "condition Z9", "", "   ", None, "PRD R9", "PRD R8"]
    a2["constraints"] = ["Create only #menu", "Do not change other sections"]
    snap["audits"] = [{"results": [{"id": "F-1", "criterion": "The page opens", "note": "it parsed"}]}]
    deps = snap["atoms_v1"][a2["id"]]
    SITE = "its section, anchor and selectors were listed in this builder's prompt"
    deps["depends_on"] = [{"atom": w["atoms"][0]["id"], "why": SITE if i >= 4 else f"why {i}"} for i, w in enumerate(wfs[:6])]
    snap["atoms_v1"][wfs[4]["atoms"][0]["id"]]["output"] = [{"verb": "make", "name": "S", "form": "doc", "ref": "prd.md#Success metrics"}]
    snap["atoms_v1"][wfs[5]["atoms"][0]["id"]]["output"] = [{"verb": "make", "name": "S", "form": "page", "ref": "index.html#menu"}]
    long_r9 = ("The pre-order form lets coffee lovers and office staff nearby choose items and a pickup time and send the order "
               "to the shop so that it is ready when they arrive and nothing is paid online at all")
    deps["used_by"] = deps["depends_on"][:2]
    deps["output"] = [{"verb": "make", "name": "hld.md", "form": "doc", "ref": "hld.md"},
                      {"verb": "make", "name": "A section", "form": "doc", "ref": "hld.md#Sec"},
                      {"verb": "make", "name": "site/v9.md", "form": "doc", "ref": "site/v9.md"}]
    hld = {c["id"]: c["text"] for c in j2.load_conditions("hld.json")}
    refs = {"conditions": hld, "rubric": {i: t for o in j2.OUTPUTS.values() for i, t in o.get("rubric") or ()}}
    out = render_node(snap, docs={"docs": {"prd": {"asked": True, "text": "- **R1**: One\n- **R9**: " + long_r9}}, "refs": refs})
    opened = render_node(snap, docs={"docs": {"prd": {"asked": True, "text": "- **R1**: One\n- **R9**: " + long_r9}}, "refs": refs,
                                     "open": [f"blk-{a2['id']}-dep", f"blk-{a2['id']}-used", f"ctx-{a2['id']}-R9",
                                              f"blk-{a2['id']}-context", f"blk-{a2['id']}-constraints", f"blk-{a2['id']}-evidence"]})

    def card_of(o, i, wf):
        c = o["cards"][wf["id"]]
        return c[c.index(f'data-key="atom-{i}"'):].split("Run details")[0]
    card = lambda i, wf: card_of(out, i, wf)
    c1, c2 = card(a1["id"], w1), card(a2["id"], w2)
    # headings and glosses are the original ones
    for lab in ("Goal", "How", "Constraints", "Context", "Defined output", "Verification", "Notes", "Workflow output", "Result"):
        assert f"<b>{lab}</b>" in c1, lab
    assert '<b>Context</b> <span class="empty">What it needs to know</span>' in c1 and 'Who builds it' in c1
    assert "Depends on:" in c2 and "Used by:" in c2 and "What this workflow was given" not in c2
    # a revise atom: its own builder name and what it revises
    how = c2.split('data-block="how"')[1].split("data-block=")[0]
    assert f"AI builds it: {w2['id']} {w2['name']} builder as a revision of {w1['id']} {w1['name']}." in how.replace("&amp;", "&"), how
    assert f"AI builds it: {w1['id']} {w1['name']} builder." in c1.split('data-block="how"')[1].split("data-block=")[0].replace("&amp;", "&")
    assert "as a revision" not in c1.split('data-block="how"')[1].split("data-block=")[0]
    import html as htmllib
    ctx = htmllib.unescape(c2.split('data-block="context"')[1].split("data-block=")[0])
    row = lambda t: re.findall(rf'<li data-ctx="[\w-]+"><b class="ctype">{t}</b> ([^<]*(?:<[^>]*>[^<]*)*?)</li>', ctx)
    assert f'<li data-ctx="revises"><b class="ctype">Revises</b> {w1["id"]} {w1["name"]}: why 0.</li>' in ctx, ctx
    assert f'<b class="ctype">From workflow</b> {w3["id"]} {w3["name"]}: why 2.' in ctx, ctx
    assert '<b class="ctype">Section</b>' not in ctx, "a workflow id is never a Section"
    assert f'<b class="ctype">Condition</b> H5: {hld["H5"]}</li>' in ctx, ctx
    assert f'<b class="ctype">Finding</b> D-5: {refs["rubric"]["D-5"]}</li>' in ctx, ctx
    assert '<b class="ctype">Finding</b> F-1: The page opens: it parsed</li>' in ctx, ctx
    assert '<b class="ctype">Finding</b> X-9 <span class="empty">no source on record</span>' in ctx
    assert '<b class="ctype">Condition</b> Z9 <span class="empty">no source on record</span>' in ctx
    assert '<b class="ctype">PRD</b> R8 <span class="empty">no source on record</span>' in ctx
    # the Context, Constraints and Verification toggles: hidden "show" by default, shown "hide" once open, text present either way
    oc = card_of(opened, a2["id"], w2)
    for key, text in (("context", "Condition</b> H5"), ("constraints", "Only build"), ("evidence", "")):
        for card_html, label, hid in ((c2, "show", " hidden"), (oc, "hide", "")):
            blk = card_html.split(f'data-foldlist="blk-{a2["id"]}-{key}"')[1]
            assert blk.startswith(f'><button type="button" class="linkbtn" data-fold="blk-{a2["id"]}-{key}" aria-expanded="{str(hid == "").lower()}">{label}</button>'
                                  f'<div class="foldbody"{hid}>'), (key, label, blk[:200])
            assert text in blk.split("data-block=")[0] and "<details" not in blk.split("data-block=")[0], (key, label)
    assert "Agent output" not in c2 and "Workflow output</button>" in c2 and 'data-link="agent"' in c2
    # the type label sits right before its text: one space, no padding
    assert '<b class="ctype">Condition</b> H5: ' in ctx and not re.search(r"\.ctype\{[^}]*(?:min-width|width|padding|margin)", html_page())
    # a long line is cut at a word boundary, with "show more"; open, the full text and "show less"
    m = re.search(r'<b class="ctype">PRD</b> <a [^>]*>R9</a>: ([^<]*)… <button[^>]*data-fold="ctx-[^"]*"[^>]*>show more</button>', ctx)
    assert m and long_r9.startswith(m.group(1)) and long_r9[len(m.group(1))] == " ", ("cut inside a word", m and m.group(1))
    octx = htmllib.unescape(card_of(opened, a2["id"], w2).split('data-block="context"')[1].split("data-block=")[0])
    assert long_r9 in octx and ">show less</button>" in octx and "… " not in octx.split("R9")[1].split("</li>")[0], octx
    # a workflow shown from Context is not repeated from depends_on
    assert ctx.count(f'{w1["id"]} {w1["name"]}') == 1 and ctx.count(f'{w3["id"]} {w3["name"]}') == 1
    # never a blank row: empty references are skipped, every row has text after its label
    rows = re.findall(r'<li data-ctx="[\w-]+">(.*?)</li>', ctx)
    assert rows and all(re.sub(r"<[^>]+>", "", r).replace(re.search(r'<b class="ctype">([^<]*)</b>', r).group(1), "", 1).strip() for r in rows), rows
    # collapsed dependency list: 6 rows hidden behind "show" with a count; "hide" and the rows once opened; 2 rows stay inline
    dep = c2.split("Depends on:")[1].split("Used by:")[0]
    assert re.search(r'6 atoms \(<button[^>]*data-fold="blk-[^"]*-dep"[^>]*aria-expanded="false">show</button>\)<span class="foldbody" hidden>', dep), dep
    assert "why 3" in dep, "rows are in the page text, hidden"
    odep = card_of(opened, a2["id"], w2).split("Depends on:")[1].split("Used by:")[0]
    assert re.search(r'6 atoms \(<button[^>]*aria-expanded="true">hide</button>\)<span class="foldbody">', odep) and "why 3" in odep, odep
    used = c2.split("Used by:")[1]
    assert "atoms (" not in used.split("Workflow output")[0] and "why 1" in used
    # reason text: a document section builder's edge names the document and section; the website wording stays
    assert f"{wfs[4]['atoms'][0]['id']}: the PRD Success metrics section was in this builder's prompt." in odep, odep
    assert f"{wfs[5]['atoms'][0]['id']}: {SITE}." in odep, odep
    # outputs: the section first, files on a "Saved in" line
    wo = c2.split('data-block="agent_output"')[1].split("data-block=")[0]
    assert wo.index("A section") < wo.index("Depends on") and "A section" in wo.split("Saved in")[0]
    assert 'Saved in: hld.md, site/v9.md' in wo and "site/v9.md" not in wo.split("Saved in")[0] and "hld.md</span>" not in wo.split("Saved in")[0]
    assert wo.index("A section") < wo.index("Saved in")
    # W-13's own data (a stored run, when present): nine rows, none blank
    here = os.path.dirname(os.path.abspath(__file__))
    name = "j2-20261016-160511"
    if os.path.isdir(os.path.join(here, "runs", name)):
        keep_dir, j2.RUNS_DIR = j2.RUNS_DIR, os.path.join(here, "runs")       # read-only: the stored run is evidence
        try:
            s13 = j2.saved_snapshot(name, 0)
        finally:
            j2.RUNS_DIR = keep_dir
        R = os.path.join(here, "runs", name)
        d13 = {"docs": {k: {"asked": True, "text": open(os.path.join(R, k + ".md"), encoding="utf-8").read()} for k in ("prd", "hld", "lld")}, "refs": refs}
        c13 = render_node(s13, docs=d13)["cards"]["W-13"]
        c13 = c13[c13.index('data-key="atom-W-13.a1"'):].split("Run details")[0]
        x13 = c13.split('data-block="context"')[1].split("data-block=")[0]
        r13 = re.findall(r'<li data-ctx="([\w-]+)">(.*?)</li>', x13)
        assert len(r13) == 8 and [t for t, _ in r13].count("from-workflow") == 5, r13          # W-4 and its revision W-8 carry the same text: one row
        assert all(re.sub(r"<[^>]+>", "", b).split(" ", 1)[-1].strip() for _, b in r13), r13
        assert "<li></li>" not in x13 and "<li data-ctx=\"\">" not in x13
        assert re.search(r'6 atoms \(<button[^>]*>show</button>\)<span class="foldbody" hidden>', c13.split("Depends on:")[1]), "W-13 depends on 6: collapsed"
        assert "the HLD Purpose and scope section was in this builder's prompt" in c13 and SITE not in c13, "document wording in a document run"
        assert "W-13 HLD Hosting and deployment section builder" in c13
    print("card content: original headings, revise label, workflow/condition/finding rows, no source noted, no blank rows, "
          "collapsed dependencies, outputs vs Saved in ok")


def charter_tab():
    """UI: the read-only Charter tab (kept apart from the Goal tab): header, note, brief, items with source and Used by,
    business type and answers, the empty case, and "Charter item N" links that open it at that item."""
    import shutil
    if not shutil.which("node"):
        print("charter tab test skipped: node not found")
        return
    run = docs_run(DESIGN_BRIEF, 900)
    assert run.outcome.startswith("Done"), run.outcome
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
    w1, w2 = snap["workflows"][0], snap["workflows"][1]
    for w in snap["workflows"]:
        for at in w["atoms"]:
            at["context"] = []
    w1["atoms"][0]["context"] = ["charter item 1", "section Problem"]
    w2["atoms"][0]["context"] = ["charter items 1 and 2", "PRD R1"]
    w1["atoms"][0]["goal"] = "Per charter item 2 the page opens"
    g = snap["goal"]
    g["charter_items"] = ["The page shows the menu", "The page shows the hours", "Nobody cites this"]
    g["business_type"] = "coffee shop"
    g["requester_answers"] = ["Pickup only", "", "Open 7 to 9"]
    snap["chat"] = [{"who": "user", "text": "A website for Brew Lab", "id": 1}]
    out = render_node(snap, docs={"docs": {}, "chfocus": 2})
    c = out["charter"]
    assert c.startswith('<p class="chnote empty">The charter is the list of things the Requester confirmed. Every agent works from it.</p>'), c[:200]
    assert c.index("data-brief") < c.index('id="charterblock"') and "A website for Brew Lab" in c, "the brief is at the top"
    head = re.search(r'<header class="dochead" style="background:var\(--charter\)"><b>Charter</b> <span class="code">CHARTER</span><span class="docmeta">([^<]*)</span></header>', c)
    assert head and head.group(1) == "3 items · Recorded by Identity", head
    assert re.findall(r'<li value="(\d)" id="charter-item-\d" class="dl( hit)?"', c) == [("1", ""), ("2", " hit"), ("3", "")], "numbered in order, item 2 highlighted"
    for n, t in ((1, "The page shows the menu"), (2, "The page shows the hours")):
        assert f'data-item="{n}">{t}' in c
    # no run stores where an item came from: nothing is shown or guessed; a stored source (charter_sources) is shown
    assert "data-chsrc" not in c and "Came from" not in c
    g["charter_sources"] = {"1": "Brief: \"menu with prices\""}
    c2 = render_node(snap, docs={"docs": {}})["charter"]
    assert c2.count("data-chsrc") == 1 and 'Came from: Brief: &quot;menu with prices&quot;</div>' in c2.split('data-item="2"')[0]
    # Used by: the workflows whose Context cites the item, as links that open the card
    used = lambda n: re.search(rf'data-item="{n}">.*?<div class="chused" data-chused>Used by: (.*?)</div>', c, re.S).group(1)
    assert f'<a href="#" data-wfcard="{w1["id"]}">{w1["id"]} {w1["name"]}</a>' in used(1) and f'data-wfcard="{w2["id"]}"' in used(1), used(1)
    assert f'data-wfcard="{w2["id"]}"' in used(2) and f'data-wfcard="{w1["id"]}"' not in used(2), used(2)
    assert used(3) == "no workflow cites it"
    # type, then answers under their own headings; blank answers skipped
    assert "<h4>Business type</h4><p>coffee shop</p>" in c and "<h4>The Requester's confirmed answers</h4>" in c
    assert c.count("<li>Pickup only</li>") == 1 and "<li>Open 7 to 9</li>" in c and "<li></li>" not in c
    assert c.index("Used by") < c.index("<h4>Business type</h4>") < c.index("confirmed answers")
    # the Goal tab is still there
    html = html_page()
    assert 'data-tab="goal"' in html and 'data-tab="charter"' in html and 'id="pane-goal"' in html and 'id="pane-charter"' in html
    # empty: no charter yet
    empty = json.loads(json.dumps(snap))
    empty["goal"] = None
    assert render_node(empty, docs={"docs": {}})["charter"].endswith('<p class="empty">No charter recorded yet.</p>')
    # without business type and answers, no empty headings
    bare = json.loads(json.dumps(snap))
    bare["goal"].pop("business_type"), bare["goal"].pop("requester_answers")
    assert "<h4>Business type" not in render_node(bare, docs={"docs": {}})["charter"]
    # "Charter item N" on a card is a link that opens the tab at that item (Context rows and plain text)
    cards = render_node(snap, docs={"docs": {}})["cards"]
    card1 = cards[w1["id"]].split('data-key="atom-' + w1["atoms"][0]["id"] + '"')[1].split("Run details")[0]
    assert '<b class="ctype">Charter</b> <a href="#" data-charter="1">item 1</a>: The page shows the menu' in card1
    assert 'Per charter item <a href="#" data-charter="2">2</a> the page opens' in card1.replace("Per charter item", "Per charter item", 1) or \
        'charter item <a href="#" data-charter="2">2</a>' in card1, card1[:1500]
    card2 = cards[w2["id"]].split('data-key="atom-' + w2["atoms"][0]["id"] + '"')[1].split("Run details")[0]
    assert 'data-charter="1">item 1</a>' in card2 and 'data-charter="2">item 2</a>' in card2
    print("charter tab: header, note, brief, numbered items with highlight, source only when stored, Used by links, business type "
          "and answers, empty case, item links on cards, jump-bar entry, Goal tab kept ok")


def fixes_line():
    """UI: a revise atom's "Fixes: ids. Before: status. After: status." line from stored audits only; a status nobody recorded
    says "not recorded"; an atom that revises nothing has no line. Also the /api/audits endpoint, on a stored run."""
    import shutil
    import threading
    import urllib.request
    import j2_server
    if not shutil.which("node"):
        print("fixes line test skipped: node not found")
        return
    run = docs_run(DESIGN_BRIEF, 900)
    assert run.outcome.startswith("Done"), run.outcome
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
    w1, w2, w3 = snap["workflows"][:3]
    a1, a2, a3 = w1["atoms"][0], w2["atoms"][0], w3["atoms"][0]
    w2["revises"], w2["cycle"], w2["fixes"] = w1["id"], 3, ["F-1"]
    a2["context"] = ["finding F-1", "condition H5", "charter item 1"]
    w3["revises"], w3["cycle"], w3["fixes"] = w1["id"], 3, []
    a3["context"] = ["finding F-1"]
    for a in (a1,):
        a["context"] = ["finding F-1"]                       # w1 revises nothing: no line, whatever its Context says
    w1.pop("revises", None)
    snap["audits"] = [{"cycle": 2, "results": [{"id": "F-1", "verdict": "NEEDS_REVISION", "pass": False}], "conditions": [{"id": "H5", "met": False}]},
                      {"cycle": 3, "results": [{"id": "F-1", "verdict": "PASS", "pass": True}], "conditions": [{"id": "H5", "met": True}]}]

    def line(sn, a, w, stage=None):
        c = render_node(sn, docs={"docs": {}, "stage_audits": stage or []})["cards"][w["id"]]
        blk = c[c.index(f'data-key="atom-{a["id"]}"'):].split("Run details")[0].split('data-block="agent_output"')[1].split("data-block=")[0]
        m = re.search(r'<span data-note="fixes">(.*?)</span>(?=<br>|</div>|$)', blk.replace('<span class="empty">not recorded</span>', "NOTREC"), re.S)
        return m.group(1).replace("NOTREC", "not recorded") if m else None
    # both recorded
    assert line(snap, a2, w2) == "Fixes: F-1, H5. Before: F-1 needs revision, H5 not met. After: F-1 passed, H5 met.", line(snap, a2, w2)
    # a single id: one status for the whole line
    assert line(snap, a3, w3) == "Fixes: F-1. Before: needs revision. After: passed."
    # only the before is recorded: the after says so, and nothing is guessed
    only_before = json.loads(json.dumps(snap))
    only_before["audits"] = only_before["audits"][:1]
    assert line(only_before, a3, w3) == "Fixes: F-1. Before: needs revision. After: not recorded."
    # only the after is recorded
    only_after = json.loads(json.dumps(snap))
    only_after["audits"] = only_after["audits"][1:]
    assert line(only_after, a3, w3) == "Fixes: F-1. Before: not recorded. After: passed."
    assert line(only_after, a2, w2) == "Fixes: F-1, H5. Before: not recorded. After: F-1 passed, H5 met."
    # nothing recorded
    none = json.loads(json.dumps(snap))
    none["audits"] = []
    assert line(none, a3, w3) == "Fixes: F-1. Before: not recorded. After: not recorded."
    # earlier stages come from /api/audits: the stage audits fill in what the current stage's audits lack
    stage = [{"cycle": 2, "results": [{"id": "F-1", "verdict": "FAIL"}], "conditions": []}]
    assert line(none, a3, w3, stage) == "Fixes: F-1. Before: failed. After: not recorded."
    # no ids in Context: the stored fixes list; neither: "not recorded"
    a3["context"] = ["charter item 1"]
    w3["fixes"] = ["F-1"]
    assert line(snap, a3, w3).startswith("Fixes: F-1. Before:")
    w3["fixes"] = []
    assert line(snap, a3, w3) == "Fixes: not recorded."
    # the line sits right under what the atom made, above Saved in, Depends on and Used by
    blk = render_node(snap, docs={"docs": {}})["cards"][w2["id"]].split('data-key="atom-' + a2["id"] + '"')[1].split("Run details")[0].split('data-block="agent_output"')[1]
    assert blk.index("make · ") < blk.index('data-note="fixes"') < blk.index("Depends on:") < blk.index("Used by:"), "Fixes under the section line"
    assert 'data-note="saved"' not in blk or blk.index('data-note="fixes"') < blk.index('data-note="saved"')
    # a non-revise atom has no Fixes line
    assert line(snap, a1, w1) is None
    # the endpoint reads the stage files, nothing else
    here = os.path.dirname(os.path.abspath(__file__))
    name = "j2-20261016-160511"
    if os.path.isdir(os.path.join(here, "runs", name)):
        keep_dir, j2.RUNS_DIR = j2.RUNS_DIR, os.path.join(here, "runs")             # read-only: the stored run is evidence
        srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            got = json.load(urllib.request.urlopen(f"http://127.0.0.1:{srv.server_address[1]}/api/audits?run={name}"))
            real = j2.saved_snapshot(name, 0)
        finally:
            j2.RUNS_DIR = keep_dir
            srv.shutdown()
        assert sorted({a["cycle"] for a in got}) == [1, 2, 3, 4, 5, 6, 7, 8, 9] and all("note" not in r for a in got for r in a["results"]), got[:1]
        w8 = next(w for w in real["workflows"] if w["id"] == "W-8")
        c = render_node(real, docs={"docs": {}, "stage_audits": got})["cards"]["W-8"]
        assert re.search(r'Fixes: D-5\. Before: (passed|passed with a warning|needs revision|failed)\. After: ', c), "W-8 records D-5 before and after"
        c15 = render_node(real, docs={"docs": {}, "stage_audits": got})["cards"]["W-15"]
        assert "Fixes: D-5, H5. Before:" in c15 and "not recorded" not in c15.split("Fixes:")[1].split("</span>")[0] or "Fixes: D-5, H5" in c15, c15[:50]
    print("fixes line: both statuses, one status, none recorded, stage audits, stored fixes fallback, no line without a revision, /api/audits ok")


def fresh_page_run():
    """A fresh page: the live run has an id but no folder, because no prompt was sent. That is normal and empty: every endpoint
    answers (by the live path or by the run's own id), and the page's state renders with no error and no leftover content."""
    import shutil
    import threading
    import urllib.request
    import j2_server
    keep = j2_server.RUN
    live = j2.J2Run()
    assert not os.path.exists(live.dir), "a fresh run has no folder yet"
    j2_server.RUN = live
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    try:
        rid = os.path.basename(live.dir)
        for q in ("", f"run={rid}"):                               # the page asks without ?run=, or with the live run's own id
            def get(path):
                return json.load(urllib.request.urlopen(base + path + (("&" if "?" in path else "?") + q if q else "")))
            st = get("/api/state?since=0")
            assert st.get("status") == "idle" and st["chat"] == [] and not st.get("error") and st["workflows"] == [], (q, list(st)[:6])
            assert get("/api/audits") == [] and all(v["text"] is None and not v["asked"] for v in get("/api/docs").values()), q
            assert get("/api/preview?atom=W-1.a1") == {"anchor": None, "html": None}
            assert "conditions" in get("/api/refs") and get("/api/runs")["live"] == rid
        assert not os.path.exists(live.dir), "reading must not create the run folder"
        # the page's render of that state: Charter, Documents and cards say "nothing yet", no error text, no blank rows
        if shutil.which("node"):
            docs = {"docs": get("/api/docs"), "stage_audits": get("/api/audits")}
            out = render_node(st, docs=docs)
            assert out["charter"].endswith("No charter recorded yet.</p>") and out["cards"] == {}, out["charter"][-80:]
            assert out["docs"].count("Not asked for in this run.") == 3 and "undefined" not in out["docs"] + out["charter"], out["docs"][:300]
            assert "Something went wrong" not in out["simple"] + out["docs"] + out["charter"]
        if j2.j2_check.find_chrome() and (BROWSER or not FAST):       # the real page: no error text in the page's own elements
            import subprocess
            for q in ("?view=detail", f"?view=detail&run={rid}", "?view=simple"):
                dom = chrome_dom([j2.j2_check.find_chrome(), "--headless=new", "--disable-gpu", "--virtual-time-budget=6000",
                                      "--dump-dom", base + "/" + q], capture_output=True, text=True, timeout=90).stdout
                assert re.search(r'id="sstatus">\s*</', dom) and 'class="oops"' not in dom.split("</script>")[-1], q
    finally:
        j2_server.RUN = keep
        srv.shutdown()
    print("fresh page: live run with an id and no folder: every endpoint answers empty (with and without its own id), the folder "
          "is not created, Charter and Documents render 'nothing yet' with no error ok")


def htmllib_unescape(t):
    import html
    return html.unescape(t)


def output_tags():
    """UI: output tags on the workflow list and cards (from atoms.json refs), the filter, group lines, legend colours and the
    run summary. Expected values are computed here from the stored atoms, not from the page's own code."""
    import shutil
    if not shutil.which("node"):
        print("output tags test skipped: node not found")
        return
    run = docs_run(DESIGN_BRIEF, 900)
    assert run.outcome.startswith("Done"), run.outcome
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
    stored = json.load(open(os.path.join(run.dir, "atoms.json"), encoding="utf-8"))["atoms"]
    by_atom = {a["id"]: a for a in stored}

    def expect(wf, views):                                    # tags a workflow should show, in order, from its atoms' output refs
        tags = []
        for a in wf["atoms"]:
            for o in views[a["id"]]["output"]:
                ref = o["ref"]
                m = re.match(r"(prd|hld|lld)\.md(#|$)", ref)
                if m:
                    t = m.group(1).upper()
                elif o["verb"] == "make" and re.fullmatch(r"[\w./-]+\.(md|html|docx|css|js|json)", ref):
                    continue
                elif o["verb"] == "make" and "#" in ref and re.search(r"\.html?#", ref):
                    t = "Site " + ref.split("#")[1]
                else:
                    t = "Other"
                if t not in tags:
                    tags.append(t)
        return tags

    def runs(ids):
        nums = sorted(int(i[2:]) for i in ids)
        out, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            out += [f"W-{nums[i]} to W-{nums[j]}"] if j - i >= 2 else [f"W-{n}" for n in nums[i:j + 1]]
            i = j + 1
        return ", ".join(out)
    views = snap["atoms_v1"]
    w_first, w_multi, w_none = snap["workflows"][0], snap["workflows"][1], snap["workflows"][2]
    views[w_multi["atoms"][0]["id"]]["output"] = [{"verb": "make", "name": "x", "form": "doc", "ref": "prd.md#Users"},
                                                  {"verb": "make", "name": "y", "form": "page", "ref": "index.html#menu"},
                                                  {"verb": "do", "name": "z", "form": "behaviour", "ref": "check-1"},
                                                  {"verb": "make", "name": "f", "form": "doc", "ref": "site/v3.md"}]
    views[w_none["atoms"][0]["id"]]["output"] = []
    out = render_node(snap, docs={"docs": {}})["workflows"]
    row = lambda o, wid: o.split(f"<tr><td>{wid}</td>")[1].split("</tr>")[0]
    card = lambda o, wid: o.split(f'data-key="wf-{wid}"')[1].split("<summary>")[0] and o.split(f'data-key="wf-{wid}"')[1].split("</summary>")[1].split("<details")[0]
    for w in snap["workflows"]:
        want = expect(w, views)
        got = re.findall(r'<button type="button" class="otag[^"]*" data-otag="\w+" style="[^"]*" title="[^"]*">([^<]*)</button>', row(out, w["id"]))
        assert got == want, (w["id"], got, want)
        assert (not want) == ("no output recorded" in row(out, w["id"])), w["id"]
        cardtags = re.findall(r'data-otag="\w+"[^>]*>([^<]*)</button>', card(out, w["id"]).split('<div class="ablock"')[0].split('<details class="atom"')[0])
        assert cardtags == want, (w["id"], cardtags, want)
    assert expect(w_multi, views) == ["PRD", "Site menu", "Other"], "several outputs, several tags (a saved-in file is not one)"
    assert "no output recorded" in row(out, w_none["id"]) and "no output recorded" in card(out, w_none["id"])
    # group lines: one per type, consecutive ids compressed, what has no output listed too
    lines = re.findall(r"<div>([^<]*)</div>", out.split('data-groups>')[1].split("<h3>")[0].replace("</div></div>", "</div>"))
    for label, key in (("PRD", "PRD"), ("HLD", "HLD"), ("LLD", "LLD"), ("Site", "Site"), ("Other", "Other")):
        ids = [w["id"] for w in snap["workflows"] if any(t == key or t.startswith(key + " ") for t in expect(w, views))]
        assert (f"{label}: {runs(ids)}" in lines) == bool(ids), (label, ids, lines)
    none_ids = [w["id"] for w in snap["workflows"] if not expect(w, views)]
    assert w_none["id"] in none_ids and f"No output recorded: {runs(none_ids)}" in lines, (none_ids, lines)
    # legend: one colour per type, the Documents headers use the same ones; status stays a word
    legend = out.split("data-legend>")[1].split("</div>")[0]
    assert re.findall(r'data-otag="(\w+)" style="background:var\(--(\w+)\)"', legend) == [(k, k) for k in ("prd", "hld", "lld", "site", "other")]
    page = html_page()
    assert all(re.search(rf"--{k}:var\(--\w+\)", page) for k in ("prd", "hld", "lld", "site", "other"))
    assert 'style="background:var(--prd)"><b>Product requirements' in render_node(snap, docs={"docs": {"prd": {"asked": True, "text": "- **R1**: x"}}})["docs"]
    assert "Status is always a word as well as a colour" in legend and '<span class="pill ' in out
    # filter: only workflows with that type; the same click again (filter cleared) is the full list
    n_all = out.count('<details data-key="wf-')
    for k in ("hld", "site"):
        f = render_node(snap, docs={"docs": {}, "otag": k})["workflows"]
        keep = [w["id"] for w in snap["workflows"] if any(t.lower().startswith(k) for t in expect(w, views))]
        assert re.findall(r'<details data-key="wf-([\w-]+)"', f) == keep and 0 < len(keep) < n_all or k == "site", (k, keep)
        assert f"Showing {len(keep)} of {n_all} workflows" in f and 'otag on" data-otag="' + k + '"' in f
    assert render_node(snap, docs={"docs": {}, "otag": ""})["workflows"] == out
    # summary: the numbers match atoms.json
    summ = out.split("data-summary>")[1].split('<div class="legend"')[0]
    closed = sum(a["status"] == "closed" for a in stored)
    failed = sum(a["status"] in ("verify_failed", "abandoned") for a in stored)
    assert f"{closed} closed · {failed} failed · {len(stored) - closed - failed} open" in summ, summ
    assert f"{snap['spent']} of {snap['budget']} used" in summ and htmllib_unescape(snap["outcome"].replace("done-when", "checklist")[:60]) in htmllib_unescape(summ), summ
    assert "<b>Built:</b> PRD, HLD, LLD, website" in summ, summ
    # a stored run, when present: the document groups are the documents' own workflows
    here = os.path.dirname(os.path.abspath(__file__))
    name = "j2-20261016-160511"
    if os.path.isdir(os.path.join(here, "runs", name)):
        keep_dir, j2.RUNS_DIR = j2.RUNS_DIR, os.path.join(here, "runs")
        try:
            real = j2.saved_snapshot(name, 0)
        finally:
            j2.RUNS_DIR = keep_dir
        ro = render_node(real, docs={"docs": {}})["workflows"]
        assert "PRD: W-1 to W-8" in ro and "HLD: W-9 to W-17" in ro and "LLD: W-18 to W-25" in ro, re.findall(r"<div>((?:PRD|HLD|LLD)[^<]*)</div>", ro)
        assert "31 closed · 0 failed · 0 open" in ro
        assert "show more</button>" in ro.split("data-summary>")[1].split("<b>Credits")[0], "a long outcome is cut at a word with show more"
    print("output tags: tags per workflow (several, none), saved-in files not tagged, filter and clear, group lines with compressed ids, "
          "legend colours = Documents colours, summary numbers = atoms.json ok")


def lineage_rows():
    """UI: a website atom in a documents run shows its document lineage in Context: LLD section text, the HLD component it
    implements, the PRD requirements it serves, each a link to the Documents tab; slice and summary text on From workflow rows;
    an unknown id says "not found in the documents"; a run without documents gets no such rows."""
    import shutil
    import html as htmllib
    if not shutil.which("node"):
        print("lineage rows test skipped: node not found")
        return
    run = docs_run(DESIGN_BRIEF, 900)
    assert run.outcome.startswith("Done"), run.outcome
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
    wfs = snap["workflows"]
    w, wd1, wd2 = wfs[0], wfs[1], wfs[2]
    a, ad1, ad2 = w["atoms"][0], wd1["atoms"][0], wd2["atoms"][0]
    a["context"] = ["charter item 1", "section S1", "section S9"]
    snap["goal"]["charter_items"] = ["The page shows the menu"]
    snap["atoms_v1"][a["id"]]["depends_on"] = [
        {"atom": ad1["id"], "why": "LLD slice S1: its pages and sections were in this builder's prompt"},
        {"atom": ad2["id"], "why": "the HLD component summary was in this builder's prompt"}]
    snap["atoms_v1"][ad1["id"]]["output"] = [{"verb": "make", "name": "x", "form": "doc", "ref": "lld.md#Pages and sections"}]
    snap["atoms_v1"][ad2["id"]]["output"] = [{"verb": "make", "name": "y", "form": "doc", "ref": "hld.md#Components"}]
    long_slice = "the menu with each item and its price so visitors see what things cost before they visit the shop in Indiranagar today"
    docs = {"docs": {
        "prd": {"asked": True, "text": "- **R1**: The page shows the menu\n- **R5**: One page only"},
        "hld": {"asked": True, "text": "## Components\n- **C1** Menu display: shows each menu item\n- **C2** Form: the pre-order form\n"},
        "lld": {"asked": True, "text": "## Pages and sections\n- **S1** `#menu`: " + long_slice + ". Implements C1, C9. Serves R1, R5, R6.\n- **S2** `#form`: other\n"}}}

    def ctx_of(sn, d):
        c = render_node(sn, docs=d)["cards"][w["id"]]
        return htmllib.unescape(c[c.index(f'data-key="atom-{a["id"]}"'):].split("Run details")[0].split('data-block="context"')[1].split("data-block=")[0])
    ctx = ctx_of(snap, docs)
    rows = re.findall(r'<li data-ctx="([\w-]+)">', ctx)
    assert rows == ["charter", "lld", "lld", "hld", "prd", "from-workflow", "from-workflow"], rows
    # LLD S1 with its own text (cut at a word, show more), linked
    assert re.search(r'<b class="ctype">LLD</b> <a href="#" data-doc="lld" data-item="S1">S1</a>: #menu: the menu with each item[^<]*… <button', ctx), ctx
    # HLD C1 (from "Implements C1, C9"): its text and a link; C9 is not in the HLD: no link, said plainly
    assert '<b class="ctype">HLD</b> <a href="#" data-doc="hld" data-item="C1">C1</a>, C9 <span class="empty">not found in the documents</span>: Menu display: shows each menu item</li>' in ctx, ctx
    # PRD R1, R5, R6 (from "Serves"): the two found are linked with their text; R6 is not found
    prd = re.search(r'<b class="ctype">PRD</b> (.*?)</li>', ctx).group(1)
    assert '<a href="#" data-doc="prd" data-item="R1">R1</a>, <a href="#" data-doc="prd" data-item="R5">R5</a>, R6 <span class="empty">not found in the documents</span>: R1 The page shows the menu; R5 One page only' == prd.replace("&amp;", "&"), prd
    # an unknown section id in a documents run: no link, "not found in the documents"
    assert '<b class="ctype">LLD</b> S9 <span class="empty">not found in the documents</span></li>' in ctx, ctx
    # From workflow rows show the slice and the summary themselves
    # the slice is the same text as the LLD S1 row above: a short pointer that keeps the link, not the text twice
    assert f'<b class="ctype">From workflow</b> <a href="#" data-wfcard="{wd1["id"]}">{wd1["id"]}</a>: same as <a href="#" data-doc="lld" data-item="S1">LLD S1</a> above</li>' in ctx, ctx
    assert ctx.count("the menu with each item") == 1, "the same text is not shown twice"
    assert f'{wd2["id"]} {wd2["name"]}: HLD component summary: C1 Menu display: shows each menu item C2 Form: the pre-order form</li>' in ctx, ctx
    assert "was in this builder" not in ctx, "the text, not the prompt wording"
    # an original and its revision with the same text: one row, the text once, both ids visible and linked
    wd3 = wfs[3]
    ad3 = wd3["atoms"][0]
    wd3["revises"] = wd2["id"]
    snap["atoms_v1"][ad3["id"]]["output"] = [{"verb": "make", "name": "y2", "form": "doc", "ref": "hld.md#Components"}]
    snap["atoms_v1"][a["id"]]["depends_on"].append({"atom": ad3["id"], "why": "the HLD component summary was in this builder's prompt"})
    mctx = ctx_of(snap, docs)
    merged = re.search(r'<li data-ctx="from-workflow"><b class="ctype">From workflow</b> (HLD Components summary, made by .*?)</li>', mctx).group(1)
    assert merged.startswith(f'HLD Components summary, made by <a href="#" data-wfcard="{wd2["id"]}">{wd2["id"]}</a>, revised by <a href="#" data-wfcard="{wd3["id"]}">{wd3["id"]}</a>: C1 Menu display'), merged
    assert mctx.count("C2 Form: the pre-order form") == 1 and mctx.count("HLD component summary:") == 0, "the text once"
    assert len(re.findall(r'<li data-ctx="from-workflow">', mctx)) == 2, "pointer row and merged row"
    # a long merged text keeps the show-more rule
    long_docs = json.loads(json.dumps(docs))
    long_docs["docs"]["hld"]["text"] = "## Components\n" + "- **C1** Menu display: " + "shows each menu item with its price " * 8 + "\n"
    assert re.search(r"made by .*?: C1 Menu display[^<]*… <button[^>]*data-fold=\"ctx-[^\"]*-m0\"[^>]*>show more</button>", ctx_of(snap, long_docs))
    # no duplicates: different texts stay separate rows in the usual form, and Depends on lists every atom as before
    nd = json.loads(json.dumps(snap))
    nd["atoms_v1"][a["id"]]["depends_on"] = [{"atom": ad1["id"], "why": "LLD slice S2: its pages and sections were in this builder's prompt"},
                                             {"atom": ad2["id"], "why": "the HLD component summary was in this builder's prompt"}]
    nctx = ctx_of(nd, docs)
    assert f'{wd1["id"]} {wd1["name"]}: LLD slice S2: S2 #form: other</li>' in nctx and f'{wd2["id"]} {wd2["name"]}: HLD component summary: C1 Menu display' in nctx
    assert "same as" not in nctx and "made by" not in nctx
    dep_card = render_node(snap, docs=docs)["cards"][w["id"]].split(f'data-key="atom-{a["id"]}"')[1].split("Run details")[0]
    assert f'{ad1["id"]}: LLD slice S1' in dep_card.split("Depends on:")[1] and f'{ad3["id"]}: the HLD component summary was in this builder' in dep_card.split("Depends on:")[1], "Depends on is unchanged"
    # a slice or summary missing from the documents is said so, not guessed
    gone = json.loads(json.dumps(docs))
    gone["docs"]["lld"]["text"] = "## Other\n- **S1** `#menu`: x\n"
    assert "LLD slice S1 <span class=\"empty\">not found in the documents</span>" in ctx_of(snap, gone)
    # every LLD / HLD / PRD row is a link to the Documents tab at that item
    assert all(re.search(r'data-doc="(prd|hld|lld)" data-item="[A-Z]\d+"', r) for r in re.findall(r'<li data-ctx="(?:prd|hld|lld)">.*?</li>', ctx) if "not found" not in r.split(":")[0])
    # a plain website atom in a run without documents: no lineage rows, nothing new
    plain = json.loads(json.dumps(snap))
    plain["atoms_v1"][a["id"]]["depends_on"] = []
    pa = next(x for x in plain["workflows"] if x["id"] == w["id"])["atoms"][0]
    pa["context"] = ["charter item 1", "section menu"]
    pc = ctx_of(plain, {"docs": {}})
    assert re.findall(r'<li data-ctx="([\w-]+)">', pc) == ["charter", "section"], pc
    assert '<b class="ctype">Section</b> menu <span class="empty">no source on record</span>' in pc and "not found in the documents" not in pc
    print("lineage rows: LLD section text, HLD component, PRD requirements (all linked), slice and summary text, unknown ids 'not found "
          "in the documents', plain website run unchanged ok")


def missing_run_endpoints():
    """Every GET endpoint that reads a run folder gives an empty result in its usual shape (404 for a file) when the folder
    is not there, for a saved run id and for a live run whose folder was never written; never a dropped connection."""
    import threading
    import urllib.request
    import urllib.error
    import j2_server
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    keep, keep_dir = j2_server.RUN, None

    def hit(u):
        try:
            r = urllib.request.urlopen(base + u)
            return r.status, r.read().decode("utf-8")
        except urllib.error.HTTPError as e:
            return e.code, e.read().decode("utf-8")
    empty_docs = {k: {"asked": False, "text": None, "workflows": []} for k in ("prd", "hld", "lld")}
    try:
        live = j2.J2Run()
        j2_server.RUN = live
        keep_dir, live.dir = live.dir, os.path.join(tempfile.mkdtemp(prefix="j2-missing-"), "j2-never-written")
        for q in ("run=j2-20261020-144628", "run=j2-nope", "run=..%2F..", ""):         # saved ids that do not exist, a bad id, the live run
            def u(path):
                return path + (("&" if "?" in path else "?") + q if q else "")
            assert json.loads(hit(u("/api/audits"))[1]) == [], q
            code, body = hit(u("/api/docs"))
            assert code == 200 and json.loads(body) == empty_docs, (q, code, body)
            code, body = hit(u("/api/preview?atom=W-1.a1"))
            assert code == 200 and json.loads(body) == {"anchor": None, "html": None}, (q, code, body)
            code, body = hit(u("/api/preview?atom=W-1.a1&x=1"))
            assert code == 200, (q, code)
            assert hit(u("/api/refs"))[0] == 200 and "conditions" in json.loads(hit(u("/api/refs"))[1])
            for f in ("/shot/1", "/file/prd.md", "/file/hld.docx", "/file/site.html"):
                assert hit(u(f))[0] == 404, (q, f)
            code, body = hit(u("/api/call?id=1"))
            assert code == 404 and "error" in json.loads(body), (q, code)
            assert hit(u("/site/latest"))[0] == 200 and hit(u("/site/v1"))[0] == 200, q
            if q:
                code, body = hit(u("/api/state"))
                assert code == 404 and "error" in json.loads(body), (q, code)
            assert hit(u("/api/runs"))[0] == 200
    finally:
        j2_server.RUN = keep
        if keep_dir:
            live.dir = keep_dir
        srv.shutdown()
    print("missing run folder: /api/audits, /api/docs, /api/preview give empty results; /api/refs, /api/runs, /site ok; /shot, /file, "
          "/api/call, /api/state answer 404 with no exception ok")


def replay_size():
    """D: the replay script reads a saved run (no model call, nothing written) and shows the slim prompts are smaller."""
    import subprocess
    here = os.path.dirname(os.path.abspath(__file__))
    if not os.path.isdir(os.path.join(here, "runs", "j2-20261015-114723")):
        print("replay test skipped: runs/j2-20261015-114723 not found")
        return
    r = subprocess.run([sys.executable, os.path.join(here, "j2_replay_size.py")], capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, r.stderr[-600:]
    m = re.search(r"new vs old: ([+-][\d.]+)%", r.stdout)
    assert m and float(m.group(1)) < 0 and "TOTAL" in r.stdout and "the fields cost less than they save: yes" in r.stdout, r.stdout[-500:]
    print(f"replay of runs/j2-20261015-114723: slim prompts {m.group(1)}% vs old, no model call ok")


CAMPUS = (" Users: students and teachers. Problem: students lose track of their notes. Must-have features: upload notes, "
          "search notes. Constraints: free to use. Success: 50 uploads per week.")
BREW_FACTS = (" Users: coffee lovers and office staff. Problem: people can't see the menu or order ahead. Must-have "
              "features: see the menu with prices, pre-order a coffee, find the shop, read the opening hours. "
              "Constraints: one page only. Success: 30 pre-orders per week.")
DESIGN_BRIEF = BREW_LAB + " Write the design docs first." + BREW_FACTS
PLAIN_CHECKPOINT = "13ccccd"   # plain_unchanged compares the working tree with this commit; move it when the plain path changes on purpose
CHECKPOINT = "0657d93"         # the commit before the design-document work: a plain brief must behave exactly as there


def docs_run(text, budget=900, patch=None):
    """A mock run of a brief; patch(run, orig_call) returns a replacement for run.call."""
    run = j2.J2Run()
    run.budget = budget
    if patch:
        run.call = patch(run, run.call)
    say(run, text)
    return run


def docs_files(run):
    return sorted(f for f in os.listdir(run.dir) if re.fullmatch(r"(?:prd|hld|lld)\.(?:md|docx)", f))


def docs_detection():
    """1: documents only when asked, by explicit terms. Set requests mean all three; naming some makes exactly those."""
    d = j2.detect_documents
    assert d(BREW_LAB) == [] and d("") == [] and d(None) == []
    for t in ("design docs", "Please write the design documents", "design documentation for my app", "I need full documentation",
              "DESIGN DOCS, please"):
        assert d(t) == ["prd", "hld", "lld"], t
    assert d("write a PRD") == ["prd"] and d("product requirements document") == ["prd"]
    assert d("HLD and LLD") == ["hld", "lld"] and d("an LLD and a PRD") == ["prd", "lld"]
    assert d("the high-level design") == ["hld"] and d("architecture overview") == ["hld"]
    assert d("a low-level design") == ["lld"] and d("detailed design of the forms") == ["lld"]
    for t in ("a modern design", "design a logo", "I like the design of your site", "design principles", "designing a menu",
              "clean design, warm colours", "a designer will review it"):
        assert d(t) == [], t
    assert j2.wants_site("docs for my website") and not j2.wants_site("design docs for the app")
    # a plain website brief creates none
    run = docs_run(BREW_LAB, 200)
    assert run.outcome.startswith("Done") and not run.docset and not run.stages and not docs_files(run), (run.outcome, docs_files(run))
    assert not os.path.exists(os.path.join(run.dir, "documents.json")) and not run.past_workflows
    # "design docs" creates all three, in order; "write a PRD" only the PRD; "HLD and LLD" those two
    run = docs_run("Design docs for Campus Notes." + CAMPUS, 400)
    assert run.outcome.startswith("Done") and list(run.docs) == ["prd", "hld", "lld"], (run.outcome, list(run.docs))
    assert docs_files(run) == ["hld.docx", "hld.md", "lld.docx", "lld.md", "prd.docx", "prd.md"], docs_files(run)
    run = docs_run("Write a PRD for Campus Notes." + CAMPUS, 200)
    assert run.outcome.startswith("Done") and list(run.docs) == ["prd"] and docs_files(run) == ["prd.docx", "prd.md"], docs_files(run)
    run = docs_run("Write the HLD and LLD for Campus Notes." + CAMPUS, 400)
    assert run.outcome.startswith("Done") and list(run.docs) == ["hld", "lld"] and docs_files(run) == \
        ["hld.docx", "hld.md", "lld.docx", "lld.md"], (run.outcome, docs_files(run))
    # answers count too, before the charter exists
    run = j2.J2Run()
    run.request = "A site for a cafe"
    run.intake = [{"asked": ["anything else?"], "answer": "Yes, an HLD as well"}]
    run._plan_stages()
    assert run.docset == ["hld"] and run.output == "hld"
    print("documents on request only: set requests, single requests, other senses of 'design', a plain brief (none), "
          "'design docs' (3), 'write a PRD' (1), 'HLD and LLD' (2) ok")


def docs_chain():
    """2: PRD, HLD, LLD in a chain: each reads the one before by its agent_output; H1-H6 and L1-L6 by code; chain
    gaps flagged before building; budget short drops documents, the website's essentials stay funded."""
    run = docs_run("Design docs for Campus Notes." + CAMPUS, 400)
    assert run.outcome.startswith("Done"), run.outcome
    arch = {k: json.load(open(os.path.join(run.dir, f"stage-{k}.json"), encoding="utf-8")) for k in ("prd", "hld", "lld")}
    for k, ids in (("hld", ["H1", "H2", "H3", "H4", "H5", "H6"]), ("lld", ["L1", "L2", "L3", "L4", "L5", "L6"])):
        conds = arch[k]["audits"][-1]["conditions"]
        assert [c["id"] for c in conds] == ids and all(c["met"] for c in conds), (k, conds)
    assert [c["id"] for c in arch["prd"]["audits"][-1]["conditions"]] == [f"P{i}" for i in range(1, 9)]
    docs = run.docs
    assert [docs[k]["workflows"][0] for k in docs] == ["W-1", "W-7", "W-13"], [docs[k]["workflows"] for k in docs]
    # the chain: every requirement is in a component, every component in an LLD section
    reqs = [r["id"] for r in docs["prd"]["agent_output"]["requirements"]]
    comps = docs["hld"]["agent_output"]["components"]
    secs = docs["lld"]["agent_output"]["sections"]
    assert reqs == ["R1", "R2"] and {r for c in comps for r in c["covers"]} == set(reqs)
    assert {c for s in secs for c in s["implements"]} == {c["id"] for c in comps}
    # each document is built by workflows whose atoms carry both outputs
    atoms = [a for w in run._all_wfs() for a in w["atoms"]]
    assert atoms and all(a.get("status") == "passed" and not j2.atom_done_problem(a) for a in atoms), \
        [(a["id"], a.get("status")) for a in atoms if not j2.atom_done_problem(a) == []]
    # the HLD builders get the PRD's agent_output (requirements), never its text
    rows = calls_file(run)
    hld_prompts = [r["prompt"] for r in rows if r["fn"] in docs["hld"]["workflows"]]
    assert hld_prompts and all('"requirements"' in p and "## Success metrics" not in p and "## Out of scope" not in p
                               for p in hld_prompts), "HLD builders must read the PRD's agent_output, not its text"
    lld_prompts = [r["prompt"] for r in rows if r["fn"] in docs["lld"]["workflows"]]
    assert lld_prompts and all('"components"' in p and "## Components" not in p for p in lld_prompts)
    # documents.json and the exported files
    saved = json.load(open(os.path.join(run.dir, "documents.json"), encoding="utf-8"))
    assert list(saved) == ["prd", "hld", "lld"] and saved["hld"]["file"] == "hld.md"
    for f in ("hld.docx", "lld.docx", "prd.docx"):
        assert open(os.path.join(run.dir, f), "rb").read(2) == b"PK", f
    # the L and H checks catch what they name
    hg = j2.stage_goal("hld", run.stage_archive[0]["goal"], {"prd": docs["prd"]["agent_output"]})
    src = j2.squash(" ".join(j2.strings(hg) + ["Design docs for Campus Notes." + CAMPUS]))
    hld = open(os.path.join(run.dir, "hld.md"), encoding="utf-8").read()
    miss = hld.replace("Covers R2.", "").replace(" R2", "")
    j2.OUTPUT = "hld"
    by = lambda conds: {c["id"]: c for c in conds}
    bad = by(j2.check_conditions(j2.load_conditions("hld.json"), miss, hg, src, j2.HLD_SECTIONS))
    assert not bad["H4"]["met"] and "R2" in bad["H4"]["evidence"], bad["H4"]
    dup = hld.replace("## Risks", "## Pages")
    assert not by(j2.check_conditions(j2.load_conditions("hld.json"), dup, hg, src, j2.HLD_SECTIONS))["H2"]["met"]
    lld = open(os.path.join(run.dir, "lld.md"), encoding="utf-8").read()
    lg = j2.stage_goal("lld", run.stage_archive[0]["goal"], {**{k: v["agent_output"] for k, v in docs.items()}})
    j2.OUTPUT = "lld"
    src = j2.squash(" ".join(j2.strings(lg) + ["Design docs for Campus Notes." + CAMPUS]))
    chk = lambda text: by(j2.check_conditions(j2.load_conditions("lld.json"), text, lg, src, j2.LLD_SECTIONS))
    assert all(c["met"] for c in chk(lld).values()), [(c["id"], c["evidence"]) for c in chk(lld).values() if not c["met"]]
    assert not chk(lld.replace("`#sec-2`", "`#sec-1`"))["L2"]["met"]
    assert "E2" in chk(lld.replace("trigger click; behaviour", "behaviour", 1).replace("**E2** in S2: `#act-2` trigger click; behaviour", "**E2** in S2: `#act-2`"))["L3"]["evidence"]
    assert not chk(lld.replace(", required yes", ""))["L4"]["met"]
    assert not chk(lld.replace("Implements C2.", "").replace("shows C2.", ""))["L5"]["met"]
    # a plan that leaves a chain link unnamed is flagged before building (then retried once, as the nav plan is)
    def strip_ids(run, orig):
        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            if fn == "adaptation" and isinstance(r, dict):
                for w in r.get("workflows") or []:
                    w["completion_check"] = re.sub(r", and covers .*", "", w.get("completion_check") or "")
            return r
        return call
    run = docs_run("Write the PRD and the HLD for Campus Notes." + CAMPUS, 400, strip_ids)
    flagged = [m["text"] for m in run.chat if m["who"] == "system" and "flagged the plan before building" in m["text"]]
    assert flagged and "no workflow builds" in flagged[0], flagged
    assert any(m["text"].startswith("The plan still leaves H4") for m in run.chat if m["who"] == "system")
    assert run.outcome.startswith("Done"), run.outcome       # the build still meets H4, so it is workable
    print("chain: PRD, HLD, LLD in order; H1-H6 and L1-L6 by code; HLD reads the PRD's agent_output, LLD the HLD's; "
          "chain gaps caught (R2 missing, repeated heading, repeated id, element without trigger, field without required "
          "flag, component not in an LLD section); plan flagged before building ok")


def docs_blueprint():
    """3: with an LLD the site is built from it: coverage checked before building, each builder gets its slice, the
    page is checked against it, one revision, then escalation, Audit compares, the chain is recorded."""
    run = docs_run(DESIGN_BRIEF, 900)
    assert run.outcome.startswith("Done") and run.stage_kind == "website", run.outcome
    spec = run.blueprint
    assert [s["id"] for s in spec["sections"]] == ["S1", "S2", "S3", "S4"]
    build = [w for w in run.workflows if w.get("lld_sections")]
    assert [w["lld_sections"] for w in build] == [["S1"], ["S2"], ["S3"], ["S4"]]
    for w in build:                                     # the check that the page matches the slice is on the atom
        chk = [c for c in w["atoms"][0]["verification"] if c["type"] == "lld-slice"]
        assert len(chk) == 1 and [s["id"] for s in chk[0]["spec"]["sections"]] == w["lld_sections"], w["id"]
        rec = [r for r in run.receipts if r["atom"] == w["atoms"][0]["id"] and r["type"] == "lld-slice"]
        assert rec and rec[-1]["pass"] is True and "as written" in rec[-1]["evidence"], rec
    # each builder's prompt: its own slice and a short HLD summary, not the whole page
    rows = calls_file(run)
    mine = {r["fn"]: r["prompt"] for r in rows if r["fn"] in {w["id"] for w in build}}
    for w in build:
        p = mine[w["id"]]
        sid = w["lld_sections"][0]
        assert "YOUR SPEC" in p and f'"id": "{sid}"' in p and "HLD SUMMARY" in p and "THE WHOLE SITE FILE" not in p, w["id"]
        other = [s for s in ("S1", "S2", "S3", "S4") if s != sid]
        assert not any(f'"id": "{o}", "anchor"' in p.split("YOUR SPEC", 1)[1].split("HLD SUMMARY")[0] for o in other)
    # the Audit compared the finished page to the LLD; the final record has the chain
    cmp = run.audits[-1]["lld"]
    assert cmp == {"matches": ["S1", "S2", "S3", "S4"], "mismatches": [], "dropped": []}, cmp
    chain = json.load(open(os.path.join(run.dir, "chain.json"), encoding="utf-8"))
    assert chain[0]["requirement"] == "R1" and chain[0]["component"] == "C1" and chain[0]["lld_section"] == "S1" \
        and chain[0]["atoms"] and chain == run._chain_record(), chain[:1]
    summ = json.load(open(os.path.join(run.dir, "summary.json"), encoding="utf-8"))["design_documents"]
    assert summ["requested"] == ["prd", "hld", "lld"] and summ["chain"] == chain and summ["lld_compare"] == cmp
    assert "LLD COMPARISON" in next(r["prompt"] for r in reversed(rows) if r["fn"] == "audit")
    # J2_SLIM=1 keeps the manifest of other atoms' agent_output beside the slice
    os.environ["J2_SLIM"] = "1"
    try:
        slim_run = docs_run(DESIGN_BRIEF, 900)
    finally:
        os.environ.pop("J2_SLIM", None)
    sp = [r["prompt"] for r in calls_file(slim_run) if r["fn"] in {w["id"] for w in slim_run.workflows if w.get("lld_sections")}]
    assert slim_run.outcome.startswith("Done") and sp and all("YOUR SPEC" in p and "OTHER SECTIONS" in p
                                                              and "THE WHOLE SITE FILE" not in p for p in sp[1:])
    # coverage gap: a plan that leaves S2 without a workflow is flagged before building (then the plan goes ahead)
    def drop_s2(run, orig):
        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            if fn == "adaptation" and isinstance(r, dict) and run.blueprint:
                gone = [w for w in r.get("workflows") or [] if w.get("lld_sections") == ["S2"]]
                r["workflows"] = [w for w in r.get("workflows") or [] if w.get("lld_sections") != ["S2"]]
                for w in r["workflows"]:           # the items S2's workflow covered move to S3's, so only the section is dropped
                    if gone and w.get("lld_sections") == ["S3"]:
                        w["covers"], w["parts"] = w["covers"] + gone[0]["covers"], w["parts"] + gone[0]["parts"]
            return r
        return call
    run = docs_run(DESIGN_BRIEF, 900, drop_s2)
    flag = [m["text"] for m in run.chat if m["who"] == "system" and "flagged the plan before building" in m["text"]]
    assert flag and "no workflow implements LLD section S2" in flag[0], flag
    dropped = run.audits[-1]["lld"]["dropped"]
    assert [d["section"] for d in dropped] == ["S2"] and "isn't on the page" in dropped[0]["evidence"], run.audits[-1]["lld"]
    assert run.outcome.startswith("Done"), run.outcome       # a dropped section is a warning, not a failure
    # a workflow that cites no LLD section is flagged too
    def uncite(run, orig):
        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            if fn == "adaptation" and isinstance(r, dict) and run.blueprint:
                for w in r.get("workflows") or []:
                    if w.get("lld_sections") == ["S3"]:
                        w["lld_sections"] = []
            return r
        return call
    run = docs_run(DESIGN_BRIEF, 900, uncite)
    assert any("cites no LLD section" in m["text"] for m in run.chat if m["who"] == "system" and "flagged" in m["text"])

    # a build that doesn't match its slice fails its atom: the same failure twice escalates to Adaptation
    def lose_button(run, orig):
        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            w = run._wf(fn) if fn.startswith("W-") else None
            if w and w.get("lld_sections") == ["S1"] and isinstance(r, str):
                return r.replace('id="act-1"', 'id="wrong-1"')
            return r
        return call
    run = docs_run(DESIGN_BRIEF, 900, lose_button)
    w1 = next(w for w in run.workflows if w.get("lld_sections") == ["S1"])
    bad = [r for r in run.receipts if r["atom"] == w1["atoms"][0]["id"] and r["type"] == "lld-slice"]
    assert bad and not bad[0]["pass"] and "E1 #act-1 isn't on the page" in bad[0]["evidence"], bad[:1]
    assert [r["attempt"] for r in bad][:2] == [1, 2] and any(m["text"].startswith(f"{w1['atoms'][0]['id']} failed its checks on attempt 2")
                                                              for m in run.chat if m["who"] == "system"), "no third attempt"
    # one revision: the builder reports the conflict, the LLD section is revised once with a version note and the reason
    def conflict(times):
        def patch(run, orig):
            seen = []

            def call(fn, fmt, task, mock, **k):
                r = orig(fn, fmt, task, mock, **k)
                w = run._wf(fn) if fn.startswith("W-") else None
                if w and w.get("lld_sections") == ["S1"] and isinstance(r, str) and len(seen) < times:
                    seen.append(1)
                    head, _, rest = r.partition("\n")
                    meta = json.loads(head)
                    meta["spec_conflict"] = {"section": "S1", "reason": "the button id clashes with the hero",
                                             "use": {"#act-1": "#go-1"}}
                    built = "go-1" if times == 1 else "oops-1"      # twice: the first build doesn't match the revised spec
                    return json.dumps(meta) + "\n" + rest.replace('id="act-1"', f'id="{built}"')
                return r
            return call
        return patch
    run = docs_run(DESIGN_BRIEF, 900, conflict(1))
    assert run.outcome.startswith("Done") and len(run.lld_revisions) == 1, (run.outcome, run.lld_revisions)
    rev = run.lld_revisions[0]
    assert rev["section"] == "S1" and rev["version"] == 2 and rev["reason"] == "the button id clashes with the hero"
    lld = open(os.path.join(run.dir, "lld.md"), encoding="utf-8").read()
    assert "#go-1" in lld and "`#act-1`" not in lld and "Version 2: the button id clashes with the hero" in lld
    assert run.docs["lld"]["agent_output"]["revisions"] == run.lld_revisions
    assert any(a["agent_output"].get("revisions") for w in run.past_workflows if "Pages and sections" in w["name"]
               for a in w["atoms"]), "the LLD atom carries the reason"
    w1 = next(w for w in run.workflows if w.get("lld_sections") == ["S1"])
    assert w1["atoms"][0]["status"] == "passed"
    # a second conflict on the same section escalates to Adaptation, naming the LLD's owner atom
    run = docs_run(DESIGN_BRIEF, 900, conflict(2))
    notes = [m["text"] for m in run.chat if m["who"] == "system"]
    esc = next((t for t in notes if "second conflict with LLD section S1" in t), None)
    assert esc and "Owner atom: W-13.a1" in esc and "already revised once" in esc, esc
    assert len(run.lld_revisions) == 1
    print("blueprint: slices as specs, coverage gap and uncited workflow flagged before building, mismatch fails the atom "
          "(no third attempt), one revision with version note and reason, second conflict escalates naming the owner "
          "atom, Audit compares the page to the LLD, chain recorded, J2_SLIM keeps the manifest ok")


def plain_unchanged(checkpoint=None, slim=None, board=None):
    """3: with no document asked for, the website flow is unchanged: same prompts, same page, same result as the
    checkpoint commit (run in a copy of it). checkpoint: the commit to compare with (default PLAIN_CHECKPOINT); slim: set
    J2_SLIM to this in both runs ("0" = the old whole-page path), else it is left unset (the default); board: set
    J2_BOARD to this in both runs (the checkpoint commit predates j2_board and ignores the env var either way)."""
    checkpoint = checkpoint or PLAIN_CHECKPOINT
    import subprocess
    import shutil
    here = os.path.dirname(os.path.abspath(__file__))
    probe = ("import json, os, sys, time, tempfile, hashlib\n"
             "os.environ.update(STUDIO_MOCK='1', J2_MEMORY='0', J2_RUNS_DIR=tempfile.mkdtemp())\n"
             "sys.argv = ['x']\n"
             "import re\n"
             "import j2\n"
             + inspect.getsource(fake_tap_targets) + inspect.getsource(fake_check) +
             "j2.j2_check.tap_targets, j2.j2_check.check, j2.j2_check.find_chrome = fake_tap_targets, fake_check, lambda: None\n"
             "r = j2.J2Run()\n"
             "r.user_says(sys.stdin.read())\n"
             "while r.status != 'waiting': time.sleep(0.02)\n"
             "rows = [json.loads(l) for l in open(os.path.join(r.dir, 'calls.jsonl'), encoding='utf-8')]\n"
             "h = lambda t: hashlib.sha256(re.sub(r'\\d{4}-\\d\\d-\\d\\dT\\d\\d:\\d\\d:\\d\\d', 'T', t).encode()).hexdigest()[:12]\n"
             "print(json.dumps({'outcome': r.outcome, 'spent': r.spent, 'site': h(r.sites[-1]), 'sections': h(json.dumps(r.sections, sort_keys=True)),\n"
             "  'calls': [[c['fn'], h(c['system']), h(c['prompt'])] for c in rows],\n"
             "  'workflows': h(json.dumps([{k: v for k, v in w.items() if k not in ('t',)} for w in r.workflows], sort_keys=True))}))\n")
    old = tempfile.mkdtemp(prefix="j2-checkpoint-")
    tar = subprocess.run(["git", "archive", checkpoint, "--", "j2.py", "j2_check.py", "llm.py", "conditions", "agents"]
                         + (["j2_schema.py"] if subprocess.run(["git", "cat-file", "-e", f"{checkpoint}:j2_schema.py"], cwd=here).returncode == 0 else []), cwd=here,
                         capture_output=True)
    if tar.returncode:
        print(f"plain brief unchanged: skipped (checkpoint {checkpoint} not in this clone)")
        shutil.rmtree(old, ignore_errors=True)
        return
    subprocess.run(["tar", "-x", "-C", old], input=tar.stdout, check=True)
    for name in ("j2_check.py",):
        pass
    try:
        results = []
        for cwd in (old, here):
            env = {**os.environ, "STUDIO_MOCK": "1", "J2_MEMORY": "0", "J2_FIXES": os.environ.get("J2_FIXES", "0")}
            env.pop("J2_SLIM", None)
            if slim is not None:
                env["J2_SLIM"] = slim
            if board is not None:
                env["J2_BOARD"] = board
            r = subprocess.run([sys.executable, "-c", probe], input=BREW_LAB, cwd=cwd, capture_output=True, text=True, env=env, timeout=300)
            assert r.returncode == 0, r.stderr[-800:]
            results.append(json.loads(r.stdout.strip().splitlines()[-1]))
        a, b = results
        assert a["outcome"] == b["outcome"] and a["spent"] == b["spent"] and a["site"] == b["site"] and a["sections"] == b["sections"], (a, b)
        assert [c[:2] for c in a["calls"]] == [c[:2] for c in b["calls"]], "different agents or system prompts"
        diff = [i for i, (x, y) in enumerate(zip(a["calls"], b["calls"])) if x != y]
        assert not diff, f"prompts differ at calls {diff[:5]}"
        assert a["workflows"] == b["workflows"]
        print(f"plain brief unchanged: {len(b['calls'])} calls, every system prompt and prompt, the page, the sections, the workflows "
              f"and the outcome are identical to checkpoint {checkpoint}" + (f" with J2_SLIM={slim}" if slim is not None else "")
              + (f" with J2_BOARD={board}" if board is not None else ""))
    finally:
        shutil.rmtree(old, ignore_errors=True)


def charter_crash_fix():
    """runs/j2-20261023-115255 crashed (TypeError: unhashable type: 'dict') because done_when held dicts, and a
    rejected-then-replaced charter doubled charter_items to 18 entries (9 strings from the rejected charter, 9
    dicts from the new one), from the same root cause: nothing normalised a charter's done_when items to strings
    before storing them. (a) a charter whose done_when items are dicts registers and runs past the validator,
    normalised to strings. (b) a second charter submitted before any workflow is registered renumbers from 1, not
    accumulated onto the first's."""
    FIRST = ["Problem section states the shop has no online booking", "Users section names students and staff",
             "Requirements section lists the classes with time and price"]
    SECOND = [{"item": 1, "criterion": "Problem section states the shop has no online booking"},
              {"item": 2, "criterion": "Users section names students and staff"}]

    def make_goal(items):
        return {"statement": "A one-page site for Still Point Yoga", "site_name": "Still Point Yoga",
                "audience": "Students", "standards": ["Clear"], "rules": ["Plain HTML"], "done_when": items,
                "core": [True] * len(items), "to_be_confirmed": [], "dropped_from_brief": [], "facts": [],
                "business_type": "studio", "visitor_needs": []}

    run = j2.J2Run()
    run.budget = 300
    run.request = "A one-page site for a yoga studio called Still Point Yoga."
    calls = {"n": 0}

    def call(fn, fmt, task, mock, **k):
        calls["n"] += 1
        return {"say": "@Adaptation Charter recorded.", "enough": True, "questions": [],
                "goal": make_goal(FIRST if calls["n"] == 1 else SECOND), "next": "adaptation", "reason": "ok"}
    run.call = call
    run._identity_intake({})
    assert run.goal["done_when"] == FIRST, run.goal["done_when"]
    assert run.goal["charter_items"] == FIRST, run.goal["charter_items"]
    assert not run.workflows and not run.past_workflows, "nothing registered yet: the next charter must renumber"
    run._identity_intake({})
    expect = [x["criterion"] for x in SECOND]
    assert run.goal["done_when"] == expect, run.goal["done_when"]                 # dicts normalised to strings
    assert all(isinstance(x, str) for x in run.goal["done_when"]), run.goal["done_when"]
    assert run.goal["charter_items"] == expect, run.goal["charter_items"]         # renumbered 1..2, not 1..5
    print("charter crash fix: dict done_when items normalise to strings and register past the validator; a charter "
          "replaced before any work is registered renumbers from 1 ok")


def board_disabled():
    """Message board step 1: with J2_BOARD unset (default off), a run writes no board.json or board.jsonl, and
    artifacts.json has no board row at all (absent, the same way a mock run's missing screenshots produce no row,
    not a placeholder saying "not recorded")."""
    import j2_artifacts as A
    os.environ.pop("J2_BOARD", None)
    run = docs_run(BREW_LAB, 200)
    assert run.outcome.startswith("Done"), run.outcome
    assert not os.path.isfile(os.path.join(run.dir, "board.json"))
    assert not os.path.isfile(os.path.join(run.dir, "board.jsonl"))
    rows = A.write(run.dir)["rows"]
    assert not any(r["type"] == "board" for r in rows), rows
    print("board disabled: J2_BOARD unset writes no board.json or board.jsonl, artifacts.json has no board row ok")


def board_story():
    """Message board, Part A: the main story. J2_FORCE_QUESTION=order makes the Order workflow's first attempt ask
    Products one question (item 2); Products answers from its own recorded section, with an artifact_update, and
    its atom gets "Updated by Q-1" (item 3); Filter reuses Q-1 without asking, recording a depends_on edge (item 4),
    and Audit reuses it too (the price topic is in the items it checks), while the Viewport fix, whose work shares no
    topic with the question, does not; no Adaptation re-plan follows (exactly the 5 proposed workflows run, no
    Fix/Cover-item extras). The switch is never in calls.jsonl, and a second question from the same atom is refused."""
    os.environ["J2_FORCE_QUESTION"] = "order"
    try:
        run = docs_run(BOUTIQUE_BRIEF, 200)
    finally:
        os.environ.pop("J2_FORCE_QUESTION", None)
    assert run.outcome.startswith("Done"), run.outcome
    names = [w["name"] for w in run._all_wfs()]
    assert names == ["Scaffold and styling", "Products", "Order on WhatsApp", "Filter", "Viewport fix"], names
    assert len(run.board["entries"]) == 1, run.board["entries"]
    q1 = run.board["entries"][0]
    assert q1["asked_by"] == "W-3" and q1["to"] == "W-2.a1" and q1["status"] == "answered", q1
    assert q1["answer"] and q1["answer"] != "not recorded", q1
    assert q1["artifact_update"], "Products' answer should have corrected its own section"
    products = next(w for w in run._all_wfs() if w["id"] == "W-2")
    assert "Updated by Q-1" in products["atoms"][0]["human_output"]["text"], products["atoms"][0]["human_output"]
    filt = next(w for w in run._all_wfs() if w["id"] == "W-4")
    assert any(e["atom"] == "W-2.a1" and e["why"] == "board Q-1" for e in run.given["W-4"]["edges"]), run.given["W-4"]
    assert "W-4.a1" in q1["used_by"], q1
    with open(os.path.join(run.dir, "chat.jsonl"), encoding="utf-8") as f:
        assert json.loads(f.readline())["force_question"] == "order"
    with open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8") as f:
        calls = [json.loads(ln) for ln in f]
    for c in calls:
        seen = (c["system"] + c["prompt"]).lower()
        for bad in ("j2_force_question", "test switch", "forced question"):
            assert bad not in seen, (c["id"], c["fn"], bad)
    order_wf = next(w for w in run._all_wfs() if w["id"] == "W-3")
    run.atom = order_wf["atoms"][0]["id"]        # simulate being inside W-3's own turn again
    second = run._ask_board(order_wf, {"to": "W-2.a1", "about": "x", "text": "y"}, 1)
    assert second["status"] == "refused" and "already asked" in (
        json.loads(open(os.path.join(run.dir, "board.jsonl"), encoding="utf-8").readlines()[-1])["reason"]), second
    assert "W-5.a1" not in q1["used_by"], ("the Viewport fix shares no topic with the question, so it may not reuse "
                                           "it in mock mode", q1["used_by"])
    assert q1["asked_by_role"] == "builder" and (q1["answered_from"] or {}).get("workflow") == "W-2", q1
    print("board story: J2_FORCE_QUESTION asks once, Products answers with an artifact_update and 'Updated by Q-1', "
          "Filter reuses it (depends_on 'board Q-1', used_by) and the Viewport fix does not, no re-plan, the switch "
          "never reaches an AI call, and a second question from the same atom is refused ok")


def board_run(brief=None, demo=False):
    """The message board demo run (Part B): the Thread & Loom brief with its prices and timings, built to Done, then
    the change request "Add a section for visiting hours". J2_FIXES=1, so the change request is checked at J2.a,
    which is where Identity asks the workflow that already built the timings. Returns the finished run."""
    fixes = os.environ.get("J2_FIXES")
    os.environ["J2_FIXES"] = "1"
    os.environ["J2_FORCE_QUESTION"] = "identity,priority,audit"
    try:
        run = j2.J2Run()
        run.budget = 260
        say(run, brief or BOUTIQUE_HOURS_BRIEF)
        finish(run)
        assert run.outcome.startswith("Done"), run.outcome
        say(run, "Add a section for visiting hours")
        finish(run)
    finally:
        os.environ.pop("J2_FORCE_QUESTION", None)
        os.environ["J2_FIXES"] = fixes if fixes is not None else "0"
    if demo:
        print(f"board demo run: {os.path.basename(run.dir)} ({run.outcome})")
    return run


def board_functions():
    """Message board, Part B: the whole story with the functions asking. Thread & Loom is built to Done, then the
    Requester asks for a section of visiting hours, and (J2_FORCE_QUESTION=identity,priority,audit) exactly three
    questions are asked, each once: Q-1 Identity asks the workflow that built the timings where they are; Q-2
    Priority asks Adaptation whether the change is new work or a fix; Q-3 Audit asks the products workflow which
    attribute holds the price. Adaptation reuses Q-1 and plans a revise of that workflow's section priced as a fix,
    the revise builder reuses Q-1, and on the next audit Audit reads its own Q-3 back off the board instead of
    asking again. A question the board already answers is refused ("already answered in Q-3"). The run ends Done,
    N6 (no invented content) passes, and the switch never reaches an AI call."""
    run = board_run()
    assert run.outcome.startswith("Done"), run.outcome
    b = run.board["entries"]
    assert [e["id"] for e in b] == ["Q-1", "Q-2", "Q-3"], b                      # exactly 3, in order
    assert [e["asked_by"] for e in b] == ["identity", "priority", "audit"], b
    assert all(e["asked_by_role"] == "function" and e["status"] == "answered" for e in b), b
    assert len({e["question"] for e in b}) == 3, b                               # each asked once
    timings = b[0]["to"].split(".")[0]
    assert b[0]["to"].startswith("W-") and b[0]["answered_from"]["workflow"] == timings, b[0]
    assert isinstance(b[0]["answered_from"]["version"], int), b[0]               # the version it had on record
    assert "item-" in b[0]["answer"] and "timings" in b[0]["answer"], b[0]       # read off its recorded section
    assert b[1]["to"] == "adaptation" and b[1]["answered_from"] == {"function": "adaptation"}, b[1]
    assert b[1]["answer"].lower().startswith("a fix"), b[1]
    assert "data-price" in b[2]["answer"], b[2]
    # Reuse: Adaptation and the revise builder read Q-1, Audit reads its own Q-3 back on the next audit
    assert sum(len(e["used_by"]) for e in b) >= 3, b
    assert "adaptation" in b[0]["used_by"] and any(r.startswith("W-") for r in b[0]["used_by"]), b[0]
    assert "audit" in b[2]["used_by"], b[2]
    rev = next(w for w in run._all_wfs() if w["name"] == "Visiting hours")
    assert rev["revises"] == timings and rev["cost"] == j2.RATES["fix"], rev     # a revise, priced as a fix
    chat = [m for m in run.chat if m["who"] in ("identity", "priority", "audit", "adaptation")]
    assert any(m["who"] == "identity" and "Q-1:" in m["text"] for m in chat), "the ask is in the chat"
    assert any(m["who"] == "adaptation" and "Q-2:" in m["text"] for m in chat), "the answer is in the chat"
    assert any(m["who"] == "adaptation" and "[read Q-1]" in m["text"] for m in chat), "a function's read tag"
    assert any(m["who"] == "audit" and "[read" in m["text"] and "Q-3" in m["text"] for m in chat), "Audit's read tag"
    # A question the board already answers is refused, whoever asks it
    dup = run._ask_board("coordination", {"to": b[2]["to"], "about": "the price attribute",
                                          "text": "Which attribute holds the price?"}, 1, role="function")
    assert dup["status"] == "refused", dup
    with open(os.path.join(run.dir, "board.jsonl"), encoding="utf-8") as f:
        assert json.loads(f.readlines()[-1])["reason"] == "already answered in Q-3"
    # and a function asks at most once in a turn (the turn it asked in is the entry's attempt)
    cycle, run.cycle = run.cycle, b[2]["cycle"]
    again = run._ask_board("audit", {"to": b[0]["to"], "about": "the nav", "text": "Which section holds the nav?"},
                           b[2]["attempt"], role="function")
    run.cycle = cycle
    assert again["status"] == "refused", again
    with open(os.path.join(run.dir, "board.jsonl"), encoding="utf-8") as f:
        assert "per turn" in json.loads(f.readlines()[-1])["reason"]
    with open(os.path.join(run.dir, "conditions.json"), encoding="utf-8") as f:
        n6 = next(c for c in json.load(f) if c["id"] == "N6")
    assert n6["met"], n6                      # no invented content: the mock built from the brief only
    with open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8") as f:
        for c in (json.loads(ln) for ln in f):
            seen = (c["system"] + c["prompt"]).lower()
            for bad in ("j2_force_question", "test switch", "identity,priority,audit"):
                assert bad not in seen, (c["id"], c["fn"], bad)
    only_identity_addresses_user(run)         # the board's own messages never reach the Requester
    board_tab_render(run, b[:3], timings)
    print("board functions: Identity, Priority and Audit each ask once (a workflow, a function and an earlier "
          "workflow answer), Adaptation and the revise builder reuse Q-1, Audit reads its own Q-3 back on the next "
          "audit, a duplicate is refused, N6 passes, the run ends Done and the switch never reaches an AI call ok")


def board_tab_render(run, entries, timings):
    """Item 5, browser-free (j2_ui_render.js): the Board tab shows every entry with who asked and in what role, and
    where the answer came from (a workflow with its version, or a function). The atom card of the workflow that
    answered a function's question says so, since the asker is no atom."""
    import shutil
    if not shutil.which("node"):
        print("board tab render skipped: node not found")
        return
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
    out = render_node(snap, docs={"docs": {}, "board": entries})
    tab = out["board"]
    assert "<th>Answered from</th>" in tab and tab.count('<tr data-board=') == len(entries), tab[:400]
    assert tab.count(">function<") == len(entries), "every entry says a function asked it"
    assert "Adaptation (function)" in tab, tab          # Q-2 was answered by a function
    assert re.search(rf"{timings} v\d+", tab), tab      # Q-1 by a workflow, with the version it had on record
    card = out["cards"][timings]
    line = card.split('data-note="board">')[1].split("</span>")[0] if 'data-note="board">' in card else "(no board line)"
    assert "Answered Q-1, asked by Identity (function)" in line, line
    print("board tab: every entry shows its asker and role and where the answer came from, and the answering atom's "
          "card names the function that asked ok")


            # the demo brief of the message board's document changes: runs/j2-20261026-114456's first message
YOGA_BRIEF = (
    "Create the design documents (PRD, HLD and LLD) and then build the website for Still Point Yoga, a small yoga "
    "studio at 4th Floor, 22 Church Street, Bengaluru 560001.\n\nThe page shows the name Still Point Yoga and the "
    "address as text, with no embedded map.\n\nThe weekly class list has five classes, each with day, time and price "
    "in rupees:\nMonday 7:00 am Gentle Flow 500\nTuesday 6:30 pm Strength Yoga 600\nThursday 7:00 am Breath and "
    "Stretch 450\nSaturday 9:00 am Beginners Basics 400\nSunday 8:00 am Slow Sunday 500\n\nThe page has a class "
    "booking form with the fields name, email, class and preferred date. The form sends the booking to "
    "bookings@stillpoint.example. Payment is made at the studio. No online payment is taken.\n\nThe teachers' names "
    "are not decided yet, and the studio has no phone number. Do not show teacher names or a phone number, and do "
    "not make any other detail up.\n\nThe page must work on phones. The documents and the website must agree with "
    "each other.")


def doc_board_run(force="lld", patch=None, budget=1400, demo=False):
    """The document-change demo run: Still Point Yoga, the PRD, HLD and LLD, then the website, with
    J2_FORCE_QUESTION=<force> so one builder asks the document it was built from."""
    os.environ["J2_FORCE_QUESTION"] = force
    try:
        run = docs_run(YOGA_BRIEF, budget, patch=patch)
        finish(run)
    finally:
        os.environ.pop("J2_FORCE_QUESTION", None)
    if demo:
        print(f"board doc demo run: {os.path.basename(run.dir)} ({run.outcome})")
    return run


def board_doc_change():
    """Message board, part C (a): the LLD's Interactive elements builder asks the HLD's Components atom how the
    booking is sent; the HLD owner answers from the brief's address and corrects its own section. Code applies it
    (only the owner may), bumps the HLD's version, re-runs H1-H6 on the new text, re-checks every atom the recorded
    edges reach (a receipt with "trigger": "Q-1" for each), and the run still ends Done."""
    run = doc_board_run()
    assert run.outcome.startswith("Done"), run.outcome
    assert [e["id"] for e in run.board["entries"]] == ["Q-1"], run.board["entries"]
    e = run.board["entries"][0]
    assert e["asked_by"] == "W-14" and e["to"] == "W-9.a1" and e["status"] == "answered", e
    assert "mailto" in e["answer"] and "bookings@stillpoint.example" in e["answer"], e
    up = e["artifact_update"]
    assert up == {"doc": "hld", "section": "Components", "from_version": up["from_version"],
                  "to_version": up["from_version"] + 1}, up
    # the document itself: the new text, the version note, the export and the records that read it
    hld = open(os.path.join(run.dir, "hld.md"), encoding="utf-8").read()
    comps = j2.md_section(hld, "Components")
    assert "mailto to bookings@stillpoint.example" in comps, comps[-300:]
    assert f"Version {up['to_version']}: Q-1," in comps, comps[-300:]
    assert os.path.isfile(os.path.join(run.dir, "hld.docx"))
    assert os.path.isfile(os.path.join(run.dir, "site", f"hld-v{up['to_version']}.md"))
    assert "mailto" in json.dumps(run.docs["hld"]["agent_output"]), "documents.json carries the new components"
    conds = json.load(open(os.path.join(run.dir, "conditions-hld.json"), encoding="utf-8"))
    assert [c["id"] for c in conds] == ["H1", "H2", "H3", "H4", "H5", "H6"] and all(c["met"] for c in conds), conds
    # the impact: recorded edges only, each with its own re-check receipt
    impact = {x["atom"]: x["result"] for x in e["impact"]}
    assert impact and all(r == "passed" for r in impact.values()), e["impact"]
    assert "W-13.a1" in impact, "the LLD atoms that used HLD Components are re-checked"
    assert any(a.startswith(("W-19", "W-20", "W-21", "W-22", "W-23")) for a in impact), "the website atoms too"
    recs = [json.loads(ln) for ln in open(os.path.join(run.dir, "receipts.jsonl"), encoding="utf-8")]
    triggered = {r["atom"] for r in recs if r.get("trigger") == "Q-1"}
    assert set(impact) | {"W-9.a1"} <= triggered, (sorted(triggered), sorted(impact))
    rechecks = [a for a in json.load(open(os.path.join(run.dir, "audits.json"), encoding="utf-8"))
                if a.get("kind") == "recheck"]
    assert len(rechecks) == 1 and rechecks[0]["trigger"] == "Q-1" and rechecks[0]["doc"] == "hld" \
        and rechecks[0]["version"] == up["to_version"] and rechecks[0]["results"], rechecks
    assert all(r["pass"] for r in rechecks[0]["results"]), rechecks[0]
    atoms = {a["id"]: a for a in json.load(open(os.path.join(run.dir, "atoms.json"), encoding="utf-8"))["atoms"]}
    assert atoms["W-9.a1"]["rechecked"] == "Q-1" and "Updated by Q-1" in atoms["W-9.a1"]["notes"], atoms["W-9.a1"]
    assert all(atoms[a]["rechecked"] == "Q-1" for a in impact), [a for a in impact if not atoms[a].get("rechecked")]
    assert not any(a.get("stale") for a in atoms.values()), "a re-checked atom is not stale"
    rows = {r["type"]: r for r in json.load(open(os.path.join(run.dir, "artifacts.json"), encoding="utf-8"))["rows"]
            if r["kind"] == "type"}
    assert rows["hld"]["changed_by"] == ["Q-1"] and rows["hld"]["version"] == up["to_version"], rows["hld"]
    assert rows["lld"]["rechecked_after"] == ["Q-1"] and rows["website"]["rechecked_after"] == ["Q-1"], rows
    m = json.load(open(os.path.join(run.dir, "map.json"), encoding="utf-8"))
    assert {"from": "work:lld", "to": "work:hld", "kind": "asks", "label": "Q-1", "n": None,
            "basis": "recorded", "source": {"file": "board.json", "field": "entries", "value": "Q-1"}} in m["edges"], [x for x in m["edges"] if x["kind"] == "asks"]
    hld_node = next(n for n in m["nodes"] if n["id"] == "work:hld")
    assert hld_node["version"] == up["to_version"] and "re-checked after Q-1" in hld_node["detail"], hld_node
    line = next((x for x in run.chat if x["who"] == "system" and f"HLD v{up['from_version']} -> "
                 f"v{up['to_version']} (Q-1)" in x["text"]), None)
    assert line and "Re-checked:" in line["text"] and "sent to Adaptation" in line["text"], line
    assert any(m["who"] == "W-14" and "Q-1:" in m["text"] for m in run.chat), "the ask is in the chat"
    assert any(m["who"] == "W-9" and "(changed the HLD)" in m["text"] for m in run.chat), "the answer is in the chat"
    import shutil
    if shutil.which("node"):
        snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
        out = render_node(snap, docs={"docs": {"hld": {"text": hld, "asked": True, "workflows": ["W-9"]}}})
        assert f"Version {up['to_version']}: Q-1," in out["docs"], "the Documents tab shows the version note"
        for bad in ("Q-1", "re-check", "HLD v", "board"):
            assert bad not in out["simple"], (bad, "the Simple view is unchanged by a board change")
    print("board doc change: the LLD builder asks, the HLD owner corrects its own section, the HLD is re-exported at "
          f"v{up['to_version']} with its version note, H1-H6 pass again, every atom the recorded edges reach is "
          "re-checked with a Q-1 receipt, the records and the map carry it, and the run ends Done ok")


def board_doc_not_owner():
    """Message board, part C (b): a change to a section the answering atom doesn't own is refused with the reason,
    and the document is untouched."""
    run = doc_board_run()
    before = open(os.path.join(run.dir, "hld.md"), encoding="utf-8").read()
    run.call = lambda fn, fmt, task, mock, **k: {
        "answer": "The components are in the HLD.",
        "artifact_update": {"doc": "hld", "section": "Components", "text": "- **C1** anything.", "reason": "mine now"}}
    e = run._ask_board("audit", {"to": "W-12.a1", "about": "the components", "text": "Who owns the component list?"},
                       1, role="function")
    assert e["status"] == "refused" and e["artifact_update"] is None, e
    with open(os.path.join(run.dir, "board.jsonl"), encoding="utf-8") as f:
        assert json.loads(f.readlines()[-1])["reason"] == "W-12.a1 doesn't own the HLD's Components section"
    assert open(os.path.join(run.dir, "hld.md"), encoding="utf-8").read() == before, "the HLD is untouched"
    print("board doc not owner: a change to a section the answerer doesn't own is refused and nothing is written ok")


def board_doc_prd():
    """Message board, part C (c): the HLD's Components builder asks the PRD's Requirements atom which requirement
    covers the booking form. A PRD change is never applied from the board: the entry records that it went to
    Identity, the PRD file is unchanged, and Identity gets it on the existing path."""
    run = doc_board_run(force="hld")
    assert run.outcome.startswith("Done"), run.outcome
    e = run.board["entries"][0]
    assert e["asked_by"] == "W-9" and e["to"] == "W-4.a1" and e["status"] == "answered", e
    assert e["artifact_update"] is None and e["sent_to_identity"]["doc"] == "prd", e
    assert e["sent_to_identity"]["section"] == "Requirements", e
    assert not e.get("impact"), "nothing is propagated from a change that was not applied"
    prd = open(os.path.join(run.dir, "prd.md"), encoding="utf-8").read()
    assert "Version " not in prd and not os.path.isfile(os.path.join(run.dir, "conditions-prd.json")), "the PRD is unchanged"
    assert not [a for a in json.load(open(os.path.join(run.dir, "audits.json"), encoding="utf-8"))
                if a.get("kind") == "recheck"], "no re-check: nothing changed"
    assert any(m["who"] == "system" and "never applied from the board" in m["text"] for m in run.chat), "the chat says so"
    calls = [json.loads(ln) for ln in open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8")]
    asked = [c for c in calls if c["fn"] == "identity" and "Q-1: the PRD's Requirements section needs changing" in c["prompt"]]
    assert len(asked) == 1, "Identity got the change request on the existing path, once"
    assert any(m["who"] == "identity" and m["text"].startswith("@W-4 Builder") for m in run.chat), \
        "and answered the atom that asked for the change"
    print("board doc prd: a PRD change from the board is not applied; it goes to Identity as a change request and the "
          "PRD file stays as it was ok")


def board_doc_no_edge():
    """Message board, part C (d): a change to a section no recorded edge leads away from propagates to nothing: the
    entry's impact is [] and no other atom is re-checked or marked."""
    run = doc_board_run()
    v0 = run._doc_version("lld")
    rules = j2.md_section(open(os.path.join(run.dir, "lld.md"), encoding="utf-8").read(),
                          "Responsive and accessibility rules")
    run.call = lambda fn, fmt, task, mock, **k: {
        "answer": "Every field has a visible label, as recorded.",
        "artifact_update": {"doc": "lld", "section": "Responsive and accessibility rules",
                            "text": rules.strip() + "\n- Every control keeps its visible label.",
                            "reason": "the label rule was not explicit"}}
    before = {a["id"]: dict(a) for w in run._all_wfs() for a in w.get("atoms") or [] if a["id"] != "W-17.a1"}
    e = run._ask_board("audit", {"to": "W-17.a1", "about": "the label rule", "text": "Where is the label rule kept?"},
                       1, role="function")
    assert e["status"] == "answered" and e["artifact_update"]["to_version"] == v0 + 1, e
    assert e["impact"] == [], e["impact"]
    assert all(before[a["id"]] == a for w in run._all_wfs() for a in w.get("atoms") or []
               if a["id"] in before), "no other atom is touched"
    print("board doc no edge: a change the recorded edges lead away from nowhere propagates to nothing (impact []) ok")


def board_doc_recheck_fails():
    """Message board, part C (e): the owner's own checks fail after its change, so code hands the atom to Adaptation
    on the existing atom_failed path (a revise priced as a fix) instead of leaving a stale document behind."""
    def patch(run, orig):
        def call(fn, fmt, task, mock, **k):
            if "ASKS YOU" in task and k.get("answer_only"):
                return {"answer": "mailto, no server.",
                        "artifact_update": {"doc": "hld", "section": "Components", "text": "To be decided.",
                                            "reason": "the component list is not settled"}}
            return orig(fn, fmt, task, mock, **k)
        return call
    run = doc_board_run(patch=patch)
    e = run.board["entries"][0]
    assert e["status"] == "answered" and e["artifact_update"]["doc"] == "hld", e
    rechecks = [a for a in json.load(open(os.path.join(run.dir, "audits.json"), encoding="utf-8"))
                if a.get("kind") == "recheck"]
    assert rechecks and any(r["pass"] is False and r["atom"] == "W-9.a1" for r in rechecks[0]["results"]), rechecks
    assert any(m["who"] == "system" and "failed after the board change Q-1" in m["text"] for m in run.chat), \
        [m["text"] for m in run.chat if m["who"] == "system"][-3:]
    assert any(m.get("act") in ("CHECK_WRONG", "BUILD_WRONG") and m["who"] == "adaptation" for m in run.chat), \
        "Adaptation read the receipts and decided"
    atoms = {a["id"]: a for a in json.load(open(os.path.join(run.dir, "atoms.json"), encoding="utf-8"))["atoms"]}
    assert atoms["W-9.a1"].get("stale") == "Q-1" or atoms["W-9.a1"].get("rechecked") == "Q-1", atoms["W-9.a1"]
    print("board doc recheck fails: a change that breaks the owner's own checks goes to Adaptation on the existing "
          "atom_failed path, and the atom stays marked ok")


def board_doc_off():
    """Message board, part C (f): with J2_BOARD unset the same brief and switch run byte for byte the same (every
    call, the documents and the outcome), and no board or re-check record is written."""
    os.environ.pop("J2_BOARD", None)
    plain, switched = doc_board_run(force=""), doc_board_run(force="lld")
    for run in (plain, switched):
        assert run.outcome.startswith("Done"), run.outcome
        assert not run.board["entries"] and not run.doc_changes and not run.rechecks
        for name in ("board.json", "board.jsonl", "conditions-hld.json"):
            assert not os.path.isfile(os.path.join(run.dir, name)), name
        assert not [a for a in json.load(open(os.path.join(run.dir, "audits.json"), encoding="utf-8"))
                    if a.get("kind") == "recheck"]
        rows = json.load(open(os.path.join(run.dir, "artifacts.json"), encoding="utf-8"))["rows"]
        assert not any("changed_by" in r or "rechecked_after" in r for r in rows), rows
        atoms = json.load(open(os.path.join(run.dir, "atoms.json"), encoding="utf-8"))["atoms"]
        assert not any("stale" in a or "rechecked" in a for a in atoms)

    def fingerprint(run):
        calls = [json.loads(ln) for ln in open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8")]
        h = lambda t: hashlib.sha256(re.sub(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d", "T", t).encode()).hexdigest()[:12]
        return ([[c["fn"], h(c["system"]), h(c["prompt"])] for c in calls], run.outcome,
                [h(open(os.path.join(run.dir, f"{k}.md"), encoding="utf-8").read()) for k in ("prd", "hld", "lld")])
    a, b = fingerprint(plain), fingerprint(switched)
    assert a[1:] == b[1:] and a[0] == b[0], f"{len(a[0])} vs {len(b[0])} calls; first difference: " + str(
        next((i for i, (x, y) in enumerate(zip(a[0], b[0])) if x != y), None))
    for c in (json.loads(ln) for ln in open(os.path.join(switched.dir, "calls.jsonl"), encoding="utf-8")):
        for bad in ("BOARD (", "board_question", "board_used", "artifact_update"):
            assert bad not in c["system"] + c["prompt"], (c["id"], c["fn"], bad)
    print(f"board doc off: with J2_BOARD unset the switch changes nothing: {len(a[0])} calls, the three documents and "
          "the outcome are identical, and no board or re-check record is written ok")


def board_functions_off():
    """Message board, Part B: the same brief and change request with J2_BOARD unset write no board file, ask
    nothing, and no prompt carries the board fields or the BOARD block (the flag gates all of it; plain_unchanged
    guards the prompts byte for byte)."""
    os.environ.pop("J2_BOARD", None)
    run = board_run()
    assert run.outcome.startswith("Done"), run.outcome
    assert not run.board["entries"] and not os.path.isfile(os.path.join(run.dir, "board.json"))
    assert not os.path.isfile(os.path.join(run.dir, "board.jsonl"))
    with open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8") as f:
        for c in (json.loads(ln) for ln in f):
            for bad in ("BOARD (", "board_question", "board_used"):
                assert bad not in c["system"] + c["prompt"], (c["id"], c["fn"], bad)
    print("board off: the same brief and change request with J2_BOARD unset write no board file, ask nothing, and "
          "no prompt mentions the board ok")


def board_not_mine():
    """Message board, Part A: "not mine" (item 3). The target's reply isn't its area, so the entry is unanswered and
    the asking workflow continues normally (the existing Adaptation path; no follow-up question, no crash)."""
    os.environ["J2_FORCE_QUESTION"] = "order"

    def patch(run, orig):
        def call(fn, fmt, task, mock, **k):
            if "ASKS YOU" in task:
                return {"answer": "not mine", "artifact_update": None}
            return orig(fn, fmt, task, mock, **k)
        return call
    try:
        run = docs_run(BOUTIQUE_BRIEF, 200, patch=patch)
    finally:
        os.environ.pop("J2_FORCE_QUESTION", None)
    assert run.outcome.startswith("Done"), run.outcome
    assert len(run.board["entries"]) == 1
    q1 = run.board["entries"][0]
    assert q1["status"] == "unanswered" and not q1["answer"] or q1["answer"].lower() == "not mine", q1
    assert q1["artifact_update"] is None and q1["used_by"] == [], q1
    print("board not mine: an unanswered question leaves the asking workflow on the existing path, no follow-up ok")


def docs_views():
    """4: Detailed view: document atoms get the eight-block card with a preview and .md/.docx links. Simple view: plain
    words for the finished set; a plain brief's Simple view is identical to the checkpoint's."""
    import subprocess
    import shutil
    import threading
    import urllib.request
    import html as htmllib
    run = docs_run(DESIGN_BRIEF, 900)
    assert run.outcome.startswith("Done")
    snap = j2.saved_snapshot(os.path.basename(run.dir), 0)
    assert snap["docset"] == ["prd", "hld", "lld"] and [d["kind"] for d in snap["docset_files"]] == ["prd", "hld", "lld"]
    assert snap["files"] == ["prd.md", "prd.docx", "hld.md", "hld.docx", "lld.md", "lld.docx"], snap["files"]
    here = os.path.dirname(os.path.abspath(__file__))
    srv = None
    import j2_server
    keep, j2_server.RUN = j2_server.RUN, run
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    try:
        for k in ("prd", "hld", "lld"):
            for ext in ("md", "docx"):
                for q in ("", f"?run={os.path.basename(run.dir)}"):
                    r = urllib.request.urlopen(f"{base}/file/{k}.{ext}{q}")
                    body = r.read()
                    assert (body[:2] == b"PK") == (ext == "docx") and (ext == "docx" or body.startswith(b"# ")), (k, ext, q)
        doc_atom = next(a for w in run._all_wfs() for a in w["atoms"] if "Components" in w["name"])
        got = json.load(urllib.request.urlopen(f"{base}/api/preview?atom={doc_atom['id']}"))
        assert got["html"] and "<h2>Components</h2>" in got["html"] and "C1" in got["html"], got
        saved = json.load(urllib.request.urlopen(f"{base}/api/preview?atom={doc_atom['id']}&run={os.path.basename(run.dir)}"))
        assert saved["html"] == got["html"]
    finally:
        j2_server.RUN = keep
        srv.shutdown()
    import shutil as _sh
    if not _sh.which("node"):
        print("document views (render) skipped: node not found")
        return
    previews = {}
    for w in run._all_wfs():
        for a in w["atoms"]:
            body = j2.atom_section_body(run.dir, run._all_wfs(), a["id"])
            previews[a["id"]] = j2.md_html(body, "Section") if body else ""
    out = render_node(snap, previews)
    cards = 0
    for w in snap["workflows"]:
        if w["id"] not in ("W-1", "W-7", "W-13"):
            continue
        card = out["cards"][w["id"]]
        a = w["atoms"][0]
        card = card[card.index(f'data-key="atom-{a["id"]}"'):].split("Run details")[0]
        assert re.findall(r'data-block="([a-z_]+)"', card) == ["goal", "how", "constraints", "context", "expected_output",
                                                              "verification", "human_output", "agent_output", "result"], w["id"]
        pre = card.split('data-block="result"')[1]
        assert 'data-preview="frame"' in pre and "sandbox" in pre and "srcdoc=" in pre, pre[:200]
        src = htmllib.unescape(re.search(r'srcdoc="([^"]*)"', pre).group(1))
        assert src == previews[a["id"]] and "<h2>" in src
        base_name = a["agent_output"]["file"][:-3]
        assert f'data-link="md" href="/file/{base_name}.md"' in card and f'data-link="docx" href="/file/{base_name}.docx"' in card
        assert 'data-link="site"' not in card and 'data-link="agent"' in card and 'data-link="evidence"' in card
        cards += 1
    assert cards == 3
    # Simple view: the finished set in plain words, no ids or check codes
    text = re.sub(r"<[^>]+>", " ", htmllib.unescape(out["simple"]))
    assert "Your design documents are ready" in text and "Download the product requirements (markdown)" in text \
        and "Download the low-level design (Word file)" in text, text[-500:]
    bad = SIMPLE_FORBIDDEN.findall(text) + SIMPLE_FORBIDDEN_CASE.findall(text) + re.findall(r"\b(?:[PHLN]\d|R\d|W-\d+)\b", text)
    assert not bad, bad
    # a plain brief's Simple view is identical to the checkpoint's page, on the same state
    plain = docs_run(BREW_LAB, 200)
    ps = j2.saved_snapshot(os.path.basename(plain.dir), 0)
    old = subprocess.run(["git", "show", f"{CHECKPOINT}:j2_ui.html"], capture_output=True, text=True, cwd=here)
    if old.returncode == 0:
        p = os.path.join(tempfile.mkdtemp(prefix="j2-old-ui-"), "j2_ui_old.html")
        open(p, "w", encoding="utf-8").write(old.stdout)
        before, after = render_node(ps, html=p), render_node(ps)
        assert before["simple"] == after["simple"] and before["sprog"] == after["sprog"] and after["simple"] and "ready" in after["simple"]
        # and a PRD-only run keeps its own ready card
        prd = docs_run("Write a PRD for Campus Notes." + CAMPUS, 200)
        ss = j2.saved_snapshot(os.path.basename(prd.dir), 0)
        assert render_node(ss, html=p)["simple"] == render_node(ss)["simple"]
    print("document views: .md and .docx downloads (live and saved), preview of a document atom's section, eight-block "
          "cards with download links, Simple view says 'Your design documents are ready' in plain words, a plain "
          "brief's Simple view identical to the checkpoint's ok")



def stage_question(run):
    return next(m for m in reversed(run.chat) if m.get("act") == "ASK_REQUESTER")["text"]


def docs_funding():
    """The work credits can't cover every document and the site: nothing is skipped by code. The existing decision
    question (add credits, leave out, finish as it is) lists each as an item; the Requester's answer is what is applied."""
    # at the start: asked before anything is planned or built
    run = docs_run(DESIGN_BRIEF, 120)
    assert run.paused and run.paused["kind"] == "limit" and not run.workflows and not run.docs, run.outcome
    assert run.stages == ["prd", "hld", "lld", "website"] and run.docset == ["prd", "hld", "lld"], run.stages
    text = stage_question(run)
    for item in ("the product requirements", "the high-level design", "the low-level design", "the website"):
        assert item in text, (item, text)
    assert text.count("\n- ") == 4 and text.count("needs at least") == 5 and "not covered" in text and "Still not right" not in text, text
    opts = run.paused["questions"][0]["options"]
    assert opts[0].startswith("Add ") and opts[1].startswith("Leave out the high-level design, the low-level design and the website") \
        and len(opts) == 2, opts            # nothing is written yet, so "finish as it is" isn't offered
    assert not any("stays funded" in m["text"] or "is left out so" in m["text"] for m in run.chat), "silent skip"
    body = text.replace("@Requester ", "", 1)
    assert not SIMPLE_FORBIDDEN.search(body) and not re.search(r"\b(?:HLD|LLD|PRD|W-\d+)\b", body), body
    # leave out what doesn't fit: the PRD alone is made, the decision is recorded
    run.budget = 120
    say(run, "Option 2")
    assert run.outcome.startswith("Done") and run.stages == ["prd"] and list(run.docs) == ["prd"], (run.outcome, run.stages)
    assert run.decisions[-1]["decision"] == "drop" and run.decisions[-1]["left_out"] == ["hld", "lld", "website"]
    assert any(m["text"].startswith("Left out on the Requester's decision") for m in run.chat if m["who"] == "system")
    # add the credits, keep going: asked again between documents when the rest still doesn't fit
    run = docs_run(DESIGN_BRIEF, 100)
    assert run.paused and "Add 100 work credits (the most allowed)" in stage_question(run)
    say(run, "Option 1")
    assert run.budget == 200 and run.paused and run.stage_kind == "prd" and list(run.docs) == ["prd"], (run.budget, run.outcome)
    text = stage_question(run)
    assert "the low-level design (needs at least 88 work credits; not covered)" in text and "the high-level design (needs at least 98 work credits)" in text, text
    opts = run.paused["questions"][0]["options"]
    assert [o.split(":")[0] for o in opts] == ["Leave out the low-level design and the website", "Finish as it is"] and "hld.md" not in os.listdir(run.dir)
    assert not any("is left out so" in m["text"] for m in run.chat)
    # finish as it is: what is written stays, the rest is named as not made
    fin = docs_run(DESIGN_BRIEF, 100)
    say(fin, "Option 1")
    say(fin, "Option 2")
    assert fin.outcome.startswith("Done with open items") and "Not made: the high-level design, the low-level design and the website" in fin.outcome, fin.outcome
    assert "prd.md" in os.listdir(fin.dir) and "hld.md" not in os.listdir(fin.dir)
    # leave out what doesn't fit (at the calibrated minimums the LLD and the site): the HLD is made, the chain ends
    say(run, "Option 1")
    assert run.stages == ["prd", "hld"] and not run.paused and run.outcome.startswith("Done") and list(run.docs) == ["prd", "hld"], \
        (run.stages, run.outcome)
    # enough credits: no question at all
    run = docs_run(DESIGN_BRIEF, 900)
    assert not any(m.get("act") == "ASK_REQUESTER" for m in run.chat) and run.outcome.startswith("Done")
    print("funding: short credits ask the decision question with each document and the site as items, nothing is "
          "skipped by code; leave out, add and keep going, finish as it is, and no question when the credits cover it ok")


def docs_repeat_failure():
    """The same document check fails twice in a chain: "Finish as it is" accepts that document and the chain goes on,
    "Stop the whole run here" is its own option, the words are about a document (never "site" or "missing parts"), and
    Identity's message states check results from code (never "passes" while a necessary condition is unmet)."""
    orig_cm = j2.condition_met
    holder = []

    def failing_p4(cond, *a, **k):
        if holder and cond.get("check") == "no-invented-facts" and holder[0].stage_kind == "prd":
            return False, "not in the charter or the Requester's answers, and not marked Assumption: This"
        return orig_cm(cond, *a, **k)

    def lying(run, orig):
        def call(fn, fmt, task, mock, **k):
            r = orig(fn, fmt, task, mock, **k)
            if fn == "identity" and "LIMIT:" in task:
                r["say"] = "The document has all its parts, and each one has been checked and passes."
            return r
        return call

    def start(text, budget=900):
        run = j2.J2Run()
        run.budget = budget
        run.call = lying(run, run.call)
        holder[:] = [run]
        say(run, text)
        return run
    j2.condition_met = failing_p4
    try:
        brief = "Design docs for Campus Notes." + CAMPUS
        run = start(brief)
        assert run.paused and run.paused["kind"] == "limit" and run.stage_kind == "prd", run.outcome
        h, q = run.paused["handoff"], run.paused["questions"][0]
        assert "ship" in h["choices"] and "stop" in h["choices"], h["choices"]
        text = stage_question(run) + json.dumps(q)
        # item 2: the message is code's, with counts, and says which condition is unmet; no pass claim
        assert re.search(r"\d+ of \d+ required conditions are met", text) and "Not met: No invented facts" in text, text
        assert not re.search(r"checked and passes|all its parts", text), text
        # item 3: words about a document
        assert not re.search(r"\b(?:web)?site\b|missing parts", text, re.I), text
        ship = next(o for o in q["options"] if o.startswith("Finish as it is"))
        assert "the product requirements" in ship and "goes on with the high-level design and the low-level design" in ship, ship
        assert any(o.startswith("Stop the whole run here") and "not made" in o for o in q["options"]), q["options"]
        say(run, str(h["choices"].index("ship") + 1))        # finish as it is: this document is accepted, the chain goes on
        assert list(run.docs) == ["prd", "hld", "lld"] and run.stage_kind == "lld", (list(run.docs), run.stage_kind, run.outcome)
        assert run.outcome.startswith("Done with open items") and "the product requirements" in run.outcome, run.outcome
        assert "P4" not in run.outcome or "No invented facts" in run.outcome
        assert set(docs_files(run)) == {"hld.docx", "hld.md", "lld.docx", "lld.md", "prd.docx", "prd.md"}, docs_files(run)
        # stop the whole run: the rest is named as not made
        run = start(brief)
        say(run, str(run.paused["handoff"]["choices"].index("stop") + 1))
        assert run.outcome.startswith("Done with open items") and "Not made: the high-level design and the low-level design" in run.outcome, run.outcome
        assert list(run.docs) == [] or list(run.docs) == ["prd"], list(run.docs)
        # a single document: "finish as it is" ends the run, and there is no "stop the whole run" option
        run = start("Write a PRD for Campus Notes." + CAMPUS)
        h, q = run.paused["handoff"], run.paused["questions"][0]
        assert run.paused["kind"] == "limit" and "stop" not in h["choices"], h["choices"]
        assert not any(o.startswith("Stop") for o in q["options"]) and "goes on" not in " ".join(q["options"]), q["options"]
        say(run, str(h["choices"].index("ship") + 1))
        assert run.outcome.startswith("Done with open items") and docs_files(run) == ["prd.docx", "prd.md"], run.outcome
    finally:
        j2.condition_met = orig_cm
    # item 2 at the source: a necessary condition unmet is never reported as passing
    audit = {"results": [{"pass": True}] * 9, "rubric": [{"pass": True}] * 3,
             "conditions": [{"id": f"P{i}", "text": f"Condition {i}", "met": i != 4} for i in range(1, 9)]}
    msg = j2.check_summary(audit, "the product requirements")
    assert "7 of 8 required conditions are met" in msg and "Not met: Condition 4." in msg and "9 of 9 checklist items pass" in msg, msg
    assert not re.search(r"checked and passes|all .* pass", msg), msg
    print("repeat failure in a chain: finish as it is accepts the document and goes on, stop the whole run is its own option, "
          "a single document ends, wording is about the document, the message's check results are code's and never claim a pass over an unmet condition ok")


def charter_objects():
    """A live model copied the STATE's shape into its charter: done_when as {"item", "criterion"} objects, not strings (run
    j2-20261023-102259 stopped with "unhashable type: 'dict'" right after Priority's GRANT of the PRD workflows). The charter is
    read as the text and the run goes on: documents, website and a test email, past the GRANT, with the deed atom created."""
    import j2_send
    keep = {k: os.environ.get(k) for k in ("J2_TEST_EMAIL", "J2_SEND_MODE")}
    os.environ["J2_TEST_EMAIL"] = "me@example.invalid"
    os.environ.pop("J2_SEND_MODE", None)
    given, read_as = [], []

    def objects(run, orig):
        facts = run._fact_problems
        def read(goal, change):      # the charter as the intake reads it, before anything else sees it
            read_as.append(list(goal["done_when"]))
            return facts(goal, change)
        run._fact_problems = read
        def call(fn, *a, **k):
            r = orig(fn, *a, **k)
            g = r.get("goal") if isinstance(r, dict) else None
            if fn == "identity" and isinstance(g, dict) and g.get("done_when") and all(isinstance(x, str) for x in g["done_when"]):
                given.append(list(g["done_when"]))      # out of order on purpose: the item number, not the list position, says where each goes
                g["done_when"] = [{"item": i, "criterion": x} for i, x in reversed(list(enumerate(g["done_when"], 1)))]
            return r
        return call
    try:
        run = docs_run(DESIGN_BRIEF + " When the documents are finished, send a test email to me with a short summary of the PRD.", 900, patch=objects)
    finally:
        for k_, v in keep.items():
            os.environ.pop(k_, None)
            if v is not None:
                os.environ[k_] = v
    chat = [json.loads(l) for l in open(os.path.join(run.dir, "chat.jsonl"), encoding="utf-8") if l.strip() and not l.startswith('{"meta"')]
    assert run.error is None and not any(m.get("who") == "system" and str(m.get("text", "")).startswith("Stopped:") for m in chat), run.error
    assert any(m.get("who") == "priority" and m.get("decision") == "grant" for m in chat), "Priority granted"
    assert any(m.get("who") == "validator" and m.get("register") for m in chat), "the run went past the GRANT: workflows were registered"
    g = json.load(open(os.path.join(run.dir, "goal.json"), encoding="utf-8"))
    assert all(isinstance(c, str) for c in g["done_when"] + g["charter_items"] + g.get("core_items", [])), "the charter holds the item text, never the object"
    assert given and read_as == given, ("item n is read as charter item n, with its criterion text unchanged", given, read_as)
    assert all(run._no(c) == i for i, c in enumerate(given[-1], 1)), "each criterion keeps its item number"
    deed = [w for w in run.past_workflows + run.workflows if w["name"] == j2_send.DEEDS["send_email"]["label"]]
    assert deed and deed[0]["atoms"], "the send deed's atom was created"
    print("charter given as objects (out of order): each item number keeps its criterion text, done_when read as text, the run goes past Priority's GRANT, documents and website built, the test email deed atom created ok")


def milestone_emails():
    """Milestone emails (code, dry run): with J2_MILESTONE_EMAILS=1, or a brief that asks for updates by email, one email each
    time a goal of the run is reached (PRD, HLD, LLD filed and checked; the site workable), only for the run's own goals. The
    recipient is only J2_TEST_EMAIL; receipts are masked; an email never stops the run; the map hangs each under its goal."""
    import j2_send
    import j2_map as M
    keep = {k: os.environ.get(k) for k in ("J2_TEST_EMAIL", "J2_SEND_MODE", "J2_MILESTONE_EMAILS")}
    addr = "me@example.invalid"

    def env(**kw):
        for k_ in keep:
            os.environ.pop(k_, None)
        os.environ.update({k_: v for k_, v in kw.items() if v is not None})

    def receipts(run):
        path = os.path.join(run.dir, "receipts.jsonl")
        return [r for r in (json.loads(x) for x in open(path, encoding="utf-8") if x.strip()) if r.get("deed") == "milestone_email"] if os.path.exists(path) else []

    def everything(run):
        return "".join(open(os.path.join(dp, f), encoding="utf-8", errors="ignore").read() for dp, _, fs in os.walk(run.dir) for f in fs if f.endswith((".json", ".jsonl", ".md", ".txt")))
    try:
        assert j2_send.asks_updates("Send me updates by email as you go") and j2_send.asks_updates("email me when each part is done")
        assert j2_send.asks_updates("keep me posted by email") and not j2_send.asks_updates("an email field in the form", "show our email address", BREW_LAB)
        # on, documents and a site: one email per goal reached, in order, each a dry-run deed under its goal
        env(J2_TEST_EMAIL=addr, J2_MILESTONE_EMAILS="1")
        run = docs_run(DESIGN_BRIEF, 900)
        assert [x["goal"] for x in run.milestone_log] == ["PRD", "HLD", "LLD", "Website"] and run.outcome.startswith("Done"), (run.milestone_log, run.outcome)
        rs = receipts(run)
        assert [r["subject"] for r in rs] == [f"J2 · Brew Lab · {g} done" for g in ("PRD", "HLD", "LLD", "Website")], [r["subject"] for r in rs]
        assert all(r["mode"] == "dry-run" and r["sent"] is False and r["to_masked"] == j2_send.mask(addr) and r["to_hash"] == j2_send.digest(addr)[:16] for r in rs), rs
        assert addr not in everything(run), "the address is never written in full"
        wfs = {w["id"]: w for w in run.past_workflows + run.workflows}
        mails = [wfs[x["atom"].split(".")[0]] for x in run.milestone_log]
        assert all(w["cost"] == j2.RATES_DEED().get("send_email") == 5 and w["atoms"][0]["results"][0]["pass"] is None for w in mails), "5 credits each; dry run leaves the atom open"
        mp = M.build(run.dir)
        st = {e["to"]: e["from"] for e in mp["edges"] if e["kind"] == "step" and e["from"] != "work:charter"}      # a step's recorded input (the Charter feeds every team too)
        brs = [n for n in mp["nodes"] if n["kind"] == "work" and n.get("lane") == "branch" and n["tag"] == "Send"]
        assert [(st[st[n["id"]]], n["label"], n["status"]) for n in brs] == [(f"work:{k}", f"Email · {g} done", "left_open") for k, g in (("prd", "PRD"), ("hld", "HLD"), ("lld", "LLD"), ("website", "Website"))], brs
        out = render_node(j2.saved_snapshot(os.path.basename(run.dir), 0), docs={"docs": {}, "map": mp})
        wide = out["map"][out["map"].index('<svg class="mapsvg"'):out["map"].index("</svg>")]
        assert wide.count('class="mapedge branch') == 8 and all(f">Email · {g} done · dry run<" in wide for g in ("PRD", "HLD", "LLD", "Website")), "one branch per email, under its goal"
        assert "Finished. 4 emails waiting (dry run)." in out["needs"], out["needs"]
        # a PRD only: one email, never one for a goal the run does not have
        run2 = docs_run("Write a PRD for Campus Notes." + CAMPUS, 400)
        assert [x["goal"] for x in run2.milestone_log] == ["PRD"], run2.milestone_log
        # off: none at all (the plain path is also guarded by plain_unchanged)
        env(J2_TEST_EMAIL=addr)
        run3 = docs_run(DESIGN_BRIEF, 900)
        assert run3.milestone_log == [] and receipts(run3) == [], "off by default"
        # the brief asks for updates by email: on, without the setting
        run4 = docs_run("Write a PRD for Campus Notes. Send me updates by email as you go." + CAMPUS, 400)
        assert [x["goal"] for x in run4.milestone_log] == ["PRD"], run4.milestone_log
        # no test address: refused and recorded, the run goes on
        env(J2_MILESTONE_EMAILS="1")
        run5 = docs_run("Write a PRD for Campus Notes." + CAMPUS, 400)
        assert run5.outcome.startswith("Done") and receipts(run5)[0]["refused"] and receipts(run5)[0]["to_masked"] is None, receipts(run5)
        # the credits can't cover it: skipped with the reason, the run goes on
        env(J2_TEST_EMAIL=addr, J2_MILESTONE_EMAILS="1")
        keep_rate, j2.RATES_DEED = j2.RATES_DEED, (lambda: {"send_email": 10 ** 6})
        try:
            run6 = docs_run("Write a PRD for Campus Notes." + CAMPUS, 400)
        finally:
            j2.RATES_DEED = keep_rate
        assert run6.outcome.startswith("Done") and run6.milestone_log[0]["result"] == "skipped" and "work credits" in run6.milestone_log[0]["why"] and receipts(run6) == [], run6.milestone_log
    finally:
        env(**{k_: v for k_, v in keep.items()})
    print("milestone emails: one per goal reached (PRD, HLD, LLD, website) and only the run's own goals, off by default, on when the brief asks "
          "for updates by email, only to the test address and masked, dry run leaves them open, skipped (never stopping) when credits are short "
          "or refused when no address is set, one map branch under each goal, outcome line counts them ok")


def fake_smtp(behave="ok", password="pw-never-in-a-file-123"):
    """A local SMTP server on 127.0.0.1 (no real network): EHLO, STARTTLS with its own certificate (made here by openssl),
    AUTH PLAIN, MAIL, RCPT, DATA. behave: "ok", "auth" (535 to the login) or "refuse" (550 to RCPT). Its error replies carry
    server text that must never reach a run file. Returns (server, port, got, cert): got lists each message it accepted."""
    import base64
    import socketserver
    import ssl
    import subprocess
    import threading
    d = tempfile.mkdtemp(prefix="j2-smtp-")
    cert, key = os.path.join(d, "c.pem"), os.path.join(d, "k.pem")
    subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-keyout", key, "-out", cert, "-days", "1",
                    "-subj", "/CN=127.0.0.1", "-addext", "subjectAltName=IP:127.0.0.1"], check=True, capture_output=True)
    ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    ctx.load_cert_chain(cert, key)
    got = []

    class H(socketserver.StreamRequestHandler):
        def handle(self):
            sock, tls = self.request, False
            f = sock.makefile("rb")
            w = lambda t: sock.sendall((t + "\r\n").encode())
            w("220 fake ESMTP")
            while True:
                line = f.readline()
                if not line:
                    return
                cmd = line.decode().strip()
                up = cmd.upper()
                if up.startswith(("EHLO", "HELO")):
                    ext = ["fake", "AUTH PLAIN" if tls else "STARTTLS", "OK"]
                    for i, e in enumerate(ext):
                        w(f"250{' ' if i == len(ext) - 1 else '-'}{e}")
                elif up == "STARTTLS":
                    w("220 go ahead")
                    sock = ctx.wrap_socket(sock, server_side=True)
                    f, tls = sock.makefile("rb"), True
                elif up.startswith("AUTH PLAIN"):
                    parts = cmd.split()
                    pw = base64.b64decode(parts[2]).split(b"\0")[-1].decode() if len(parts) > 2 else ""
                    w("535 5.7.8 server-secret-text bad credentials" if behave == "auth" or pw != password else "235 ok")
                elif up.startswith("MAIL"):
                    w("250 ok")
                elif up.startswith("RCPT"):
                    w("550 5.1.1 server-secret-text no such user" if behave == "refuse" else "250 ok")
                elif up == "DATA":
                    w("354 go ahead")
                    data = []
                    while (ln := f.readline()) not in (b".\r\n", b""):
                        data.append(ln)
                    got.append({"data": b"".join(data).decode(), "tls": tls})
                    w("250 2.0.0 queued")
                elif up == "QUIT":
                    w("221 bye")
                    return
                else:
                    w("502 not here")

    srv = socketserver.ThreadingTCPServer(("127.0.0.1", 0), H)
    srv.daemon_threads = True
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, srv.server_address[1], got, cert


def live_sends():
    """Live sending (J2_SEND_MODE=live) against a fake local SMTP server: asked once per run (Yes / No), STARTTLS and login,
    one key never sent twice, J2_SEND_CAP, missing settings refused and recorded, auth / connection / refused failures leave the
    atom open with the error class; receipts and every run file hold no password, no full address, no server text."""
    import email
    import email.policy
    import socket
    import ssl
    import j2_send
    import j2_map as M
    names = ("J2_SEND_MODE", "J2_SEND_CAP", "J2_TEST_EMAIL", "J2_MILESTONE_EMAILS", *j2_send.SMTP_VARS)
    keep, keep_tls = {k: os.environ.get(k) for k in names}, j2_send.tls_context
    addr, frm, pw = "me@example.invalid", "j2-from@example.invalid", "pw-never-in-a-file-123"
    servers, shown = [], []

    def env(port, **kw):
        for k in names:
            os.environ.pop(k, None)
        os.environ.update({"J2_SEND_MODE": "live", "J2_TEST_EMAIL": addr, "J2_MILESTONE_EMAILS": "1", "J2_SMTP_HOST": "127.0.0.1",
                           "J2_SMTP_PORT": str(port), "J2_SMTP_USER": "j2-user", "J2_SMTP_PASSWORD": pw, "J2_SMTP_FROM": frm})
        for k, v in kw.items():
            os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)

    def server(behave="ok"):
        srv, port, got, cert = fake_smtp(behave, pw)
        servers.append(srv)
        j2_send.tls_context = lambda: ssl.create_default_context(cafile=cert)      # trust the fake server's own certificate only
        return port, got

    def recs(run):
        return j2_send.records(run.dir)

    def asks(run):
        return [m for m in run.chat if m.get("pause") == "send_confirm"]

    def clean(run):
        text = "".join(open(os.path.join(dp, f), encoding="utf-8", errors="ignore").read() for dp, _, fs in os.walk(run.dir) for f in fs)
        text += json.dumps(run.chat) + json.dumps(run.snapshot(), default=str)
        for bad in (pw, addr, frm, "server-secret-text", "bad credentials"):
            assert bad not in text, f"{bad!r} is in a run file, the chat or the state"
        shown.extend(recs(run))

    def prd(answer, brief="Write a PRD for Campus Notes." + CAMPUS):
        run = docs_run(brief, 400)
        assert run.paused and run.paused["kind"] == "send_confirm" and not recs(run), (run.outcome, run.paused)
        q = stage_question(run)
        assert "Send real emails to your test address during this run?" in q and "Yes" in q and "No" in q, q
        say(run, answer)
        return run

    def box(run):
        return next(n for n in M.build(run.dir)["nodes"] if str(n.get("label", "")).startswith("Email · "))
    try:
        # success: asked once, then STARTTLS, login, one message to the test address; the atom closes; map and outcome say sent
        port, got = server()
        env(port)
        run = prd("Yes")
        r = recs(run)
        assert [(x["mode"], x["sent"], x["smtp_code"], x["result"], x["error"]) for x in r] == [("live", True, 250, "sent", None)], r
        assert r[0]["to_masked"] == "m***@example.invalid" and r[0]["to_hash"] == j2_send.digest(addr)[:16] and r[0]["key"].endswith("/Email: PRD done") and r[0]["t"]
        m = email.message_from_string(got[0]["data"], policy=email.policy.default) if len(got) == 1 else {}
        assert got[0]["tls"] and m["To"] == addr and m["From"] == frm and m["Subject"] == "J2 · Campus Notes · PRD done", (got, dict(m))
        assert run.outcome.startswith("Done") and len(asks(run)) == 1, run.outcome
        v = next(v for v in json.load(open(os.path.join(run.dir, "atoms.json")))["atoms"] if v["id"] == r[0]["atom"])
        assert v["status"] == "closed" and v["notes"] == "Sent to your test address (m***@example.invalid)." and v["details"]["mode"] == "live"
        assert box(run)["status"] == "filed", box(run)
        out = render_node(j2.saved_snapshot(os.path.basename(run.dir), 0), docs={"docs": {}, "map": M.build(run.dir)})
        assert ">Email · PRD done · sent<" in out["map"] and "Finished. Email · PRD done · sent." in out["needs"], out["needs"]
        clean(run)
        # no duplicates: the same key again (a retry, or the run resumed from its folder) sends nothing and says so
        again = j2_send.send("milestone_email", "prd.md", "", None, built={"subject": "s", "body": "b"}, run_dir=run.dir, key=r[0]["key"], confirmed=True)
        assert again["sent"] is False and again["error"] == "duplicate" and again["result"] == "skipped" and len(got) == 1, again
        # the brief's own test email, with milestones on: one question for both sends, two messages, two keys
        got.clear()
        run = prd("Option 1", "Write a PRD for Campus Notes. Send me an email with it." + CAMPUS)
        r = recs(run)
        assert [x["deed"] for x in r] == ["milestone_email", "send_email"] and all(x["sent"] for x in r) and len(got) == 2 and len(asks(run)) == 1, r
        assert len({x["key"] for x in r}) == 2 and run.outcome.startswith("Done"), run.outcome
        clean(run)
        # the cap: four goals, J2_SEND_CAP=2: two go out, two are skipped and recorded; still asked only once
        got.clear()
        env(port, J2_SEND_CAP="2")
        run = docs_run(DESIGN_BRIEF, 900)
        say(run, "yes")
        r = recs(run)
        assert [(x["sent"], x["error"]) for x in r] == [(True, None), (True, None), (False, "cap"), (False, "cap")] and len(got) == 2, r
        assert "J2_SEND_CAP" in r[2]["why"] and len(asks(run)) == 1 and run.outcome.startswith("Done"), run.outcome
        wfs = {w["id"]: w for w in run.past_workflows + run.workflows}
        assert [(wfs[x["atom"].split(".")[0]]["status"], wfs[x["atom"].split(".")[0]].get("send_error")) for x in r] == [
            ("done", None), ("done", None), ("skipped", "cap"), ("skipped", "cap")], "a skipped send is never done"
        v = next(v for v in json.load(open(os.path.join(run.dir, "atoms.json")))["atoms"] if v["id"] == r[2]["atom"])
        assert v["status"] == "open" and v["notes"].startswith("Skipped (cap): the cap of 2 live emails"), v["notes"]
        out = render_node(j2.saved_snapshot(os.path.basename(run.dir), 0), docs={"docs": {}, "map": M.build(run.dir)})
        assert ">Email · LLD done · skipped (cap)<" in out["map"] and "Email · Website done · skipped (cap)." in out["needs"], out["needs"]
        assert "checks did not pass" not in out["map"], "a send not sent is not a failed check"
        assert [x["goal"] for x in run.milestone_log] == ["PRD", "HLD", "LLD", "Website"]
        clean(run)
        # No: every send of the run stays a dry run; nothing reaches the server; never asked again
        got.clear()
        env(port)
        run = docs_run(DESIGN_BRIEF, 900)
        say(run, "No")
        r = recs(run)
        assert len(r) == 4 and all(x["mode"] == "dry-run" and not x["sent"] for x in r) and not got and len(asks(run)) == 1, r
        assert box(run)["detail"] == "Not sent (dry run)." and run.outcome.startswith("Done")
        clean(run)
        # auth failure: 535; not sent, the atom stays open with the reason; the map says failed; no server text kept
        port, got = server("auth")
        env(port)
        run = prd("Yes")
        r = recs(run)[0]
        assert (r["sent"], r["error"], r["smtp_code"], r["result"]) == (False, "auth", 535, "failed") and not got, r
        v = next(v for v in json.load(open(os.path.join(run.dir, "atoms.json")))["atoms"] if v["id"] == r["atom"])
        assert v["status"] == "open" and v["notes"] == "Failed (auth): the SMTP server did not accept the login.", v
        w = next(w for w in run.past_workflows + run.workflows if w["id"] == r["atom"].split(".")[0])
        assert (w["status"], w["send_error"]) == ("failed", "auth") and box(run)["status"] == "left_open" and run.outcome.startswith("Done"), run.outcome
        out = render_node(j2.saved_snapshot(os.path.basename(run.dir), 0), docs={"docs": {}, "map": M.build(run.dir)})
        assert ">Email · PRD done · failed (auth)<" in out["map"] and "Finished. Email · PRD done · failed (auth)." in out["needs"], out["needs"]
        clean(run)
        # refused: 550 to the recipient
        port, got = server("refuse")
        env(port)
        r = recs(prd("Yes"))[0]
        assert (r["sent"], r["error"], r["smtp_code"]) == (False, "refused", 550) and not got, r
        # connection failure: nothing listens on the port
        with socket.socket() as sk:
            sk.bind(("127.0.0.1", 0))
            dead = sk.getsockname()[1]
        env(dead)
        run = prd("Yes")
        r = recs(run)[0]
        assert (r["sent"], r["error"], r["smtp_code"]) == (False, "connection", None) and run.outcome.startswith("Done"), r
        clean(run)
        # missing settings: refused and recorded by name (never a value), the run goes on
        env(port, J2_SMTP_HOST=None, J2_SMTP_PASSWORD="")
        run = prd("Yes")
        r = recs(run)[0]
        assert (r["sent"], r["error"]) == (False, "settings") and r["why"] == "J2_SMTP_HOST, J2_SMTP_PASSWORD not set, so nothing is sent", r
        assert run.outcome.startswith("Done") and box(run)["detail"].startswith("Failed (settings): J2_SMTP_HOST")
        # a brief that asks for updates by email gets milestone emails only, never a test-email deed as well (run 130930 got both)
        port, got = server()
        env(port, J2_MILESTONE_EMAILS=None)
        run = prd("Yes", "Write a PRD for Campus Notes. Send me updates by email as you go." + CAMPUS)
        assert run.stages == ["prd"] and [x["deed"] for x in recs(run)] == ["milestone_email"] and len(got) == 1, (run.stages, recs(run))
        body = email.message_from_string(got[0]["data"], policy=email.policy.default).get_content()
        summ = re.search(r"^Summary: (.*)$", body, re.M).group(1).strip()
        assert len(summ) <= 200 and (summ.endswith((".", "!", "?")) or summ.endswith("...")), summ
        clean(run)
        clean(run)
    finally:
        for srv in servers:
            srv.shutdown()
        j2_send.tls_context = keep_tls
        for k, v in keep.items():
            os.environ.pop(k, None)
            if v is not None:
                os.environ[k] = v
    if "--receipts" in sys.argv:
        for x in shown:
            print(json.dumps({k: x.get(k) for k in ("deed", "mode", "sent", "result", "smtp_code", "error", "why", "to_masked", "to_hash", "key", "t")}))
    print("live sends (fake local SMTP server): asked once per run (yes sends, no keeps every send a dry run), STARTTLS and login, "
          "one key never sent twice, J2_SEND_CAP skips and records, auth / refused / connection failures and missing settings leave the "
          "atom open with the error class, the map and outcome line say sent or failed, no password, full address or server text in any file ok")


def send_deeds():
    """Send deeds, phase 1 (no network): the brief's send words trigger a trailing deed stage; the recipient comes only
    from J2_TEST_EMAIL / J2_TEST_SLACK_TARGET; dry-run leaves the atom open, mock records one message; the proof check
    reads the receipt; funding goes through the stage gate; a plain brief is untouched."""
    import j2_schema as S
    import j2_send
    FAKE = {"J2_TEST_EMAIL": "me@example.invalid", "J2_TEST_SLACK_TARGET": "#j2-test",
            "J2_SMTP_PASSWORD": "pw-should-never-appear", "J2_SLACK_WEBHOOK_URL": "https://hooks.invalid/never-in-a-file"}
    keep = {k: os.environ.get(k) for k in (*FAKE, "J2_SEND_MODE")}

    def env(mode=None, **kw):
        for k in keep:
            os.environ.pop(k, None)
        os.environ.update({**FAKE, **{k: v for k, v in kw.items() if v is not None}})
        for k, v in kw.items():
            if v is None:
                os.environ.pop(k, None)
        if mode:
            os.environ["J2_SEND_MODE"] = mode
        j2_send.SENT.clear()

    def deed_of(run, kind="send_email"):
        w = next(w for w in run.past_workflows + run.workflows if w["name"] == j2_send.DEEDS[kind]["label"])
        return w, w["atoms"][0]

    def run_text(run):
        return "".join(open(os.path.join(run.dir, f), encoding="utf-8").read() for f in os.listdir(run.dir)
                       if f.endswith((".json", ".jsonl", ".md")))
    try:
        # 1: a send verb near email or slack triggers; a mere mention does not
        d = j2_send.detect_deeds
        assert d("Send me an email with the PRD") == ["send_email"] and d("email it to me") == ["send_email"]
        assert d("send the PRD by email") == ["send_email"] and d("forward it by e-mail") == ["send_email"]
        assert d("Post it to Slack") == ["send_slack_message"] and d("Slack me the PRD") == ["send_slack_message"]
        assert d("send an email and post to Slack") == ["send_email", "send_slack_message"]
        assert d("Write a PRD. Send me updates by email as you go.") == [] and d("Send me a test email.") == ["send_email"]      # updates: milestone emails only
        assert d("Send me updates by email as you go. Also email me the PRD.") == ["send_email"], "another sentence still asks for the deed"
        # the milestone summary never ends mid-word: the last full sentence that fits, else whole words and "..."
        long_ = "Students who miss a lecture have no easy way to get good notes; they ask around on group chats and often get nothing. " \
                "The students affected are undergraduate students at one university. Campus Notes is where they share and find them."
        assert j2_send.clip(long_) == long_[:long_.index("university.") + 11] and j2_send.clip("short") == "short"
        assert j2_send.clip("word " * 60) == ("word " * 39).strip() + "..." and len(j2_send.clip("word " * 60)) <= 200
        assert j2_send.clip("x" * 250).endswith("...") and len(j2_send.clip("x" * 250)) <= 200
        assert "Summary: " + j2_send.clip(long_) in j2_send.milestone("P", "PRD", 1, (1, 1), long_)["body"]
        for t in ("", None, BREW_LAB, "a Slack link in the footer", "integrate with Slack", "an email field in the form",
                  "show our email address", "Slack is where the team chats", "design a logo"):
            assert d(t) == [], t
        # 2: the allowlist: everything refused if unset; one recipient; no line breaks; nothing else gets through
        env(J2_TEST_EMAIL=None, J2_TEST_SLACK_TARGET=None)
        ok, why = j2_send.allowed("email", "me@example.invalid")
        assert not ok and "J2_TEST_EMAIL" in why and "example.invalid" not in why, why
        assert not j2_send.allowed("slack", "#j2-test")[0]
        env()
        assert j2_send.allowed("email", "me@example.invalid")[0] and j2_send.allowed("email", " ME@Example.invalid ")[0]
        assert j2_send.allowed("slack", "#j2-test")[0]
        for ch, to in (("email", "other@example.invalid"), ("email", "me@example.invalid, other@example.invalid"),
                       ("email", "me@example.invalid\nBcc: x@example.invalid"), ("email", "me@example.invalid\r"),
                       ("email", ["me@example.invalid"]), ("email", None), ("email", ""), ("slack", "#general"),
                       ("slack", "#j2-test #general"), ("slack", "#j2-test\n"), ("fax", "me@example.invalid")):
            assert not j2_send.allowed(ch, to)[0], (ch, to)
        env("mock")
        for msg in ({"channel": "email", "to": "other@example.invalid", "subject": "s", "body": "b"},
                    {"channel": "slack", "to": "#general", "subject": None, "body": "b"},
                    {"channel": "email", "to": "me@example.invalid", "subject": "s", "body": "b", "cc": "x@example.invalid"},
                    {"channel": "email", "to": "me@example.invalid", "subject": "s", "body": "b", "bcc": "x@example.invalid"}):
            r = j2_send.deliver(msg)
            assert not r["sent"] and r["refused"] and not j2_send.SENT, (msg, r)
        # 3: dry-run is the default: nothing is sent, the atom stays open, the proof check does not pass
        env()
        run = docs_run("Write a PRD for Campus Notes. Send me an email with it." + CAMPUS, 400)
        assert run.stages == ["prd", "send_email"] and run.outcome.startswith("Done"), (run.stages, run.outcome)
        assert "Not sent (dry run)" in run.outcome and not j2_send.SENT
        w, a = deed_of(run)
        assert a["status"] == "could not run" and a["results"][0]["pass"] is None and a["results"][0]["evidence"] == "Not sent (dry run)"
        assert a["human_output"]["text"] == "Not sent (dry run)." and w["status"] == "done" and w["cost"] == 5
        views = json.load(open(os.path.join(run.dir, "atoms.json")))["atoms"]
        v = next(v for v in views if v["id"] == a["id"])
        assert v["status"] == "open" and v["notes"] == "Not sent (dry run)." and v["verification"][0]["met"] is None
        # the atom shape: verb do, form send, scope external, one deed-send check, edges to the PRD's atoms
        assert v["output"] == [{"verb": "do", "name": a["results"][0]["check"], "form": "send", "ref": f"{a['id']}:1"}], v["output"]
        assert v["details"]["scope"] == "external" and v["details"]["mode"] == "dry-run" and v["details"]["sent"] is False
        assert S.atom_kind(v) == "do" and len(v["verification"]) == 1 and v["verification"][0]["method"].startswith("deed-send")
        assert {e["atom"] for e in v["depends_on"]} == set(run.docs["prd"]["atoms"]) and v["owner"] == "code"
        assert all(a["id"] in [u["atom"] for u in x["used_by"]] for x in views if x["id"] in run.docs["prd"]["atoms"])
        assert not any(a["id"] in x for x in json.load(open(os.path.join(run.dir, "atoms.json")))["warnings"]), "warned about a send"
        rec = [json.loads(x) for x in open(os.path.join(run.dir, "receipts.jsonl"), encoding="utf-8")]
        rec = next(r for r in rec if r.get("type") == "deed-send-record")
        for k in ("atom", "deed", "channel", "mode", "to_masked", "to_hash", "subject", "body_sha256", "body_chars", "artifact",
                  "artifact_sha256", "sent", "result", "refused", "t", "by"):
            assert k in rec, k
        assert rec["mode"] == "dry-run" and rec["sent"] is False and rec["to_masked"] == "m***@example.invalid" and rec["by"] == "code"
        assert "me@example.invalid" not in run_text(run) and not any(x in run_text(run) for x in
                                                                    (FAKE["J2_SMTP_PASSWORD"], FAKE["J2_SLACK_WEBHOOK_URL"]))
        # 4: mock: one message recorded, to the configured target only; the check passes and the atom closes
        env("mock")
        run = docs_run("Write a PRD for Campus Notes. Send me an email with it." + CAMPUS, 400)
        w, a = deed_of(run)
        assert len(j2_send.SENT) == 1 and j2_send.SENT[0]["to"] == "me@example.invalid" and j2_send.SENT[0]["channel"] == "email"
        assert "prd.md" in j2_send.SENT[0]["subject"] and "Sections:" in j2_send.SENT[0]["body"]
        assert a["status"] == "passed" and a["results"][0]["pass"] is True and "mock" in a["results"][0]["evidence"]
        v = next(v for v in json.load(open(os.path.join(run.dir, "atoms.json")))["atoms"] if v["id"] == a["id"])
        assert v["status"] == "closed" and v["details"]["mode"] == "mock" and v["details"]["sent"] is True
        assert "me@example.invalid" not in run_text(run) and FAKE["J2_SMTP_PASSWORD"] not in run_text(run)
        # the proof check only reads: running it again sends nothing more, and it fails if the document changed
        c = a["verification"][0]
        assert j2_send.check(c, run.dir)["pass"] is True and len(j2_send.SENT) == 1
        with open(os.path.join(run.dir, "prd.md"), "a", encoding="utf-8") as f:
            f.write("\n## Extra\nchanged\n")
        assert j2_send.check(c, run.dir)["pass"] is False and "changed" in j2_send.check(c, run.dir)["evidence"]
        assert j2_send.check({**c, "atom": "W-99.a1"}, run.dir)["pass"] is False      # no receipt
        env("mock", J2_TEST_EMAIL="someone-else@example.invalid")      # the configured target moved: the receipt no longer matches
        assert j2_send.check(c, run.dir)["pass"] is False
        # a Slack send, and both together, in order
        env("mock")
        run = docs_run("Write a PRD for Campus Notes. Email it to me and post it to Slack." + CAMPUS, 400)
        assert run.stages == ["prd", "send_email", "send_slack_message"] and run.outcome.startswith("Done"), run.stages
        assert [(m["channel"], m["to"]) for m in j2_send.SENT] == [("email", "me@example.invalid"), ("slack", "#j2-test")]
        assert deed_of(run, "send_slack_message")[1]["status"] == "passed"
        # target not configured: nothing is sent, the atom is verify_failed, the outcome says so
        env("mock", J2_TEST_EMAIL=None)
        run = docs_run("Write a PRD for Campus Notes. Send me an email with it." + CAMPUS, 400)
        w, a = deed_of(run)
        assert not j2_send.SENT and a["status"] == "failed" and w["status"] == "failed"
        assert run.outcome.startswith("Done with open items") and "J2_TEST_EMAIL" in run.outcome, run.outcome
        v = next(v for v in json.load(open(os.path.join(run.dir, "atoms.json")))["atoms"] if v["id"] == a["id"])
        assert v["status"] == "verify_failed"
        # 5: the brief names another recipient: told in plain words, the send still goes only to the configured target
        env("mock")
        run = docs_run("Write a PRD for Campus Notes. Send me an email with it, to boss@elsewhere.invalid." + CAMPUS, 400)
        note = next(m["text"] for m in run.chat if m["who"] == "system" and "different address" in m["text"])
        assert "Test mode sends only to the configured test address" in note and "boss@elsewhere.invalid" not in note, note
        assert [m["to"] for m in j2_send.SENT] == ["me@example.invalid"], j2_send.SENT      # never the one in the brief
        assert not any("different" in m["text"] for m in docs_run("Write a PRD for Campus Notes. Email it to me@example.invalid." + CAMPUS, 400).chat
                       if m["who"] == "system")
        # 6: budget: 5 credits each from rates_deed; J2_RATES=0 makes them free; short credits go through the stage gate
        env()
        base = docs_run("Write a PRD for Campus Notes." + CAMPUS, 400)
        with_email = docs_run("Write a PRD for Campus Notes. Send me an email with it." + CAMPUS, 400)
        assert with_email.spent - base.spent == 5 == json.load(open(os.path.join(j2.J2_DIR, "rates.json")))["rates_deed"]["send_email"]
        rates_on = j2.RATES_ON
        j2.RATES_ON = False      # J2_RATES=0: the deed's price is 0 (a whole PRD run needs the rate card, so the price is read directly)
        try:
            assert j2.RATES_DEED() == {} and with_email._stage_cost("send_email") == 0 and with_email._stage_cost("send_slack_message") == 0
        finally:
            j2.RATES_ON = rates_on
        assert with_email._stage_cost("send_email") == 5 and with_email._stage_cost("send_slack_message") == 5
        run = docs_run("Write a PRD for Campus Notes. Send me an email with it." + CAMPUS, 100)
        assert run.paused and run.paused["kind"] == "limit" and run.stages == ["prd", "send_email"], run.outcome
        text = stage_question(run)
        assert "the test email (needs at least 5 work credits; not covered)" in text and "the product requirements (needs at least 98" in text, text
        say(run, "Option 1")      # add the credits: the PRD and the email are both done
        assert run.outcome.startswith("Done") and deed_of(run)[0]["cost"] == 5, run.outcome
        run = docs_run("Write a PRD for Campus Notes. Send me an email with it." + CAMPUS, 100)
        say(run, "Option 2")      # leave the email out: no send workflow exists
        assert run.stages == ["prd"] and run.outcome.startswith("Done") and not any(
            w["name"] == "Send a test email" for w in run.past_workflows + run.workflows)
        # 10: no send words: no deed stage, nothing recorded; send words with no document: told, and the site is built as before
        run = docs_run("Write a PRD for Campus Notes." + CAMPUS, 400)
        assert run.stages == ["prd"] and not run.deed_log and "deed-send-record" not in run_text(run)
        run = docs_run(BREW_LAB + " Email it to me.", 200)
        assert not run.stages and run.outcome.startswith("Done") and not j2_send.SENT and "deed-send-record" not in run_text(run)
        assert any("A send needs a document" in m["text"] for m in run.chat if m["who"] == "system")
    finally:
        for k, v in keep.items():
            os.environ.pop(k, None)
            if v is not None:
                os.environ[k] = v
        j2_send.SENT.clear()
    print("send deeds: trigger needs a send verb, allowlist refuses anything but the two configured targets, dry-run stays open "
          "(not sent), mock records one message, proof check only reads the receipt, no secret or full address in any file, "
          "another recipient in the brief is told and ignored, 5 credits each and free with J2_RATES=0, stage gate funds it, "
          "plain briefs untouched ok")


def artifact_types():
    """artifacts/types.json mirrors the Python for prd, hld, lld and website, field by field: sections, file names,
    trigger words, groups, tags, labels, colours, conditions file, builder, rate card, minimum cost and what each reads.
    The same comparison fails when either side drifts (checked on edited copies)."""
    import collections
    import copy
    import j2_artifacts as A
    T = A.load_types()
    ui = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "j2_ui.html"), encoding="utf-8").read()
    kinds = {k: (tag, label) for k, tag, label in re.findall(r'\["(prd|hld|lld)", "([A-Z]+)", "([^"]+)"\]', ui)}
    site_tag = re.search(r'\["site", "([^"]+)"\]', ui).group(1)
    corpus = ["write a PRD", "the HLD", "an LLD please", "product requirements", "high-level design", "low level design",
              "architecture overview", "detailed design", "design docs", "a modern design", "BRD", "HLDs", "LLDs and PRDs"]
    run = j2.J2Run()

    def problems(T):
        by, bad = {t["name"]: t for t in T["types"]}, []
        chk = lambda ok, what: bad.append(what) if not ok else None
        types = [t for t in T["types"] if t["kind"] == "type"]
        chk([t["name"] for t in types] == [*j2.DOC_SET, "website"], "type order")
        chk(T["site_words"] == j2._SITE_WORDS.pattern, "site words")
        g = T["groups"]["design-docs"]
        chk(g["pattern"] == j2._DOC_SET_WORDS.pattern and g["members"] == list(j2.DOC_SET), "design-docs group")
        for k in j2.DOC_SET:
            t, o = by[k], j2.OUTPUTS[k]
            chk(t["workflow"]["sections"] == list(o["sections"]), f"{k} sections")
            chk(t["file"] == o["file"] and t["export"] == "docx" and bool(o["export"]) and t["form"] == "doc", f"{k} file")
            chk(t["conditions"] == o["conditions"] and os.path.isfile(os.path.join(j2.CONDITIONS_DIR, t["conditions"])), f"{k} conditions")
            chk(getattr(j2, t["workflow"]["builder"].split(":", 1)[1], None) is o["builder"], f"{k} builder")
            chk(t["cost"]["card"] == o["rates"], f"{k} rate card")
            chk(t["cost"]["parts"] == {"doc_section": len(o["sections"]) - 1, "requirements_section": 1}, f"{k} cost parts")
            card = A_rates(t["cost"]["card"])
            chk(sum(card.get(p, 0) * n for p, n in t["cost"]["parts"].items()) == run._stage_cost(k), f"{k} minimum cost")
            chk(t["triggers"] == [j2._DOC_WORDS[k].pattern], f"{k} trigger words")
            rx = re.compile("|".join(t["triggers"]), re.I)
            chk(all(bool(rx.search(x)) == bool(j2._DOC_WORDS[k].search(x)) for x in corpus), f"{k} trigger behaviour")
            chk(t["groups"] == ["design-docs"] and k in g["members"], f"{k} groups")
            chk((t["tag"], t["label"]) == (kinds[k][0], kinds[k][1]) and t["label"].lower() == j2.DOC_NAMES[k], f"{k} tag and label")
            chk(t["color"] == f"--{k}" and f"--{k}:" in ui, f"{k} colour")
        w, o = by["website"], j2.OUTPUTS["website"]
        chk(w["file"] == "site/index.html" and w["form"] == "page" and w["triggers"] == [] and w["also_when"] == "site-words", "website file and triggers")
        chk(w["conditions"] == o["conditions"] and w["workflow"]["builder"] == o["builder"] and w["cost"]["card"] == o["rates"], "website conditions, builder, card")
        chk(w["cost"]["parts"] == dict(collections.Counter(j2.J2Run.SITE_ESSENTIALS)), "website cost parts")
        chk(sum(A_rates(w["cost"]["card"]).get(p, 0) * n for p, n in w["cost"]["parts"].items()) == run._stage_cost("website"), "website minimum cost")
        chk((w["tag"], w["color"]) == (site_tag, "--site") and "--site:" in ui and w["label"] == o["label"], "website tag, colour, label")
        return bad

    def A_rates(card):
        return {k: int(v) for k, v in json.load(open(os.path.join(j2.J2_DIR, "rates.json"))).get(card, {}).items()}
    assert problems(T) == [], problems(T)
    # drift on either side is caught
    for path, val in (("types.0.workflow.sections", ["Problem"]), ("types.1.color", "--x"), ("types.2.triggers", ["\\bLLD\\b"]),
                      ("types.3.cost.parts", {"form": 9}), ("types.0.file", "p.md"), ("types.1.tag", "H"), ("types.3.file", "x.html"),
                      ("types.2.conditions", "none.json"), ("groups.design-docs.members", ["prd"])):
        bent, node, keys = copy.deepcopy(T), None, path.split(".")
        node = bent
        for kx in keys[:-1]:
            node = node[int(kx)] if kx.isdigit() else node[kx]
        node[keys[-1]] = val
        assert problems(bent), path
    keep = j2._DOC_WORDS["hld"]
    j2._DOC_WORDS["hld"] = re.compile(r"\bHLD\b", re.I)
    try:
        assert problems(T), "Python trigger words changed and the list did not notice"
    finally:
        j2._DOC_WORDS["hld"] = keep
    keep = j2.OUTPUTS["lld"]["sections"]
    j2.OUTPUTS["lld"]["sections"] = keep[:-1]
    try:
        assert problems(T), "Python sections changed and the list did not notice"
    finally:
        j2.OUTPUTS["lld"]["sections"] = keep
    # records: every record is a kind record, not a deliverable, with a file pattern; a full run keeps the ones a mock run writes
    recs = [t for t in T["types"] if t["kind"] == "record"]
    assert [t["name"] for t in recs] == ["screenshots", "atoms", "audits", "receipts", "conditions", "chain", "goal", "workflows",
                                         "summary", "report", "chat", "board", "realization", "workflow_files",
                                     "workflow_files_in_folders"] and not any(t["deliverable"] for t in recs)
    # depends_on matches the edges a real chain run records between the documents and the site
    run = docs_run(DESIGN_BRIEF, 900)
    owner = {a: k for k, v in json.load(open(os.path.join(run.dir, "documents.json"))).items() for a in v["atoms"]}
    for w in run._all_wfs():
        for a in w["atoms"]:
            owner.setdefault(a["id"], "website")
    seen = collections.defaultdict(set)
    for v in json.load(open(os.path.join(run.dir, "atoms.json")))["atoms"]:
        for e in v["depends_on"]:
            if owner[e["atom"]] != owner[v["id"]]:
                seen[owner[v["id"]]].add(owner[e["atom"]])
    by = {t["name"]: t for t in T["types"]}
    for k in (*j2.DOC_SET, "website"):
        assert set(by[k]["depends_on"]) == seen.get(k, set()), (k, by[k]["depends_on"], seen.get(k))
    rows = A.rows(run.dir)
    docs = json.load(open(os.path.join(run.dir, "documents.json")))
    for k in j2.DOC_SET:      # the workflows a row names are the workflows the run recorded for that document
        assert next(r for r in rows if r["type"] == k)["workflows"] == sorted(docs[k]["workflows"], key=lambda x: int(x[2:])), k
    for t in recs:
        # a mock run takes no screenshots; board.json only exists under J2_BOARD=1; the realization plan and workflow files
        # only exist under J2_REALIZATION=1
        if t["name"] not in ("screenshots", "board", "realization", "workflow_files", "workflow_files_in_folders"):
            assert os.path.isfile(os.path.join(run.dir, t["file"])), t["file"]
    print("artifact types: prd, hld, lld and website match the Python field by field (sections, files, triggers, groups, tags, "
          "colours, conditions, builder, rates, minimum cost, depends_on); drift on either side fails; 15 record entries ok")


def artifacts_tab():
    """Artifacts: artifacts.json written at the end of a run (additive), the same rows rebuilt from an old folder, the
    endpoints serve only files the list names, and the tab renders two headings with tags, workflows, dependencies and an
    Open button per form (document text, website preview, screenshot image, data text)."""
    import shutil
    import threading
    import urllib.error
    import urllib.request
    import j2_artifacts as A
    import j2_server
    run = docs_run(DESIGN_BRIEF, 900)
    name = os.path.basename(run.dir)
    saved = json.load(open(os.path.join(run.dir, "artifacts.json")))
    assert saved["schema"] == "artifacts/1" and all(set(r) == {"name", "type", "kind", "form", "file", "deliverable", "workflows",
                                                                "depends_on", "version", "sha256", "time"} for r in saved["rows"])
    by = {r["type"]: r for r in saved["rows"]}
    assert [r["type"] for r in saved["rows"] if r["deliverable"]] == ["prd", "hld", "lld", "website"]
    assert by["prd"]["workflows"][0] == "W-1" and by["lld"]["depends_on"] == ["prd", "hld"] and by["website"]["depends_on"] == ["hld", "lld"]
    assert by["prd"]["version"] >= 1 and by["website"]["file"] == "site/index.html"
    assert by["prd"]["sha256"] == hashlib.sha256(open(os.path.join(run.dir, "prd.md"), "rb").read()).hexdigest()
    now = A.rows(run.dir)      # an old run has no artifacts.json: the scan gives the same rows (the chat log grew after the file was written)
    assert [r for r in now if r["file"] != "chat.jsonl"] == [r for r in saved["rows"] if r["file"] != "chat.jsonl"]
    assert {k: v for k, v in next(r for r in now if r["file"] == "chat.jsonl").items() if k not in ("sha256", "time")} == \
        {k: v for k, v in next(r for r in saved["rows"] if r["file"] == "chat.jsonl").items() if k not in ("sha256", "time")}
    # a plain website run: artifacts.json lists the site and the records, no documents
    plain = docs_run(BREW_LAB, 200)
    prow = json.load(open(os.path.join(plain.dir, "artifacts.json")))["rows"]
    assert [r["type"] for r in prow if r["deliverable"]] == ["website"] and prow[0]["depends_on"] == [] and prow[0]["workflows"]
    # an old run: the same folder with no artifacts.json, and a screenshot added by hand
    old = os.path.join(j2.RUNS_DIR, "j2-20200101-000000")
    shutil.copytree(run.dir, old)
    os.remove(os.path.join(old, "artifacts.json"))
    os.makedirs(os.path.join(old, "site", "shots"), exist_ok=True)
    open(os.path.join(old, "site", "shots", "v4.png"), "wb").write(b"\x89PNG\r\n\x1a\n" + b"0" * 16)
    lst = A.listing(old)
    assert not lst["deliverables"][0]["saved"] and [r["type"] for r in lst["deliverables"]] == ["prd", "hld", "lld", "website"]
    assert any(r["form"] == "image" and r["file"] == "site/shots/v4.png" and r["open"] == "image" for r in lst["records"])
    assert not os.path.exists(os.path.join(old, "artifacts.json")), "listing an old run must not write to it"
    # a real stored run from before artifacts existed, when present (read-only evidence)
    here = os.path.dirname(os.path.abspath(__file__))
    real = "j2-20261016-160511"
    if os.path.isdir(os.path.join(here, "runs", real)):
        keep_dir, j2.RUNS_DIR = j2.RUNS_DIR, os.path.join(here, "runs")
        try:
            assert not os.path.exists(os.path.join(here, "runs", real, "artifacts.json"))
            got = A.listing(j2.saved_dir(real))
            assert got["deliverables"] and not got["deliverables"][0]["saved"] and [r["type"] for r in got["deliverables"]][0] == "prd"
            assert not os.path.exists(os.path.join(here, "runs", real, "artifacts.json"))
        finally:
            j2.RUNS_DIR = keep_dir
    empty = A.listing(None)
    assert empty == {"deliverables": [], "records": []} and A.rows(os.path.join(run.dir, "nope")) == []
    # endpoints: live and saved, serving only the files the list names
    keep, j2_server.RUN = j2_server.RUN, run
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    get = lambda path: urllib.request.urlopen(base + path)
    try:
        for q in ("", f"?run={name}", "?run=j2-20200101-000000"):
            got = json.load(get("/api/artifacts" + q))
            assert [r["type"] for r in got["deliverables"]] == ["prd", "hld", "lld", "website"], q
            assert {"summary", "atoms", "receipts", "chat"} <= {r["type"] for r in got["records"]}, q
            row = got["deliverables"][0]
            assert row["label"] == "Product requirements" and row["tag"] == "PRD" and row["color"] == "--prd" and row["open"] == "text"
            assert got["deliverables"][3]["open"] == "preview" and got["deliverables"][3]["depends_on_labels"] == ["High-level design", "Low-level design"]
        r = get("/api/artifact?name=prd.md")
        assert r.read().decode().startswith("# ") and r.headers["Content-Type"].startswith("text/plain")
        assert get(f"/api/artifact?name=site/index.html&run={name}").read().decode().lstrip().lower().startswith("<!doctype")
        assert get("/api/artifact?name=site/shots/v4.png&run=j2-20200101-000000").read()[:4] == b"\x89PNG"
        assert json.loads(get("/api/artifact?name=atoms.json").read())["schema"] == "atom/1"
        for bad in ("prd.docx", "../j2.py", "/etc/passwd", "site/v1.html", "hld-v1.md", "site/hld-v1.md", "nothing.json", "", "calls.jsonl",
                    "artifacts.json", "prd.md/../goal.json"):
            try:
                get("/api/artifact?name=" + urllib.request.quote(bad))
                assert False, f"served {bad!r}"
            except urllib.error.HTTPError as e:
                assert e.code == 404, (bad, e.code)
        try:
            get("/api/artifact?name=prd.md&run=j2-20990101-000000")
            assert False
        except urllib.error.HTTPError as e:
            assert e.code == 404
        assert json.load(get("/api/artifacts?run=j2-20990101-000000")) == {"deliverables": [], "records": []}
        for f in ("documents.json",):      # the Documents tab's own endpoints are untouched
            assert get("/file/prd.md").read().startswith(b"# ")
    finally:
        j2_server.RUN = keep
        srv.shutdown()
        shutil.rmtree(old, ignore_errors=True)
    if not shutil.which("node"):
        print("artifacts tab render skipped: node not found")
        return
    snap = j2.saved_snapshot(name, 0)
    lst = A.listing(run.dir)
    html = render_node(snap, docs={"docs": {}, "artifacts": lst})["artifacts"]
    assert html.index("<h3>Deliverables</h3>") < html.index("<h3>Evidence and records</h3>") and html.count("<h3>") == 2
    top, rest = html.split("<h3>Evidence and records</h3>")
    assert top.count('class="artrow"') == 4 and rest.count('class="artrow"') == len(lst["records"])
    for r in lst["deliverables"]:
        row = top[top.index(f'data-art="{r["file"]}"'):].split('<div class="artrow"')[0]
        assert f"<b>{r['label']}</b>" in row and f'style="background:var({r["color"]})">{r["tag"]}</span>' in row and "Open</button>" in row
        assert (f"made by {j2_ui_idruns(r['workflows'])}" in row) and r["form"] in row
    assert "depends on nothing in this run" in top.split('data-art="hld.md"')[0] and "depends on Product requirements" in top.split('data-art="hld.md"')[1].split('data-art="lld.md"')[0]
    assert "depends on High-level design, Low-level design" in top.split('data-art="site/index.html"')[1]
    assert "Open</button>" in rest and "iframe" not in html and "<img" not in html and "<pre" not in html      # nothing open by default
    # Open: a document shows readable text, the site a preview frame, a screenshot an image, a data file plain text
    lst["records"].append({**lst["records"][0], "name": "Screenshot v4", "type": "screenshots", "form": "image", "file": "site/shots/v4.png",
                           "open": "image", "tag": "Other", "color": "--other"})
    for f, mode, mark in (("prd.md", "text", '<pre class="arttext">'), ("site/index.html", "preview", '<iframe class="artframe" sandbox=""'),
                          ("site/shots/v4.png", "image", '<img class="artimg"'), ("atoms.json", "text", '<pre class="arttext">')):
        out = render_node(snap, docs={"docs": {}, "artifacts": lst, "artopen": f, "arttext": {f: "# Product\n<b>x</b>"}})["artifacts"]
        assert out.count(mark) == 1 and "Close</button>" in out, (f, mode)
        if mode == "text":
            assert "# Product\n&lt;b&gt;x&lt;/b&gt;" in out      # escaped, never run as markup
        assert f"/api/artifact?name={urllib.request.quote(f, safe='')}" in out or mode == "text"
    out = render_node(snap, docs={"docs": {}, "artifacts": {"deliverables": [], "records": []}})["artifacts"]
    assert "Nothing built yet." in out and "Nothing recorded yet." in out
    # the Documents and Charter tabs still render as before
    both = render_node(snap, docs={"docs": {k: {"asked": True, "text": open(os.path.join(run.dir, k + ".md"), encoding="utf-8").read()}
                                            for k in ("prd", "hld", "lld")}})
    assert 'id="docblock-prd"' in both["docs"] and "Charter" in both["charter"]
    print("artifacts: artifacts.json written at run end (additive), an old run gets the same rows from its folder, endpoints serve "
          "only listed files (live, saved, missing run), the tab shows two headings with tag, workflows, dependencies and an Open "
          "button per form ok")


def j2_ui_idruns(ids):
    """The page's idRuns(): W-1, W-2, W-3, W-5 -> "W-1 to W-3, W-5"."""
    nums, out, i = sorted(int(x[2:]) for x in ids), [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out += [f"W-{nums[i]} to W-{nums[j]}"] if j - i >= 2 else [f"W-{n}" for n in nums[i:j + 1]]
        i = j + 1
    return ", ".join(out)


def by_fn(fn, rng, calls):
    return sum(1 for c in calls if c["fn"] == fn and rng["from"] <= c["id"] <= rng["to"])


def map_logic():
    """Every run under runs/ with workflows.json: every line has a record (source) that exists in the run's files and points producer ->
    consumer; every recorded read between two drawn boxes is a line; every workflow is in exactly one stage's circle or under Not recorded;
    no deliverable sits above one it reads in the Tree; every label comes from its file, its type or its workflow. One row per run."""
    import j2_map as M
    import j2_artifacts as A
    repo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "runs")
    names = sorted(n for n in os.listdir(repo) if os.path.isfile(os.path.join(repo, n, "workflows.json"))) if os.path.isdir(repo) else []
    types = {t["file"]: t["label"] for t in A.load_types()["types"] if "*" not in t["file"]}
    by_type = {t["name"]: t["label"] for t in A.load_types()["types"]}
    rows_out, bad = [], []
    keep_dir, j2.RUNS_DIR = j2.RUNS_DIR, repo      # read only: building and drawing a map never writes to a run
    try:
        for name in names:
            d = os.path.join(repo, name)
            rd = lambda f, dflt=None: json.load(open(os.path.join(d, f), encoding="utf-8")) if os.path.isfile(os.path.join(d, f)) else dflt
            wfs = {w["id"]: w for w in rd("workflows.json", []) if isinstance(w, dict) and w.get("id")}
            named = {k: w["id"] for w in wfs.values() for k in (w.get("proposal_id"), w.get("name"), w["id"]) if k}
            plan = {p_.get("name"): p_ for p_ in (((rd("realization.json") or {}).get("plan") or {}).get("workflows") or []) if isinstance(p_, dict)}
            atoms = {a["id"]: a for a in (rd("atoms.json") or {}).get("atoms", [])}
            m = M.build(d)
            by = {n["id"]: n for n in m["nodes"]}
            holds = {}
            for n in m["nodes"]:
                if n["kind"] == "work":
                    for w in n.get("workflows") or []:
                        holds.setdefault(w, set()).add(n["id"])
            def proven(src):      # the record named by an edge is really in the run's files
                f, wid, field, val = src.get("file"), src.get("workflow"), src.get("field"), src.get("value")
                if f == "live run" or not f or not os.path.isfile(os.path.join(d, f)):
                    return False
                if wid and wid not in wfs:
                    return False
                w = wfs.get(wid) or {}
                if f == "workflows.json" and field in ("dependencies", "artifact", "stage", "section", "style_guide"):
                    v = w.get(field)
                    return val in v if isinstance(v, list) else v == val
                if f == "realization.json" and field == "inputs_from":
                    return val in ((plan.get(w.get("proposal_id")) or plan.get(w.get("name")) or {}).get("inputs_from") or [])
                if f == "atoms.json" and field == "depends_on":
                    return any(e.get("atom") == val for a in atoms.values() if a.get("workflow") == wid for e in a.get("depends_on") or [])
                return True      # a whole-file record (goal.json, chat.jsonl, calls.jsonl, audits, artifacts.json, summary.json)
            no_basis, wrong_way = 0, 0
            for e in m["edges"]:
                src = e.get("source")
                if not src or not proven(src):
                    no_basis += 1
                    continue
                if e["kind"] == "handoff":      # producer -> consumer: the reader makes the target; what it reads made the source
                    reader, read = src.get("workflow"), src.get("value")
                    read = named.get(read) or str(read).split(".")[0]
                    if reader not in by[e["to"]].get("workflows", []) or read not in by[e["from"]].get("workflows", []):
                        wrong_way += 1
            # every recorded read between two drawn boxes is a line
            have = {(e["from"], e["to"]) for e in m["edges"] if e["kind"] == "handoff"}
            want = set()
            for wid, w in wfs.items():
                p_ = plan.get(w.get("proposal_id")) or plan.get(w.get("name")) or {}
                srcs = [named.get(x) for x in (w.get("dependencies") or []) + (p_.get("inputs_from") or [])]
                srcs += [str(e.get("atom")).split(".")[0] for a in atoms.values() if a.get("workflow") == wid for e in a.get("depends_on") or []]
                for s_ in filter(None, srcs):
                    for a_ in holds.get(s_, ()):
                        for b_ in holds.get(wid, ()):
                            if a_ != b_ and s_ != wid:
                                want.add((a_, b_))
            missing = len(want - have)
            # the Tree: each workflow in exactly one stage's circle (or listed as not placed); no box above one it reads
            t = render_node({"workflows": [], "chat": []}, docs={"docs": {}, "map": m})["tree"]
            circles = [w for n in t["nodes"] if n["kind"] == "stage" for w in n["workflows"]]
            listed = {x for s_ in m["not_recorded"] if s_.startswith("workflows not placed on the map: ") for x in s_.split(": ", 1)[1].split(", ")}
            once = sum(1 for w in wfs if circles.count(w) + (w in listed) != 1)
            row = {n["id"]: n["row"] for n in t["nodes"]}
            misplaced = sum(1 for l in t["links"] if l["kind"] == "reads" and row[l["to"]] <= row[l["from"]]) + once
            # labels: from the file, the file's type, or the workflow (its name, or the field that records it)
            labels = 0
            for n in m["nodes"]:
                if n["kind"] != "work":
                    continue
                ok = {types.get(n.get("file")), M.file_label(n.get("file")) if n.get("file") else None, by_type.get(n["id"].split(":")[1])}
                if not n.get("file"):
                    ok.add(n.get("output"))      # planned, nothing written yet: named by its stage's output
                ok |= {(wfs.get(w) or {}).get("name") for w in n.get("workflows") or []}
                ok |= {M.file_label(f) for w in n.get("workflows") or [] for f, v in (wfs.get(w) or {}).items() if v is True}
                labels += n["label"] not in ok
            rows_out.append((name, len(m["edges"]), no_basis + wrong_way, missing, misplaced, labels, t["crossings"]))
            if no_basis or wrong_way or missing or misplaced or labels:
                bad.append(rows_out[-1])
    finally:
        j2.RUNS_DIR = keep_dir
    if "--map-logic" in sys.argv or bad:
        print(f"{'run':<22} {'edges':>5} {'no basis':>8} {'missing reads':>13} {'misplaced':>9} {'labels':>6} {'crossings':>9}")
        for r in rows_out:
            print(f"{r[0]:<22} {r[1]:>5} {r[2]:>8} {r[3]:>13} {r[4]:>9} {r[5]:>6} {r[6]:>9}")
    assert not bad, bad
    print(f"map logic: {len(rows_out)} runs, every line has a record in the run's files and points producer -> consumer, no recorded read is "
          "missing, every workflow in one stage's circle, nothing below what reads it, every label from its file, type or workflow ok")


def map_tree():
    """The Map's Tree view (default; "Line" is the metro drawing): Charter, the stages (one per team), their deliverables, a deliverable
    that reads others below them. Same run, same positions; every workflow sits in the tree or its stage's circle; no two labels
    overlap; the crossing count is reported. A clicked stage gets its workflows on a circle round it and the rest fades; Esc returns."""
    import html as htmllib
    import shutil
    import threading
    import j2_map as M
    import j2_server
    repo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "runs")
    names = [n for n in ("j2-20261026-174314", "j2-20261026-155610", "j2-20261027-130930") if os.path.isdir(os.path.join(repo, n))]
    if not names:
        print("map tree: runs 174314, 155610 and 130930 not in runs/, the tree checks on them are skipped")
    report = []
    for name in names:
        if not os.path.exists(os.path.join(j2.RUNS_DIR, name)):
            shutil.copytree(os.path.join(repo, name), os.path.join(j2.RUNS_DIR, name))      # read only: the test's own copy
        mp = M.build(os.path.join(j2.RUNS_DIR, name))
        snap = j2.saved_snapshot(name, 0)
        a, b = render_node(snap, docs={"docs": {}, "map": mp}), render_node(snap, docs={"docs": {}, "map": mp})
        t = a["tree"]
        assert t == b["tree"] and a["map"] == b["map"], f"{name}: the same run draws the same tree"
        assert 'data-view="tree"' in a["map"] and '<svg class="treesvg"' in a["map"] and 'data-mapview="tree" aria-pressed="true"' in a["map"], "Tree is the default"
        rows = {n["id"]: n["row"] for n in t["nodes"]}
        files = {n["id"]: n.get("file") for n in mp["nodes"] if n.get("tag") == "Spec"}
        for n in t["nodes"]:      # a spec is named from its file: artifacts/hint-method.md is "Hint method"
            if n["id"] in files:
                base = os.path.splitext(os.path.basename(files[n["id"]]))[0].replace("-", " ").replace("_", " ")
                assert n["label"] == base[:1].upper() + base[1:], (name, n["id"], n["label"], files[n["id"]])
        if name == "j2-20261026-174314":
            lab = {n["id"]: n["label"] for n in t["nodes"]}
            assert (lab["work:W-1"], lab["work:W-3"], lab["work:W-4"]) == ("Product spec", "Copy guide", "Hint method"), lab
            assert {"Curriculum", "Content plan"} <= {n["sub"].split(" · ")[0] for n in t["nodes"]}, "the stage output stays as the second line"
        if "work:style" in rows:      # style.json is written by the first build workflow and read by the later ones: an input of the site
            assert rows["work:style"] < rows["work:website"] and {"from": "work:style", "to": "work:website", "kind": "reads"} in t["links"], (name, rows)
            assert not any(l["from"] == "work:website" and l["to"] == "work:style" for l in t["links"]), "never below the site"
            reads_ = [l["from"] for l in t["links"] if l["to"] == "work:style" and l["kind"] == "reads"]
            assert rows["work:style"] == max([rows[x] for x in reads_] + [1]) + 1, ("right below the files its maker reads", rows)
        assert rows["work:charter"] == 0 and all(n["row"] == 1 for n in t["nodes"] if n["kind"] == "stage") and all(n["row"] >= 2 for n in t["nodes"] if n["kind"] == "deck")
        for l in t["links"]:
            if l["kind"] == "reads":
                assert rows[l["to"]] > rows[l["from"]], (name, l, "a deliverable sits below every deliverable it reads")
        boxes = [n["box"] for n in t["nodes"]]
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                p, q = boxes[i], boxes[j]
                assert p[2] <= q[0] or q[2] <= p[0] or p[3] <= q[1] or q[3] <= p[1], (name, t["nodes"][i]["id"], t["nodes"][j]["id"], "labels overlap")
        same_row = {}
        for n in t["nodes"]:
            same_row.setdefault(n["row"], []).append(n)
        for r in same_row.values():
            r.sort(key=lambda n: n["x"])
            assert all(r[k + 1]["box"][0] - r[k]["box"][2] >= 24 - 0.2 for k in range(len(r) - 1)), (name, "at least 24px between labels")
        # every workflow: in the tree (a deliverable's workflows) and in its stage's circle
        wids = [w["id"] for w in json.load(open(os.path.join(j2.RUNS_DIR, name, "workflows.json"), encoding="utf-8"))]
        stages = [n for n in t["nodes"] if n["kind"] == "stage"]
        circled = set()
        for st in stages:
            o = render_node(snap, docs={"docs": {}, "map": mp, "tree": {"open": st["id"]}})
            ot = o["tree"]
            assert ot["open"] == st["id"] and sorted(x["id"] for x in ot["ring"]) == sorted(st["workflows"]), (name, st["id"])
            c, rx, ry = ot["c"], ot["rx"], ot["ry"]
            assert all(abs(((x["x"] - c["x"]) / rx) ** 2 + ((x["y"] - c["y"]) / ry) ** 2 - 1) < 0.02 for x in ot["ring"]), "evenly on one ring round the stage"
            svg = o["map"][o["map"].index('<svg class="treesvg"'):o["map"].index("</svg>", o["map"].index('<svg class="treesvg"'))]
            for n in t["nodes"]:
                g = re.search(r'<g class="tnode ([^"]*)" data-id="%s"' % re.escape(htmllib.escape(n["id"])), svg)
                assert g and (("tdim" in g.group(1) and "thide" in g.group(1)) == (n["id"] != st["id"])), (name, st["id"], n["id"], "every other node fades, its words hidden")
            w_ = st["box"][2] - st["box"][0]
            shown = [x["box"] for x in ot["ring"]] + [[c["x"] - w_ / 2, c["y"] + st["r"] + 4, c["x"] + w_ / 2, c["y"] + st["r"] + 33]]      # the only words left: the circle's and the stage's
            for i in range(len(shown)):
                for j in range(i + 1, len(shown)):
                    p, q = shown[i], shown[j]
                    assert p[2] <= q[0] or q[2] <= p[0] or p[3] <= q[1] or q[3] <= p[1], (name, st["id"], "labels in the circle overlap", i, j)
            assert svg.count('class="tnode twf') == len(st["workflows"]) and svg.count('class="tspoke"') == len(st["workflows"])
            circled |= {x["id"] for x in ot["ring"]}
        placed = circled | {w for n in t["nodes"] for w in n["workflows"]}
        missing = [w for w in wids if w not in placed]
        assert not missing or all(f"workflows not placed on the map: " in x and all(w in x for w in missing) for x in mp["not_recorded"] if "not placed" in x), (name, missing)
        assert not missing, (name, "every workflow is in the tree or its stage's circle", missing)
        # a workflow picked in the circle: its panel (delivers, reads, read by, checks, status, Open) and its own links
        st = max(stages, key=lambda x: len(x["workflows"]))
        flows = {w["id"]: w for w in mp["workflows"]}
        wid = next((w for w in st["workflows"] if flows[w]["reads"]), st["workflows"][0])
        o = render_node(snap, docs={"docs": {}, "map": mp, "tree": {"open": st["id"], "wf": wid}})["map"]
        pn = o[o.index('data-wfpanel="%s"' % wid):]
        pn = pn[:pn.index("</aside>")]
        assert all(k in pn for k in ("status: ", "delivers: ", "reads: ", "read by: ", "checks: ", f'data-wfcard="{wid}">Open<')), pn
        grp = re.search(r'<g class="twflinks on" data-wf="%s">(.*?)</g>' % re.escape(wid), o).group(1)      # the picked one's links are drawn; the others' show on hover
        assert grp.count('class="tlink reads"') == len(flows[wid]["reads"]) and grp.count('class="tlink readby"') == sum(1 for w in flows.values() if wid in w["reads"]), (name, wid)
        assert o.count('class="twflinks on"') == 1, "only the picked workflow's links are drawn at rest"
        # the Line toggle: the metro drawing, as before
        ln = render_node(snap, docs={"docs": {}, "map": mp, "mapview": "line"})["map"]
        assert 'data-view="line"' in ln and '<svg class="mapsvg"' in ln and 'data-mapview="line" aria-pressed="true"' in ln
        report.append(f"{name[3:]} {t['crossings']} crossing{'s' if t['crossings'] != 1 else ''}")
    chrome = j2.j2_check.find_chrome()
    if chrome and BROWSER and names:      # the real page at 1440x900: the flow on hover, a stage's circle, one panel inside the map card, Esc, Line
        ui = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "j2_ui.html"), encoding="utf-8").read()
        name = names[0]

        def probe(js, flags=()):
            tmp = tempfile.mkdtemp(prefix="j2-tree-")
            open(os.path.join(tmp, "j2_ui.html"), "w", encoding="utf-8").write(ui.replace("</body>", "<script>setTimeout(() => { try { " + js + " } catch (e) { document.body.setAttribute('data-r', JSON.stringify({error: String(e.stack || e)})); } }, 2500)</script></body>"))
            keep_here, j2_server.HERE = j2_server.HERE, tmp
            srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            try:
                dom = chrome_dom([chrome, "--headless=new", "--disable-gpu", "--window-size=1440,900", "--virtual-time-budget=10000", *flags, "--dump-dom",
                                  f"http://127.0.0.1:{srv.server_address[1]}/?view=detail&run={name}"], capture_output=True, text=True, timeout=180).stdout
            finally:
                j2_server.HERE = keep_here
                srv.shutdown()
            got = json.loads(htmllib.unescape(re.search(r'data-r="([^"]*)"', dom).group(1)))
            assert "error" not in got, got
            return got
        base = ("const q = s => [...document.querySelectorAll(s)], r = {}, ev = (el, t, rel) => el.dispatchEvent(new MouseEvent(t, {bubbles: true, relatedTarget: rel || null})),"
                " vis = el => getComputedStyle(el).display !== 'none', touching = id => q('.tedge').filter(e => vis(e) && (e.dataset.from === id || e.dataset.to === id)),"
                " busiest = q('.tnodes > .tnode.tdeck').sort((a, b) => touching(b.dataset.id).length - touching(a.dataset.id).length)[0], hid = busiest.dataset.id;")
        js = (base + "const st = q('.tnode.tstage').sort((a, b) => b.querySelectorAll('*').length - a.querySelectorAll('*').length || +b.dataset.x - +a.dataset.x)[0], id = st.dataset.id; r.id = id;"
              "const s0 = document.querySelector('.treesvg').getBoundingClientRect(); r.fits = [s0.bottom <= innerHeight, s0.top >= 0, document.documentElement.scrollHeight <= innerHeight + 1];"
              "r.d3 = typeof d3; r.tree = getComputedStyle(document.querySelector('.treesvg')).display; r.line = getComputedStyle(document.querySelector('.mapsvg:not(.mapvert)')).display;"
              "r.rest = q('.tflow').length; r.hid = hid; ev(busiest, 'mouseover');"
              "r.flow = q('.tflow').map(c => [c.dataset.from, c.dataset.to, c.classList.contains('fin'), c.querySelector('animateMotion').getAttribute('path'), c.querySelector('animateMotion').getAttribute('dur')]);"
              "r.edges = touching(hid).map(e => [e.dataset.from, e.dataset.to, e.getAttribute('d'), e.classList.contains('hot')]); ev(busiest, 'mouseout'); r.after0 = q('.tflow').length;"
              "ev(st, 'click');"
              "setTimeout(() => { const g = q('.tnode.tstage').find(x => x.dataset.id === id), svg = document.querySelector('.treesvg');"
              " r.open = g.classList.contains('open'); r.at = g.getAttribute('transform'); r.mid = [+svg.dataset.w / 2, +svg.dataset.h / 2]; r.ring = g.querySelectorAll('.twf').length;"
              " r.dim = q('.tnodes > .tnode.tdim').length; r.all = q('.tnodes > .tnode').length; r.back = !document.querySelector('[data-treeback]').hidden;"
              " r.words = q('.tnodes > .tnode.thide .tlab').filter(x => x.getBoundingClientRect().width > 0).length;"
              " r.gone = q('.tedge').filter(e => vis(e) && e.dataset.from !== id && e.dataset.to !== id).length; r.stagePanel = q('#pane-map .mappanel').map(x => x.dataset.stagepanel || '');"
              " const w = g.querySelector('.twf'); ev(w, 'mouseover'); r.wflow = q('.tflow').length; r.wlinks = q('.twflinks.hov .tlink').length; ev(w, 'mouseout');"
              " ev(w, 'click'); const box = r_ => { const b = r_.getBoundingClientRect(); return [b.left, b.top, b.right, b.bottom]; };"
              " r.panels = q('#pane-map .mappanel').length; r.pb = box(document.querySelector('.treepanel')); r.card = box(document.querySelector('.treebox')); r.chat = box(document.querySelector('#dview .chat')); r.needs = box(document.getElementById('needs'));"
              " r.wfcard = !!document.querySelector('.treepanel [data-wfcard]');"
              " ev(q('.tnodes > .tnode.tdeck')[0], 'click'); r.second = q('#pane-map .mappanel').map(x => x.dataset.panel || x.dataset.wfpanel || x.dataset.stagepanel);"
              " document.dispatchEvent(new KeyboardEvent('keydown', {key: 'Escape', bubbles: true})); r.esc1 = q('#pane-map .mappanel').map(x => x.dataset.stagepanel || x.dataset.panel || x.dataset.wfpanel);"
              " document.dispatchEvent(new KeyboardEvent('keydown', {key: 'Escape', bubbles: true}));"
              " setTimeout(() => { const h = q('.tnode.tstage').find(x => x.dataset.id === id); r.after = [q('.tnode.open').length, q('.tdim').length, h.getAttribute('transform') === 'translate(' + h.dataset.x + ',' + h.dataset.y + ')', q('.twf').length, q('#pane-map .mappanel').length];"
              "  document.querySelector('[data-mapview=line]').click(); r.line2 = getComputedStyle(document.querySelector('.mapsvg:not(.mapvert)')).display; r.tree2 = getComputedStyle(document.querySelector('.treebox')).display;"
              "  document.querySelector('[data-mapview=tree]').click(); document.body.setAttribute('data-r', JSON.stringify(r)); }, 900); }, 900);")
        got = probe(js)
        assert got["d3"] == "object" and got["tree"] == "block" and got["line"] == "none", got
        assert got["fits"] == [True, True, True], ("the whole tree fits 1440x900 at load, no scrolling", got)
        # the flow: nothing moves at rest; on hover one dot per lit line, along the line's own path (producer -> consumer), about 1.5 s,
        # "reads" (into the node) and "read by" (out of it) told apart; leaving the node stops it
        assert got["rest"] == 0 and got["after0"] == 0, got
        lines = {(a, b): (d, hot) for a, b, d, hot in got["edges"]}
        assert got["flow"] and len(got["flow"]) == len(lines) and all(h for _, h in lines.values()), got
        for a, b, fin, path, dur in got["flow"]:
            assert (a, b) in lines and path == lines[(a, b)][0] and fin == (b == got["hid"]) and dur == "1.5s", (a, b, fin, got["hid"])
        assert {f[2] for f in got["flow"]} == {True, False} or len(got["flow"]) == 1, "lines in and lines out, when it has both"
        # a stage: the middle, its circle, every other node faded with its words hidden, only its own lines drawn, its panel the only one
        x, y = (float(v) for v in re.match(r"translate\(([\d.]+),([\d.]+)\)", got["at"]).groups())
        assert got["open"] and abs(x - got["mid"][0]) < 1 and abs(y - got["mid"][1]) < 1 and got["ring"] > 0 and got["back"], got
        assert got["dim"] == got["all"] - 1 and got["words"] == 0 and got["gone"] == 0 and got["stagePanel"] == [got["id"]], got
        assert got["wflow"] == got["wlinks"], ("a circle workflow's links flow on hover", got)
        # the workflow panel: inside the map card, never over the Chat or the Needs you strip; a second item replaces it
        pb, card, chat, needs = got["pb"], got["card"], got["chat"], got["needs"]
        inside = pb[0] >= card[0] - 0.5 and pb[1] >= card[1] - 0.5 and pb[2] <= card[2] + 0.5 and pb[3] <= card[3] + 0.5
        meets = lambda a, b: a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]
        assert got["panels"] == 1 and got["wfcard"] and inside and not meets(pb, chat) and not meets(pb, needs) and pb[3] <= 900, got
        assert len(got["second"]) == 1 and got["second"][0].startswith("work:"), ("one panel at a time", got)
        assert got["esc1"] == [got["id"]] and got["after"] == [0, 0, True, 0, 0], ("Esc closes the panel, then the stage", got)
        assert got["line2"] != "none" and got["tree2"] == "none", "the Line toggle shows the metro map"
        # reduced motion: no dots, small arrowheads on the lit lines instead
        rm = probe(base + "ev(busiest, 'mouseover'); r.rm = [q('.tflow').length, touching(hid).filter(e => e.getAttribute('marker-end')).length, touching(hid).length];"
                   "document.body.setAttribute('data-r', JSON.stringify(r));", ("--force-prefers-reduced-motion",))
        assert rm["rm"][0] == 0 and rm["rm"][1] == rm["rm"][2] > 0, ("reduced motion: arrowheads, nothing moves", rm)
    print("map tree: Charter, stages, deliverables below what they read; the same run draws the same; every workflow in the tree or its stage's "
          "circle; no labels overlap, 24px apart; a stage's circle round it with the rest faded, a picked workflow's panel and links, Esc returns; "
          "Line shows the metro map ok (" + ", ".join(report) + ")")


def map_every_workflow():
    """Map: every workflow in workflows.json is on the map (by stage, then section, then output refs) or listed by id under
    not recorded; a realization run's specs are deliverables with their stage's role and "given to" lines to the build
    workflows that read them; a workflow set to done whose atom did not pass shows "done · checks did not pass" (amber),
    counted in the team label; the style guide is a small deliverable."""
    import j2_map as M
    repo = os.path.join(os.path.dirname(os.path.abspath(__file__)), "runs")      # read-only: building a map never writes to a run
    runs = [os.path.join(repo, n) for n in sorted(os.listdir(repo)) if n.startswith("j2-")] if os.path.isdir(repo) else []
    seen = 0
    for d in runs:
        wfs = j2._read(d, "workflows.json", None)
        if not isinstance(wfs, list):
            continue
        m = M.build(d)
        on = {x for n in m["nodes"] for x in n.get("workflows") or []}
        listed = {x for t in m["not_recorded"] if t.startswith("workflows not placed on the map: ") for x in t.split(": ", 1)[1].split(", ")}
        miss = [w["id"] for w in wfs if w["id"] not in on | listed]
        assert not miss, (d, miss)
        assert not on & listed, (d, on & listed)
        seen += 1
    r = os.path.join(repo, "j2-20261026-155610")
    if not os.path.isdir(r):
        print(f"map: every workflow placed or listed in {seen} saved runs (run 155610 not here: its checks skipped) ok")
        return
    m = M.build(r)
    by = {n["id"]: n for n in m["nodes"]}
    work = [n["id"] for n in m["nodes"] if n["kind"] == "work" and n["lane"] == "trunk"]
    assert work == ["work:charter", "work:W-1", "work:W-2", "work:W-3", "work:style", "work:website"], work      # Charter -> specs -> style guide -> Website
    assert [(by[k]["label"], by[k].get("output"), by[k]["file"], by[k]["status"]) for k in work[1:4]] == [("Product spec", "Product plan", "artifacts/product-spec.md", "filed"),
        ("Design guide", "Design guide", "artifacts/design-guide.md", "filed"), ("Architecture", "Architecture (HLD + LLD)", "artifacts/architecture.md", "filed")]
    assert [by[f"step:{a}>{b}"]["team"] for a, b in (("charter", "W-1"), ("W-1", "W-2"), ("charter", "W-3"), ("W-3", "style"), ("style", "website"))] == \
        ["Product Manager", "Design Lead", "Solution Architect", "Web Developer", "Web Developer"]      # the architecture reads the product spec, not the design guide
    given = {(e["from"], e["to"]): e["label"] for e in m["edges"] if e["kind"] == "handoff"}
    assert given[("work:W-1", "work:W-2")] == "given to W-2" and given[("work:W-1", "work:W-3")] == "given to W-3" and ("work:W-2", "work:W-3") not in given, given
    assert given[("work:style", "work:website")].startswith("given to W-5, ") and ("work:website", "work:style") not in given, "the style guide feeds the site"
    site, step = by["work:website"], by["step:style>website"]
    assert {"W-7", "W-9"} <= set(site["workflows"]) and site["status"] == "done_unchecked" and step["unchecked"] == ["W-7", "W-9"], (site, step)
    assert step["detail"].endswith("· 2 done, checks did not pass (W-7, W-9)"), step["detail"]
    assert (by["work:style"]["file"], by["work:style"]["workflows"], by["work:style"]["lane"]) == ("style.json", ["W-4"], "trunk")
    keep_dir, j2.RUNS_DIR = j2.RUNS_DIR, repo
    try:
        out = render_node(j2.saved_snapshot("j2-20261026-155610", 0), docs={"docs": {}, "map": m})["map"]
    finally:
        j2.RUNS_DIR = keep_dir
    wide = out[out.index('<svg class="mapsvg"'):out.index("</svg>")]
    assert all(t in wide for t in ("Product Manager", "Design Lead", "Solution Architect", "Web Developer</tspan> · 14 workflows · 2 done, checks did not pass")), wide[:400]
    assert 'stroke="var(--validator)" stroke-width="4" stroke-dasharray="2 3"' in wide and "done · checks did not pass" in out, "amber and dotted, never colour alone"
    print(f"map: every workflow placed or listed in {seen} saved runs; run 155610 Charter -> three specs (Product Manager, Design Lead, "
          "Solution Architect) -> Website, given-to lines, W-7 and W-9 done · checks did not pass (amber), style guide ok")


def map_tab():
    """Map v1 (map/1): map.json written beside atoms.json (additive, after the summary and again after the report), the
    same map from a rebuilt folder, an old run, a live run, /api/map, the SVG for every status (never colour alone), the
    card list for small screens, plain words, and (with Chrome) no sideways scroll at 500px."""
    import copy
    import shutil
    import threading
    import urllib.error
    import urllib.request
    import collections
    import html as htmllib
    import j2_map as M
    import j2_server
    run = docs_run(DESIGN_BRIEF, 900)
    name = os.path.basename(run.dir)
    saved = json.load(open(os.path.join(run.dir, "map.json")))
    atoms = json.load(open(os.path.join(run.dir, "atoms.json")))["atoms"]
    assert os.path.exists(os.path.join(run.dir, "atoms.json")) and saved["schema"] == "map/1" and saved["run"] == name and not saved["live"]
    # golden structure: the Requester, then the work in an order where each box follows what it reads, a team (step) before each box.
    # A step's id names the box before it only when the box reads it (else the Charter): the style guide reads the HLD, not the LLD.
    assert [n["id"] for n in saved["nodes"]] == ["person:requester", "fn:identity", "fn:adaptation", "fn:priority", "fn:coordination", "fn:audit", "work:charter", "step:charter>prd", "work:prd", "step:prd>hld",
        "work:hld", "step:hld>lld", "work:lld", "step:charter>style", "work:style", "step:style>website", "work:website"], [n["id"] for n in saved["nodes"]]
    assert [(n["kind"], n["row"]) for n in saved["nodes"]] == [("person", "people")] + [("function", "functions")] * 5 + [(k, "work") for k in
        ("work", "engine", "work", "engine", "work", "engine", "work", "engine", "work", "engine", "work")]
    by = {n["id"]: n for n in saved["nodes"]}
    assert [by[f"work:{k}"]["status"] for k in ("charter", "prd", "hld", "lld", "website")] == ["filed"] * 5
    assert [by[f"work:{k}"]["label"] for k in ("prd", "hld", "lld", "website")] == ["Product requirements", "High-level design", "Low-level design", "Website"]
    assert [by[f"work:{k}"]["tag"] for k in ("prd", "hld", "lld", "website")] == ["PRD", "HLD", "LLD", "Site"]
    assert [by[f"work:{k}"]["color"] for k in ("prd", "hld", "lld", "website")] == ["--prd", "--hld", "--lld", "--site"]
    rows = {r["type"]: r for r in json.load(open(os.path.join(run.dir, "artifacts.json")))["rows"]}
    for k in ("prd", "hld", "lld", "website"):      # versions and files come from artifacts.json, never invented
        assert by[f"work:{k}"]["version"] == rows[k]["version"] and by[f"work:{k}"]["file"] == rows[k]["file"], k
    assert by["work:charter"]["version"] is None and by["work:charter"]["file"] == "goal.json"
    assert [by[f"step:{a}>{b}"]["label"] for a, b in (("charter", "prd"), ("prd", "hld"), ("hld", "lld"), ("style", "website"))] == ["Write", "Write", "Write", "Build"]
    for k, step in (("prd", "charter>prd"), ("hld", "prd>hld"), ("lld", "hld>lld"), ("website", "style>website")):      # counts equal the atoms
        mine = [a for a in atoms if a["workflow"] in rows[k]["workflows"]]
        ok = sum(1 for a in mine for x in a["verification"] if x["met"] is True)
        tot = sum(1 for a in mine for x in a["verification"] if x["met"] is not None)
        assert by[f"step:{step}"]["detail"] == f"{len(rows[k]['workflows'])} workflows · {ok} of {tot} checks passed", (k, by[f"step:{step}"]["detail"])
    assert {(e["from"], e["to"], e["basis"]) for e in saved["edges"] if e["kind"] == "handoff"} == {      # recorded reads only, producer -> consumer
        ("work:prd", "work:hld", "recorded"), ("work:prd", "work:lld", "recorded"), ("work:hld", "work:lld", "recorded"), ("work:hld", "work:style", "recorded"),
        ("work:hld", "work:website", "recorded"), ("work:lld", "work:website", "recorded"), ("work:style", "work:website", "recorded")}
    assert all(e.get("source") for e in saved["edges"]), "every line carries its record"
    assert sum(1 for e in saved["edges"] if e["kind"] == "step") == 5 * 2 + 3      # the Charter into each of 5 teams, each team into its box, 3 recorded reads of the box before
    ck = [x for a in atoms for x in a["verification"] if x["met"] is not None]
    assert saved["status"] == {"state": "Done", "state_status": "idle", "checks": {"passed": sum(x["met"] is True for x in ck), "run": len(ck)}}
    assert "version of Charter" in saved["not_recorded"] and "which other function worked on Website (only Audit is recorded)" in saved["not_recorded"] and \
        any("Knows row" in x for x in saved["not_recorded"]) and not any("Functions row" in x for x in saved["not_recorded"]) and \
        not any(n["kind"] == "know" for n in saved["nodes"]) and "messages to and from workflow agents (W-n) are not drawn" in saved["not_recorded"]
    assert [l["key"] for l in saved["legend"] if not l["drawn"]] == []
    assert saved == M.build(run.dir), "map.json is what the folder builds"
    # the Functions row, golden: calls, seconds and tokens per function from calls.jsonl; a finished run is never "running"
    calls = [json.loads(x) for x in open(os.path.join(run.dir, "calls.jsonl"), encoding="utf-8")]
    chat = [json.loads(x) for x in open(os.path.join(run.dir, "chat.jsonl"), encoding="utf-8")]
    fnn = [n for n in saved["nodes"] if n["kind"] == "function"]
    assert [(n["label"], n["tab"], n["color"]) for n in fnn] == [("Identity", "charter", "--identity"), ("Adaptation", "workflows", "--adaptation"),
        ("Priority", "conflicts", "--priority"), ("Coordination", "timeline", "--coordination"), ("Audit", "audit", "--audit")]
    for n in fnn:
        fid, mine = n["id"][3:], [c for c in calls if c["fn"] == n["id"][3:]]
        said = [m for m in chat if m.get("who") == fid]
        assert n["calls"] == len(mine) and n["seconds"] == round(sum(c["seconds"] or 0 for c in mine), 1) and n["tokens"] == {
            "input": sum((c["tokens"] or {}).get("input", 0) for c in mine), "output": sum((c["tokens"] or {}).get("output", 0) for c in mine)}, n["id"]
        assert n["status"] == "idle" and n["detail"] == "last seen: " + said[-1]["text"][:140], n["id"]
    assert not any(n["status"] in ("running", "waiting") for n in fnn) and "running" not in json.dumps([n["status"] for n in saved["nodes"] if n["row"] != "work"])
    ud = [(e["from"], e["to"], e["n"]) for e in saved["edges"] if e["kind"] == "speaks_to" and not e["to"].startswith("work:")]      # speaks to: counts from the chat
    want = collections.Counter((m["who"], m["next"]) for m in chat if m.get("who") in ("identity", "adaptation", "priority", "coordination", "audit")
                               and m.get("next") in ("identity", "adaptation", "priority", "coordination", "audit") and m["who"] != m["next"])
    assert ud[:2] == [("person:requester", "fn:identity", sum(1 for m in chat if m.get("who") == "user")),
                      ("fn:identity", "person:requester", sum(1 for m in chat if m.get("who") == "identity" and m.get("to_user")))]
    assert {(a[3:], b[3:]): c for a, b, c in ud[2:]} == dict(want) and all(a.startswith("fn:") for a, _, _ in ud[2:])
    docs_rec = json.load(open(os.path.join(run.dir, "documents.json")))      # who made each step: only where documents.json has the call range
    for k, step in (("prd", "charter>prd"), ("hld", "prd>hld"), ("lld", "hld>lld")):
        rng, st = docs_rec[k]["calls"], next(n for n in saved["nodes"] if n["id"] == f"step:{step}")
        per = {f: sum(1 for c in calls if c["fn"] == f and rng["from"] <= c["id"] <= rng["to"]) for f in ("identity", "adaptation", "priority", "coordination", "audit")}
        assert st["made_by"] == {"calls": rng, "functions": {f: n for f, n in per.items() if n}, "basis": "recorded"}, step
        assert {(e["from"][3:], e["n"], e["label"]) for e in saved["edges"] if e["kind"] == "speaks_to" and e["to"] == f"work:{k}"} == {(f, n, "calls") for f, n in per.items() if n}
    assert next(n for n in saved["nodes"] if n["id"] == "step:style>website")["made_by"] is None      # the site: not recorded, never guessed
    # function to work, golden: recorded data only. Identity to the Charter; each function to a document by its calls in the recorded range;
    # the site and a deed get only Audit (its audits in stage-*.json); nothing is a "made" edge any more
    lines = [(e["from"][3:], e["to"][5:], e["n"], e["label"], e["basis"]) for e in saved["edges"] if e["kind"] == "speaks_to" and e["to"].startswith("work:")]
    stage_web = json.load(open(os.path.join(run.dir, "stage-website.json")))
    exp = [("identity", "charter", None, "wrote the charter", "recorded")]
    for k in ("prd", "hld", "lld"):
        rng = docs_rec[k]["calls"]
        exp += [(f, k, by_fn(f, rng, calls), "calls", "recorded") for f in ("identity", "adaptation", "priority", "coordination", "audit") if by_fn(f, rng, calls)]
    exp += [("audit", "website", len(stage_web["audits"]), "audits", "recorded")]
    assert lines == exp, (lines, exp)
    assert not any(e["kind"] == "made" for e in saved["edges"]) and not [e for e in saved["edges"] if e["kind"] == "speaks_to" and e["to"].startswith("step:")]
    assert not [x for x in lines if x[1] == "website" and x[0] != "audit"], "the site gets only the Audit line"
    for k in ("prd", "hld", "lld"):      # every document with an audit in its stage file has an Audit line
        assert json.load(open(os.path.join(run.dir, f"stage-{k}.json")))["audits"] and any(x[0] == "audit" and x[1] == k for x in lines)
    assert by["person:requester"]["detail"] == "1 message from you" and by["person:requester"]["status"] == "idle"
    # a plain website run has the Charter and the site only; a PRD-only run has no site
    plain = docs_run(BREW_LAB, 200)
    assert [n["id"] for n in json.load(open(os.path.join(plain.dir, "map.json")))["nodes"]] == ["person:requester", "fn:identity", "fn:adaptation", "fn:priority", "fn:coordination", "fn:audit", "work:charter", "step:charter>style", "work:style", "step:style>website", "work:website"]
    # an old run: no artifacts.json, atoms.json or map.json; the same boxes, and the gaps are named
    old = os.path.join(j2.RUNS_DIR, "j2-20200102-000000")
    shutil.copytree(run.dir, old)
    for f in ("artifacts.json", "atoms.json", "map.json"):
        os.remove(os.path.join(old, f))
    m_old = M.build(old)
    box_ = lambda mp: sorted((n["id"], n["status"], n["version"]) for n in mp["nodes"] if n["kind"] != "engine")
    assert box_(m_old) == box_(saved), "the same boxes"
    assert not any(e["kind"] == "handoff" for e in m_old["edges"]), "no atoms.json: no recorded reads, so no read lines (never guessed from run order)"
    assert "artifacts.json (rows rebuilt from the folder)" in m_old["not_recorded"] and "atoms.json (atoms rebuilt from workflows.json)" in m_old["not_recorded"]
    assert not any(os.path.exists(os.path.join(old, f)) for f in ("artifacts.json", "atoms.json", "map.json")), "building a map never writes to an old run"
    os.remove(os.path.join(old, "calls.jsonl"))
    m2 = M.build(old)
    assert all(n["calls"] == 0 for n in m2["nodes"] if n["kind"] == "function") and "AI calls per function (calls.jsonl)" in m2["not_recorded"]
    assert all(n["detail"].startswith("last seen: ") for n in m2["nodes"] if n["kind"] == "function"), "the chat still says what each did"
    os.remove(os.path.join(old, "chat.jsonl"))
    m3 = M.build(old)
    f3 = [n for n in m3["nodes"] if n["kind"] == "function"]
    assert len(f3) == 5 and all(n["detail"] == "no message recorded" and n["status"] == "idle" for n in f3)
    assert "what Identity did (chat)" in m3["not_recorded"] and next(n for n in m3["nodes"] if n["id"] == "person:requester")["detail"] == "not recorded"
    assert not [e for e in m3["edges"] if e["kind"] == "speaks_to" and not e["to"].startswith("work:")] and not [e for e in m3["edges"] if e["kind"] == "made"]
    # an old run with no call ranges (no documents.json): only the Charter line and Audit's audits are drawn, the rest is listed as not recorded
    for f in [x for x in os.listdir(old) if x == "documents.json"]:
        os.remove(os.path.join(old, f))
    m4 = M.build(old)
    l4 = [(e["from"][3:], e["to"][5:], e["label"]) for e in m4["edges"] if e["kind"] == "speaks_to" and e["to"].startswith("work:")]
    assert l4 == [("identity", "charter", "wrote the charter"), ("audit", "prd", "audits"), ("audit", "hld", "audits"), ("audit", "lld", "audits"), ("audit", "website", "audits")], l4
    for lab in ("Product requirements", "High-level design", "Low-level design", "Website"):
        assert f"which other function worked on {lab} (only Audit is recorded)" in m4["not_recorded"], lab
    for f in [x for x in os.listdir(old) if x.startswith("stage-")]:
        os.remove(os.path.join(old, f))
    m5 = M.build(old)
    l5 = [(e["from"][3:], e["to"][5:]) for e in m5["edges"] if e["kind"] == "speaks_to" and e["to"].startswith("work:")]
    assert l5 == [("identity", "charter"), ("audit", "website")] and "which function worked on Product requirements: not recorded" in m5["not_recorded"], l5      # the site's audits are still in audits.json
    assert M.build(None)["nodes"] == [] and M.build(os.path.join(run.dir, "nope"))["not_recorded"] == ["run not found"]
    # statuses: a finished run says "left open", never "running"; fault and not made
    edit = os.path.join(j2.RUNS_DIR, "j2-20200103-000000")
    shutil.copytree(run.dir, edit)
    doc = json.load(open(os.path.join(edit, "atoms.json")))
    pick = lambda k: next(a for a in doc["atoms"] if a["workflow"] in rows[k]["workflows"])
    pick("hld")["status"], pick("lld")["status"] = "open", "verify_failed"
    json.dump(doc, open(os.path.join(edit, "atoms.json"), "w"))
    ewf = json.load(open(os.path.join(edit, "workflows.json")))
    json.dump([{**w, "status": "done"} for w in ewf], open(os.path.join(edit, "workflows.json"), "w"))
    st = {n["id"]: n["status"] for n in M.build(edit)["nodes"]}      # set to done but the atom did not pass: amber, never filed and checked
    assert st["work:hld"] == st["work:lld"] == "done_unchecked", st
    for w in ewf:      # the workflows that own the open and the failed atom are not done: left open and fault
        w["status"] = {pick("hld")["workflow"]: "running", pick("lld")["workflow"]: "failed"}.get(w["id"], w["status"])
    json.dump(ewf, open(os.path.join(edit, "workflows.json"), "w"))
    json.dump({"schema": "artifacts/1", "rows": [r for r in rows.values() if r["type"] != "website"] + [r for r in json.load(open(os.path.join(edit, "artifacts.json")))["rows"] if r["kind"] == "record"]}, open(os.path.join(edit, "artifacts.json"), "w"))
    me = M.build(edit)
    st = {n["id"]: n["status"] for n in me["nodes"]}
    assert st["work:hld"] == "left_open" and st["work:lld"] == "fault" and st["work:website"] == "not_made" and "running" not in json.dumps(st)      # the summary still lists the site as planned
    # a live run: open means running now, the pill names who is working, a pending question is waiting on you, planned stages are not made
    snap = j2.saved_snapshot(name, 0)
    snap["atoms_v1"], snap["workflows"] = {a["id"]: a for a in doc["atoms"]}, ewf
    lane = {"id": "identity", "name": "Identity", "status": "working"}
    live = M.build(edit, {"snapshot": {**snap, "status": "running", "lanes": [lane], "paused": None}, "stages": ["prd", "hld", "lld", "send_email", "website"],
                          "stage_words": {"send_email": "the test email"}})
    ls = {n["id"]: n for n in live["nodes"]}
    assert live["live"] and ls["work:hld"]["status"] == "running" and ls["work:lld"]["status"] == "fault" and live["status"]["state"] == "Identity running"
    assert ls["work:website"]["status"] == "filed"      # built: a live run reads the folder as it is
    prd_only = docs_run("Write a PRD for Campus Notes." + CAMPUS, 400)      # planned but not built: hollow, "not started", the deed named from the stage words
    pl = M.build(prd_only.dir, {"snapshot": {**j2.saved_snapshot(os.path.basename(prd_only.dir), 0), "status": "running", "lanes": [], "paused": None},
                                "stages": ["prd", "hld", "send_email", "website"], "stage_words": {"send_email": "the test email"}})
    ps = {n["id"]: n for n in pl["nodes"]}
    assert [n["id"] for n in pl["nodes"] if n["kind"] == "work"] == ["work:charter", "work:prd", "work:hld", "work:send_email", "work:website"], list(ps)
    assert all(ps[k]["status"] == "not_made" and ps[k]["detail"] == "not started" and ps[k]["file"] is None for k in ("work:hld", "work:send_email", "work:website"))
    assert ps["work:send_email"]["label"] == "The test email" and ps["work:hld"]["label"] == "High-level design" and ps["work:prd"]["status"] == "filed"
    ask = M.build(edit, {"snapshot": {**snap, "status": "waiting", "lanes": [], "paused": {"kind": "limit"}}, "stages": [], "stage_words": {}})
    assert ask["status"] == {**ask["status"], "state": "Waiting on you", "state_status": "waiting_on_you"} and \
        {n["id"]: n["status"] for n in ask["nodes"]}["person:requester"] == "waiting_on_you"
    # live lanes: working = running now, waiting on a function, waiting on you, blocked, idle (last seen)
    lanes = [{"id": "identity", "name": "Identity", "status": "working", "doing": "writing the charter"},
             {"id": "adaptation", "name": "Adaptation", "status": "waiting", "on": "Priority", "doing": "proposed three workflows"},
             {"id": "priority", "name": "Priority", "status": "waiting", "on": "Requester", "doing": "asked a question"},
             {"id": "coordination", "name": "Coordination", "status": "blocked", "doing": "W-3 is held"},
             {"id": "audit", "name": "Audit", "status": "idle", "doing": "checked site v2"}]
    lv = M.build(edit, {"snapshot": {**snap, "status": "running", "lanes": lanes, "paused": None, "call_log": calls, "chat": chat}, "stages": [], "stage_words": {}})
    f = {n["id"]: n for n in lv["nodes"] if n["kind"] == "function"}
    assert (f["fn:identity"]["status"], f["fn:identity"]["detail"]) == ("running", "writing the charter") and lv["status"]["state"] == "Identity running"
    assert (f["fn:adaptation"]["status"], f["fn:adaptation"]["wait_on"]) == ("waiting", "Priority")
    assert (f["fn:priority"]["status"], f["fn:priority"]["wait_on"]) == ("idle", None), "waiting on you only while a question to you is open (the same source as Needs you)"
    lq = M.build(edit, {"snapshot": {**snap, "status": "waiting", "lanes": lanes, "paused": {"kind": "info", "questions": [{"question": "Which day?", "options": []}], "msg": 1}, "call_log": calls, "chat": chat}, "stages": [], "stage_words": {}})
    assert {n["id"]: n for n in lq["nodes"]}["fn:priority"]["status"] == "waiting_on_you", "with a question open it does wait on you"
    # the Map and Needs you always agree on "waiting on you": same source (an open question), checked for each kind of state
    for paused_, idn_lane in ((None, "waiting"), (None, "idle"), ({"kind": "info", "questions": [{"question": "Which day?", "options": []}], "msg": 1, "label": "a fact"}, "waiting")):
        ln_ = [{"id": "identity", "name": "Identity", "status": idn_lane, "on": "Requester", "doing": "asked"}]
        mm_ = M.build(edit, {"snapshot": {**snap, "status": "waiting", "lanes": ln_, "paused": paused_, "call_log": calls, "chat": chat}, "stages": [], "stage_words": {}})
        on_map = any(n["status"] == "waiting_on_you" for n in mm_["nodes"] if n["kind"] in ("function", "person"))
        cnt = render_node({**j2.saved_snapshot(name, 0), "paused": paused_, "chat": chat}, docs={"docs": {}, "map": mm_})["needcount"]
        assert on_map == (cnt > 0), (paused_, idn_lane, on_map, cnt)
    assert (f["fn:coordination"]["status"], f["fn:coordination"]["detail"]) == ("fault", "W-3 is held")
    assert (f["fn:audit"]["status"], f["fn:audit"]["detail"]) == ("idle", "last seen: checked site v2")
    assert f["fn:identity"]["calls"] == sum(1 for c in calls if c["fn"] == "identity")
    pz = M.build(edit, {"snapshot": {**snap, "status": "waiting", "paused": {"kind": "info"}, "call_log": calls, "chat": chat,
                                     "lanes": [{"id": "identity", "name": "Identity", "status": "waiting", "on": "Requester", "doing": "asked a question"}]},
                        "stages": [], "stage_words": {}})
    pf = {n["id"]: n for n in pz["nodes"]}
    assert pf["fn:identity"]["status"] == "waiting_on_you" and pf["person:requester"]["status"] == "waiting_on_you" and pz["status"]["state"] == "Waiting on you"
    assert all(pf[f"fn:{x}"]["status"] == "idle" for x in ("adaptation", "priority", "coordination", "audit"))
    # a saved run that ended on a question: Identity and the Requester say "waiting on you", the rest "last seen"
    ended = os.path.join(j2.RUNS_DIR, "j2-20200104-000000")
    shutil.copytree(run.dir, ended)
    with open(os.path.join(ended, "chat.jsonl"), "a", encoding="utf-8") as fh:
        fh.write(json.dumps({"id": 9999, "who": "identity", "text": "@Requester Add credits?", "to_user": True, "pause": "limit", "next": "user"}) + "\n")
    sv = {n["id"]: n for n in M.build(ended)["nodes"]}
    assert sv["fn:identity"]["status"] == "waiting_on_you" and sv["person:requester"]["status"] == "waiting_on_you" and M.build(ended)["status"]["state"] == "Waiting on you"
    assert all(sv[f"fn:{x}"]["status"] == "idle" and sv[f"fn:{x}"]["detail"].startswith("last seen: ") for x in ("adaptation", "priority", "coordination", "audit"))
    shutil.rmtree(ended, ignore_errors=True)
    # /api/map: live (idle run reads its folder), saved, old, missing
    keep, j2_server.RUN = j2_server.RUN, run
    srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    get = lambda path: json.load(urllib.request.urlopen(base + path))
    try:
        for q in ("", f"?run={name}"):
            assert get("/api/map" + q) == saved, q
        assert box_(get("/api/map?run=j2-20200102-000000")) == box_(saved)
        assert get("/api/map?run=j2-20990101-000000")["nodes"] == [] and get("/api/map?run=../x")["nodes"] == []
        assert get("/api/artifacts")["deliverables"] and urllib.request.urlopen(base + "/api/artifact?name=goal.json").read()      # the box's Open view
    finally:
        j2_server.RUN = keep
        srv.shutdown()
        for f in ("j2-20200102-000000", "j2-20200103-000000"):
            shutil.rmtree(os.path.join(j2.RUNS_DIR, f), ignore_errors=True)
    if not shutil.which("node"):
        print("map render skipped: node not found")
        return
    # the SVG: every status has its own shape and mark, boxes link to Artifacts, full text in a tooltip, plain words
    nodes = [{"id": "person:requester", "kind": "person", "row": "people", "label": "You", "status": "waiting_on_you", "version": None, "detail": "2 messages from you", "file": None, "tag": None, "color": "--user", "workflows": []}]
    for i, (stt, color) in enumerate((("filed", "--prd"), ("running", "--hld"), ("left_open", "--lld"), ("fault", "--site"), ("not_made", "--other"), ("waiting_on_you", "--charter"))):
        nodes.append({"id": f"work:t{i}", "kind": "work", "row": "work", "label": f"Box {i} with a very long name", "status": stt, "version": i + 1 if i != 5 else None,
                      "detail": "Every done-when item passes inside its envelope", "file": f"f{i}.md", "tag": "X", "color": color, "workflows": []})
        if i < 5:
            nodes.append({"id": f"step:{i}", "kind": "engine", "row": "work", "label": "Write", "status": stt, "version": None, "detail": "3 workflows · 4 of 4 checks passed",
                          "file": None, "tag": None, "color": None, "workflows": []})
    edges = [{"from": "work:t0", "to": "work:t3", "kind": "handoff", "label": "x", "n": 1, "basis": "recorded"}, {"from": "work:t1", "to": "work:t5", "kind": "handoff", "label": "x", "n": 1, "basis": "declared"}]
    fn = lambda k, label, stt, detail, on=None, calls=3: {"id": f"fn:{k}", "kind": "function", "row": "functions", "label": label, "status": stt, "version": None, "tag": None,
                                                           "color": f"--{k}", "file": None, "workflows": [], "tab": {"identity": "charter", "adaptation": "workflows", "priority": "conflicts", "coordination": "timeline", "audit": "audit"}[k],
                                                           "wait_on": on, "detail": detail, "calls": calls, "seconds": 12.5, "tokens": {"input": 100, "output": 20}}
    fnodes = [fn("identity", "Identity", "running", "writing the charter"), fn("adaptation", "Adaptation", "waiting", "proposed workflows", "Priority"),
              fn("priority", "Priority", "waiting_on_you", "asked a question", "you", 1), fn("coordination", "Coordination", "fault", "W-3 is held"),
              fn("audit", "Audit", "idle", "last seen: @Identity all criteria pass; done-when items met", None, 0)]
    spk = lambda a, b, n: {"from": a, "to": b, "kind": "speaks_to", "label": None, "n": n, "basis": "recorded"}
    few = [spk("person:requester", "fn:identity", 2), spk("fn:identity", "person:requester", 5), spk("fn:identity", "fn:adaptation", 4), spk("fn:adaptation", "fn:priority", 7)]
    mp = {"schema": "map/1", "status": {"state": "Identity running", "state_status": "running", "checks": {"passed": 3, "run": 4}}, "nodes": nodes, "edges": edges, "legend": saved["legend"], "not_recorded": ["version of Charter"]}
    mp2 = {**mp, "nodes": [nodes[0], *fnodes, *nodes[1:]], "edges": edges + few}
    out2 = render_node(snap, docs={"docs": {}, "map": mp2})["map"]
    svg2 = out2[out2.index("<svg class=\"mapsvg\""):out2.index("</svg>", out2.index("<svg class=\"mapsvg\"")) + 6]
    ng = lambda s, i: next(x for x in s.split('<g class="mapnode')[1:] if f'data-id="{htmllib.escape(i)}"' in x)      # one dot: its tooltip, rings and label
    fbox = lambda k: ng(svg2, f"fn:{k}")
    PN = {"identity": "Client Lead (Identity)", "adaptation": "Strategist (Adaptation)", "priority": "Prioritiser (Priority)", "coordination": "Scheduler (Coordination)", "audit": "Quality Reviewer (Audit)"}
    rowl = [">Management layer</text>", ">Team layer and deliverables</text>"]
    sv2v = out2[out2.index('<svg class="mapsvg mapvert"'):out2.index("</svg>", out2.index('<svg class="mapsvg mapvert"')) + 6]      # the narrow drawing (under 900px)
    rows_ = lambda h: re.findall(r'class="rowlab"[^>]*>([^<]+)<', h)
    assert rows_(sv2v) == ["You", "Management layer", "Team layer and deliverables"] and rows_(svg2) == [], "the wide drawing has no bands: one row of chips above the trunk"
    chip_ = lambda i: tuple(float(v) for v in re.findall(r'<rect class="shape" x="([\d.-]+)" y="([\d.-]+)"', ng(svg2, i))[-1])      # the inner outline (a running chip has two)
    assert chip_("person:requester")[1] == chip_("fn:identity")[1] and chip_("person:requester")[0] < chip_("fn:identity")[0] < chip_("fn:adaptation")[0] < chip_("fn:audit")[0], "You first, then the five functions, in one row"
    assert 'class="youdot"' in ng(svg2, "person:requester") and re.search(r'class="sub chipn"[^>]*>2<', ng(svg2, "person:requester")) and re.search(r'class="sub chipn"[^>]*>3<', fbox("identity")), "each chip shows a count (messages from you; calls)"
    assert ">FUNCTIONS<" not in svg2 and ">WORK<" not in svg2 and svg2.count("data-maptab=") == 0, "links live in the panel and the cards"
    assert svg2.count('<g class="mapnode') == 1 + 5 + 5 + 6 and re.findall(r'data-fn="(fn:\w+)"', svg2) == ["fn:" + k for k in PN], "a dot for every node, the five functions in order"
    assert fbox("identity").count('<rect class="shape"') == 2 and ">Client Lead<" in fbox("identity") and 'aria-label="Client Lead (Identity): details"' in fbox("identity") and "writing the charter" in fbox("identity").split("</title>")[0] and ">3 calls<" not in fbox("identity")      # double ring = running now; the call count and the state sit in the tooltip
    assert 'stroke="var(--validator)" stroke-width="4"' in fbox("adaptation") and ">?</text>" in fbox("adaptation") and "waiting on Priority" in fbox("adaptation").split("</title>")[0] and ">Strategist<" in fbox("adaptation")
    assert 'stroke="var(--validator)" stroke-width="4"' in fbox("priority") and "waiting on you" in fbox("priority").split("</title>")[0] and "1 AI call," in fbox("priority").split("</title>")[0] and ">Prioritiser<" in fbox("priority")
    assert 'stroke="var(--bad)" stroke-width="4"' in fbox("coordination") and ">!</text>" in fbox("coordination") and "blocked" in fbox("coordination").split("</title>")[0] and ">Scheduler<" in fbox("coordination")
    assert 'stroke="var(--line)"' in fbox("audit") and ", idle, " in fbox("audit") and "last seen: " in fbox("audit") and "@Identity" not in fbox("audit") and not re.search(r"done.when", fbox("audit")) and ">Quality Reviewer<" in fbox("audit")
    assert "<title>Client Lead (Identity), running now, writing the charter, 3 AI calls, 12.5 s, 100 tokens in, 20 out</title>" in fbox("identity")      # the full text and numbers in a tooltip
    assert "2 from you · 5 to you" in sv2v and sv2v.count('stroke-dasharray="1 3"') >= 3, "the narrow drawing draws a few function lines"
    assert "Client Lead (Identity) speaks to Strategist (Adaptation): 4 messages" in sv2v and "Strategist (Adaptation) speaks to Prioritiser (Priority): 7 messages" in sv2v
    assert svg2.count('<path class="mapedge speak') == 1 and '<details class="mapdd"><summary>Messages between functions (speaks to)</summary>' in out2, "the chip row draws only You to Client Lead; messages between functions are in the table"
    many = [spk("fn:identity", "fn:adaptation", 4), spk("fn:adaptation", "fn:priority", 7), spk("fn:priority", "fn:coordination", 2), spk("fn:coordination", "fn:audit", 9), spk("fn:audit", "fn:identity", 1)]
    out3 = render_node(snap, docs={"docs": {}, "map": {**mp2, "edges": edges + many}})["map"]
    svg3 = out3[out3.index("<svg class=\"mapsvg\""):out3.index("</svg>", out3.index("<svg class=\"mapsvg\"")) + 6]
    assert '<details class="mapdd"><summary>Messages between functions (speaks to)</summary>' in out3 and "speaks to Strategist (Adaptation): 4" not in svg3 and svg3.count('stroke-dasharray="1 3"') == 0, "too many lines go in a table"
    assert out3.count("<tr><td>") >= 5 and "<td>Scheduler (Coordination)</td><td>Quality Reviewer (Audit)</td><td>9</td>" in out3
    assert not re.search(r'<details class="mapdd"[^>]* open', out3), "the table starts closed, a click opens it"
    cards2 = out2.split('<div class="mapcards">')[1].split('<div class="maplegend"')[0]      # small screens: the functions stacked as cards
    assert cards2.count('data-node="fn:') == 5 and cards2.count('class="mapcard st-') == 1 + 5 + 6 and "speaks to Strategist (Adaptation) · 4" in cards2 and "speaks to Prioritiser (Priority) · 7" in cards2
    assert "waiting on Priority" in cards2 and "blocked" in cards2 and "speaks to Client Lead (Identity) · 2" in cards2 and "Client Lead (Identity) speaks to you · 5" in cards2 and "<b>?</b>" in cards2
    assert 'data-maptab="charter"' in cards2 and not re.search(r"envelope|done.when", out2 + out3, re.I)
    # function-to-work lines: drawn under the boxes, faint, highlighted per function (hover, click), only on demand when there are many
    wk = lambda f, t, n, lab="calls": {"from": f, "to": t, "kind": "speaks_to", "label": lab, "n": n, "basis": "recorded"}
    few_w = [wk("fn:identity", "work:t0", None, "wrote the charter"), wk("fn:identity", "work:t1", 2), wk("fn:audit", "work:t1", 1, "audits"), wk("fn:coordination", "work:t3", 4)]
    o4 = render_node(snap, docs={"docs": {}, "map": {**mp2, "edges": edges + few + few_w}, "mappin": "fn:identity", "mapview": "line"})["map"]
    s4 = o4[o4.index("<svg class=\"mapsvg"):o4.index("</svg>", o4.index("<svg class=\"mapsvg")) + 6]
    assert re.findall(r'<path class="mapln[^"]*" data-from="([^"]+)" data-to="([^"]+)"', s4) == [("fn:identity", "work:t0"), ("fn:identity", "work:t1"), ("fn:audit", "work:t1"), ("fn:coordination", "work:t3")]
    assert s4.startswith('<svg class="mapsvg" data-pin="fn:identity"') and ">wrote it<" in s4 and re.findall(r'class="maplbl[^"]*" data-from="[^"]+" data-to="[^"]+" x="[^"]+" y="[^"]+">([^<]+)<', s4) == ["wrote it", "2", "1", "4"]
    for need in (".mapln{opacity:0}", ".maplbl{opacity:0", ".mapsvg .hot.mapln", ".mapsvg .mapnode.dim", ".mapsvg .hot.maplbl"):      # function-to-deliverable lines are hidden until a dot is lit; lit when hot; the rest fade when another dot is lit
        assert need in s4, need
    s4v = o4[o4.index('<svg class="mapsvg mapvert"'):]
    assert "hot" in ng(s4v, "fn:adaptation").split(">")[0] and "dim" in ng(s4, "fn:adaptation").split(">")[0], "a function line is a neighbour where it is drawn (narrow); the wide chip row keeps those in the table"
    assert "hot" in ng(s4, "fn:identity").split(">")[0] and "hot" in ng(s4, "person:requester").split(">")[0] and "hot" in ng(s4, "work:t0").split(">")[0] and "hot" in ng(s4, "work:t1").split(">")[0], "the pinned dot and its direct neighbours are lit"
    assert "dim" in ng(s4, "fn:audit").split(">")[0] and "dim" in ng(s4, "work:t3").split(">")[0] and "dim" in ng(s4, "fn:priority").split(">")[0], "the rest fade"



    assert len(re.findall(r'<g class="mapnode[^"]*"[^>]*aria-pressed="true"', s4)) == 1 and 'data-id="fn:identity" data-fn="fn:identity" tabindex="0" role="button" aria-pressed="true"' in s4 and 'data-id="fn:audit" data-fn="fn:audit" tabindex="0" role="button" aria-pressed="false"' in s4
    assert 'class="mapln hot" data-from="fn:identity"' in s4 and 'class="mapln dim" data-from="fn:coordination"' in s4 and 'class="mapln dim" data-from="fn:audit"' in s4, "lines that touch the pinned dot are lit, the others fade"
    assert s4.index('class="mapln') < s4.index('<g class="mapnode') and s4.rindex('<path class="map') < s4.index('<g class="mapnode'), "lines are drawn under the dots"
    ends = [re.search(r'd="M[\d.]+ ([\d.]+) C[^"]* ([\d.]+)"', x).groups() for x in re.findall(r'<path class="mapln[^"]*"[^>]*>', s4)]
    assert len({a for a, _ in ends}) == 1, "function-to-deliverable lines all start under the function labels"
    tops_ = {}      # every node's centre and top in the wide drawing
    for g_ in s4.split('<g class="mapnode')[1:]:
        id_ = htmllib.unescape(re.search(r'data-id="([^"]+)"', g_).group(1))
        m1 = re.search(r'points="([\d.]+),([\d.]+) ([\d.]+),([\d.]+)', g_) or None
        m2 = re.search(r'class="teamhit" cx="([\d.]+)" cy="([\d.]+)" r="([\d.]+)"', g_)
        if m1: tops_[id_] = (float(m1.group(1)), float(m1.group(2)))
        elif m2: tops_[id_] = (float(m2.group(1)), float(m2.group(2)) - float(m2.group(3)))
    for lp in re.findall(r'<path class="mapln[^"]*" data-from="[^"]+" data-to="([^"]+)" d="([^"]+)"', s4):      # a chip's line ends just above the team that made the deliverable, or the deliverable
        ex_, ey_ = (float(v) for v in re.findall(r"-?[\d.]+", lp[1])[-2:])
        assert any(abs(ex_ - x) < 0.6 and -0.6 <= y - ey_ <= 8.6 for x, y in tops_.values()), lp      # (a running node has an outer outline 4 wider)
    for tip in ("Client Lead (Identity) wrote the Box 0 with a very long name", "Client Lead (Identity) made 2 AI calls on Box 1 with a very long name", "Quality Reviewer (Audit) audited Box 1 with a very long name 1 time", "Scheduler (Coordination) made 4 AI calls on Box 3 with a very long name"):
        assert f"<title>{tip}</title>" in s4, tip
    assert "many" not in s4.split(">")[0], "no 'many' mode any more: the lines are always hover-only"
    lots = [wk(f"fn:{f}", f"work:t{i}", i + 1) for i in range(5) for f in ("identity", "adaptation", "priority")]      # 15 lines
    s5 = (lambda o: o[o.index("<svg class=\"mapsvg"):o.index("</svg>", o.index("<svg class=\"mapsvg")) + 6])(render_node(snap, docs={"docs": {}, "map": {**mp2, "edges": edges + lots}})["map"])
    assert s5.startswith('<svg class="mapsvg"') and s5.count('<path class="mapln') == 15, "however many lines there are: drawn, but shown only on hover or click"
    c4 = o4.split('<div class="mapcards">')[1].split('<div class="maplegend"')[0]      # small screens: no lines, plain sentences
    assert "mapln" not in c4 and "made with: Client Lead (Identity), " not in c4 and "made with: Client Lead (Identity) 2, Quality Reviewer (Audit) 1" in c4 and "which function worked on this: not recorded" in c4
    assert "worked on: Box 0 with a very long name, Box 1 with a very long name 2" in c4 and "worked on: Box 1 with a very long name 1" in c4 and "worked on: Box 3 with a very long name 4" in c4
    assert not re.search(r"envelope|done.when", o4 + s5, re.I)
    out = render_node(snap, docs={"docs": {}, "map": mp})["map"]
    svg = out[out.index("<svg class=\"mapsvg\""):out.index("</svg>", out.index("<svg class=\"mapsvg\"")) + 6]
    assert out.count("data-mapopen=") == 6 and "data-mapopen" not in svg and 'class="mapnode' in svg      # the Open links are in the cards and in the click panel, not on the canvas
    pn0 = render_node(snap, docs={"docs": {}, "map": mp, "mappin": "work:t0", "mapview": "line"})["map"]
    assert pn0.count("data-mapopen=") == 7 and 'class="mappanel" data-panel="work:t0"' in pn0 and 'data-mapopen="f0.md">Open Box 0 with a very long name</a>' in pn0 and "data-mapclose" in pn0 and "v1" in pn0.split('class="mappanel"')[1], "the panel has the link and the details"
    assert "Identity running" in out and "3 of 4 checks passed" in out and 'class="mappill running"' in out and 'class="mappill checks"' in out
    nc = render_node(snap, docs={"docs": {}, "map": {**mp, "status": {**mp["status"], "state": "Crashed", "state_status": "fault", "checks": {"passed": 0, "run": 0}}}})["map"]
    assert ">No checks yet<" in nc and 'class="mappill nochecks"' in nc and 'class="mappill checks"' not in nc and "0 of 0" not in nc, "a run with no checks is not shown as a green chip"
    box = lambda i: ng(svg, f"work:t{i}")
    assert 'stroke-dasharray="6 4"' in box(2) and "stroke-dasharray" not in box(0)                       # dashed = left open
    assert box(1).count("<polygon") == 2 and box(0).count("<polygon") == 1 and "<circle class=\"shape\"" not in box(0)      # deliverables are diamonds; a double outline = running now
    assert 'stroke="var(--bad)" stroke-width="4"' in box(3) and ">!</text>" in box(3)                       # thick red edge plus "!"
    assert 'fill="none" stroke="var(--line)"' in box(4) and ">!</text>" not in box(4)                      # hollow = not made
    assert 'stroke="var(--validator)" stroke-width="4"' in box(5) and ">?</text>" in box(5) and "version not recorded" in box(5)
    assert "color-mix(in srgb, var(--ok) 14%, var(--panel))" in box(0) and ", version 1, filed and checked," in box(0) and re.search(r'class="sub"[^>]*>v1(<| · )', box(0)) and ">filed<" not in box(0)      # green = filed; the short fact line under it
    assert not re.search(r"var\(--(prd|hld|lld|site|charter|other|identity|adaptation|priority|coordination|audit|user)\)", svg), "colours follow the status only, never the kind of document or function"
    assert "left open" in box(2).split("</title>")[0] and "running now" in box(1).split("</title>")[0] and "fault" in box(3).split("</title>")[0] and "not made" in box(4).split("</title>")[0] and "waiting on you" in box(5).split("</title>")[0]
    assert "<title>Box 3 with a very long name, version 4, fault, Every checklist item passes inside its limits</title>" in box(3)      # full text in a tooltip
    assert not re.search(r"envelope|done.when", out, re.I), "plain words"
    assert "<circle" in svg and svg.count("<path") >= 2 and 'stroke-dasharray="2 3"' in svg and 'marker-end="url(#maparrow)"' in svg
    assert 'role="img"' in svg and 'viewBox="0 0' in svg and 'width="' not in svg.split(">")[0], "scales to the screen"
    # small screens: the same boxes as cards, the work as a top-to-bottom list with "hands off to" lines
    cards = out.split('<div class="mapcards">')[1].split('<div class="maplegend"')[0]
    assert cards.count('class="mapcard st-') == 7 and cards.count('class="mapstep"') == 5 and "given to Box 3 with a very long name (which step used it: not recorded)" in cards and "hands off" not in cards
    assert 'class="mapcard st-fault"' in cards and "<b>!</b>" in cards and "<b>?</b>" in cards and "left open" in cards and "running now" in cards
    ui = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "j2_ui.html"), encoding="utf-8").read()
    assert "@media (max-width:600px){.mapsvgwrap{display:none}.mapcards{display:block}}" in ui and ".mapcards{display:none}" in ui
    assert ".mapsvg{width:100%;height:auto" in ui and "st-running{border-style:double" in ui and "st-left_open{border-style:dashed" in ui
    assert ui.index('data-tab="map"') < ui.index('id="railsub"') < ui.index('data-tab="docs"') < ui.index('data-tab="artifacts"') < ui.index('data-tab="style"'), "Map is the first rail entry; every other tab sits in the Internals folder"
    empty = render_node(snap, docs={"docs": {}, "map": {"nodes": [], "edges": []}})["map"]
    assert "Nothing to draw yet." in empty
    real = render_node(snap, docs={"docs": {}, "map": saved})["map"]
    assert real.count('<g class="mapnode') == 2 * len(saved["nodes"]) and real.split('<svg class="mapsvg mapvert"')[0].count('<g class="mapnode') == len(saved["nodes"]) and "Product requirements" in real and "v6" in real and f"Not recorded ({len(saved['not_recorded'])})" in real
    assert not re.search(r"envelope|done.when", real, re.I)
    assert real.count('data-maptab=') >= 5 and ">Management layer</text>" in real and '<details class="mapdd"><summary>Messages between functions (speaks to)</summary>' in real      # six pairs: a table, not crossing lines
    assert '<details class="mapdd"><summary>Who made each step (calls by function)</summary>' in real and "<td>Web Developer (building Website)</td><td>not recorded</td>" in real
    assert not re.search(r'<details class="mapdd"[^>]* open', real), "both tables start closed on the real map too"
    assert "<td>Product Analyst (drafting Product requirements)</td><td>%d</td>" % by_fn("identity", docs_rec["prd"]["calls"], calls) in real
    rc = real.split('<div class="mapcards">')[1].split('<div class="maplegend"')[0]
    lab = {k: by[f"work:{k}"]["label"] for k in ("prd", "hld", "lld", "website")}
    mw = lambda k: ", ".join(f"{PN[f]} {n}" for f, t, n, lb, _ in lines if t == k)
    assert f"made with: {mw('prd')}" in rc and f"made with: {mw('lld')}" in rc and "made with: Quality Reviewer (Audit) 2" in rc and "made with: Client Lead (Identity)<" in rc.replace("</div>", "<")
    assert rc.count("which other function worked on this: not recorded") == 1 and "worked on: Charter, " + lab["prd"] + " " + str(by_fn("identity", docs_rec["prd"]["calls"], calls)) in rc
    assert "worked on: " + lab["prd"] + " 1, " + lab["hld"] + " 1, " + lab["lld"] + " 1, " + lab["website"] + " 2" in rc      # Audit
    assert real.split("<svg")[1].split(">")[0].startswith(' class="mapsvg"') and real.count('<path class="mapln"') == 2 * len([x for x in lines]), "17 lines in each drawing, shown only on hover or click"
    assert "which function made this step: not recorded" in rc and "made by Client Lead (Identity) 2, Strategist (Adaptation) 1, Prioritiser (Priority) 1, Scheduler (Coordination) 6, Quality Reviewer (Audit) 1" in rc and rc.count('data-node="fn:') == 5
    # a browser: no sideways scroll at 500px, cards shown, SVG hidden; the tab sits after Artifacts and the pane reads plain
    chrome = j2.j2_check.find_chrome()
    if chrome and BROWSER:
        import subprocess
        real_line = render_node(snap, docs={"docs": {}, "map": saved, "mapview": "line"})["map"]      # the metro drawing (the Line view)
        css = re.search(r"<style>(.*?)</style>", ui, re.S).group(1)
        page = os.path.join(tempfile.mkdtemp(prefix="j2-map-"), "map.html")
        open(page, "w", encoding="utf-8").write(f"<html><head><meta name=viewport content='width=device-width'><style>{css}</style></head><body>{real_line}"
            "<script>document.body.setAttribute('data-m', [document.documentElement.scrollWidth, innerWidth, getComputedStyle(document.querySelector('.mapsvgwrap')).display, getComputedStyle(document.querySelector('.mapcards')).display, [...document.querySelectorAll('.mapln')].every(x => x.getBoundingClientRect().width === 0)].join('|'))</script></body></html>")
        for width, want in ((500, "none|block|true"), (1100, "block|none|false")):
            dom = chrome_dom([chrome, "--headless=new", "--disable-gpu", f"--window-size={width},800", "--virtual-time-budget=2000", "--dump-dom", "file://" + page],
                                 capture_output=True, text=True, timeout=120).stdout
            sw, iw, a, b, c = re.search(r'data-m="([^"]+)"', dom).group(1).split("|")
            assert int(sw) <= int(iw) and f"{a}|{b}|{c}" == want, (width, sw, iw, a, b, c)      # at 500px the function-to-work lines take no space
    if chrome and BROWSER:      # the real page: hover, click and click again on function boxes highlight and release their lines
        import subprocess
        tmp = tempfile.mkdtemp(prefix="j2-maphl-")
        probe = ("<script>setTimeout(() => { const q = s => [...document.querySelectorAll(s)], op = s => q(s).map(x => getComputedStyle(x).opacity), svg = () => document.querySelector('.mapsvg'), r = {},"
                 "fire = (el, t) => el.dispatchEvent(new MouseEvent(t, {bubbles: true})), fn = id => document.querySelector('.mapfn[data-fn=\"fn:' + id + '\"]');"
                 "r.many = svg().classList.contains('many'); r.idle = op('.mapln'); r.idleLbl = op('.maplbl');"
                 "fire(fn('coordination'), 'mouseover'); r.hoverOwn = op('.mapln[data-from=\"fn:coordination\"]'); r.hoverOther = op('.mapln[data-from=\"fn:identity\"]'); r.lblOwn = op('.maplbl[data-from=\"fn:coordination\"]'); r.lblOther = op('.maplbl[data-from=\"fn:audit\"]');"
                 "fire(fn('coordination'), 'mouseout'); r.after = op('.mapln[data-from=\"fn:coordination\"]');"
                 "fire(fn('coordination'), 'click'); r.pin = svg().dataset.pin; r.pinOwn = op('.mapln[data-from=\"fn:coordination\"]'); r.pressed = fn('coordination').getAttribute('aria-pressed'); r.pinOther = op('.mapln[data-from=\"fn:identity\"]');"
                 "fire(fn('audit'), 'click'); r.repin = svg().dataset.pin; fire(fn('audit'), 'click'); r.unpin = svg().dataset.pin;"
                 "showTab('map'); fire(fn('audit'), 'click'); r.panel = !!document.querySelector('.mappanel[data-panel=\"fn:audit\"]'); document.dispatchEvent(new KeyboardEvent('keydown', {key: 'Escape', bubbles: true}));"
                 "r.esc = svg().dataset.pin + '|' + !!document.querySelector('.mappanel');"
                 "fn('audit').dispatchEvent(new FocusEvent('focusin', {bubbles: true})); r.focusOwn = op('.mapln[data-from=\"fn:audit\"]'); r.focusOther = op('.mapln[data-from=\"fn:identity\"]');"
                 "document.body.setAttribute('data-r', JSON.stringify(r)); }, 4500)</script>")
        open(os.path.join(tmp, "j2_ui.html"), "w", encoding="utf-8").write(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "j2_ui.html"), encoding="utf-8").read().replace("</body>", probe + "</body>"))
        keep_here, j2_server.HERE = j2_server.HERE, tmp
        keep, j2_server.RUN = j2_server.RUN, run
        srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            dom = chrome_dom([chrome, "--headless=new", "--disable-gpu", "--virtual-time-budget=12000", "--dump-dom",
                                  f"http://127.0.0.1:{srv.server_address[1]}/?view=detail&run={name}"], capture_output=True, text=True, timeout=180).stdout
        finally:
            j2_server.HERE, j2_server.RUN = keep_here, keep
            srv.shutdown()
        got = json.loads(htmllib.unescape(re.search(r'data-r="([^"]*)"', dom).group(1)))
        assert got["many"] is False and set(got["idle"]) == {"0"} and set(got["idleLbl"]) == {"0"}, got      # 17 lines, and always: nothing drawn until a function is hovered or clicked (no "many" mode any more)
        assert got["hoverOwn"] and set(got["hoverOwn"]) == {"1"} and set(got["hoverOther"]) == {"0"} and set(got["lblOwn"]) == {"1"} and set(got["lblOther"]) == {"0"}, got
        assert set(got["after"]) == {"0"} and got["pin"] == "fn:coordination" and set(got["pinOwn"]) == {"1"} and got["pressed"] == "true" and set(got["pinOther"]) == {"0"}, got
        assert got["repin"] == "fn:audit" and got["unpin"] == "", got      # another click moves the highlight; clicking it again releases it
        assert got["panel"] is True and got["esc"] == "|false", got      # a click opens the panel; Esc closes it and releases the dot
        assert got["focusOwn"] and set(got["focusOwn"]) == {"1"} and set(got["focusOther"]) == {"0"}, got      # keyboard focus lights a dot like hover
    if chrome and BROWSER:      # the real page, a saved run: the Map tab is after Artifacts and its pane reads in plain words
        import subprocess
        keep, j2_server.RUN = j2_server.RUN, run
        srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            dom = chrome_dom([chrome, "--headless=new", "--disable-gpu", "--virtual-time-budget=8000", "--dump-dom",
                                  f"http://127.0.0.1:{srv.server_address[1]}/?view=detail&run={name}"], capture_output=True, text=True, timeout=120).stdout
        finally:
            j2_server.RUN = keep
            srv.shutdown()
        pane = dom.split('id="pane-map"')[1].split('id="pane-style"')[0]
        text = htmllib.unescape(re.sub(r"<[^>]+>", " ", pane))
        assert "Product requirements" in text and re.search(r"\d+ of \d+ checks passed", text) and "Done" in text, text[:300]
        assert 'data-mapopen="prd.md"' in pane and 'data-maptab="charter"' in pane and "Management layer" in pane and "Product Analyst" in pane and "Web Developer" in pane and "Specialists" not in pane and not re.search(r"envelope|done.when", text, re.I), "plain words in the Map tab"
        boxes = re.findall(r'<g class="mapnode mapfn.*?(?=<g class="mapnode)', pane, re.S)[:5]      # the five function dots in the page (the legend may say "running now")
        assert len(boxes) == 5 and all("last seen: " in b and "running now" not in b and "waiting on" not in b for b in boxes), "a saved run never says running"
        assert dom.index('data-tab="map"') < dom.index('id="railsub"') < dom.index('data-tab="artifacts"')
    # "given to": each hand-off line ends on the step that was given the document, never on a later document box
    rg = render_node(snap, docs={"docs": {}, "map": saved})
    gv, byid = rg["given"], {n["id"]: n for n in saved["nodes"]}
    assert gv and all(g["known"] and byid[g["to"]]["kind"] == "engine" and g["to"].endswith(">" + g["box"].split(":", 1)[1]) for g in gv), gv
    wide = lambda h: h.split('<svg class="mapsvg mapvert"')[0]      # the wide drawing (the narrow one draws the same lines)
    drawn = re.findall(r'<path class="mapedge given[^"]*" data-from="([^"]+)" data-to="([^"]+)"', wide(rg["map"]))
    assert drawn and all(b.startswith("step:") for _, b in drawn) and len(drawn) == len([g for g in gv if not g["chain"]]), drawn
    assert "work:" not in " ".join(b for _, b in drawn) and "which step used" not in rg["map"], "every consumer is known on the real map"
    cut = {**saved, "edges": [e for e in saved["edges"] if not (e["kind"] == "step" and e["to"] == "work:lld")]}      # no step recorded for the LLD: the line stays on the box and the gap is listed
    oc = render_node(snap, docs={"docs": {}, "map": cut})
    gc = [g for g in oc["given"] if g["box"] == "work:lld"]
    assert gc and all(not g["known"] and g["to"] == "work:lld" for g in gc), gc
    dc = re.findall(r'<path class="mapedge given[^"]*" data-from="([^"]+)" data-to="([^"]+)"', wide(oc["map"]))
    assert ("work:prd", "work:lld") in dc and all(b.startswith("step:") for a, b in dc if b != "work:lld")
    assert f"Not recorded ({len(saved['not_recorded']) + len(gc)})" in oc["map"] and "which step used Product requirements given to Low-level design: not recorded" in oc["map"]
    assert "given to Low-level design (which step used it: not recorded)" in oc["map"].split('class="mapcards"')[1], "the cards say the same"
    # every function dot reads "Persona (Function)"
    assert re.findall(r'data-fn="fn:\w+"[^>]*aria-label="([^"]+)"', wide(rg["map"])) == [PN[k] + ": details" for k in PN], "Persona (Function) on every function dot"
    assert ">Identity<" not in rg["map"].split("<svg")[1].split("</svg>")[0] and ">(Identity)<" in rg["map"], "the bare function name is only in brackets"
    # a step is Drafting, Building or Sending by what it makes, wherever it sits in the row
    def tiny(order):
        kinds = {"d1": ("PRD", "Write"), "d2": ("HLD", "Write"), "d3": ("Site", "Build"), "d4": ("Send", "Send"), "d5": ("LLD", "Make")}
        ns, es, prev = [], [], None
        for i in order:
            tag, verb = kinds[i]
            ns.append({"id": f"work:{i}", "kind": "work", "row": "work", "label": f"Deliverable {i}", "tag": tag, "color": "--other", "status": "filed", "version": 1, "detail": "", "file": None, "workflows": []})
            if prev:
                sid = f"step:{prev}>{i}"
                ns.insert(len(ns) - 1, {"id": sid, "kind": "engine", "row": "work", "label": kinds[i][1], "status": "filed", "version": None, "detail": "2 workflows", "file": None, "workflows": ["W-1", "W-2"], "tag": None, "color": None})
                es += [{"from": f"work:{prev}", "to": sid, "kind": "step", "label": None, "n": 2, "basis": "recorded"}, {"from": sid, "to": f"work:{i}", "kind": "step", "label": None, "n": 2, "basis": "recorded"}]
            prev = i
        return {"schema": "map/1", "status": {"state": "Done", "state_status": "idle", "checks": {"passed": 0, "run": 0}}, "nodes": ns, "edges": es, "legend": saved["legend"], "not_recorded": []}
    for order in (["d1", "d2", "d3", "d4", "d5"], ["d3", "d5", "d1", "d4", "d2"]):
        mo = tiny(order)["nodes"]
        out_t = render_node(snap, docs={"docs": {}, "map": tiny(order)})["map"]
        svg_t = out_t[out_t.index("<svg class=\"mapsvg"):out_t.index("</svg>")]
        for n in mo:
            if n["kind"] != "engine":
                continue
            tgt = n["id"].split(">")[1]
            who_, verb = {"d1": ("Product Analyst", "drafting"), "d2": ("Solution Architect", "drafting"), "d3": ("Web Developer", "building"), "d4": ("Messenger", "sending"), "d5": ("Technical Designer", "drafting")}[tgt]
            words_ = " ".join(re.sub(r"<[^>]+>", " ", ng(svg_t, n["id"]).split("</title>")[1]).split())
            assert words_ == f"{who_} · 2 workflows" and f"{who_}, {verb} Deliverable {tgt} (2 workflows)" in ng(svg_t, n["id"]), (order, n["id"], words_)
            svt_v = out_t[out_t.index('<svg class="mapsvg mapvert"'):]
            assert f">{who_}<" in ng(svt_v, n["id"]) and f">{verb} · 2 workflows<" in ng(svt_v, n["id"]), "the narrow drawing keeps one line each"
            assert f"{who_}, {verb} Deliverable {tgt} (2 workflows)" in out_t.split('class="mapcards"')[1], "the cards use the same words"
        assert ">Write<" not in svg_t and ">Build<" not in svg_t and ">Send<" not in svg_t and ">Drafting<" not in svg_t
    # ---- Map and Chat are the whole app: the rail, the Needs you strip, and the Map as the navigation
    rail = ui[ui.index('<nav class="rail tabs"'):ui.index("</nav>", ui.index('<nav class="rail tabs"'))]
    top, inside = rail.split('id="railsub"')
    old_tabs = ["preview", "goal", "workflows", "conflicts", "timeline", "audit", "charter", "docs", "artifacts", "board", "style", "lessons", "playbooks", "report", "rundetails", "runsettings"]
    assert re.findall(r'data-tab="(\w+)"', top) == ["map"] and 'id="gochat"' in top and 'id="internals"' in top and re.findall(r"<span>([^<]+)</span>", top) == ["Map", "Chat", "Internals"], "the rail has Map, Chat and one Internals folder"
    assert re.findall(r'data-tab="(\w+)"', inside) == old_tabs, "every other tab is under Internals, in the old order"
    assert top.count("<svg") == 3 and "<svg" not in inside and ' hidden' in rail.split('id="railsub"')[1].split(">")[0], "one small icon each, none in the folder, folder starts collapsed"
    dvh = ui[ui.index('<div id="dview">'):ui.index('<dialog id="detail">')]
    panes = re.findall(r'<div class="pane[^"]*" id="pane-(\w+)"( hidden)?', dvh)
    assert [n for n, h in panes if not h] == ["map"] and {n for n, _ in panes} == set(old_tabs) | {"map"}, "the home shows only the Map pane; every old pane still exists"
    assert dvh.index('id="needs"') < dvh.index('<section class="work">') < dvh.index('<section class="chat">'), "Needs you, then the Map, then the Chat"
    assert "@media (max-width:900px){#dview .dcol{overflow:auto}#dview main{grid-template-columns:minmax(0,1fr)" in ui and "@media (max-width:600px){.mapsvgwrap{display:none}.mapcards{display:block}}" in ui, "stacked below 900px; map cards at 600px"
    # Needs you: every open question, a plain line each, a button into the Chat; nothing waiting says so
    n0 = render_node(snap, docs={"docs": {}, "map": saved})
    assert n0["needcount"] == 0 and "Nothing needs you right now." in n0["needs"] and 'id="needcount">0<' in n0["needs"] and "data-needmsg" not in n0["needs"]
    qs = [{"question": "Add 25 more credits?", "options": ["Yes", "No"]}, {"question": "Leave out the lookbook?", "options": []}]
    ask = {"id": 9001, "who": "identity", "text": "@Requester Add credits? Leave out the lookbook?", "to_user": True, "pause": "limit", "next": "user"}
    n1 = render_node({**snap, "paused": {"kind": "limit", "label": "credits", "msg": 9001, "questions": qs}, "chat": snap["chat"] + [ask]}, docs={"docs": {}, "map": saved})
    assert n1["needcount"] == 2 and 'id="needcount">2<' in n1["needs"] and n1["needs"].count('data-needmsg="9001"') == 2 and ">Add 25 more credits?<" in n1["needs"] and ">Leave out the lookbook?<" in n1["needs"] and "Nothing needs you" not in n1["needs"]
    n2 = render_node({**snap, "paused": {"kind": "info", "label": "one fact", "msg": 9001, "questions": []}, "chat": snap["chat"] + [ask]}, docs={"docs": {}, "map": saved})
    assert n2["needcount"] == 1 and ">one fact<" in n2["needs"] and 'data-needmsg="9001"' in n2["needs"]
    n3 = render_node({**snap, "paused": None, "chat": snap["chat"] + [ask]}, docs={"docs": {}, "map": saved})      # a saved run that ended on a question: the same item a "waiting on you" node comes from
    assert n3["needcount"] == 1 and "Add credits? Leave out the lookbook?" in n3["needs"] and "@Requester" not in n3["needs"] and 'data-needmsg="9001"' in n3["needs"]
    assert not re.search(r"envelope|done.when", n1["needs"] + n3["needs"], re.I)
    # the Open button of every kind of node says where it goes
    want = {"person:requester": "Show what needs you", "fn:identity": "Open Chat", "fn:adaptation": "Show messages and calls", "fn:audit": "Open Audit findings", "step:prd>hld": "Open these workflows",
            "work:prd": "Open the document", "work:website": "Open the site preview", "work:charter": "Open the Charter"}
    for i, lab in want.items():
        o = render_node(snap, docs={"docs": {}, "map": saved, "mappin": i, "mapview": "line"})["map"]
        assert f'data-mapgo="{htmllib.escape(i)}">{lab}</button>' in o and o.count("data-mapgo=") == 1, (i, lab)
    for i in ("fn:priority", "person:requester"):      # a node that waits on you goes to its question in the Chat
        o = render_node(snap, docs={"docs": {}, "map": mp2, "mappin": i, "mapview": "line"})["map"]
        assert f'data-mapgo="{i}">Answer in Chat</button>' in o, i
    # a step opens the Workflows view filtered to its workflows; a function opens the tables under the Map filtered to it
    stp = by["step:prd>hld"]
    wf_all = [w["id"] for w in snap["workflows"]]
    fw = render_node(snap, docs={"docs": {}, "map": saved, "wfstep": stp["workflows"]})["workflows"]
    assert "data-stepfilter" in fw and "data-wfclear" in fw and all(f'data-key="wf-{w}"' in fw for w in stp["workflows"]) and not any(f'data-key="wf-{w}"' in fw for w in wf_all if w not in stp["workflows"]), "only that step's workflows"
    fa = render_node(snap, docs={"docs": {}, "map": saved})["workflows"]
    assert "data-stepfilter" not in fa and all(f'data-key="wf-{w}"' in fa for w in wf_all), "no filter: every workflow"
    fo = render_node(snap, docs={"docs": {}, "map": saved, "mapfilt": "fn:adaptation"})["map"]
    blk = fo.split("<summary>Messages between functions (speaks to)</summary>")[1].split("</details>")[0]
    rws = re.findall(r"<tr><td>([^<]*)</td><td>([^<]*)</td>", blk)
    assert "Showing messages and calls for Strategist (Adaptation)." in fo and "data-mapfiltclear" in fo and fo.count('<details class="mapdd" open>') == 2 and rws and all("Strategist (Adaptation)" in (a, b) for a, b in rws), rws
    assert "<th>Strategist (Adaptation)</th>" in fo and "<th>Prioritiser (Priority)</th>" not in fo, "the calls table keeps only that function"
    assert "data-mapfiltclear" not in real and '<details class="mapdd" open>' not in real, "unfiltered: both tables closed"
    # ---- the home is simpler: the header is only the department name; the controls are in Run settings; one Map heading
    assert '<header class="dtop">' not in ui and ui.count("J2 Department") == 1 and "<title>J2 Department</title>" in ui, "no department heading on the page; the title stays"
    assert ui.split('<div id="dview">')[1].split('<div class="dcol">')[1].lstrip().startswith('<section class="needs"'), "the page starts with the Needs you strip"
    rs = ui[ui.index('id="pane-runsettings"'):ui.index("</section>", ui.index('id="pane-runsettings"'))]
    assert all(f'id="{i}"' in rs for i in ("outsel", "runsel", "reset", "to-simple")) and "Back to simple view" in rs and 'href="?view=simple"' in rs, "the old controls live in Run settings"
    assert all(f'id="{i}"' not in ui[:ui.index('id="pane-runsettings"')].split('<div id="dview">')[1] for i in ("outsel", "runsel", "reset", "to-simple")), "and nowhere else in the Detailed view"
    assert real.count("<h3>Map</h3>") == 1 and real.count(">Map<") == 1, "one Map heading, next to the status chips"
    assert 'class="maphead"' in real and real.index("<h3>Map</h3>") < real.index('class="mappill'), "the heading comes first, then the chips"
    # the outcome is one plain line in the Needs you strip, with a link to the Report; the banner text is in the Report
    no = render_node(snap, docs={"docs": {}, "map": saved})["needs"]
    assert "data-outcome" in no and 'data-gotab="report">See report</button>' in no and re.search(r"Finished\. ", no) and "Viewing saved run" not in no, no
    ui_i = ui.index('id="pane-report"'); assert 'id="banner"' in ui[ui_i:ui_i + 200], "the full banner text lives in the Report tab"
    assert render_node({**snap, "outcome": "Done with open items: shipped as is."}, docs={"docs": {}, "map": saved})["needs"].count("Finished. Some items were left open.") == 1
    assert "data-outcome" not in render_node({**snap, "outcome": ""}, docs={"docs": {}, "map": saved})["needs"], "a run with no outcome yet has no outcome line"
    # no two lines cross, on three runs (computed from the drawn SVG), and every team node has a persona from the list
    def svg_edges(h):
        out = []
        for m_ in re.finditer(r'<path class="mapedge[^"]*" data-from="([^"]+)" data-to="([^"]+)" d="([^"]+)"', h):
            pts, cx, cy = [], 0.0, 0.0
            for k_, a_ in re.findall(r"([MLC])([^MLC]*)", m_.group(3)):
                v = [float(x) for x in a_.split()]
                if k_ != "C":
                    cx, cy = v[0], v[1]; pts.append((cx, cy)); continue
                pts += [((1 - t) ** 3 * cx + 3 * (1 - t) ** 2 * t * v[0] + 3 * (1 - t) * t * t * v[2] + t ** 3 * v[4], (1 - t) ** 3 * cy + 3 * (1 - t) ** 2 * t * v[1] + 3 * (1 - t) * t * t * v[3] + t ** 3 * v[5]) for t in (i / 60 for i in range(1, 61))]
                cx, cy = v[4], v[5]
            out.append((htmllib.unescape(m_.group(1)), htmllib.unescape(m_.group(2)), pts))
        return out
    def crossings(h):
        es, n = svg_edges(h), []
        o = lambda a, b, c: (b[0] - a[0]) * (c[1] - a[1]) - (b[1] - a[1]) * (c[0] - a[0])
        for i in range(len(es)):
            for j in range(i + 1, len(es)):
                P, Q = es[i][2], es[j][2]
                if any(o(P[a], P[a + 1], Q[b]) * o(P[a], P[a + 1], Q[b + 1]) < -1e-6 and o(Q[b], Q[b + 1], P[a]) * o(Q[b], Q[b + 1], P[a + 1]) < -1e-6 for a in range(len(P) - 1) for b in range(len(Q) - 1)):
                    n.append((es[i][:2], es[j][:2]))
        return n
    TEAM = {"Product Analyst", "Solution Architect", "Technical Designer", "Specialist", "Web Developer", "Messenger"}
    repo_runs, real_names = os.path.join(os.path.dirname(os.path.abspath(__file__)), "runs"), []
    for n_ in ("j2-20261016-160511", "j2-20261022-123744"):      # the two saved runs, copied (read only) into the test's runs folder so the page and the server can open them
        if os.path.isdir(os.path.join(repo_runs, n_)) and not os.path.exists(os.path.join(j2.RUNS_DIR, n_)):
            shutil.copytree(os.path.join(repo_runs, n_), os.path.join(j2.RUNS_DIR, n_))
            real_names.append(n_)
    keep_mail = os.environ.get("J2_TEST_EMAIL"); os.environ["J2_TEST_EMAIL"] = "me@example.invalid"
    try:
        send_run = docs_run(DESIGN_BRIEF + " When the documents are finished, send a test email to me with a short summary of the PRD.", 900)      # a mock run with a send step
    finally:
        os.environ.pop("J2_TEST_EMAIL", None)
        if keep_mail is not None:
            os.environ["J2_TEST_EMAIL"] = keep_mail
    send_name = os.path.basename(send_run.dir)
    cross_runs = [("mock run", run.dir), ("mock run with a send", send_run.dir)] + [(n_, os.path.join(j2.RUNS_DIR, n_)) for n_ in real_names]
    assert len(cross_runs) == 4, "both saved runs are in runs/"
    sm_ = M.build(send_run.dir)
    wo_ = [n_["id"] for n_ in sm_["nodes"] if n_["row"] == "work"]
    assert wo_.index("work:lld") < wo_.index("work:website") and any(n_.get("tag") == "Send" and n_["lane"] == "branch" for n_ in sm_["nodes"]), "the send sits off the trunk, the site follows the LLD"
    # the outcome line names what is left open; the chip says when the latest Audit wants a revision
    ns_ = render_node(j2.saved_snapshot(send_name, 0), docs={"docs": {}, "map": sm_})["needs"]
    assert "Finished. 1 item left open: test email (dry run)." in ns_, ns_
    ss_ = j2.saved_snapshot(send_name, 0)
    chip_ok = render_node({**ss_, "workable": []}, docs={"docs": {}, "map": sm_})["map"]
    chip_bad = render_node({**ss_, "workable": ["necessary condition N8 not met: a nav reaches every section"]}, docs={"docs": {}, "map": sm_})["map"]
    assert 'class="mappill checks"' in chip_ok and "needs revision" not in chip_ok, "every check passed and Audit agrees: green"
    assert 'class="mappill checks needsrev"' in chip_bad and "checks passed · Audit: needs revision<" in chip_bad, "the checks passed but the latest Audit says not workable: amber"
    # two side jobs under one parent (run 130930: W-7 milestone email and W-8 test email, both under the PRD) fan out side by side
    fan_ = "j2-20261027-130930"
    if os.path.isdir(os.path.join(repo_runs, fan_)):
        if not os.path.exists(os.path.join(j2.RUNS_DIR, fan_)):
            shutil.copytree(os.path.join(repo_runs, fan_), os.path.join(j2.RUNS_DIR, fan_))
        fo_ = render_node(j2.saved_snapshot(fan_, 0), docs={"docs": {}, "map": M.build(os.path.join(j2.RUNS_DIR, fan_))})
        fs_ = fo_["map"][fo_["map"].index('<svg class="mapsvg'):fo_["map"].index("</svg>", fo_["map"].index('<svg class="mapsvg')) + 6]
        assert crossings(fs_) == [] and "could not be drawn without crossing" not in fo_["map"], crossings(fs_)
        lab_ = {}
        for id_ in ("work:W-7", "work:W-8"):
            g_ = ng(fs_, id_)
            x_, t_ = re.search(r'<text[^>]* x="([\d.]+)"[^>]*>(?:<tspan[^>]*>)?([^<]+)', g_.split("</title>")[1]).groups()
            lab_[id_] = (float(x_), len(t_) * 5.8 + 6, t_)
        (x7, w7, t7), (x8, w8, t8) = lab_["work:W-7"], lab_["work:W-8"]
        assert (t7, t8) == ("Email · PRD done · sent", "Test email · skipped (cap)") and abs(x8 - x7) >= (w7 + w8) / 2, lab_
        assert "Finished. Email · PRD done · sent. Test email · skipped (cap)." in fo_["needs"], fo_["needs"]
    pages_ = {}
    for label_, dir_ in cross_runs:
        mp_ = M.build(dir_)
        sn_ = j2.saved_snapshot(os.path.basename(dir_), 0)
        o_ = render_node(sn_, docs={"docs": {}, "map": mp_})["map"]
        sv_ = o_[o_.index("<svg class=\"mapsvg"):o_.index("</svg>", o_.index("<svg class=\"mapsvg")) + 6]
        assert crossings(sv_) == [], (label_, crossings(sv_))
        assert "could not be drawn without crossing" not in o_, label_
        steps_ = [n_["id"] for n_ in mp_["nodes"] if n_["kind"] == "engine"]
        by_ = {n_["id"]: n_ for n_ in mp_["nodes"]}
        assert steps_, label_
        for sid in steps_:
            words_ = " ".join(re.sub(r"<[^>]+>", " ", ng(sv_, sid).split("</title>")[1]).replace("!", " ").replace("?", " ").split())
            assert any(words_.startswith(t_ + " ") for t_ in TEAM), (label_, sid, words_)
            g_ = ng(sv_, sid); assert g_.count('class="teamdot"') == 5 and 'class="teamhit"' in g_, "a team persona is a ring of five dots in a round hit area"
            px_, py_ = (float(v) for v in re.search(r'class="teamhit" cx="([\d.]+)" cy="([\d.]+)"', g_).groups())
            tgt_ = "work:" + htmllib.unescape(sid).split(">")[1]
            if by_[tgt_].get("lane") == "branch":
                continue
            src_ = "work:" + htmllib.unescape(sid).split(">")[0].split(":")[1]
            dx_, dy_ = (float(v) for v in re.search(r'points="([\d.]+),[\d.]+ [\d.]+,([\d.]+)', ng(sv_, tgt_)).groups())
            sx_ = float(re.search(r'points="([\d.]+),', ng(sv_, src_)).group(1))
            assert abs(py_ - dy_) < 0.6 and sx_ < px_ < dx_, (label_, sid, "the team sits on the stretch of line leading into the deliverable it made")
        # side jobs: never in the trunk; the site's step takes the last document; drawn below the trunk
        stp_ = [(e["from"], e["to"]) for e in mp_["edges"] if e["kind"] == "step"]
        trunk_ = [n_["id"] for n_ in mp_["nodes"] if n_["kind"] == "work" and n_.get("lane") == "trunk"]
        for d_ in [x for x in mp_["nodes"] if x["kind"] == "work" and x.get("tag") == "Send"]:
            assert d_["lane"] == "branch" and not any(f == d_["id"] for f, _ in stp_), (label_, d_["id"], "a send is a side job: nothing is built from it")
            ty_ = float(re.search(r'points="[\d.]+,[\d.]+ [\d.]+,([\d.]+)', ng(sv_, "work:website")).group(1))
            ey_ = float(re.search(r'<rect class="shape" x="[\d.-]+" y="([\d.-]+)"', ng(sv_, d_["id"])).group(1))
            assert ey_ > ty_, (label_, "the side job is drawn below the trunk")
        if "work:website" in trunk_:
            into_web = [f for f, t in stp_ if t.endswith(">website") and t.startswith("step:") and f != "work:charter"]      # the Charter feeds every team
            prev_ = trunk_[trunk_.index("work:website") - 1]
            read_ = any(e["kind"] == "handoff" and e["from"] == prev_ and e["to"] == "work:website" for e in mp_["edges"])
            assert into_web == ([prev_] if read_ else []), (label_, into_web, "the site's step takes the box before it as input only when it reads it")
        for d_ in [x for x in mp_["nodes"] if x["kind"] == "work"]:
            if d_.get("tag") == "Send":
                assert '<rect class="shape"' in ng(sv_, d_["id"]) and "<polygon" not in ng(sv_, d_["id"]), (label_, d_["id"], "a send ends in a letter icon")
                continue
            assert '<polygon class="shape"' in ng(sv_, d_["id"]) and '<circle class="shape"' not in ng(sv_, d_["id"]), (label_, d_["id"], "a deliverable is a diamond")
        sv_v = o_[o_.index('<svg class="mapsvg mapvert"'):o_.index("</svg>", o_.index('<svg class="mapsvg mapvert"'))]
        assert crossings(sv_v) == [], (label_, "narrow drawing", crossings(sv_v))
        pages_[label_] = [(htmllib.unescape(a_["from"]), htmllib.unescape(a_["to"]), a_["page"]) for a_ in render_node(sn_, docs={"docs": {}, "map": mp_})["pages"]]
        # label sides, no exceptions: a deliverable's name and fact line below its diamond, a team's line above the trunk
        for d_ in [x for x in mp_["nodes"] if x["kind"] == "work" and x.get("lane") == "trunk"]:
            g_ = ng(sv_, d_["id"]); cy_ = float(re.search(r'points="[\d.]+,[\d.]+ [\d.]+,([\d.]+)', g_).group(1))
            ys_ = [float(y) for y in re.findall(r'<text[^>]* y="([\d.]+)"', g_.split("</title>")[1]) if True]
            ys_ = [y for y, t_ in zip(ys_, re.findall(r'<text[^>]*>([^<]*)<', g_.split("</title>")[1])) if t_ not in ("!", "?")]
            assert len(ys_) == 2 and all(y > cy_ for y in ys_), (label_, d_["id"], "name and fact line below the diamond")
        for sid in [x["id"] for x in mp_["nodes"] if x["kind"] == "engine" and x.get("lane") == "trunk"]:
            g_ = ng(sv_, sid); cy_ = float(re.search(r'class="teamhit" cx="[\d.]+" cy="([\d.]+)"', g_).group(1))
            ys_ = [float(y) for y in re.findall(r'<text class="sub"[^>]* y="([\d.]+)"', g_)]
            assert len(ys_) == 1 and ys_[0] < cy_, (label_, sid, "the team's line above the trunk")
        # the legend: in a collapsed row, and only what this map draws
        lg_ = o_.split('<details class="mapdd mapleg"><summary>Legend</summary>')[1].split("</details>")[0]
        keys_ = re.findall(r'data-legend="([^"]+)"', lg_)
        assert "asks" not in keys_ and "uses" not in keys_ and ("send" in keys_) == any(x.get("tag") == "Send" for x in mp_["nodes"]), (label_, keys_)
        assert ("given_to" in keys_) == ('class="mapedge given' in sv_) and ("left_open" in keys_) == any(x["status"] == "left_open" for x in mp_["nodes"]), (label_, keys_)
        assert not re.search(r'<details class="mapdd mapleg"[^>]* open', o_), "the legend starts closed"
    # three given-to loops that all interleave cannot sit on two sides: the one that cannot be placed is drawn round the outside and listed, never silently crossed
    ds = [f"d{i}" for i in range(1, 7)]
    tn, te = [], []
    for i, d_ in enumerate(ds):
        tn.append({"id": f"work:{d_}", "kind": "work", "row": "work", "label": f"Deliverable {d_}", "tag": "PRD", "color": "--other", "status": "filed", "version": 1, "detail": "", "file": None, "workflows": []})
        if i + 1 < len(ds):
            sid = f"step:{d_}>{ds[i + 1]}"
            tn.append({"id": sid, "kind": "engine", "row": "work", "label": "Write", "status": "filed", "version": None, "detail": "2 workflows", "file": None, "workflows": ["W-1", "W-2"], "tag": None, "color": None})
            te += [{"from": f"work:{d_}", "to": sid, "kind": "step", "label": None, "n": 2, "basis": "recorded"}, {"from": sid, "to": f"work:{ds[i + 1]}", "kind": "step", "label": None, "n": 2, "basis": "recorded"}]
    te += [{"from": f"work:{a_}", "to": f"work:{b_}", "kind": "handoff", "label": "x", "n": 1, "basis": "recorded"} for a_, b_ in (("d1", "d4"), ("d2", "d5"), ("d3", "d6"))]      # [0,5] [2,7] [4,9]: each pair interleaves
    tm = {"schema": "map/1", "status": {"state": "Done", "state_status": "idle", "checks": {"passed": 0, "run": 0}}, "nodes": tn, "edges": te, "legend": saved["legend"], "not_recorded": []}
    ot = render_node(snap, docs={"docs": {}, "map": tm})
    tsv = ot["map"][ot["map"].index('<svg class="mapsvg'):ot["map"].index("</svg>")]
    assert len(crossings(tsv)) >= 1 and ot["map"].count("could not be drawn without crossing") == 1 and "Not recorded (1)" in ot["map"], "the unavoidable crossing is listed, once"
    two = {**tm, "edges": [e for e in te if e["kind"] != "handoff" or e["from"] != "work:d3"]}      # two interleaving loops fit on two sides
    o2 = render_node(snap, docs={"docs": {}, "map": two})["map"]
    assert crossings(o2[o2.index('<svg class="mapsvg'):o2.index("</svg>")]) == [] and "could not be drawn" not in o2, "two loops go on opposite sides: no crossing"
    # the real page: a saved run that ended on a question, and a plain saved run
    if chrome and BROWSER:
        import subprocess
        ended = os.path.join(j2.RUNS_DIR, "j2-20200105-000000")
        shutil.copytree(run.dir, ended)
        os.remove(os.path.join(ended, "map.json"))
        with open(os.path.join(ended, "chat.jsonl"), "a", encoding="utf-8") as fh:
            fh.write(json.dumps({"id": 9999, "who": "identity", "text": "@Requester Add credits?", "to_user": True, "pause": "limit", "next": "user"}) + "\n")
        def home_probe(run_name, js, size="1280,900"):
            tmp = tempfile.mkdtemp(prefix="j2-home-")
            open(os.path.join(tmp, "j2_ui.html"), "w", encoding="utf-8").write(ui.replace("</body>", "<script>setTimeout(() => { const r = {}; try { " + js + " } catch (e) { r.error = String(e); } document.body.setAttribute('data-r', JSON.stringify(r)); }, 4500)</script></body>"))
            keep_here, j2_server.HERE = j2_server.HERE, tmp
            keep, j2_server.RUN = j2_server.RUN, run
            srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), j2_server.Handler)
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            try:
                dom = chrome_dom([chrome, "--headless=new", "--disable-gpu", f"--window-size={size}", "--virtual-time-budget=12000", "--dump-dom",
                                      f"http://127.0.0.1:{srv.server_address[1]}/?view=detail&run={run_name}&mapview=line"], capture_output=True, text=True, timeout=180).stdout
            finally:
                j2_server.HERE, j2_server.RUN = keep_here, keep
                srv.shutdown()
            got = json.loads(htmllib.unescape(re.search(r'data-r="([^"]*)"', dom).group(1)))
            assert "error" not in got, got
            return got
        base_js = ("const q = s => [...document.querySelectorAll(s)], vis = () => q('#dview .pane').filter(p => !p.hidden).map(p => p.id.slice(5)).join(','), rect = s => document.querySelector(s).getBoundingClientRect(),"
                   "click = id => { document.querySelector('.mapnode[data-id=\"' + id + '\"]').dispatchEvent(new MouseEvent('click', {bubbles: true})); document.querySelector('[data-mapgo]').click(); };")
        try:
            a = home_probe("j2-20200105-000000", base_js +
                "r.home = vis(); r.needs = document.getElementById('needs').innerText.replace(/\\s+/g, ' '); r.count = document.getElementById('needcount').textContent;"
                "r.layout = [rect('#needs').bottom <= rect('#pane-map').top + 1, rect('#dview .chat').left >= rect('#dview .work').right - 1, rect('#dview .chat').height > 200];"
                "r.rail = q('.rail.tabs > button').map(b => b.textContent.trim()); r.sub = document.getElementById('railsub').hidden; r.back0 = document.getElementById('backmap').hidden;"
                "document.querySelector('[data-needmsg]').click(); const f = document.activeElement; r.focus = f.classList.contains('msg') + '|' + f.dataset.id;"
                "r.tabs = {}; q('#railsub [data-tab]').forEach(b => { b.click(); r.tabs[b.dataset.tab] = vis() + '|' + document.getElementById('backmap').hidden + '|' + document.getElementById('railsub').hidden; });"
                "document.getElementById('backmap').click(); r.back = vis() + '|' + document.getElementById('backmap').hidden;"
                "let asked = []; const fq = focusQuestion; focusQuestion = id => { asked.push(String(id)); return fq(id); }; r.nodes = {};"
                "for (const id of ['person:requester', 'fn:identity', 'fn:adaptation', 'fn:audit', 'step:prd>hld', 'work:prd', 'work:website', 'work:charter']) { showTab('map'); asked = []; click(id);"
                " r.nodes[id] = {pane: vis(), asked: asked.join(','), wf: WFSTEP.join(','), filt: MAPFILT, art: ARTOPEN}; }"
                "showTab('workflows'); r.wfnote = !!document.querySelector('[data-stepfilter]'); document.querySelector('[data-wfclear]').click(); r.wfcleared = !document.querySelector('[data-stepfilter]') && WFSTEP.length;")
            b = home_probe(name, base_js +
                "let chats = 0; const fc = focusChat; focusChat = () => { chats++; return fc(); }; r.none = document.getElementById('needs').innerText.replace(/\\s+/g, ' ') + '|' + document.getElementById('needcount').textContent;"
                "click('person:requester'); r.person = document.activeElement.id; showTab('map'); click('fn:identity'); r.chats = chats; r.pane = vis();"
                "showTab('map'); window.j2mark = 1; const where = location.href, rs = () => [document.body.classList.contains('simple'), rect('#sview').width > 0, rect('.rail.tabs').width > 0, rect('#dview .dcol').width, document.getElementById('gochat').getAttribute('aria-current'), getComputedStyle(document.querySelector('.rail [data-tab=map]')).fontWeight].join('|'),"
                " msgs = () => document.getElementById('msgs').children.length + '/' + document.getElementById('smsgs').children.length + '|' + document.getElementById('needs').textContent.replace(/\\s+/g, ' ') + '|' + document.getElementById('runsel').value;"
                "r.m0 = msgs(); r.home0 = rs(); document.getElementById('gochat').click(); r.railchat = chats; r.simple = rs(); r.railleft = rect('.rail.tabs').right <= rect('#sview').left + 1; r.m1 = msgs();"
                "document.querySelector('.rail [data-tab=map]').click(); r.back = rs() + '|' + vis(); r.m2 = msgs(); document.getElementById('gochat').click(); r.m3 = msgs(); document.querySelector('.rail [data-tab=map]').click(); r.m4 = msgs(); r.same = (window.j2mark === 1) + '|' + (location.href === where);"
                "document.getElementById('internals').click(); r.toggled = document.getElementById('railsub').hidden;"
                "showTab('runsettings'); document.getElementById('to-simple').click(); r.tosimple = rs(); document.querySelector('.rail [data-tab=map]').click();")
        finally:
            shutil.rmtree(ended, ignore_errors=True)
        assert a["home"] == "map" and a["count"] == "1" and "Needs you" in a["needs"] and "Add credits?" in a["needs"] and "Answer in Chat" in a["needs"], a
        assert a["layout"] == [True, True, True] and a["rail"] == ["Map", "Chat", "Internals"] and a["sub"] is True and a["back0"] is True, a      # strip above the map, chat docked on the right, no Back link on the home
        assert a["focus"] == "true|9999", a      # the strip's button focuses the question in the Chat
        assert a["tabs"] == {t: f"{t}|false|false" for t in old_tabs}, a      # every old tab opens from Internals, with a Back to map link
        assert a["back"] == "map|true", a
        sw = by["step:prd>hld"]["workflows"]
        n = a["nodes"]
        assert n["person:requester"]["asked"] == "9999" and n["fn:identity"]["asked"] == "9999", n      # waiting on you: straight to the question in the Chat
        assert n["fn:adaptation"] == {"pane": "map", "asked": "", "wf": "", "filt": "fn:adaptation", "art": ""}, n      # the other functions: the tables under the Map, filtered to them
        assert n["fn:audit"]["pane"] == "audit" and n["step:prd>hld"]["pane"] == "workflows" and n["step:prd>hld"]["wf"] == ",".join(sw), n
        assert n["work:prd"]["pane"] == "artifacts" and n["work:prd"]["art"] == "prd.md" and n["work:website"]["pane"] == "preview" and n["work:charter"]["pane"] == "charter", n
        assert a["wfnote"] is True and a["wfcleared"] == 0, a      # "Show all" clears the step filter
        assert b["none"].startswith("Needs you 0 Nothing needs you right now.") and b["none"].endswith("See report|0") and b["person"] == "needs" and b["chats"] == 1 and b["railchat"] == 1 and b["toggled"] is False, b
        # the rail's Chat opens the Simple view with the rail kept and Chat lit; Map goes back to the unchanged home; same run, same messages, no reload
        assert b["home0"] == "false|false|true|" + b["home0"].split("|")[3] + "|false|600" and float(b["home0"].split("|")[3]) > 0, b
        assert b["simple"] == "true|true|true|0|page|500" and b["railleft"] is True, b
        assert b["back"] == b["home0"] + "|map" and b["tosimple"] == b["simple"], b      # Back to simple view does what Chat does
        assert b["m0"] == b["m1"] == b["m2"] == b["m3"] == b["m4"] and b["same"] == "true|true" and int(b["m0"].split("/")[0]) > 0, b
        # the header is the name only; Map has one visible heading; the outcome line opens the Report; the banner text is there; nothing overlaps
        vis_js = "const seen = e => e.getClientRects().length > 0 && getComputedStyle(e).visibility !== 'hidden';"
        c = home_probe(name, base_js + vis_js +
            "r.head = q('#dview h1').filter(seen).length + '|' + document.getElementById('dview').innerText.includes('J2 Department'); r.heads = q('#dview h1, #dview h2, #dview h3').filter(seen).map(h => h.textContent.trim()).filter(t => t === 'Map').length;"
            "r.outcome = !!document.querySelector('#needs [data-outcome]'); document.querySelector('#needs [data-gotab=\"report\"]').click(); r.afterLink = vis(); r.banner = document.getElementById('banner').textContent; r.bannerIn = !!document.getElementById('banner').closest('#pane-report'); r.bannerShown = seen(document.getElementById('banner'));"
            "showTab('runsettings'); r.settings = ['outsel', 'runsel', 'reset', 'to-simple'].map(i => seen(document.getElementById(i))).join(','); r.runsel = document.getElementById('runsel').value; r.opts = document.getElementById('runsel').options.length;"
            "showTab('map');")
        assert c["head"] == "0|false" and c["heads"] == 1, c      # no department heading; exactly one visible "Map" heading
        assert c["outcome"] is True and c["afterLink"] == "report" and c["bannerIn"] is True and c["bannerShown"] is True and c["banner"].startswith("Viewing saved run " + name + " (read only)."), c
        assert c["settings"] == "true,true,true,true" and c["opts"] >= 2, c      # the controls are all there in Run settings
        newest = os.path.join(j2.RUNS_DIR, "j2-29990101-000001")      # sorts first, so the selector (newest 60) always lists it
        shutil.copytree(run.dir, newest)
        assert home_probe("j2-29990101-000001", base_js + "showTab('runsettings'); r.val = document.getElementById('runsel').value;")["val"] == "j2-29990101-000001", "the selector shows the run that is open"
        fx = ("const sv = [...document.querySelectorAll('.mapsvg')].find(x => x.getClientRects().length); const lines = [...sv.querySelectorAll('path.mapedge')].filter(p => parseFloat(getComputedStyle(p).opacity) > 0.01).map(p => { const L = p.getTotalLength(), o = []; for (let i = 0; i <= L; i += 2) { const t = p.getPointAtLength(i); o.push([t.x, t.y]); } return {id: p.dataset.from + ' > ' + p.dataset.to, o}; });"
              "const texts = [...sv.querySelectorAll('text')].filter(t => parseFloat(getComputedStyle(t).opacity) > 0.01 && t.textContent.trim()).map(t => ({t: t.textContent.trim(), b: t.getBBox()}));"
              "r.n = texts.length; r.hits = []; for (const x of texts) { for (const l of lines) if (l.o.some(([px, py]) => px > x.b.x - 1 && px < x.b.x + x.b.width + 1 && py > x.b.y && py < x.b.y + x.b.height)) r.hits.push(x.t + ' on ' + l.id); }"
              "for (let i = 0; i < texts.length; i++) for (let j = i + 1; j < texts.length; j++) { const a = texts[i].b, b = texts[j].b; if (a.x < b.x + b.width - 1 && b.x < a.x + a.width - 1 && a.y < b.y + b.height - 1 && b.y < a.y + a.height - 1) r.hits.push(texts[i].t + ' over ' + texts[j].t); }")
        for rn in (name, send_name, *real_names):
            d_ = home_probe(rn, base_js + fx)
            assert d_["n"] > 20 and d_["hits"] == [], (rn, d_["hits"])      # no label sits on a line or on another label
        # the wide drawing fits the map pane at 1440x900 with no scroll inside it; under 900px the narrow drawing is shown instead
        fit_js = ("const p = document.getElementById('pane-map'), w = document.querySelector('.mapsvg:not(.mapvert)'), v = document.querySelector('.mapsvg.mapvert');"
                  "r.fit = [p.scrollHeight, p.clientHeight, w.getClientRects().length > 0, v.getClientRects().length > 0];")
        for rn in (name, send_name, *real_names):
            f1 = home_probe(rn, base_js + fit_js, "1440,900")["fit"]
            assert f1[2] is True and f1[3] is False and f1[0] <= f1[1] + 1, (rn, f1)      # wide drawing shown, and it fits
        f2 = home_probe(real_names[0], base_js + fit_js, "860,900")["fit"]
        assert f2[2] is False and f2[3] is True, f2      # under 900px: the vertical drawing
        # phone width: nothing wider than 360px on the Simple view or the home, and the rail is along the bottom. Headless Chrome's
        # window is at least 500px wide, so the page is measured in a 360px iframe served from the same origin.
        def phone_probe(run_name, view):
            js = ("setTimeout(() => { const d = document.querySelector('iframe').contentDocument, W = d.documentElement.clientWidth, r = {W, sw: d.documentElement.scrollWidth, out: []};"
                  "d.querySelectorAll('body *').forEach(e => { const b = e.getBoundingClientRect(); if (b.width && b.height && (b.right > W + 0.5 || b.left < -0.5)) r.out.push((e.id || e.className || e.tagName) + ' ' + Math.round(b.left) + '-' + Math.round(b.right)); });"
                  "const rail = d.querySelector('.rail.tabs').getBoundingClientRect(), main = d.querySelector(" + json.dumps("#sview" if view == "simple" else "#dview .dcol") + ").getBoundingClientRect();"
                  "r.rail = [rail.top >= main.bottom - 1, Math.abs(rail.bottom - d.documentElement.clientHeight) <= 1, rail.height > 0];"
                  "document.body.setAttribute('data-r', JSON.stringify(r)); }, 6000)")
            page = (f'<!doctype html><body style="margin:0"><iframe src="/?view={view}&run={run_name}" style="width:360px;height:800px;border:0;display:block"></iframe>'
                    f'<script>{js}</script></body>')

            class Wrap(j2_server.Handler):
                def do_GET(self):
                    return self._send(200, page, "text/html") if self.path == "/wrap" else super().do_GET()
            tmp = tempfile.mkdtemp(prefix="j2-phone-")
            open(os.path.join(tmp, "j2_ui.html"), "w", encoding="utf-8").write(ui)
            keep_here, j2_server.HERE = j2_server.HERE, tmp
            keep, j2_server.RUN = j2_server.RUN, run
            srv = j2_server.http.server.ThreadingHTTPServer(("127.0.0.1", 0), Wrap)
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            try:
                dom = chrome_dom([chrome, "--headless=new", "--disable-gpu", "--window-size=600,800", "--virtual-time-budget=10000", "--dump-dom",
                                  f"http://127.0.0.1:{srv.server_address[1]}/wrap"], capture_output=True, text=True, timeout=180).stdout
            finally:
                j2_server.HERE, j2_server.RUN = keep_here, keep
                srv.shutdown()
            return json.loads(htmllib.unescape(re.search(r'data-r="([^"]*)"', dom).group(1)))
        for v_ in ("simple", "detail"):
            p_ = phone_probe(name, v_)
            assert p_["W"] == 360 and p_["sw"] == 360 and p_["out"] == [], (v_, p_)      # no sideways scroll, nothing past either edge
            assert p_["rail"] == [True, True, True], (v_, p_)      # the rail is along the bottom, below the screen's content
        # Run settings: New run posts to /api/reset; choosing a saved run opens it
        e_ = home_probe(name, base_js + "let posted = ''; window.fetch = (u, o) => { if (String(u).includes('/api/reset')) posted = (o && o.method) || 'GET'; return new Promise(() => {}); }; showTab('runsettings'); document.getElementById('reset').click(); r.posted = posted;")
        assert e_["posted"] == "POST", e_
        f_ = home_probe(name, base_js + "if (sessionStorage.step !== '2') { sessionStorage.step = '2'; showTab('runsettings'); const sel = document.getElementById('runsel'); const o = [...sel.options].find(x => x.value && x.value !== '" + name + "'); sel.value = o ? o.value : ''; r.picked = sel.value; sel.dispatchEvent(new Event('change')); } else { r.after = location.search; r.val = document.getElementById('runsel').value; sessionStorage.clear(); }")
        assert f_.get("after", "").startswith("?run=j2-") and f_["val"] and f_["val"] != name, f_      # the selector switches runs
        shutil.rmtree(newest, ignore_errors=True)
    for n_ in real_names:
        shutil.rmtree(os.path.join(j2.RUNS_DIR, n_), ignore_errors=True)
    print("map: map.json written beside atoms.json and equal to the folder's build, golden nodes/edges/statuses, counts from the atoms, "
          "old run rebuilt with its gaps named, left open vs running, fault, not made, waiting on you, /api/map live/saved/old/missing, "
          "SVG shapes and marks per status, links to Artifacts, tooltips, plain words, cards for small screens ok")


def atom_schema():
    """atom/1: the adapter maps a stored run (makes and dos), edges are what each builder's prompt actually held
    (manifest, LLD slice, earlier document), used_by is the exact inverse, warnings never change a run, atoms.json is
    written, and a builder given the whole page has no attributable edge."""
    import j2_schema as S
    keys = ["schema", "id", "workflow", "name", "owner", "status", "attempt", "goal", "how", "constraints", "context",
            "defined_output", "verification", "output", "depends_on", "used_by", "notes", "details"]
    # the adapter, on a stored website run: standard fields only, website words only in details
    os.environ["J2_SLIM"] = "0"          # the old whole-page path: no edge can be attributed
    try:
        run = j2.J2Run()
        say(run, BREW_LAB)
    finally:
        os.environ.pop("J2_SLIM", None)
    wfs = json.loads(json.dumps(run.workflows))
    views, warns = S.run_views(wfs)
    assert views and all(list(v) == keys and v["schema"] == "atom/1" for v in views), views[0].keys()
    a, v = next((a, v) for w in wfs for a in w["atoms"] for v in views if v["id"] == a["id"] and a.get("agent_output", {}).get("anchor"))
    assert v["notes"] == a["human_output"]["text"] and v["defined_output"] == a["expected_output"] and v["owner"] == a["how"]["agent"]
    assert v["status"] == {"passed": "closed"}.get(a["status"], v["status"]) and v["goal"] == a["goal"] and v["workflow"] == a["id"].split(".")[0]
    sec = next(o for o in v["output"] if o["form"] == "page" and "#" in o["ref"])
    assert sec == {"verb": "make", "name": a["agent_output"]["section_names"][0], "form": "page",
                   "ref": f"{a['agent_output']['files'][-1]}#{a['agent_output']['anchor']}"}, sec
    assert any(o["form"] == "page" and o["ref"].endswith(".html") and o["verb"] == "make" for o in v["output"])
    assert v["details"]["selectors"] == a["agent_output"]["selectors"] and "selectors" not in v
    assert len(v["verification"]) == len(a["results"]) and all(set(x) == {"check", "applies_to", "method", "met", "evidence"}
                                                                  for x in v["verification"])
    assert {x["check"] for x in v["verification"]} <= {"exists", "valid", "meets the defined output", "receipt"}
    assert all(x["applies_to"] in {o["name"] for o in v["output"]} for x in v["verification"] if v["output"]), v["verification"]
    # a behaviour step is a do item of form "behaviour" (scope self), checked by one "receipt" check: the receipt id
    eff = json.loads(json.dumps(a))
    eff["results"] = [{"atom": a["id"], "check": "Total updates", "type": "behavior-steps", "target": [], "pass": True,
                       "evidence": "step 3 saw 12", "attempt": 1}]
    ev = S.atom_view(eff)
    assert [o for o in ev["output"] if o["verb"] == "do"] == [{"verb": "do", "name": "Total updates", "form": "behaviour",
                                                              "ref": f"{a['id']}:1"}], ev["output"]
    assert ev["verification"][0]["check"] == "receipt" and ev["verification"][0]["evidence"] == f"{a['id']}:1" \
        and ev["verification"][0]["met"] is True and ev["verification"][0]["applies_to"] == "Total updates"
    assert S.atom_kind(ev) == "mixed" and S.atom_kind(v) == "make" and not any("kind" in x or "type" in x for x in ev["output"])
    assert any(x.startswith("info: ") and "mixed atom" in x for x in S.validate_views(S.link_views([ev]))), "mixed is flagged"
    assert not S.validate_atom(ev), "a mixed atom is valid"
    assert ev["details"]["scope"] == "self" and "scope" not in v["details"]
    # a document atom: its output is a document section, ref = file and heading
    prd = docs_run("Write a PRD for Campus Notes." + CAMPUS, 200)
    d = next(x for w in prd.workflows for x in w["atoms"])
    dv = S.atom_view(d)
    assert [o for o in dv["output"] if o["form"] == "doc" and "#" in o["ref"]] == [
        {"verb": "make", "name": d["agent_output"]["section_names"][0], "form": "doc",
         "ref": f"prd.md#{d['expected_output'][0]['target']}"}], dv["output"]
    # whole page given and nothing attributable: no edge is guessed; dependencies only when the workflow declares them
    assert all(g["whole"] and not g["edges"] for g in run.given.values()), run.given
    assert not any(e for w in wfs if not w.get("dependencies") for x in S.run_views([w])[0] for e in x["depends_on"])
    # edges from a manifest run: every edge is an atom named in that builder's prompt, and no manifest atom is left out
    try:
        slim_run = j2.J2Run()
        say(slim_run, BREW_LAB)
    finally:
        os.environ.pop("J2_SLIM", None)
    rows = calls_file(slim_run)
    n_edges, n_old = 0, sum(len(x["depends_on"]) for x in views)
    for wid, g in slim_run.given.items():
        p = [r["prompt"] for r in rows if r["fn"] == wid and "OTHER SECTIONS" in r["prompt"]][-1]
        listed = set(re.findall(r'"atom": "([^"]+)"', p.split("OTHER SECTIONS")[1].split("\n")[0]))
        assert {e["atom"] for e in g["edges"]} == listed and not g["whole"], (wid, listed, g)
        n_edges += len(g["edges"])
    assert n_edges, "later builders see earlier atoms"
    # edges from an LLD-slice run and a chain run (earlier document), and used_by is the exact inverse of depends_on
    full = docs_run(DESIGN_BRIEF, 900)
    assert full.outcome.startswith("Done"), full.outcome
    rows, fv, fw = calls_file(full), S.run_views(full._all_wfs(), full.given)[0], {w["id"]: w for w in full._all_wfs()}
    lld_atoms = {x for x in (full.docs["lld"]["atoms"])}
    hld_atoms = {x for x in (full.docs["hld"]["atoms"])}
    prd_req = set(full._doc_atoms("prd", "Requirements"))
    build = [w for w in full.workflows if w.get("lld_sections")]
    assert build
    for w in build:
        edges = full.given[w["id"]]["edges"]
        spec_e = [e for e in edges if "LLD slice" in e["why"] or "HLD" in e["why"]]      # the manifest adds site atoms too
        assert {e["atom"] for e in spec_e} & lld_atoms and all(e["atom"] in lld_atoms | hld_atoms for e in spec_e), (w["id"], edges)
        assert any("LLD slice" in e["why"] and w["lld_sections"][0] in e["why"] for e in edges)
        p = [r["prompt"] for r in rows if r["fn"] == w["id"]][-1]
        assert "YOUR SPEC" in p and "HLD SUMMARY" in p and "THE WHOLE SITE FILE" not in p
    hld_w = [x for x in fw.values() if x["id"] in full.docs["hld"]["workflows"]]
    assert all(set(e["atom"] for e in full.given[x["id"]]["edges"] if "PRD" in e["why"]) == prd_req for x in hld_w), "HLD builders were given the PRD"
    assert any("PRD" in e["why"] for x in hld_w for e in full.given[x["id"]]["edges"])
    pairs = {(v["id"], e["atom"]) for v in fv for e in v["depends_on"]}
    assert pairs and pairs == {(u["atom"], v["id"]) for v in fv for u in v["used_by"]}, "used_by is the inverse of depends_on"
    assert all(e["why"] for v in fv for e in v["depends_on"] + v["used_by"])
    # status: exactly four values, mapped from today's names; kinds: from the open list, none called "section"
    st = lambda s, ws=None, ended=False, att=1: S.atom_status({"status": s, "attempts": att}, {"status": ws}, ended)
    assert S.STATUSES == ("open", "closed", "verify_failed", "abandoned")
    assert [st("passed", "done"), st("escalated", "done"), st("failed", "failed"), st("failed", "done"), st("failed", "running", att=3),
            st("failed", "running", att=1), st(None, "registered"), st(None, "held"), st(None, "running"), st("could not run", "done"),
            st(None, "held", ended=True), st(None, "registered", ended=True), st(None, "unfunded"), st(None, "done", ended=True),
            st("could not run", "held", ended=True)] == \
        ["closed", "verify_failed", "verify_failed", "verify_failed", "verify_failed", "open", "open", "open", "open", "open",
         "abandoned", "abandoned", "abandoned", "open", "open"]
    seen = {o["form"] for v in fv for o in v["output"]}
    assert all(o["verb"] in ("make", "do") and "type" not in o and "kind" not in o for v in fv + views for o in v["output"])
    assert all(v["status"] in S.STATUSES for v in fv + views) and "section" not in seen and "document-section" not in seen, seen
    assert {"page", "doc"} <= seen and all(o["form"] in S.MAKE_FORMS for v in fv for o in v["output"] if o["verb"] == "make"), seen
    odd = S.atom_view({"id": "X.a1"}); odd["output"].append({"verb": "make", "name": "n", "form": "diagram", "ref": "d"})
    assert S.validate_atom(odd) == [], "an unknown kind is recorded, never rejected"
    f_done = json.load(open(os.path.join(full.dir, "atoms.json")))
    assert all(v["status"] in S.STATUSES for v in f_done["atoms"])
    print("atom/1 statuses: four values; forms seen in the full chain: " + ", ".join(sorted(seen)))
    # a run with behaviour steps (a calculator): do items of form "behaviour", each with a receipt check that passes
    calc = j2.J2Run()
    say(calc, BREW_LAB.replace("and where to find us", "where to find us, and a price calculator for an order"))
    finish(calc)
    cv = S.run_views(calc._all_wfs(), calc.given)[0]
    dos = [(v, o) for v in cv for o in v["output"] if o["verb"] == "do"]
    assert dos and all(o["form"] == "behaviour" and o["ref"].startswith(v["id"] + ":") for v, o in dos), dos
    chk = [x for v, o in dos for x in v["verification"] if x["applies_to"] == o["name"] and x["check"] == "receipt"]
    assert len(chk) == len(dos) and all(x["evidence"] == o["ref"] for (v, o), x in zip(dos, chk)), chk
    if not BROWSER or j2.j2_check.find_chrome():
        assert all(x["met"] is True for x in chk), chk
    assert {S.atom_kind(v) for v, _ in dos} == {"mixed"} and all(v["status"] in S.STATUSES for v in cv)
    assert any(x.startswith("info: ") and "mixed atom" in x for x in S.validate_views(cv)), "mixed is flagged in the run's warnings"
    atoms_json = json.load(open(os.path.join(calc.dir, "atoms.json")))
    assert any(x.startswith("info: ") for x in atoms_json["warnings"]), atoms_json["warnings"]
    kinds = lambda vs: {k: sum(S.atom_kind(v) == k for v in vs) for k in ("make", "do", "mixed")}
    import shutil
    if shutil.which("node"):                       # the card: verb, noun, form and ref; "mixed" shown only on a mixed atom
        card = render_node(j2.saved_snapshot(os.path.basename(calc.dir), 0))["cards"]
        mixed_ids = {v["workflow"] for v, _ in dos}
        for w in calc._all_wfs():
            c = card[w["id"]]
            assert ('data-note="mixed"' in c) == (w["id"] in mixed_ids), w["id"]
        mc = card[next(iter(mixed_ids))]
        assert "make · artifact · page" in mc and "do · deed · behaviour" in mc and "effect" not in mc.lower(), mc[:200]
    print(f"atom/1 verbs: full chain {kinds(fv)}, behaviour run {kinds(cv)} (make-only, do-only, mixed); behaviour run has "
          f"{len(dos)} do item(s) of form behaviour, each with a receipt check")
    # warnings only: dangling edge, cycle, orphan, missing type
    mk = lambda i, out=None, dep=(): {**S.atom_view({"id": i, "status": "passed"}), "output": out or [], "depends_on": [{"atom": d, "why": "x"} for d in dep]}
    final = [{"verb": "make", "name": "p", "form": "page", "ref": "site/index.html#p"}]
    side = [{"verb": "make", "name": "f", "form": "file", "ref": "f.txt"}]
    w = S.validate_views(S.link_views([mk("A.a1", final, ["ZZ.a1"])]))
    assert any("ZZ.a1" in x and "not an atom" in x for x in w), w
    w = S.validate_views(S.link_views([mk("A.a1", final, ["B.a1"]), mk("B.a1", final, ["A.a1"])]))
    assert any(x.startswith("cycle: ") for x in w), w
    w = S.validate_views(S.link_views([mk("A.a1", side), mk("B.a1", final)]))
    assert w == ["A.a1: its output is never used and is not a final output"], w     # no info line: not mixed
    w = S.validate_views(S.link_views([mk("A.a1", [{"name": "x", "form": "file", "ref": "x"}] + final)]))
    assert any("no verb" in x for x in w), w
    # atoms.json: written at the end of a run with the atom/1 views; existing run files are unchanged in name
    f = json.load(open(os.path.join(full.dir, "atoms.json")))
    assert f["schema"] == "atom/1" and [v["id"] for v in f["atoms"]] == [v["id"] for v in fv] and "warnings" in f
    assert f["atoms"][-1]["depends_on"] and set(os.listdir(full.dir)) >= {"summary.json", "workflows.json", "atoms.json"}
    assert json.load(open(os.path.join(full.dir, "workflows.json")))[0]["atoms"][0].keys() >= {"expected_output", "human_output", "agent_output"}
    print(f"atom/1 edges in a plain website run: J2_SLIM=0 {n_old}, default (slim) {n_edges}")
    print(f"atom/1: adapter (make and do items, details, notes), edges from a manifest run ({n_edges}), an LLD-slice run "
          f"and a chain run ({len(pairs)} in the full chain), used_by the exact inverse, warnings for dangling, cycle, orphan "
          "and missing verb, atoms.json written ok")


def docs_intake_and_compare():
    """A documents-plus-website run asks for the website's details once, before the first document stage, and the answers
    reach every stage's charter; the website stage doesn't ask again. And lld_compare follows the revised LLD section: a
    revision that removes a field is compared as removed."""
    ANSWER = "Menu: Filter coffee 90, Latte 180. Hours: Monday to Sunday 8 am to 9 pm. Address: 12, 100 Feet Road. Form: name, phone, items."
    seen = []

    def asker(run, orig):
        def call(fn, fmt, task, mock, **k):
            if fn == "identity" and "EARLIER QUESTIONS AND ANSWERS" in task:
                seen.append((run.stage_kind, bool(run.intake), "Ask for them once, now" in task, "asked for the website's details before them" in task))
                if "Ask for them once, now" in task and not run.intake:
                    return {"say": "To build the page I need a few real details only you know. Please send them as you'd like them to appear.",
                            "enough": False, "next": "Requester",
                            "questions": [{"question": "Please send these: the menu and prices, the opening hours, the address, "
                                                       "and what the pre-order form should ask for.", "options": []}]}
            return orig(fn, fmt, task, mock, **k)
        return call
    run = docs_run(DESIGN_BRIEF, 900, asker)
    assert run.paused and run.paused["kind"] == "info" and run.stage_kind == "prd" and not run.workflows and not run.docs, run.outcome
    say(run, ANSWER)
    assert run.outcome.startswith("Done") and run.stages == ["prd", "hld", "lld", "website"], run.outcome
    asks = [m for m in run.chat if m.get("act") == "ASK_REQUESTER" and m["who"] == "identity"]
    assert len(asks) == 1 and "details only you know" in asks[0]["text"], [m["text"][:80] for m in asks]
    first_doc = next(i for i, m in enumerate(run.chat) if m.get("workflow") or "registered" in m.get("text", "") and m["who"] == "priority")
    assert run.chat.index(asks[0]) < first_doc
    assert [x for x in seen if x[2]] and all(x[3] for x in seen if x[0] == "website"), seen        # asked first; the website stage knows
    assert [x for x in seen if x[0] == "website"] and not [x for x in seen if x[0] == "website" and x[2]], seen
    for g in run.stage_archive[1:3]:                                  # the HLD's and the LLD's charters carry the answer
        assert g["goal"]["requester_answers"] == [ANSWER], g["stage"]
    hld = next(w for w in run._all_wfs() if w["id"] in run.docs["hld"]["workflows"])
    assert any(ANSWER in c["prompt"] for c in calls_file(run) if c["fn"] == hld["id"]), "the HLD builder sees the details"
    # a run without a website, and a plain brief, are not asked this way
    tasks = []

    def spy(run2, orig):
        def call(fn, fmt, task, mock, **k):
            tasks.append(task)
            return orig(fn, fmt, task, mock, **k)
        return call
    docs_run("Write a PRD for Campus Notes." + CAMPUS, 200, spy)
    docs_run(BREW_LAB, 200, spy)
    assert tasks and not any("Ask for them once, now" in t or "asked for the website's details before them" in t for t in tasks)
    # lld_compare follows the revised section: a revision that removes a field
    full = docs_run(DESIGN_BRIEF, 900)
    assert full.outcome.startswith("Done"), full.outcome
    f = full.blueprint["fields"][0]
    page = "<html><body>" + j2.lld_markup(full.blueprint) + "</body></html>"
    page = re.sub(rf"<label>[^<]*<input name=\"{f['name']}\"[^>]*></label>", "", page)
    assert f["name"] not in page and any(m["section"] == f["section"] for m in full._lld_compare(page)["mismatches"])
    w = next(x for x in full.workflows if f["section"] in (x.get("lld_sections") or []))
    full._revise_lld(w, f["section"], {"reason": "the charter has no such field", "use": {}, "fields": {f["id"]: None}})
    cmp = full._lld_compare(page)
    assert f["section"] in cmp["matches"] and not cmp["mismatches"], cmp
    assert all(x["id"] != f["id"] for x in full.blueprint["fields"]) and all(x["id"] != f["id"] for x in full.docs["lld"]["agent_output"]["fields"])
    assert f"**{f['id']}**" not in open(os.path.join(full.dir, "lld.md"), encoding="utf-8").read()
    assert full.lld_revisions[-1]["fields"] == {f["id"]: None}
    print("documents run: the website's details are asked once before the PRD and reach every stage; lld_compare follows "
          "the revised section (a removed field is a match) ok")


if __name__ == "__main__":
    if "--board-doc-demo" in sys.argv:
        # python3 j2_flow_test.py --fast --board-doc-demo: the document-change demo (board_doc_change's own run),
        # saved to runs/ as evidence instead of the temp folder. Mock mode, the browser measurement stubbed.
        os.environ["J2_BOARD"] = "1"
        j2.RUNS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "runs")
        demo = doc_board_run(demo=True)
        e = (demo.board["entries"] or [{}])[0]
        print(f"Board: {e.get('id')} {e.get('asked_by')} -> {e.get('to')} ({e.get('status')}); "
              f"{json.dumps(e.get('artifact_update'))}; impact {len(e.get('impact') or [])} atom(s)")
        print(f"http://127.0.0.1:8773/?run={os.path.basename(demo.dir)}&view=detail")
        sys.exit(0)
    if "--board-demo" in sys.argv:
        # python3 j2_flow_test.py --fast --board-demo: the message board demo (board_functions' own run), saved to
        # runs/ as evidence instead of the temp folder. Mock mode, the browser measurement stubbed as --fast does.
        os.environ["J2_BOARD"] = "1"
        j2.RUNS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "runs")
        demo = board_run(demo=True)
        print(f"Board: {len(demo.board['entries'])} question(s), "
              + ", ".join(f"{e['id']} {e['asked_by']} -> {e['to']} ({e['status']})" for e in demo.board["entries"]))
        print(f"http://127.0.0.1:8773/?run={os.path.basename(demo.dir)}&view=detail")
        sys.exit(0)
    if FAST:
        print("--fast: browser measurement stubbed" + ("" if BROWSER else "; tests that drive the page in Chrome are "
                                                       "skipped below (add --browser to run them)"))
    for mode in ("0", "1"):
        os.environ["J2_FIXES"] = mode
        vague(mode == "1")
        vague_stories(mode == "1")
        full(mode == "1")
        policy_budget(mode == "1")
        below_floor()
        thread_and_loom_core()
        deferred_twice()
        connections()
        build_before_asking()
        print(f"J2_FIXES={mode}: vague brief (both answers), full brief, budget question rejected and re-asked, below-floor proposal with left_out, one decision question after the ruling, core marks from J2.a, deferred twice, connections at three budgets, funded work built before the question ok")
        workflow_agents()
        run_20261015_fixes()
        atom_outputs()
        slim_prompts()
        atom_cards_render()
        replay_size()
        docs_detection()
        docs_chain()
        docs_funding()
        send_deeds()
        if mode == "0":
            live_sends()
        charter_objects()
        milestone_emails()
        artifact_types()
        artifacts_tab()
        map_tab()
        map_every_workflow()
        if mode == "0":
            map_tree()
            map_logic()
        docs_repeat_failure()
        docs_intake_and_compare()
        plain_unchanged("f6938b2", slim="0")      # J2_SLIM=0 is the old path, prompt for prompt
        plain_unchanged()                         # the default (slim) path against the commit made right after the flip
        plain_unchanged(board="0")                 # J2_BOARD=0 explicitly set: still identical, prompt for prompt
        charter_crash_fix()
        board_disabled()
        board_functions_off()
        os.environ["J2_BOARD"] = "1"
        board_story()
        board_functions()
        board_not_mine()
        if mode == "0":      # the board's document changes don't branch on J2_FIXES: four stages, run once
            board_doc_change()
            board_doc_not_owner()
            board_doc_prd()
            board_doc_no_edge()
            board_doc_recheck_fails()
        os.environ.pop("J2_BOARD", None)
        if mode == "0":
            board_doc_off()
        atom_schema()
        docs_blueprint()
        docs_views()
        docs_tab()
        card_content()
        charter_tab()
        fixes_line()
        output_tags()
        lineage_rows()
        fresh_page_run()
        missing_run_endpoints()
        atom_checks()
        behavior_steps()
        charter_facts()
        print(f"J2_FIXES={mode}: behaviour steps: interpreter with a fake driver, validator rules (selectors, charter "
              "numbers, date control, 8 steps, action goals), one re-ask, mock behaviour items ok")
        leave_out_named()
        rubric_and_credits()
        charter_numbers()
        chrome_unavailable()
        unique_short_names()
        same_failure()
        cycle_limit()
        to_be_confirmed()
        bargain_time()
        memory()
        forced_fail()
        test_lessons()
        style_guide()
        rubric_loop()
        lone_rubric_failure()
        pricing()
        unpaid_items()
        intake_needs()
        playbooks()
        summary_before_review()
        prd_run()
        prd_checks()
        prd_groups_metrics_assumptions()
        prd_recovery()
        website_conditions()
        print(f"J2_FIXES={mode}: website necessary conditions: N1 to N8 each reported not met with evidence, nav plan "
              "flagged before building, workable needs every condition, no None or internal ids for the Requester ok")
        print(f"J2_FIXES={mode}: PRD output: mock PRD brief Done workable with prd.md and prd.docx, document checks pass "
              "and fail (duplicate R2 rejected), invented facts-present rejected, no browser check, no website "
              "conditions, P7 user groups, P8 window warning, assumptions labelled and listed, short names, heading targets "
              "normalised, wrong target is CHECK_WRONG, unusable reply twice gets a code fix, sections-only question ok")
        print(f"J2_FIXES={mode}: B workflow agents, atom checks (rules, receipts, third failure to Adaptation), leave-out question names its items, rubric in briefs, tap targets, no \"units\", charter numbers for agents, no Chrome = could not run, unique short names, C workable loop and limits, D memory, F forced audit failure (hidden from agents) and test lessons, summary before review, H style guide, rubric loop, needs, playbooks, I pricing and unpaid items ok")
    ruling_over_budget()
    redirect()
    workable_unit()
    playbook_matching()
    limit_answers()
    nine_workflows()
    budget_cap()
    repeats()
    adaptation_idle()
    identity_decisions()
    two_audiences()
    audiences_everywhere()
    limit_wording()
    open_list()
    covers_by_text()
    readable_decision()
    simple_messages()
    owner_override()
    form_submit()
    print("second pass: two audiences, simple messages, owner override, form submit ok")
    print("limit answers (turn cap, nine workflows, own amount, drop with reason, done with open items), budget cap, repeats, idle Adaptation, decisions applied by code (msgs 24, 50) ok")
    behavior_browser()
    page()
    prd_page()
    print("PRD page: output picked before the first message and shown, preview from prd.md, .md and .docx downloads ok")
    print("E behind-the-scenes view (lanes, detail, timeline, saved runs, redaction, lessons) ok")
    generalist_units()
    detail_stays_open()
    work_atoms()
    work_atom_cards()
    print("work atoms: validator names a missing field and retries once, atoms stored and tagged, cards render ok")
    print("Detailed view: opened screenshot, event log and message detail stay open while the run updates ok")
    simple_view()
    plain_words()
    identity_prompt_is_plain()
    print("G simple view (plain words at the source and as fallback, buttons once, phone toggle, ready card) ok")
    assert fingerprint(REAL_MEMORY) == REAL_BEFORE, "the tests changed the real memory/ files"
    print("playbooks: same business type only, mock memory kept from live runs; real memory/ untouched ok")
    print("j2 flow checks passed")
